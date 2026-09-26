"""Workbook edit operations (openpyxl) shared by sheet_edit.py and sheet_create.py.

Each op is a dict with "op" plus arguments; see references/ops.md. Structural ops (insert/delete rows and columns,
rename/delete sheets) rewrite formula references everywhere they occur: cell formulas on every sheet, defined names,
conditional formats, data validations, tables, merged cells, print areas, chart series, hyperlinks and anchors.
"""

from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path
from typing import Any, Callable

from _a1 import MAX_COL, MAX_ROW, col_index, col_letter, parse_cell, parse_cols, parse_range, parse_rows, quote_sheet, range_name, split_sheet
from _common import SkillError
from _formula import (
    FormulaSyntaxError, add_prefixes, drop_sheet_refs, parse, rename_sheet_refs, shift_area, shift_refs, translate,
)

OPS: dict[str, Callable[["Ctx", dict[str, Any]], str]] = {}
ALIASES = {
    "set_value": "set", "set_values": "set", "write": "set", "formula": "set", "set_formula": "set",
    "insert_row": "insert_rows", "delete_row": "delete_rows", "insert_col": "insert_cols", "insert_column": "insert_cols",
    "insert_columns": "insert_cols", "delete_col": "delete_cols", "delete_column": "delete_cols", "delete_columns": "delete_cols",
    "format": "style", "number_format": "style", "width": "column_width", "col_width": "column_width", "height": "row_height",
    "replace": "find_replace", "conditional": "conditional_format", "data_validation": "validation", "define_name": "name",
    "named_range": "name", "add_table": "table", "add_chart": "chart", "add_comment": "comment", "link": "hyperlink",
    "freeze_panes": "freeze", "filter": "autofilter", "rename": "rename_sheet", "copy": "copy_sheet", "move": "move_sheet",
    "reorder_sheets": "reorder", "hide_sheet": "sheet_state", "fill_down": "fill", "fill_right": "fill", "add_image": "image",
    "page_setup": "print", "unmerge_cells": "unmerge", "merge_cells": "merge", "autofit_columns": "autofit",
}


def op(name: str) -> Callable[[Callable[["Ctx", dict[str, Any]], str]], Callable[["Ctx", dict[str, Any]], str]]:
    def deco(f: Callable[["Ctx", dict[str, Any]], str]) -> Callable[["Ctx", dict[str, Any]], str]:
        OPS[name] = f
        return f

    return deco


class Ctx:
    """State shared by the ops of one run."""

    def __init__(self, wb: Any, default_sheet: str | None = None) -> None:
        self.wb = wb
        self.default_sheet = default_sheet
        self.formulas_changed = False
        self.new_cells: list[tuple[Any, Any]] = []  # (worksheet, cell) of formulas written by ops; cells move with edits
        self.warnings: list[str] = []

    @property
    def new_formulas(self) -> set[tuple[str, int, int]]:
        """Current (sheet, row, col) of formulas written by ops that are still formulas."""
        out = set()
        for ws, cell in self.new_cells:
            if ws.title in self.wb.sheetnames and isinstance(cell.value, str) and cell.value.startswith("=") and ws._cells.get((cell.row, cell.column)) is cell:
                out.add((ws.title, cell.row, cell.column))
        return out

    def ws(self, name: str | None) -> Any:
        if name is None:
            name = self.default_sheet
        if name is None:
            return self.wb.active or self.wb.worksheets[0]
        if name in self.wb.sheetnames:
            return self.wb[name]
        for n in self.wb.sheetnames:
            if n.lower() == str(name).lower():
                return self.wb[n]
        raise SkillError(f"no sheet named '{name}'; sheets: {', '.join(self.wb.sheetnames)}")

    def area(self, o: dict[str, Any], key: str = "range", alt: tuple[str, ...] = ("cell", "cells", "at")) -> tuple[Any, int, int, int, int]:
        """(worksheet, r1, c1, r2, c2) from o[key] (or an alternative key), honouring 'Sheet!A1:B2' and o['sheet']."""
        spec = o.get(key)
        if spec is None:
            for k in alt:
                if o.get(k) is not None:
                    spec = o[k]
                    break
        if spec is None:
            raise SkillError(f"{o.get('op')}: needs '{key}'")
        sheet, rest = split_sheet(str(spec))
        ws = self.ws(sheet or o.get("sheet"))
        try:
            r1, c1, r2, c2 = parse_range(rest)
        except ValueError as e:
            raise SkillError(f"{o.get('op')}: {e}") from e
        if r2 == MAX_ROW:
            r2 = max(ws.max_row, r1)
        if c2 == MAX_COL:
            c2 = max(ws.max_column, c1)
        return ws, r1, c1, r2, c2


def apply_ops(ctx: Ctx, ops: list[dict[str, Any]]) -> list[str]:
    log = []
    for i, o in enumerate(ops, 1):
        if not isinstance(o, dict) or "op" not in o:
            raise SkillError(f"operation {i} must be an object with an 'op' key")
        name = str(o["op"]).lower().strip()
        name = ALIASES.get(name, name)
        f = OPS.get(name)
        if f is None:
            raise SkillError(f"operation {i}: unknown op '{o['op']}'. Known: {', '.join(sorted(OPS))}")
        o = dict(o)
        o["op"] = name
        try:
            msg = f(ctx, o)
        except SkillError as e:
            raise SkillError(f"operation {i} ({name}): {e}") from None
        except (ValueError, KeyError, TypeError) as e:
            raise SkillError(f"operation {i} ({name}): {type(e).__name__}: {e}") from None
        log.append(f"{i}. {msg}")
    return log


# ── values ──────────────────────────────────────────────────────────────

_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_DT = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?$")


class LiteralText(str):
    """Text written as-is even when it starts with '=' (not a formula)."""


def to_cell_value(v: Any, raw: bool = False) -> tuple[Any, str | None]:
    """A JSON value as (cell value, number format to apply or None). '=' strings are formulas; ISO dates become dates."""
    if isinstance(v, dict):
        if "formula" in v:
            return normalize_formula(str(v["formula"])), v.get("format")
        if "value" in v:
            val, fmt = to_cell_value(v["value"], raw or bool(v.get("text")))
            return val, v.get("format", fmt)
        if "date" in v:
            return _dt.date.fromisoformat(str(v["date"])), v.get("format", "yyyy-mm-dd")
        raise SkillError(f"cannot write {v!r} into a cell")
    if isinstance(v, str) and raw:
        return (LiteralText(v) if v.startswith("=") else v), None
    if isinstance(v, str) and not raw:
        s = v
        if s.startswith("=") and len(s) > 1:
            return normalize_formula(s), None
        if _ISO_DATE.match(s):
            try:
                return _dt.date.fromisoformat(s), "yyyy-mm-dd"
            except ValueError:
                return s, None
        if _ISO_DT.match(s):
            try:
                return _dt.datetime.fromisoformat(s.replace(" ", "T")), "yyyy-mm-dd hh:mm"
            except ValueError:
                return s, None
        return s, None
    if isinstance(v, list):
        raise SkillError("use 'values' (a 2-D list) to write several cells")
    return v, None


def normalize_formula(f: str) -> str:
    """Validates a formula and adds the prefixes Excel needs for newer functions."""
    text = f if f.startswith("=") else "=" + f
    try:
        parse(text)
    except FormulaSyntaxError as e:
        raise SkillError(f"formula {text!r} is not valid: {e}") from e
    return add_prefixes(text)


def write_cell(ctx: Ctx, ws: Any, r: int, c: int, value: Any, fmt: str | None = None) -> None:
    cell = ws.cell(row=r, column=c)
    if cell.__class__.__name__ == "MergedCell":
        raise SkillError(f"{col_letter(c)}{r} is inside a merged range; write to its top-left cell")
    if type(value) is LiteralText:
        if cell.data_type == "f":
            ctx.formulas_changed = True
        cell.value = str(value)
        cell.data_type = "s"
        if fmt:
            cell.number_format = fmt
        return
    if isinstance(value, str) and value.startswith("="):
        ctx.formulas_changed = True
        ctx.new_cells.append((ws, cell))
    elif hasattr(cell, "data_type") and cell.data_type == "f":
        ctx.formulas_changed = True
    cell.value = value
    if fmt:
        cell.number_format = fmt


# ── writing cells ───────────────────────────────────────────────────────


@op("set")
def op_set(ctx: Ctx, o: dict[str, Any]) -> str:
    raw = bool(o.get("text") or o.get("raw"))
    if "values" in o:
        ws, r1, c1, _, _ = ctx.area(o)
        values = o["values"]
        if not isinstance(values, list):
            raise SkillError("'values' must be a list of rows")
        if values and not isinstance(values[0], list):
            values = [values] if not o.get("down") else [[v] for v in values]
        n = 0
        for i, row in enumerate(values):
            for j, v in enumerate(row):
                val, fmt = to_cell_value(v, raw)
                write_cell(ctx, ws, r1 + i, c1 + j, val, o.get("number_format") or fmt)
                n += 1
        if o.get("style"):
            apply_style(ws, r1, c1, r1 + len(values) - 1, c1 + max((len(r) for r in values), default=1) - 1, o["style"])
        return f"set {n} cells from {col_letter(c1)}{r1} on {ws.title}"
    ws, r1, c1, r2, c2 = ctx.area(o)
    if "formula" in o:
        val, fmt = normalize_formula(str(o["formula"])), None
    elif "value" in o:
        val, fmt = to_cell_value(o["value"], raw)
    else:
        raise SkillError("set needs 'value', 'formula' or 'values'")
    n = 0
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            v = val
            if isinstance(val, str) and val.startswith("=") and (r, c) != (r1, c1):
                v = translate(val, r - r1, c - c1)
            write_cell(ctx, ws, r, c, v, o.get("number_format") or fmt)
            n += 1
    if o.get("style"):
        apply_style(ws, r1, c1, r2, c2, o["style"])
    what = "formula" if isinstance(val, str) and val.startswith("=") else "value"
    return f"set {what} in {range_name(r1, c1, r2, c2)} on {ws.title}" + (f" ({n} cells)" if n > 1 else "")


@op("fill")
def op_fill(ctx: Ctx, o: dict[str, Any]) -> str:
    """Fills a range like Excel's Ctrl+D / Ctrl+R: with 'formula' or 'value', every cell gets it (relative references
    shift); without, each column is filled down from its top cell (a single row: rightwards from its first cell)."""
    ws, r1, c1, r2, c2 = ctx.area(o)
    if "formula" in o:
        base: Any = normalize_formula(str(o["formula"]))
    elif "value" in o:
        base = to_cell_value(o["value"])[0]
    elif "series" in o:
        s = o["series"]
        start, step = s.get("start", 1), s.get("step", 1)
        k = 0
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                write_cell(ctx, ws, r, c, start + k * step)
                k += 1
        return f"filled a series into {range_name(r1, c1, r2, c2)} on {ws.title}"
    else:
        copy_style = o.get("copy_style", True)
        down = r2 > r1 or c1 == c2
        sources = [(r1, c) for c in range(c1, c2 + 1)] if down else [(r1, c1)]
        filled, empty = 0, []
        for (sr, sc) in sources:
            src = ws.cell(sr, sc)
            v0 = src.value
            if v0 is None:
                empty.append(f"{col_letter(sc)}{sr}")
                continue
            style = src._style if copy_style else None
            targets = [(r, sc) for r in range(r1 + 1, r2 + 1)] if down else [(r1, c) for c in range(c1 + 1, c2 + 1)]
            for r, c in targets:
                v = translate(v0, r - sr, c - sc) if isinstance(v0, str) and v0.startswith("=") else v0
                write_cell(ctx, ws, r, c, v)
                if style is not None:
                    ws.cell(r, c)._style = style
                filled += 1
        if filled == 0:
            raise SkillError(f"{', '.join(empty) or col_letter(c1) + str(r1)} is empty; give 'formula' or 'value'")
        if empty:
            ctx.warnings.append(f"fill {range_name(r1, c1, r2, c2)}: {', '.join(empty[:8])} empty, so those columns were left as they were")
        how = "down from row " + str(r1) if down else "right from column " + col_letter(c1)
        return f"filled {range_name(r1, c1, r2, c2)} on {ws.title} ({how})"
    src_style = ws.cell(r1, c1)._style if o.get("copy_style", True) else None
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            v = base
            if isinstance(base, str) and base.startswith("="):
                v = translate(base, r - r1, c - c1)
            write_cell(ctx, ws, r, c, v)
            if src_style is not None and (r, c) != (r1, c1):
                ws.cell(r, c)._style = src_style
    return f"filled {range_name(r1, c1, r2, c2)} on {ws.title}"


@op("clear")
def op_clear(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.styles import DEFAULT_FONT  # noqa: F401
    from openpyxl.cell.cell import Cell

    ws, r1, c1, r2, c2 = ctx.area(o)
    what = o.get("what", "contents")
    n = 0
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws._cells.get((r, c))
            if cell is None or cell.__class__.__name__ == "MergedCell":
                continue
            if what in ("contents", "all", "values"):
                if cell.data_type == "f":
                    ctx.formulas_changed = True
                cell.value = None
            if what in ("formats", "all", "styles"):
                cell.style = "Normal"
            n += 1
    return f"cleared {what} of {range_name(r1, c1, r2, c2)} on {ws.title}"


@op("copy_range")
def op_copy_range(ctx: Ctx, o: dict[str, Any]) -> str:
    ws, r1, c1, r2, c2 = ctx.area(o, "from", ("range",))
    to_sheet, to_ref = split_sheet(str(o.get("to")))
    dest = ctx.ws(to_sheet or o.get("to_sheet") or ws.title)
    tr, tc = parse_cell(to_ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            src = ws.cell(r, c)
            v = src.value
            if isinstance(v, str) and v.startswith("="):
                v = translate(v, tr - r1, tc - c1)
            if hasattr(v, "text") and v.__class__.__name__ == "ArrayFormula":
                v = translate(v.text, tr - r1, tc - c1)
            write_cell(ctx, dest, tr + r - r1, tc + c - c1, v)
            if o.get("styles", True):
                dest.cell(tr + r - r1, tc + c - c1)._style = src._style
    return f"copied {range_name(r1, c1, r2, c2)} to {dest.title}!{col_letter(tc)}{tr}"


@op("move_range")
def op_move_range(ctx: Ctx, o: dict[str, Any]) -> str:
    ws, r1, c1, r2, c2 = ctx.area(o)
    rows, cols = int(o.get("rows", 0)), int(o.get("cols", 0))
    ws.move_range(range_name(r1, c1, r2, c2), rows=rows, cols=cols, translate=True)
    ctx.formulas_changed = True
    return f"moved {range_name(r1, c1, r2, c2)} by {rows} rows, {cols} columns on {ws.title}"


# ── structure: rows, columns, sheets ────────────────────────────────────


def _each_formula_cell(wb: Any) -> Any:
    from openpyxl.worksheet.formula import ArrayFormula

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    yield ws, cell, v, None
                elif isinstance(v, ArrayFormula):
                    yield ws, cell, v.text or "", v


def _rewrite_all(ctx: Ctx, fn: Callable[[str, str], str]) -> int:
    """Applies fn(formula, formula_sheet) → new formula to every formula-bearing part of the workbook."""
    wb = ctx.wb
    n = 0
    for ws, cell, text, arr in _each_formula_cell(wb):
        new = fn(text, ws.title)
        if new != text:
            n += 1
            if arr is not None:
                arr.text = new
            else:
                cell.value = new
    # defined names (workbook and sheet scope)
    for dn in list(_all_defined_names(wb)):
        owner, name, d = dn
        if d.attr_text:
            new = fn("=" + d.attr_text, owner.title if owner is not wb else None)[1:] if not d.attr_text.startswith("=") else fn(d.attr_text, None)
            if new != d.attr_text:
                d.attr_text = new
                n += 1
    for ws in wb.worksheets:
        # conditional formatting rule formulas
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.formula:
                    rule.formula = [fn("=" + f, ws.title)[1:] for f in rule.formula]
        for dv in ws.data_validations.dataValidation:
            for attr in ("formula1", "formula2"):
                f = getattr(dv, attr)
                if f and not f.startswith('"'):
                    setattr(dv, attr, fn("=" + f, ws.title)[1:])
        for ch in getattr(ws, "_charts", []):
            _rewrite_chart(ch, lambda f: fn("=" + f, ws.title)[1:])
    return n


def _all_defined_names(wb: Any) -> Any:
    for name, d in list(wb.defined_names.items()):
        yield wb, name, d
    for ws in wb.worksheets:
        for name, d in list(ws.defined_names.items()):
            yield ws, name, d


def _rewrite_chart(ch: Any, fn: Callable[[str], str]) -> None:
    def fix(ref_holder: Any) -> None:
        if ref_holder is None:
            return
        for kind in ("numRef", "strRef"):
            r = getattr(ref_holder, kind, None)
            if r is not None and getattr(r, "f", None):
                try:
                    r.f = fn(r.f)
                except Exception:  # noqa: BLE001 — leave odd chart references alone
                    pass

    for s in getattr(ch, "series", []) or []:
        fix(getattr(s, "val", None))
        fix(getattr(s, "cat", None))
        fix(getattr(s, "xVal", None))
        fix(getattr(s, "yVal", None))
        tx = getattr(s, "tx", None)
        if tx is not None and getattr(tx, "strRef", None) is not None and tx.strRef.f:
            try:
                tx.strRef.f = fn(tx.strRef.f)
            except Exception:  # noqa: BLE001
                pass


def _shift_multi(sqref: Any, axis: str, at: int, count: int) -> str | None:
    parts = []
    for rng in str(sqref).split():
        try:
            r1, c1, r2, c2 = parse_range(rng)
        except ValueError:
            parts.append(rng)
            continue
        new = shift_area(r1, c1, r2, c2, axis, at, count)
        if new is not None:
            parts.append(range_name(*new))
    return " ".join(parts) if parts else None


def structural(ctx: Ctx, ws: Any, axis: str, at: int, count: int) -> dict[str, int]:
    """Inserts (count > 0) or deletes (count < 0) rows/columns and rewrites every dependent reference."""
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.worksheet.cell_range import MultiCellRange

    wb = ctx.wb
    target = ws.title
    stats = {"formulas": 0, "merged": 0, "validations": 0, "conditional_formats": 0, "tables": 0, "names": 0, "charts": 0}
    # 1. formulas everywhere (and names, cf/dv formulas, charts)
    stats["formulas"] = _rewrite_all(ctx, lambda f, sheet: _safe_shift(f, sheet, target, axis, at, count))
    # array formula ranges on the target sheet
    from openpyxl.worksheet.formula import ArrayFormula

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, ArrayFormula) and cell.value.ref:
                try:
                    new = shift_area(*parse_range(cell.value.ref), axis, at, count)
                except ValueError:
                    new = None
                if new is not None:
                    cell.value.ref = range_name(*new)
    # 2. merged cells (unmerge before moving, re-merge after)
    merged = [tuple(parse_range(str(m))) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    # 3. move the cells
    if axis == "row":
        if count > 0:
            ws.insert_rows(at, count)
        else:
            ws.delete_rows(at, -count)
    else:
        if count > 0:
            ws.insert_cols(at, count)
        else:
            ws.delete_cols(at, -count)
    for m in merged:
        new = shift_area(*m, axis, at, count)
        if new is not None and (new[0], new[1]) != (new[2], new[3]):
            ws.merge_cells(range_name(*new))
            stats["merged"] += 1
    # 4. conditional formatting ranges
    old_cf = ws.conditional_formatting
    new_cf = ConditionalFormattingList()
    for cf in old_cf:
        sq = _shift_multi(cf.sqref, axis, at, count)
        if sq:
            for rule in cf.rules:
                new_cf.add(sq, rule)
            stats["conditional_formats"] += 1
    ws.conditional_formatting = new_cf
    # 5. data validations
    keep = []
    for dv in ws.data_validations.dataValidation:
        sq = _shift_multi(dv.sqref, axis, at, count)
        if sq:
            dv.sqref = MultiCellRange(sq)
            keep.append(dv)
            stats["validations"] += 1
    ws.data_validations.dataValidation = keep
    # 6. tables and autofilter
    for name, tbl in list(dict.items(ws.tables)):
        r1, c1, r2, c2 = parse_range(tbl.ref)
        new = shift_area(r1, c1, r2, c2, axis, at, count)
        if new is None:
            del ws.tables[name]
            ctx.warnings.append(f"table {name} was deleted with its cells")
            continue
        if axis == "col" and new[3] - new[1] != c2 - c1:
            _resize_table_columns(ws, tbl, r1, c1, c2, at, count)
        header = 0 if tbl.headerRowCount == 0 else 1
        if axis == "row" and count > 0 and r1 + header <= at <= r2 - (tbl.totalsRowCount or 0):
            _fill_table_rows(ctx, ws, new, header, at, count)
        tbl.ref = range_name(*new)
        if tbl.autoFilter is not None:
            tbl.autoFilter.ref = range_name(new[0], new[1], new[2] - (tbl.totalsRowCount or 0), new[3])
        stats["tables"] += 1
    if ws.auto_filter.ref:
        sq = _shift_multi(ws.auto_filter.ref, axis, at, count)
        ws.auto_filter.ref = sq
    # 7. print area and titles
    if ws.print_area:
        pa = []
        for part in str(ws.print_area).split(","):
            s, rng = split_sheet(part)
            try:
                new = shift_area(*parse_range(rng), axis, at, count)
            except ValueError:
                new = None
            if new is not None:
                pa.append(range_name(*new))
        ws.print_area = ",".join(pa) if pa else None
    # 8. row heights / column widths move with their rows and columns
    _shift_dimensions(ws, axis, at, count)
    # 9. chart and image anchors on this sheet
    for obj in list(getattr(ws, "_charts", [])) + list(getattr(ws, "_images", [])):
        _shift_anchor(obj, axis, at, count)
    # 10. hyperlinks keep their cell; refresh their ref
    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                cell.hyperlink.ref = cell.coordinate
    ctx.formulas_changed = True
    return stats


def _fill_table_rows(ctx: Ctx, ws: Any, area: tuple[int, int, int, int], header: int, at: int, count: int) -> None:
    """Rows inserted inside a table get its calculated-column formulas and the formatting of the row above, as in
    Excel."""
    import copy

    r1, c1, r2, c2 = area
    ref_row = at - 1 if at - 1 >= r1 + header else at + count
    if ref_row > r2:
        return
    for c in range(c1, c2 + 1):
        src = ws.cell(ref_row, c)
        v = src.value
        formula = isinstance(v, str) and v.startswith("=")
        if formula:
            # a calculated column: the other data rows hold the same formula, shifted
            other = next((r for r in range(r1 + header, r2 + 1) if r != ref_row and not (at <= r < at + count)), None)
            if other is not None:
                ov = ws.cell(other, c).value
                try:
                    formula = isinstance(ov, str) and translate(v, other - ref_row, 0) == ov
                except FormulaSyntaxError:
                    formula = False
        for r in range(at, at + count):
            cell = ws.cell(r, c)
            if formula and cell.value is None:
                write_cell(ctx, ws, r, c, translate(v, r - ref_row, 0))
            if src.has_style:
                cell._style = copy.copy(src._style)


def _safe_shift(f: str, sheet: str | None, target: str, axis: str, at: int, count: int) -> str:
    try:
        return shift_refs(f, sheet, target, axis, at, count)
    except FormulaSyntaxError:
        return f


def _resize_table_columns(ws: Any, tbl: Any, r1: int, c1: int, c2: int, at: int, count: int) -> None:
    from openpyxl.worksheet.table import TableColumn

    cols = list(tbl.tableColumns)
    if count > 0:
        pos = at - c1
        if 0 < pos <= len(cols):
            existing = {c.name for c in cols}
            new_cols = []
            for k in range(count):
                n = len(cols) + k + 1
                name = f"Column{n}"
                while name in existing:
                    n += 1
                    name = f"Column{n}"
                existing.add(name)
                new_cols.append(TableColumn(id=0, name=name))
                ws.cell(r1, at + k).value = name
            cols = cols[:pos] + new_cols + cols[pos:]
    else:
        n = -count
        lo = max(at, c1) - c1
        hi = min(at + n - 1, c2) - c1
        cols = [c for i, c in enumerate(cols) if not (lo <= i <= hi)]
    for i, c in enumerate(cols, 1):
        c.id = i
    tbl.tableColumns = cols


def _shift_dimensions(ws: Any, axis: str, at: int, count: int) -> None:
    import copy

    if axis == "row":
        dims = {k: v for k, v in ws.row_dimensions.items()}
        moved = {}
        for k, d in dims.items():
            if k < at:
                continue
            new = k + count if count > 0 else (None if k < at - count else k + count)
            moved[k] = (new, d)
        for k in moved:
            del ws.row_dimensions[k]
        for k, (new, d) in moved.items():
            if new is not None and new >= 1:
                nd = copy.copy(d)
                nd.index = new
                ws.row_dimensions[new] = nd
    else:
        dims = {}
        for key, d in list(ws.column_dimensions.items()):
            try:
                lo, hi = col_index(d.min and col_letter(d.min) or key), col_index(d.max and col_letter(d.max) or key)
            except ValueError:
                continue
            dims[key] = (lo, hi, d)
        changed = []
        for key, (lo, hi, d) in dims.items():
            if hi < at:
                continue
            changed.append((key, lo, hi, d))
        for key, *_ in changed:
            del ws.column_dimensions[key]
        for key, lo, hi, d in changed:
            new = shift_area(1, lo, 1, hi, "col", at, count)
            if new is None:
                continue
            nd = copy.copy(d)
            nd.index = col_letter(new[1])
            nd.min, nd.max = new[1], new[3]
            ws.column_dimensions[col_letter(new[1])] = nd


def _shift_anchor(obj: Any, axis: str, at: int, count: int) -> None:
    anchor = getattr(obj, "anchor", None)
    if anchor is None:
        return
    if isinstance(anchor, str):
        try:
            r, c = parse_cell(anchor)
        except ValueError:
            return
        if axis == "row" and r >= at:
            r = max(1, r + count)
        elif axis == "col" and c >= at:
            c = max(1, c + count)
        obj.anchor = f"{col_letter(c)}{r}"
        return
    for part in ("_from", "to"):
        m = getattr(anchor, part, None)
        if m is None:
            continue
        if axis == "row" and m.row + 1 >= at:
            m.row = max(0, m.row + count)
        elif axis == "col" and m.col + 1 >= at:
            m.col = max(0, m.col + count)


def _at_count(o: dict[str, Any], axis: str) -> tuple[int, int]:
    if axis == "row":
        if "rows" in o:
            a, b = parse_rows(str(o["rows"]))
            return a, b - a + 1
        at = int(o.get("at", o.get("row", 0)))
    else:
        if "columns" in o or "cols" in o:
            a, b = parse_cols(str(o.get("columns", o.get("cols"))))
            return a, b - a + 1
        v = o.get("at", o.get("column", o.get("col")))
        at = col_index(v) if isinstance(v, str) and not v.isdigit() else int(v or 0)
    count = int(o.get("count", 1))
    if at < 1 or count < 1:
        raise SkillError("needs 'at' (a row number or column letter) and a positive 'count'")
    return at, count


def _structural_op(ctx: Ctx, o: dict[str, Any], axis: str, sign: int) -> str:
    ws = ctx.ws(o.get("sheet"))
    at, count = _at_count(o, axis)
    stats = structural(ctx, ws, axis, at, sign * count)
    what = "rows" if axis == "row" else "columns"
    where = f"{at}" if axis == "row" else col_letter(at)
    verb = "inserted" if sign > 0 else "deleted"
    extra = ", ".join(f"{v} {k.replace('_', ' ')}" for k, v in stats.items() if v)
    return f"{verb} {count} {what} at {where} on {ws.title}" + (f"; updated {extra}" if extra else "")


@op("insert_rows")
def op_insert_rows(ctx: Ctx, o: dict[str, Any]) -> str:
    return _structural_op(ctx, o, "row", 1)


@op("delete_rows")
def op_delete_rows(ctx: Ctx, o: dict[str, Any]) -> str:
    return _structural_op(ctx, o, "row", -1)


@op("insert_cols")
def op_insert_cols(ctx: Ctx, o: dict[str, Any]) -> str:
    return _structural_op(ctx, o, "col", 1)


@op("delete_cols")
def op_delete_cols(ctx: Ctx, o: dict[str, Any]) -> str:
    return _structural_op(ctx, o, "col", -1)


@op("add_sheet")
def op_add_sheet(ctx: Ctx, o: dict[str, Any]) -> str:
    name = str(o.get("name") or o.get("sheet") or "")
    if not name:
        raise SkillError("add_sheet needs 'name'")
    _check_sheet_name(ctx, name)
    idx = o.get("index")
    ws = ctx.wb.create_sheet(name, int(idx) - 1 if idx is not None else None)
    return f"added sheet {ws.title}"


def _check_sheet_name(ctx: Ctx, name: str) -> None:
    if len(name) > 31 or re.search(r"[\[\]:*?/\\]", name) or not name.strip() or name.startswith("'") or name.endswith("'"):
        raise SkillError(f"'{name}' is not a valid sheet name (max 31 characters, none of [ ] : * ? / \\)")
    if any(n.lower() == name.lower() for n in ctx.wb.sheetnames):
        raise SkillError(f"a sheet named '{name}' already exists")


@op("rename_sheet")
def op_rename_sheet(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    new = str(o.get("to") or o.get("name") or "")
    if not new:
        raise SkillError("rename_sheet needs 'to'")
    old = ws.title
    if new.lower() != old.lower():
        _check_sheet_name(ctx, new)
    n = _rewrite_all(ctx, lambda f, _s: rename_sheet_refs(f, old, new))
    ws.title = new
    ctx.formulas_changed = True
    pivots = 0
    seen = set()
    for w in ctx.wb.worksheets:
        for pt in getattr(w, "_pivots", []) or []:
            src = getattr(getattr(getattr(pt, "cache", None), "cacheSource", None), "worksheetSource", None)
            if src is not None and id(src) not in seen and getattr(src, "sheet", None) and src.sheet.lower() == old.lower():
                seen.add(id(src))
                src.sheet = new
                pivots += 1
    return f"renamed sheet {old} to {new}; updated {n} formulas and names" + (f" and {pivots} pivot cache source{'s' if pivots != 1 else ''}" if pivots else "")


@op("copy_sheet")
def op_copy_sheet(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    new = ctx.wb.copy_worksheet(ws)
    name = o.get("to") or o.get("name")
    if name:
        _check_sheet_name(ctx, str(name))
        new.title = str(name)
    if o.get("index") is not None:
        ctx.wb.move_sheet(new, int(o["index"]) - 1 - ctx.wb.index(new))
    for row in new.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                ctx.formulas_changed = True
                break
    if getattr(ws, "_charts", None):
        ctx.warnings.append(f"charts on {ws.title} are not copied (openpyxl limitation); add them with the chart op")
    return f"copied sheet {ws.title} to {new.title}"


@op("delete_sheet")
def op_delete_sheet(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet") or o.get("name"))
    if len(ctx.wb.worksheets) == 1:
        raise SkillError("a workbook needs at least one sheet")
    name = ws.title
    ctx.wb.remove(ws)
    n = _rewrite_all(ctx, lambda f, _s: drop_sheet_refs(f, name))
    ctx.formulas_changed = True
    if n:
        ctx.warnings.append(f"{n} formulas referred to {name} and now contain #REF!")
    return f"deleted sheet {name}" + (f"; {n} formulas now show #REF!" if n else "")


@op("move_sheet")
def op_move_sheet(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    idx = int(o.get("index", o.get("to", 1)))
    ctx.wb.move_sheet(ws, idx - 1 - ctx.wb.index(ws))
    return f"moved sheet {ws.title} to position {idx}"


@op("reorder")
def op_reorder(ctx: Ctx, o: dict[str, Any]) -> str:
    order = o.get("order")
    if not isinstance(order, list):
        raise SkillError("reorder needs 'order': [sheet names]")
    sheets = [ctx.ws(n) for n in order]
    rest = [w for w in ctx.wb._sheets if w not in sheets]
    ctx.wb._sheets = sheets + rest
    return "reordered sheets: " + ", ".join(w.title for w in ctx.wb._sheets)


@op("sheet_state")
def op_sheet_state(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    state = o.get("state", "hidden")
    state = {"very_hidden": "veryHidden", "veryhidden": "veryHidden"}.get(state, state)
    if state not in ("visible", "hidden", "veryHidden"):
        raise SkillError("state is visible, hidden or veryHidden")
    ws.sheet_state = state
    if o.get("tab_color"):
        ws.sheet_properties.tabColor = _hex(o["tab_color"])
    return f"sheet {ws.title} is now {state}"


@op("tab_color")
def op_tab_color(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    ws.sheet_properties.tabColor = _hex(o.get("color", "#1F4E78"))
    return f"set the tab color of {ws.title}"


# ── formatting ──────────────────────────────────────────────────────────

_NAMED_COLORS = {
    "black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "00B050", "blue": "0070C0", "yellow": "FFFF00",
    "orange": "FFC000", "gray": "808080", "grey": "808080", "lightgray": "D9D9D9", "lightgrey": "D9D9D9", "darkblue": "1F4E78",
    "lightblue": "DDEBF7", "lightgreen": "E2EFDA", "lightred": "FCE4E4", "lightyellow": "FFF2CC", "purple": "7030A0", "navy": "1F3864",
}


def _hex(c: Any) -> str:
    s = str(c).strip()
    low = s.lower()
    if low in _NAMED_COLORS:
        return "FF" + _NAMED_COLORS[low]
    s = s.lstrip("#")
    if re.fullmatch(r"[0-9A-Fa-f]{3}", s):
        s = "".join(ch * 2 for ch in s)
    if re.fullmatch(r"[0-9A-Fa-f]{6}", s):
        return "FF" + s.upper()
    if re.fullmatch(r"[0-9A-Fa-f]{8}", s):
        return s.upper()
    raise SkillError(f"bad color {c!r} (use #RRGGBB or a name like red, lightgray)")


def make_font(spec: Any, base: Any = None) -> Any:
    from openpyxl.styles import Font

    if spec is True:
        spec = {"bold": True}
    if not isinstance(spec, dict):
        raise SkillError("font must be an object like {\"bold\": true, \"color\": \"#C00000\"}")
    kw: dict[str, Any] = {}
    if base is not None:
        kw = {"name": base.name, "sz": base.sz, "b": base.b, "i": base.i, "u": base.u, "strike": base.strike, "color": base.color, "vertAlign": base.vertAlign}
    for k, v in spec.items():
        if k == "bold":
            kw["b"] = bool(v)
        elif k == "italic":
            kw["i"] = bool(v)
        elif k == "underline":
            kw["u"] = ("single" if v is True else v) if v else None
        elif k in ("strike", "strikethrough"):
            kw["strike"] = bool(v)
        elif k == "size":
            kw["sz"] = float(v)
        elif k == "name":
            kw["name"] = str(v)
        elif k == "color":
            kw["color"] = _hex(v)
        elif k in ("superscript", "subscript"):
            kw["vertAlign"] = k if v else None
    return Font(**kw)


def make_fill(spec: Any) -> Any:
    from openpyxl.styles import PatternFill

    if spec is None or spec is False or spec == "none":
        return PatternFill(fill_type=None)
    if isinstance(spec, str):
        return PatternFill("solid", fgColor=_hex(spec), bgColor=_hex(spec))
    if isinstance(spec, dict):
        color = _hex(spec.get("color", "FFFFFF"))
        return PatternFill(spec.get("pattern", "solid"), fgColor=color, bgColor=_hex(spec.get("background", spec.get("color", "FFFFFF"))))
    raise SkillError("fill is a color string or {\"color\": …, \"pattern\": \"solid\"}")


def make_side(spec: Any) -> Any:
    from openpyxl.styles import Side

    if spec is None or spec is False or spec == "none":
        return Side(style=None)
    if isinstance(spec, str):
        parts = spec.split()
        style = parts[0]
        color = _hex(parts[1]) if len(parts) > 1 else "FF000000"
        return Side(style=style, color=color)
    return Side(style=spec.get("style", "thin"), color=_hex(spec.get("color", "#000000")))


def make_alignment(spec: dict[str, Any], base: Any = None) -> Any:
    from openpyxl.styles import Alignment

    kw: dict[str, Any] = {}
    if base is not None:
        kw = {"horizontal": base.horizontal, "vertical": base.vertical, "wrap_text": base.wrap_text, "indent": base.indent, "text_rotation": base.text_rotation, "shrink_to_fit": base.shrink_to_fit}
    for k, v in spec.items():
        if k in ("horizontal", "h"):
            kw["horizontal"] = v
        elif k in ("vertical", "v"):
            kw["vertical"] = v
        elif k in ("wrap", "wrap_text"):
            kw["wrap_text"] = bool(v)
        elif k == "indent":
            kw["indent"] = int(v)
        elif k in ("rotation", "text_rotation"):
            kw["text_rotation"] = int(v)
        elif k == "shrink":
            kw["shrink_to_fit"] = bool(v)
    return Alignment(**kw)


def apply_style(ws: Any, r1: int, c1: int, r2: int, c2: int, st: dict[str, Any]) -> int:
    """Applies font/fill/border/align/number_format to every cell of an area (borders: all, outline, inner, or sides)."""
    from openpyxl.styles import Border

    n = 0
    border = st.get("border")
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if cell.__class__.__name__ == "MergedCell" and not border:
                continue
            if "font" in st or "bold" in st or "italic" in st or "color" in st:
                spec = dict(st.get("font") or {})
                for k in ("bold", "italic", "color", "size"):
                    if k in st and k not in spec:
                        spec[k] = st[k]
                cell.font = make_font(spec, cell.font)
            if "fill" in st or "background" in st:
                cell.fill = make_fill(st.get("fill", st.get("background")))
            if "align" in st or "alignment" in st or "wrap" in st:
                spec = dict(st.get("align") or st.get("alignment") or {})
                if "wrap" in st:
                    spec["wrap"] = st["wrap"]
                cell.alignment = make_alignment(spec, cell.alignment)
            if "number_format" in st or "format" in st:
                cell.number_format = str(st.get("number_format", st.get("format")))
            if border:
                cell.border = _border_for(cell.border, border, r, c, r1, c1, r2, c2)
            if "locked" in st:
                from openpyxl.styles import Protection

                cell.protection = Protection(locked=bool(st["locked"]), hidden=cell.protection.hidden)
            n += 1
    return n


def _border_for(current: Any, spec: Any, r: int, c: int, r1: int, c1: int, r2: int, c2: int) -> Any:
    from openpyxl.styles import Border

    left, right, top, bottom = current.left, current.right, current.top, current.bottom
    if isinstance(spec, str):
        spec = {"all": spec}
    if not isinstance(spec, dict):
        raise SkillError("border is a style like 'thin' or {\"outline\": \"medium\", \"inner\": \"thin\", \"bottom\": \"double\"}")
    if "all" in spec:
        s = make_side(spec["all"])
        left = right = top = bottom = s
    if "inner" in spec:
        s = make_side(spec["inner"])
        if c > c1:
            left = s
        if c < c2:
            right = s
        if r > r1:
            top = s
        if r < r2:
            bottom = s
    if "outline" in spec:
        s = make_side(spec["outline"])
        if c == c1:
            left = s
        if c == c2:
            right = s
        if r == r1:
            top = s
        if r == r2:
            bottom = s
    for side in ("left", "right", "top", "bottom"):
        if side in spec:
            s = make_side(spec[side])
            edge = {"left": c == c1, "right": c == c2, "top": r == r1, "bottom": r == r2}[side]
            if edge:
                if side == "left":
                    left = s
                elif side == "right":
                    right = s
                elif side == "top":
                    top = s
                else:
                    bottom = s
    return Border(left=left, right=right, top=top, bottom=bottom)


@op("style")
def op_style(ctx: Ctx, o: dict[str, Any]) -> str:
    ws, r1, c1, r2, c2 = ctx.area(o)
    st = {k: v for k, v in o.items() if k not in ("op", "range", "cell", "sheet")}
    if "format" in st and "number_format" not in st:
        st["number_format"] = st.pop("format")
    if not st:
        raise SkillError("style needs font, fill, border, align or number_format")
    apply_style(ws, r1, c1, r2, c2, st)
    return f"styled {range_name(r1, c1, r2, c2)} on {ws.title} ({', '.join(sorted(st))})"


@op("column_width")
def op_column_width(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    spec = o.get("columns", o.get("column", o.get("cols")))
    if spec is None:
        raise SkillError("column_width needs 'columns' (e.g. 'A:C') and 'width'")
    widths = o.get("widths")
    a, b = parse_cols(str(spec))
    for i, c in enumerate(range(a, b + 1)):
        w = widths[i] if isinstance(widths, list) and i < len(widths) else o.get("width")
        if w is None:
            raise SkillError("column_width needs 'width'")
        ws.column_dimensions[col_letter(c)].width = float(w)
        if "hidden" in o:
            ws.column_dimensions[col_letter(c)].hidden = bool(o["hidden"])
    return f"set width of columns {col_letter(a)}:{col_letter(b)} on {ws.title}"


@op("hide_columns")
def op_hide_columns(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    a, b = parse_cols(str(o.get("columns", o.get("column"))))
    for c in range(a, b + 1):
        ws.column_dimensions[col_letter(c)].hidden = bool(o.get("hidden", True))
    return f"{'hid' if o.get('hidden', True) else 'showed'} columns {col_letter(a)}:{col_letter(b)} on {ws.title}"


@op("hide_rows")
def op_hide_rows(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    a, b = parse_rows(str(o.get("rows", o.get("row"))))
    for r in range(a, b + 1):
        ws.row_dimensions[r].hidden = bool(o.get("hidden", True))
    return f"{'hid' if o.get('hidden', True) else 'showed'} rows {a}:{b} on {ws.title}"


@op("group")
def op_group(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    hidden = bool(o.get("collapsed", False))
    if "rows" in o:
        a, b = parse_rows(str(o["rows"]))
        ws.row_dimensions.group(a, b, outline_level=int(o.get("level", 1)), hidden=hidden)
        return f"grouped rows {a}:{b} on {ws.title}"
    a, b = parse_cols(str(o.get("columns")))
    ws.column_dimensions.group(col_letter(a), col_letter(b), outline_level=int(o.get("level", 1)), hidden=hidden)
    return f"grouped columns {col_letter(a)}:{col_letter(b)} on {ws.title}"


@op("row_height")
def op_row_height(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    a, b = parse_rows(str(o.get("rows", o.get("row"))))
    for r in range(a, b + 1):
        ws.row_dimensions[r].height = float(o["height"])
    return f"set height of rows {a}:{b} on {ws.title}"


def display_width(value: Any, number_format: str | None, bold: bool = False, size: float = 11) -> float:
    """Approximate width in Excel character units of a cell's displayed text."""
    from _numfmt import format_value

    if value is None:
        return 0
    if isinstance(value, str) and value.startswith("="):
        return 0
    try:
        text = format_value(value, number_format or "General")[0]
    except Exception:  # noqa: BLE001
        text = str(value)
    lines = str(text).split("\n")
    longest = max(lines, key=len)
    w = 0.0
    for ch in longest:
        if ch in "il.,:;|!'`":
            w += 0.5
        elif ch in "mwMW@%":
            w += 1.3
        elif ch.isupper():
            w += 1.15
        elif ord(ch) > 0x2E80:
            w += 2.0
        else:
            w += 1.0
    w *= (1.1 if bold else 1.0) * (size / 11)
    return w


def _header_cells(ws: Any) -> set[tuple[int, int]]:
    """Header cells that show a filter button: Excel tables' header rows and the autofilter's first row."""
    out: set[tuple[int, int]] = set()
    refs = []
    try:
        for t in dict.values(ws.tables):
            if (t.headerRowCount is None or t.headerRowCount) and t.ref:
                refs.append(t.ref)
    except Exception:  # noqa: BLE001
        pass
    if ws.auto_filter is not None and ws.auto_filter.ref:
        refs.append(ws.auto_filter.ref)
    for ref in refs:
        try:
            r1, c1, _r2, c2 = parse_range(ref)
        except ValueError:
            continue
        for c in range(c1, c2 + 1):
            out.add((r1, c))
    return out


def autofit(ws: Any, cols: tuple[int, int] | None = None, min_width: float = 6, max_width: float = 60, values: dict[tuple[int, int], Any] | None = None) -> list[str]:
    """Sets column widths from the widest displayed value (formula cells use computed values when given).

    Header cells with a filter button (tables, autofilter) get room for the button and are measured bold (table
    styles make them bold); wrapped cells are at least as wide as their longest word, so words never break."""
    widths: dict[int, float] = {}
    merged = set()
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                merged.add((r, c))
    headers = _header_cells(ws)
    for row in ws.iter_rows():
        for cell in row:
            c = cell.column
            if cols and not (cols[0] <= c <= cols[1]):
                continue
            if (cell.row, c) in merged:
                continue
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and values is not None:
                v = values.get((cell.row, c))
            if v is None or (isinstance(v, str) and v.startswith("=") and cell.data_type == "f"):
                continue
            f = cell.font
            is_header = (cell.row, c) in headers
            bold = bool(f and f.b) or is_header
            size = float(f.sz) if f and f.sz else 11
            w = display_width(v, cell.number_format, bold, size)
            if cell.alignment is not None and cell.alignment.wrap_text and isinstance(v, str):
                words = [x for x in re.split(r"\s+", v) if x]
                longest = max((display_width(x, None, bold, size) for x in words), default=0)
                w = max(min(w, 40), longest)
            elif cell.alignment is not None and cell.alignment.wrap_text:
                w = min(w, 40)
            if is_header:
                w += 2.2  # the filter button
            widths[c] = max(widths.get(c, 0), w)
    done = []
    for c, w in widths.items():
        ws.column_dimensions[col_letter(c)].width = round(min(max(w + 2, min_width), max_width), 1)
        done.append(col_letter(c))
    return done


@op("autofit")
def op_autofit(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    spec = o.get("columns", "all")
    cols = None if spec in (None, "all") else parse_cols(str(spec))
    done = autofit(ws, cols, float(o.get("min", 6)), float(o.get("max", 60)))
    ctx.autofit_pending = getattr(ctx, "autofit_pending", []) + [(ws, cols, float(o.get("min", 6)), float(o.get("max", 60)))]
    return f"auto-fitted {len(done)} columns on {ws.title}"


@op("freeze")
def op_freeze(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    if o.get("cell") in (None, "", "none", False) and "rows" not in o and "cols" not in o:
        ws.freeze_panes = None
        return f"unfroze panes on {ws.title}"
    if "cell" in o:
        cell = str(o["cell"]).upper()
    else:
        cell = f"{col_letter(int(o.get('cols', 0)) + 1)}{int(o.get('rows', 0)) + 1}"
    ws.freeze_panes = cell
    return f"froze panes at {cell} on {ws.title}"


@op("unfreeze")
def op_unfreeze(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    ws.freeze_panes = None
    return f"unfroze panes on {ws.title}"


@op("autofilter")
def op_autofilter(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    if not o.get("range"):
        ws.auto_filter.ref = None
        return f"removed the filter on {ws.title}"
    _, r1, c1, r2, c2 = ctx.area(o)
    ws.auto_filter.ref = range_name(r1, c1, r2, c2)
    return f"added a filter on {ws.title}!{ws.auto_filter.ref}"


@op("merge")
def op_merge(ctx: Ctx, o: dict[str, Any]) -> str:
    ws, r1, c1, r2, c2 = ctx.area(o)
    ws.merge_cells(range_name(r1, c1, r2, c2))
    if o.get("align", True):
        from openpyxl.styles import Alignment

        tl = ws.cell(r1, c1)
        tl.alignment = Alignment(horizontal=o.get("horizontal", "center"), vertical=o.get("vertical", "center"), wrap_text=tl.alignment.wrap_text)
    return f"merged {range_name(r1, c1, r2, c2)} on {ws.title}"


@op("unmerge")
def op_unmerge(ctx: Ctx, o: dict[str, Any]) -> str:
    ws, r1, c1, r2, c2 = ctx.area(o)
    n = 0
    for m in list(ws.merged_cells.ranges):
        if not (m.max_row < r1 or m.min_row > r2 or m.max_col < c1 or m.min_col > c2):
            ws.unmerge_cells(str(m))
            n += 1
    return f"unmerged {n} ranges in {range_name(r1, c1, r2, c2)} on {ws.title}"


# ── data operations ─────────────────────────────────────────────────────


def _sort_key(v: Any) -> tuple:
    if v is None or v == "":
        return (3, 0)
    if isinstance(v, bool):
        return (2, v)
    if isinstance(v, (int, float)):
        return (0, v)
    if isinstance(v, (_dt.date, _dt.datetime)):
        return (0, v.toordinal() if isinstance(v, _dt.date) and not isinstance(v, _dt.datetime) else v.timestamp() / 86400)
    return (1, str(v).lower())


@op("sort")
def op_sort(ctx: Ctx, o: dict[str, Any]) -> str:
    """Sorts the rows of a range by one or more columns; formulas move with their rows (relative references shift)."""
    ws, r1, c1, r2, c2 = ctx.area(o)
    if o.get("header"):
        r1 += 1
    by = o.get("by") or [{"column": col_letter(c1)}]
    if isinstance(by, (str, dict)):
        by = [by]
    keys = []
    for b in by:
        if isinstance(b, str):
            b = {"column": b}
        col = b.get("column", col_letter(c1))
        ci = col_index(col) if isinstance(col, str) and not col.isdigit() else c1 + int(col) - 1
        if not (c1 <= ci <= c2):
            raise SkillError(f"sort column {col} is outside {range_name(r1, c1, r2, c2)}")
        desc = str(b.get("order", "asc")).lower().startswith("desc") or bool(b.get("descending"))
        keys.append((ci, desc))
    rows = []
    for r in range(r1, r2 + 1):
        cells = [ws.cell(r, c) for c in range(c1, c2 + 1)]
        rows.append((r, [(c.value, c._style, c.hyperlink, c.comment) for c in cells]))
    values_for_sort = {}
    for r, cells in rows:
        values_for_sort[r] = [cells[ci - c1][0] for ci, _ in keys]
    order = [r for r, _ in rows]
    for idx in range(len(keys) - 1, -1, -1):
        ci, desc = keys[idx]
        blanks = [r for r in order if values_for_sort[r][idx] in (None, "")]
        filled = [r for r in order if values_for_sort[r][idx] not in (None, "")]
        filled.sort(key=lambda r: _sort_key(values_for_sort[r][idx]), reverse=desc)
        order = filled + blanks  # blanks always last, like Excel
    data = dict(rows)
    for new_r, old_r in zip(range(r1, r2 + 1), order):
        for j, (v, st, link, com) in enumerate(data[old_r]):
            if isinstance(v, str) and v.startswith("="):
                v = translate(v, new_r - old_r, 0)
                ctx.formulas_changed = True
            cell = ws.cell(new_r, c1 + j)
            cell.value = v
            cell._style = st
            cell.hyperlink = link
            cell.comment = com
    return f"sorted {range_name(r1, c1, r2, c2)} on {ws.title} by " + ", ".join(f"{col_letter(ci)}{' desc' if d else ''}" for ci, d in keys)


@op("find_replace")
def op_find_replace(ctx: Ctx, o: dict[str, Any]) -> str:
    find = o.get("find")
    if find is None or find == "":
        raise SkillError("find_replace needs 'find'")
    repl = str(o.get("replace", ""))
    flags = 0 if o.get("match_case") else re.I
    pat = re.compile(str(find) if o.get("regex") else re.escape(str(find)), flags)
    where = o.get("in", "values")
    whole = bool(o.get("whole_cell"))
    if o.get("range"):
        ws, r1, c1, r2, c2 = ctx.area(o)
        targets = [(ws, (r1, c1, r2, c2))]
    elif o.get("sheet"):
        targets = [(ctx.ws(o["sheet"]), None)]
    else:
        targets = [(w, None) for w in ctx.wb.worksheets]
    n = cells = 0
    for ws, area in targets:
        for row in _iter(ws, area):
            for cell in row:
                v = cell.value
                if not isinstance(v, str):
                    if isinstance(v, (int, float)) and not isinstance(v, bool) and where in ("values", "all") and whole and pat.fullmatch(str(v)):
                        cell.value, _ = to_cell_value(repl)
                        n += 1
                        cells += 1
                    continue
                is_formula = v.startswith("=")
                if is_formula and where not in ("formulas", "all"):
                    continue
                if not is_formula and where not in ("values", "all"):
                    continue
                if whole:
                    if pat.fullmatch(v):
                        new = pat.sub(repl, v) if o.get("regex") else repl
                        cell.value = new if is_formula else to_cell_value(new)[0]
                        n += 1
                        cells += 1
                    continue
                new, k = pat.subn(repl, v)
                if k:
                    if is_formula:
                        new = normalize_formula(new)
                        ctx.formulas_changed = True
                    cell.value = new
                    n += k
                    cells += 1
    return f"replaced {n} occurrence{'s' if n != 1 else ''} of {find!r} in {cells} cell{'s' if cells != 1 else ''}"


def _iter(ws: Any, area: tuple[int, int, int, int] | None) -> Any:
    if area is None:
        return ws.iter_rows()
    r1, c1, r2, c2 = area
    return ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)


# ── rules: conditional formats, validation, names, tables ───────────────


def _dxf(o: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.styles import Font, PatternFill

    kw: dict[str, Any] = {}
    fill = o.get("fill", o.get("background"))
    if fill:
        c = _hex(fill)
        kw["fill"] = PatternFill(start_color=c, end_color=c, fill_type="solid")
    font = o.get("font")
    if font or o.get("color") or o.get("bold"):
        spec = dict(font or {})
        if o.get("color"):
            spec.setdefault("color", o["color"])
        if o.get("bold"):
            spec.setdefault("bold", True)
        kw["font"] = make_font(spec)
    if not kw:
        c = _hex("#FFC7CE")
        kw["fill"] = PatternFill(start_color=c, end_color=c, fill_type="solid")
        kw["font"] = Font(color="FF9C0006")
    return kw


_CF_OPS = {">": "greaterThan", ">=": "greaterThanOrEqual", "<": "lessThan", "<=": "lessThanOrEqual", "=": "equal", "==": "equal", "!=": "notEqual", "<>": "notEqual", "between": "between", "not_between": "notBetween"}


@op("conditional_format")
def op_conditional_format(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule
    from openpyxl.styles.differential import DifferentialStyle

    ws, r1, c1, r2, c2 = ctx.area(o)
    rng = range_name(r1, c1, r2, c2)
    t = str(o.get("type", "cell")).lower()
    first = f"{col_letter(c1)}{r1}"
    if t in ("cell", "cell_is", "value"):
        oper = _CF_OPS.get(str(o.get("operator", ">")), o.get("operator", "greaterThan"))
        vals = o.get("values", o.get("value"))
        if not isinstance(vals, list):
            vals = [vals]
        formula = [_cf_value(v) for v in vals]
        rule = CellIsRule(operator=oper, formula=formula, **_dxf(o))
    elif t == "formula":
        f = str(o.get("formula", "")).lstrip("=")
        if not f:
            raise SkillError("formula rule needs 'formula' (relative to the first cell, e.g. \"$C2>100\")")
        rule = FormulaRule(formula=[add_prefixes("=" + f)[1:]], **_dxf(o))
    elif t in ("color_scale", "colorscale", "scale"):
        colors = o.get("colors") or ["#F8696B", "#FFEB84", "#63BE7B"]
        if len(colors) == 2:
            rule = ColorScaleRule(start_type="min", start_color=_hex(colors[0]), end_type="max", end_color=_hex(colors[1]))
        else:
            rule = ColorScaleRule(start_type="min", start_color=_hex(colors[0]), mid_type="percentile", mid_value=50, mid_color=_hex(colors[1]), end_type="max", end_color=_hex(colors[2]))
    elif t in ("data_bar", "databar", "bar"):
        rule = DataBarRule(start_type="min", end_type="max", color=_hex(o.get("color", "#638EC6"))[2:], showValue=True)
    elif t in ("icon_set", "icons"):
        rule = IconSetRule(o.get("icons", "3TrafficLights1"), "percent", [0, 33, 67], showValue=True, reverse=bool(o.get("reverse")))
    elif t in ("top", "bottom", "top10"):
        rank = int(o.get("rank", o.get("count", 10)))
        rule = Rule(type="top10", rank=rank, bottom=t == "bottom", percent=bool(o.get("percent")), dxf=DifferentialStyle(**_dxf(o)))
    elif t in ("above_average", "below_average"):
        rule = Rule(type="aboveAverage", aboveAverage=t == "above_average", dxf=DifferentialStyle(**_dxf(o)))
    elif t in ("duplicates", "duplicate"):
        rule = Rule(type="duplicateValues", dxf=DifferentialStyle(**_dxf(o)))
    elif t == "unique":
        rule = Rule(type="uniqueValues", dxf=DifferentialStyle(**_dxf(o)))
    elif t in ("text", "contains"):
        txt = str(o.get("text", o.get("value", "")))
        esc = txt.replace('"', '""')
        rule = Rule(type="containsText", operator="containsText", text=txt, formula=[f'NOT(ISERROR(SEARCH("{esc}",{first})))'], dxf=DifferentialStyle(**_dxf(o)))
    elif t == "blanks":
        rule = Rule(type="containsBlanks", formula=[f"LEN(TRIM({first}))=0"], dxf=DifferentialStyle(**_dxf(o)))
    elif t == "errors":
        rule = Rule(type="containsErrors", formula=[f"ISERROR({first})"], dxf=DifferentialStyle(**_dxf(o)))
    else:
        raise SkillError(f"unknown conditional format type '{t}' (cell, formula, color_scale, data_bar, icon_set, top, bottom, above_average, below_average, duplicates, unique, text, blanks, errors)")
    if o.get("stop"):
        rule.stopIfTrue = True
    ws.conditional_formatting.add(rng, rule)
    return f"added a {t} conditional format on {ws.title}!{rng}"


def _cf_value(v: Any) -> str:
    if isinstance(v, str):
        if v.startswith("="):
            return add_prefixes(v)[1:]
        try:
            float(v)
            return v
        except ValueError:
            return '"' + v.replace('"', '""') + '"'
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return str(v)


@op("validation")
def op_validation(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.worksheet.datavalidation import DataValidation

    ws, r1, c1, r2, c2 = ctx.area(o)
    rng = range_name(r1, c1, r2, c2)
    t = str(o.get("type", "list")).lower()
    t = {"integer": "whole", "number": "decimal", "text_length": "textLength", "length": "textLength"}.get(t, t)
    kw: dict[str, Any] = {"type": t, "allow_blank": bool(o.get("allow_blank", True))}
    if t == "list":
        if o.get("source"):
            kw["formula1"] = str(o["source"]).lstrip("=")
        else:
            vals = o.get("values")
            if not isinstance(vals, list) or not vals:
                raise SkillError("list validation needs 'values' or 'source' (a range like $H$1:$H$5)")
            joined = ",".join(str(v) for v in vals)
            if len(joined) > 255:
                raise SkillError("list values are longer than Excel's 255-character limit; put them in cells and use 'source'")
            kw["formula1"] = '"' + joined.replace('"', '""') + '"'
    elif t == "custom":
        kw["formula1"] = str(o.get("formula", "")).lstrip("=")
    elif t in ("whole", "decimal", "date", "time", "textLength"):
        oper = str(o.get("operator", "between"))
        oper = {">": "greaterThan", ">=": "greaterThanOrEqual", "<": "lessThan", "<=": "lessThanOrEqual", "=": "equal", "<>": "notEqual"}.get(oper, oper)
        kw["operator"] = oper
        lo = o.get("min", o.get("value"))
        hi = o.get("max")
        if t == "date":
            from _numfmt import date_to_serial

            conv = lambda v: str(int(date_to_serial(_dt.date.fromisoformat(v)))) if isinstance(v, str) and _ISO_DATE.match(v) else str(v)  # noqa: E731
        else:
            conv = lambda v: str(v).lstrip("=")  # noqa: E731
        if oper in ("between", "notBetween"):
            if lo is None or hi is None:
                raise SkillError("between needs 'min' and 'max'")
            kw["formula1"], kw["formula2"] = conv(lo), conv(hi)
        else:
            if lo is None:
                lo = hi
            kw["formula1"] = conv(lo)
    else:
        raise SkillError("validation type is list, whole, decimal, date, time, text_length or custom")
    dv = DataValidation(**kw)
    if o.get("prompt"):
        dv.prompt = str(o["prompt"])
        dv.promptTitle = str(o.get("prompt_title", ""))[:32] or None
        dv.showInputMessage = True
    if o.get("error"):
        dv.error = str(o["error"])
        dv.errorTitle = str(o.get("error_title", "Invalid value"))[:32]
    dv.showErrorMessage = True
    dv.errorStyle = o.get("error_style", "stop")
    ws.add_data_validation(dv)
    dv.add(rng)
    return f"added {t} validation on {ws.title}!{rng}"


@op("name")
def op_name(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.workbook.defined_name import DefinedName

    name = str(o.get("name", ""))
    ref = o.get("ref", o.get("refers_to", o.get("range")))
    if not re.fullmatch(r"[A-Za-z_\\][A-Za-z0-9_.]*", name) or re.fullmatch(r"[A-Za-z]{1,3}\d+", name) or name.upper() in ("R", "C"):
        raise SkillError(f"'{name}' is not a valid name (letters, digits, _ and ., not a cell address)")
    if ref is None:
        raise SkillError("name needs 'ref' (e.g. \"Inputs!$B$2\" or a formula)")
    ref = str(ref).lstrip("=")
    sheet, rest = split_sheet(ref)
    if sheet is None and re.fullmatch(r"\$?[A-Za-z]{1,3}\$?\d+(:\$?[A-Za-z]{1,3}\$?\d+)?", ref):
        ws = ctx.ws(o.get("sheet"))
        ref = f"{quote_sheet(ws.title)}!{_absolute(ref)}"
    elif sheet is not None and re.fullmatch(r"\$?[A-Za-z]{1,3}\$?\d+(:\$?[A-Za-z]{1,3}\$?\d+)?", rest):
        ref = f"{quote_sheet(sheet)}!{_absolute(rest)}"
    d = DefinedName(name, attr_text=add_prefixes("=" + ref)[1:])
    if o.get("scope"):
        ws = ctx.ws(o["scope"])
        ws.defined_names[name] = d
    else:
        if name in ctx.wb.defined_names:
            del ctx.wb.defined_names[name]
        ctx.wb.defined_names[name] = d
    ctx.formulas_changed = True
    return f"defined {name} = {d.attr_text}"


def _absolute(ref: str) -> str:
    return ":".join(re.sub(r"^\$?([A-Za-z]{1,3})\$?(\d+)$", r"$\1$\2", p) for p in ref.split(":"))


@op("delete_name")
def op_delete_name(ctx: Ctx, o: dict[str, Any]) -> str:
    name = str(o.get("name", ""))
    if name in ctx.wb.defined_names:
        del ctx.wb.defined_names[name]
        ctx.formulas_changed = True
        return f"deleted name {name}"
    for ws in ctx.wb.worksheets:
        if name in ws.defined_names:
            del ws.defined_names[name]
            ctx.formulas_changed = True
            return f"deleted name {name} ({ws.title})"
    raise SkillError(f"no name {name}")


_TOTAL_FUNCS = {"sum": 109, "average": 101, "avg": 101, "count": 103, "counta": 103, "max": 104, "min": 105, "stdev": 107, "var": 110}


@op("table")
def op_table(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

    ws, r1, c1, r2, c2 = ctx.area(o)
    existing = {t.lower() for w in ctx.wb.worksheets for t in dict.keys(w.tables)}
    name = str(o.get("name") or f"Table{len(existing) + 1}")
    name = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not re.match(r"[A-Za-z_]", name):
        name = "T_" + name
    while name.lower() in existing:
        name += "_2"
    headers = []
    seen: set[str] = set()
    for c in range(c1, c2 + 1):
        v = ws.cell(r1, c).value
        h = str(v).strip() if v not in (None, "") else f"Column{c - c1 + 1}"
        base, k = h, 2
        while h.lower() in seen:
            h = f"{base}{k}"
            k += 1
        seen.add(h.lower())
        if v in (None, "") or str(v) != h:
            ws.cell(r1, c).value = h
        headers.append(h)
    totals = o.get("totals")
    ref_r2 = r2
    cols = [TableColumn(id=i + 1, name=h) for i, h in enumerate(headers)]
    if totals:
        tr = r2 + 1
        if any(ws.cell(tr, c).value not in (None, "") for c in range(c1, c2 + 1)):
            raise SkillError(f"row {tr} below the table is not empty; cannot add a totals row")
        spec = totals if isinstance(totals, dict) else {headers[-1]: "sum"}
        label_done = False
        for i, h in enumerate(headers):
            fn = spec.get(h) or spec.get(h.lower())
            cell = ws.cell(tr, c1 + i)
            if fn:
                code = _TOTAL_FUNCS.get(str(fn).lower())
                if code is None:
                    raise SkillError(f"totals function {fn} (use sum, average, count, max, min, stdev, var)")
                colref = h.replace("'", "''").replace("[", "'[").replace("]", "']").replace("#", "'#")
                write_cell(ctx, ws, tr, c1 + i, f"=SUBTOTAL({code},{name}[{colref}])", ws.cell(r2, c1 + i).number_format if code not in (103,) else None)
                cols[i].totalsRowFunction = {101: "average", 103: "count", 104: "max", 105: "min", 107: "stdDev", 109: "sum", 110: "var"}[code]
            elif not label_done and i == 0:
                cell.value = "Total"
                cols[i].totalsRowLabel = "Total"
                label_done = True
            cell.font = make_font({"bold": True}, cell.font)
        ref_r2 = tr
    t = Table(displayName=name, ref=range_name(r1, c1, ref_r2, c2), tableColumns=cols)
    if totals:
        t.totalsRowCount = 1
        from openpyxl.worksheet.filters import AutoFilter

        t.autoFilter = AutoFilter(ref=range_name(r1, c1, r2, c2))
    t.tableStyleInfo = TableStyleInfo(name=o.get("style", "TableStyleMedium2"), showRowStripes=bool(o.get("banded_rows", True)), showColumnStripes=bool(o.get("banded_columns", False)), showFirstColumn=False, showLastColumn=False)
    ws.add_table(t)
    ctx.formulas_changed = True
    return f"added table {name} over {ws.title}!{t.ref}" + (" with a totals row" if totals else "")


# ── charts, comments, links, images, protection, print, properties ──────


def _ref_obj(ctx: Ctx, spec: str, default_ws: Any) -> tuple[Any, int, int, int, int]:
    sheet, rest = split_sheet(str(spec))
    ws = ctx.ws(sheet) if sheet else default_ws
    r1, c1, r2, c2 = parse_range(rest)
    return ws, r1, c1, r2, c2


def _column_range(ctx: Ctx, ws: Any, name: str, table_area: tuple[int, int, int, int] | None) -> str | None:
    """Resolves a header name to the column's range (header included) inside a data area."""
    if table_area is None:
        return None
    r1, c1, r2, c2 = table_area
    for c in range(c1, c2 + 1):
        v = ws.cell(r1, c).value
        if v is not None and str(v).strip().lower() == name.strip().lower():
            return range_name(r1, c, r2, c)
    return None


@op("chart")
def op_chart(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, RadarChart, Reference, ScatterChart, Series

    ws = ctx.ws(o.get("sheet"))
    kind = str(o.get("type", "column")).lower()
    data_area = None
    if o.get("data_area"):
        _, a1, b1, a2, b2 = _ref_obj(ctx, o["data_area"], ws)
        data_area = (a1, b1, a2, b2)
    values = o.get("values", o.get("data"))
    if values is None:
        raise SkillError("chart needs 'values' (a range with the series, header row first) and usually 'categories'")
    if isinstance(values, str):
        values = [values]
    val_refs = []
    for v in values:
        rng = _column_range(ctx, ws, v, data_area) if data_area and not re.search(r"\d", str(v).split("!")[-1]) else None
        val_refs.append(_ref_obj(ctx, rng or v, ws))
    cats = o.get("categories")
    cat_ref = None
    if cats:
        rng = _column_range(ctx, ws, cats, data_area) if data_area and not re.search(r"\d", str(cats).split("!")[-1]) else None
        cat_ref = _ref_obj(ctx, rng or cats, ws)
    titles_from_data = o.get("titles_from_data")
    if titles_from_data is None:
        first = val_refs[0]
        titles_from_data = isinstance(first[0].cell(first[1], first[2]).value, str)
    if kind in ("column", "bar"):
        ch = BarChart()
        ch.type = "col" if kind == "column" else "bar"
        if o.get("stacked"):
            ch.grouping = "percentStacked" if o.get("stacked") == "percent" else "stacked"
            ch.overlap = 100
    elif kind == "line":
        ch = LineChart()
        if o.get("stacked"):
            ch.grouping = "stacked"
    elif kind == "area":
        ch = AreaChart()
        if o.get("stacked"):
            ch.grouping = "stacked"
    elif kind == "pie":
        ch = PieChart()
    elif kind == "doughnut":
        ch = DoughnutChart()
    elif kind == "radar":
        ch = RadarChart()
    elif kind == "scatter":
        ch = ScatterChart()
        ch.style = 13
    else:
        raise SkillError("chart type is column, bar, line, area, pie, doughnut, radar or scatter")
    if kind == "scatter":
        if cat_ref is None:
            raise SkillError("a scatter chart needs 'categories' (the x values)")
        cws, a1, b1, a2, b2 = cat_ref
        xs = Reference(cws, min_col=b1, min_row=a1 + (1 if titles_from_data else 0), max_col=b2, max_row=a2)
        for vws, r1, c1, r2, c2 in val_refs:
            for c in range(c1, c2 + 1):
                ys = Reference(vws, min_col=c, min_row=r1, max_col=c, max_row=r2)
                s = Series(ys, xs, title_from_data=bool(titles_from_data))
                s.marker.symbol = o.get("marker", "circle")
                if not o.get("lines"):
                    s.graphicalProperties.line.noFill = True
                ch.series.append(s)
    else:
        for vws, r1, c1, r2, c2 in val_refs:
            ref = Reference(vws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
            ch.add_data(ref, titles_from_data=bool(titles_from_data))
        if cat_ref is not None:
            cws, a1, b1, a2, b2 = cat_ref
            ch.set_categories(Reference(cws, min_col=b1, min_row=a1 + (1 if titles_from_data and cws.cell(a1, b1).value is not None and isinstance(cws.cell(a1, b1).value, str) and a2 > a1 and a1 == val_refs[0][1] else 0), max_col=b2, max_row=a2))
    if o.get("title"):
        ch.title = str(o["title"])
    if hasattr(ch, "x_axis") and o.get("x_title"):
        ch.x_axis.title = str(o["x_title"])
    if hasattr(ch, "y_axis") and o.get("y_title"):
        ch.y_axis.title = str(o["y_title"])
    if hasattr(ch, "x_axis") and kind not in ("pie", "doughnut"):
        # openpyxl hides axes by default in recent versions; show them
        try:
            ch.x_axis.delete = False
            ch.y_axis.delete = False
        except AttributeError:
            pass
    legend = o.get("legend", "right" if len(val_refs) > 1 or (val_refs and val_refs[0][4] > val_refs[0][2]) or kind in ("pie", "doughnut") else "none")
    if legend in ("none", None, False):
        ch.legend = None
    else:
        ch.legend.position = {"right": "r", "left": "l", "top": "t", "bottom": "b"}.get(str(legend), "r")
    if o.get("data_labels"):
        from openpyxl.chart.label import DataLabelList

        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        if kind in ("pie", "doughnut") and o.get("data_labels") == "percent":
            ch.dataLabels.showVal = False
            ch.dataLabels.showPercent = True
    if o.get("style"):
        ch.style = int(o["style"])
    ch.width = float(o.get("width", 16))
    ch.height = float(o.get("height", 8))
    anchor = str(o.get("anchor", "")).upper()
    if not anchor:
        anchor = f"{col_letter(ws.max_column + 2)}2"
    ws.add_chart(ch, anchor)
    return f"added a {kind} chart at {ws.title}!{anchor}"


@op("comment")
def op_comment(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.comments import Comment

    ws, r1, c1, _, _ = ctx.area(o)
    text = o.get("text")
    cell = ws.cell(r1, c1)
    if text in (None, ""):
        cell.comment = None
        return f"removed the comment on {ws.title}!{cell.coordinate}"
    cm = Comment(str(text), str(o.get("author", "Desk")))
    cm.width = int(o.get("width", 220))
    cm.height = int(o.get("height", 90))
    cell.comment = cm
    return f"added a comment on {ws.title}!{cell.coordinate}"


@op("hyperlink")
def op_hyperlink(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.worksheet.hyperlink import Hyperlink

    ws, r1, c1, _, _ = ctx.area(o)
    cell = ws.cell(r1, c1)
    if o.get("url"):
        cell.hyperlink = str(o["url"])
    elif o.get("location"):
        loc = str(o["location"])
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=loc, display=str(o.get("text") or loc))
    else:
        cell.hyperlink = None
        return f"removed the link on {ws.title}!{cell.coordinate}"
    if o.get("tooltip") and cell.hyperlink is not None:
        cell.hyperlink.tooltip = str(o["tooltip"])
    if o.get("text") is not None:
        cell.value = str(o["text"])
    elif cell.value is None:
        cell.value = str(o.get("url") or o.get("location"))
    if o.get("style", True):
        cell.font = make_font({"color": "#0563C1", "underline": "single"}, cell.font)
    return f"linked {ws.title}!{cell.coordinate} to {o.get('url') or o.get('location')}"


@op("image")
def op_image(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.drawing.image import Image

    ws, r1, c1, _, _ = ctx.area(o)
    path = Path(str(o.get("path", ""))).expanduser()
    if not path.is_file():
        raise SkillError(f"image {path} does not exist")
    img = Image(str(path))
    w, h = img.width, img.height
    if o.get("width") and not o.get("height"):
        ratio = float(o["width"]) / w
        img.width, img.height = float(o["width"]), h * ratio
    elif o.get("height") and not o.get("width"):
        ratio = float(o["height"]) / h
        img.width, img.height = w * ratio, float(o["height"])
    elif o.get("width") and o.get("height"):
        img.width, img.height = float(o["width"]), float(o["height"])
    ws.add_image(img, f"{col_letter(c1)}{r1}")
    return f"placed {path.name} at {ws.title}!{col_letter(c1)}{r1}"


@op("protect")
def op_protect(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    ws.protection.sheet = True
    if o.get("password"):
        ws.protection.password = str(o["password"])
    for k in ("sort", "autoFilter", "formatCells", "formatColumns", "formatRows", "insertRows", "deleteRows"):
        if k in o.get("allow", []):
            setattr(ws.protection, k, False)
    if o.get("unlock"):
        _, r1, c1, r2, c2 = ctx.area({"range": o["unlock"], "sheet": ws.title, "op": "protect"})
        apply_style(ws, r1, c1, r2, c2, {"locked": False})
    return f"protected sheet {ws.title}" + (" with a password" if o.get("password") else "")


@op("unprotect")
def op_unprotect(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    ws.protection.sheet = False
    ws.protection.password = None
    return f"removed the protection of {ws.title}"


@op("print")
def op_print(ctx: Ctx, o: dict[str, Any]) -> str:
    ws = ctx.ws(o.get("sheet"))
    done = []
    if o.get("orientation"):
        ws.page_setup.orientation = o["orientation"]
        done.append(o["orientation"])
    if o.get("paper"):
        ws.page_setup.paperSize = {"letter": 1, "legal": 5, "a4": 9, "a3": 8}.get(str(o["paper"]).lower(), 9)
        done.append(str(o["paper"]))
    if "fit_width" in o or "fit_height" in o:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = int(o.get("fit_width", 1))
        ws.page_setup.fitToHeight = int(o.get("fit_height", 0))
        done.append("fit to page")
    if o.get("print_area"):
        ws.print_area = str(o["print_area"])
        done.append(f"area {o['print_area']}")
    if o.get("repeat_rows"):
        ws.print_title_rows = str(o["repeat_rows"])
        done.append(f"repeat rows {o['repeat_rows']}")
    if o.get("gridlines") is not None:
        ws.print_options.gridLines = bool(o["gridlines"])
    if o.get("header"):
        ws.oddHeader.center.text = str(o["header"])
    if o.get("footer"):
        ws.oddFooter.center.text = str(o["footer"])
    return f"page setup for {ws.title}: {', '.join(done) or 'updated'}"


@op("properties")
def op_properties(ctx: Ctx, o: dict[str, Any]) -> str:
    p = ctx.wb.properties
    done = []
    for k in ("title", "subject", "creator", "keywords", "description", "category", "lastModifiedBy"):
        if k in o:
            setattr(p, k, str(o[k]))
            done.append(k)
    if "author" in o:
        p.creator = str(o["author"])
        done.append("creator")
    return "set properties " + ", ".join(done)


@op("calc")
def op_calc(ctx: Ctx, o: dict[str, Any]) -> str:
    from openpyxl.workbook.properties import CalcProperties

    cp = ctx.wb.calculation or CalcProperties()
    if "iterate" in o:
        cp.iterate = bool(o["iterate"])
    if "iterate_count" in o:
        cp.iterateCount = int(o["iterate_count"])
    if "iterate_delta" in o:
        cp.iterateDelta = float(o["iterate_delta"])
    if "full_calc_on_load" in o:
        cp.fullCalcOnLoad = bool(o["full_calc_on_load"])
    if "mode" in o:
        cp.calcMode = str(o["mode"])
    ctx.wb.calculation = cp
    ctx.formulas_changed = True
    return "updated calculation settings"
