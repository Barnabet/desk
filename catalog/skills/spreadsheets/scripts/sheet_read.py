#!/usr/bin/env python3
"""Read cell values from any spreadsheet as Markdown, CSV or JSON, with formulas, display formats and styles on demand.

Big sheets (more than 5 000 data rows) get a map by default: the data block, header, column types, empty share,
ranges, first and last rows. Then find cells with --find/--grep (addresses and their rows) or read rows with
--rows 10001-10500. Reads of big sheets come from a typed Parquet copy cached by file content (built once).
"""

from __future__ import annotations

import json
import os
import shlex
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import DEFAULT_MAX_CHARS, UsageError, add_format, input_file, parser, run_main  # noqa: E402

EPILOG = """examples:
  python3 scripts/sheet_read.py sales.xlsx                          # first sheet (a map when it has over 5,000 rows)
  python3 scripts/sheet_read.py sales.xlsx --sheet Q3 --range A1:F40 --formulas
  python3 scripts/sheet_read.py big.xlsx --find "Zelda"             # every cell containing the text: Orders!F100001 …
  python3 scripts/sheet_read.py big.xlsx --grep "^INV-\\d{6}$"       # a regular expression over cell text
  python3 scripts/sheet_read.py big.xlsx --sheet Orders --rows 10001-10500 --format csv
  python3 scripts/sheet_read.py report.xlsx --range B2:E9 --display      # numbers as Excel shows them ($1,234.50, 12%)
  python3 scripts/sheet_read.py report.xlsx --range A1:D3 --styles       # fonts, fills, borders, number formats
  python3 scripts/sheet_read.py legacy.xls --all --limit 20

Output stops at --max-chars (default 60000) on a row boundary and ends with the exact command for the next part.
"""

MD_ROWS = 500  # default page for Markdown
AUTO_CSV_ROWS = 200  # more rows than this print as CSV by default (fewer tokens than a Markdown table)


def main() -> int:
    p = parser("Read a sheet (.xlsx .xlsm .xltx .xls .xlsb .ods .csv .tsv) as a table, a map, or search results.", EPILOG)
    p.add_argument("file")
    p.add_argument("--sheet", help="sheet name or 1-based number (default: the first visible sheet)")
    p.add_argument("--all", action="store_true", help="read every sheet")
    p.add_argument("--range", help="A1-style area, e.g. A1:F200 or 'Sheet 2'!B3:D9")
    p.add_argument("--rows", help="sheet rows to read, e.g. 10001-10500 (with the header row shown above them)")
    p.add_argument("--header", choices=["auto", "yes", "no"], default="auto", help="treat the first row as column names (default auto-detect)")
    p.add_argument("--offset", type=int, default=0, help="skip this many data rows (paging)")
    p.add_argument("--limit", type=int, default=None, help=f"rows to show (default {MD_ROWS} for Markdown, all that fit for csv/json)")
    p.add_argument("--map", action="store_true", help="print the map (columns, types, head and tail) even for a small sheet")
    p.add_argument("--find", help="find cells whose text contains this (case-insensitive); prints addresses and their rows")
    p.add_argument("--grep", help="like --find with a regular expression")
    p.add_argument("--case", action="store_true", help="--find/--grep are case-sensitive")
    p.add_argument("--formulas", action="store_true", help="also list formulas (xlsx family, .ods), grouping filled-down runs; with --find, search formula text too")
    p.add_argument("--display", action="store_true", help="show values formatted with each cell's number format (xlsx family)")
    p.add_argument("--styles", action="store_true", help="list fonts, fills, borders, alignment and number formats for the range (xlsx family)")
    p.add_argument("--no-coords", action="store_true", help="plain table without row numbers and column letters")
    p.add_argument("--max-chars", type=int, default=DEFAULT_MAX_CHARS, help="cap on the output; it ends with the command for the next part (default 60000)")
    p.add_argument("--no-cache", action="store_true", help="do not use or store the cached copy of big sheets")
    add_format(p, ("auto", "md", "csv", "tsv", "json"), default="auto")
    a = p.parse_args()
    if a.no_cache:
        os.environ["DESK_NO_CACHE"] = "1"
    path = input_file(a.file)
    from _book import Workbook, area_arg

    sheet_from_range, area = area_arg(a.range)
    sheet = sheet_from_range or a.sheet
    if a.offset < 0:
        raise UsageError("--offset must be 0 or more")
    rows_span = None
    if a.rows:
        rows_span = _rows_arg(a.rows)
        if area is not None:
            raise UsageError("give --rows or --range, not both")
    with Workbook(path) as wb:
        kind = wb.kind
        if kind in ("csv", "json"):
            targets: list[str | None] = [None]
        elif a.all or ((a.find or a.grep or a.map) and not sheet):
            targets = [m.name for m in wb.sheets() if m.kind == "worksheet"]
        else:
            targets = [wb.resolve_sheet(sheet).name]
        if a.find or a.grep:
            return do_find(wb, path, targets, a)
        sizes = {t: _size_hint(wb, path, t) for t in targets}
        others = [m.name for m in wb.sheets() if m.name not in targets] if kind not in ("csv", "json") else []
    explicit = bool(area or rows_span or a.offset or a.limit is not None or a.formulas or a.styles or a.display)
    big = [t for t in targets if sizes[t] > _big_rows()]
    if a.map or (big and not explicit):
        return do_map(path, targets, a, "--map" if a.map and not big else f"{'these sheets have' if len(big) > 1 else 'the sheet has'} over {_big_rows():,} rows: a map instead of a dump", others)
    results = []
    with Workbook(path) as wb:
        for t in targets:
            if sizes[t] > _big_rows() and not (a.formulas or a.styles or a.display) and (area is None or _store_covers(path, t, a, area)):
                results.append(read_big(path, t, area, rows_span, a))
            else:
                results.append(read_one(wb, path, t, area, rows_span, a, big=sizes[t] > _big_rows()))
    return emit_results(results, a, path)


def _big_rows() -> int:
    try:
        return int(os.environ.get("DESK_SHEET_BIG_ROWS", "5000"))
    except ValueError:
        return 5000


def _rows_arg(spec: str) -> tuple[int, int]:
    import re

    m = re.fullmatch(r"\s*(\d+)\s*(?:[-:]\s*(\d+)?)?\s*", spec)
    if not m:
        raise UsageError(f"--rows takes sheet row numbers like 101-600, got {spec!r}")
    a = int(m.group(1))
    b = int(m.group(2)) if m.group(2) else (a if "-" not in spec and ":" not in spec else 1_048_576)
    if a < 1 or b < a:
        raise UsageError(f"--rows {spec}: rows run from 1 upwards")
    return a, b


def _size_hint(wb: Any, path: Path, sheet: str | None) -> int:
    """About how many rows a sheet has, without reading it all where the format allows."""
    if wb.kind == "csv":
        size = path.stat().st_size
        if size < 256 << 10:
            return 0
        with open(path, "rb") as f:
            sample = f.read(1 << 20)
        lines = sample.count(b"\n")
        return int(size / max(len(sample), 1) * lines) if lines else 0
    if wb.kind == "json":
        return 0 if path.stat().st_size < 4 << 20 else 1 << 30
    if wb.kind == "xlsx":
        meta = wb.resolve_sheet(sheet)
        st = wb.scan(meta)
        if st and st.get("max_row"):
            from _scan import is_sparse

            if is_sparse(st):
                return 0
            return st["max_row"] - st["min_row"] + 1
        return 0
    if wb.kind in ("xls", "xlsb", "ods"):
        meta = wb.resolve_sheet(sheet)
        if meta.kind != "worksheet":
            return 0
        try:
            return int(wb._cal.get_sheet_by_name(meta.name).height)
        except Exception:  # noqa: BLE001
            return 0
    return 0


def _store_covers(path: Path, sheet: str | None, a: Any, area: tuple[int, int, int, int]) -> bool:
    from _store import covers, sheet_store

    meta = sheet_store(path, sheet, a.header)
    return covers(meta, area) or (meta.get("header_row") is not None and area[0] == meta["header_row"] and covers(meta, (meta["first_row"], area[1], area[2], area[3])))


def _cmd(path: Path, a: Any, sheet: str | None, extra: list[str]) -> str:
    """The command line that reads the next part."""
    parts = ["python3", "scripts/sheet_read.py", a.file]
    if sheet and not a.all:
        parts += ["--sheet", sheet]
    for flag in ("header",):
        v = getattr(a, flag)
        if v != "auto":
            parts += [f"--{flag}", v]
    for flag in ("formulas", "display", "no_coords"):
        if getattr(a, flag):
            parts.append("--" + flag.replace("_", "-"))
    if a.format not in ("auto",):
        parts += ["--format", a.format]
    parts += extra
    return " ".join(shlex.quote(str(x)) for x in parts)


# ── the map ─────────────────────────────────────────────────────────────


def do_map(path: Path, targets: list[str | None], a: Any, why: str, others: list[str] | None = None) -> int:
    from _store import render_map, sheet_store

    metas = [sheet_store(path, t, a.header) for t in targets]
    if a.format == "json":
        slim = [{k: v for k, v in m.items() if k not in ("parquet", "dir", "source", "version")} for m in metas]
        print(_cap_json({"file": str(path), "map": slim}, a.max_chars))
        return 0
    text = render_map(path, metas, why)
    if others:
        text += "\n\nOther sheets (not mapped here; --sheet NAME --map, or --map alone for all): " + ", ".join(others)
    first = metas[0]
    sh = first.get("sheet") if first.get("kind") not in ("csv", "json") else None
    nxt = ["", "Next:"]
    if first.get("rows"):
        r1 = first["first_row"]
        nxt.append(f"- rows: {_cmd(path, a, sh, ['--rows', f'{r1}-{r1 + 499}', '--format', 'csv'] if a.format == 'auto' else ['--rows', f'{r1}-{r1 + 499}'])}")
    nxt.append(f"- search: {_cmd(path, a, None, ['--find', 'TEXT'])}")
    table = _sql_name(path, first)
    nxt.append(f"- questions over all rows: python3 scripts/sheet_query.py {shlex.quote(a.file)} --sql \"SELECT … FROM {table} …\" (every table has a _row column: the sheet row)")
    print(_budget_text(text + "\n" + "\n".join(nxt), a.max_chars, "narrow with --sheet"))
    return 0


def _sql_name(path: Path, meta: dict[str, Any]) -> str:
    from _store import sanitize

    if meta.get("kind") in ("csv", "json"):
        return sanitize(path.stem)
    return sanitize(meta.get("sheet") or "sheet")


def _budget_text(text: str, max_chars: int, hint: str) -> str:
    if not max_chars or len(text) <= max_chars:
        return text
    cut = text.rfind("\n", 0, max_chars)
    cut = cut if cut > 0 else max_chars
    return text[:cut] + f"\n[… {len(text) - cut:,} more characters: {hint}]"


def _cap_json(obj: Any, max_chars: int) -> str:
    text = json.dumps(obj, ensure_ascii=False, default=str, separators=(",", ":"))
    if not max_chars or len(text) <= max_chars:
        return text
    return json.dumps({"error": f"the JSON is {len(text):,} characters, over --max-chars {max_chars:,}", "hint": "narrow with --sheet, --rows or --range, or raise --max-chars"}, indent=1)


# ── search ──────────────────────────────────────────────────────────────


def do_find(wb: Any, path: Path, targets: list[str | None], a: Any) -> int:
    import re

    from _book import text_value

    needle = a.grep if a.grep else a.find
    if not needle:
        raise UsageError("--find needs some text")
    rx = None
    if a.grep:
        try:
            rx = re.compile(a.grep, 0 if a.case else re.I)
        except re.error as e:
            raise UsageError(f"--grep: {e}") from e
    limit = a.limit or 200
    hits: list[dict[str, Any]] = []
    total = 0
    notes = []
    for t in targets:
        if _size_hint(wb, path, t) > _big_rows():
            h, n = _find_big(path, t, a, needle, limit - len(hits))
        else:
            h, n = _find_small(wb, path, t, a, needle, rx, limit - len(hits))
        hits.extend(h)
        total += n
        if len(hits) >= limit:
            if len(targets) > 1:
                notes.append("stopped at --limit; narrow with --sheet")
            break
    if a.format == "json":
        print(_cap_json({"query": needle, "regex": bool(a.grep), "matches": total, "shown": len(hits), "hits": hits}, a.max_chars))
        return 0
    what = f"/{needle}/" if a.grep else repr(needle)
    lines = [f"{total:,} cell{'s' if total != 1 else ''} match {what}" + (f"; showing {len(hits)}" if total > len(hits) else "") + ":"]
    for h in hits:
        ctx = " · ".join(f"{k}={v}" for k, v in h.get("context", []))
        where = f"{h['sheet']}!{h['cell']}" if h.get("sheet") else h["cell"]
        lines.append(f"- {where}" + (f" [{h['column']}]" if h.get("column") else "") + f": {h['value']}" + (f"  (row {h['row']}: {ctx})" if ctx else ""))
    if total > len(hits):
        lines.append(f"More: raise --limit (now {limit}), or narrow with --sheet; for counts and groups use sheet_query.py with WHERE … LIKE.")
    lines.extend(notes)
    if not hits:
        lines.append("No cell matches." + ("" if a.formulas else " (Formula text is not searched; add --formulas.)" if wb.kind == "xlsx" else ""))
    print(_budget_text("\n".join(lines), a.max_chars, "raise --max-chars or narrow the search"))
    del text_value
    return 0


def _context(names: list[str], letters: list[str], row: list[Any], j: int) -> list[tuple[str, str]]:
    from _store import _short

    keep = sorted({0, 1, j - 1, j, j + 1} & set(range(len(row))))
    out = []
    for k in keep:
        v = row[k]
        if v is None or v == "":
            continue
        label = f"{letters[k]} {names[k]}" if names[k] and names[k] != letters[k] else letters[k]
        out.append((label, _short(v, 40)))
    return out[:6]


def _find_big(path: Path, sheet: str | None, a: Any, needle: str, limit: int) -> tuple[list[dict[str, Any]], int]:
    from _store import connect, find, page, sheet_store

    meta = sheet_store(path, sheet, a.header)
    con = connect()
    try:
        hits, total = find(con, meta, needle, regex=bool(a.grep), limit=max(limit, 0), case=a.case)
        names = [c["name"] for c in meta["columns"]]
        letters = [c["letter"] for c in meta["columns"]]
        rows = {}
        for r in sorted({h["row"] for h in hits}):
            got = page(con, meta, r, r)
            if got:
                rows[r] = got[0][1]
        # the header row and rows above it are not in the Parquet copy
        extra = []
        low = needle if a.case else needle.lower()
        import re

        rx = re.compile(needle, 0 if a.case else re.I) if a.grep else None
        if meta.get("header_row"):
            for j, nm in enumerate(names):
                txt = nm if a.case else nm.lower()
                if (rx.search(nm) if rx else low in txt):
                    extra.append({"cell": f"{letters[j]}{meta['header_row']}", "row": meta["header_row"], "column": "", "value": nm, "context": [("header", "")]})
    finally:
        con.close()
    out = []
    for h in hits:
        row = rows.get(h["row"], [])
        out.append({"sheet": meta.get("sheet") if meta.get("kind") not in ("csv", "json") else None, "cell": h["cell"], "row": h["row"], "column": h["column"], "value": h["value"], "context": _context(names, letters, row, h["col"]) if row else []})
    for e in extra:
        e["sheet"] = meta.get("sheet") if meta.get("kind") not in ("csv", "json") else None
    return extra + out, total + len(extra)


def _find_small(wb: Any, path: Path, sheet: str | None, a: Any, needle: str, rx: Any, limit: int) -> tuple[list[dict[str, Any]], int]:
    from _a1 import col_letter
    from _book import text_value

    from _store import find_header_row

    g = wb.grid(sheet)
    low = needle if a.case else needle.lower()
    hits: list[dict[str, Any]] = []
    total = 0
    h = find_header_row(g.rows[:30]) if g.rows else None
    header_i = h if h is not None else -1
    header = g.rows[h] if h is not None else []
    names = [text_value(v) for v in header]
    letters = [col_letter(g.col1 + j) for j in range(g.width)]
    sheet_label = g.sheet if wb.kind not in ("csv", "json") else None

    def match(s: str) -> bool:
        return bool(rx.search(s)) if rx is not None else (low in (s if a.case else s.lower()))

    formulas: dict[tuple[int, int], str] = {}
    if a.formulas and wb.kind == "xlsx" and sheet:
        from _xlsx import load_book

        book, pkg, _ = load_book(path, sheets=[sheet])
        pkg.close()
        sh = book.sheet(sheet)
        if sh is not None:
            formulas = {k: "=" + fc.text.lstrip("=") for k, fc in sh.formulas.items()}
    for i, row in enumerate(g.rows):
        for j, v in enumerate(row):
            r, c = g.row1 + i, g.col1 + j
            f = formulas.get((r, c))
            if v is None and f is None:
                continue
            s = text_value(v) if v is not None else ""
            if match(s) or (f is not None and match(f)):
                total += 1
                if len(hits) < limit:
                    below = i > header_i
                    hits.append({"sheet": sheet_label, "cell": f"{col_letter(c)}{r}", "row": r, "column": names[j] if below and j < len(names) and names[j] else "", "value": s + (f"  ({f})" if f else ""), "context": _context((names if below else []) + [""] * (len(row) - (len(names) if below else 0)), letters, row, j) if i != header_i else [("header", "")]})
    for addr, v in g.outliers:
        s = text_value(v)
        if match(s):
            total += 1
            if len(hits) < limit:
                hits.append({"sheet": sheet_label, "cell": addr, "row": 0, "column": "", "value": s, "context": [("outside the data block", "")]})
    return hits, total


# ── reading rows ────────────────────────────────────────────────────────


def read_big(path: Path, sheet: str | None, area: Any, rows_span: Any, a: Any) -> dict[str, Any]:
    """A page of a big sheet from its cached Parquet copy."""
    from _a1 import MAX_COL, MAX_ROW
    from _book import json_value
    from _store import area_columns, connect, page, sheet_store

    meta = sheet_store(path, sheet, a.header)
    if not meta.get("rows"):
        return {"sheet": meta.get("sheet"), "range": "", "header": None, "first_row": 0, "rows": [], "total_rows": 0, "_rows": [], "_letters": [], "_numbers": []}
    r_first, r_last = meta["first_row"], meta["last_row"]
    cols = None
    if area is not None:
        r1, c1, r2, c2 = area
        cols = area_columns(meta, c1, c2 if c2 < MAX_COL else meta["last_col"])
        lo, hi = max(r1, r_first), min(r2 if r2 < MAX_ROW else r_last, r_last)
    elif rows_span is not None:
        lo, hi = max(rows_span[0], r_first), min(rows_span[1], r_last)
    else:
        lo, hi = r_first, r_last
    limit = a.limit if a.limit is not None else (MD_ROWS if a.format == "md" else 100_000)
    con = connect()
    try:
        # --offset counts data rows from the first one shown
        rc = meta["row_column"]
        from _store import ident, sql_str

        nums = [r[0] for r in con.execute(f"SELECT {ident(rc)} FROM read_parquet({sql_str(meta['parquet'])}) WHERE {ident(rc)} BETWEEN ? AND ? ORDER BY 1 LIMIT ? OFFSET ?", [lo, hi, limit, a.offset]).fetchall()]
        got = page(con, meta, nums[0], nums[-1], cols) if nums else []
        total = con.execute(f"SELECT count(*) FROM read_parquet({sql_str(meta['parquet'])}) WHERE {ident(rc)} BETWEEN ? AND ?", [lo, hi]).fetchone()[0]
    finally:
        con.close()
    names = [c["name"] for c in meta["columns"]]
    letters = [c["letter"] for c in meta["columns"]]
    if cols is not None:
        names = [names[j] for j in cols]
        letters = [letters[j] for j in cols]
    header = names if meta.get("header_row") else None
    out = {
        "sheet": meta.get("sheet") if meta.get("kind") not in ("csv", "json") else Path(path).stem,
        "header": header, "header_row": meta.get("header_row"),
        "first_row": got[0][0] if got else lo,
        "rows": [[json_value(v) for v in r] for _, r in got],
        "row_numbers": [n for n, _ in got],
        "total_rows": total,
        "source": "cached copy",
        "_letters": letters, "_sheet_arg": meta.get("sheet") if meta.get("kind") not in ("csv", "json") else None,
    }
    shown_last = got[-1][0] if got else lo - 1
    if total > a.offset + len(got):
        out["next_rows"] = f"{shown_last + 1}-{min(hi, shown_last + max(len(got), 1))}"
    if meta.get("excluded"):
        out["excluded"] = meta["excluded"]
    return out


def read_one(wb: Any, path: Path, sheet: str | None, area: Any, rows_span: Any, a: Any, big: bool = False) -> dict[str, Any]:
    """A block of one sheet. big: a sheet over the map threshold read in full detail (--formulas, --styles,
    --display): without --range or --rows it is read a page at a time (the first rows are streamed from the XML)."""
    from _a1 import MAX_COL, col_letter
    from _book import detect_header, json_value, text_value, trim_grid

    whole_big = big and area is None and rows_span is None
    if rows_span is not None:
        area = (rows_span[0], 1, rows_span[1], MAX_COL)
    limit = a.limit if a.limit is not None else (MD_ROWS if a.format == "md" or whole_big else None)
    fetch = None
    if limit is not None:
        fetch = a.offset + limit + 2
    fsheet = None
    if a.formulas and wb.kind in ("xlsx", "ods"):
        # every formula cell comes from the file's XML: a formula stored without a result has no value, so the
        # value grid alone would leave it out
        fsheet = _formula_sheet(path, wb.resolve_sheet(sheet).name, wb.kind)
    g = wb.grid(sheet, area, max_rows=fetch)
    loaded_all = fetch is None or len(g.rows) < fetch
    rows = g.rows if area is not None and rows_span is None else trim_grid(g.rows)
    if rows_span is not None and rows:
        # keep the columns of the used range only
        width = max((len(r) for r in rows), default=0)
        rows = [r[:width] for r in rows]
    if fsheet is not None and fsheet.formulas:
        want = area or ((rows_span[0], 1, rows_span[1], MAX_COL) if rows_span is not None else None)
        rows = _cover_formulas(g, rows, fsheet.formulas, want, loaded_all, fetch)
    header = False
    head_row = None
    header_row_no = None
    if rows_span is not None:
        # the header row shown above a --rows page
        if a.header != "no" and rows_span[0] > 1:
            hg = wb.grid(sheet, (1, g.col1, min(30, rows_span[0] - 1), g.col1 + max((len(r) for r in rows), default=1) - 1))
            from _store import find_header_row

            h = find_header_row(hg.rows, a.header)
            if h is not None:
                head_row = hg.rows[h]
                header_row_no = hg.row1 + h
                header = True
        body = rows
        first_data_row = g.row1
    else:
        header = a.header == "yes" or (a.header == "auto" and detect_header(rows))
        head_row = rows[0] if header and rows else None
        header_row_no = g.row1 if header and rows else None
        body = rows[1:] if header else rows
        first_data_row = g.row1 + (1 if header else 0)
    start = min(a.offset, len(body))
    shown = body[start : start + limit] if limit is not None else body[start:]
    total_body = (g.total_rows - (1 if header and rows_span is None else 0)) if area is None and g.total_rows else len(body)
    if area is not None:
        total_body = len(body)
    display_note = None
    if a.display and wb.kind == "xlsx":
        shown, head_row, display_note = _display(path, g.sheet, shown, head_row, first_data_row + start, g.col1)
    width = max((len(r) for r in rows), default=0)
    out: dict[str, Any] = {
        "sheet": g.sheet,
        "header": [text_value(v) for v in head_row] if head_row else None,
        "header_row": header_row_no,
        "first_row": first_data_row + start,
        "rows": [[json_value(v) for v in r] for r in shown],
        "row_numbers": [first_data_row + start + i for i in range(len(shown))],
        "total_rows": total_body,
        "_letters": [col_letter(g.col1 + j) for j in range(width)],
        "_sheet_arg": g.sheet if wb.kind not in ("csv", "json") else None,
    }
    if start + len(shown) < total_body:
        out["next_offset"] = start + len(shown)
        if whole_big:
            nxt = first_data_row + start + len(shown)  # continue with --rows: found without reading the rows above
            out["next_rows"] = f"{nxt}-{nxt + (limit or MD_ROWS) - 1}"
    if g.merged:
        out["merged"] = g.merged
    if g.outlier_count:
        from _store import _short

        out["outliers"] = {"count": g.outlier_count, "used_range": g.extent, "cells": [[addr, _short(v)] for addr, v in g.outliers[:20]]}
    if a.formulas:
        if wb.kind in ("xlsx", "ods"):
            # every formula of the requested range (the whole sheet by default) is listed on exactly one page: the
            # rows of this page, from the top of the range on the first page, to its bottom on the last
            from _a1 import MAX_ROW

            want = area or ((rows_span[0], 1, rows_span[1], MAX_COL) if rows_span is not None else (1, 1, MAX_ROW, MAX_COL))
            top = want[0] if start == 0 else first_data_row + start
            bottom = first_data_row + start + len(shown) - 1 if "next_offset" in out else want[2]
            out["formulas"] = read_formulas(path, g.sheet, (top, want[1], max(top, bottom), want[3]), wb.kind, g, fsheet)
            missing = sum(f.get("without_value", 0) for f in out["formulas"])
            if missing:
                out["formulas_without_values"] = missing
                out["formulas_note"] = _no_values_note(a.file, wb.kind, missing)
        else:
            out["formulas_note"] = f"formulas are only readable from the xlsx family and .ods; convert {path.suffix} to .xlsx with LibreOffice (sheet_convert.py) to see them"
    if a.styles:
        if wb.kind == "xlsx":
            r_first = first_data_row + start - (1 if header and rows_span is None else 0)
            r_last = first_data_row + start + len(shown) - 1
            out["styles"] = read_styles(path, g.sheet, (r_first, g.col1, r_last, g.col1 + max(width, 1) - 1))
        else:
            out["styles_note"] = "styles are only readable from the xlsx family"
    if display_note:
        out["display_note"] = display_note
    return out


# ── output with budgets ─────────────────────────────────────────────────


def emit_results(results: list[dict[str, Any]], a: Any, path: Path) -> int:
    from _book import csv_text, text_value
    from _common import md_escape_cell

    fmt = a.format
    if fmt == "auto":
        n = sum(len(r["rows"]) for r in results)
        fmt = "csv" if n > AUTO_CSV_ROWS else "md"
    budget = a.max_chars or 0
    if fmt == "json":
        payload = []
        for r in results:
            d = {k: v for k, v in r.items() if not k.startswith("_")}
            payload.append(d)
        def dump(obj: Any) -> str:
            return json.dumps(obj, ensure_ascii=False, default=str, separators=(",", ":"))

        text = dump(payload if a.all else payload[0])
        if budget and len(text) > budget:
            # keep whole rows: drop rows from the end until it fits, then say how to continue
            r = payload[-1]
            lo, hi = 0, len(r["rows"])
            full = r["rows"]
            nums = r.get("row_numbers") or []
            while lo < hi:
                mid = (lo + hi + 1) // 2
                r["rows"], r["row_numbers"] = full[:mid], nums[:mid]
                if len(dump(payload if a.all else payload[0])) + 300 <= budget:
                    lo = mid
                else:
                    hi = mid - 1
            r["rows"], r["row_numbers"] = full[:lo], nums[:lo]
            if lo < len(full):
                src = results[-1]
                nxt = _next_args(src, lo, a)
                r["truncated"] = f"{len(full) - lo:,} more rows did not fit in --max-chars {budget:,}"
                r["next"] = _cmd(path, a, src.get("_sheet_arg"), nxt)
            text = dump(payload if a.all else payload[0])
        print(text)
        return 0
    out_lines: list[str] = []
    used = 0
    for k, r in enumerate(results):
        letters = r.get("_letters") or []
        nums = r.get("row_numbers") or []
        rows = r["rows"]
        head = r.get("header")
        block: list[str] = []
        if fmt == "md":
            title = f"## {r['sheet']}"
            if rows and nums and letters:
                top = r.get("header_row") if head and r.get("header_row") else nums[0]
                if top < nums[0] - 1:
                    title += f" ({letters[0]}{nums[0]}:{letters[-1]}{nums[-1]}; header in row {top})"
                else:
                    title += f" ({letters[0]}{min(top, nums[0])}:{letters[-1]}{nums[-1]})"
            block.append(title)
            block.append("")
            if not rows and not head:
                block.append("_(empty sheet)_")
        elif a.all:
            block.append(f"# {r['sheet']}")
        width = max([len(letters)] + [len(x) for x in rows[:1]]) if (rows or letters) else 0
        if fmt == "md":
            coords = not a.no_coords
            names = [text_value(v) for v in (head or [])] + [""] * max(0, width - len(head or []))
            heads = [f"{letters[j]}: {names[j]}" if coords and j < len(letters) and names[j] else (letters[j] if coords and j < len(letters) else names[j]) for j in range(width)]
            if coords:
                heads = ["row"] + heads
            table_head = ["| " + " | ".join(md_escape_cell(h) for h in heads) + " |", "|" + "|".join("---" for _ in heads) + "|"] if (rows or head) else []
            block.extend(table_head)
            row_lines = []
            for i, row in enumerate(rows):
                cells = [md_escape_cell(text_value(v)) for v in row] + [""] * (width - len(row))
                if coords:
                    cells = [str(nums[i] if i < len(nums) else "")] + cells
                row_lines.append("| " + " | ".join(cells) + " |")
        else:
            delim = "\t" if fmt == "tsv" else ","
            coords = not a.no_coords
            hdr = (["row"] if coords else []) + list(head) if head else None
            if hdr is None and coords:
                hdr = ["row"] + letters[:width]
            if hdr:
                block.append(csv_text([hdr], delim).rstrip("\n"))
            row_lines = [csv_text([([nums[i]] if coords and i < len(nums) else []) + list(row)], delim).rstrip("\n") for i, row in enumerate(rows)]
        # notes after the rows
        notes = []
        if r.get("excluded"):
            notes.append("Left out (not data): " + "; ".join(f"row {e['row']} ({e['why']})" for e in r["excluded"][:5]))
        if r.get("outliers"):
            o = r["outliers"]
            notes.append(f"Used range {o['used_range']}: {o['count']} cell(s) sit far outside this block: " + ", ".join(f"{c} = {v}" for c, v in o["cells"][:10]) + (" …" if o["count"] > 10 else "") + " (read them with --range)")
        if r.get("merged") and fmt == "md":
            notes.append("Merged: " + ", ".join(r["merged"][:30]) + (" …" if len(r["merged"]) > 30 else ""))
        if r.get("display_note") and fmt == "md":
            notes.append(r["display_note"])
        if r.get("formulas") is not None:
            notes.append(render_formulas(r["formulas"]))
        if r.get("formulas_note"):
            notes.append(r["formulas_note"])
        if r.get("styles") is not None:
            notes.append(render_styles(r["styles"]))
        if r.get("styles_note"):
            notes.append(r["styles_note"])
        tail_len = sum(len(x) + 2 for x in notes) + 220
        head_text = "\n".join(block)
        room = (budget - used - len(head_text) - tail_len) if budget else None
        kept = len(row_lines)
        if room is not None:
            acc = 0
            for i, ln in enumerate(row_lines):
                acc += len(ln) + 1
                if acc > room:
                    kept = max(i, 1)  # always show at least one row, so paging moves forward
                    break
        body = "\n".join(row_lines[:kept])
        more_msg = None
        if kept < len(row_lines):
            nxt = _next_args(r, kept, a)
            more_msg = f"[… {len(row_lines) - kept:,} more rows here did not fit in --max-chars {budget:,}. Next: {_cmd(path, a, r.get('_sheet_arg'), nxt)}]"
        else:
            total = r.get("total_rows") or 0
            shown_n = len(rows)
            if r.get("next_offset") is not None or r.get("next_rows"):
                nxt = _next_args(r, shown_n, a)
                more_msg = f"Showing {shown_n:,} of {total:,} rows. Next: {_cmd(path, a, r.get('_sheet_arg'), nxt)}"
        section = head_text + ("\n" + body if body else "")
        if more_msg:
            section += "\n" + more_msg
        if notes:
            section += "\n\n" + "\n".join(notes) if fmt == "md" else "\n" + "\n".join("# " + n.replace("\n", "\n# ") for n in notes)
        out_lines.append(section)
        used += len(section) + 2
        if budget and used >= budget and k < len(results) - 1:
            rest = [x["sheet"] for x in results[k + 1 :]]
            out_lines.append(f"[… {len(rest)} more sheet(s) did not fit: {', '.join(rest)}; read them with --sheet]")
            break
    print(("\n\n" if fmt == "md" else "\n").join(out_lines))
    return 0


def _next_args(r: dict[str, Any], kept: int, a: Any) -> list[str]:
    """Arguments that continue after the first `kept` rows of a result."""
    nums = r.get("row_numbers") or []
    if a.rows or r.get("next_rows") is not None or (nums and "next_offset" not in r and a.range is None):
        if kept < len(nums):
            start = nums[kept]
        else:
            start = (nums[-1] + 1) if nums else 1
        page_len = max(kept, 50) if kept else 500
        return ["--rows", f"{start}-{start + page_len - 1}"]
    if a.range:
        from _a1 import MAX_COL, MAX_ROW, col_letter, parse_range, quote_sheet, split_sheet

        sh, rest = split_sheet(a.range)
        r1, c1, r2, c2 = parse_range(rest)
        start = nums[kept] if kept < len(nums) else (nums[-1] + 1 if nums else r1)
        end = "" if r2 >= MAX_ROW else str(r2)
        rng = f"{col_letter(c1)}{start}:{col_letter(c2) if c2 < MAX_COL else 'XFD'}{end or MAX_ROW}"
        return ["--range", (quote_sheet(sh) + "!" if sh else "") + rng]
    off = a.offset + kept
    extra = ["--offset", str(off)]
    if a.limit is not None:
        extra += ["--limit", str(a.limit)]
    return extra


def _formula_sheet(path: Path, sheet: str, kind: str) -> Any:
    """The formula cells of one sheet, read from the file itself (an engine SheetData: .formulas maps (row, col) to
    a FormulaCell with its text and cached value). .ods formulas are translated to Excel syntax and get their
    values from the grid later (read_formulas)."""
    if kind == "ods":
        import _ods
        from _formula import Book

        book = Book()
        bsh = book.add_sheet(sheet)
        for (r, c), f in _ods.read_formulas(path).get(sheet, {}).items():
            book.set_formula(bsh, r, c, f.lstrip("="))
        return bsh
    from _xlsx import load_book

    book, pkg, _extra = load_book(path, sheets=[sheet])
    pkg.close()
    return book.sheet(sheet)


def _cover_formulas(g: Any, rows: list[list[Any]], formulas: dict[tuple[int, int], Any], want: tuple[int, int, int, int] | None, loaded_all: bool, fetch: int | None) -> list[list[Any]]:
    """Widens the value grid so the table shows the formula cells of the requested area (or sheet) too: a formula
    stored without a result is blank to the value reader, which trims it away. Updates g (row1, col1, total_rows,
    rows). Formula cells far from the data (a sparse sheet) leave the table as it is; they are still listed."""
    from _a1 import MAX_COL, MAX_ROW

    w1, v1, w2, v2 = want or (1, 1, MAX_ROW, MAX_COL)
    keys = [k for k in formulas if w1 <= k[0] <= w2 and v1 <= k[1] <= v2]
    if not keys:
        return rows
    fr1, fr2 = min(k[0] for k in keys), max(k[0] for k in keys)
    fc1, fc2 = min(k[1] for k in keys), max(k[1] for k in keys)
    width = max((len(r) for r in rows), default=0)
    row1, col1 = (g.row1, g.col1) if rows else (fr1, fc1)  # no values at all: start at the first formula
    n1, m1 = min(row1, fr1), min(col1, fc1)
    m2 = max(col1 + width - 1, fc2)
    last = row1 + len(rows) - 1
    n2 = max(last, fr2) if loaded_all else last
    if fetch is not None:
        n2 = min(n2, max(last, n1 + fetch - 1))
    if (n2 - n1 + 1) * (m2 - m1 + 1) > max(2 * len(rows) * width, 10_000):
        return rows
    out = []
    for r in range(n1, n2 + 1):
        i = r - row1
        src = rows[i] if 0 <= i < len(rows) else []
        out.append([src[c - col1] if 0 <= c - col1 < len(src) else None for c in range(m1, m2 + 1)])
    if want is None:
        extent_last = row1 + (g.total_rows or len(rows)) - 1
        g.total_rows = max(extent_last, fr2) - n1 + 1
    g.rows, g.row1, g.col1 = out, n1, m1
    return out


def _no_values_note(file: str, kind: str, missing: int) -> str:
    """What to do about formula cells stored without a result (files written by libraries have none)."""
    what = f"{missing} formula cell(s) here have no stored result (the file was saved without calculated values, as libraries like openpyxl do), so their values show blank."
    if kind == "xlsx":
        p = Path(file)
        return f"{what} Compute and store them: python3 scripts/sheet_recalc.py {shlex.quote(file)} --out {shlex.quote(p.stem + '-calc' + p.suffix)}, then read that file."
    return f"{what} Convert the file to .xlsx (sheet_convert.py) and run sheet_recalc.py on it to compute them."


def read_formulas(path: Path, sheet: str, area: tuple[int, int, int, int], kind: str = "xlsx", grid: Any = None, fsheet: Any = None) -> list[dict[str, Any]]:
    """Formulas in an area, with filled-down (or across) runs grouped: [{'range', 'formula', 'value'}]. Entries
    whose cells have no stored result carry 'without_value' (how many cells).

    .ods formulas are translated from OpenFormula to Excel syntax; their values come from the grid."""
    from _a1 import col_letter
    from _formula import FormulaSyntaxError, strip_prefixes, translate

    sh = fsheet if fsheet is not None else _formula_sheet(path, sheet, kind)
    if sh is None:
        return []
    if kind == "ods" and grid is not None:
        for (r, c), fc in sh.formulas.items():
            i, j = r - grid.row1, c - grid.col1
            if 0 <= i < len(grid.rows) and 0 <= j < len(grid.rows[i]):
                fc.cached = grid.rows[i][j]
    r1, c1, r2, c2 = area
    cells = sorted((k, fc) for k, fc in sh.formulas.items() if r1 <= k[0] <= r2 and c1 <= k[1] <= c2)
    by_col: dict[int, list[tuple[int, Any]]] = {}
    for (r, c), fc in cells:
        by_col.setdefault(c, []).append((r, fc))
    out: list[dict[str, Any]] = []
    for c in sorted(by_col):
        run: list[tuple[int, Any]] = []

        def flush() -> None:
            if not run:
                return
            first_r, first = run[0]
            last_r = run[-1][0]
            text = "=" + strip_prefixes(first.text.lstrip("="))
            entry: dict[str, Any] = {"range": f"{col_letter(c)}{first_r}" + (f":{col_letter(c)}{last_r}" if last_r != first_r else ""), "formula": text}
            if first.array_ref:
                a1, b1, a2, b2 = first.array_ref
                entry["array"] = f"{col_letter(b1)}{a1}:{col_letter(b2)}{a2}"
            if first.dynamic:
                entry["dynamic_array"] = True
            from _book import json_value

            if last_r == first_r:
                entry["value"] = json_value(_cached(first))
            else:
                entry["filled"] = last_r - first_r + 1
            empty = sum(1 for _r, fc in run if fc.cached is None)
            if empty:
                entry["without_value"] = empty
            out.append(entry)

        prev_r = None
        for r, fc in by_col[c]:
            if run and prev_r is not None and r == prev_r + 1:
                base_r, base = run[0]
                try:
                    same = translate("=" + base.text.lstrip("="), r - base_r, 0) == "=" + fc.text.lstrip("=")
                except FormulaSyntaxError:
                    same = False
                if same and not fc.array_ref and not base.array_ref:
                    run.append((r, fc))
                    prev_r = r
                    continue
            flush()
            run = [(r, fc)]
            prev_r = r
        flush()
    out.sort(key=lambda e: _sort_key(e["range"]))
    return out


def _cached(fc: Any) -> Any:
    v = fc.cached
    if hasattr(v, "code"):
        return v.code
    return v


def _sort_key(ref: str) -> tuple[int, int]:
    from _a1 import parse_range

    r1, c1, _, _ = parse_range(ref)
    return (r1, c1)


def render_formulas(formulas: list[dict[str, Any]]) -> str:
    if not formulas:
        return "No formulas in this range."
    lines = ["Formulas:"]
    for f in formulas[:400]:
        extra = ""
        missing = f.get("without_value", 0)
        if f.get("filled"):
            none = "; no stored results" if missing == f["filled"] else (f"; {missing} without a stored result" if missing else "")
            extra = f"  (filled down {f['filled']} rows; relative references shift{none})"
        elif missing:
            extra = "  → (no stored result)"
        elif "value" in f:
            v = f["value"]
            extra = f"  → {v!r}" if not isinstance(v, str) else f"  → {v}"
        if f.get("array"):
            extra += f"  [array over {f['array']}]"
        if f.get("dynamic_array"):
            extra += "  [dynamic array]"
        lines.append(f"- {f['range']}: {f['formula']}{extra}")
    if len(formulas) > 400:
        lines.append(f"- … {len(formulas) - 400} more; narrow with --range")
    return "\n".join(lines)


def _display(path: Path, sheet: str, shown: list[list[Any]], head_row: list[Any] | None, first_row: int, col1: int) -> tuple[list[list[Any]], list[Any] | None, str | None]:
    """Values formatted with their cell number formats (from the file's styles)."""
    from _numfmt import format_value
    from _xlsx import Package, area_cells, stream_area

    with Package(path) as pkg:
        si = pkg.sheet(sheet)
        last_row = first_row + len(shown) - 1
        width = max((len(r) for r in shown), default=0)
        fmts: dict[tuple[int, int], str] = {}
        if si.path and pkg.has(si.path):
            data = stream_area(pkg, si.path, first_row, last_row)[1]
            for r, c, style, _a, _i in area_cells(data, first_row, last_row, col1, col1 + max(0, width - 1)):
                if style:
                    fmts[(r, c)] = pkg.number_format(style)
        date1904 = pkg.date1904
    out = []
    for i, row in enumerate(shown):
        new = []
        for j, v in enumerate(row):
            code = fmts.get((first_row + i, col1 + j), "General")
            if v is None:
                new.append(None)
            elif hasattr(v, "code"):
                new.append(v)
            else:
                new.append(format_value(v, code, date1904)[0])
        out.append(new)
    return out, head_row, "Values shown with their number formats (--display)."


def read_styles(path: Path, sheet: str, area: tuple[int, int, int, int]) -> list[dict[str, Any]]:
    import openpyxl

    from _a1 import col_letter

    r1, c1, r2, c2 = area
    if (r2 - r1 + 1) * (c2 - c1 + 1) > 2000:
        r2 = r1 + max(0, 2000 // max(1, c2 - c1 + 1) - 1)
    if path.stat().st_size > (2 << 20):
        return read_styles_xml(path, sheet, (r1, c1, r2, c2))
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        ws = wb[sheet]
        out = []
        seen: dict[str, list[str]] = {}
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(r, c)
                desc = describe_style(cell)
                if not desc:
                    continue
                key = repr(sorted(desc.items()))
                seen.setdefault(key, []).append(f"{col_letter(c)}{r}")
                if len(seen[key]) == 1:
                    out.append({"cells": seen[key], **desc})
        for e in out:
            e["cells"] = _compress_cells(e["cells"])
        widths = {}
        for c in range(c1, c2 + 1):
            d = ws.column_dimensions.get(col_letter(c))
            if d is not None and d.width:
                widths[col_letter(c)] = round(d.width, 2)
        if widths:
            out.append({"column_widths": widths})
        heights = {}
        for r in range(r1, r2 + 1):
            d = ws.row_dimensions.get(r)
            if d is not None and d.height:
                heights[str(r)] = d.height
        if heights:
            out.append({"row_heights": heights})
        return out
    finally:
        wb.close()


def read_styles_xml(path: Path, sheet: str, area: tuple[int, int, int, int]) -> list[dict[str, Any]]:
    """read_styles for big workbooks: the cells' style indexes straight from the sheet XML, described through
    styles.xml, without loading the workbook."""
    import re
    from types import SimpleNamespace

    from openpyxl.styles.alignment import Alignment
    from openpyxl.styles.protection import Protection

    from _a1 import col_letter
    from _patch import StyleBook
    from _xlsx import Package, area_cells, row_attrs_in, stream_area

    r1, c1, r2, c2 = area
    with Package(path) as pkg:
        si = pkg.sheet(sheet)
        if not si.path or not pkg.has(si.path):
            return []
        # the sheet part is streamed: only the requested rows (and the columns' widths before them) are kept
        head, data = stream_area(pkg, si.path, r1, r2)
        styles_path = next((t for (typ, t) in pkg._rels.values() if typ.endswith("/styles")), None)
        sb = StyleBook(pkg.read(styles_path)) if styles_path and pkg.has(styles_path) else None
        cells = [(r, c, s) for r, c, s, _a, _i in area_cells(data, r1, r2, c1, c2)]
        rows = row_attrs_in(data, r1, r2)
    memo: dict[int, dict[str, Any]] = {}

    def desc(xf: int) -> dict[str, Any]:
        if xf not in memo:
            if sb is None:
                memo[xf] = {}
            else:
                attrs, inner = sb.xf_attrs(xf)
                al = re.search(rb"<(?:\w+:)?alignment\b[^>]*/>", inner)
                pr = re.search(rb"<(?:\w+:)?protection\b[^>]*/>", inner)
                cell = SimpleNamespace(
                    font=sb.font(int(attrs.get("fontId", "0") or 0)),
                    fill=sb.fill(int(attrs.get("fillId", "0") or 0)),
                    border=sb.border(int(attrs.get("borderId", "0") or 0)),
                    alignment=sb._obj(Alignment, al.group(0)) if al else None,
                    protection=sb._obj(Protection, pr.group(0)) if pr else None,
                    number_format=sb.number_format(xf),
                )
                memo[xf] = describe_style(cell)
        return memo[xf]

    out: list[dict[str, Any]] = []
    seen: dict[str, list[str]] = {}
    for r, c, xf in cells:
        d = desc(xf)
        if not d:
            continue
        key = repr(sorted(d.items()))
        seen.setdefault(key, []).append(f"{col_letter(c)}{r}")
        if len(seen[key]) == 1:
            out.append({"cells": seen[key], **d})
    for e in out:
        e["cells"] = _compress_cells(e["cells"])
    widths = {}
    for m in re.finditer(rb"<(?:\w+:)?col\b([^>]*)/?>", head):
        a = dict(re.findall(rb'\b(\w+)="([^"]*)"', m.group(1)))
        try:
            lo, hi, w = int(a.get(b"min", b"0")), int(a.get(b"max", b"0")), float(a.get(b"width", b"0"))
        except ValueError:
            continue
        for c in range(max(lo, c1), min(hi, c2) + 1):
            if w:
                widths[col_letter(c)] = round(w, 2)
    if widths:
        out.append({"column_widths": widths})
    heights = {}
    for r, attrs in rows.items():
        m = re.search(rb'\bht="([\d.]+)"', attrs)
        if m and b'customHeight="1"' in attrs:
            heights[str(r)] = float(m.group(1))
    if heights:
        out.append({"row_heights": heights})
    return out


def _compress_cells(cells: list[str]) -> str:
    if len(cells) <= 6:
        return ", ".join(cells)
    return ", ".join(cells[:5]) + f" … ({len(cells)} cells)"


def _color(c: Any) -> str | None:
    if c is None:
        return None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str):
            v = c.rgb
            return "#" + (v[2:] if len(v) == 8 else v)
        if c.type == "theme":
            return f"theme {c.theme}" + (f" tint {c.tint:+.2f}" if c.tint else "")
        if c.type == "indexed":
            return f"indexed {c.indexed}"
    except Exception:  # noqa: BLE001
        return None
    return None


def describe_style(cell: Any) -> dict[str, Any]:
    d: dict[str, Any] = {}
    f = cell.font
    font = {}
    if f is not None:
        if f.b:
            font["bold"] = True
        if f.i:
            font["italic"] = True
        if f.u:
            font["underline"] = f.u
        if f.strike:
            font["strike"] = True
        if f.sz and float(f.sz) != 11:
            font["size"] = float(f.sz)
        if f.name and f.name not in ("Calibri",):
            font["name"] = f.name
        col = _color(f.color)
        if col and col not in ("#000000", "theme 1"):
            font["color"] = col
    if font:
        d["font"] = font
    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) not in (None, "none"):
        col = _color(fill.fgColor) or _color(fill.start_color)
        d["fill"] = col if fill.fill_type == "solid" else f"{fill.fill_type} {col}"
    b = cell.border
    if b is not None:
        sides = {}
        for side in ("left", "right", "top", "bottom"):
            s = getattr(b, side)
            if s is not None and s.style:
                sides[side] = s.style + (f" {_color(s.color)}" if _color(s.color) and _color(s.color) != "#000000" else "")
        if sides:
            if len(set(sides.values())) == 1 and len(sides) == 4:
                d["border"] = next(iter(sides.values())) + " all sides"
            else:
                d["border"] = sides
    al = cell.alignment
    if al is not None:
        ad = {}
        if al.horizontal and al.horizontal != "general":
            ad["horizontal"] = al.horizontal
        if al.vertical and al.vertical != "bottom":
            ad["vertical"] = al.vertical
        if al.wrap_text:
            ad["wrap"] = True
        if al.indent:
            ad["indent"] = al.indent
        if al.text_rotation:
            ad["rotation"] = al.text_rotation
        if ad:
            d["align"] = ad
    if cell.number_format and cell.number_format != "General":
        d["number_format"] = cell.number_format
    if cell.protection is not None and cell.protection.locked is False:
        d["unlocked"] = True
    return d


def render_styles(styles: list[dict[str, Any]]) -> str:
    if not styles:
        return "No formatting in this range (all default)."
    lines = ["Styles:"]
    for s in styles:
        if "column_widths" in s:
            lines.append("- column widths: " + ", ".join(f"{k} {v}" for k, v in s["column_widths"].items()))
            continue
        if "row_heights" in s:
            lines.append("- row heights: " + ", ".join(f"{k} {v}" for k, v in s["row_heights"].items()))
            continue
        parts = []
        for k in ("font", "fill", "border", "align", "number_format", "unlocked"):
            if k in s:
                v = s[k]
                if isinstance(v, dict):
                    v = ", ".join(f"{a}={b}" if b is not True else a for a, b in v.items())
                parts.append(f"{k} {v}" if k != "unlocked" else "unlocked")
        lines.append(f"- {s['cells']}: " + "; ".join(parts))
    return "\n".join(lines)


if __name__ == "__main__":
    run_main(main)
