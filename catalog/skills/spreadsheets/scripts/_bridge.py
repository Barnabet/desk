"""openpyxl ↔ formula engine: evaluate an in-memory workbook, and the save → recalculate pipeline used by
sheet_create.py and sheet_edit.py."""

from __future__ import annotations

import datetime as _dt
import os
import tempfile
from pathlib import Path
from typing import Any

from _a1 import col_letter, parse_range
from _common import SkillError
from _formula import (
    DYNAMIC_FUNCS, Book, Ctx as EvalCtx, Engine, FormulaSyntaxError, TableDef, XLError, functions_used, parse,
)
from _numfmt import date_to_serial


def book_from_openpyxl(wb: Any) -> Book:
    """An engine Book with the values and formulas of an openpyxl workbook (formulas loaded, not data_only)."""
    from openpyxl.worksheet.formula import ArrayFormula

    book = Book()
    try:
        from openpyxl.utils.datetime import CALENDAR_MAC_1904

        book.date1904 = wb.epoch == CALENDAR_MAC_1904
    except (ImportError, AttributeError):
        book.date1904 = False
    calc = getattr(wb, "calculation", None)
    if calc is not None and calc.iterate:
        book.iterate = True
        book.iterate_count = int(calc.iterateCount or 100)
        book.iterate_delta = float(calc.iterateDelta or 0.001)
    for ws in wb.worksheets:
        book.add_sheet(ws.title)
    for ws in wb.worksheets:
        sh = book.sheet(ws.title)
        assert sh is not None
        areas: list[tuple[int, int, tuple[int, int, int, int]]] = []
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                r, c = cell.row, cell.column
                if isinstance(v, ArrayFormula):
                    text = (v.text or "").lstrip("=")
                    try:
                        area = parse_range(v.ref) if v.ref else (r, c, r, c)
                    except ValueError:
                        area = (r, c, r, c)
                    dyn = _uses_dynamic(text)
                    book.set_formula(sh, r, c, text, array_ref=None if dyn else area, dynamic=dyn)
                    areas.append((r, c, area))
                elif isinstance(v, str) and v.startswith("=") and len(v) > 1 and cell.data_type == "f":
                    book.set_formula(sh, r, c, v[1:])
                elif isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
                    book.set_value(sh, r, c, date_to_serial(v, book.date1904))
                elif isinstance(v, _dt.timedelta):
                    book.set_value(sh, r, c, v.total_seconds() / 86400)
                elif isinstance(v, (int, float, str, bool)):
                    book.set_value(sh, r, c, v)
                else:
                    book.set_value(sh, r, c, str(v))
        # values inside an array formula's range are its cached results (openpyxl loads them as constants)
        for r, c, (a1, b1, a2, b2) in areas:
            for rr in range(a1, a2 + 1):
                for cc in range(b1, b2 + 1):
                    if (rr, cc) != (r, c) and (rr, cc) not in sh.formulas:
                        sh.values.pop((rr, cc), None)
    for name, d in wb.defined_names.items():
        if d.attr_text and not name.startswith("_xlnm."):
            book.names[(None, name.upper())] = d.attr_text
    for i, ws in enumerate(wb.worksheets):
        for name, d in ws.defined_names.items():
            if d.attr_text and not name.startswith("_xlnm."):
                book.names[(i, name.upper())] = d.attr_text
        for tname, t in dict.items(ws.tables):
            try:
                area = parse_range(t.ref)
            except ValueError:
                continue
            cols = [tc.name for tc in t.tableColumns]
            book.tables[t.displayName.lower()] = TableDef(t.displayName, ws.title, area, cols, int(t.headerRowCount if t.headerRowCount is not None else 1), int(t.totalsRowCount or 0))
    return book


def _uses_dynamic(text: str) -> bool:
    try:
        return bool(functions_used(parse("=" + text)) & DYNAMIC_FUNCS)
    except FormulaSyntaxError:
        return False


def evaluate(wb: Any, now: Any = None) -> tuple[Book, Engine, dict[str, Any]]:
    book = book_from_openpyxl(wb)
    eng = Engine(book, now=now)
    report = eng.recalc()
    return book, eng, report


def upgrade_to_array_formulas(wb: Any, book: Book, eng: Engine, cells: set[tuple[str, int, int]]) -> list[str]:
    """New formulas that fail under legacy implicit intersection but work as array formulas (e.g.
    INDEX/MATCH with multiplied conditions, SUM(A1:A9*B1:B9)) are stored as single-cell array formulas, which is what
    Excel 365 does when such a formula is typed."""
    from openpyxl.worksheet.formula import ArrayFormula

    changed = []
    for sheet, r, c in sorted(cells):
        sh = book.sheet(sheet)
        if sh is None:
            continue
        fc = sh.formulas.get((r, c))
        if fc is None or fc.array_ref is not None or fc.dynamic or fc.ast is None:
            continue
        if type(fc.value) is not XLError or fc.value.code not in ("#VALUE!", "#N/A"):
            continue
        try:
            f = eng.compiler.compile(fc.ast, sh, r, c, True)
            v = f(EvalCtx(eng, sh, r, c))
        except Exception:  # noqa: BLE001 — keep the formula as it is
            continue
        from _formula import Array, Ref

        if type(v) is Ref:
            if not v.is_cell():
                continue
            v = eng.cell(v.sheet, v.r1, v.c1)
        if type(v) is Array:
            if v.height != 1 or v.width != 1:
                continue
            v = v.rows[0][0]
        if type(v) is XLError:
            continue
        ws = wb[sheet]
        cell = ws.cell(r, c)
        text = cell.value if isinstance(cell.value, str) else None
        if not text:
            continue
        ref = f"{col_letter(c)}{r}"
        cell.value = ArrayFormula(ref, text)
        changed.append(f"{sheet}!{ref}")
    return changed


def computed_values(book: Book, sheet: str) -> dict[tuple[int, int], Any]:
    sh = book.sheet(sheet)
    out: dict[tuple[int, int], Any] = {}
    if sh is None:
        return out
    for key, fc in sh.formulas.items():
        v = fc.value
        out[key] = v.code if type(v) is XLError else v
    return out


def save_and_recalc(wb: Any, out: Path, ctx: Any, recalc: bool = True, now: Any = None, template: bool = False) -> dict[str, Any]:
    """Evaluates new formulas (array upgrade, auto-fit with computed values), saves with openpyxl, then stores
    cached values with the engine. Returns the recalculation report (or {} when recalc is off)."""
    from _ops import autofit

    report: dict[str, Any] = {}
    pending = getattr(ctx, "autofit_pending", [])
    need_eval = recalc and (ctx.new_formulas or pending)
    if need_eval:
        book, eng, _ = evaluate(wb, now)
        upgraded = upgrade_to_array_formulas(wb, book, eng, ctx.new_formulas) if ctx.new_formulas else []
        if upgraded:
            report["array_formulas"] = upgraded
        for ws, cols, lo, hi in pending:
            if ws.title in wb.sheetnames:
                autofit(ws, cols, lo, hi, computed_values(book, ws.title))
    suffix = out.suffix.lower()
    if template or suffix in (".xltx", ".xltm"):
        wb.template = True
    fd, tmp_name = tempfile.mkstemp(prefix=".desk-sheet-", suffix=suffix if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm") else ".xlsx", dir=str(out.parent))
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        try:
            wb.save(str(tmp))
        except Exception as e:  # noqa: BLE001 — openpyxl raises many types
            raise SkillError(f"could not save the workbook: {type(e).__name__}: {e}") from e
        if recalc:
            from _xlsx import recalc_file

            rep = recalc_file(tmp, out, now=now)
            rep.pop("_book", None)
            report.update(rep)
        else:
            from _xlsx import publish

            publish(tmp, out)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    return report


def recalc_summary(rep: dict[str, Any], limit: int = 20) -> list[str]:
    """Short lines for the agent: formulas computed, errors by cell, cycles, unsupported functions."""
    if not rep:
        return []
    lines = []
    counts = rep.get("error_counts") or {}
    n = rep.get("formulas", 0)
    if n:
        lines.append(f"Recalculated {n} formulas" + (": no errors." if not counts else ": " + ", ".join(f"{k} ×{v}" for k, v in counts.items()) + "."))
    for e in [e for e in rep.get("errors", []) if e.get("error")][:limit]:
        lines.append(f"- {e['sheet']}!{e['cell']} {e['error']}: {e['formula']}" + (f" — {e['why']}" if e.get("why") else ""))
    for cyc in rep.get("cycles", [])[:5]:
        lines.append("- circular reference: " + " → ".join(cyc))
    for k, cells in (rep.get("unsupported") or {}).items():
        lines.append(f"- not evaluated ({k}): {', '.join(cells[:6])}")
    if rep.get("array_formulas"):
        lines.append("Stored as array formulas (they need array evaluation in Excel): " + ", ".join(rep["array_formulas"][:10]))
    if rep.get("spills"):
        lines.append("Dynamic arrays: " + ", ".join(f"{s['sheet']}!{s['cell']} → {s['range']}" for s in rep["spills"][:10]))
    return lines
