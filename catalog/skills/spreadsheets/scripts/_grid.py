"""Built-in spreadsheet renderer: draws a range the way a spreadsheet window shows it (Pillow only).

Column letters and row numbers, gridlines, fills (theme and indexed colors with tints), fonts (family, size, bold,
italic, underline, strike, color), borders (all styles), horizontal/vertical alignment, indent, wrapping and overflow
into empty neighbours, number formats (with [Red]-style section colors), merged cells, column widths and row heights
(including automatic heights for wrapped text), hidden rows/columns, frozen-pane lines, comments, conditional formats
(cell rules, formulas, color scales, data bars, icon sets, top/bottom, averages, duplicates, text), images, and charts
drawn from their data. Large ranges are split into pages no larger than VISION_EDGE.
"""

from __future__ import annotations

import colorsys
import datetime as _dt
import io
import math
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from _a1 import col_index, col_letter, parse_range, split_sheet
from _numfmt import date_to_serial, format_value, general

PX_PER_PT = 96 / 72
EMU_PER_PX = 9525
GRID = (225, 225, 225)
HEADER_BG = (243, 243, 243)
HEADER_LINE = (200, 200, 200)
HEADER_TEXT = (90, 90, 90)
PALETTE = ["#4472C4", "#ED7D31", "#A5A5A5", "#FFC000", "#5B9BD5", "#70AD47", "#264478", "#9E480E", "#636363", "#997300", "#255E91", "#43682B"]
DEFAULT_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72"]


# ── colors ──────────────────────────────────────────────────────────────


def hex_rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    if len(h) == 8:
        h = h[2:]
    if len(h) != 6:
        return (0, 0, 0)
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def theme_colors(wb: Any) -> list[str]:
    """Theme palette in Excel's index order (lt1, dk1, lt2, dk2, accent1-6, hlink, folHlink)."""
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return DEFAULT_THEME
    try:
        text = raw.decode("utf-8", "replace") if isinstance(raw, bytes) else str(raw)
    except Exception:  # noqa: BLE001
        return DEFAULT_THEME
    # str.find, not "<a:clrScheme.*?</a:clrScheme>": a lazy regex rescans the rest of the part for every unclosed
    # opening tag, and a theme holding thousands of them stalled every render for minutes.
    m = re.search(r"<a:clrScheme\b", text)
    end = text.find("</a:clrScheme>", m.end()) if m else -1
    if m is None or end < 0:
        return DEFAULT_THEME
    scheme = text[m.start() : end]
    found: dict[str, str] = {}
    for name in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"):
        start = scheme.find(f"<a:{name}>")
        stop = scheme.find(f"</a:{name}>", start) if start >= 0 else -1
        if stop >= 0:
            c = re.search(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', scheme[start:stop])
            if c:
                found[name] = c.group(1).upper()
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
    return [found.get(n, DEFAULT_THEME[i]) for i, n in enumerate(order)]


def _tint(rgb: tuple[int, int, int], tint: float) -> tuple[int, int, int]:
    if not tint:
        return rgb
    h, l, s = colorsys.rgb_to_hls(*(x / 255 for x in rgb))
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return (round(r * 255), round(g * 255), round(b * 255))


def resolve_color(c: Any, theme: list[str], default: tuple[int, int, int] | None) -> tuple[int, int, int] | None:
    if c is None:
        return default
    try:
        t = c.type
    except AttributeError:
        return default
    try:
        if t == "rgb":
            v = c.rgb
            if not isinstance(v, str):
                return default
            if len(v) == 8 and v[:2] == "00" and v[2:] == "000000" and default is not None:
                return default
            return hex_rgb(v)
        if t == "theme":
            idx = int(c.theme)
            base = hex_rgb(theme[idx]) if 0 <= idx < len(theme) else (0, 0, 0)
            return _tint(base, float(c.tint or 0))
        if t == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX

            i = int(c.indexed)
            if i == 64:
                return default if default is not None else (0, 0, 0)
            if i == 65:
                return default if default is not None else (255, 255, 255)
            if 0 <= i < len(COLOR_INDEX):
                return hex_rgb(COLOR_INDEX[i])
        if t == "auto":
            return default
    except (ValueError, TypeError):
        return default
    return default


# ── sizes ───────────────────────────────────────────────────────────────


DEFAULT_WIDTH = 9.140625  # the stored width of Excel's default 8.43-character column (64 px)


def chars_to_width(chars: float) -> float:
    """Stored column width (what files hold) for a width in characters as Excel's UI shows it."""
    return math.trunc((chars * 7 + 5) / 7 * 256) / 256


def col_px(width: float | None, zoom: float) -> int:
    """Pixels of a column from its stored width (Calibri 11: 7 px per digit)."""
    if width is None:
        width = DEFAULT_WIDTH
    if width <= 0:
        return 0
    px = math.trunc(((256 * width + math.trunc(128 / 7)) / 256) * 7)
    return max(2, round(px * zoom))


def row_px(height_pt: float | None, zoom: float) -> int:
    if height_pt is None:
        height_pt = 15
    return max(2, round(height_pt * PX_PER_PT * zoom))


# ── data model ──────────────────────────────────────────────────────────


class Cell:
    __slots__ = ("value", "text", "color", "font", "fill", "borders", "h", "v", "wrap", "indent", "rotation", "numeric", "comment", "bar", "icon", "shrink")

    def __init__(self) -> None:
        self.value: Any = None
        self.text = ""
        self.color: tuple[int, int, int] = (0, 0, 0)
        self.font: tuple[str, float, bool, bool, Any, bool] = ("Calibri", 11.0, False, False, None, False)
        self.fill: tuple[int, int, int] | None = None
        self.borders: dict[str, tuple[str, tuple[int, int, int]]] = {}
        self.h = "general"
        self.v = "bottom"
        self.wrap = False
        self.indent = 0
        self.rotation = 0
        self.numeric = False
        self.comment = False
        self.bar: tuple[float, float, tuple[int, int, int]] | None = None
        self.icon: tuple[int, int, int] | None = None
        self.shrink = False


class View:
    """Everything needed to draw a block of one sheet."""

    def __init__(self, sheet: str) -> None:
        self.sheet = sheet
        self.rows: list[int] = []
        self.cols: list[int] = []
        self.col_w: dict[int, int] = {}
        self.row_h: dict[int, int] = {}
        self.cells: dict[tuple[int, int], Cell] = {}
        self.merges: list[tuple[int, int, int, int]] = []
        self.freeze: tuple[int, int] = (0, 0)
        self.gridlines = True
        self.images: list[tuple[Any, int, int, int, int, int, int]] = []  # (PIL image, r, c, dx, dy, w, h)
        self.charts: list[dict[str, Any]] = []
        self.notes: list[str] = []
        self.zoom = 1.0
        self.total_rows = 0
        self.total_cols = 0


# ── loading ─────────────────────────────────────────────────────────────


def _trimmed_copy(path: Path, sheet: str, r1: int, r2: int) -> Path | None:
    """A temp copy whose big sheets are cut down to what the view needs (openpyxl then loads fast)."""
    from _xlsx import Package

    with Package(path) as pkg:
        infos = {s.path: s for s in pkg.sheets if s.path}
        sizes = {p: pkg.zip.getinfo(p).file_size for p in infos if pkg.has(p)}
        if sum(sizes.values()) < 12_000_000:
            return None
        fd, name = tempfile.mkstemp(prefix="desk-render-", suffix=".xlsx")
        import os

        os.close(fd)
        out = Path(name)
        row_re = re.compile(rb"<(?:\w+:)?row\b[^>]*?\br=\"(\d+)\"[^>]*?(?:/>|>.*?</(?:\w+:)?row>)", re.S)
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for info in pkg.zip.infolist():
                data = pkg.zip.read(info.filename)
                si = infos.get(info.filename)
                if si is not None and sizes.get(info.filename, 0) > 1_000_000:
                    m = re.search(rb"<(?:\w+:)?sheetData\b[^>]*>(.*)</(?:\w+:)?sheetData>", data, re.S)
                    if m:
                        if si.name == sheet:
                            keep = [mm.group(0) for mm in row_re.finditer(m.group(1)) if r1 <= int(mm.group(1)) <= r2]
                            body = b"".join(keep)
                        else:
                            body = b""
                        data = data[: m.start(1)] + body + data[m.end(1) :]
                z.writestr(info, data)
        return out


def load_view(path: Path, sheet: str | None, area: tuple[int, int, int, int] | None, max_rows: int, max_cols: int, zoom: float, show_formulas: bool = False) -> View:
    from _book import kind_of

    kind = kind_of(path)
    if kind != "xlsx":
        return _values_view(path, sheet, area, max_rows, max_cols, zoom)
    return _xlsx_view(path, sheet, area, max_rows, max_cols, zoom, show_formulas)


def _values_view(path: Path, sheet: str | None, area: Any, max_rows: int, max_cols: int, zoom: float) -> View:
    """Unformatted files (csv, xls, xlsb, ods): values with default styling."""
    from _book import Workbook, trim_grid

    with Workbook(path) as wb:
        g = wb.grid(sheet, area)
        name = g.sheet
    rows = g.rows if area else trim_grid(g.rows)
    v = View(name)
    v.zoom = zoom
    r1, c1 = (area[0], area[1]) if area else (g.row1, g.col1)
    nrows = len(rows)
    ncols = max((len(r) for r in rows), default=0)
    v.total_rows, v.total_cols = nrows, ncols
    v.rows = list(range(r1, r1 + max(1, min(nrows, max_rows))))
    v.cols = list(range(c1, c1 + max(1, min(ncols, max_cols))))
    for c in v.cols:
        v.col_w[c] = col_px(None, zoom)
    for r in v.rows:
        v.row_h[r] = row_px(15, zoom)
    widths: dict[int, float] = {}
    for i, row in enumerate(rows[: len(v.rows)]):
        for j, val in enumerate(row[: len(v.cols)]):
            if val is None:
                continue
            cell = Cell()
            cell.value = val
            _set_text(cell, val, "General", False)
            v.cells[(r1 + i, c1 + j)] = cell
            widths[c1 + j] = max(widths.get(c1 + j, 0), len(cell.text) * 1.05 + 1)
    for c, w in widths.items():
        v.col_w[c] = col_px(chars_to_width(min(max(8.43, w), 50)), zoom)
    for m in g.merged:
        try:
            v.merges.append(parse_range(m))
        except ValueError:
            pass
    v.notes.append("no cell formatting in this format: values drawn with default styling")
    return v


def _set_text(cell: Cell, val: Any, fmt: str | None, date1904: bool, width_chars: int | None = None) -> None:
    color = None
    if val is None:
        cell.text = ""
        return
    if hasattr(val, "code"):
        cell.text = val.code
        cell.numeric = False
        return
    if isinstance(val, bool):
        cell.text = "TRUE" if val else "FALSE"
        return
    if isinstance(val, (_dt.datetime, _dt.date, _dt.time)):
        cell.numeric = True
        try:
            text, color = format_value(val, fmt or "General", date1904, width_chars)
        except Exception:  # noqa: BLE001
            text = str(val)
        cell.text = text
    elif isinstance(val, (int, float)):
        cell.numeric = True
        try:
            text, color = format_value(val, fmt or "General", date1904, width_chars)
        except Exception:  # noqa: BLE001
            text = general(val, width_chars or 11)
        cell.text = text
    else:
        s = str(val)
        if fmt and fmt != "General" and "@" in fmt:
            try:
                s, color = format_value(s, fmt, date1904)
            except Exception:  # noqa: BLE001
                pass
        cell.text = s
    if color:
        cell.color = hex_rgb(color)


class Values:
    """Cell values to draw: the results cached in the file, or the built-in engine's when formulas have none."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self._cached: dict[str, dict[tuple[int, int], Any]] = {}
        self._missing: dict[str, bool] = {}
        self._book: Any = None
        self._engine: Any = None
        self.recalculated = False

    def load(self, sheet: str, upto_row: int | None = None) -> None:
        """Streams one sheet's stored values (stopping after upto_row)."""
        if sheet in self._cached:
            return
        from _xlsx import Package

        vals: dict[tuple[int, int], Any] = {}
        missing = False
        with Package(self.path) as pkg:
            si = pkg.sheet(sheet)
            gen = pkg.iter_cells(si)
            try:
                for r, c, val, f, _s in gen:
                    if upto_row is not None and r > upto_row:
                        break
                    if f is not None and val is None and f["type"] != "dataTable":
                        missing = True
                    if val is not None:
                        vals[(r, c)] = val
            finally:
                gen.close()
        self._cached[sheet] = vals
        self._missing[sheet] = missing

    def missing(self, sheet: str) -> bool:
        return self._missing.get(sheet, False)

    def book(self) -> Any:
        """The engine Book (formulas evaluated on demand, or all of them after recalc())."""
        if self._book is None:
            from _formula import Engine
            from _xlsx import load_book

            book, pkg, _ = load_book(self.path)
            pkg.close()
            self._book = book
            self._engine = Engine(book)
        return self._book

    def recalc(self) -> None:
        self.book()
        if not self.recalculated:
            self._engine.recalc()
            self.recalculated = True

    def get(self, sheet: str, r: int, c: int) -> Any:
        if not self.recalculated:
            if sheet not in self._cached:
                self.load(sheet)
            return self._cached[sheet].get((r, c))
        sh = self._book.sheet(sheet)
        if sh is None:
            return None
        return self._engine.cell(sh, r, c)

    def evaluate(self, formula: str, sheet: str, r: int, c: int) -> Any:
        from _formula import Ctx, Ref, parse

        book = self.book()
        eng = self._engine
        sh = book.sheet(sheet)
        f = eng.compiler.compile(parse(formula), sh, r, c, False)
        val = f(Ctx(eng, sh, r, c))
        if type(val) is Ref:
            val = eng.cell(val.sheet, val.r1, val.c1) if val.is_cell() else eng.ref_array(val)
        if hasattr(val, "rows"):
            val = val.rows[0][0]
        return val


def _sheet_target(pkg: Any, sheet: str | None) -> Any:
    names = [s.name for s in pkg.sheets]
    if not names:
        raise ValueError("the workbook has no sheets")
    if sheet is None:
        return next((s for s in pkg.sheets if s.path and s.state == "visible"), None) or next((s for s in pkg.sheets if s.path), pkg.sheets[0])
    if isinstance(sheet, str) and sheet.isdigit() and sheet not in names:
        i = int(sheet) - 1
        if not (0 <= i < len(names)):
            raise ValueError(f"sheet {sheet} is out of range (the workbook has {len(names)})")
        return pkg.sheets[i]
    for s in pkg.sheets:
        if s.name == sheet:
            return s
    for s in pkg.sheets:
        if s.name.lower() == str(sheet).lower():
            return s
    raise ValueError(f"no sheet named '{sheet}'; sheets: {', '.join(names)}")


def _xlsx_view(path: Path, sheet: str | None, area: Any, max_rows: int, max_cols: int, zoom: float, show_formulas: bool) -> View:
    import warnings

    import openpyxl

    from _xlsx import Package

    with Package(path) as pkg:
        si = _sheet_target(pkg, sheet)
        target = si.name
        dim = None
        if si.path:
            with pkg.zip.open(si.path) as fh:
                head = fh.read(4096)
            m = re.search(rb'<(?:\w+:)?dimension\b[^>]*\bref="([^"]*)"', head)
            if m:
                try:
                    dim = parse_range(m.group(1).decode())
                except ValueError:
                    dim = None
        aliases = _title_aliases([s.name for s in pkg.sheets])
    vals = Values(path)
    if not si.path:
        if si.state != "chart":
            raise ValueError(f"'{target}' is not a worksheet (dialog or macro sheet); nothing to draw")
        copy = _aliased_copy(path, aliases) if aliases else None
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = openpyxl.load_workbook(copy or path, rich_text=False)
        finally:
            _unlink(copy)
        v = View(target)
        v.zoom = zoom
        for ch in getattr(wb[aliases.get(target, target)], "_charts", []):
            _add_chart(v, ch, vals, target, zoom)
        if not v.charts:
            v.notes.append("this chart sheet's chart could not be read")
        return v
    r_lo = area[0] if area else 1
    r_hi = area[2] if area and area[2] < 1_048_576 else (dim[2] if dim else 1_048_576)
    upto = min(r_hi, r_lo + max_rows + 50)
    trimmed = _trimmed_copy(path, target, r_lo, upto)
    vals.load(target, upto if trimmed is not None else None)
    extent = None
    if trimmed is not None:
        from _scan import scan_cached

        st = scan_cached(path, si.path) or {}
        if st.get("max_row"):
            extent = (int(st["max_row"]), int(st.get("max_col") or 1))
    copy = None
    try:
        if aliases:
            copy = _aliased_copy(trimmed or path, aliases)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(copy or trimmed or path, rich_text=False)
        return _build_view(vals, wb, target, area, max_rows, max_cols, zoom, show_formulas, extent, ws_title=aliases.get(target))
    finally:
        _unlink(trimmed)
        _unlink(copy)


def _unlink(p: Path | None) -> None:
    if p is not None:
        try:
            p.unlink()
        except OSError:
            pass


#: characters openpyxl refuses in a sheet name (Excel does too, but other programs write them)
_BAD_TITLE = re.compile(r"[\\*?:/\[\]]")


def _title_aliases(names: list[str]) -> dict[str, str]:
    """{sheet name: stand-in} for the sheet names openpyxl cannot load (empty, or with \\ * ? : / [ ])."""
    taken = {n.lower() for n in names}
    out: dict[str, str] = {}
    for i, name in enumerate(names, 1):
        if name and not _BAD_TITLE.search(name):
            continue
        alias, k = f"Sheet{i}", 1
        while alias.lower() in taken:
            k += 1
            alias = f"Sheet{i}_{k}"
        taken.add(alias.lower())
        out[name] = alias
    return out


def _aliased_copy(path: Path, aliases: dict[str, str]) -> Path:
    """A temp copy of the workbook whose sheets named in aliases carry their stand-in names in workbook.xml, so
    openpyxl loads it. Cells, formulas and charts keep the real names; the renderer reads values by the real ones."""
    import html
    import os

    from _xlsx import Package, rewrite_zip

    fd, name = tempfile.mkstemp(prefix="desk-render-names-", suffix=".xlsx")
    os.close(fd)
    out = Path(name)

    def rename(m: re.Match[bytes]) -> bytes:
        el = m.group(0)
        nm = re.search(rb'\sname="([^"]*)"', el)
        if nm is None:
            return el
        real = html.unescape(nm.group(1).decode("utf-8", "replace"))
        if real not in aliases:
            return el
        return el[: nm.start(1)] + aliases[real].encode() + el[nm.end(1) :]

    with Package(path) as pkg:
        wb_part = pkg.workbook_path
        xml = re.sub(rb"<(?:\w+:)?sheet\b[^>]*>", rename, pkg.read(wb_part))
        rewrite_zip(pkg.zip, out, lambda n: xml if n == wb_part else None)
    return out


def _build_view(vals: Values, wb: Any, target: str, area: Any, max_rows: int, max_cols: int, zoom: float, show_formulas: bool, extent: tuple[int, int] | None = None, ws_title: str | None = None) -> View:
    """extent: the sheet's real (last row, last column) when wb is a copy trimmed to the rows drawn. ws_title: the
    sheet's stand-in name in wb when its real name (target) is one openpyxl cannot load."""
    from openpyxl.worksheet.formula import ArrayFormula

    ws = wb[ws_title or target]
    theme = theme_colors(wb)
    v = View(target if ws_title else ws.title)
    v.zoom = zoom
    date1904 = False
    try:
        from openpyxl.utils.datetime import CALENDAR_MAC_1904

        date1904 = wb.epoch == CALENDAR_MAC_1904
    except (ImportError, AttributeError):
        pass
    # area
    if area:
        r1, c1, r2, c2 = area
        r2 = min(r2, max(ws.max_row, r1))
        c2 = min(c2, max(ws.max_column, c1))
    else:
        r1, c1 = max(1, ws.min_row), max(1, ws.min_column)
        r2, c2 = max(ws.max_row, r1), max(ws.max_column, c1)
        r1, c1 = 1, 1  # a window shows the sheet from A1
        for ch in list(getattr(ws, "_charts", [])) + list(getattr(ws, "_images", [])):
            pos = _anchor_cells(ch)
            if pos:
                r2 = max(r2, min(pos[2], r1 + max_rows - 1))
                c2 = max(c2, min(pos[3], c1 + max_cols - 1))
    v.total_rows, v.total_cols = r2 - r1 + 1, c2 - c1 + 1
    # hidden rows / columns and sizes
    default_w = float(ws.sheet_format.defaultColWidth) if ws.sheet_format.defaultColWidth else chars_to_width(float(ws.sheet_format.baseColWidth or 8) + 0.43)
    default_h = float(ws.sheet_format.defaultRowHeight or 15)
    col_widths: dict[int, float | None] = {}
    hidden_cols: set[int] = set()
    for key, d in ws.column_dimensions.items():
        try:
            lo = d.min or col_index(key)
            hi = d.max or lo
        except ValueError:
            continue
        for c in range(lo, min(hi, c2) + 1):
            if d.hidden:
                hidden_cols.add(c)
            if d.width:
                col_widths[c] = float(d.width)
    cols = [c for c in range(c1, c2 + 1) if c not in hidden_cols][:max_cols]
    rows = []
    custom_h: dict[int, float] = {}
    for r in range(r1, r2 + 1):
        d = ws.row_dimensions.get(r)
        if d is not None and d.hidden:
            continue
        if d is not None and d.height is not None:
            custom_h[r] = float(d.height)
        rows.append(r)
        if len(rows) >= max_rows:
            break
    v.rows, v.cols = rows, cols
    last_row = r2
    if extent is not None:
        last_row = max(r2, min(area[2], extent[0]) if area else extent[0])
    if last_row > r2 or len(rows) < v.total_rows - sum(1 for r in range(r1, r2 + 1) if (ws.row_dimensions.get(r) is not None and ws.row_dimensions[r].hidden)):
        v.notes.append(f"showing rows {rows[0]}-{rows[-1]} of {r1}-{last_row}; use --range or --max-rows for more")
    if len(cols) < (c2 - c1 + 1) - len(hidden_cols & set(range(c1, c2 + 1))):
        v.notes.append(f"showing columns {col_letter(cols[0])}-{col_letter(cols[-1])} of {col_letter(c1)}-{col_letter(c2)}; use --range or --max-cols for more")
    for c in cols:
        v.col_w[c] = col_px(col_widths.get(c, default_w), zoom)
    try:
        v.gridlines = ws.sheet_view.showGridLines is not False
    except AttributeError:
        pass
    fr = ws.freeze_panes
    if fr:
        try:
            fr_r, fr_c = parse_range(fr)[:2]
            v.freeze = (fr_r - 1, fr_c - 1)
        except ValueError:
            pass
    # values: the results cached in the file, or the engine's when some formula has none stored
    if vals.missing(target) and not show_formulas:
        vals.recalc()
        v.notes.append("formula results computed by the built-in engine (the file stores none for some formulas)")
    row_set, col_set = set(rows), set(cols)
    has_cf = bool(ws.conditional_formatting)

    merged_tl: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for mr in ws.merged_cells.ranges:
        rect = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        if rect[2] < rows[0] or rect[0] > rows[-1] or rect[3] < cols[0] or rect[1] > cols[-1]:
            continue
        v.merges.append(rect)
        merged_tl[(rect[0], rect[1])] = rect
    # cells
    for r in rows:
        for c in cols:
            src = ws._cells.get((r, c))
            if src is None:
                # spilled results have no cell object when the file stores none
                val = vals.get(target, r, c)
                if val is None:
                    continue
                cell = Cell()
                cell.value = val
                _set_text(cell, val, "General", date1904)
                v.cells[(r, c)] = cell
                continue
            cell = Cell()
            raw = src.value
            is_formula = isinstance(raw, ArrayFormula) or (isinstance(raw, str) and src.data_type == "f")
            if show_formulas and is_formula:
                val = raw.text if isinstance(raw, ArrayFormula) else raw
                if not str(val).startswith("="):
                    val = "=" + str(val)
            else:
                val = vals.get(target, r, c)
                if val is None and is_formula:
                    val = 0 if vals.recalculated else None
            if val is not None and not isinstance(val, (str, int, float, bool, _dt.date, _dt.datetime, _dt.time)) and not hasattr(val, "code"):
                val = str(val)
            cell.value = val
            f = src.font
            if f is not None:
                cell.font = (f.name or "Calibri", float(f.sz or 11), bool(f.b), bool(f.i), f.u, bool(f.strike))
                col = resolve_color(f.color, theme, (0, 0, 0))
                cell.color = col or (0, 0, 0)
            fill = src.fill
            if fill is not None and fill.fill_type not in (None, "none"):
                if fill.fill_type == "solid":
                    cell.fill = resolve_color(fill.fgColor, theme, None) or resolve_color(fill.start_color, theme, None)
                elif getattr(fill, "fill_type", None) == "gradient" or fill.__class__.__name__ == "GradientFill":
                    stops = getattr(fill, "stop", None)
                    if stops:
                        cell.fill = resolve_color(stops[0].color, theme, None)
                else:
                    fg = resolve_color(fill.fgColor, theme, None)
                    if fg:
                        cell.fill = tuple(round(x * 0.5 + 255 * 0.5) for x in fg)  # type: ignore[assignment]
            elif fill is not None and fill.__class__.__name__ == "GradientFill":
                stops = getattr(fill, "stop", None)
                if stops:
                    cell.fill = resolve_color(stops[0].color, theme, None)
            b = src.border
            if b is not None:
                for side in ("left", "right", "top", "bottom"):
                    s = getattr(b, side)
                    if s is not None and s.style:
                        cell.borders[side] = (s.style, resolve_color(s.color, theme, (0, 0, 0)) or (0, 0, 0))
            al = src.alignment
            if al is not None:
                cell.h = al.horizontal or "general"
                cell.v = al.vertical or "bottom"
                cell.wrap = bool(al.wrap_text)
                cell.indent = int(al.indent or 0)
                cell.rotation = int(al.text_rotation or 0)
                cell.shrink = bool(al.shrink_to_fit)
            cell.comment = src.comment is not None
            fmt = src.number_format
            if show_formulas and isinstance(val, str) and val.startswith("="):
                cell.text = val
            else:
                span = merged_tl.get((r, c))
                wpx = sum(v.col_w.get(cc, 0) for cc in range(c, (span[3] if span else c) + 1))
                chars = max(1, int((wpx - 5) / (7 * zoom)))
                _set_text(cell, val, fmt, date1904, chars if (fmt in (None, "General") and isinstance(val, float)) else None)
            if cell.text or cell.fill or cell.borders or cell.comment:
                v.cells[(r, c)] = cell
    try:
        _apply_table_styles(ws, v, theme)
    except Exception as e:  # noqa: BLE001 — table styling is a nicety
        v.notes.append(f"table styles not drawn ({type(e).__name__})")
    if has_cf:
        try:
            _apply_conditional(ws, v, vals, target, wb, theme, date1904)
        except Exception as e:  # noqa: BLE001 — conditional formats are a nicety
            v.notes.append(f"conditional formats not drawn ({type(e).__name__})")
    # row heights (auto-fit wrapped text and large fonts where no custom height is set)
    for r in rows:
        if r in custom_h:
            v.row_h[r] = row_px(custom_h[r], zoom)
            continue
        need = default_h
        for c in cols:
            cell = v.cells.get((r, c))
            if cell is None or not cell.text:
                continue
            pt = cell.font[1]
            lines = 1
            if cell.wrap:
                span = merged_tl.get((r, c))
                if span and span[2] > span[0]:
                    continue
                wpx = sum(v.col_w.get(cc, 0) for cc in range(c, (span[3] if span else c) + 1))
                lines = len(_wrap_lines(cell.text, _font_for(cell, zoom)[0], max(4, wpx - 6)))
            elif "\n" in cell.text:
                lines = 1
            need = max(need, pt * 1.36 * lines + 1.5 if lines > 1 else pt * 1.36)
        v.row_h[r] = row_px(min(need, 409), zoom)
    # images and charts
    for img in getattr(ws, "_images", []):
        try:
            _add_image(v, img, zoom)
        except Exception:  # noqa: BLE001
            v.notes.append("an image could not be drawn")
    for ch in getattr(ws, "_charts", []):
        try:
            _add_chart(v, ch, vals, target, zoom)
        except Exception as e:  # noqa: BLE001
            v.notes.append(f"a chart could not be drawn ({type(e).__name__}: {e})")
    if ws.__class__.__name__ != "Chartsheet":
        inside = set(rows), set(cols)
        outside = [ch for ch in v.charts if ch["box"][0] not in inside[0] or ch["box"][1] not in inside[1]]
        if outside:
            v.charts = [ch for ch in v.charts if ch not in outside]
            where = ", ".join(f"{col_letter(ch['box'][1])}{ch['box'][0]}" + (f" '{ch['title']}'" if ch.get("title") else "") for ch in outside[:6])
            v.notes.append(f"{len(outside)} chart(s) anchored outside the drawn range: {where}")
        v.images = [im for im in v.images if im[1] in inside[0] and im[2] in inside[1]]
    if not area:
        _extend_for_drawings(v, ws, hidden_cols, col_widths, default_w, custom_h, default_h, max_rows, max_cols, zoom)
    return v


def _extend_for_drawings(v: View, ws: Any, hidden_cols: set[int], col_widths: dict[int, float | None], default_w: float, custom_h: dict[int, float], default_h: float, max_rows: int, max_cols: int, zoom: float) -> None:
    """Adds columns and rows so charts and images anchored in the view are drawn whole (within the limits)."""
    boxes = [(r, c, dx, dy, w, h) for _im, r, c, dx, dy, w, h in v.images] + [ch["box"] for ch in v.charts]
    for r, c, dx, dy, w, h in boxes:
        if w <= 0 or h <= 0 or c not in v.col_w or r not in v.rows:
            continue
        need, got = dx + w, 0
        k = c
        while got < need:
            if k not in v.col_w:
                if len(v.cols) >= max_cols or k > 16384:
                    break
                if k in hidden_cols:
                    k += 1
                    continue
                v.cols.append(k)
                v.col_w[k] = col_px(col_widths.get(k, default_w), zoom)
            got += v.col_w[k]
            k += 1
        need, got = dy + h, 0
        k = r
        while got < need:
            if k not in v.row_h and k not in v.rows:
                if len(v.rows) >= max_rows or k > 1048576:
                    break
                d = ws.row_dimensions.get(k)
                if d is not None and d.hidden:
                    k += 1
                    continue
                v.rows.append(k)
                v.row_h[k] = row_px(float(d.height) if d is not None and d.height else default_h, zoom)
            got += v.row_h.get(k, 0)
            k += 1
    v.cols.sort()
    v.rows.sort()


# ── conditional formatting ──────────────────────────────────────────────


# ── Excel table styles (TableStyleLight1…21, Medium1…28, Dark1…11), approximated ──


def _table_palette(style: str, theme: list[str]) -> dict[str, Any] | None:
    m = re.fullmatch(r"TableStyle(Light|Medium|Dark)(\d+)", style or "")
    if not m:
        return None
    family, n = m.group(1), int(m.group(2))
    idx = (n - 1) % 7  # 0 = neutral (black/grey), 1-6 = accent1-6
    base = (0, 0, 0) if idx == 0 else hex_rgb(theme[3 + idx] if 3 + idx < len(theme) else "4472C4")
    light = (217, 217, 217) if idx == 0 else _tint(base, 0.8)
    mid = (166, 166, 166) if idx == 0 else _tint(base, 0.6)
    white = (255, 255, 255)
    band = (n - 1) // 7
    if family == "Light":
        if band == 0:
            return {"head_font": base if idx else (0, 0, 0), "head_border": base, "stripe": light, "line": base}
        if band == 1:
            return {"head_fill": base, "head_font": white, "outline": base}
        return {"head_font": (0, 0, 0), "stripe": light, "grid": base}
    if family == "Medium":
        if band == 0:
            return {"head_fill": base, "head_font": white, "stripe": light}
        if band == 1:
            return {"head_fill": base, "head_font": white, "stripe": mid, "row": light, "grid": white}
        if band == 2:
            return {"head_fill": base, "head_font": white, "stripe": light, "grid": base}
        return {"head_fill": light, "head_font": (0, 0, 0), "stripe": light}
    dark = _tint(base, -0.5) if idx else (64, 64, 64)
    if band == 0:
        return {"head_fill": (0, 0, 0), "head_font": white, "stripe": dark, "row": _tint(dark, 0.2), "body_font": white}
    return {"head_fill": dark, "head_font": white, "stripe": mid, "row": light}


def _apply_table_styles(ws: Any, v: View, theme: list[str]) -> None:
    visible_r, visible_c = set(v.rows), set(v.cols)
    for tbl in dict.values(ws.tables):
        info = tbl.tableStyleInfo
        pal = _table_palette(getattr(info, "name", None) or "TableStyleMedium2", theme) if info is not None else None
        if pal is None:
            continue
        r1, c1, r2, c2 = parse_range(tbl.ref)
        head = 0 if tbl.headerRowCount == 0 else 1
        tot = int(tbl.totalsRowCount or 0)
        stripes = bool(getattr(info, "showRowStripes", True))
        for r in range(r1, r2 + 1):
            if r not in visible_r:
                continue
            for c in range(c1, c2 + 1):
                if c not in visible_c:
                    continue
                cell = v.cells.get((r, c))
                if cell is None:
                    cell = Cell()
                    v.cells[(r, c)] = cell
                own_fill = cell.fill is not None
                if head and r == r1:
                    if pal.get("head_fill") and not own_fill:
                        cell.fill = pal["head_fill"]
                    if cell.color == (0, 0, 0):
                        cell.color = pal.get("head_font", (0, 0, 0))
                    name, sz, _b, it, u, st = cell.font
                    cell.font = (name, sz, True, it, u, st)
                    if pal.get("head_border"):
                        cell.borders.setdefault("bottom", ("thin", pal["head_border"]))
                elif tot and r == r2:
                    name, sz, _b, it, u, st = cell.font
                    cell.font = (name, sz, True, it, u, st)
                    cell.borders.setdefault("top", ("double", pal.get("head_fill") or pal.get("line") or (0, 0, 0)))
                else:
                    k = r - r1 - head
                    fill = pal.get("stripe") if stripes and k % 2 == 0 else pal.get("row")
                    if fill and not own_fill:
                        cell.fill = fill
                    if pal.get("body_font") and cell.color == (0, 0, 0):
                        cell.color = pal["body_font"]
                if pal.get("grid"):
                    for side in ("left", "right", "top", "bottom"):
                        cell.borders.setdefault(side, ("thin", pal["grid"]))
                if pal.get("outline"):
                    for side, edge in (("left", c == c1), ("right", c == c2), ("top", r == r1), ("bottom", r == r2)):
                        if edge:
                            cell.borders.setdefault(side, ("thin", pal["outline"]))
                if pal.get("line") and r == r2:
                    cell.borders.setdefault("bottom", ("thin", pal["line"]))


def _cf_values(v: View, vals: Values, sheet: str, ranges: Any, limit: tuple[int, int]) -> list[tuple[int, int, Any]]:
    out = []
    for rng in str(ranges).split():
        try:
            r1, c1, r2, c2 = parse_range(rng)
        except ValueError:
            continue
        r2, c2 = min(r2, limit[0]), min(c2, limit[1])
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.append((r, c, vals.get(sheet, r, c)))
    return out


def _dxf_apply(cell: Cell, dxf: Any, theme: list[str]) -> None:
    if dxf is None:
        return
    fill = getattr(dxf, "fill", None)
    if fill is not None:
        col = resolve_color(getattr(fill, "bgColor", None), theme, None) or resolve_color(getattr(fill, "fgColor", None), theme, None) or resolve_color(getattr(fill, "start_color", None), theme, None)
        if col:
            cell.fill = col
    font = getattr(dxf, "font", None)
    if font is not None:
        col = resolve_color(font.color, theme, None)
        if col:
            cell.color = col
        name, sz, b, i, u, st = cell.font
        cell.font = (name, sz, bool(font.b) if font.b is not None else b, bool(font.i) if font.i is not None else i, font.u or u, bool(font.strike) if font.strike is not None else st)


def _apply_conditional(ws: Any, v: View, vals: Values, sheet: str, wb: Any, theme: list[str], date1904: bool) -> None:
    from _formula import XLError, compare, translate

    dxfs = getattr(wb, "_differential_styles", None)
    visible = {(r, c) for r in v.rows for c in v.cols}
    items = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            items.append((rule.priority or 0, cf, rule))
    items.sort(key=lambda t: t[0])
    done: set[tuple[int, int]] = set()
    for _prio, cf, rule in items:
        cells = _cf_values(v, vals, sheet, cf.sqref, (max(ws.max_row, 1), max(ws.max_column, 1)))
        if not cells:
            continue
        first = min((r, c) for r, c, _ in cells)
        dxf = rule.dxf
        if dxf is None and rule.dxfId is not None and dxfs is not None:
            try:
                dxf = dxfs[rule.dxfId]
            except (IndexError, TypeError):
                dxf = None
        t = rule.type
        nums = [val for _, _, val in cells if isinstance(val, (int, float)) and not isinstance(val, bool)]

        def cell_for(r: int, c: int) -> Cell | None:
            if (r, c) not in visible:
                return None
            cell = v.cells.get((r, c))
            if cell is None:
                cell = Cell()
                v.cells[(r, c)] = cell
            return cell

        if t == "colorScale" and rule.colorScale is not None and nums:
            cs = rule.colorScale
            colors = [resolve_color(c, theme, (255, 255, 255)) or (255, 255, 255) for c in cs.color]
            stops = [_cfvo(x, nums, vals, sheet, first) for x in cs.cfvo][: len(colors)]
            if len(stops) < 2 or len(stops) != len(colors):
                continue
            for r, c, val in cells:
                if not isinstance(val, (int, float)) or isinstance(val, bool):
                    continue
                cell = cell_for(r, c)
                if cell is None:
                    continue
                cell.fill = _scale_color(val, stops, colors)
            continue
        if t == "dataBar" and rule.dataBar is not None and nums:
            col = resolve_color(rule.dataBar.color, theme, (99, 142, 198)) or (99, 142, 198)
            cf = list(rule.dataBar.cfvo or [])
            lo = _cfvo(cf[0], nums, vals, sheet, first) if len(cf) >= 1 else min(nums)
            hi = _cfvo(cf[1], nums, vals, sheet, first) if len(cf) >= 2 else max(nums)
            lo, hi = min(lo, 0.0), max(hi, 0.0)
            axis = 0.0 if lo >= 0 else -lo / (hi - lo) if hi != lo else 0.0
            for r, c, val in cells:
                if not isinstance(val, (int, float)) or isinstance(val, bool):
                    continue
                cell = cell_for(r, c)
                if cell is None:
                    continue
                pos = (min(max(val, lo), hi) - lo) / (hi - lo) if hi != lo else 1.0
                if lo >= 0:
                    pos = max(pos, 0.1)
                cell.bar = (axis, pos, col if val >= 0 else (255, 0, 0))
            continue
        if t == "iconSet" and nums and rule.iconSet is not None:
            ics = rule.iconSet
            cf = list(ics.cfvo or [])
            n = len(cf) or 3
            bounds = [_cfvo(x, nums, vals, sheet, first) for x in cf[1:]] if cf else [min(nums) + (max(nums) - min(nums)) * k / 3 for k in (1, 2)]
            palette = _ICON_COLORS.get(n, _ICON_COLORS[3])
            if ics.reverse:
                palette = palette[::-1]
            for r, c, val in cells:
                if not isinstance(val, (int, float)) or isinstance(val, bool):
                    continue
                cell = cell_for(r, c)
                if cell is None:
                    continue
                level = sum(1 for bnd in bounds if val >= bnd)
                cell.icon = palette[min(level, len(palette) - 1)]
                if ics.showValue is False:
                    cell.text = ""
            continue
        test = None
        if t == "cellIs" and rule.formula:
            try:
                targets = [vals.evaluate("=" + f, sheet, first[0], first[1]) for f in rule.formula]
            except Exception:  # noqa: BLE001
                targets = [_literal(f) for f in rule.formula]
            targets = [x.rows[0][0] if hasattr(x, "rows") else x for x in targets]
            op = rule.operator or "equal"

            def test(r: int, c: int, val: Any, op: str = op, targets: list[Any] = targets) -> bool:
                if val is None:
                    val = 0
                if type(val) is XLError:
                    return False
                cmp0 = compare(val, targets[0])
                if op == "between" and len(targets) > 1:
                    return compare(val, min(targets[:2], key=_key)) >= 0 and compare(val, max(targets[:2], key=_key)) <= 0
                if op == "notBetween" and len(targets) > 1:
                    return not (compare(val, min(targets[:2], key=_key)) >= 0 and compare(val, max(targets[:2], key=_key)) <= 0)
                return {"equal": cmp0 == 0, "notEqual": cmp0 != 0, "greaterThan": cmp0 > 0, "lessThan": cmp0 < 0, "greaterThanOrEqual": cmp0 >= 0, "lessThanOrEqual": cmp0 <= 0}.get(op, False)
        elif t in ("expression", "containsText", "notContainsText", "beginsWith", "endsWith", "containsBlanks", "notContainsBlanks", "containsErrors", "notContainsErrors") and rule.formula:
            formula = "=" + rule.formula[0]

            def test(r: int, c: int, val: Any, formula: str = formula) -> bool:
                try:
                    res = vals.evaluate(translate(formula, r - first[0], c - first[1]), sheet, r, c)
                except Exception:  # noqa: BLE001
                    return False
                if hasattr(res, "rows"):
                    res = res.rows[0][0]
                return res is True or (isinstance(res, (int, float)) and not isinstance(res, bool) and res != 0)
        elif t == "top10" and nums:
            k = int(rule.rank or 10)
            if rule.percent:
                k = max(1, int(len(nums) * k / 100))
            ordered = sorted(nums, reverse=not rule.bottom)
            limit = ordered[min(k, len(ordered)) - 1]

            def test(r: int, c: int, val: Any, limit: float = limit, bottom: bool = bool(rule.bottom)) -> bool:
                return isinstance(val, (int, float)) and not isinstance(val, bool) and (val <= limit if bottom else val >= limit)
        elif t == "aboveAverage" and nums:
            avg = sum(nums) / len(nums)
            above = rule.aboveAverage is not False

            def test(r: int, c: int, val: Any, avg: float = avg, above: bool = above) -> bool:
                return isinstance(val, (int, float)) and not isinstance(val, bool) and (val > avg if above else val < avg)
        elif t in ("duplicateValues", "uniqueValues"):
            counts: dict[Any, int] = {}
            for _, _, val in cells:
                if val is not None:
                    k = val.lower() if isinstance(val, str) else val
                    counts[k] = counts.get(k, 0) + 1
            dup = t == "duplicateValues"

            def test(r: int, c: int, val: Any, counts: dict[Any, int] = counts, dup: bool = dup) -> bool:
                if val is None:
                    return False
                k = val.lower() if isinstance(val, str) else val
                return counts.get(k, 0) > 1 if dup else counts.get(k, 0) == 1
        if test is None:
            continue
        for r, c, val in cells:
            if (r, c) not in visible or ((r, c) in done and rule.stopIfTrue):
                continue
            try:
                ok = test(r, c, val)
            except Exception:  # noqa: BLE001
                ok = False
            if ok:
                cell = cell_for(r, c)
                if cell is not None:
                    _dxf_apply(cell, dxf, theme)
                    if rule.stopIfTrue:
                        done.add((r, c))


_ICON_COLORS = {
    3: [(248, 105, 107), (255, 192, 0), (99, 190, 123)],
    4: [(248, 105, 107), (255, 192, 0), (180, 200, 90), (99, 190, 123)],
    5: [(248, 105, 107), (250, 150, 60), (255, 192, 0), (180, 200, 90), (99, 190, 123)],
}


def _cfvo(x: Any, nums: list[float], vals: Values, sheet: str, first: tuple[int, int]) -> float:
    """The number a conditional-format threshold (min, max, num, percent, percentile, formula) stands for."""
    lo, hi = min(nums), max(nums)
    t = getattr(x, "type", None) or "min"
    raw = getattr(x, "val", None)
    try:
        if t in ("min", "autoMin"):
            return lo
        if t in ("max", "autoMax"):
            return hi
        if t == "formula" or (isinstance(raw, str) and not _is_number(raw)):
            got = vals.evaluate("=" + str(raw).lstrip("="), sheet, first[0], first[1])
            return float(got) if isinstance(got, (int, float)) and not isinstance(got, bool) else lo
        v = float(raw) if raw is not None else 0.0
        if t == "num":
            return v
        if t == "percent":
            return lo + (hi - lo) * v / 100
        if t == "percentile":
            ordered = sorted(nums)
            k = (len(ordered) - 1) * v / 100
            f = math.floor(k)
            return ordered[f] + (ordered[min(f + 1, len(ordered) - 1)] - ordered[f]) * (k - f)
    except (TypeError, ValueError):
        pass
    return lo


def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def _scale_color(val: float, stops: list[float], colors: list[tuple[int, int, int]]) -> tuple[int, int, int]:
    if val <= stops[0]:
        return colors[0]
    for i in range(1, len(stops)):
        if val <= stops[i]:
            span = stops[i] - stops[i - 1]
            f = 0.0 if span == 0 else (val - stops[i - 1]) / span
            a, b = colors[i - 1], colors[i]
            return (round(a[0] + (b[0] - a[0]) * f), round(a[1] + (b[1] - a[1]) * f), round(a[2] + (b[2] - a[2]) * f))
    return colors[-1]


def _key(x: Any) -> tuple:
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        return (0, x)
    return (1, str(x))


def _literal(f: str) -> Any:
    s = f.strip()
    if s.startswith('"') and s.endswith('"'):
        return s[1:-1].replace('""', '"')
    try:
        return float(s)
    except ValueError:
        return s


# ── images and charts ───────────────────────────────────────────────────


def _anchor_cells(obj: Any) -> tuple[int, int, int, int] | None:
    """(row1, col1, row2, col2) covered by a drawing anchor (1-based)."""
    a = getattr(obj, "anchor", None)
    if a is None:
        return None
    if isinstance(a, str):
        try:
            from _a1 import parse_cell

            r, c = parse_cell(a)
        except ValueError:
            return None
        w = float(getattr(obj, "width", 15) or 15)
        h = float(getattr(obj, "height", 7.5) or 7.5)
        if obj.__class__.__name__ == "Image":
            return r, c, r + int(h / 20) + 1, c + int(w / 64) + 1
        return r, c, r + int(h / 0.53) + 1, c + int(w / 1.7) + 1
    fr = getattr(a, "_from", None)
    to = getattr(a, "to", None)
    if fr is None:
        return None
    if to is not None:
        return fr.row + 1, fr.col + 1, to.row + 1, to.col + 1
    ext = getattr(a, "ext", None)
    if ext is not None:
        return fr.row + 1, fr.col + 1, fr.row + 1 + int(ext.height / EMU_PER_PX / 20) + 1, fr.col + 1 + int(ext.width / EMU_PER_PX / 64) + 1
    return fr.row + 1, fr.col + 1, fr.row + 15, fr.col + 8


def _anchor_box(obj: Any, zoom: float, default_w: float, default_h: float) -> tuple[int, int, int, int, int, int] | None:
    """(row, col, dx, dy, width px, height px) of a drawing."""
    a = getattr(obj, "anchor", None)
    if a is None:
        return None
    if isinstance(a, str):
        from _a1 import parse_cell

        try:
            r, c = parse_cell(a)
        except ValueError:
            return None
        return r, c, 0, 0, round(default_w * zoom), round(default_h * zoom)
    fr = getattr(a, "_from", None)
    if fr is None:
        pos, ext = getattr(a, "pos", None), getattr(a, "ext", None)
        if pos is None and ext is None:
            return None
        x = round((getattr(pos, "x", 0) or 0) / EMU_PER_PX * zoom)
        y = round((getattr(pos, "y", 0) or 0) / EMU_PER_PX * zoom)
        w = round((getattr(ext, "width", 0) or 0) / EMU_PER_PX * zoom) or round(default_w * zoom)
        h = round((getattr(ext, "height", 0) or 0) / EMU_PER_PX * zoom) or round(default_h * zoom)
        return 1, 1, x, y, w, h
    dx, dy = round(fr.colOff / EMU_PER_PX * zoom), round(fr.rowOff / EMU_PER_PX * zoom)
    ext = getattr(a, "ext", None)
    to = getattr(a, "to", None)
    if ext is not None and getattr(ext, "width", None):
        return fr.row + 1, fr.col + 1, dx, dy, round(ext.width / EMU_PER_PX * zoom), round(ext.height / EMU_PER_PX * zoom)
    if to is not None:
        return fr.row + 1, fr.col + 1, dx, dy, -(to.col + 1), -(to.row + 1)  # resolved against the grid later
    return fr.row + 1, fr.col + 1, dx, dy, round(default_w * zoom), round(default_h * zoom)


def _add_image(v: View, img: Any, zoom: float) -> None:
    from PIL import Image

    data = img._data()
    im = Image.open(io.BytesIO(data))
    im.load()
    box = _anchor_box(img, zoom, float(img.width or im.width), float(img.height or im.height))
    if box is None:
        return
    r, c, dx, dy, w, h = box
    if w > 0 and h > 0:
        pass
    v.images.append((im, r, c, dx, dy, w, h))


def _ref_values(vals: Values, ref_text: str | None, default_sheet: str) -> list[Any]:
    if not ref_text:
        return []
    sheet, rest = split_sheet(ref_text.replace("$", ""))
    sheet = sheet or default_sheet
    try:
        r1, c1, r2, c2 = parse_range(rest)
    except ValueError:
        return []
    if (r2 - r1 + 1) * (c2 - c1 + 1) > 20000:
        r2 = r1 + max(0, 20000 // (c2 - c1 + 1) - 1)
    try:
        vals.load(sheet)
    except (KeyError, ValueError):
        return []
    if vals.missing(sheet):
        vals.recalc()
    out = []
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            val = vals.get(sheet, r, c)
            if hasattr(val, "code"):
                val = None
            out.append(val)
    while out and out[-1] is None:
        out.pop()
    return out


def _rich_text(t: Any) -> str:
    if t is None:
        return ""
    if isinstance(t, str):
        return t
    try:
        tx = t.tx
        if tx is not None and tx.rich is not None:
            paras = []
            for p in tx.rich.p:
                text = "".join(run.t or "" for run in p.r or []).strip()
                if text:
                    paras.append(text)
            return " · ".join(paras) if len(paras) > 1 else "".join(paras)
        if tx is not None and tx.strRef is not None:
            return ""
    except AttributeError:
        pass
    return ""


def _series_ref(holder: Any) -> str | None:
    if holder is None:
        return None
    for kind in ("numRef", "strRef"):
        r = getattr(holder, kind, None)
        if r is not None and getattr(r, "f", None):
            return r.f
    return None


def _series_cache(holder: Any) -> list[Any]:
    """The values the file caches with a series reference (numCache/strCache, or the first level of a multi-level
    category cache). Pivot charts point at pivot-table fields ("pt@data 0"), so these are their only data."""
    if holder is None:
        return []
    for kind, cname in (("numRef", "numCache"), ("strRef", "strCache")):
        ref = getattr(holder, kind, None)
        cache = getattr(ref, cname, None) if ref is not None else None
        if cache is not None and cache.pt:
            return _cache_points(cache, kind == "numRef")
    ml = getattr(holder, "multiLvlStrRef", None)
    cache = getattr(ml, "multiLvlStrCache", None) if ml is not None else None
    if cache is not None and cache.lvl:
        lvl = cache.lvl[0]
        n = cache.ptCount or 0
        out: list[Any] = [None] * max(n, max((p.idx for p in lvl.pt), default=-1) + 1)
        for p in lvl.pt:
            out[p.idx] = p.v
        return out
    return []


def _cache_points(cache: Any, numeric: bool) -> list[Any]:
    n = max(cache.ptCount or 0, max((p.idx for p in cache.pt), default=-1) + 1)
    out: list[Any] = [None] * min(n, 20000)
    for p in cache.pt:
        if 0 <= p.idx < len(out):
            try:
                out[p.idx] = float(p.v) if numeric else p.v
            except (TypeError, ValueError):
                out[p.idx] = None
    return out


def _series_data(vals: Values, holder: Any, sheet: str) -> list[Any]:
    """A series' values from the cells it references, else from the file's cache, else its literal values."""
    got = _ref_values(vals, _series_ref(holder), sheet) if _looks_like_ref(_series_ref(holder)) else []
    if any(x is not None for x in got):
        return got
    return _series_cache(holder) or _series_literal(holder)


def _looks_like_ref(f: str | None) -> bool:
    return bool(f) and "@" not in str(f) and "!" in str(f) or bool(f and re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?", str(f)))


def _series_color(s: Any) -> str | None:
    """An explicit RGB fill (bars, areas, pie) or line colour set on a series, as #RRGGBB."""
    sp = getattr(s, "spPr", None)
    if sp is None:
        return None
    for choice in (getattr(sp, "solidFill", None), getattr(getattr(sp, "ln", None), "solidFill", None)):
        rgb = getattr(choice, "srgbClr", None) if choice is not None else None
        val = getattr(rgb, "val", rgb) if rgb is not None else None
        if isinstance(val, str) and re.fullmatch(r"[0-9A-Fa-f]{6}", val):
            return "#" + val.upper()
    return None


def _series_literal(holder: Any) -> list[Any]:
    if holder is None:
        return []
    for kind in ("numLit", "strLit"):
        lit = getattr(holder, kind, None)
        if lit is not None:
            try:
                return [float(p.v) if kind == "numLit" else p.v for p in lit.pt]
            except (AttributeError, ValueError):
                return []
    return []


def _add_chart(v: View, ch: Any, vals: Values, sheet: str, zoom: float) -> None:
    kind = ch.__class__.__name__.replace("Chart", "").lower()
    info: dict[str, Any] = {"kind": kind, "title": _rich_text(ch.title) if ch.title is not None else "", "series": []}
    if kind == "bar":
        info["kind"] = "column" if getattr(ch, "type", "col") == "col" else "bar"
        info["grouping"] = getattr(ch, "grouping", "clustered")
    elif kind in ("line", "area"):
        info["grouping"] = getattr(ch, "grouping", "standard")
    try:
        info["x_title"] = _rich_text(ch.x_axis.title) if getattr(ch, "x_axis", None) is not None and ch.x_axis.title is not None else ""
        info["y_title"] = _rich_text(ch.y_axis.title) if getattr(ch, "y_axis", None) is not None and ch.y_axis.title is not None else ""
    except AttributeError:
        pass
    info["legend"] = ch.legend is not None
    if ch.legend is not None:
        info["legend_pos"] = getattr(ch.legend, "position", None) or "r"
    for s in ch.series:
        name = ""
        if s.tx is not None:
            if s.tx.strRef is not None and (s.tx.strRef.f or s.tx.strRef.strCache is not None):
                got = _series_data(vals, s.tx, sheet)
                name = str(got[0]) if got and got[0] is not None else ""
            elif getattr(s.tx, "v", None):
                name = s.tx.v
        if kind in ("scatter", "bubble"):
            xs = _series_data(vals, s.xVal, sheet)
            ys = _series_data(vals, s.yVal, sheet)
            info["series"].append({"name": name, "x": xs, "y": ys})
        else:
            got = _series_data(vals, s.val, sheet)
            cats = _series_data(vals, s.cat, sheet)
            info["series"].append({"name": name, "values": got, "categories": cats})
        col = _series_color(s)
        if col:
            info["series"][-1]["color"] = col
    box = _anchor_box(ch, zoom, float(ch.width or 15) / 2.54 * 96, float(ch.height or 7.5) / 2.54 * 96)
    if box is None:
        return
    info["box"] = box
    to = getattr(getattr(ch, "anchor", None), "to", None)
    if box[4] < 0 and to is not None:
        info["to_off"] = (round((to.colOff or 0) / EMU_PER_PX * zoom), round((to.rowOff or 0) / EMU_PER_PX * zoom))
    v.charts.append(info)


# ── drawing ─────────────────────────────────────────────────────────────


def _font_for(cell: Cell, zoom: float) -> tuple[Any, bool]:
    import _fonts

    name, pt, bold, italic, _u, _s = cell.font
    return _fonts.load(name, round(pt * PX_PER_PT * zoom), bold, italic)


def _text_w(font: Any, s: str) -> float:
    try:
        return font.getlength(s)
    except AttributeError:
        return font.getsize(s)[0]


def _wrap_lines(text: str, font: Any, width: float) -> list[str]:
    lines: list[str] = []
    for para in text.split("\n"):
        words = para.split(" ")
        cur = ""
        for w in words:
            cand = w if not cur else cur + " " + w
            if _text_w(font, cand) <= width or not cur:
                if not cur and _text_w(font, w) > width:
                    # break a long word
                    piece = ""
                    for ch in w:
                        if _text_w(font, piece + ch) > width and piece:
                            lines.append(piece)
                            piece = ch
                        else:
                            piece += ch
                    cur = piece
                else:
                    cur = cand
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def _dash(draw: Any, p1: tuple[float, float], p2: tuple[float, float], color: Any, width: int, pattern: tuple[int, int]) -> None:
    (x1, y1), (x2, y2) = p1, p2
    length = math.hypot(x2 - x1, y2 - y1)
    if length == 0:
        return
    on, off = pattern
    pos = 0.0
    while pos < length:
        end = min(pos + on, length)
        a = pos / length
        b = end / length
        draw.line([(x1 + (x2 - x1) * a, y1 + (y2 - y1) * a), (x1 + (x2 - x1) * b, y1 + (y2 - y1) * b)], fill=color, width=width)
        pos = end + off


def _border(draw: Any, style: str, color: Any, p1: tuple[int, int], p2: tuple[int, int], zoom: float, horizontal: bool) -> None:
    w = {"hair": 1, "thin": 1, "dotted": 1, "dashed": 1, "dashDot": 1, "dashDotDot": 1, "medium": 2, "mediumDashed": 2, "mediumDashDot": 2, "mediumDashDotDot": 2, "slantDashDot": 2, "thick": 3, "double": 1}.get(style, 1)
    w = max(1, round(w * zoom))
    if style == "double":
        off = max(1, round(zoom))
        if horizontal:
            draw.line([(p1[0], p1[1] - off), (p2[0], p2[1] - off)], fill=color, width=1)
            draw.line([(p1[0], p1[1] + off), (p2[0], p2[1] + off)], fill=color, width=1)
        else:
            draw.line([(p1[0] - off, p1[1]), (p2[0] - off, p2[1])], fill=color, width=1)
            draw.line([(p1[0] + off, p1[1]), (p2[0] + off, p2[1])], fill=color, width=1)
        return
    if style in ("dotted", "hair"):
        _dash(draw, p1, p2, color, w, (1, 2) if style == "dotted" else (1, 1))
    elif "dash" in style.lower():
        _dash(draw, p1, p2, color, w, (4, 2) if "Dot" not in style else (5, 2))
    else:
        draw.line([p1, p2], fill=color, width=w)


def draw_page(v: View, rows: list[int], cols: list[int], title: str | None = None, headers: bool = True) -> Any:
    """Draws the given rows and columns of a view; returns a PIL image."""
    from PIL import Image, ImageDraw

    import _fonts

    zoom = v.zoom
    hdr_font, _ = _fonts.load("Calibri", round(10 * PX_PER_PT * zoom), False, False)
    rh_w = round(max(34 * zoom, _text_w(hdr_font, str(rows[-1] if rows else 1)) + 14 * zoom)) if headers else 0
    ch_h = round(20 * zoom) if headers else 0
    title_h = round(26 * zoom) if title else 0
    xs: dict[int, int] = {}
    x = rh_w
    for c in cols:
        xs[c] = x
        x += v.col_w.get(c, col_px(None, zoom))
    total_w = x + 1
    ys: dict[int, int] = {}
    y = title_h + ch_h
    for r in rows:
        ys[r] = y
        y += v.row_h.get(r, row_px(15, zoom))
    total_h = y + 1
    img = Image.new("RGB", (max(total_w, 40), max(total_h, 30)), "white")
    d = ImageDraw.Draw(img)
    row_set, col_set = set(rows), set(cols)
    last_x = {c: xs[c] + v.col_w.get(c, 0) for c in cols}
    last_y = {r: ys[r] + v.row_h.get(r, 0) for r in rows}

    def rect_of(r1: int, c1: int, r2: int, c2: int) -> tuple[int, int, int, int] | None:
        rr = [r for r in rows if r1 <= r <= r2]
        cc = [c for c in cols if c1 <= c <= c2]
        if not rr or not cc:
            return None
        return xs[cc[0]], ys[rr[0]], last_x[cc[-1]], last_y[rr[-1]]

    merged_cover: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for m in v.merges:
        for r in range(m[0], m[2] + 1):
            for c in range(m[1], m[3] + 1):
                merged_cover[(r, c)] = m
    # gridlines
    if v.gridlines:
        for c in cols:
            d.line([(last_x[c], title_h + ch_h), (last_x[c], total_h - 1)], fill=GRID)
        for r in rows:
            d.line([(rh_w, last_y[r]), (total_w - 1, last_y[r])], fill=GRID)
    # fills (merged rects first so their interior gridlines disappear)
    for m in v.merges:
        box = rect_of(*m)
        if box is None:
            continue
        cell = v.cells.get((m[0], m[1]))
        fill = cell.fill if cell and cell.fill else (255, 255, 255)
        d.rectangle([box[0] + 1, box[1] + 1, box[2] - 1, box[3] - 1], fill=fill)
    for (r, c), cell in v.cells.items():
        if r not in row_set or c not in col_set or (r, c) in merged_cover:
            continue
        if cell.fill:
            d.rectangle([xs[c], ys[r], last_x[c], last_y[r]], fill=cell.fill)
            # fills cover the gridlines of their own cell but not the neighbours'
        if cell.bar:
            axis, pos, col = cell.bar
            inner = last_x[c] - xs[c] - 4
            xa, xb = xs[c] + 2 + inner * axis, xs[c] + 2 + inner * pos
            light = tuple(round(x * 0.35 + 255 * 0.65) for x in col)
            d.rectangle([min(xa, xb), ys[r] + 3, max(xa, xb), last_y[r] - 3], fill=light, outline=col)
            if axis > 0:
                _dash(d, (xs[c] + 2 + inner * axis, ys[r] + 1), (xs[c] + 2 + inner * axis, last_y[r] - 1), (0, 0, 0), 1, (1, 2))
    # text
    occupied = {(r, c) for (r, c), cell in v.cells.items() if cell.text}
    for (r, c), cell in sorted(v.cells.items()):
        if not cell.text or r not in row_set or c not in col_set:
            continue
        m = merged_cover.get((r, c))
        if m is not None and (r, c) != (m[0], m[1]):
            continue
        box = rect_of(*m) if m else (xs[c], ys[r], last_x[c], last_y[r])
        if box is None:
            continue
        _draw_text(img, d, v, cell, box, r, c, cols, xs, last_x, occupied, merged_cover)
        if cell.icon:
            cy = (box[1] + box[3]) // 2
            rad = max(3, round(4 * zoom))
            d.ellipse([box[0] + 3, cy - rad, box[0] + 3 + 2 * rad, cy + rad], fill=cell.icon)
    # icons in cells without text
    for (r, c), cell in v.cells.items():
        if cell.icon and not cell.text and r in row_set and c in col_set:
            cy = (ys[r] + last_y[r]) // 2
            rad = max(3, round(4 * zoom))
            d.ellipse([xs[c] + 3, cy - rad, xs[c] + 3 + 2 * rad, cy + rad], fill=cell.icon)
    # borders
    for (r, c), cell in v.cells.items():
        if not cell.borders or r not in row_set or c not in col_set:
            continue
        m = merged_cover.get((r, c))
        x0, y0, x1, y1 = xs[c], ys[r], last_x[c], last_y[r]
        for side, (style, color) in cell.borders.items():
            if m is not None:
                if side == "left" and c != m[1] or side == "right" and c != m[3] or side == "top" and r != m[0] or side == "bottom" and r != m[2]:
                    continue
            if side == "top":
                _border(d, style, color, (x0, y0), (x1, y0), zoom, True)
            elif side == "bottom":
                _border(d, style, color, (x0, y1), (x1, y1), zoom, True)
            elif side == "left":
                _border(d, style, color, (x0, y0), (x0, y1), zoom, False)
            else:
                _border(d, style, color, (x1, y0), (x1, y1), zoom, False)
    # comments
    for (r, c), cell in v.cells.items():
        if cell.comment and r in row_set and c in col_set:
            s = max(4, round(6 * zoom))
            x1 = last_x[c]
            d.polygon([(x1 - s, ys[r] + 1), (x1, ys[r] + 1), (x1, ys[r] + 1 + s)], fill=(220, 0, 0))
    # images and charts: placed from the whole view's geometry, so a drawing cut by a page edge is drawn on every
    # page it overlaps, each showing its own part of the same picture
    geo = _geometry(v)
    ox, oy = rh_w - geo[0][cols[0]], title_h + ch_h - geo[2][rows[0]]
    drawings = [((r, c, dx, dy, w, h), None, lambda W, H, im=im: im.convert("RGBA").resize((max(1, W), max(1, H)))) for im, r, c, dx, dy, w, h in v.images]
    drawings += [(ch["box"], ch.get("to_off"), lambda W, H, ch=ch: draw_chart(ch, W, H, zoom)) for ch in v.charts]
    for box, to_off, make in drawings:
        rect = drawing_rect(v, box, to_off, geo)
        if rect is None:
            continue
        x0, y0, x1, y1 = rect[0] + ox, rect[1] + oy, rect[2] + ox, rect[3] + oy
        if x1 <= rh_w or y1 <= title_h + ch_h or x0 >= total_w or y0 >= total_h:
            continue
        tile = make(x1 - x0, y1 - y0)
        if tile is None:
            continue
        img.paste(tile, (x0, y0), tile if tile.mode == "RGBA" else None)
    # frozen panes
    fr, fc = v.freeze
    if fr and fr in row_set and fr + 1 in row_set:
        d.line([(rh_w, last_y[fr]), (total_w - 1, last_y[fr])], fill=(110, 110, 110), width=max(1, round(zoom)))
    if fc and fc in col_set and fc + 1 in col_set:
        d.line([(last_x[fc], title_h + ch_h), (last_x[fc], total_h - 1)], fill=(110, 110, 110), width=max(1, round(zoom)))
    # headers
    if headers:
        d.rectangle([0, title_h, total_w - 1, title_h + ch_h - 1], fill=HEADER_BG)
        d.rectangle([0, title_h, rh_w - 1, total_h - 1], fill=HEADER_BG)
        for c in cols:
            label = col_letter(c)
            tw = _text_w(hdr_font, label)
            cx = xs[c] + (v.col_w.get(c, 0) - tw) / 2
            d.text((cx, title_h + (ch_h - hdr_font.size) / 2 - 1), label, fill=HEADER_TEXT, font=hdr_font)
            d.line([(last_x[c], title_h), (last_x[c], title_h + ch_h - 1)], fill=HEADER_LINE)
        for r in rows:
            label = str(r)
            tw = _text_w(hdr_font, label)
            h = v.row_h.get(r, 0)
            if h >= hdr_font.size * 0.8:
                d.text(((rh_w - tw) / 2, ys[r] + (h - hdr_font.size) / 2 - 1), label, fill=HEADER_TEXT, font=hdr_font)
            d.line([(0, last_y[r]), (rh_w - 1, last_y[r])], fill=HEADER_LINE)
        d.line([(0, title_h + ch_h - 1), (total_w - 1, title_h + ch_h - 1)], fill=HEADER_LINE)
        d.line([(rh_w - 1, title_h), (rh_w - 1, total_h - 1)], fill=HEADER_LINE)
    if title:
        tf, _ = _fonts.load("Calibri", round(11 * PX_PER_PT * zoom), True, False)
        d.rectangle([0, 0, total_w - 1, title_h - 1], fill=(33, 115, 70))
        d.text((8 * zoom, (title_h - tf.size) / 2 - 1), title, fill="white", font=tf)
    return img


def _geometry(v: View) -> tuple[dict[int, int], int, dict[int, int], int]:
    """The view laid out as one picture: x of each column's left edge and y of each row's top (from 0), with the
    total width and height. draw_page draws a slice of it."""
    gx: dict[int, int] = {}
    x = 0
    for c in v.cols:
        gx[c] = x
        x += v.col_w.get(c, col_px(None, v.zoom))
    gy: dict[int, int] = {}
    y = 0
    for r in v.rows:
        gy[r] = y
        y += v.row_h.get(r, row_px(15, v.zoom))
    return gx, x, gy, y


def _edge(pos: dict[int, int], total: int, keys: list[int], k: int, default: int) -> int:
    """The left (top) edge of column (row) k in the view's picture, also for one hidden or beyond the view."""
    if k in pos:
        return pos[k]
    if not keys:
        return 0
    if k > keys[-1]:
        return total + (k - keys[-1] - 1) * default
    if k < keys[0]:
        return -(keys[0] - k) * default
    return pos[next(x for x in keys if x > k)]


def drawing_rect(v: View, box: tuple[int, int, int, int, int, int], to_off: tuple[int, int] | None = None, geo: Any = None) -> tuple[int, int, int, int] | None:
    """(x0, y0, x1, y1) of a drawing anchored at box = (row, col, dx, dy, width, height) in the view's picture;
    a negative width and height name the cell of the far corner (two-cell anchors), to_off its offsets."""
    gx, tw, gy, th = geo or _geometry(v)
    r, c, dx, dy, w, h = box
    dw, dh = col_px(None, v.zoom), row_px(15, v.zoom)
    x0 = _edge(gx, tw, v.cols, c, dw) + dx
    y0 = _edge(gy, th, v.rows, r, dh) + dy
    if w < 0 and h < 0:
        ox, oy = to_off or (0, 0)
        x1 = max(x0 + 40, _edge(gx, tw, v.cols, -w, dw) + ox)
        y1 = max(y0 + 30, _edge(gy, th, v.rows, -h, dh) + oy)
        return x0, y0, x1, y1
    if w <= 0 or h <= 0:
        return None
    return x0, y0, x0 + w, y0 + h


def _index_span(pos: dict[int, int], keys: list[int], sizes: dict[int, int], lo: int, hi: int) -> tuple[int, int] | None:
    """Indexes (into keys) of the first and last column/row that the pixel span lo..hi overlaps."""
    hit = [i for i, k in enumerate(keys) if pos[k] < hi and pos[k] + sizes[k] > lo]
    return (hit[0], hit[-1]) if hit else None


def drawing_spans(v: View) -> list[tuple[tuple[int, int], tuple[int, int]] | None]:
    """For each chart, then each image: ((first, last) row index, (first, last) column index) into v.rows and
    v.cols, or None when it lies outside the view."""
    geo = _geometry(v)
    gx, _tw, gy, _th = geo
    cw = {c: v.col_w.get(c, col_px(None, v.zoom)) for c in v.cols}
    rh = {r: v.row_h.get(r, row_px(15, v.zoom)) for r in v.rows}
    boxes = [(ch["box"], ch.get("to_off")) for ch in v.charts] + [((r, c, dx, dy, w, h), None) for _im, r, c, dx, dy, w, h in v.images]
    out: list[tuple[tuple[int, int], tuple[int, int]] | None] = []
    for box, to_off in boxes:
        rect = drawing_rect(v, box, to_off, geo)
        rs = _index_span(gy, v.rows, rh, rect[1], rect[3]) if rect else None
        cs = _index_span(gx, v.cols, cw, rect[0], rect[2]) if rect else None
        out.append((rs, cs) if rs and cs else None)
    return out


def chart_pages(v: View, pages: list[tuple[list[int], list[int]]]) -> list[list[int]]:
    """For each chart of the view, the indexes of the pages (from paginate) that show part of it."""
    spans = drawing_spans(v)[: len(v.charts)]
    out = []
    for sp in spans:
        hold = []
        if sp is not None:
            (r0, r1), (c0, c1) = sp
            rows, cols = set(v.rows[r0 : r1 + 1]), set(v.cols[c0 : c1 + 1])
            hold = [i for i, (pr, pc) in enumerate(pages) if rows & set(pr) and cols & set(pc)]
        out.append(hold)
    return out


def _draw_text(img: Any, d: Any, v: View, cell: Cell, box: tuple[int, int, int, int], r: int, c: int, cols: list[int], xs: dict[int, int], last_x: dict[int, int], occupied: set[tuple[int, int]], merged_cover: dict[tuple[int, int], Any]) -> None:
    zoom = v.zoom
    font, fake_bold = _font_for(cell, zoom)
    x0, y0, x1, y1 = box
    pad = max(2, round(3 * zoom))
    width = x1 - x0 - 2 * pad
    text = cell.text
    h = cell.h
    if h == "general":
        h = "right" if cell.numeric else ("center" if (text in ("TRUE", "FALSE") or text.startswith("#")) and not isinstance(cell.value, str) else "left")
    if h in ("centerContinuous", "distributed"):
        h = "center"
    if h == "fill":
        if text:
            reps = max(1, int(width // max(1, _text_w(font, text))))
            text = text * reps
        h = "left"
    indent_px = round(cell.indent * 9 * zoom)
    if cell.rotation in (90, 180) or cell.rotation == 255:
        _draw_rotated(img, cell, text, font, box, zoom)
        return
    lines = [text]
    if cell.wrap or h == "justify":
        lines = _wrap_lines(text, font, max(4, width - indent_px))
    elif "\n" in text:
        lines = [text.split("\n")[0]]
    # numbers that do not fit show ####
    if cell.numeric and not cell.wrap and _text_w(font, text) > width + 1 and not cell.shrink:
        n = max(1, int(width // max(1, _text_w(font, "#"))))
        lines = ["#" * n]
    # overflow into empty neighbours for plain text
    left_lim, right_lim = x0 + pad, x1 - pad
    if not cell.wrap and not cell.numeric and (r, c) not in merged_cover:
        need = _text_w(font, lines[0]) + indent_px
        if need > width:
            pos = cols.index(c)
            if h in ("left", "center"):
                k = pos
                while need > right_lim - left_lim and k + 1 < len(cols) and (r, cols[k + 1]) not in occupied and (r, cols[k + 1]) not in merged_cover:
                    k += 1
                    right_lim = last_x[cols[k]] - pad
            if h in ("right", "center"):
                k = pos
                while need > right_lim - left_lim and k - 1 >= 0 and (r, cols[k - 1]) not in occupied and (r, cols[k - 1]) not in merged_cover:
                    k -= 1
                    left_lim = xs[cols[k]] + pad
    if cell.shrink and not cell.wrap:
        avail = right_lim - left_lim
        tw = _text_w(font, lines[0])
        if tw > avail and tw > 0:
            import _fonts

            name, pt, bold, italic, _u, _s = cell.font
            font, fake_bold = _fonts.load(name, max(6, round(pt * PX_PER_PT * zoom * avail / tw)), bold, italic)
    try:
        asc, desc = font.getmetrics()
    except AttributeError:
        asc, desc = font.size, 2
    line_h = asc + desc
    block_h = line_h * len(lines)
    if cell.v == "top":
        ty = y0 + pad // 2
    elif cell.v in ("center", "distributed", "justify"):
        ty = y0 + (y1 - y0 - block_h) / 2
    else:
        ty = y1 - block_h - max(1, round(2 * zoom))
    for line in lines:
        # clip to the allowed width
        avail = right_lim - left_lim - indent_px
        if _text_w(font, line) > avail and not cell.numeric:
            lo, hi = 0, len(line)
            while lo < hi:
                mid = (lo + hi + 1) // 2
                if _text_w(font, line[:mid]) <= avail:
                    lo = mid
                else:
                    hi = mid - 1
            line = line[:lo]
        tw = _text_w(font, line)
        if h == "right":
            tx = right_lim - tw - indent_px
        elif h == "center":
            tx = (left_lim + right_lim - tw) / 2
        else:
            tx = left_lim + indent_px
        if ty + line_h > y0 - 1 or len(lines) == 1:
            kw: dict[str, Any] = {"fill": cell.color, "font": font}
            if fake_bold:
                kw["stroke_width"] = max(1, round(0.4 * zoom))
                kw["stroke_fill"] = cell.color
            d.text((tx, ty), line, **kw)
            if cell.font[4]:
                uy = ty + asc + max(1, round(1 * zoom))
                d.line([(tx, uy), (tx + tw, uy)], fill=cell.color, width=max(1, round(zoom)))
                if cell.font[4] == "double":
                    d.line([(tx, uy + 2), (tx + tw, uy + 2)], fill=cell.color, width=1)
            if cell.font[5]:
                sy = ty + asc * 0.62
                d.line([(tx, sy), (tx + tw, sy)], fill=cell.color, width=max(1, round(zoom)))
        ty += line_h
        if ty > y1 + line_h:
            break


def _draw_rotated(img: Any, cell: Cell, text: str, font: Any, box: tuple[int, int, int, int], zoom: float) -> None:
    from PIL import Image, ImageDraw

    x0, y0, x1, y1 = box
    if cell.rotation == 255:
        text = "\n".join(text)
    tw = int(max((_text_w(font, ln) for ln in text.split("\n")), default=1)) + 4
    try:
        asc, desc = font.getmetrics()
    except AttributeError:
        asc, desc = font.size, 2
    th = (asc + desc) * len(text.split("\n")) + 2
    layer = Image.new("RGBA", (max(1, tw), max(1, th)), (0, 0, 0, 0))
    ImageDraw.Draw(layer).multiline_text((2, 0), text, fill=cell.color, font=font, spacing=0)
    if cell.rotation == 90:
        layer = layer.rotate(90, expand=True)
    elif cell.rotation == 180:
        layer = layer.rotate(-90, expand=True)
    w, h = layer.size
    px = int(x0 + (x1 - x0 - w) / 2)
    py = int(y1 - h - 2) if cell.rotation != 180 else int(y0 + 2)
    crop = layer.crop((0, max(0, y0 - py), w, min(h, y1 - py)))
    img.paste(crop, (px, max(py, y0)), crop)


# ── charts ──────────────────────────────────────────────────────────────


def _nice(lo: float, hi: float, n: int = 5) -> tuple[float, float, float]:
    if hi == lo:
        hi = lo + 1 if lo >= 0 else lo + abs(lo) * 0.1 + 1
    span = hi - lo
    raw = span / n
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    return math.floor(lo / step) * step, math.ceil(hi / step) * step, step


def _num(x: Any) -> float | None:
    if isinstance(x, bool):
        return float(x)
    if isinstance(x, (int, float)):
        return float(x) if math.isfinite(x) else None
    if isinstance(x, (_dt.date, _dt.datetime)):
        return date_to_serial(x)
    return None


def _label(x: Any) -> str:
    from _book import text_value

    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return text_value(x) if x is not None else ""


def _tick(v: float, step: float) -> str:
    if abs(v) >= 1e6:
        return f"{v / 1e6:g}M"
    if abs(v) >= 1e4:
        return f"{v / 1e3:g}K"
    dec = 0  # as many decimals as the step has: 2.5 steps read 2.5, 5.0, 7.5 (not 2, 5, 8)
    while dec < 6 and abs(step * 10**dec - round(step * 10**dec)) > 1e-9 * max(1.0, step * 10**dec):
        dec += 1
    return f"{v:,.{dec}f}"


def draw_chart(ch: dict[str, Any], W: int, H: int, zoom: float) -> Any:
    """A simple, faithful drawing of a chart from its data: column/bar (clustered, stacked), line, area, pie/doughnut,
    scatter and radar (as lines)."""
    from PIL import Image, ImageDraw

    import _fonts

    W, H = max(80, W), max(60, H)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W - 1, H - 1], outline=(191, 191, 191))
    f_small, _ = _fonts.load("Calibri", max(8, round(9 * PX_PER_PT * zoom)), False, False)
    f_title, _ = _fonts.load("Calibri", max(9, round(13 * PX_PER_PT * zoom)), False, False)
    kind = ch["kind"]
    series = [s for s in ch["series"] if s.get("values") or s.get("y")]
    top = round(8 * zoom)
    if ch.get("title"):
        tw = _text_w(f_title, ch["title"])
        d.text(((W - tw) / 2, top), ch["title"], fill=(64, 64, 64), font=f_title)
        top += f_title.size + round(8 * zoom)
    if not series:
        d.text((10, H / 2), "(no data)", fill=(128, 128, 128), font=f_small)
        return img
    colors = [hex_rgb(c) for c in PALETTE]
    if kind not in ("pie", "doughnut") and any(s.get("color") for s in series):
        colors = [hex_rgb(s["color"]) if s.get("color") else colors[i % len(colors)] for i, s in enumerate(series)]
    names = [s.get("name") or f"Series{i + 1}" for i, s in enumerate(series)]
    if kind in ("pie", "doughnut"):
        cats = [_label(x) for x in (series[0].get("categories") or [])] or [str(i + 1) for i in range(len(series[0]["values"]))]
        legend_items = cats
    else:
        legend_items = names if (len(series) > 1 or ch.get("legend")) else []
    if ch.get("legend") is False:  # the chart has no legend element: Excel shows none
        legend_items = []
    plot = [round(10 * zoom), top, W - round(12 * zoom), H - round(10 * zoom)]
    if legend_items:
        img.info["legend_box"] = _legend(d, [str(x)[:30] for x in legend_items[:20]], colors, ch.get("legend_pos") or "r", plot, W, H, top, f_small, zoom)
    if kind in ("pie", "doughnut"):
        vals = [max(0.0, _num(x) or 0.0) for x in series[0]["values"]]
        total = sum(vals) or 1.0
        size = min(plot[2] - plot[0], plot[3] - plot[1]) - 6
        cx, cy = (plot[0] + plot[2]) / 2, (plot[1] + plot[3]) / 2
        bb = [cx - size / 2, cy - size / 2, cx + size / 2, cy + size / 2]
        ang = -90.0
        for i, val in enumerate(vals):
            sweep = 360 * val / total
            if sweep > 0:
                d.pieslice(bb, ang, ang + sweep, fill=colors[i % len(colors)], outline="white")
                mid = math.radians(ang + sweep / 2)
                if sweep > 12:
                    lab = f"{100 * val / total:.0f}%"
                    lx = cx + math.cos(mid) * size * 0.33 - _text_w(f_small, lab) / 2
                    ly = cy + math.sin(mid) * size * 0.33 - f_small.size / 2
                    d.text((lx, ly), lab, fill="white", font=f_small)
            ang += sweep
        if kind == "doughnut":
            hole = size * 0.25
            d.ellipse([cx - hole, cy - hole, cx + hole, cy + hole], fill="white")
        return img
    if kind in ("scatter", "bubble"):
        pts = [[(_num(x), _num(y)) for x, y in zip(s.get("x") or range(1, len(s["y"]) + 1), s["y"])] for s in series]
        allx = [p[0] for ps in pts for p in ps if p[0] is not None and p[1] is not None]
        ally = [p[1] for ps in pts for p in ps if p[0] is not None and p[1] is not None]
        if not allx:
            return img
        xlo, xhi, xstep = _nice(min(allx), max(allx))
        ylo, yhi, ystep = _nice(min(0, min(ally)), max(ally))
        area = _axes(d, plot, ylo, yhi, ystep, f_small, zoom, ch)
        ax0, ay0, ax1, ay1 = area
        k = xlo
        while k <= xhi + 1e-9:
            px = ax0 + (k - xlo) / (xhi - xlo) * (ax1 - ax0)
            lab = _tick(k, xstep)
            d.text((px - _text_w(f_small, lab) / 2, ay1 + 3), lab, fill=(89, 89, 89), font=f_small)
            k += xstep
        for i, ps in enumerate(pts):
            col = colors[i % len(colors)]
            r = max(2, round(3 * zoom))
            for x, y in ps:
                if x is None or y is None:
                    continue
                px = ax0 + (x - xlo) / (xhi - xlo) * (ax1 - ax0)
                py = ay1 - (y - ylo) / (yhi - ylo) * (ay1 - ay0)
                d.ellipse([px - r, py - r, px + r, py + r], fill=col)
        return img
    cats = series[0].get("categories") or []
    n = max(len(s["values"]) for s in series)
    labels = [_label(cats[i]) if i < len(cats) else str(i + 1) for i in range(n)]
    data = [[_num(s["values"][i]) if i < len(s["values"]) else None for i in range(n)] for s in series]
    stacked = ch.get("grouping") in ("stacked", "percentStacked")
    percent = ch.get("grouping") == "percentStacked"
    if percent:
        totals = [sum(abs(row[i] or 0) for row in data) or 1 for i in range(n)]
        data = [[(row[i] or 0) / totals[i] for i in range(n)] for row in data]
    if stacked:
        pos = [sum(max(0, row[i] or 0) for row in data) for i in range(n)]
        neg = [sum(min(0, row[i] or 0) for row in data) for i in range(n)]
        lo, hi = min(neg + [0]), max(pos + [0])
    else:
        vals = [x for row in data for x in row if x is not None]
        lo, hi = min(vals + [0]), max(vals + [0])
    ylo, yhi, ystep = _nice(lo, hi)
    if kind == "bar":
        # horizontal bars: categories down the left
        lab_w = max((_text_w(f_small, lab) for lab in labels), default=0) + 6
        plot = [plot[0] + lab_w, plot[1], plot[2], plot[3] - f_small.size - 6]
        ax0, ay0, ax1, ay1 = plot
        d.line([(ax0, ay0), (ax0, ay1)], fill=(166, 166, 166))
        k = ylo
        while k <= yhi + 1e-9:
            px = ax0 + (k - ylo) / (yhi - ylo) * (ax1 - ax0)
            d.line([(px, ay0), (px, ay1)], fill=(217, 217, 217))
            lab = _tick(k, ystep) if not percent else f"{k * 100:.0f}%"
            d.text((px - _text_w(f_small, lab) / 2, ay1 + 3), lab, fill=(89, 89, 89), font=f_small)
            k += ystep
        band = (ay1 - ay0) / max(1, n)
        zero = ax0 + (0 - ylo) / (yhi - ylo) * (ax1 - ax0)
        for i in range(n):
            by = ay0 + band * i
            d.text((ax0 - lab_w, by + band / 2 - f_small.size / 2), labels[i][:24], fill=(89, 89, 89), font=f_small)
            acc_p = acc_n = 0.0
            m = len(series)
            for j, row in enumerate(data):
                val = row[i] or 0
                col = colors[j % len(colors)]
                if stacked:
                    start = acc_p if val >= 0 else acc_n
                    end = start + val
                    if val >= 0:
                        acc_p = end
                    else:
                        acc_n = end
                    x_a = ax0 + (start - ylo) / (yhi - ylo) * (ax1 - ax0)
                    x_b = ax0 + (end - ylo) / (yhi - ylo) * (ax1 - ax0)
                    d.rectangle([min(x_a, x_b), by + band * 0.2, max(x_a, x_b), by + band * 0.8], fill=col)
                else:
                    h_each = band * 0.7 / m
                    y_a = by + band * 0.15 + j * h_each
                    x_b = ax0 + (val - ylo) / (yhi - ylo) * (ax1 - ax0)
                    d.rectangle([min(zero, x_b), y_a, max(zero, x_b), y_a + h_each - 1], fill=col)
        return img
    area = _axes(d, plot, ylo, yhi, ystep, f_small, zoom, ch, percent)
    ax0, ay0, ax1, ay1 = area
    band = (ax1 - ax0) / max(1, n)
    step_lab = max(1, int(math.ceil(n / max(1, (ax1 - ax0) / max(20, max((_text_w(f_small, lab) for lab in labels), default=10) + 6)))))
    for i in range(0, n, step_lab):
        lab = labels[i][:14]
        cx = ax0 + band * (i + 0.5)
        d.text((cx - _text_w(f_small, lab) / 2, ay1 + 3), lab, fill=(89, 89, 89), font=f_small)

    def yp(val: float) -> float:
        return ay1 - (val - ylo) / (yhi - ylo) * (ay1 - ay0)

    zero = yp(0) if ylo <= 0 <= yhi else (ay1 if ylo > 0 else ay0)
    if kind == "column":
        m = len(series)
        for i in range(n):
            acc_p = acc_n = 0.0
            for j, row in enumerate(data):
                val = row[i]
                if val is None:
                    continue
                col = colors[j % len(colors)]
                if stacked:
                    start = acc_p if val >= 0 else acc_n
                    end = start + val
                    if val >= 0:
                        acc_p = end
                    else:
                        acc_n = end
                    x_a = ax0 + band * (i + 0.2)
                    d.rectangle([x_a, min(yp(start), yp(end)), x_a + band * 0.6, max(yp(start), yp(end))], fill=col)
                else:
                    w_each = band * 0.7 / m
                    x_a = ax0 + band * (i + 0.15) + j * w_each
                    d.rectangle([x_a, min(zero, yp(val)), x_a + w_each - 1, max(zero, yp(val))], fill=col)
        return img
    # line, area, radar (drawn as lines)
    acc = [0.0] * n
    for j, row in enumerate(data):
        col = colors[j % len(colors)]
        pts = []
        for i in range(n):
            val = row[i]
            if val is None:
                pts.append(None)
                continue
            if stacked:
                acc[i] += val
                val = acc[i]
            pts.append((ax0 + band * (i + 0.5), yp(val)))
        seg = [p for p in pts if p is not None]
        if kind == "area" and len(seg) >= 2:
            poly = [(seg[0][0], zero)] + seg + [(seg[-1][0], zero)]
            light = tuple(round(x * 0.75 + 255 * 0.25) for x in col)
            d.polygon(poly, fill=light)
        if len(seg) >= 2:
            d.line(seg, fill=col, width=max(2, round(2 * zoom)), joint="curve")
        r = max(2, round(2.5 * zoom))
        if kind == "line" and n <= 40:
            for p in seg:
                d.ellipse([p[0] - r, p[1] - r, p[0] + r, p[1] + r], fill=col)
    return img


def _legend(d: Any, items: list[str], colors: list[Any], pos: str, plot: list[int], W: int, H: int, top: int, font: Any, zoom: float) -> tuple[int, int, int, int]:
    """Draws the legend where the chart puts it (legendPos r, l, t, b or tr), shrinks the plot area (in place) to
    leave it room, and returns the legend's box (x0, y0, x1, y1)."""
    sw, gap, pad = 9 * zoom, 13 * zoom, round(10 * zoom)
    line_h = font.size + round(5 * zoom)

    def entry(i: int, name: str, x: float, y: float) -> None:
        d.rectangle([x, y + 3, x + sw, y + 3 + sw], fill=colors[i % len(colors)])
        d.text((x + gap, y), name, fill=(64, 64, 64), font=font)

    if pos in ("t", "b"):
        # one or more centred rows, like Excel's top and bottom legends
        widths = [gap + _text_w(font, name) + 14 * zoom for name in items]
        lines: list[list[int]] = [[]]
        used = 0.0
        for i, w in enumerate(widths):
            if lines[-1] and used + w > W - 2 * pad:
                lines.append([])
                used = 0.0
            lines[-1].append(i)
            used += w
        lines = lines[:3]
        block = len(lines) * line_h
        y = top + round(2 * zoom) if pos == "t" else H - pad - block
        if pos == "t":
            plot[1] = round(y + block + 4 * zoom)
        else:
            plot[3] = round(y - 4 * zoom)
        box_y = y
        widest = max(sum(widths[i] for i in line) for line in lines)
        for line in lines:
            x = (W - sum(widths[i] for i in line) + 14 * zoom) / 2
            for i in line:
                entry(i, items[i], x, y)
                x += widths[i]
            y += line_h
        return round((W - widest + 14 * zoom) / 2), round(box_y), round((W + widest) / 2), round(box_y + block)
    # a column on the right (r, tr) or the left (l)
    legend_w = min(W // 3, int(max(_text_w(font, name) for name in items) + 26 * zoom))
    block = len(items) * line_h
    y = top + 4 if pos == "tr" else max(top + 4, (top + H - pad) / 2 - block / 2)
    if pos == "l":
        x = pad
        plot[0] = pad + legend_w
    else:
        x = W - legend_w
        plot[2] = W - legend_w - round(12 * zoom)
    box = (round(x), round(y), round(x + legend_w), round(y + block))
    for i, name in enumerate(items):
        entry(i, name, x, y)
        y += line_h
    return box


def _axes(d: Any, plot: list[int], ylo: float, yhi: float, ystep: float, font: Any, zoom: float, ch: dict[str, Any], percent: bool = False) -> tuple[float, float, float, float]:
    labels = []
    k = ylo
    while k <= yhi + 1e-9:
        labels.append((k, _tick(k, ystep) if not percent else f"{k * 100:.0f}%"))
        k += ystep
    lab_w = max(_text_w(font, lab) for _, lab in labels) + 6
    ax0 = plot[0] + lab_w
    ay0 = plot[1] + font.size / 2
    ax1 = plot[2]
    ay1 = plot[3] - font.size - 6
    for val, lab in labels:
        py = ay1 - (val - ylo) / (yhi - ylo) * (ay1 - ay0)
        d.line([(ax0, py), (ax1, py)], fill=(217, 217, 217))
        d.text((ax0 - _text_w(font, lab) - 4, py - font.size / 2 - 1), lab, fill=(89, 89, 89), font=font)
    d.line([(ax0, ay1), (ax1, ay1)], fill=(166, 166, 166))
    return ax0, ay0, ax1, ay1


# ── pages ───────────────────────────────────────────────────────────────


def paginate(v: View, max_edge: int) -> list[tuple[list[int], list[int]]]:
    """Splits rows and columns into pages whose drawings stay within max_edge pixels. A break that would cut a
    chart or image is moved before it when the drawing fits on one page; one that does not fit is drawn in part
    on each page it overlaps (draw_page)."""
    zoom = v.zoom
    spans = [sp for sp in drawing_spans(v) if sp is not None]
    col_groups = _groups(v.cols, [v.col_w.get(c, 64) for c in v.cols], max_edge, round(40 * zoom), [sp[1] for sp in spans])
    row_groups = _groups(v.rows, [v.row_h.get(r, 20) for r in v.rows], max_edge, round(46 * zoom), [sp[0] for sp in spans])
    return [(rg, cg) for rg in row_groups for cg in col_groups] or [(v.rows, v.cols)]


def _groups(keys: list[int], sizes: list[int], max_edge: int, head: int, spans: list[tuple[int, int]]) -> list[list[int]]:
    """Consecutive runs of keys whose sizes add up to at most max_edge - head; spans (index ranges of drawings)
    move a break to before a drawing that would be cut and fits in one run."""
    groups: list[list[int]] = []
    start, used, i = 0, head, 0
    while i < len(keys):
        if i > start and used + sizes[i] > max_edge:
            cut = i
            fits = [a for a, b in spans if start < a < i <= b and head + sum(sizes[a : b + 1]) <= max_edge]
            if fits:
                cut = min(fits)
            groups.append(keys[start:cut])
            start, used = cut, head + sum(sizes[cut:i])
            continue
        used += sizes[i]
        i += 1
    if start < len(keys):
        groups.append(keys[start:])
    return groups


def render_workbook_pages(path: Path, outdir: Path, sheet: str | None = None, area: tuple[int, int, int, int] | None = None, zoom: float = 1.0, max_rows: int = 500, max_cols: int = 60) -> list[Path]:
    """Every visible worksheet (or one sheet, or one area) drawn as pages, for the built-in PDF export."""
    from _book import Workbook

    with Workbook(path) as wb:
        sheets = [wb.resolve_sheet(sheet).name] if sheet else [m.name for m in wb.sheets() if m.kind == "worksheet" and m.visible == "visible"]
    out: list[Path] = []
    for name in sheets:
        v = load_view(path, name, area, max_rows, max_cols, zoom)
        if not v.rows or not v.cols:
            continue
        for rows, cols in paginate(v, 1400):
            img = draw_page(v, rows, cols, title=f"{v.sheet} — {col_letter(cols[0])}{rows[0]}:{col_letter(cols[-1])}{rows[-1]}")
            p = outdir / f"{len(out) + 1:03d}.png"
            img.save(p)
            out.append(p)
    return out
