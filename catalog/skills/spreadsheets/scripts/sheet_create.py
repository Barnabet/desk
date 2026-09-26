#!/usr/bin/env python3
"""Create a styled, recalculated workbook from a JSON spec, CSV/TSV, JSON records or Markdown tables."""

from __future__ import annotations

import datetime as _dt
import re
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import SkillError, UsageError, add_format, emit, input_file, load_json_arg, output_path, parser, run_main  # noqa: E402

EPILOG = """examples:
  python3 scripts/sheet_create.py sales.xlsx --from sales.csv --table --totals
  python3 scripts/sheet_create.py report.xlsx --from summary.md                 # every pipe table becomes a sheet
  python3 scripts/sheet_create.py model.xlsx --spec model.json
  python3 scripts/sheet_create.py quick.xlsx --spec '{"sheets":[{"name":"Q1","columns":[{"header":"Item"},
      {"header":"Qty","format":"#,##0"},{"header":"Price","format":"$#,##0.00"},
      {"header":"Total","formula":"=B{row}*C{row}","format":"$#,##0.00","total":"sum"}],
      "rows":[["Pens",10,1.5],["Paper",4,6]],"charts":[{"type":"column","categories":"Item","values":["Total"]}]}]}'

Spec keys (full reference: references/spec.md): sheets[] with name, columns[] (header, width, format, formula with
{row}, total, align, style), rows (lists or objects), start, table (true or {name, style}), header_style, freeze,
autofilter, banded, totals, conditional[], validations[], charts[], cells{}, merge[], ops[] (any sheet_edit op);
plus workbook-level names{}, properties{}, ops[].
"""


def main() -> int:
    p = parser("Create an .xlsx/.xlsm/.xltx workbook from a JSON spec or from data files.", EPILOG)
    p.add_argument("out", help="the workbook to write")
    p.add_argument("--spec", help="JSON spec (inline, a .json file, or - for stdin)")
    p.add_argument("--from", dest="sources", action="append", default=[], help="a .csv/.tsv/.json/.jsonl/.md file (repeatable; one sheet per table)")
    p.add_argument("--sheet-name", action="append", default=[], help="sheet name for each --from (in order)")
    p.add_argument("--table", action="store_true", help="make each data block an Excel table (with --from)")
    p.add_argument("--totals", action="store_true", help="add a totals row that sums the numeric columns (with --from)")
    p.add_argument("--chart", choices=["column", "bar", "line", "area", "pie", "scatter"], help="add a chart of the numeric columns against the first column (with --from)")
    p.add_argument("--no-style", action="store_true", help="plain header, no freeze or filter (with --from)")
    p.add_argument("--allow-formulas", action="store_true", help="with --from: text starting with = becomes a formula (off by default: CSV formula injection)")
    p.add_argument("--no-recalc", action="store_true", help="skip computing cached values")
    p.add_argument("--now", help="date/time for TODAY()/NOW()")
    p.add_argument("--force", action="store_true", help="overwrite the output if it exists")
    add_format(p)
    a = p.parse_args()
    out = output_path(a.out, [s for s in a.sources], a.force)
    if out.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        raise UsageError("the output must be .xlsx (or .xlsm/.xltx/.xltm); convert afterwards with sheet_convert.py for other formats")
    if not a.spec and not a.sources:
        raise UsageError("give --spec or at least one --from file")
    spec: dict[str, Any] = {"sheets": []}
    if a.spec:
        loaded = load_json_arg(a.spec)
        if isinstance(loaded, list):
            loaded = {"sheets": loaded}
        if not isinstance(loaded, dict):
            raise UsageError("the spec must be a JSON object with 'sheets'")
        spec = loaded
        spec.setdefault("sheets", [])
    for i, src in enumerate(a.sources):
        spec["sheets"].extend(sheets_from_file(input_file(src), a, a.sheet_name[i] if i < len(a.sheet_name) else None))
    if not spec["sheets"]:
        raise UsageError("nothing to write: the spec has no sheets and the inputs no tables")
    now = _dt.datetime.fromisoformat(a.now) if a.now else None
    wb, ctx, log = build(spec)
    from _bridge import save_and_recalc

    report = save_and_recalc(wb, out, ctx, recalc=not a.no_recalc, now=now)
    for sd in spec["sheets"]:
        if isinstance(sd, dict) and sd.get("_formula_text"):
            ctx.warnings.append(f"{sd['_formula_text']} cells of {sd.get('name')} start with '=' and were kept as text (a data file's text is not a formula); --allow-formulas makes them formulas")
        if isinstance(sd, dict) and sd.get("_chart_note"):
            ctx.warnings.append(sd["_chart_note"])
    result = {"output": str(out), "sheets": [{"name": ws.title, "range": ws.dimensions} for ws in wb.worksheets], "log": log, "warnings": ctx.warnings, "recalc": report}
    emit(result, a.format, render)
    return 0


# ── inputs → sheet specs ────────────────────────────────────────────────


def sheets_from_file(path: Path, a: Any, name: str | None) -> list[dict[str, Any]]:
    from _book import Currency, Percent, json_blocks, kind_of, markdown_tables, read_csv

    ext = path.suffix.lower()
    blocks: list[tuple[str, list[list[Any]]]] = []
    if ext in (".md", ".markdown", ".txt") and "|" in path.read_text(encoding="utf-8", errors="replace")[:200000]:
        tables = markdown_tables(path.read_text(encoding="utf-8", errors="replace"))
        if not tables:
            raise SkillError(f"{path.name} has no Markdown pipe tables")
        for k, (heading, rows) in enumerate(tables, 1):
            blocks.append((heading or f"Table {k}", rows))
    elif ext in (".json", ".jsonl", ".ndjson"):
        blocks.extend(json_blocks(path))
    elif kind_of(path) in ("html", "xlsx", "xls", "xlsb", "ods"):
        from _book import Workbook, trim_grid

        with Workbook(path) as wb:
            for m in wb.sheets():
                if m.kind == "worksheet":
                    blocks.append((m.name, trim_grid(wb.grid(m.name).rows)))
    else:
        rows, _ = read_csv(path)
        blocks.append((path.stem, rows))
    out = []
    for k, (title, rows) in enumerate(blocks):
        if not rows:
            continue
        header = [str(h) if h is not None else f"Column{j + 1}" for j, h in enumerate(rows[0])]
        body = rows[1:]
        cols: list[dict[str, Any]] = [{"header": h} for h in header]
        for j, col in enumerate(cols):
            vals = [r[j] for r in body if j < len(r) and r[j] is not None]
            if vals and all(isinstance(v, Percent) for v in vals if isinstance(v, float)) and any(isinstance(v, Percent) for v in vals):
                col["format"] = "0.0%" if any(round(v * 100, 1) != round(v * 100) for v in vals if isinstance(v, float)) else "0%"
            elif vals and any(isinstance(v, Currency) for v in vals):
                sym = next(v.symbol for v in vals if isinstance(v, Currency))
                col["format"] = f'"{sym}"#,##0.00' if sym != "$" else "$#,##0.00"
            elif vals and all(isinstance(v, _dt.datetime) for v in vals):
                col["format"] = "yyyy-mm-dd hh:mm"
            if a.totals and vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals) and j > 0:
                col["total"] = "average" if any(isinstance(v, Percent) for v in vals) else "sum"
        clean = [[float(v) if isinstance(v, (Percent, Currency)) else v for v in r] for r in body]
        guarded = 0
        if not getattr(a, "allow_formulas", False):
            # data from files is data: '=HYPERLINK(…)' in a CSV must not become a live formula
            for r in clean:
                for j, v in enumerate(r):
                    if isinstance(v, str) and v.startswith("=") and len(v) > 1:
                        r[j] = {"value": v, "text": True}
                        guarded += 1
            for j, h in enumerate(header):
                if h.startswith("=") and len(h) > 1:
                    cols[j]["header"] = h
                    cols[j]["header_text"] = True
        sheet_name = name if name and len(blocks) == 1 else (f"{name} {k + 1}" if name else title)
        sd: dict[str, Any] = {"name": _safe_sheet_name(sheet_name), "columns": cols, "rows": clean}
        if guarded:
            sd["_formula_text"] = guarded
        if a.table:
            sd["table"] = True
        if a.no_style:
            sd.update({"header_style": {"font": {"bold": True}}, "freeze": False, "autofilter": False})
        if a.chart:
            numeric = [(j, c) for j, c in enumerate(cols) if j > 0 and any(isinstance(r[j], (int, float)) and not isinstance(r[j], bool) for r in clean if j < len(r))]
            # one value axis: plot the columns on the scale of the biggest one (percentages next to currency vanish)
            def peak(j: int) -> float:
                return max((abs(float(r[j])) for r in clean if j < len(r) and isinstance(r[j], (int, float)) and not isinstance(r[j], bool)), default=0.0)

            pct = [c["header"] for j, c in numeric if str(c.get("format", "")).endswith("%")]
            plain = [(j, c) for j, c in numeric if not str(c.get("format", "")).endswith("%")]
            pick = plain or numeric
            top = max((peak(j) for j, _ in pick), default=0.0)
            values = [c["header"] for j, c in pick if top == 0 or peak(j) >= top / 50][:6]
            left_out = [c["header"] for j, c in numeric if c["header"] not in values]
            if values:
                title = (" and ".join(values) if len(values) <= 2 else ", ".join(values[:-1]) + " and " + values[-1]) + f" by {cols[0]['header']}"
                sd["charts"] = [{"type": a.chart, "categories": cols[0]["header"], "values": values, "title": title}]
                if left_out:
                    sd["_chart_note"] = f"chart leaves out {', '.join(left_out)} ({'percentages' if set(left_out) <= set(pct) else 'much smaller values'} on another scale); add a second chart with --spec for them"
        out.append(sd)
    return out


def _safe_sheet_name(name: str) -> str:
    s = re.sub(r"[\[\]:*?/\\]", "-", str(name)).strip().strip("'") or "Sheet"
    return s[:31]


# ── spec → workbook ─────────────────────────────────────────────────────

DEFAULT_HEADER = {"font": {"bold": True, "color": "#FFFFFF"}, "fill": "#1F4E78", "align": {"horizontal": "center", "vertical": "center", "wrap": True}}


def build(spec: dict[str, Any]) -> tuple[Any, Any, list[str]]:
    import openpyxl

    from _ops import Ctx, apply_ops

    wb = openpyxl.Workbook()
    ctx = Ctx(wb)
    log: list[str] = []
    used: set[str] = set()
    for i, sd in enumerate(spec["sheets"]):
        if not isinstance(sd, dict):
            raise UsageError(f"sheet {i + 1} must be an object")
        name = _safe_sheet_name(sd.get("name") or f"Sheet{i + 1}")
        base, k = name, 2
        while name.lower() in used:
            name = f"{base[:28]} {k}"
            k += 1
        used.add(name.lower())
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        log.append(build_sheet(ctx, ws, sd))
    names = spec.get("names") or {}
    name_ops = [{"op": "name", "name": n, "ref": r} for n, r in names.items()] if isinstance(names, dict) else [dict(n, op="name") for n in names]
    extra = name_ops
    if spec.get("properties"):
        extra.append(dict(spec["properties"], op="properties"))
    if spec.get("calc"):
        extra.append(dict(spec["calc"], op="calc"))
    extra.extend(spec.get("ops") or [])
    if extra:
        log.extend(apply_ops(ctx, extra))
    if spec.get("active"):
        wb.active = wb.sheetnames.index(ctx.ws(spec["active"]).title)
    return wb, ctx, log


def build_sheet(ctx: Any, ws: Any, sd: dict[str, Any]) -> str:
    from _a1 import col_letter, parse_cell, range_name
    from _ops import apply_ops, apply_style, autofit, to_cell_value, write_cell

    r0, c0 = parse_cell(str(sd.get("start", "A1")))
    columns = sd.get("columns")
    rows = sd.get("rows", sd.get("data")) or []
    if columns is None and rows and isinstance(rows[0], dict):
        keys: list[str] = []
        for r in rows:
            for k in r:
                if k not in keys:
                    keys.append(k)
        columns = [{"header": k} for k in keys]
    if columns is None and sd.get("header"):
        columns = [{"header": h} for h in sd["header"]]
    columns = [c if isinstance(c, dict) else {"header": c} for c in (columns or [])]
    has_header = any(c.get("header") is not None for c in columns)
    table = sd.get("table")
    ncols = max(len(columns), max((len(r) for r in rows if isinstance(r, list)), default=0))
    row = r0
    if has_header:
        for j in range(ncols):
            h = columns[j].get("header") if j < len(columns) else None
            from _ops import LiteralText

            hv = h if h is not None else f"Column{j + 1}"
            write_cell(ctx, ws, row, c0 + j, LiteralText(hv) if isinstance(hv, str) and hv.startswith("=") else hv)
        row += 1
    first_data = row
    for rv in rows:
        if isinstance(rv, dict):
            vals = [rv.get(c.get("header")) for c in columns]
        elif isinstance(rv, list):
            vals = list(rv)
        else:
            vals = [rv]
        for j in range(ncols):
            col = columns[j] if j < len(columns) else {}
            v = vals[j] if j < len(vals) else None
            fmt = None
            if v is None and col.get("formula"):
                f = str(col["formula"]).replace("{row}", str(row)).replace("{r}", str(row))
                v, fmt = to_cell_value(f if f.startswith("=") else "=" + f)
            elif v is not None:
                v, fmt = to_cell_value(v)
            if v is not None:
                write_cell(ctx, ws, row, c0 + j, v, col.get("format") or fmt)
            elif col.get("format"):
                ws.cell(row, c0 + j).number_format = col["format"]
        row += 1
    last_data = row - 1
    n_data = last_data - first_data + 1
    c_last = c0 + max(ncols, 1) - 1
    # column formats, alignment and styles
    for j, col in enumerate(columns):
        c = c0 + j
        if n_data > 0 and (col.get("align") or col.get("style")):
            st = dict(col.get("style") or {})
            if col.get("align"):
                st["align"] = col["align"] if isinstance(col["align"], dict) else {"horizontal": col["align"]}
            apply_style(ws, first_data, c, last_data, c, st)
    # header style
    if has_header and not (table and sd.get("header_style") is None):
        hs = sd.get("header_style", DEFAULT_HEADER)
        if hs:
            apply_style(ws, r0, c0, r0, c_last, hs)
    elif has_header and table:
        apply_style(ws, r0, c0, r0, c_last, {"align": {"horizontal": "center", "vertical": "center", "wrap": True}})
    if sd.get("banded") and not table and n_data > 1:
        fill = sd["banded"] if isinstance(sd["banded"], str) else "#F2F2F2"
        for r in range(first_data + 1, last_data + 1, 2):
            apply_style(ws, r, c0, r, c_last, {"fill": fill})
    notes = []
    totals = {c.get("header"): c.get("total") for c in columns if c.get("total")}
    if isinstance(sd.get("totals"), dict):
        totals.update(sd["totals"])
    elif sd.get("totals") is True and not totals:
        for j, col in enumerate(columns[1:], 1):
            vals = [r[j] for r in rows if isinstance(r, list) and j < len(r)]
            if vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals if v is not None):
                totals[col.get("header")] = "sum"
    ops: list[dict[str, Any]] = []
    if table:
        t = table if isinstance(table, dict) else {}
        top = {"op": "table", "sheet": ws.title, "range": range_name(r0, c0, max(last_data, r0 + 1), c_last)}
        top.update({k: v for k, v in t.items() if k in ("name", "style", "banded_rows", "banded_columns")})
        if not t.get("name"):
            top["name"] = re.sub(r"\W", "_", ws.title) or "Table"
        if totals:
            top["totals"] = totals
        ops.append(top)
    elif totals and n_data > 0:
        tr = last_data + 1
        fn_codes = {"sum": 109, "average": 101, "avg": 101, "count": 103, "max": 104, "min": 105}
        label_col = None
        for j, col in enumerate(columns):
            fn = totals.get(col.get("header"))
            if fn:
                code = fn_codes.get(str(fn).lower())
                if code is None:
                    raise SkillError(f"total '{fn}' for {col.get('header')} (use sum, average, count, max or min)")
                L = col_letter(c0 + j)
                write_cell(ctx, ws, tr, c0 + j, f"=SUBTOTAL({code},{L}{first_data}:{L}{last_data})", col.get("format"))
            elif label_col is None:
                label_col = c0 + j
        if label_col is not None and label_col < c0 + next((j for j, c in enumerate(columns) if totals.get(c.get("header"))), 0):
            write_cell(ctx, ws, tr, label_col, "Total")
        apply_style(ws, tr, c0, tr, c_last, {"font": {"bold": True}, "border": {"top": "thin", "bottom": "double"}})
        notes.append("totals row")
    header_row = r0 if has_header else None
    freeze = sd.get("freeze", True if has_header and n_data > 0 else False)
    if freeze:
        cell = freeze if isinstance(freeze, str) else f"{col_letter(c0)}{r0 + 1}"
        ws.freeze_panes = cell
    if sd.get("autofilter", has_header and not table and n_data > 0):
        ws.auto_filter.ref = range_name(r0, c0, max(last_data, r0), c_last)
    for k, v in (sd.get("cells") or {}).items():
        ops.append({"op": "set", "sheet": ws.title, "cell": k, "value": v})
    for m in sd.get("merge") or []:
        ops.append({"op": "merge", "sheet": ws.title, "range": m})
    data_area = range_name(r0, c0, max(last_data, r0), c_last) if header_row else None
    for cf in sd.get("conditional") or sd.get("conditional_formats") or []:
        ops.append(_resolve_cols(dict(cf, op="conditional_format", sheet=ws.title), ws, r0, first_data, last_data, c0, columns))
    for v in sd.get("validations") or []:
        ops.append(_resolve_cols(dict(v, op="validation", sheet=ws.title), ws, r0, first_data, last_data, c0, columns))
    right = c_last
    for k in (sd.get("cells") or {}):
        try:
            right = max(right, parse_cell(str(k).split("!")[-1])[1])
        except ValueError:
            pass
    for n, ch in enumerate(sd.get("charts") or []):
        o = dict(ch, op="chart", sheet=ws.title)
        if data_area:
            o.setdefault("data_area", data_area)
        if not o.get("anchor"):
            o["anchor"] = f"{col_letter(right + 2)}{r0 + 1 + 17 * n}"
        ops.append(o)
    for o in sd.get("ops") or []:
        o = dict(o)
        o.setdefault("sheet", ws.title)
        ops.append(o)
    msgs = apply_ops(ctx, ops) if ops else []
    # widths: explicit ones win; others auto-fit (formula columns again after evaluation)
    explicit = {c0 + j: col["width"] for j, col in enumerate(columns) if col.get("width")}
    if sd.get("autofit", True):
        autofit(ws, (c0, c_last) if ncols else None)
        ctx.autofit_pending = getattr(ctx, "autofit_pending", []) + [(ws, (c0, c_last) if ncols else None, 6.0, 60.0)]
    for c, w in explicit.items():
        ws.column_dimensions[col_letter(c)].width = float(w)
    if explicit:
        ctx.autofit_pending = [(w_, cols, lo, hi) for (w_, cols, lo, hi) in getattr(ctx, "autofit_pending", []) if w_ is not ws] + [
            (ws, (c, c), 6.0, 60.0) for c in range(c0, c_last + 1) if c not in explicit
        ]
    for k, w in (sd.get("widths") or {}).items():
        ws.column_dimensions[k].width = float(w)
    summary = f"sheet {ws.title}: {n_data} data rows × {ncols} columns"
    if table:
        summary += ", Excel table"
    if notes:
        summary += ", " + ", ".join(notes)
    if msgs:
        summary += "; " + "; ".join(m.split(". ", 1)[-1] for m in msgs)
    return summary


def _resolve_cols(o: dict[str, Any], ws: Any, r0: int, first: int, last: int, c0: int, columns: list[dict[str, Any]]) -> dict[str, Any]:
    """'column': 'Revenue' → the data range of that column."""
    from _a1 import col_letter

    col = o.pop("column", None)
    if col is not None and "range" not in o:
        names = [str(c.get("header", "")).lower() for c in columns]
        if str(col).lower() in names:
            j = names.index(str(col).lower())
        else:
            raise SkillError(f"no column named {col!r}")
        L = col_letter(c0 + j)
        o["range"] = f"{L}{first}:{L}{max(last, first)}"
    return o


def render(r: dict[str, Any]) -> str:
    from _bridge import recalc_summary

    lines = [f"Wrote {r['output']}: " + ", ".join(f"{s['name']} ({s['range']})" for s in r["sheets"])]
    lines.extend(f"- {m}" for m in r["log"])
    for w in r.get("warnings", []):
        lines.append(f"warning: {w}")
    lines.extend(recalc_summary(r.get("recalc") or {}))
    lines.append("Check it: sheet_render.py <file> and view_image, or sheet_read.py --formulas.")
    return "\n".join(lines)


if __name__ == "__main__":
    run_main(main)
