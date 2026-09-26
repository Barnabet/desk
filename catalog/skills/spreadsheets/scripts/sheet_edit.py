#!/usr/bin/env python3
"""Edit a workbook with a list of JSON operations; formulas are recalculated and references kept consistent."""

from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import SkillError, UsageError, add_format, emit, input_file, load_json_arg, output_path, parser, run_main  # noqa: E402

EPILOG = """examples:
  python3 scripts/sheet_edit.py in.xlsx --out out.xlsx --ops '[{"op":"set","cell":"B2","value":42}]'
  python3 scripts/sheet_edit.py in.xlsx --out out.xlsx --ops edits.json
  python3 scripts/sheet_edit.py in.xlsx --out out.xlsx --sheet Data --ops '[
      {"op":"insert_rows","at":5,"count":2},
      {"op":"set","range":"A5","values":[["West",12,3.5],["East",8,4]]},
      {"op":"fill","range":"D2:D30","formula":"=B2*C2"},
      {"op":"style","range":"A1:D1","font":{"bold":true},"fill":"#DDEBF7","border":{"bottom":"medium"}},
      {"op":"conditional_format","range":"D2:D30","type":"data_bar"},
      {"op":"chart","type":"column","values":"D1:D30","categories":"A2:A30","anchor":"F2","title":"Revenue"}]'
  python3 scripts/sheet_edit.py model.xlsm --out model-v2.xlsm --mode patch --ops '[{"op":"set","cell":"Inputs!C4","value":0.07}]'

Operations (details and every option in references/ops.md):
  cells     set, fill, clear, copy_range, move_range, find_replace, sort
  structure insert_rows, delete_rows, insert_cols, delete_cols (formulas, names, merges, formats, validations,
            tables, charts and print areas are rewritten on every sheet), add_sheet, rename_sheet, copy_sheet,
            delete_sheet, move_sheet, reorder, sheet_state, tab_color
  format    style (font, fill, border, align, number_format), column_width, row_height, autofit, hide_rows,
            hide_columns, group, merge, unmerge, freeze, unfreeze, autofilter
  rules     conditional_format, validation, name, delete_name, table
  objects   chart, image, comment, hyperlink
  workbook  protect, unprotect, print, properties, calc

Modes: auto (default) writes set, fill, clear and style straight into the sheet XML ("patch": everything else in
the file is kept byte for byte, and a one-cell edit of a 40 MB workbook takes seconds) when the file is bigger than
3 MB or holds features openpyxl would drop; other operations, and small files, go through openpyxl (every operation;
drops shapes, sparklines, slicers, form controls and x14 extensions if the file has them, and says so).
"""



def _big_edit_bytes() -> int:
    """Files at least this big get cell edits patched in place (DESK_EDIT_PATCH_MB overrides the 3 MB default)."""
    import os

    try:
        return int(float(os.environ.get("DESK_EDIT_PATCH_MB", "3")) * (1 << 20))
    except ValueError:
        return 3 << 20


def main() -> int:
    p = parser("Apply JSON edit operations to an .xlsx/.xlsm/.xltx workbook and save to a new file.", EPILOG)
    p.add_argument("file")
    p.add_argument("--out", required=True, help="output workbook (never the input)")
    p.add_argument("--ops", required=True, help="JSON list of operations, a .json file, or - for stdin")
    p.add_argument("--sheet", help="default sheet for operations without a sheet (default: the active sheet)")
    p.add_argument("--mode", choices=["auto", "openpyxl", "patch"], default="auto", help="auto = patch when only set/fill/clear/style are used on a big file or one with features openpyxl would drop")
    p.add_argument("--no-recalc", action="store_true", help="do not recalculate (cached values are then missing for formulas)")
    p.add_argument("--now", help="date/time for TODAY()/NOW() during recalculation")
    p.add_argument("--force", action="store_true", help="overwrite --out if it exists")
    add_format(p)
    a = p.parse_args()
    src = input_file(a.file, {".xlsx", ".xlsm", ".xltx", ".xltm"})
    from _book import kind_of

    if kind_of(src) != "xlsx":
        raise SkillError(f"{src.name} is not an .xlsx-family workbook; convert it with sheet_convert.py --to xlsx first")
    out = output_path(a.out, [src], a.force)
    ops = load_json_arg(a.ops)
    if isinstance(ops, dict):
        ops = ops.get("ops", [ops])
    if not isinstance(ops, list) or not ops:
        raise UsageError("--ops must be a non-empty JSON list of operations")
    now = _dt.datetime.fromisoformat(a.now) if a.now else None
    from _ops import ALIASES

    from _patch import PATCH_OPS

    names = {ALIASES.get(str(o.get("op", "")).lower(), str(o.get("op", "")).lower()) for o in ops if isinstance(o, dict)}
    mode = a.mode
    losses = _openpyxl_losses(src)
    big = src.stat().st_size >= _big_edit_bytes()
    notes: list[str] = []
    if mode == "auto":
        mode = "patch" if names <= PATCH_OPS and (losses or big) else "openpyxl"
        if mode == "openpyxl" and big:
            cells = _cell_count(src)
            what = ", ".join(sorted(names - PATCH_OPS)) or "these operations"
            cost = _slow_cost(cells, not a.no_recalc)
            if cells > _slow_cells():
                raise SkillError(f"{what} on this workbook ({cells:,} cells) needs the full openpyxl path: loading and saving everything takes {cost}. "
                                 "Run again with --mode openpyxl to go ahead. Cell edits (set, fill, clear, style) alone are patched in place in seconds, "
                                 "so split them into their own call; for a new layout, sheet_query.py --out can write a reshaped copy of the data.")
            notes.append(f"{what} needed the full openpyxl path: the whole workbook ({cells:,} cells) was loaded and saved ({cost}). Cell edits (set, fill, clear, style) alone are patched in place in seconds.")
    if mode == "patch":
        if not names <= PATCH_OPS:
            raise UsageError(f"--mode patch supports only {', '.join(sorted(PATCH_OPS))}; use --mode openpyxl for {', '.join(sorted(names - PATCH_OPS))}")
        result = run_patch(src, out, ops, a.sheet, not a.no_recalc, now)
    else:
        result = run_openpyxl(src, out, ops, a.sheet, not a.no_recalc, now)
        if losses:
            result["warnings"].append("rewritten with openpyxl, which dropped: " + ", ".join(sorted(losses)) + " (edit with set/fill/clear/style only to keep them: they are patched in place)")
    result["warnings"] = notes + result["warnings"]
    result["mode"] = mode
    emit(result, a.format, render)
    return 0


def _slow_cells() -> int:
    """Above this many cells a structural edit is refused until --mode openpyxl is given (DESK_EDIT_SLOW_CELLS)."""
    import os

    try:
        return int(os.environ.get("DESK_EDIT_SLOW_CELLS", "2000000"))
    except ValueError:
        return 2_000_000


def _slow_cost(cells: int, recalc: bool) -> str:
    """Measured on a 6-million-cell workbook: about 19 µs and 340 bytes per cell to load and save with openpyxl,
    1.5 µs more per cell to recalculate."""
    secs = cells * (19e-6 + (1.5e-6 if recalc else 0))
    mem = cells * 340 / (1 << 30)
    t = f"about {max(2, round(secs))} s" if secs < 90 else f"about {secs / 60:.0f} min"
    return f"{t} and {mem:.1f} GB of memory" if mem >= 0.5 else t


def _openpyxl_losses(src: Path) -> set[str]:
    from _xlsx import Package, openpyxl_losses

    with Package(src) as pkg:
        return openpyxl_losses(pkg)


def _cell_count(src: Path) -> int:
    from _scan import scan_cached
    from _xlsx import Package

    with Package(src) as pkg:
        parts = [si.path for si in pkg.sheets if si.path and pkg.has(si.path)]
    return sum((scan_cached(src, part) or {}).get("cells", 0) for part in parts)


def run_openpyxl(src: Path, out: Path, ops: list[dict[str, Any]], sheet: str | None, recalc: bool, now: Any) -> dict[str, Any]:
    import warnings

    import openpyxl

    from _bridge import save_and_recalc
    from _ops import Ctx, apply_ops

    has_vba = False
    try:
        import zipfile

        with zipfile.ZipFile(src) as z:
            has_vba = any(n.lower().endswith("vbaproject.bin") for n in z.namelist())
    except zipfile.BadZipFile as e:
        raise SkillError(f"{src.name} is not a valid workbook") from e
    keep_vba = has_vba and out.suffix.lower() in (".xlsm", ".xltm")
    notes = []
    if has_vba and not keep_vba:
        notes.append(f"{src.name} contains VBA macros; they are not kept in {out.suffix} (write to .xlsm to keep them)")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            wb = openpyxl.load_workbook(src, keep_vba=keep_vba, keep_links=True, rich_text=True)
        except Exception as e:  # noqa: BLE001
            raise SkillError(f"openpyxl cannot open {src.name}: {type(e).__name__}: {e}") from e
    ctx = Ctx(wb, sheet)
    if sheet:
        ctx.ws(sheet)
    log = apply_ops(ctx, ops)
    notes.extend(_pivot_notes(wb))
    report = save_and_recalc(wb, out, ctx, recalc=recalc, now=now)
    return {"output": str(out), "ops": log, "warnings": ctx.warnings + notes, "recalc": report}


def _pivot_notes(wb: Any) -> list[str]:
    """Pivot tables are kept, not refreshed: say when data now lies outside a pivot table's source range."""
    from _a1 import parse_range

    out = []
    seen = set()
    for ws in wb.worksheets:
        for pt in getattr(ws, "_pivots", []) or []:
            src = getattr(getattr(getattr(pt, "cache", None), "cacheSource", None), "worksheetSource", None)
            if src is None or not getattr(src, "ref", None) or not getattr(src, "sheet", None) or id(src) in seen:
                continue
            seen.add(id(src))
            if src.sheet not in wb.sheetnames:
                out.append(f"pivot table {pt.name} reads {src.sheet}!{src.ref}, a sheet that no longer exists")
                continue
            try:
                r1, c1, r2, c2 = parse_range(src.ref)
            except ValueError:
                continue
            data_ws = wb[src.sheet]
            below = []  # data rows appended right under the source range (a contiguous block)
            for r in range(r2 + 1, min(data_ws.max_row, r2 + 100_000) + 1):
                if not any(data_ws.cell(r, c).value is not None for c in range(c1, c2 + 1)):
                    break
                below.append(r)
            if below:
                out.append(f"pivot table {pt.name} reads {src.sheet}!{src.ref}, but {'row ' + str(below[0]) if len(below) == 1 else f'rows {below[0]}-{below[-1]}'} below it now hold{'s' if len(below) == 1 else ''} data: widen its source and refresh it in Excel (pivot tables are kept, not recomputed)")
            else:
                out.append(f"pivot table {pt.name} ({src.sheet}!{src.ref}) is kept as it was; refresh it in Excel or LibreOffice to see the edits")
    return out


def run_patch(src: Path, out: Path, ops: list[dict[str, Any]], sheet: str | None, recalc: bool, now: Any) -> dict[str, Any]:
    """set/fill/clear/style written directly into the sheet XML and styles.xml; then the engine refreshes cached values."""
    import time

    from _patch import apply
    from _xlsx import publish, recalc_file, temp_beside

    t0 = time.time()
    tmp = temp_beside(out, prefix=".desk-patch-", suffix=out.suffix or ".xlsx")
    try:
        plan = apply(src, tmp, ops, sheet, full_calc=not recalc)
        report: dict[str, Any] = {}
        if recalc and plan.values_changed:
            report = recalc_file(tmp, out, now=now)
            report.pop("_book", None)
        else:
            publish(tmp, out)
    finally:
        if tmp.exists():
            tmp.unlink()
    return {"output": str(out), "ops": plan.log, "warnings": sorted(set(plan.warnings)), "recalc": report, "seconds": round(time.time() - t0, 2)}


def render(r: dict[str, Any]) -> str:
    from _bridge import recalc_summary

    lines = [f"Wrote {r['output']} ({r['mode']} mode" + (f", {r['seconds']} s" if r.get("seconds") else "") + ")."]
    lines.extend(r["ops"])
    for w in r.get("warnings", []):
        lines.append(f"warning: {w}")
    lines.extend(recalc_summary(r.get("recalc") or {}))
    lines.append("Check it: sheet_read.py (values, --formulas) and sheet_render.py + view_image (layout).")
    return "\n".join(lines)


if __name__ == "__main__":
    run_main(main)
