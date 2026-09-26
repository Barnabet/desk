#!/usr/bin/env python3
"""Convert between spreadsheet formats (xlsx xlsm xls xlsb ods csv tsv json jsonl md html parquet pdf)."""

from __future__ import annotations

import datetime as _dt
import html
import json
import shutil
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import SkillError, UsageError, add_format, emit, input_file, output_dir, output_path, parser, run_main  # noqa: E402

EPILOG = """examples:
  python3 scripts/sheet_convert.py legacy.xls --out legacy.xlsx          # LibreOffice when installed (formulas, formatting); else values
  python3 scripts/sheet_convert.py book.xlsx --out book.ods              # built-in ODS writer keeps formulas and formats
  python3 scripts/sheet_convert.py book.xlsx --out data.csv --sheet Orders
  python3 scripts/sheet_convert.py book.xlsx --to csv --out-dir csv/     # one file per sheet
  python3 scripts/sheet_convert.py book.xlsx --out book.json             # {sheet: [records]} with header detection
  python3 scripts/sheet_convert.py book.xlsx --out book.md --display      # Markdown tables with Excel's number formats
  python3 scripts/sheet_convert.py data.csv --out data.xlsx              # styled header, types, widths (like sheet_create)
  python3 scripts/sheet_convert.py big.csv --out big.parquet             # DuckDB, streaming: a 1M-row CSV in seconds
  python3 scripts/sheet_convert.py book.xlsx --out book.pdf              # LibreOffice print layout, else built-in A4 pages

--engine auto uses LibreOffice for workbook→workbook and PDF conversions when it is installed, else the built-in path:
.xls/.xlsb → .xlsx keep values only (formulas need LibreOffice); .ods → .xlsx keeps formulas, names, number formats,
fonts, fills, borders, alignment, widths and merges; .xlsx → .ods keeps formulas (structured references become
ranges), values, number formats, bold/italic, fills, merges and widths, and lists what it drops (Excel tables,
conditional formats, validations, charts, comments).
Data targets (csv, tsv, json, jsonl, parquet) hold the data rows: totals rows (of Excel tables, or a last row labelled
Total) and title lines above the header are left out and listed; --keep-totals keeps totals rows. CSV text that
starts with = stays text in .xlsx output unless --allow-formulas (formula injection from untrusted files).
"""

WORKBOOK = {"xlsx", "xlsm", "xltx", "xltm", "xls", "xlsb", "ods"}
TEXT = {"csv", "tsv", "json", "jsonl", "md", "html", "parquet"}


def main() -> int:
    p = parser("Convert a spreadsheet to another format.", EPILOG)
    p.add_argument("file")
    p.add_argument("--out", help="output file (its extension picks the format)")
    p.add_argument("--to", help="target format when using --out-dir: xlsx, ods, xls, csv, tsv, json, jsonl, md, html, parquet, pdf")
    p.add_argument("--out-dir", help="write one file per sheet into this folder (csv, tsv, json, jsonl, md, html, parquet)")
    p.add_argument("--sheet", help="only this sheet (name or number)")
    p.add_argument("--range", help="only this area, e.g. A1:F200 or 'Sheet 2'!B3:D9")
    p.add_argument("--header", choices=["auto", "yes", "no"], default="auto", help="first row holds column names (json/parquet; default auto)")
    p.add_argument("--display", action="store_true", help="text formats show values with Excel number formats (xlsx family)")
    p.add_argument("--engine", choices=["auto", "builtin", "libreoffice"], default="auto")
    p.add_argument("--delimiter", help="CSV delimiter (default , for csv, TAB for tsv)")
    p.add_argument("--bom", action="store_true", help="write a UTF-8 BOM in CSV (helps Excel on Windows detect UTF-8)")
    p.add_argument("--keep-totals", action="store_true", help="data targets: keep totals rows as data")
    p.add_argument("--allow-formulas", action="store_true", help="CSV/JSON → xlsx: text starting with = becomes a formula (off by default: formula injection)")
    p.add_argument("--no-cache", action="store_true", help="do not use or store cached copies of big sheets")
    p.add_argument("--force", action="store_true")
    add_format(p)
    a = p.parse_args()
    if a.no_cache:
        import os

        os.environ["DESK_NO_CACHE"] = "1"
    src = input_file(a.file)
    from _book import kind_of

    kind = kind_of(src)
    if a.out_dir:
        target = (a.to or "csv").lower().lstrip(".")
        if target not in TEXT:
            raise UsageError("--out-dir writes one file per sheet: use --to csv, tsv, json, jsonl, md, html or parquet")
        result = per_sheet(src, Path(a.out_dir), target, a)
    else:
        if not a.out:
            raise UsageError("give --out <file> (or --out-dir with --to)")
        out = output_path(a.out, [src], a.force)
        target = (a.to or out.suffix.lstrip(".")).lower()
        if target in WORKBOOK:
            result = to_workbook(src, kind, out, target, a)
        elif target == "pdf":
            result = to_pdf(src, kind, out, a)
        elif target in TEXT or target == "txt":
            result = to_text(src, out, "csv" if target == "txt" else target, a)
        else:
            raise UsageError(f"unsupported target '.{target}'")
    emit(result, a.format, render)
    return 0


# ── workbook targets ────────────────────────────────────────────────────


def to_workbook(src: Path, kind: str, out: Path, target: str, a: Any) -> dict[str, Any]:
    import _lo

    if kind in ("csv", "json"):
        if target in ("xlsx", "xlsm", "xltx", "xltm"):
            return from_data(src, out, a)
        if target == "ods":
            tmp = out.with_suffix(".tmp.xlsx")
            try:
                from_data(src, tmp, a)
                return xlsx_to_ods(tmp, out, note="from delimited text")
            finally:
                tmp.unlink(missing_ok=True)
    src_kind = "xlsx" if kind == "xlsx" else kind
    if src_kind == "xlsx" and target in ("xlsx", "xlsm", "xltx", "xltm") and out.suffix.lower() == src.suffix.lower():
        if a.sheet or a.range:
            raise UsageError("to keep part of an .xlsx, use sheet_edit.py (delete_sheet ops keep formatting) or convert to csv/json")
        shutil.copyfile(src, out)
        return {"output": str(out), "engine": "copy", "notes": ["same format: copied unchanged"]}
    use_lo = a.engine == "libreoffice" or (a.engine == "auto" and _lo.available() and not (src_kind == "xlsx" and target == "ods"))
    if target == "xls" and not _lo.available():
        raise SkillError("writing .xls needs LibreOffice; write .xlsx instead (every modern tool reads it)")
    if use_lo:
        if not _lo.available():
            raise SkillError("LibreOffice is not installed; use --engine builtin")
        fmt = {"xlsx": "xlsx:Calc MS Excel 2007 XML", "xlsm": "xlsm:Calc MS Excel 2007 VBA XML", "xls": "xls:MS Excel 97", "ods": "ods", "xltx": "xltx:Calc MS Excel 2007 XML Template", "xltm": "xlsm:Calc MS Excel 2007 VBA XML"}.get(target, target)
        notes = []
        if target == "xls" and src_kind == "xlsx":
            risky = _xls_risks(src)
            if risky:
                notes.append("Excel 97 (.xls) cannot store " + risky + ": LibreOffice writes such formulas as =NA() or as values. Keep the .xlsx as the master copy.")
        produced = _lo.convert(src, fmt)
        try:
            shutil.move(str(produced), str(out))
        finally:
            shutil.rmtree(produced.parent, ignore_errors=True)
        notes.insert(0, "LibreOffice converted the whole workbook (formatting and charts where the target format supports them).")
        if target in ("xlsx", "xlsm"):
            from _xlsx import load_book

            try:
                book, pkg, _ = load_book(out)
                pkg.close()
                total = sum(len(sh.formulas) for sh in book.sheets)
                na = [f"{sh.name}!{_addr(r, c)}" for sh in book.sheets for (r, c), fc in sh.formulas.items() if fc.text.strip().upper().lstrip("=") in ("NA()", "_XLFN.NA()")]
                notes.append(f"{total} formulas in the result." if total else "The result holds no formulas (values only).")
                if na:
                    notes.append(f"{len(na)} formulas are =NA() (the source could not express them: an .xls holds no SUMIFS inside tables, structured references or newer functions): " + ", ".join(na[:8]) + (" …" if len(na) > 8 else "") + ". Check them with sheet_recalc.py --check.")
                from _formula import Engine

                rep = Engine(book).recalc()
                if rep.get("unsupported"):
                    notes.append("Some functions are not evaluated by the built-in engine: " + ", ".join(rep["unsupported"]))
            except Exception:  # noqa: BLE001 — the conversion itself succeeded
                pass
        return {"output": str(out), "engine": "libreoffice", "notes": notes}
    if target in ("xlsx", "xlsm", "xltx", "xltm"):
        return builtin_to_xlsx(src, kind, out, a)
    if target == "ods":
        if kind == "xlsx":
            return xlsx_to_ods(src, out)
        tmp = out.with_suffix(".tmp.xlsx")
        try:
            builtin_to_xlsx(src, kind, tmp, a)
            return xlsx_to_ods(tmp, out, note=f"from .{kind} values")
        finally:
            tmp.unlink(missing_ok=True)
    raise SkillError(f"cannot write .{target} without LibreOffice")


def from_data(src: Path, out: Path, a: Any = None) -> dict[str, Any]:
    from types import SimpleNamespace

    allow = bool(getattr(a, "allow_formulas", False))
    if src.stat().st_size > (4 << 20) or _csv_rows_hint(src) > 50_000:
        return stream_to_xlsx(src, out, a)
    from _bridge import save_and_recalc
    from sheet_create import build, sheets_from_file

    opts = SimpleNamespace(table=False, totals=False, chart=None, no_style=False, allow_formulas=allow)
    sheets = sheets_from_file(src, opts, None)
    spec = {"sheets": sheets}
    wb, ctx, _log = build(spec)
    save_and_recalc(wb, out, ctx)
    notes = ["typed values, styled header, frozen first row, filter, fitted widths"]
    kept_text = sum(sd.get("_formula_text", 0) for sd in sheets)
    if kept_text:
        notes.append(f"{kept_text} cells starting with '=' were kept as text (formula injection guard); --allow-formulas makes them formulas")
    return {"output": str(out), "engine": "builtin", "notes": notes}


def _csv_rows_hint(src: Path) -> int:
    size = src.stat().st_size
    if size < 256 << 10:
        return 0
    with open(src, "rb") as f:
        sample = f.read(1 << 20)
    return int(size / max(len(sample), 1) * max(sample.count(b"\n"), 1))


def stream_to_xlsx(src: Path, out: Path, a: Any = None) -> dict[str, Any]:
    """A big CSV/JSON → .xlsx without an object model: typed once into the cached Parquet copy, then written row by
    row (styled header, number and date formats, widths, frozen header, filter)."""
    from _fastxlsx import write_table
    from _store import connect, ident, sheet_store, sql_str

    meta = sheet_store(src, None, getattr(a, "header", "auto") if a is not None else "auto")
    cols = meta["columns"]
    if len(cols) > 16384:
        raise SkillError(f"{src.name} has {len(cols):,} columns; a worksheet holds 16,384")
    if meta["rows"] > 1_048_575:
        raise SkillError(f"{src.name} has {meta['rows']:,} data rows; a worksheet holds 1,048,575 under its header. Split it (sheet_query.py with LIMIT/OFFSET) or write .parquet/.csv")
    fmts = []
    for c in cols:
        f = c.get("format")
        if not f:
            f = {"DATE": "yyyy-mm-dd", "TIMESTAMP": "yyyy-mm-dd hh:mm", "TIME": "hh:mm:ss"}.get(c["type"])
        fmts.append(f)
    con = connect()
    rel = con.sql(f"SELECT {', '.join(ident(c['name']) for c in cols)} FROM read_parquet({sql_str(meta['parquet'])}) ORDER BY {ident(meta['row_column'])}")

    def gen() -> Any:
        while True:
            batch = rel.fetchmany(20000)
            if not batch:
                break
            yield from batch

    header = [c["name"] for c in cols] if meta.get("header_row") else [c["letter"] for c in cols]
    n = write_table(out, _sheet_title(src.stem), header, gen(), fmts)
    con.close()
    notes = [f"{n:,} rows written by the streaming writer (typed columns, styled header, frozen first row, filter, widths)"]
    if any(c["type"] == "VARCHAR" for c in cols):
        notes.append("text is written as text: cells starting with '=' stay text")
    return {"output": str(out), "engine": "builtin", "rows": n, "notes": notes}


def _sheet_title(name: str) -> str:
    import re

    return (re.sub(r"[\[\]:*?/\\]", "-", name).strip("'") or "Sheet")[:31]


def _addr(r: int, c: int) -> str:
    from _a1 import col_letter

    return f"{col_letter(c)}{r}"


_NEWER = {"IFERROR", "SUMIFS", "COUNTIFS", "AVERAGEIFS", "IFS", "MAXIFS", "MINIFS", "SWITCH", "XLOOKUP", "XMATCH", "TEXTJOIN", "CONCAT",
          "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "LET", "LAMBDA", "TAKE", "DROP", "VSTACK", "HSTACK", "AGGREGATE", "IFNA"}


def _xls_risks(src: Path) -> str:
    """What an .xlsx uses that Excel 97 cannot store (for the .xls warning)."""
    from _book import Workbook, SheetMeta  # noqa: F401
    from _xlsx import Package

    funcs: set[str] = set()
    tables = 0
    with Workbook(src) as wb:
        for m in wb.sheets():
            st = wb.scan(m) or {}
            funcs |= set((st.get("functions") or {}).keys())
    with Package(src) as pkg:
        tables = len(pkg.tables())
    parts = []
    newer = sorted(f for f in funcs if f in _NEWER)
    if newer:
        parts.append("newer functions (" + ", ".join(newer[:8]) + ")")
    if tables:
        parts.append(f"Excel tables and their structured references ({tables} tables)")
    return " or ".join(parts)


def _apply_ods_format(ws: Any, r: int, c: int, f: dict[str, Any], cache: dict[Any, Any]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    cell = ws.cell(r, c)
    if cell.__class__.__name__ == "MergedCell":
        return
    key = repr(sorted(f.items()))
    got = cache.get(key)
    if got is None:
        font = None
        if any(k in f for k in ("bold", "italic", "underline", "color", "size", "font")):
            font = Font(name=f.get("font"), size=f.get("size"), bold=f.get("bold"), italic=f.get("italic"), underline="single" if f.get("underline") else None, color=("FF" + f["color"][1:]) if f.get("color") else None)
        fill = PatternFill("solid", fgColor="FF" + f["fill"][1:]) if f.get("fill") else None
        align = Alignment(horizontal=f.get("halign"), vertical=f.get("valign"), wrap_text=f.get("wrap") or None) if any(f.get(k) for k in ("halign", "valign", "wrap")) else None
        border = None
        if f.get("borders"):
            sides = {}
            for side, b in f["borders"].items():
                if b:
                    sides[side] = Side(style=b[0], color=("FF" + b[1][1:]) if b[1] else "FF000000")
            if sides:
                border = Border(**sides)
        got = (font, fill, align, border, f.get("numfmt"))
        cache[key] = got
    font, fill, align, border, numfmt = got
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    if numfmt and cell.data_type != "s":
        cell.number_format = numfmt


def builtin_to_xlsx(src: Path, kind: str, out: Path, a: Any) -> dict[str, Any]:
    """Values (and, from .ods, formulas, names and formatting) into a fresh .xlsx with dates, merges and widths."""
    import openpyxl

    from _a1 import parse_range
    from _book import Workbook
    from _ops import autofit

    formulas: dict[str, dict[tuple[int, int], str]] = {}
    names: list[tuple[str, str, str | None]] = []
    notes = []
    fixed_arity = 0
    if kind == "ods":
        try:
            import _ods

            formulas = _ods.read_formulas(src)
            names = _ods.read_names(src)
            for sheet_f in formulas.values():
                for k, f in list(sheet_f.items()):
                    g = _ods.fix_lo_arity(f)
                    if g != f:
                        sheet_f[k] = g
                        fixed_arity += 1
        except Exception as e:  # noqa: BLE001
            notes.append(f"formulas could not be read ({e}); values only")
    elif kind in ("xls", "xlsb"):
        notes.append(f"values only: .{kind} formulas are not readable without LibreOffice (install it to keep formulas and formatting)")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n_formulas = 0
    ods_formats: dict[str, Any] = {}
    styled_cells = [0]
    if kind == "ods":
        try:
            import _ods

            limits = {}
            with Workbook(src) as book:
                for m in book.sheets():
                    if m.kind == "worksheet":
                        g0 = book.grid(m.name)
                        fm0 = formulas.get(m.name, {})
                        max_r = max([g0.row1 + len(g0.rows) - 1] + [k[0] for k in fm0])
                        max_c = max([g0.col1 + g0.width - 1] + [k[1] for k in fm0])
                        limits[m.name] = (max_r + 2, max_c + 2)
            ods_formats = _ods.read_formats(src, limits)
        except Exception as e:  # noqa: BLE001 — formatting is a bonus; values and formulas still convert
            notes.append(f"formatting could not be read ({type(e).__name__}); values and formulas only")
    with Workbook(src) as book:
        for m in book.sheets():
            if m.kind != "worksheet":
                continue
            if a.sheet and m.name.lower() != str(a.sheet).lower() and str(m.index + 1) != str(a.sheet):
                continue
            ws = wb.create_sheet(m.name[:31])
            if m.visible != "visible":
                ws.sheet_state = "hidden"
            g = book.grid(m.name)
            fm = formulas.get(m.name, {})
            for i, row in enumerate(g.rows):
                for j, v in enumerate(row):
                    r, c = g.row1 + i, g.col1 + j
                    f = fm.get((r, c))
                    if f:
                        from _formula import add_prefixes

                        ws.cell(r, c).value = add_prefixes(f)
                        n_formulas += 1
                        continue
                    if v is None:
                        continue
                    if hasattr(v, "code"):
                        v = v.code  # openpyxl stores error literals such as "#N/A" as error cells
                    cell = ws.cell(r, c)
                    cell.value = v
                    if isinstance(v, _dt.datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss" if v.second else "yyyy-mm-dd hh:mm"
                    elif isinstance(v, _dt.date):
                        cell.number_format = "yyyy-mm-dd"
                    elif isinstance(v, _dt.time):
                        cell.number_format = "hh:mm:ss"
                    elif isinstance(v, _dt.timedelta):
                        cell.value = v.total_seconds() / 86400
                        cell.number_format = "[h]:mm:ss"
            for (r, c), f in fm.items():
                if ws.cell(r, c).value is None:
                    from _formula import add_prefixes

                    ws.cell(r, c).value = add_prefixes(f)
                    n_formulas += 1
            merges = set()
            for mr in g.merged:
                try:
                    merges.add(parse_range(mr))
                except ValueError:
                    pass
            fmt = ods_formats.get(m.name) if kind == "ods" else None
            if fmt:
                merges |= set(fmt["merges"])
            for r1, c1, r2, c2 in sorted(merges):
                try:
                    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
                except ValueError:
                    pass
            if fmt:
                cache: dict[Any, Any] = {}
                for (r, c), f in fmt["cells"].items():
                    _apply_ods_format(ws, r, c, f, cache)
                    styled_cells[0] += 1
                from _a1 import col_letter as _cl

                for c, w in fmt["widths"].items():
                    ws.column_dimensions[_cl(c)].width = round(w, 2)
                for r, h in fmt["heights"].items():
                    ws.row_dimensions[r].height = round(h, 2)
                for c in fmt["hidden_cols"]:
                    ws.column_dimensions[_cl(c)].hidden = True
                for r in fmt["hidden_rows"]:
                    ws.row_dimensions[r].hidden = True
                if not fmt["widths"]:
                    autofit(ws)
            else:
                autofit(ws)
    if not wb.worksheets:
        raise SkillError("no worksheet to convert")
    if names:
        from openpyxl.workbook.defined_name import DefinedName

        from _formula import add_prefixes

        for name, formula, local in names:
            dn = DefinedName(name, attr_text=add_prefixes("=" + formula).lstrip("="))
            try:
                if local and local[:31] in wb.sheetnames:
                    wb[local[:31]].defined_names[name] = dn
                else:
                    wb.defined_names[name] = dn
            except (ValueError, KeyError):
                notes.append(f"name {name} could not be kept")
        notes.append(f"{len(names)} named ranges kept")
    from types import SimpleNamespace

    from _bridge import save_and_recalc

    ctx = SimpleNamespace(new_formulas=set(), autofit_pending=[], warnings=[])
    rep = save_and_recalc(wb, out, ctx, recalc=n_formulas > 0)
    if n_formulas:
        notes.append(f"{n_formulas} formulas kept and recalculated" + (f" ({sum(rep.get('error_counts', {}).values())} errors)" if rep.get("error_counts") else ""))
    if fixed_arity:
        notes.append(f"{fixed_arity} LibreOffice-only formula forms made Excel-valid (ROUND(x) → ROUND(x,0), CEILING(x) → CEILING(x,1))")
    if kind == "ods":
        if ods_formats:
            notes.append(f"formatting kept: number formats, fonts, fills, borders, alignment, widths, heights and merges ({styled_cells[0]:,} styled cells); conditional formats, charts and images are not (LibreOffice keeps them)")
    else:
        notes.append("formatting is not carried over from ." + kind + " by the built-in path (fonts, fills, borders); LibreOffice keeps it")
    return {"output": str(out), "engine": "builtin", "notes": notes}


def xlsx_to_ods(src: Path, out: Path, note: str | None = None) -> dict[str, Any]:
    """Built-in .xlsx → .ods: values, formulas (OpenFormula) with cached results, number formats, bold/italic, fills,
    alignment, merged cells and column widths."""
    import openpyxl

    import _ods
    from _a1 import col_index
    from _xlsx import load_book
    from _formula import Engine

    from _formula import structured_to_a1

    book, pkg, _extra = load_book(src)
    pkg.close()
    Engine(book).recalc()
    wb = openpyxl.load_workbook(src, rich_text=False)
    wb_names = list(wb.defined_names.keys())
    sheets = []
    untranslated = 0
    structured = 0
    dropped: dict[str, int] = {}
    try:
        for ws in wb.worksheets:
            osh = _ods.OdsSheet(ws.title)
            bsh = book.sheet(ws.title)
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if cell.__class__.__name__ == "MergedCell":
                        continue
                    f = None
                    if hasattr(v, "text") and v.__class__.__name__ == "ArrayFormula":
                        v = v.text
                    if isinstance(v, str) and v.startswith("="):
                        text = v
                        if "[" in v or any(t.lower() in v.lower() for t in book.tables):
                            s2 = structured_to_a1(v, book.tables, ws.title, cell.row, cell.column)
                            if s2 is not None and s2 != v:
                                text = s2
                                structured += 1
                        f = _ods.excel_to_odf(text)
                        fc = bsh.formulas.get((cell.row, cell.column)) if bsh else None
                        val = fc.value if fc is not None else None
                        if f is None:
                            untranslated += 1
                        v = val
                    if v is None and f is None and not (cell.font and cell.font.b) and not (cell.fill and cell.fill.fill_type == "solid"):
                        continue
                    fill = None
                    if cell.fill is not None and cell.fill.fill_type == "solid":
                        rgb = getattr(cell.fill.fgColor, "rgb", None)
                        if isinstance(rgb, str) and len(rgb) in (6, 8) and cell.fill.fgColor.type == "rgb":
                            fill = "#" + rgb[-6:]
                    oc = _ods.OdsCell(v, f, cell.number_format, bool(cell.font and cell.font.b), bool(cell.font and cell.font.i), fill, cell.alignment.horizontal if cell.alignment else None)
                    osh.cells[(cell.row, cell.column)] = oc
            if bsh is not None:
                for (r, c), fc in bsh.spilled.items():
                    if (r, c) not in osh.cells and fc.array_value is not None:
                        osh.cells[(r, c)] = _ods.OdsCell(fc.array_value.rows[r - fc.row][c - fc.col])
            for key, d in ws.column_dimensions.items():
                if d.width:
                    try:
                        lo, hi = (d.min or col_index(key)), (d.max or col_index(key))
                    except ValueError:
                        continue
                    for c in range(lo, min(hi, lo + 200) + 1):
                        osh.widths[c] = float(d.width)
            for m in ws.merged_cells.ranges:
                osh.merges.append((m.min_row, m.min_col, m.max_row, m.max_col))
            for label, n in (("Excel tables (their cells are kept)", len(ws.tables)), ("conditional formats", sum(len(cf.rules) for cf in ws.conditional_formatting)),
                             ("data validations", len(ws.data_validations.dataValidation)), ("charts", len(getattr(ws, "_charts", []))), ("images", len(getattr(ws, "_images", []))),
                             ("comments", sum(1 for row in ws.iter_rows() for c in row if c.comment is not None) if ws.max_row * ws.max_column < 200_000 else 0)):
                if n:
                    dropped[label] = dropped.get(label, 0) + n
            sheets.append(osh)
    finally:
        wb.close()
    names = [(n, f) for (scope, n), f in book.names.items() if scope is None]
    names = [(next((dn for dn in wb_names if dn.upper() == n), n), f) for n, f in names]
    _ods.write_ods(out, sheets, names)
    notes = ["built-in ODS writer: values, formulas with cached results, number formats, bold/italic, fills, merges and widths"]
    if note:
        notes.append(note)
    if structured:
        notes.append(f"{structured} formulas with structured references (Table[Column]) were rewritten as ranges")
    if untranslated:
        notes.append(f"{untranslated} formulas (external links or 3-D references) were stored as values")
    if dropped:
        notes.append("not written to .ods by the built-in writer: " + ", ".join(f"{n} {label}" for label, n in dropped.items()) + " (LibreOffice keeps them: --engine libreoffice)")
    return {"output": str(out), "engine": "builtin", "notes": notes}


# ── PDF ─────────────────────────────────────────────────────────────────


def to_pdf(src: Path, kind: str, out: Path, a: Any) -> dict[str, Any]:
    import _lo

    if a.engine != "builtin" and _lo.available():
        import tempfile

        notes = ["LibreOffice print layout (print areas, page setup, headers and footers)"]
        tmpdir = Path(tempfile.mkdtemp(prefix="desk-sheetpdf-"))
        try:
            work = src
            if (a.sheet or a.range) and kind == "xlsx":
                from _book import area_arg
                from sheet_render import _print_copy

                sheet_from_range, _area = area_arg(a.range)
                work = _print_copy(src, tmpdir / "print.xlsx", a.sheet or sheet_from_range, a.range, False)
            elif a.sheet or a.range:
                notes.append(f"--sheet/--range apply to .xlsx files only here: the whole .{kind} was printed")
            produced = _lo.convert(work, "pdf")
            try:
                shutil.move(str(produced), str(out))
            finally:
                shutil.rmtree(produced.parent, ignore_errors=True)
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)
        return {"output": str(out), "engine": "libreoffice", "notes": notes}
    if a.engine == "libreoffice":
        raise SkillError("LibreOffice is not installed; use --engine builtin")
    from _book import area_arg

    sheet_from_range, area = area_arg(a.range)
    pages = print_pages(src, a.sheet or sheet_from_range, area)
    if not pages:
        raise SkillError("nothing to print")
    try:
        pages[0].save(out, "PDF", resolution=float(PRINT_DPI), save_all=True, append_images=pages[1:])
    finally:
        for im in pages:
            im.close()
    return {"output": str(out), "engine": "builtin", "pages": len(pages), "notes": [
        f"{len(pages)} A4 page(s) drawn by the built-in renderer at {PRINT_DPI} dpi (cells, formats, borders, merges and charts, no gridlines or headings). "
        "The pages are images: text cannot be selected or searched, and print areas, page breaks, headers and footers are not applied. "
        "Install LibreOffice for Excel's print layout with real text."]}


PRINT_DPI = 150


def print_pages(src: Path, sheet: str | None, area: Any) -> list[Any]:
    """Sheets drawn as printable A4 pages (portrait, or landscape for wide sheets), down then across."""
    from PIL import Image, ImageDraw

    import _fonts
    from _book import Workbook
    from _grid import draw_page, load_view

    zoom = PRINT_DPI / 96
    margin = PRINT_DPI // 2
    portrait = (round(8.27 * PRINT_DPI), round(11.69 * PRINT_DPI))
    with Workbook(src) as wb:
        names = [wb.resolve_sheet(sheet).name] if sheet else [m.name for m in wb.sheets() if m.kind == "worksheet" and m.visible == "visible"] if wb.kind not in ("csv", "json") else [None]
    pages: list[Any] = []
    font, _ = _fonts.load("Calibri", round(8 * PRINT_DPI / 72), False, False)
    for name in names:
        v = load_view(src, name, area, 4000, 120, zoom)
        if not v.rows or not v.cols:
            continue
        v.gridlines = False
        width = sum(v.col_w.get(c, 64) for c in v.cols)
        size = portrait
        if width > portrait[0] - 2 * margin:
            size = (portrait[1], portrait[0])
        cw, ch = size[0] - 2 * margin, size[1] - 2 * margin - round(0.25 * PRINT_DPI)
        # up to 2.5 pages wide: shrink to one page wide (like "fit all columns on one page"), so charts and
        # tables are not cut between pages; wider sheets print across several pages
        fit = cw / width if cw < width <= 2.5 * cw else 1.0
        col_groups: list[list[int]] = []
        cur: list[int] = []
        w = 0
        for c in v.cols:
            x = v.col_w.get(c, 64)
            if cur and w + x > cw / fit:
                col_groups.append(cur)
                cur, w = [], 0
            cur.append(c)
            w += x
        if cur:
            col_groups.append(cur)
        cap = ch / fit
        heights = {r: v.row_h.get(r, 20) for r in v.rows}
        spans = []  # rows each drawing covers: a page break is moved above a drawing that would be cut
        for _im, br, _bc, _dx, dy, _w, bh in [(None, *ch_["box"]) for ch_ in v.charts] + list(v.images):
            acc, end = -dy, br
            for r in v.rows:
                if r < br:
                    continue
                acc += heights[r]
                end = r
                if acc >= bh:
                    break
            spans.append((br, end))
        row_groups: list[list[int]] = []
        cur, h = [], 0
        for r in v.rows:
            y = heights[r]
            if cur and h + y > cap:
                cut = len(cur)
                for s0, e0 in spans:
                    if s0 in cur and s0 != cur[0] and e0 >= r and sum(heights.get(x, 0) for x in range(s0, e0 + 1)) <= cap:
                        cut = min(cut, cur.index(s0))
                row_groups.append(cur[:cut])
                cur = cur[cut:]
                h = sum(heights[x] for x in cur)
            cur.append(r)
            h += y
        if cur:
            row_groups.append(cur)
        drawn = []
        for cols in col_groups:
            for rows in row_groups:
                img = draw_page(v, rows, cols, title=None, headers=False)
                if img.convert("L").getextrema()[0] >= 250:  # nothing on it (the space around a chart, say)
                    img.close()
                    continue
                if img.width > cw or img.height > ch + round(0.25 * PRINT_DPI):
                    scale = min(cw / img.width, (ch + round(0.25 * PRINT_DPI)) / img.height)
                    img = img.resize((max(1, int(img.width * scale)), max(1, int(img.height * scale))), Image.LANCZOS)
                drawn.append(img)
        for k, img in enumerate(drawn, 1):
            page = Image.new("RGB", size, "white")
            page.paste(img, (margin, margin))
            img.close()
            d = ImageDraw.Draw(page)
            label = f"{v.sheet} — page {k} of {len(drawn)}"
            d.text((margin, size[1] - margin + round(0.1 * PRINT_DPI)), label, fill=(120, 120, 120), font=font)
            pages.append(page)
    return pages


# ── text targets ────────────────────────────────────────────────────────


def _sheets(src: Path, a: Any) -> list[tuple[str, list[list[Any]], int, int]]:
    """(sheet, rows, first row, first col) for the chosen sheet(s)."""
    from _book import Workbook, area_arg, trim_grid

    sheet_from_range, area = area_arg(a.range)
    want = sheet_from_range or a.sheet
    out = []
    with Workbook(src) as wb:
        if wb.kind in ("csv", "json"):
            g = wb.grid(None, area)
            out.append((g.sheet, g.rows if area else trim_grid(g.rows), g.row1, g.col1))
            return out
        metas = [wb.resolve_sheet(want)] if want else [m for m in wb.sheets() if m.kind == "worksheet"]
        for m in metas:
            g = wb.grid(m.name, area)
            rows = g.rows if area else trim_grid(g.rows)
            if a.display and wb.kind == "xlsx" and rows:
                rows = _display_rows(src, m.name, rows, g.row1, g.col1)
            out.append((m.name, rows, g.row1, g.col1))
    return out


def _display_rows(src: Path, sheet: str, rows: list[list[Any]], r1: int, c1: int) -> list[list[Any]]:
    from _numfmt import format_value
    from _xlsx import Package

    last = r1 + len(rows) - 1
    fmts: dict[tuple[int, int], str] = {}
    with Package(src) as pkg:
        si = pkg.sheet(sheet)
        for r, c, _v, _f, style in pkg.iter_cells(si, want_formulas=False):
            if r > last:
                break
            if style:
                fmts[(r, c)] = pkg.number_format(style)
        d1904 = pkg.date1904
    out = []
    for i, row in enumerate(rows):
        new = []
        for j, v in enumerate(row):
            code = fmts.get((r1 + i, c1 + j))
            if v is None or hasattr(v, "code") or code in (None, "General"):
                new.append(v)
            else:
                new.append(format_value(v, code, d1904)[0])
        out.append(new)
    return out


DATA_TARGETS = {"csv", "tsv", "json", "jsonl", "parquet"}


def _data_sheets(src: Path, a: Any) -> list[str | None]:
    from _book import Workbook, area_arg

    sheet_from_range, _area = area_arg(a.range)
    want = sheet_from_range or a.sheet
    with Workbook(src) as wb:
        if wb.kind in ("csv", "json"):
            return [None]
        if want:
            return [wb.resolve_sheet(want).name]
        return [m.name for m in wb.sheets() if m.kind == "worksheet"]


def data_export(src: Path, out: Path, target: str, a: Any, sheets: list[str | None] | None = None) -> dict[str, Any]:
    """csv/tsv/json/jsonl/parquet from the typed Parquet copy of each sheet (DuckDB writes the file, streaming)."""
    import json as _json

    from _book import area_arg
    from _store import connect, ident, sheet_store, sql_str

    _s, area = area_arg(a.range)
    sheets = sheets if sheets is not None else _data_sheets(src, a)
    if target in ("csv", "tsv", "parquet") and len(sheets) > 1:
        raise UsageError(f"{src.name} has {len(sheets)} sheets: pick one with --sheet, or use --out-dir to write one file per sheet")
    metas = [sheet_store(src, s, a.header, area, a.keep_totals) for s in sheets]
    con = connect()
    notes: list[str] = []
    rows = 0

    def select(m: dict[str, Any], text: bool) -> str:
        cols = []
        for c in m["columns"]:
            q = ident(c["name"])
            if text and c["type"] == "DOUBLE":
                cols.append(f"CASE WHEN {q} IS NULL THEN NULL ELSE printf('%.15g', {q}) END AS {q}")
            elif text and c["type"] == "BOOLEAN":
                cols.append(f"CASE WHEN {q} IS NULL THEN NULL WHEN {q} THEN 'TRUE' ELSE 'FALSE' END AS {q}")
            else:
                cols.append(q)
        return f"SELECT {', '.join(cols) or 'NULL AS empty'} FROM read_parquet({sql_str(m['parquet'])}) ORDER BY {ident(m['row_column'])}"

    for m in metas:
        rows += m["rows"]
        label = f"{m['sheet']}: " if len(metas) > 1 else ""
        if m.get("preamble"):
            notes.append(f"{label}rows {m['preamble'][0][0]}-{m['preamble'][-1][0]} above the header (titles, notes) are not data and were left out; --range includes them")
        if m.get("excluded"):
            notes.append(f"{label}left out: " + "; ".join(f"row {e['row']} ({e['why']})" for e in m["excluded"][:5]) + (" — --keep-totals keeps totals rows" if not a.keep_totals else ""))
        if m.get("outlier_count"):
            notes.append(f"{label}{m['outlier_count']} cell(s) far outside the data block ({', '.join(x for x, _ in m['outliers'][:5])}) were not exported")
    try:
        if target in ("csv", "tsv"):
            m = metas[0]
            delim = a.delimiter or ("\t" if target == "tsv" else ",")
            header = "true" if m.get("header_row") else "false"
            dest = out
            if a.bom:
                import tempfile

                fd, tmpname = tempfile.mkstemp(suffix=".csv", dir=str(out.parent))
                import os

                os.close(fd)
                dest = Path(tmpname)
            con.execute(f"COPY ({select(m, True)}) TO {sql_str(dest)} (FORMAT csv, HEADER {header}, DELIMITER {sql_str(delim)}, QUOTE '\"')")
            if a.bom:
                with open(out, "wb") as fo, open(dest, "rb") as fi:
                    fo.write("\ufeff".encode("utf-8"))
                    shutil.copyfileobj(fi, fo, 4 << 20)
                dest.unlink()
        elif target == "parquet":
            con.execute(f"COPY ({select(metas[0], False)}) TO {sql_str(out)} (FORMAT parquet, COMPRESSION zstd)")
        elif len(metas) == 1 and metas[0].get("header_row"):
            fmt = "(FORMAT json, ARRAY true)" if target == "json" else "(FORMAT json)"
            con.execute(f"COPY ({select(metas[0], False)}) TO {sql_str(out)} {fmt}")
        else:
            from _book import json_value

            payload: dict[str, Any] = {}
            lines = []
            for m in metas:
                rel = con.sql(select(m, False))
                cols = rel.columns
                data = rel.fetchall()
                if m.get("header_row"):
                    recs: list[Any] = [{k: json_value(v) for k, v in zip(cols, r)} for r in data]
                else:
                    recs = [[json_value(v) for v in r] for r in data]
                payload[m.get("sheet") or src.stem] = recs
                for rec in recs:
                    lines.append(_json.dumps({"sheet": m.get("sheet"), **rec} if isinstance(rec, dict) and len(metas) > 1 else rec, ensure_ascii=False, default=str))
            if target == "json":
                data_out = next(iter(payload.values())) if len(metas) == 1 else payload
                out.write_text(_json.dumps(data_out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
            else:
                out.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    finally:
        con.close()
    return {"output": str(out), "engine": "builtin", "sheets": [m.get("sheet") or src.stem for m in metas], "rows": rows, "notes": notes}


def to_text(src: Path, out: Path, target: str, a: Any) -> dict[str, Any]:
    if target in DATA_TARGETS and not a.display:
        return data_export(src, out, target, a)
    sheets = _sheets(src, a)
    if target in ("csv", "tsv", "parquet") and len(sheets) > 1:
        raise UsageError(f"{src.name} has {len(sheets)} sheets: pick one with --sheet, or use --out-dir to write one file per sheet")
    write_text(out, target, sheets, a)
    return {"output": str(out), "engine": "builtin", "sheets": [s[0] for s in sheets], "rows": sum(len(s[1]) for s in sheets)}


def per_sheet(src: Path, folder: Path, target: str, a: Any) -> dict[str, Any]:
    from _book import file_stem, unique_stem

    d = output_dir(folder)
    written = []
    stems: set[str] = set()
    if target in DATA_TARGETS and not a.display:
        notes: list[str] = []
        for name in _data_sheets(src, a):
            stem = unique_stem(file_stem(name or src.stem), stems)
            path = output_path(d / f"{stem}.{target}", [src], a.force)
            r = data_export(src, path, target, a, [name])
            written.append(str(path))
            notes.extend(r.get("notes", []) if len(r.get("sheets", [])) != 1 else [f"{name}: {n}" if name else n for n in r.get("notes", [])])
        return {"outputs": written, "engine": "builtin", "notes": notes}
    for name, rows, r1, c1 in _sheets(src, a):
        stem = unique_stem(file_stem(name), stems)
        path = output_path(d / f"{stem}.{target}", [src], a.force)
        write_text(path, target, [(name, rows, r1, c1)], a)
        written.append(str(path))
    return {"outputs": written, "engine": "builtin"}


def write_text(out: Path, target: str, sheets: list[tuple[str, list[list[Any]], int, int]], a: Any) -> None:
    from _book import csv_text, detect_header, json_value, md_grid, text_value, unique_headers, Grid

    if target in ("csv", "tsv"):
        name, rows, _, _ = sheets[0]
        delim = a.delimiter or ("\t" if target == "tsv" else ",")
        text = csv_text(rows, delim)
        out.write_text(("\ufeff" if a.bom else "") + text, encoding="utf-8", newline="")
        return
    if target in ("json", "jsonl"):
        payload: dict[str, Any] = {}
        lines = []
        for name, rows, _, _ in sheets:
            header = a.header == "yes" or (a.header == "auto" and detect_header(rows))
            if header and rows:
                keys = unique_headers(rows[0])
                recs = [{k: json_value(v) for k, v in zip(keys, r + [None] * (len(keys) - len(r)))} for r in rows[1:]]
            else:
                recs = [[json_value(v) for v in r] for r in rows]
            payload[name] = recs
            for rec in recs:
                lines.append(json.dumps({"sheet": name, **rec} if isinstance(rec, dict) and len(sheets) > 1 else rec, ensure_ascii=False))
        if target == "json":
            data = payload[sheets[0][0]] if len(sheets) == 1 else payload
            out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        else:
            out.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
        return
    if target == "md":
        parts = []
        for name, rows, r1, c1 in sheets:
            header = a.header == "yes" or (a.header == "auto" and detect_header(rows))
            parts.append(f"## {name}\n\n" + md_grid(Grid(name, rows, r1, c1), header, show_coords=False))
        out.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
        return
    if target == "html":
        parts = ['<!DOCTYPE html><html><head><meta charset="utf-8"><title>' + html.escape(sheets[0][0]) + "</title>",
                 "<style>body{font-family:system-ui,sans-serif;margin:24px}table{border-collapse:collapse;margin-bottom:28px}"
                 "td,th{border:1px solid #d0d7de;padding:4px 8px;font-size:13px}th{background:#f3f4f6;text-align:left}"
                 "td.n{text-align:right;font-variant-numeric:tabular-nums}</style></head><body>"]
        for name, rows, _r1, _c1 in sheets:
            header = a.header == "yes" or (a.header == "auto" and detect_header(rows))
            parts.append(f"<h2>{html.escape(name)}</h2><table>")
            for i, r in enumerate(rows):
                tag = "th" if header and i == 0 else "td"
                cells = []
                for v in r:
                    cls = ' class="n"' if tag == "td" and isinstance(v, (int, float)) and not isinstance(v, bool) else ""
                    cells.append(f"<{tag}{cls}>{html.escape(text_value(v))}</{tag}>")
                parts.append("<tr>" + "".join(cells) + "</tr>")
            parts.append("</table>")
        parts.append("</body></html>")
        out.write_text("\n".join(parts), encoding="utf-8")
        return
    if target == "parquet":
        import tempfile

        from sheet_query import _load_rows

        name, rows, _, _ = sheets[0]
        tmp = Path(tempfile.mkdtemp(prefix="desk-pq-"))
        try:
            from _store import connect

            con = connect()  # memory-capped, no extension downloads, temp dirs only
            _load_rows(con, "t", rows, tmp, a.header)
            con.execute("COPY t TO ? (FORMAT parquet)", [str(out)])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        return
    raise UsageError(f"unsupported target .{target}")


def render(r: dict[str, Any]) -> str:
    if "outputs" in r:
        return f"Wrote {len(r['outputs'])} files:\n" + "\n".join(f"- {p}" for p in r["outputs"]) + "".join(f"\n- {n}" for n in r.get("notes", []))
    lines = [f"Wrote {r['output']} ({r.get('engine', 'builtin')})."]
    if r.get("sheets"):
        lines.append(f"Sheets: {', '.join(str(x) for x in r['sheets'])}; {r.get('rows', 0):,} rows.")
    for n in r.get("notes", []):
        lines.append(f"- {n}")
    return "\n".join(lines)


if __name__ == "__main__":
    run_main(main)
