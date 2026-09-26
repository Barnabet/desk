#!/usr/bin/env python3
"""Self-test for the spreadsheets skill: builds fixtures, runs every script as an agent would, checks the results.

Runs the LibreOffice-free path always (DESK_SOFFICE=none) and the LibreOffice path too when it is installed. The
formula battery holds 316 formulas whose expected results were cross-checked against LibreOffice (where LibreOffice
and Excel disagree, the Excel result is expected). Prints one summary line ("ok: N checks in S s") and exits
non-zero on any failure.

Example:
  python3 scripts/selftest.py
  python3 scripts/selftest.py --keep            # keep the temp folder and print its path
  python3 scripts/selftest.py --only render     # only the tests whose name contains "render"
"""

from __future__ import annotations

import ast
import datetime as dt
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
PY = sys.executable
sys.dont_write_bytecode = True  # never leave __pycache__ inside the skill (its helpers are imported in-process too)
NO_LO = {"DESK_SOFFICE": "none"}
SCRIPTS = ["sheet_info.py", "sheet_read.py", "sheet_query.py", "sheet_create.py", "sheet_edit.py", "sheet_recalc.py", "sheet_render.py", "sheet_convert.py"]
HEAVY = {"openpyxl", "duckdb", "PIL", "python_calamine", "pypdfium2"}

# formula <TAB> expected result (Python literal); Data sheet layout in battery_spec()
BATTERY = """
SUM(Data!A1:A10)	55
SUM(Data!A:A)	55
SUMPRODUCT(Data!A1:A5,Data!B1:B5)	121
PRODUCT(Data!A1:A4)	24
SUMSQ(1,2,3)	14
ABS(-3.5)	3.5
SIGN(-2)	-1
SQRT(16)	4
POWER(2,10)	1024
EXP(1)	2.71828182846
LN(10)	2.30258509299
LOG(100)	2
LOG(8,2)	3
LOG10(1000)	3
MOD(10,3)	1
MOD(-10,3)	2
MOD(10,-3)	-2
MOD(5.5,2)	1.5
QUOTIENT(-10,3)	-3
INT(-2.5)	-3
INT(2.5)	2
TRUNC(-2.567,2)	-2.56
ROUND(2.675,2)	2.68
ROUND(-2.5,0)	-3
ROUND(1234.5678,-2)	1200
ROUNDUP(3.14159,3)	3.142
ROUNDDOWN(-3.14159,2)	-3.14
MROUND(10,3)	9
MROUND(-10,-3)	-9
CEILING(2.5,1)	3
CEILING(-2.5,2)	-2
FLOOR(2.5,1)	2
FLOOR(-2.5,-2)	-2
CEILING.MATH(-2.5)	-2
CEILING.MATH(-2.5,1,1)	-3
FLOOR.MATH(-2.5)	-3
FLOOR.MATH(-2.5,1,1)	-2
EVEN(3)	4
EVEN(-1.5)	-2
ODD(2)	3
ODD(-2.1)	-3
FACT(5)	120
FACTDOUBLE(7)	105
COMBIN(10,3)	120
PERMUT(10,3)	720
GCD(24,36,48)	12
LCM(4,6,8)	24
PI()	3.14159265359
SIN(1)	0.841470984808
COS(1)	0.540302305868
TAN(1)	1.55740772465
ASIN(0.5)	0.523598775598
ACOS(0.5)	1.0471975512
ATAN(1)	0.785398163397
ATAN2(1,1)	0.785398163397
SINH(1)	1.17520119364
COSH(1)	1.54308063482
TANH(0.5)	0.46211715726
DEGREES(PI())	180
RADIANS(180)	3.14159265359
AVERAGE(Data!A1:A10)	5.5
AVERAGEA(Data!A1:C3)	2.11111111111
COUNT(Data!A1:C10)	15
COUNTA(Data!A1:C10)	25
COUNTBLANK(Data!A1:C10)	5
MAX(Data!A1:A10)	10
MIN(Data!A1:A10)	1
MEDIAN(Data!A1:A10)	5.5
MODE(1,2,2,3,3,3)	3
STDEV(Data!A1:A10)	3.0276503541
STDEV.P(Data!A1:A10)	2.87228132327
VAR(Data!A1:A10)	9.16666666667
VAR.P(Data!A1:A10)	8.25
PERCENTILE(Data!A1:A10,0.3)	3.7
PERCENTILE.EXC(Data!A1:A10,0.3)	3.3
QUARTILE(Data!A1:A10,1)	3.25
QUARTILE.EXC(Data!A1:A10,3)	8.25
RANK(Data!A3,Data!A1:A10)	8
RANK(Data!A3,Data!A1:A10,1)	3
RANK.AVG(3,{1,3,3,5})	2.5
LARGE(Data!A1:A10,2)	9
SMALL(Data!A1:A10,3)	3
CORREL(Data!A1:A5,Data!B1:B5)	0.996915159745
SLOPE(Data!B1:B5,Data!A1:A5)	2.2
INTERCEPT(Data!B1:B5,Data!A1:A5)	-8.881784197e-16
RSQ(Data!B1:B5,Data!A1:A5)	0.993839835729
FORECAST(6,Data!B1:B5,Data!A1:A5)	13.2
GEOMEAN(1,2,4,8)	2.82842712475
HARMEAN(1,2,4)	1.71428571429
AVEDEV(1,2,3,4,10)	2.4
DEVSQ(1,2,3,4)	5
KURT(1,2,3,4,10,2)	4.055859375
SKEW(1,2,3,4,10)	1.69705627485
TRIMMEAN(Data!A1:A10,0.2)	5.5
NORM.DIST(1,0,1,TRUE)	0.841344746069
NORM.DIST(1,0,1,FALSE)	0.241970724519
NORM.INV(0.9,10,2)	12.5631031311
NORM.S.INV(0.975)	1.95996398454
NORMSDIST(1.5)	0.933192798731
STANDARDIZE(5,3,2)	1
PERCENTRANK(Data!A1:A10,4)	0.333
EXPON.DIST(1,2,TRUE)	0.864664716763
POISSON.DIST(3,2.5,TRUE)	0.757576133133
BINOM.DIST(3,10,0.3,FALSE)	0.266827932
T.DIST(1.5,10,TRUE)	0.917746336777
T.DIST.2T(2,10)	0.0733880347707
T.INV(0.9,5)	1.47588404882
T.INV.2T(0.05,10)	2.22813885199
CHISQ.DIST(3,4,TRUE)	0.442174599629
CHISQ.INV.RT(0.05,3)	7.81472790325
F.DIST(2,3,10,TRUE)	0.821992592625
F.DIST.RT(2,3,10)	0.178007407375
CONFIDENCE.NORM(0.05,2,50)	0.55436152974
GAMMALN(5.5)	3.95781396762
SUMIF(Data!C1:C10,"x",Data!A1:A10)	18
SUMIF(Data!A1:A10,">3")	49
SUMIFS(Data!A1:A10,Data!C1:C10,"x",Data!A1:A10,">2")	17
COUNTIF(Data!C1:C10,"x*")	4
COUNTIF(Data!C1:C10,"<>x")	6
COUNTIF(Data!A1:A10,"")	0
COUNTIFS(Data!C1:C10,"x",Data!A1:A10,"<=5")	2
AVERAGEIF(Data!C1:C10,"y",Data!A1:A10)	5.33333333333
AVERAGEIFS(Data!A1:A10,Data!C1:C10,"?")	4.75
MAXIFS(Data!A1:A10,Data!C1:C10,"x")	8
MINIFS(Data!A1:A10,Data!C1:C10,"y")	2
COUNTIF(Data!D1:D10,">="&DATE(2024,3,1))	3
IF(Data!A1>1,"a","b")	'b'
IFS(Data!A1>5,"a",Data!A1>0,"b")	'b'
IFERROR(1/0,"e")	'e'
IFNA(NA(),"n")	'n'
SWITCH(2,1,"one",2,"two","other")	'two'
AND(TRUE,1,Data!E1:E3)	True
OR(FALSE,0)	False
XOR(TRUE,TRUE,TRUE)	True
NOT(0)	True
VLOOKUP("y",Data!C1:D10,2,FALSE)	45336
VLOOKUP(4.5,Data!A1:B10,2)	9
VLOOKUP(4.5,Data!A1:B10,2,TRUE)	9
HLOOKUP("b",Data!F1:H2,2,FALSE)	2
MATCH("x",Data!C1:C10,0)	1
MATCH(5,Data!A1:A10,1)	5
MATCH(5,{9,7,5,3},-1)	3
INDEX(Data!A1:D10,3,2)	6.5
INDEX(Data!A1:D10,0,1)	'#VALUE!'
SUM(INDEX(Data!A1:D10,0,1))	55
CHOOSE(2,"a","b","c")	'b'
ROWS(Data!A1:D10)	10
COLUMNS(Data!A1:D10)	4
ROW(Data!B7)	7
COLUMN(Data!C1)	3
ADDRESS(3,4)	'$D$3'
ADDRESS(3,4,4)	'D3'
ADDRESS(3,4,1,TRUE,"Data")	'Data!$D$3'
SUM(OFFSET(Data!A1,1,0,3,1))	9
SUM(INDIRECT("Data!A1:A3"))	6
INDIRECT("Data!B2")	4
LOOKUP(4.5,Data!A1:A10,Data!B1:B10)	9
XLOOKUP("y",Data!C1:C10,Data!A1:A10)	2
XMATCH("y",Data!C1:C10)	2
XLOOKUP(99,Data!A1:A10,Data!B1:B10,"nf")	'nf'
XLOOKUP(4.5,Data!A1:A10,Data!B1:B10,,-1)	9
XLOOKUP(4.5,Data!A1:A10,Data!B1:B10,,1)	11
LEFT("Hello",2)	'He'
RIGHT("Hello",3)	'llo'
MID("Hello",2,3)	'ell'
LEN("Hello")	5
FIND("l","Hello")	3
SEARCH("L","Hello")	3
SEARCH("l?o","Hello")	3
SUBSTITUTE("a-b-c","-","+")	'a+b+c'
SUBSTITUTE("a-b-c","-","+",2)	'a-b+c'
REPLACE("abcdef",2,3,"X")	'aXef'
TEXT(1234.567,"#,##0.00")	'1,234.57'
TEXT(0.25,"0%")	'25%'
TEXT(45322,"dddd d mmm yyyy")	'Wednesday 31 Jan 2024'
TEXT(1/3,"0.000")	'0.333'
TEXT(-5,"0;(0)")	'(5)'
VALUE("1,234.5")	1234.5
VALUE("12%")	0.12
CONCATENATE("a","b")	'ab'
TEXTJOIN("-",TRUE,Data!C1:C5)	'x-y-x-z-y'
UPPER("abc")	'ABC'
LOWER("ABC")	'abc'
PROPER("hello wORLD")	'Hello World'
TRIM("  a   b  ")	'a b'
CLEAN("a"&CHAR(7)&"b")	'ab'
REPT("ab",3)	'ababab'
EXACT("a","A")	False
CHAR(65)	'A'
CODE("A")	65
NUMBERVALUE("1.234,5",",",".")	1234.5
FIXED(1234.567,1)	'1,234.6'
YEAR(Data!D1)	2024
MONTH(Data!D1)	1
DAY(Data!D1)	31
HOUR(0.75)	18
MINUTE(0.7654)	22
SECOND(0.7654)	11
EDATE(Data!D1,1)	45351
EOMONTH(Data!D1,1)	45351
EOMONTH(Data!D1,-1)	45291
DATEDIF(Data!D1,Data!D5,"d")	329
DATEDIF(Data!D1,DATE(2025,6,15),"m")	16
DATEDIF(Data!D1,DATE(2026,6,15),"y")	2
DATEDIF(Data!D1,DATE(2025,6,15),"md")	15
DATEDIF(Data!D1,DATE(2025,6,15),"ym")	4
DATEDIF(Data!D1,DATE(2025,6,15),"yd")	135
WEEKDAY(Data!D1)	4
WEEKDAY(Data!D1,2)	3
WEEKDAY(Data!D1,3)	2
WEEKNUM(Data!D1)	5
WEEKNUM(Data!D1,2)	5
ISOWEEKNUM(Data!D1)	5
NETWORKDAYS(Data!D1,Data!D5)	236
NETWORKDAYS(Data!D1,Data!D5,Data!D2)	235
WORKDAY(Data!D1,10)	45336
WORKDAY(Data!D1,-3)	45317
DAYS(Data!D5,Data!D1)	329
DAYS360(Data!D1,Data!D5)	325
YEARFRAC(Data!D1,Data!D5)	0.902777777778
YEARFRAC(Data!D1,Data!D5,1)	0.898907103825
YEARFRAC(Data!D1,Data!D5,3)	0.901369863014
DATEVALUE("2024-03-15")	45366
TIMEVALUE("18:30")	0.770833333333
PMT(0.05/12,360,200000)	-1073.64324602
IPMT(0.05/12,1,360,200000)	-833.333333333
PPMT(0.05/12,1,360,200000)	-240.309912691
FV(0.04,10,-100,-1000)	2680.85499721
PV(0.05,10,-100)	772.173492918
NPV(0.1,-1000,300,400,500)	-19.1243767502
IRR({-1000,300,400,500})	0.0889633946933
RATE(10,-100,800)	0.0427749780351
NPER(0.05,-100,800)	10.4698484308
SLN(1000,100,5)	180
SYD(1000,100,5,2)	240
DDB(1000,100,5,2)	240
DB(1000,100,5,2)	232.839
CUMIPMT(0.05/12,360,200000,1,12,0)	-9932.98826116
CUMPRINC(0.05/12,360,200000,1,12,0)	-2950.73069113
EFFECT(0.05,12)	0.0511618978817
NOMINAL(0.05,12)	0.0488894854038
MIRR({-1000,300,400,500},0.1,0.12)	0.0981566924463
XNPV(0.1,{-1000,500,600},{45292,45474,45658})	22.1056745213
XIRR({-1000,500,600},{45292,45474,45658})	0.131822438637
ISBLANK(Data!A20)	True
ISNUMBER(Data!A1)	True
ISTEXT(Data!C1)	True
ISLOGICAL(TRUE)	True
ISERROR(1/0)	True
ISERR(NA())	False
ISNA(NA())	True
ISEVEN(4)	True
ISODD(4)	False
N("a")	0
N(TRUE)	1
TYPE("a")	2
TYPE({1,2})	64
ERROR.TYPE(1/0)	2
SUM(Data!A1:A3*Data!B1:B3)	'#VALUE!'
SUMPRODUCT((Data!C1:C10="x")*(Data!A1:A10>2))	3
SUMPRODUCT(--(Data!C1:C10="y"))	3
MAX(IF(Data!C1:C10="x",Data!A1:A10))	'#VALUE!'
SUM((Data!A1:A10>3)*1)	'#VALUE!'
"a"&1/3	'a0.333333333333333'
1/3*3=1	True
"10"+5	15
"abc"<"abd"	True
"A"="a"	True
TRUE+TRUE	2
2^-1	0.5
-2^2	4
10%	0.1
(1+2)*3	9
1E+20+1	100000000000000000000
0.1+0.2	0.3
"2024-01-31"+1	45323
AVERAGE(TRUE,1)	1
ROMAN(1999)	'MCMXCIX'
ARABIC("MCMXCIX")	1999
BASE(255,16)	'FF'
DECIMAL("FF",16)	255
DEC2BIN(10)	'1010'
DEC2HEX(-1)	'FFFFFFFFFF'
HEX2DEC("FFFFFFFFFF")	-1
BIN2DEC("1111111111")	-1
DEC2OCT(64)	'100'
BITAND(12,10)	8
BITOR(12,10)	14
BITXOR(12,10)	6
SUBTOTAL(9,Data!A1:A10)	55
SUBTOTAL(1,Data!A1:A10)	5.5
SUBTOTAL(109,Data!A1:A10)	55
AGGREGATE(9,6,Data!A1:A10)	55
AGGREGATE(14,6,Data!A1:A10,2)	9
MDETERM({1,2;3,4})	-2
SUMX2MY2({1,2},{3,4})	-20
SUMXMY2({1,2},{3,4})	8
FREQUENCY(Data!A1:A10,{2,5})	2
COVARIANCE.S(Data!A1:A5,Data!B1:B5)	5.5
STEYX(Data!B1:B5,Data!A1:A5)	0.316227766017
COMBINA(4,3)	20
SUM(SEQUENCE(4,1,2,2))	20
CONCAT("a",1,TRUE)	'a1TRUE'
"x"&TRUE	'xTRUE'
DOLLAR(-1234.5)	'($1,234.50)'
SUM("3",4)	7
PERCENTRANK.INC(Data!A1:A10,4.5)	0.388
NETWORKDAYS.INTL(Data!D1,Data!D5,11)	283
WORKDAY.INTL(Data!D1,5,"0000011")	45329
DATE(2024,2,29)	45351
DATE(2024,13,1)	45658
DATE(2024,1,0)	45291
TIME(12,30,0)	0.520833333333
DATE(2024,1,1)+0.5	45292.5
T(5)	''
INDEX(Data!A1:A10,MATCH(1,(Data!C1:C10="y")*(Data!A1:A10>3),0))	'#N/A'
"""


# ── fixtures ────────────────────────────────────────────────────────────


def battery_spec() -> dict:
    cells: dict[str, object] = {}
    for i, v in enumerate([1, 2, 3, 4, 5, 6, 7, 8, 9, 10], 1):
        cells[f"A{i}"] = v
    for i, v in enumerate([2.5, 4, 6.5, 9, 11], 1):
        cells[f"B{i}"] = v
    for i, v in enumerate(["x", "y", "x", "z", "y", "x", "", "x", "y", "zz"], 1):
        if v:
            cells[f"C{i}"] = v
    for i, v in enumerate(["2024-01-31", "2024-02-14", "2024-03-01", "2024-06-30", "2024-12-25"], 1):
        cells[f"D{i}"] = v
    cells.update({"E1": True, "E2": 1, "E3": "t", "F1": "a", "G1": "b", "H1": "c", "F2": 1, "G2": 2, "H2": 3, "B7": "bee"})
    calc = {f"B{i}": "=" + f for i, (f, _) in enumerate(battery(), 1)}
    return {"sheets": [{"name": "Calc", "cells": calc}, {"name": "Data", "cells": cells}]}


def battery() -> list[tuple[str, object]]:
    out = []
    for line in BATTERY.strip().splitlines():
        f, v = line.split("\t")
        out.append((f, ast.literal_eval(v)))
    return out


def write_fancy(path: Path) -> Path:
    """Formatting the renderer must draw: fonts, fills, borders, merges, wrap, rotation, number formats, frozen panes,
    hidden rows/columns, a comment, a chart and a chart sheet."""
    import openpyxl
    from openpyxl.chart import LineChart, PieChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fancy"
    ws["A1"] = "Quarterly report"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws["A1"].fill = PatternFill("solid", fgColor="FFC000")
    ws.merge_cells("A3:C3")
    ws["A3"] = "Merged"
    ws["A3"].alignment = Alignment(horizontal="center")
    ws["A4"] = "Wrapped text in a narrow cell that needs several lines"
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B4"] = -1234.5
    ws["B4"].number_format = "#,##0.00;[Red](#,##0.00)"
    ws["C4"] = 0.4567
    ws["C4"].number_format = "0.0%"
    ws["D4"] = dt.date(2024, 3, 15)
    ws["D4"].number_format = "yyyy-mm-dd"
    ws["F4"] = "Rotated"
    ws["F4"].alignment = Alignment(text_rotation=90)
    ws["B5"] = 3.14159
    ws["C5"] = True
    ws["D5"] = "=1/0"
    ws["E5"] = "=SUM(B4:C5)"
    ws["B8"] = "borders"
    ws["B8"].border = Border(left=Side(style="thin"), right=Side(style="medium", color="C00000"), top=Side(style="dashed"), bottom=Side(style="double"))
    ws["A9"] = "hidden row below"
    ws["A10"] = "HIDDEN"
    ws.row_dimensions[10].hidden = True
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["G"].hidden = True
    ws["G4"] = "hidden col"
    ws["H5"] = "note"
    ws["H5"].comment = Comment("a comment", "me")
    ws.freeze_panes = "B4"
    ws["A15"], ws["B15"], ws["C15"] = "Month", "North", "South"
    for i, (m, a, b) in enumerate([("Jan", 10, 5), ("Feb", 14, 7), ("Mar", 9, 12), ("Apr", 17, 11)], 16):
        ws.cell(i, 1, m)
        ws.cell(i, 2, a)
        ws.cell(i, 3, b)
    ch = LineChart()
    ch.title = "Trend"
    ch.add_data(Reference(ws, min_col=2, min_row=15, max_col=3, max_row=19), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=16, max_row=19))
    ws.add_chart(ch, "E15")
    pie = PieChart()
    pie.title = "Share"
    pie.add_data(Reference(ws, min_col=2, min_row=15, max_row=19), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=16, max_row=19))
    cs = wb.create_chartsheet("PieSheet")
    cs.add_chart(pie)
    wb.save(path)
    return path


def _patch_zip(src: Path, dest: Path, edit) -> None:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for n in names:
            data = edit(n, zin.read(n))
            if data is not None:
                zout.writestr(n, data)
        for n, data in edit.extra.items():
            zout.writestr(n, data)


def write_features(path: Path, tmp: Path) -> Path:
    """Cached values (one stale), a function the engine does not know (with its cached result), a cycle, and a
    custom XML part that only the patch mode keeps."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["A2"], ws["A3"] = 2, "=A1*2", "=A2+1"
    ws["B1"] = '=CUBEVALUE("Sales","[Measures].[Total]")'
    ws["C1"], ws["C2"] = "=C2+1", "=C1+1"
    ws["D1"] = "=TODAY()"
    ws["E1"] = "input"
    raw = tmp / "features-raw.xlsx"
    wb.save(raw)
    cached = {"A2": "5", "A3": "5", "B1": "42"}  # A2's cache is stale on purpose (2*2 = 4)

    def edit(name: str, data: bytes):
        if name == "xl/worksheets/sheet1.xml":
            text = data.decode()
            for ref, v in cached.items():
                text = re.sub(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)<v\s*/>', rf"\g<1><v>{v}</v>", text)
            return text.encode()
        if name == "[Content_Types].xml":
            return data.replace(b"</Types>", b'<Override PartName="/customXml/item1.xml" ContentType="application/xml"/></Types>')
        return data

    edit.extra = {"customXml/item1.xml": b'<?xml version="1.0"?><desk-selftest keep="yes"/>'}
    _patch_zip(raw, path, edit)
    raw.unlink()
    return path


def write_xlsm(path: Path, tmp: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 10, "=A1*3"
    raw = tmp / "macro-raw.xlsx"
    wb.save(raw)

    def edit(name: str, data: bytes):
        if name == "[Content_Types].xml":
            data = data.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml", b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            return data.replace(b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
        if name == "xl/_rels/workbook.xml.rels":
            return data.replace(b"</Relationships>", b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
        return data

    edit.extra = {"xl/vbaProject.bin": b"DESK-FAKE-VBA-PROJECT" * 64}
    _patch_zip(raw, path, edit)
    raw.unlink()
    return path


def build_fixtures(d: Path) -> dict[str, Path]:
    d.mkdir(parents=True, exist_ok=True)
    fx: dict[str, Path] = {}
    fx["csv"] = d / "sales.csv"
    fx["csv"].write_text(
        "Region,Month,Units,Price,Share,Date\n"
        "North,Jan,120,$9.50,12%,2024-01-31\n"
        "South,Jan,80,$12.00,8%,2024-01-31\n"
        "Île-de-France,Feb,140,$9.50,14%,2024-02-29\n"
        "East,Feb,60,$11.25,6%,2024-02-29\n"
        "West,Mar,95,$10.00,10%,2024-03-31\n",
        encoding="utf-8",
    )
    fx["semi"] = d / "semicolon.csv"
    fx["semi"].write_text("code;amount;note\n007;1,5;a\n042;2,25;b\n", encoding="utf-8")
    fx["targets"] = d / "targets.csv"
    fx["targets"].write_text("Region,Target\nNorth,100\nSouth,90\nEast,70\nWest,100\n", encoding="utf-8")
    fx["md"] = d / "report.md"
    fx["md"].write_text(
        "# Report\n\nSome text.\n\n| Item | Qty | Price |\n|---|---:|---:|\n| Pens | 10 | 1.5 |\n| Paper | 4 | 6 |\n\n"
        "## Staff\n\n| Name | Role |\n|---|---|\n| Ada | Lead |\n| Linus | Dev |\n",
        encoding="utf-8",
    )
    fx["json"] = d / "records.json"
    fx["json"].write_text(json.dumps([{"name": "a", "score": 1.5, "ok": True}, {"name": "b", "score": 2, "ok": False}]), encoding="utf-8")
    fx["fancy"] = write_fancy(d / "fancy.xlsx")
    fx["features"] = write_features(d / "features.xlsx", d)
    fx["xlsm"] = write_xlsm(d / "macro.xlsm", d)
    return fx


# ── harness ─────────────────────────────────────────────────────────────


class Checks:
    def __init__(self) -> None:
        import threading

        self.n = 0
        self.failures: list[str] = []
        self._lock = threading.Lock()

    def ok(self, cond: object, what: str) -> None:
        with self._lock:
            self.n += 1
            if not cond:
                self.failures.append(what)
                print(f"FAIL: {what}", file=sys.stderr)


def run(args: list, env: dict | None = None, expect: int = 0, cwd: Path | None = None, lo: bool = False, stdin: str | None = None) -> subprocess.CompletedProcess:
    """Runs scripts/<args[0]> as an agent would. LibreOffice is hidden unless lo=True. An expected failure must print
    one clean `error: …` line (no traceback)."""
    cmd = [PY, str(HERE / str(args[0])), *[str(a) for a in args[1:]]]
    e = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    if not lo:
        e.update(NO_LO)
    e.update(env or {})
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env=e, cwd=cwd, input=stdin, timeout=300)
    if r.returncode != expect:
        raise AssertionError(f"{' '.join(str(a) for a in args[:5])} … exited {r.returncode} (expected {expect}):\n{r.stderr[-1500:]}\n{r.stdout[-800:]}")
    if expect != 0 and ("Traceback" in r.stderr or "error:" not in r.stderr):
        raise AssertionError(f"{args[0]}: the error is not one clean message: {r.stderr[-800:]}")
    return r


def jrun(args: list, **kw) -> dict:
    return json.loads(run([*args, "--format", "json"], **kw).stdout)


def values(path: Path, sheet: str | None = None) -> dict[str, object]:
    """Cell → cached value, as any reader of the file sees it."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    finally:
        wb.close()


def formulas(path: Path, sheet: str | None = None) -> dict[str, str]:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        out = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if hasattr(v, "text"):
                    out[c.coordinate] = v.text
                elif isinstance(v, str) and v.startswith("="):
                    out[c.coordinate] = v
        return out
    finally:
        wb.close()


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def png_size(p: Path) -> tuple[int, int]:
    from PIL import Image

    with Image.open(p) as im:
        return im.size


def pdf_pages(p: Path) -> int:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(p))
    try:
        return len(doc)
    finally:
        doc.close()


def same(a: object, b: object) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b or a == b and type(a) is type(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9)
    if isinstance(a, dt.datetime) and isinstance(b, (int, float)):
        return False
    return str(a) == str(b)


def lo_available() -> bool:
    sys.path.insert(0, str(HERE))
    from _render import find_soffice

    return bool(find_soffice())


def main() -> int:
    import argparse
    from concurrent.futures import ThreadPoolExecutor

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--keep", action="store_true", help="keep the temp folder and print its path")
    ap.add_argument("--only", help="run only the tests whose name contains this")
    ap.add_argument("--no-lo", action="store_true", help="skip the LibreOffice tests even when it is installed")
    a = ap.parse_args()
    t0 = time.time()
    tmp = Path(tempfile.mkdtemp(prefix="spreadsheets-selftest-"))
    os.environ["DESK_FILE_CACHE"] = str(tmp / "cache")  # a fresh cache: cache hits are measured, the user's is untouched
    os.environ.pop("DESK_NO_CACHE", None)
    os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
    c = Checks()
    have_lo = False
    try:
        fx = build_fixtures(tmp / "fx")
        before = {k: sha(p) for k, p in fx.items()}
        have_lo = lo_available() and not a.no_lo
        tests = [(n, f) for n, f in globals().items() if n.startswith("test_") and callable(f) and (not a.only or a.only in n)]

        def one(item: tuple) -> None:
            name, fn = item
            work = tmp / name
            work.mkdir()
            try:
                fn(c, fx, work)
            except AssertionError as e:
                c.ok(False, f"{name}: {e}")
            except Exception as e:  # noqa: BLE001
                c.ok(False, f"{name}: {type(e).__name__}: {e}")

        plain = [t for t in tests if not t[0].startswith("test_lo_")]
        try:
            workers = max(1, int(os.environ.get("DESK_MAX_WORKERS", "4")))
        except ValueError:
            workers = 4
        with ThreadPoolExecutor(max_workers=min(workers, 4, os.cpu_count() or 2)) as pool:
            list(pool.map(one, plain))
        # LibreOffice converts one file at a time per profile: its tests run in sequence.
        if have_lo:
            for t in tests:
                if t[0].startswith("test_lo_"):
                    one(t)
        c.ok(all(sha(p) == before[k] for k, p in fx.items()), "no script modified an input file")
    finally:
        if a.keep:
            print(f"kept {tmp}", file=sys.stderr)
        else:
            shutil.rmtree(tmp, ignore_errors=True)
    took = time.time() - t0
    if c.failures:
        print(f"FAILED: {len(c.failures)} of {c.n} checks failed in {took:.1f}s", file=sys.stderr)
        return 1
    print(f"ok: {c.n} checks in {took:.1f}s" + ("" if have_lo else " (LibreOffice tests skipped: not installed or --no-lo)"))
    return 0


# ── tests ───────────────────────────────────────────────────────────────


def test_help(c: Checks, fx: dict, w: Path) -> None:
    for s in SCRIPTS:
        t = time.time()
        r = run([s, "--help"])
        took = time.time() - t
        c.ok("python3 scripts/" + s in r.stdout and "examples:" in r.stdout, f"{s} --help shows examples")
        c.ok(took < 2.0, f"{s} --help took {took:.2f}s")
        imp = subprocess.run([PY, "-X", "importtime", str(HERE / s), "--help"], capture_output=True, text=True, encoding="utf-8", env=dict(os.environ, **NO_LO), timeout=60)
        loaded = {ln.rsplit("|", 1)[-1].strip().split(".")[0] for ln in imp.stderr.splitlines() if ln.startswith("import time:")}
        c.ok(not (loaded & HEAVY), f"{s} --help imports no heavy module ({sorted(loaded & HEAVY)})")
    run(["sheet_read.py"], expect=2)
    run(["sheet_recalc.py", fx["features"]], expect=2)
    r = run(["sheet_edit.py", fx["features"], "--out", w / "x.xlsx", "--ops", '[{"op": "nope"}]'], expect=1)
    c.ok("unknown op" in r.stderr, "unknown op is named")


def test_engine_battery(c: Checks, fx: dict, w: Path) -> None:
    spec = w / "battery.json"
    spec.write_text(json.dumps(battery_spec()), encoding="utf-8")
    run(["sheet_create.py", w / "battery.xlsx", "--spec", spec, "--no-recalc"])
    info = jrun(["sheet_info.py", w / "battery.xlsx"])
    calc = next(s for s in info["sheets"] if s["name"] == "Calc")
    c.ok(calc["formulas"] == len(battery()) and calc.get("formulas_without_cached_values") == len(battery()), "battery written without cached values")
    rep = jrun(["sheet_recalc.py", w / "battery.xlsx", "--out", w / "calc.xlsx"])
    c.ok(rep["formulas"] == len(battery()) and not rep["unsupported"], f"battery recalculated, nothing unsupported: {rep['unsupported']}")
    got = values(w / "calc.xlsx", "Calc")
    for i, (f, want) in enumerate(battery(), 1):
        v = got.get(f"B{i}")
        if isinstance(v, str) and v.startswith("#") and isinstance(want, str):
            ok = v == want
        elif want == "" and v is None:
            ok = True
        else:
            ok = same(v, want)
        c.ok(ok, f"engine: ={f} gave {v!r}, expected {want!r}")
    r = run(["sheet_recalc.py", w / "calc.xlsx", "--check", "--compare"])
    c.ok("0 differ" in r.stdout, "stored results equal a fresh recalculation")


def test_engine_more(c: Checks, fx: dict, w: Path) -> None:
    """Database functions, more distributions, regular expressions and file forms of A1#, @ and LET (expected values
    checked against LibreOffice, except where Excel differs: DGET with several matches, text criteria as prefixes)."""
    db = [["Tree", "Height", "Age", "Yield", "Profit"], ["Apple", 18, 20, 14, 105], ["Pear", 12, 12, 10, 96], ["Cherry", 13, 14, 9, 105],
          ["Apple", 14, 15, 10, 75], ["Pear", 9, 8, 8, 76.8], ["Apple", 8, 9, 6, 45], ["Apricot", 10, 11, 7, 70], ["Peach", 11, 12, 8, 88]]
    crit = {"G1": "Tree", "H1": "Height", "G2": "Apple", "H2": ">10", "G3": "Pear", "G5": "Tree", "G6": "Apple", "G8": "Tree", "G9": {"value": "=Cherry", "text": True}, "I1": "Tree", "I2": "P"}
    cases = [
        ("LET(x,D2,y,x*2,x+y)", 42), ("SUM(K1#)", 201), ("@A2:A9", "Pear"),
        ('DSUM(A1:E9,"Profit",G1:H2)', 180), ("DSUM(A1:E9,4,G1:H3)", 42), ('DAVERAGE(A1:E9,"Yield",G5:G6)', 10), ('DCOUNT(A1:E9,"Age",G1:H2)', 2),
        ('DCOUNTA(A1:E9,"Tree",G5:G6)', 3), ('DMAX(A1:E9,"Profit",G1:H3)', 105), ('DSTDEV(A1:E9,"Yield",G5:G6)', 4), ('DGET(A1:E9,"Yield",G8:G9)', 9),
        ('DGET(A1:E9,"Yield",G5:G6)', "#NUM!"), ('DSUM(A1:E9,"Profit",I1:I2)', 260.8), ("FISHER(0.75)", 0.972955074527657),
        ("LOGNORM.DIST(4,3.5,1.2,TRUE)", 0.0390835557068005), ("GAMMA.DIST(10.00001131,9,2,FALSE)", 0.032639130418294), ("GAMMA.INV(0.068094,9,2)", 10.0000111914372),
        ("BETA.DIST(2,8,10,TRUE,1,3)", 0.685470581054688), ("BETA.INV(0.685470581,8,10,1,3)", 1.99999999996314), ("WEIBULL.DIST(105,20,100,TRUE)", 0.929581390069277),
        ("HYPGEOM.DIST(1,4,8,20,TRUE)", 0.465428276573787), ("NEGBINOM.DIST(10,5,0.25,TRUE)", 0.313514058478177), ("F.TEST(D2:D9,C2:C9)", 0.275757390873546),
        ("Z.TEST(D2:D9,5,2)", 7.70862895e-09), ("SERIESSUM(0.785398163,0,2,{1,-0.5,0.041666667,-0.001388889})", 0.707103215204654), ("MULTINOMIAL(2,3,4)", 1260),
        ('REGEXEXTRACT("Order 123, order 456","\\d+")', "123"), ('REGEXREPLACE("Order 123","(\\d+)","#$1")', "Order #123"), ('REGEXTEST("Order","ORDER",1)', True),
    ]
    cells = dict(crit)
    for i, row in enumerate(db, 1):
        for j, v in enumerate(row):
            cells[f"{'ABCDE'[j]}{i}"] = v
    cells["K1"] = "=SORT(E2:E3)"
    for i, (f, _) in enumerate(cases, 1):
        cells[f"M{i}"] = "=" + f
    p = w / "more.json"
    p.write_text(json.dumps({"sheets": [{"name": "DB", "cells": cells}]}), encoding="utf-8")
    out = w / "more.xlsx"
    run(["sheet_create.py", out, "--spec", p])
    v = values(out)
    for i, (f, want) in enumerate(cases, 1):
        got = v.get(f"M{i}")
        c.ok(same(got, want) or (isinstance(want, float) and isinstance(got, float) and math.isclose(got, want, rel_tol=1e-8)), f"engine: ={f} gave {got!r}, expected {want!r}")
    f = formulas(out)
    c.ok("_xlpm.x" in f.get("M1", "") and "_xlfn.ANCHORARRAY(K1)" in f.get("M2", "") and "_xlfn.SINGLE(" in f.get("M3", ""),
         f"LET parameters, A1# and @ stored in Excel's file form: {[f.get(f'M{k}') for k in (1, 2, 3)]}")
    c.ok(v.get("G9") == "=Cherry", "text starting with = stored as text when asked")
    j = jrun(["sheet_read.py", out, "--formulas", "--range", "M1:M3", "--header", "no"])
    shown = [x["formula"] for x in j["formulas"]]
    c.ok("=SUM(K1#)" in shown and "=@A2:A9" in shown and "=LET(x,D2,y,x*2,x+y)" in shown, f"formulas read back as written: {shown}")


def test_structured_refs(c: Checks, fx: dict, w: Path) -> None:
    spec = {"sheets": [{"name": "S", "columns": [{"header": "Item"}, {"header": "Qty"}, {"header": "Price"}, {"header": "Total", "formula": "=[@Qty]*[@Price]"}],
                        "rows": [["a", 2, 3], ["b", 4, 5]], "table": {"name": "Sales"},
                        "cells": {"G1": "=SUM(Sales[Total])", "G2": "=ROWS(Sales)", "G3": "=SUM(Sales[[#Totals],[Qty]])", "G4": "=INDEX(Sales,2,4)", "G5": "=COUNTA(Sales[#Headers])"}}]}
    out = w / "t.xlsx"
    rep = jrun(["sheet_create.py", out, "--spec", json.dumps(spec)])
    v = values(out)
    c.ok(v.get("D2") == 6 and v.get("D3") == 20, "[@Column] this-row references")
    c.ok(v.get("G1") == 26 and v.get("G2") == 2 and v.get("G4") == 20 and v.get("G5") == 4, f"table and column references: {[v.get(f'G{i}') for i in range(1, 6)]}")
    c.ok(v.get("G3") == "#REF!" and "G3" in json.dumps(rep.get("recalc", {})), "a totals reference without a totals row is #REF!, and reported")
    ins = jrun(["sheet_edit.py", out, "--out", w / "t2.xlsx", "--ops", '[{"op":"insert_rows","sheet":"S","at":3},{"op":"set","range":"S!A3","values":[["c",1,1]]}]'])
    v2 = values(w / "t2.xlsx")
    c.ok(v2.get("G2") == 3 and v2.get("G1") == 27 and v2.get("D3") == 1, f"a row inserted inside a table extends it: {v2.get('G1')}, {v2.get('G2')}, {v2.get('D3')}; {ins['ops']}")


def test_recalc_features(c: Checks, fx: dict, w: Path) -> None:
    rep = jrun(["sheet_recalc.py", fx["features"], "--check", "--compare", "--now", "2025-01-31"])
    c.ok(rep["cycles"] and any("C1" in " ".join(map(str, cy)) for cy in rep["cycles"]), f"cycle reported: {rep['cycles']}")
    unsupported = json.dumps(rep["unsupported"])
    c.ok("CUBEVALUE" in unsupported and "B1" in unsupported, f"unknown function reported by name and cell: {unsupported}")
    diffs = rep.get("compare", {})
    c.ok("A2" in json.dumps(diffs) and "A3" not in json.dumps(diffs.get("differences", diffs)), f"stale cached value found: {diffs}")
    run(["sheet_recalc.py", fx["features"], "--out", w / "f.xlsx", "--now", "2025-01-31"])
    v = values(w / "f.xlsx")
    c.ok(v.get("A2") == 4 and v.get("A3") == 5, "recalculated values stored")
    c.ok(v.get("B1") == 42, "unsupported function keeps its cached value")
    c.ok(v.get("D1") == 45688, f"TODAY() follows --now: {v.get('D1')}")
    with zipfile.ZipFile(w / "f.xlsx") as z:
        c.ok("xl/calcChain.xml" not in z.namelist() and "customXml/item1.xml" in z.namelist(), "recalc keeps other parts, drops the stale calc chain")
    r = run(["sheet_recalc.py", fx["features"], "--out", fx["features"]], expect=1)
    c.ok("input" in r.stderr or "same" in r.stderr or "exists" in r.stderr, "recalc refuses to overwrite its input")


def test_create_csv(c: Checks, fx: dict, w: Path) -> None:
    out = w / "sales.xlsx"
    r = run(["sheet_create.py", out, "--from", fx["csv"], "--table", "--totals", "--chart", "column"])
    c.ok("sales" in r.stdout and "view_image" in r.stdout, "create reports what it wrote and how to check it")
    info = jrun(["sheet_info.py", out])
    s = info["sheets"][0]
    c.ok(len(info["tables"]) == 1 and s.get("charts") == 1, f"table and chart created: {info['tables']}")
    c.ok(s["formulas"] >= 2 and not s.get("formulas_without_cached_values"), "totals formulas have cached values")
    j = jrun(["sheet_read.py", out, "--limit", "20"])
    rows = j["rows"]
    c.ok(j["header"] == ["Region", "Month", "Units", "Price", "Share", "Date"], f"header detected: {j['header']}")
    c.ok(rows[0][2] == 120 and rows[0][3] == 9.5 and abs(rows[0][4] - 0.12) < 1e-12 and rows[0][5] == "2024-01-31", f"typed values: {rows[0]}")
    c.ok(rows[2][0] == "Île-de-France", "UTF-8 text kept")
    total = rows[-1]
    c.ok(total[2] == 495, f"totals row sums units: {total}")
    disp = run(["sheet_read.py", out, "--display", "--range", "A1:F3"]).stdout
    c.ok("$9.50" in disp and "12%" in disp, "display formats: currency and percent kept")


def test_create_spec(c: Checks, fx: dict, w: Path) -> None:
    spec = {
        "sheets": [
            {
                "name": "Q1",
                "columns": [
                    {"header": "Item"},
                    {"header": "Qty", "format": "#,##0"},
                    {"header": "Price", "format": "$#,##0.00"},
                    {"header": "Total", "formula": "=B{row}*C{row}", "format": "$#,##0.00", "total": "sum"},
                ],
                "rows": [["Pens", 10, 1.5], ["Paper", 4, 6], ["Ink", 2, 25], ["Tape", 7, 2]],
                "conditional": [{"column": "Total", "type": "cell", "operator": ">", "value": 20, "fill": "#FFC7CE"}],
                "validations": [{"column": "Qty", "type": "whole", "operator": ">=", "value": 0}],
                "cells": {
                    "G1": "Tax",
                    "H1": 0.2,
                    "G2": "Grand total",
                    "H2": "=SUM(D2:D5)*(1+Tax)",
                    "G4": "Sorted",
                    "G5": "=SORT(A2:A5)",
                    "H5": "=XLOOKUP(\"Ink\",A2:A5,D2:D5)",
                    "I5": "=INDEX(A2:A5,MATCH(1,(B2:B5>5)*(C2:C5<2),0))",
                },
                "charts": [{"type": "column", "categories": "Item", "values": ["Total"], "title": "Totals"}],
            },
            {"name": "Notes", "cells": {"A1": "Link", "B1": "='Q1'!H2"}, "merge": ["A3:C3"]},
        ],
        "names": {"Tax": "Q1!$H$1"},
    }
    p = w / "spec.json"
    p.write_text(json.dumps(spec), encoding="utf-8")
    out = w / "q1.xlsx"
    rep = jrun(["sheet_create.py", out, "--spec", p])
    c.ok(not (rep.get("recalc") or {}).get("error_counts"), f"spec workbook computes without errors: {rep.get('recalc')}")
    v = values(out, "Q1")
    c.ok(v.get("D6") == 103 and same(v.get("H2"), 123.6), f"formulas with {{row}}, totals and a defined name: {v.get('D6')}, {v.get('H2')}")
    c.ok([v.get(f"G{r}") for r in range(5, 9)] == ["Ink", "Paper", "Pens", "Tape"], "dynamic SORT spilled into the file")
    c.ok(v.get("H5") == 50, "XLOOKUP")
    f = formulas(out, "Q1")
    c.ok(v.get("I5") == "Pens" and "I5" in f, f"array-style formula upgraded to an array formula and evaluated: {v.get('I5')}")
    c.ok(same(values(out, "Notes").get("B1"), 123.6), "cross-sheet reference")
    c.ok("_xlfn._xlws.SORT" in f.get("G5", "") and "_xlfn.XLOOKUP" in f.get("H5", ""), f"newer functions carry Excel's prefixes: {f.get('G5')}")
    with zipfile.ZipFile(out) as z:
        c.ok("xl/metadata.xml" in z.namelist(), "dynamic-array metadata written")
    info = jrun(["sheet_info.py", out])
    q1 = info["sheets"][0]
    c.ok(q1.get("conditional_formats") == 1 and q1.get("data_validations") == 1 and q1.get("charts") == 1, "conditional format, validation, chart")
    c.ok(any(n["name"] == "Tax" for n in info["names"]), "defined name")
    j = jrun(["sheet_read.py", out, "--formulas", "--range", "D2:I6"])
    fl = {x["range"]: x for x in j["formulas"]}
    c.ok(fl.get("D2:D5", {}).get("filled") == 4 and fl.get("G5", {}).get("dynamic_array"), f"formulas grouped, spills flagged: {list(fl)}")


def test_create_md_json(c: Checks, fx: dict, w: Path) -> None:
    out = w / "report.xlsx"
    run(["sheet_create.py", out, "--from", fx["md"]])
    info = jrun(["sheet_info.py", out])
    c.ok(len(info["sheets"]) == 2, f"one sheet per Markdown table: {[s['name'] for s in info['sheets']]}")
    j = jrun(["sheet_read.py", out, "--sheet", "2"])
    c.ok(j["rows"][0][0] == "Ada", f"second table: {j['rows'][:1]}")
    out2 = w / "records.xlsx"
    run(["sheet_create.py", out2, "--from", fx["json"], "--sheet-name", "Scores"])
    v = values(out2, "Scores")
    c.ok(v.get("A1") == "name" and v.get("B2") == 1.5 and v.get("C3") is False, f"JSON records: {v}")
    r = run(["sheet_create.py", out2, "--from", fx["json"]], expect=1)
    c.ok("--force" in r.stderr, "create refuses to overwrite without --force")


def test_odd_inputs(c: Checks, fx: dict, w: Path) -> None:
    """Inputs agents meet in the wild: an .xls that is really an HTML table, JSON with a sheet per key, a European CSV
    (semicolons, decimal commas) queried with SQL, and a what-if recalculation in a custom script."""
    html = w / "export.xls"
    html.write_text("<html><body><table><caption>Orders</caption><tr><th>Item</th><th>Qty</th></tr>"
                    "<tr><td>Pens</td><td>12</td></tr><tr><td colspan=2>Total 12</td></tr></table></body></html>", encoding="utf-8")
    info = jrun(["sheet_info.py", html])
    c.ok(info["sheets"][0]["name"] == "Orders" and "html" in json.dumps(info).lower(), f"an HTML table saved as .xls is recognised: {info['sheets'][:1]}")
    j = jrun(["sheet_read.py", html])
    c.ok(j["rows"][0] == ["Pens", 12] and "A3:B3" in json.dumps(j), f"HTML cells typed, colspan kept as a merge: {j['rows']} {j.get('merged')}")
    multi = w / "multi.json"
    multi.write_text(json.dumps({"North": [{"rep": "Ada", "sales": 10}], "South": [{"rep": "Bo", "sales": 7}, {"rep": "Cy", "sales": 3}]}), encoding="utf-8")
    out = w / "multi.xlsx"
    run(["sheet_create.py", out, "--from", multi])
    names = [s["name"] for s in jrun(["sheet_info.py", out])["sheets"]]
    c.ok(names == ["North", "South"] and values(out, "South").get("B3") == 3, f"JSON with a list per key gives a sheet per key: {names}")
    eu = w / "eu.csv"
    eu.write_text("Produit;Montant;Date\nA;1.234,50;2025-01-31\nB;99,95;2025-02-28\n", encoding="utf-8")
    q = jrun(["sheet_query.py", eu, "--sql", "select sum(Montant) as total from eu"])
    c.ok(abs(q["rows"][0][0] - 1334.45) < 1e-9, f"decimal commas and thousands dots typed for SQL: {q['rows']}")
    code = (
        "import sys; sys.path.insert(0, sys.argv[1])\n"
        "from _xlsx import load_book\nfrom _formula import Engine\n"
        "book, pkg, _ = load_book(sys.argv[2]); pkg.close()\n"
        "Engine(book).recalc(); a = book.sheet('Fancy').formulas[(5, 5)].value\n"
        "book.set_value(book.sheet('Fancy'), 4, 2, 1000)\n"
        "Engine(book).recalc(); print(a, book.sheet('Fancy').formulas[(5, 5)].value)\n"
    )
    r = subprocess.run([PY, "-c", code, str(HERE), str(fx["fancy"])], capture_output=True, text=True, timeout=120)
    got = r.stdout.split()
    c.ok(r.returncode == 0 and len(got) == 2 and math.isclose(float(got[1]) - float(got[0]), 1000 + 1234.5),
         f"a second Engine.recalc() after set_value recomputes (what-if): {r.stdout.strip()} {r.stderr[-300:]}")


def test_edit_structure(c: Checks, fx: dict, w: Path) -> None:
    base = w / "base.xlsx"
    spec = {
        "sheets": [
            {"name": "Data", "columns": [{"header": "k"}, {"header": "v"}], "rows": [[1, 10], [2, 20], [3, 30], [4, 40]],
             "cells": {"D1": "=SUM(B2:B5)", "D2": "=B5", "E1": "=COUNT(Total)"}, "merge": ["F2:G3"],
             "conditional": [{"range": "B2:B5", "type": "data_bar"}]},
            {"name": "Sum", "cells": {"A1": "=SUM(Data!B2:B5)", "A2": "=Data!B3"}},
        ],
        "names": {"Total": "Data!$B$2:$B$5"},
    }
    p = w / "spec.json"
    p.write_text(json.dumps(spec), encoding="utf-8")
    run(["sheet_create.py", base, "--spec", p])
    ops = [
        {"op": "insert_rows", "sheet": "Data", "at": 3, "count": 2},
        {"op": "set", "sheet": "Data", "range": "A3", "values": [[1.5, 100], [1.7, 200]]},
        {"op": "delete_cols", "sheet": "Data", "at": "C"},
        {"op": "rename_sheet", "sheet": "Data", "to": "Numbers 2025"},
    ]
    out = w / "edited.xlsx"
    rep = jrun(["sheet_edit.py", base, "--out", out, "--ops", json.dumps(ops)])
    c.ok(rep.get("mode") == "openpyxl" and len(rep["ops"]) == 4, "four operations applied")
    f = formulas(out, "Numbers 2025")
    c.ok(f.get("C1") == "=SUM(B2:B7)" and f.get("C2") == "=B7", f"formulas shifted by the row insert and column delete: {f}")
    s = formulas(out, "Sum")
    c.ok(s.get("A1") == "=SUM('Numbers 2025'!B2:B7)" and s.get("A2") == "='Numbers 2025'!B5", f"other sheets follow the rename and the shift: {s}")
    v = values(out, "Numbers 2025")
    c.ok(v.get("C1") == 400 and values(out, "Sum").get("A1") == 400, f"values recalculated after the edit: {v.get('C1')}")
    info = jrun(["sheet_info.py", out])
    c.ok(any(n["name"] == "Total" and "$B$2:$B$7" in n["refers_to"] and "Numbers 2025" in n["refers_to"] for n in info["names"]), f"defined name rewritten: {info['names']}")
    import openpyxl

    wb = openpyxl.load_workbook(out)
    ws = wb["Numbers 2025"]
    c.ok([str(m) for m in ws.merged_cells.ranges] == ["E2:F5"], f"merge grew with rows inserted inside it, moved with the deleted column: {ws.merged_cells.ranges}")
    c.ok([str(cf.sqref) for cf in ws.conditional_formatting] == ["B2:B7"], "conditional format range grew with the insert")
    wb.close()


def test_edit_ops(c: Checks, fx: dict, w: Path) -> None:
    base = w / "base.xlsx"
    run(["sheet_create.py", base, "--from", fx["csv"], "--no-style"])
    ops = [
        {"op": "style", "range": "A1:F1", "font": {"bold": True, "color": "#FFFFFF"}, "fill": "#1F4E78", "border": {"bottom": "medium"}},
        {"op": "style", "range": "D2:D6", "number_format": "€#,##0.00"},
        {"op": "conditional_format", "range": "C2:C6", "type": "color_scale"},
        {"op": "conditional_format", "range": "C2:C6", "type": "cell", "operator": ">", "value": 100, "fill": "#C6EFCE"},
        {"op": "validation", "range": "B2:B6", "type": "list", "values": ["Jan", "Feb", "Mar"]},
        {"op": "sort", "range": "A2:F6", "by": [{"column": "C", "order": "desc"}]},
        {"op": "find_replace", "find": "West", "replace": "Ouest", "range": "A1:A6"},
        {"op": "freeze", "cell": "B2"},
        {"op": "autofilter", "range": "A1:F6"},
        {"op": "chart", "type": "bar", "values": "C1:C6", "categories": "A2:A6", "anchor": "H2", "title": "Units"},
        {"op": "comment", "cell": "A1", "text": "Checked"},
        {"op": "hyperlink", "cell": "A8", "url": "https://example.com", "text": "source"},
        {"op": "copy_sheet", "sheet": "sales", "to": "Copy"},
        {"op": "add_sheet", "name": "Summary"},
        {"op": "set", "sheet": "Summary", "cell": "A1", "formula": "=SUM(sales!C2:C6)"},
        {"op": "table", "sheet": "Copy", "range": "A1:F6", "name": "CopyTable", "totals": {"Units": "sum"}},
        {"op": "column_width", "columns": "A", "width": 20},
        {"op": "print", "orientation": "landscape", "fit_width": 1},
        {"op": "protect", "sheet": "Copy"},
    ]
    out = w / "ops.xlsx"
    rep = jrun(["sheet_edit.py", base, "--out", out, "--sheet", "sales", "--ops", json.dumps(ops)])
    c.ok(len(rep["ops"]) == len(ops), "all operations applied")
    j = jrun(["sheet_read.py", out, "--sheet", "sales", "--styles", "--range", "A1:F6"])
    c.ok([r[2] for r in j["rows"] if isinstance(r[2], int)] == [140, 120, 95, 80, 60], f"sorted by units, descending: {[r[2] for r in j['rows']]}")
    c.ok(any(r[0] == "Ouest" for r in j["rows"]), "find and replace")
    st = json.dumps(j.get("styles", []))
    c.ok("1F4E78" in st.upper() and "bold" in st, "styles reported: header fill and bold")
    info = jrun(["sheet_info.py", out])
    s = {x["name"]: x for x in info["sheets"]}
    c.ok(s["sales"].get("conditional_formats") == 2 and s["sales"].get("data_validations") == 1, "rules added")
    c.ok(s["sales"].get("charts") == 1 and s["sales"].get("comments") == 1 and s["sales"].get("hyperlinks") == 1, "chart, comment, hyperlink")
    c.ok(s["sales"].get("frozen_at") == "B2" and s["sales"].get("autofilter") == "A1:F6", "freeze and filter")
    c.ok(s["Copy"].get("protected") and any(t["name"] == "CopyTable" and t["totals_row"] for t in info["tables"]), "copied sheet with a table and protection")
    c.ok(values(out, "Summary").get("A1") == 495, "new sheet formula computed")
    disp = run(["sheet_read.py", out, "--sheet", "sales", "--display", "--range", "D2:D3"]).stdout
    c.ok("€" in disp, "number format applied")
    r = run(["sheet_render.py", out, "--sheet", "sales", "--out", w / "png", "--format", "json"])
    img = json.loads(r.stdout)["images"][0]
    from PIL import Image

    with Image.open(img) as im:
        px = im.convert("RGB").getpixel((60, 28))
    c.ok(abs(px[0] - 0x1F) < 40 and abs(px[2] - 0x78) < 40, f"rendered header fill is dark blue: {px}")


def test_edit_patch_and_vba(c: Checks, fx: dict, w: Path) -> None:
    out = w / "patched.xlsx"
    rep = jrun(["sheet_edit.py", fx["features"], "--out", out, "--mode", "patch", "--ops", '[{"op":"set","cell":"A1","value":5},{"op":"set","cell":"E2","formula":"=A1*10"},{"op":"set","cell":"E3","value":"=text","text":true}]'])
    c.ok(rep.get("mode") == "patch", "patch mode used")
    with zipfile.ZipFile(fx["features"]) as a, zipfile.ZipFile(out) as b:
        c.ok(b.read("customXml/item1.xml") == a.read("customXml/item1.xml"), "patch mode keeps parts it does not touch")
    v = values(out)
    c.ok(v.get("A2") == 10 and v.get("A3") == 11 and v.get("E2") == 50 and v.get("B1") == 42, f"patched values recalculated: {v}")
    c.ok(v.get("E3") == "=text", "patch mode writes text that starts with = as text when asked")
    auto = jrun(["sheet_edit.py", fx["features"], "--out", w / "auto.xlsx", "--ops", '[{"op":"set","cell":"A1","value":3}]'])
    c.ok(auto.get("mode") == "patch", "auto mode picks patch when openpyxl would drop parts")
    m = w / "macro-v2.xlsm"
    run(["sheet_edit.py", fx["xlsm"], "--out", m, "--ops", '[{"op":"set","cell":"A1","value":7},{"op":"style","range":"A1","fill":"#FFFF00"}]'])
    with zipfile.ZipFile(fx["xlsm"]) as a, zipfile.ZipFile(m) as b:
        c.ok(b.read("xl/vbaProject.bin") == a.read("xl/vbaProject.bin"), ".xlsm keeps its VBA project")
    c.ok(values(m).get("A2") == 21, "macro workbook recalculated")
    c.ok(jrun(["sheet_info.py", m])["macros"] is True, "info reports macros")
    r = run(["sheet_edit.py", fx["features"], "--out", fx["features"], "--ops", '[{"op":"set","cell":"A1","value":1}]'], expect=1)
    c.ok("input" in r.stderr or "overwrite" in r.stderr or "same" in r.stderr, "edit never overwrites its input")
    r = run(["sheet_edit.py", fx["features"], "--out", out, "--ops", '[{"op":"set","cell":"A1","value":1}]'], expect=1)
    c.ok("--force" in r.stderr, "edit refuses an existing output without --force")


def test_read(c: Checks, fx: dict, w: Path) -> None:
    md = run(["sheet_read.py", fx["fancy"]]).stdout
    c.ok("Quarterly report" in md and "| 4 |" in md, "Markdown grid with row numbers")
    j = jrun(["sheet_read.py", fx["csv"], "--limit", "2"])
    c.ok(len(j["rows"]) == 2 and j.get("next_offset") == 2 and j["total_rows"] == 5, f"paging: {j.get('next_offset')} {j.get('total_rows')}")
    j2 = jrun(["sheet_read.py", fx["csv"], "--offset", "2", "--limit", "2"])
    c.ok(j2["rows"][0][0] == "Île-de-France" and j2["rows"][0][3] == 9.5 and abs(j2["rows"][0][4] - 0.14) < 1e-12, f"offset continues with typed values: {j2['rows'][:1]}")
    semi = jrun(["sheet_read.py", fx["semi"]])
    c.ok(semi["rows"][0][0] == "007" and semi["rows"][0][1] == 1.5, f"semicolon CSV with decimal commas, leading zeros kept: {semi['rows'][0]}")
    csvout = run(["sheet_read.py", fx["fancy"], "--range", "A15:C17", "--format", "csv"]).stdout
    c.ok(csvout.splitlines()[0].strip() == "row,Month,North,South" and "16,Jan,10,5" in csvout, f"CSV output with sheet row numbers: {csvout[:80]!r}")
    plain = run(["sheet_read.py", fx["fancy"], "--range", "A15:C17", "--format", "csv", "--no-coords"]).stdout
    c.ok(plain.splitlines()[0].strip() == "Month,North,South", "CSV output without coordinates")
    j = jrun(["sheet_read.py", fx["fancy"], "--range", "B4:E5", "--formulas"])
    c.ok(any(f["formula"] == "=SUM(B4:C5)" for f in j["formulas"]), "formulas listed")
    r = run(["sheet_read.py", fx["fancy"], "--sheet", "Nope"], expect=1)
    c.ok("Fancy" in r.stderr, "missing sheet names the sheets")
    j = jrun(["sheet_read.py", fx["fancy"], "--all", "--limit", "3"])
    c.ok(isinstance(j, (dict, list)), "all sheets")


def test_query(c: Checks, fx: dict, w: Path) -> None:
    base = w / "sales.xlsx"
    run(["sheet_create.py", base, "--from", fx["csv"]])
    t = jrun(["sheet_query.py", base, fx["targets"], "--tables"])
    names = {x["name"] for x in t["tables"]}
    c.ok({"sales_sales", "targets"} <= names or {"sales", "targets"} <= names, f"tables listed: {names}")
    tname = "sales_sales" if "sales_sales" in names else "sales"
    sql = f"select s.Region, sum(Units) as units, max(Target) as target from {tname} s join targets t using (Region) group by 1 order by units desc"
    j = jrun(["sheet_query.py", base, fx["targets"], "--sql", sql])
    c.ok(j["rows"][0][:2] == ["North", 120] and len(j["rows"]) == 4, f"join across a workbook and a CSV: {j['rows']}")
    out = w / "result.xlsx"
    run(["sheet_query.py", base, "--sql", "select Region, Units from sales where Units > 90", "--out", out])
    v = values(out)
    c.ok(v.get("A1") == "Region" and v.get("B2") in (120, 140), f"query result written as .xlsx: {v}")
    r = run(["sheet_query.py", base, "--sql", "select nope from nowhere"], expect=1)
    c.ok("nowhere" in r.stderr or "tables" in r.stderr, "SQL errors are reported with the table names")


def test_convert_builtin(c: Checks, fx: dict, w: Path) -> None:
    src = fx["fancy"]
    run(["sheet_convert.py", src, "--out", w / "f.csv", "--sheet", "Fancy"])
    c.ok("Quarterly report" in (w / "f.csv").read_text(encoding="utf-8"), "xlsx → csv")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.tsv"])
    c.ok("North\tJan\t120" in (w / "s.tsv").read_text(encoding="utf-8"), "csv → tsv")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.json"])
    recs = json.loads((w / "s.json").read_text(encoding="utf-8"))
    c.ok(recs[0]["Region"] == "North" and recs[0]["Units"] == 120, f"csv → json records: {recs[0]}")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.jsonl"])
    c.ok(len((w / "s.jsonl").read_text(encoding="utf-8").splitlines()) == 5, "csv → jsonl")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.md"])
    c.ok("| Region | Month |" in (w / "s.md").read_text(encoding="utf-8"), "csv → md")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.html"])
    c.ok("<th>Region</th>" in (w / "s.html").read_text(encoding="utf-8"), "csv → html")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.parquet"])
    j = jrun(["sheet_query.py", w / "s.parquet", "--sql", "select sum(Units) from s"])
    c.ok(j["rows"][0][0] == 495, "csv → parquet readable by sheet_query")
    run(["sheet_convert.py", fx["csv"], "--out", w / "s.xlsx"])
    c.ok(values(w / "s.xlsx").get("C2") == 120, "csv → xlsx")
    q1 = w / "q1.xlsx"
    run(["sheet_create.py", q1, "--spec", json.dumps({"sheets": [{"name": "Q1", "columns": [{"header": "a"}, {"header": "b"}, {"header": "c", "formula": "=A{row}*B{row}"}], "rows": [[1, 2], [3, 4]], "cells": {"E1": "=SUM(Rng)"}}], "names": {"Rng": "Q1!$C$2:$C$3"}})])
    r = run(["sheet_convert.py", q1, "--out", w / "q1.ods"])
    c.ok("builtin" in r.stdout, "xlsx → ods with the built-in writer")
    j = jrun(["sheet_read.py", w / "q1.ods", "--formulas"])
    c.ok([1, 2, 2] in [r[:3] for r in j["rows"]] and any(f["formula"] == "=A2*B2" for f in j["formulas"]), f"ods keeps values and formulas: {j['rows'][:2]} {j['formulas'][:2]}")
    run(["sheet_convert.py", w / "q1.ods", "--out", w / "back.xlsx"])
    f = formulas(w / "back.xlsx")
    c.ok(f.get("C3") == "=A3*B3" and f.get("E1") == "=SUM(Rng)" and values(w / "back.xlsx").get("E1") == 14, f"ods → xlsx keeps formulas and names: {f}")
    run(["sheet_convert.py", src, "--to", "csv", "--out-dir", w / "each"])
    c.ok(sorted(p.name for p in (w / "each").iterdir()) == ["Fancy.csv"], "one CSV per worksheet")
    run(["sheet_convert.py", src, "--out", w / "f.pdf"])
    c.ok((w / "f.pdf").read_bytes()[:4] == b"%PDF" and pdf_pages(w / "f.pdf") >= 1, "xlsx → pdf with the built-in renderer")
    r = run(["sheet_convert.py", src, "--out", w / "f.xls"], expect=1)
    c.ok("LibreOffice" in r.stderr, ".xls needs LibreOffice, said cleanly")


def test_render_builtin(c: Checks, fx: dict, w: Path) -> None:
    r = run(["sheet_render.py", fx["fancy"], "--all", "--out", w / "all"])
    c.ok("view_image" in r.stdout and "Fancy.png" in r.stdout and "PieSheet.png" in r.stdout, "every sheet rendered, chart sheet included")
    size = png_size(w / "all" / "Fancy.png")
    c.ok(max(size) <= 1568 and size[0] > 600, f"sized for vision: {size}")
    from PIL import Image

    with Image.open(w / "all" / "Fancy.png") as im:
        rgb = im.convert("RGB")
        c.ok(rgb.getpixel((40, 30)) == (255, 192, 0) or rgb.getpixel((60, 32)) == (255, 192, 0), "A1 fill drawn")
        colors = {col for _n, col in rgb.getcolors(maxcolors=1 << 22) or []}
    c.ok(any(abs(p[0] - 0x44) < 30 and abs(p[1] - 0x72) < 30 and abs(p[2] - 0xC4) < 30 for p in colors), "line chart drawn from its data")
    j = jrun(["sheet_render.py", fx["fancy"], "--range", "A1:F8", "--formulas", "--out", w / "f"])
    c.ok(j["sheets"][0]["range"] == "A1:F8" and any("outside" in n for n in j["notes"]), f"range with a note about the chart outside it: {j['notes']}")
    big = w / "big.csv"
    big.write_text("n,sq\n" + "".join(f"{i},{i * i}\n" for i in range(1, 400)), encoding="utf-8")
    j = jrun(["sheet_render.py", big, "--max-rows", "150", "--out", w / "big"])
    c.ok(len(j["images"]) >= 2 and all(max(png_size(Path(p))) <= 1568 for p in j["images"]), f"long ranges split into pages: {len(j['images'])}")
    r = run(["sheet_render.py", fx["fancy"], "--engine", "libreoffice", "--out", w / "lo"], expect=1)
    c.ok("LibreOffice" in r.stderr, "print-layout rendering without LibreOffice fails cleanly")
    tbl = w / "table.xlsx"
    run(["sheet_create.py", tbl, "--from", fx["csv"], "--table"])
    run(["sheet_render.py", tbl, "--out", w / "tbl"])
    with Image.open(next((w / "tbl").glob("*.png"))) as im:
        colors = {col for _n, col in im.convert("RGB").getcolors(maxcolors=1 << 22) or []}
    near = lambda want: any(all(abs(a - b) < 12 for a, b in zip(p, want)) for p in colors)  # noqa: E731
    # the file's theme decides the accent: openpyxl writes the Office 2007 theme (4F81BD), Excel 2013+ uses 4472C4
    c.ok((near((0x4F, 0x81, 0xBD)) and near((0xDC, 0xE6, 0xF2))) or (near((0x44, 0x72, 0xC4)) and near((0xD9, 0xE1, 0xF2))),
         "Excel table style drawn from the theme (header colour and row stripes)")


def test_info(c: Checks, fx: dict, w: Path) -> None:
    j = jrun(["sheet_info.py", fx["fancy"]])
    s = j["sheets"][0]
    c.ok(s["name"] == "Fancy" and s.get("merged") == 1 and s.get("frozen_at") == "B4" and s.get("charts") == 1 and s.get("comments") == 1, f"sheet details: {s}")
    c.ok(s.get("hidden", {}).get("rows") == 1 and j["sheets"][1].get("kind") == "chart", f"hidden rows and chart sheets: {j['sheets'][1]}")
    c.ok(s.get("formulas_without_cached_values") == 2 and any("sheet_recalc" in h for h in j["hints"]), "missing cached values hinted")
    c.ok(j["workbook_protected"] is False and j["macros"] is False, "no protection or macros")
    j = jrun(["sheet_info.py", fx["semi"]])
    c.ok(j["csv"].get("delimiter") == ";", f"CSV dialect: {j['csv']}")
    md = run(["sheet_info.py", fx["features"]]).stdout
    c.ok("CUBEVALUE" in md, "info names functions the engine does not know")


def test_no_lo_messages(c: Checks, fx: dict, w: Path) -> None:
    r = run(["sheet_recalc.py", fx["features"], "--out", w / "x.xlsx", "--engine", "libreoffice"], expect=1)
    c.ok("LibreOffice" in r.stderr, "LibreOffice recalculation without LibreOffice fails cleanly")
    r = run(["sheet_recalc.py", fx["features"], "--check", "--engine", "libreoffice"], expect=1)
    c.ok("LibreOffice" in r.stderr, "a LibreOffice check without LibreOffice fails cleanly (never claims it ran)")
    rep = jrun(["sheet_recalc.py", fx["features"], "--check", "--engine", "auto"])
    c.ok(rep["engine"] == "builtin" and any("not installed" in n for n in rep.get("notes", [])), f"--check --engine auto without LibreOffice names the engine that ran: {rep['engine']}")
    md = run(["sheet_recalc.py", fx["features"], "--check", "--engine", "auto"]).stdout
    c.ok("with the builtin engine" in md and "nothing written" in md, f"the check's text names the built-in engine: {md[:120]!r}")
    ods = w / "t.ods"
    run(["sheet_convert.py", fx["csv"], "--out", ods])
    j = jrun(["sheet_render.py", ods, "--out", w / "o"])
    c.ok(any("LibreOffice" in n for n in j["notes"]) and len(j["images"]) == 1, f"ods rendered from values, the note says why: {j['notes']}")


def test_lo_recalc_and_render(c: Checks, fx: dict, w: Path) -> None:
    rep = jrun(["sheet_recalc.py", fx["features"], "--out", w / "lo.xlsx", "--engine", "libreoffice", "--now", "2025-01-31"], lo=True)
    c.ok(rep["engine"] == "libreoffice", "LibreOffice recalculation")
    # --check runs LibreOffice too, on a temporary copy: its own results (CUBEVALUE is #NAME? there, A2 is 4)
    rep = jrun(["sheet_recalc.py", fx["features"], "--check", "--compare", "--engine", "libreoffice"], lo=True, cwd=w)
    got = {e["cell"]: e["error"] for e in rep.get("errors", [])}
    diff = {d["cell"]: d["computed"] for d in (rep.get("compare") or {}).get("cells", [])}
    c.ok(rep["engine"] == "libreoffice" and got.get("B1") == "#NAME?" and diff.get("Model!A2") == 4 and any("nothing was written" in n for n in rep["notes"]), f"--check --engine libreoffice reports LibreOffice's results: {got} {diff}")
    rep = jrun(["sheet_recalc.py", fx["features"], "--check", "--engine", "auto"], lo=True, cwd=w)
    got = {e["cell"]: e["error"] for e in rep.get("errors", [])}
    c.ok(rep["engine"] == "libreoffice" and got.get("B1") == "#NAME?" and any("CUBEVALUE" in n for n in rep["notes"]), f"--check --engine auto runs LibreOffice for functions the engine lacks (not a cached built-in report): {rep['engine']} {rep['notes'][:1]}")
    c.ok(sorted(p.name for p in w.iterdir()) == ["lo.xlsx"], f"a LibreOffice check writes nothing: {sorted(p.name for p in w.iterdir())}")
    v = values(w / "lo.xlsx")
    c.ok(v.get("A2") == 4 and v.get("A3") == 5, f"LibreOffice results stored: {v.get('A2')}, {v.get('A3')}")
    j = jrun(["sheet_render.py", fx["fancy"], "--engine", "libreoffice", "--range", "A1:F10", "--out", w / "print"], lo=True)
    c.ok(j["engine"].startswith("LibreOffice") and j["images"] and max(png_size(Path(j["images"][0]))) <= 1568, "print layout rendered with LibreOffice")


def test_lo_convert(c: Checks, fx: dict, w: Path) -> None:
    q = w / "q.xlsx"
    run(["sheet_create.py", q, "--spec", json.dumps({"sheets": [{"name": "S", "columns": [{"header": "a"}, {"header": "b", "formula": "=A{row}*2"}], "rows": [[1], [2], [3]]}]})])
    r = run(["sheet_convert.py", q, "--out", w / "q.xls"], lo=True)
    c.ok("libreoffice" in r.stdout, "xlsx → xls with LibreOffice")
    run(["sheet_convert.py", w / "q.xls", "--out", w / "back.xlsx"], lo=True)
    c.ok(formulas(w / "back.xlsx").get("B3") == "=A3*2" and values(w / "back.xlsx").get("B4") == 6, ".xls → .xlsx keeps formulas and values")
    j = jrun(["sheet_render.py", w / "q.xls", "--out", w / "xls"], lo=True)
    c.ok("LibreOffice" in j["engine"] and len(j["images"]) == 1, "legacy .xls rendered with its formatting through LibreOffice")
    run(["sheet_convert.py", fx["fancy"], "--out", w / "f.pdf", "--sheet", "Fancy", "--range", "A1:F10"], lo=True)
    c.ok(pdf_pages(w / "f.pdf") == 1, "PDF of one range via LibreOffice")


def _orders_book(path: Path, n: int) -> float:
    """A scaled-down big workbook (sheet Orders: n rows × 8 columns, a formula column without cached values, one
    needle) plus a small Notes sheet. Returns the Revenue total."""
    import openpyxl

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Orders")
    ws.append(["Order", "Date", "Region", "Units", "Price", "Revenue", "Code", "Score"])
    regions = ["North", "South", "East", "West"]
    base = dt.date(2024, 1, 1)
    total = 0.0
    for i in range(n):
        r = i + 2
        units, price = i % 50 + 1, round(1 + (i * 7 % 100) / 4, 2)
        total += units * price
        code = "NEEDLE-42" if r == 9876 else f"C{i:05d}"
        ws.append([100000 + i, base + dt.timedelta(days=i % 365), regions[i % 4], units, price, f"=D{r}*E{r}", code, (i * 37 % 1000) / 10])
    notes = wb.create_sheet("Notes")
    notes.append(["Read me first"])
    wb.save(path)
    return total


def test_big_files(c: Checks, fx: dict, w: Path) -> None:
    n = 12000
    raw = w / "raw.xlsx"
    total = _orders_book(raw, n)
    # recalculation: the report is cached per file content
    t = time.time()
    rep = jrun(["sheet_recalc.py", raw, "--check"])
    cold = time.time() - t
    t = time.time()
    rep2 = jrun(["sheet_recalc.py", raw, "--check"])
    warm = time.time() - t
    c.ok(rep["formulas"] == n and not rep.get("error_counts"), f"big workbook recalculated: {rep.get('formulas')} {rep.get('error_counts')}")
    c.ok(any("Reused" in x for x in rep2.get("notes", [])) and rep2["formulas"] == n, "second --check reuses the cached report")
    c.ok(cold >= 5 * warm, f"cached recalculation report at least 5x faster: {cold:.2f}s then {warm:.2f}s")
    big = w / "big.xlsx"
    run(["sheet_recalc.py", raw, "--out", big])
    # map first, with commands to go further
    md = run(["sheet_read.py", big]).stdout
    c.ok("map" in md and "Next:" in md and "--rows" in md and "Other sheets" in md and "Notes" in md, f"a big sheet reads as a map first: {md[:300]!r}")
    c.ok("| F | Revenue | number |" in md and "| G | Code | text |" in md, "the map types every column")
    # find → addresses
    f = run(["sheet_read.py", big, "--find", "needle-42"]).stdout
    c.ok("G9876" in f, f"--find gives the cell address: {f[:300]!r}")
    g = run(["sheet_read.py", big, "--grep", r"^NEEDLE-\d+$", "--format", "json"]).stdout
    c.ok("G9876" in g, "--grep with a regular expression")
    # drill down by sheet rows
    rows = run(["sheet_read.py", big, "--rows", "9001-9003", "--format", "csv"]).stdout.splitlines()
    c.ok(rows[0].startswith("row,Order,Date,Region") and rows[1].startswith("9001,108999,") and len(rows) == 4, f"--rows pages by sheet row number: {rows[:2]}")
    # budgets end on a row boundary with the next command
    md = run(["sheet_read.py", big, "--rows", "2-3000", "--max-chars", "3000"]).stdout
    c.ok(len(md) <= 3300 and "Next: python3 scripts/sheet_read.py" in md and "--rows" in md.rsplit("Next:", 1)[-1], f"--max-chars ends with the next command: {md[-240:]!r}")
    j = jrun(["sheet_read.py", big, "--rows", "2-3000", "--max-chars", "3000"])
    c.ok(len(json.dumps(j)) <= 3300 and j.get("next") and j["rows"] and len(j["rows"]) == len(j["row_numbers"]), "JSON output is capped too, whole rows and a next command")
    # SQL over every row, _row = sheet row
    q = jrun(["sheet_query.py", big, "--sql", "select count(*) n, round(sum(Revenue), 2) r, min(_row) a, max(_row) b from orders"])
    c.ok(q["rows"][0] == [n, round(total, 2), 2, n + 1], f"sheet_query over the cached columnar copy: {q['rows'][0]} vs {round(total, 2)}")
    md = run(["sheet_query.py", big, "--sql", "select * from orders", "--max-chars", "2000"]).stdout
    c.ok("Next: python3 scripts/sheet_query.py" in md and "--offset" in md, "sheet_query pages big results")
    # a small area of a huge sheet is read straight from the XML: same values as the full read
    a1 = jrun(["sheet_read.py", big, "--range", "A9000:H9003"])
    a2 = jrun(["sheet_read.py", big, "--range", "A9000:H9003"], env={"DESK_AREA_XML_MB": "0"})
    c.ok(a1["rows"] == a2["rows"] and a2["rows"][0][0] == 108998, f"area read from XML matches: {a2['rows'][:1]}")
    # renders are cached, and say how many rows the sheet really has
    t = time.time()
    r1 = jrun(["sheet_render.py", big, "--max-rows", "20", "--out", w / "r1"])
    cold = time.time() - t
    t = time.time()
    jrun(["sheet_render.py", big, "--max-rows", "20", "--out", w / "r2"])
    warm = time.time() - t
    c.ok(any(f"of 1-{n + 1}" in x for x in r1["notes"]), f"render note gives the true row count: {r1['notes']}")
    c.ok(cold >= 5 * warm and (w / "r2" / "Orders.png").exists(), f"cached render at least 5x faster: {cold:.2f}s then {warm:.2f}s")
    # surgical edits: cell ops on a big file are patched in place, the rest of the zip copied raw
    env = {"DESK_EDIT_PATCH_MB": "0.05"}
    p1 = w / "p1.xlsx"
    rep = jrun(["sheet_edit.py", big, "--out", p1, "--ops", '[{"op":"set","cell":"G9876","value":"patched"},{"op":"set","cell":"D2","value":1000},{"op":"style","range":"A1:H1","fill":"#FFF2CC","font":{"bold":true}},{"op":"style","range":"F2:F20","number_format":"#,##0.00"}]'], env=env)
    c.ok(rep.get("mode") == "patch", f"auto mode patches cell edits of a big file: {rep.get('mode')}")
    with zipfile.ZipFile(big) as za, zipfile.ZipFile(p1) as zb:
        ia, ib = za.getinfo("xl/worksheets/sheet2.xml"), zb.getinfo("xl/worksheets/sheet2.xml")
        c.ok((ia.CRC, ia.compress_size) == (ib.CRC, ib.compress_size), "untouched parts are copied without recompression")
    v = jrun(["sheet_read.py", p1, "--range", "D2:G2"])["rows"][0]
    c.ok(v[0] == 1000 and abs(v[2] - 1000 * v[1]) < 1e-6, f"patched value recalculated through the formula: {v}")
    c.ok(jrun(["sheet_read.py", p1, "--range", "G9876"])["rows"][0][0] == "patched", "patched text cell")
    st = run(["sheet_read.py", p1, "--range", "A1:B2", "--styles"]).stdout
    c.ok("#FFF2CC" in st and "bold" in st and "yyyy-mm-dd" in st, f"styles appended to styles.xml, other formats kept: {st[-300:]!r}")
    st = run(["sheet_read.py", p1, "--range", "F2:F3", "--styles"]).stdout
    c.ok("#,##0.00" in st, "number format patched in")
    # --styles and --display on a huge sheet stream its XML: only the rows asked for are kept, the rest skipped
    sys.path.insert(0, str(HERE))
    from _xlsx import Package, area_cells, stream_area

    with Package(p1) as pkg:
        part = pkg.sheet("Orders").path
        head, data = stream_area(pkg, part, 9000, 9003, chunk=4096)
        whole = pkg.read(part)
        c.ok(b"<cols" in head or b"<sheetViews" in head, "the streamed head holds the part before the rows")
        c.ok(len(data) < 4096 and list(area_cells(data, 9000, 9003)) == list(area_cells(whole, 9000, 9003)), f"only the requested rows are kept, the same cells as a full read: {len(data)} of {len(whole)} bytes")
        c.ok(list(area_cells(stream_area(pkg, part, n + 1, n + 1, chunk=4096)[1], n + 1, n + 1)) == list(area_cells(whole, n + 1, n + 1)), "the last row is found too")
    env0 = {"DESK_AREA_XML_MB": "0"}
    a = jrun(["sheet_read.py", p1, "--range", "A9000:H9001", "--styles"], env=env0)
    b = jrun(["sheet_read.py", p1, "--range", "A9000:H9001", "--styles"])
    c.ok(a["rows"] == b["rows"] and a["styles"] == b["styles"] and any("yyyy-mm-dd" in str(x) for x in a["styles"]), f"streamed --styles read matches the plain one: {a['styles']}")
    j = jrun(["sheet_read.py", p1, "--styles", "--limit", "5"], env=env0)
    k = jrun(["sheet_read.py", p1, "--styles", "--limit", "5"])
    c.ok(j["rows"] == k["rows"] and j["rows"][0][0] == 100000, f"the first page of a huge sheet is streamed with the same values: {j['rows'][:1]}")
    j = jrun(["sheet_read.py", p1, "--styles", "--format", "json"], env=env0)
    c.ok(len(j["rows"]) == 500 and j.get("next_rows") == "502-1001", f"a --styles read of a whole huge sheet is paged by sheet rows: {len(j['rows'])} {j.get('next_rows')}")
    rep = jrun(["sheet_edit.py", big, "--out", w / "p2.xlsx", "--no-recalc", "--ops", '[{"op":"insert_rows","at":5}]'], env=env)
    c.ok(rep.get("mode") == "openpyxl" and any("full openpyxl path" in x for x in rep["warnings"]), f"a structural op says it takes the slow path: {rep['warnings'][:1]}")
    r = run(["sheet_edit.py", big, "--out", w / "p3.xlsx", "--ops", '[{"op":"delete_rows","at":5}]'], env=dict(env, DESK_EDIT_SLOW_CELLS="50000"), expect=1)
    c.ok("--mode openpyxl" in r.stderr and "cells" in r.stderr, f"a structural op on a very big workbook asks first, with the cost: {r.stderr[-240:]}")


def test_fast_loader_parity(c: Checks, fx: dict, w: Path) -> None:
    for src in (fx["fancy"], fx["features"]):
        a = jrun(["sheet_recalc.py", src, "--check", "--compare", "--now", "2025-01-31", "--no-cache"])
        b = jrun(["sheet_recalc.py", src, "--check", "--compare", "--now", "2025-01-31", "--no-cache"], env={"DESK_FAST_LOAD_CELLS": "0"})
        keep = ("formulas", "error_counts", "cycles", "unsupported", "compare")
        c.ok({k: a.get(k) for k in keep} == {k: b.get(k) for k in keep}, f"the streaming loader matches openpyxl's on {src.name}")


def test_hostile_inputs(c: Checks, fx: dict, w: Path) -> None:
    import openpyxl

    bomb = w / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/worksheets/sheet1.xml", b"\0" * (17 << 20))
    for s in ("sheet_info.py", "sheet_read.py", "sheet_render.py"):
        r = run([s, bomb], expect=1)
        c.ok("zip bomb" in r.stderr and "DESK_ZIP_MAX_RATIO" in r.stderr, f"{s} refuses a zip bomb: {r.stderr[-200:]}")
    far = w / "far.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Item", "Qty"])
    for row in (["Pens", 10], ["Paper", 4], ["Ink", 7]):
        ws.append(row)
    ws["XFD1048576"] = "stray"
    wb.save(far)
    t = time.time()
    info = run(["sheet_info.py", far]).stdout
    j = jrun(["sheet_read.py", far])
    q = jrun(["sheet_query.py", far, "--sql", "select sum(Qty) from stock"])
    took = time.time() - t
    c.ok(took < 20 and q["rows"][0][0] == 21, f"a stray cell at XFD1048576 does not hang anything ({took:.1f}s)")
    c.ok(len(j["rows"]) == 3 and "XFD1048576" in json.dumps(j) and "XFD1048576" in info, "the stray cell is listed apart from the data")
    enc = w / "locked.xlsx"
    enc.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 504 + "EncryptionInfo".encode("utf-16-le") + b"\0" * 64 + "EncryptedPackage".encode("utf-16-le") + b"\0" * 512)
    r = run(["sheet_read.py", enc], expect=1)
    c.ok("password-protected" in r.stderr, f"encrypted workbooks are named as such: {r.stderr[-160:]}")
    ole = w / "ole.xlsx"
    ole.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 4088)
    c.ok("OLE2" in run(["sheet_info.py", ole], expect=1).stderr, "an OLE2 file without a workbook stream is refused")
    dtd = w / "laughs.xlsx"

    def add_dtd(name: str, data: bytes):
        if name == "xl/worksheets/sheet1.xml":
            return data.replace(b"<worksheet", b'<!DOCTYPE lol [<!ENTITY a "aaaaaaaa"><!ENTITY b "&a;&a;&a;&a;">]><worksheet', 1)
        return data

    add_dtd.extra = {}
    _patch_zip(fx["fancy"], dtd, add_dtd)
    for s in ("sheet_info.py", "sheet_edit.py"):
        args = [s, dtd] + (["--out", w / "x.xlsx", "--ops", '[{"op":"set","cell":"A1","value":1}]'] if s == "sheet_edit.py" else [])
        c.ok("DOCTYPE" in run(args, expect=1).stderr, f"{s} refuses XML with entity declarations")
    empty = w / "empty.xlsx"
    empty.write_bytes(b"")
    c.ok("empty" in run(["sheet_info.py", empty], expect=1).stderr, "an empty .xlsx is reported as empty")
    junk = w / "junk.xlsx"
    junk.write_bytes(bytes(range(256)) * 40)
    c.ok("not a spreadsheet" in run(["sheet_read.py", junk], expect=1).stderr, "random bytes are not read as CSV")
    inj = w / "inj.csv"
    inj.write_text('name,link,n\nA,"=HYPERLINK(""http://example.invalid"",""x"")",1\nB,@SUM(1),2\n', encoding="utf-8")
    def kind(path: Path, ref: str) -> str:
        wb = openpyxl.load_workbook(path)
        try:
            return wb.active[ref].data_type
        finally:
            wb.close()

    run(["sheet_create.py", w / "inj.xlsx", "--from", inj])
    c.ok(kind(w / "inj.xlsx", "B2") == "s" and values(w / "inj.xlsx").get("B2", "").startswith("=HYPERLINK"), "CSV text starting with = stays text")
    run(["sheet_create.py", w / "inj2.xlsx", "--from", inj, "--allow-formulas"])
    c.ok(kind(w / "inj2.xlsx", "B2") == "f", "--allow-formulas turns such text into formulas")
    run(["sheet_convert.py", inj, "--out", w / "inj3.xlsx"])
    c.ok(kind(w / "inj3.xlsx", "B2") == "s", "csv → xlsx conversion keeps = text as text")
    # A run of quotes once made the prefix-stripping regex backtrack exponentially (50 quotes: 6 s; 80: days).
    sys.path.insert(0, str(HERE))
    from _formula import strip_prefixes

    forms = {"SUM(_xlfn.ANCHORARRAY( 'My Sheet'!A1 ))": "SUM('My Sheet'!A1#)", "_xlfn.SINGLE('It''s'!A1:B2)+1": "@'It''s'!A1:B2+1",
             "_xlfn.ANCHORARRAY('x'''!A1)": "'x'''!A1#", "_xlfn.SINGLE('a'!A1,'b'!B1)": "SINGLE('a'!A1,'b'!B1)", "_xlfn.SINGLE('''')": "@''''",
             "_xlfn._xlws.SORT(_xlfn.ANCHORARRAY([Book1]Sheet1!A1))": "SORT([Book1]Sheet1!A1#)", "_xlfn.SINGLE('open": "SINGLE('open"}
    got = {f: strip_prefixes(f) for f in forms}
    c.ok(got == forms, f"strip_prefixes keeps its forms: {got}")
    t = time.time()
    slow = strip_prefixes("_xlfn.SINGLE(" + "'" * 100 + "x(") + strip_prefixes("_xlfn.ANCHORARRAY(" + "'" * 5001 + ")")
    c.ok(time.time() - t < 0.5 and slow.startswith("SINGLE("), f"strip_prefixes is linear on runs of quotes ({time.time() - t:.2f}s)")
    # Two more lazy regexes over workbook parts, now linear: theme colours and the print area a render sets.
    from _grid import DEFAULT_THEME, theme_colors
    from sheet_render import drop_print_area

    theme = '<a:theme><a:themeElements><a:clrScheme name="x"><a:dk1><a:sysClr val="windowText" lastClr="111111"/></a:dk1><a:accent1><a:srgbClr val="4472C4"/></a:accent1></a:clrScheme></a:themeElements></a:theme>'
    colors = theme_colors(type("WB", (), {"loaded_theme": theme.encode()})())
    t = time.time()
    hostile = theme_colors(type("WB", (), {"loaded_theme": b"<a:clrScheme>" * 30000})()) + theme_colors(type("WB", (), {"loaded_theme": (b"<a:clrScheme>" + b"<a:dk1>" * 30000 + b"</a:clrScheme>")})())
    took = time.time() - t
    c.ok(colors[1] == "111111" and colors[4] == "4472C4" and colors[0] == DEFAULT_THEME[0] and hostile == DEFAULT_THEME * 2 and took < 0.5, f"theme colours read in linear time ({took:.2f}s): {colors[:5]}")
    wbx = ('<workbook><definedNames><definedName name="_xlnm.Print_Area" localSheetId="0">S!$A$1:$B$2</definedName>'
           '<definedName localSheetId="1" name="_xlnm.Print_Area">T!$A$1:$C$3</definedName><definedName name="Rate">S!$D$1</definedName>'
           '<definedName localSheetId="10" name="_xlnm.Print_Area">U!$A$1</definedName></definedNames></workbook>')
    kept = drop_print_area(wbx, 1)
    c.ok("T!$A$1:$C$3" not in kept and all(x in kept for x in ("S!$A$1:$B$2", "S!$D$1", "U!$A$1")) and drop_print_area(wbx, 0).count("_xlnm.Print_Area") == 2,
         f"drop_print_area removes one sheet's print area only: {kept}")
    t = time.time()
    drop_print_area('<definedName name="_xlnm.Print_Area" localSheetId="0">x' * 5000, 0)
    drop_print_area("<definedName " + 'localSheetId="0" ' * 5000 + 'name="_xlnm.Print_Area" ' * 5000 + ">", 0)
    c.ok(time.time() - t < 0.5, f"drop_print_area is linear on a crafted workbook.xml ({time.time() - t:.2f}s)")
    quotes = w / "quotes.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["B1"] = "=_xlfn.SINGLE(" + "'" * 80 + "x("
    wb.save(quotes)
    t = time.time()
    r = run(["sheet_recalc.py", quotes, "--check"])
    c.ok(time.time() - t < 20 and "#NAME?" in r.stdout, f"sheet_recalc --check on a formula with 80 quotes finishes ({time.time() - t:.1f}s)")
    # Cell text coerced to numbers, times and dates: long space runs once backtracked ("(" + 500 spaces: minutes).
    from _numfmt import parse_date, parse_number, parse_time

    parsed = (parse_number(" ( $ 1,234.5 ) "), parse_number("-$5"), parse_number("12 %"), parse_number("( x"), parse_time(" 1:30 pm "), parse_time("7 a.m."),
              parse_date("2024-01-31 13:45"), parse_date("Jan 31 2024 1:30 pm"), parse_date("x\n1 12:30"))
    c.ok(parsed == (-1234.5, -5.0, 0.12, None, 0.5625, 7 / 24, 45322.572916666664, 45322.5625, None), f"number, time and date text keeps parsing: {parsed}")
    t = time.time()
    sp = " " * 30000
    parse_number("(" + sp + "x"), parse_number("-" + sp + "$" + sp + "-x"), parse_number("1" + sp + "%" + sp + "x")
    parse_time(sp + "1:2" + sp + "x"), parse_time("1" + sp + "pm" + sp + "x"), parse_date("1/2/2024" + sp + "12:30" + sp + "x")
    c.ok(time.time() - t < 0.5, f"number, time and date parsing is linear on long space runs ({time.time() - t:.2f}s)")
    from _xlsx import _ATTR, _T_ATTR

    tag = b' r="A1" s="3"  t="s" cm="1"'
    t = time.time()
    _ATTR.findall(b"_" * 200000), _T_ATTR.sub(b"", b" " * 200000 + b"x")
    c.ok(dict(_ATTR.findall(tag)) == {b"r": b"A1", b"s": b"3", b"t": b"s", b"cm": b"1"} and _T_ATTR.sub(b"", tag) == b' r="A1" s="3" cm="1"' and time.time() - t < 0.5,
         f"cell attributes are read in linear time ({time.time() - t:.2f}s)")


def test_tables_totals(c: Checks, fx: dict, w: Path) -> None:
    import openpyxl
    from openpyxl.worksheet.table import Table

    src = w / "raw.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount", "Code"])
    for row in (["N", 10.5, "007"], ["S", 20, "123"], ["E", 30, "042"]):
        ws.append(row)
    ws.append(["Total", "=SUBTOTAL(109,SalesTbl[Amount])", None])
    t = Table(displayName="SalesTbl", ref="A1:C5")
    t.totalsRowCount = 1
    ws.add_table(t)
    wb.create_sheet("Summary")["A1"] = "=SUM(SalesTbl[Amount])"
    wb.save(src)
    book = w / "t.xlsx"
    run(["sheet_recalc.py", src, "--out", book])
    q = jrun(["sheet_query.py", book, "--sql", "select sum(Amount) s, count(*) n from salestbl"])
    c.ok(q["rows"][0] == [60.5, 3], f"an Excel table is queried by name without its totals row: {q['rows']}")
    q = jrun(["sheet_query.py", book, "--sql", "select sum(Amount) s from sales"])
    c.ok(q["rows"][0][0] == 60.5, f"the sheet's table leaves the totals row out too: {q['rows']}")
    tb = run(["sheet_query.py", book, "--tables"]).stdout
    c.ok("salestbl" in tb and "Amount DOUBLE" in tb and "Code VARCHAR" in tb, f"typed columns, codes kept as text: {tb[:300]!r}")
    m = run(["sheet_read.py", book, "--map"]).stdout
    c.ok("totals row of table SalesTbl" in m, "the map says which rows were left out")
    j = jrun(["sheet_query.py", book, "--sql", "select Code from salestbl order by _row"])
    c.ok([r[0] for r in j["rows"]] == ["007", "123", "042"], f"leading zeros decided per column: {j['rows']}")
    csvf = w / "money.csv"
    csvf.write_text('Month,Revenue,Share,When\nJan,"$1,234.50",12.5%,2025-01-31\nFeb,"($150.25)",-3.1%,2025-02-28\n', encoding="utf-8")
    tb = run(["sheet_query.py", csvf, "--tables"]).stdout
    c.ok("Revenue DOUBLE" in tb and "Share DOUBLE" in tb and "When DATE" in tb, f"CSV currency, accounting negatives and percents typed as numbers: {tb!r}")
    r = run(["sheet_query.py", csvf, "--sql", "select When from money"], expect=1)
    c.ok('"When"' in r.stderr, "SQL keyword column names get a quoting hint")


def test_formats_roundtrip(c: Checks, fx: dict, w: Path) -> None:
    sys.path.insert(0, str(HERE))
    import _ods

    c.ok(_ods.fix_lo_arity("ROUND(A1)+ROUNDUP(B2,1)") == "ROUND(A1,0)+ROUNDUP(B2,1)", "LibreOffice one-argument ROUND gets Excel's second argument")
    ods = w / "f.ods"
    run(["sheet_convert.py", fx["fancy"], "--out", ods])
    back = w / "back.xlsx"
    r = run(["sheet_convert.py", ods, "--out", back])
    c.ok("formatting kept" in r.stdout, f"ods → xlsx says what formatting it kept: {r.stdout[-300:]!r}")
    import openpyxl

    wb = openpyxl.load_workbook(back)
    ws = wb["Fancy"]
    c.ok(ws["A1"].font.b and (ws["A1"].fill.fgColor.rgb or "").endswith("FFC000"), "ods → xlsx keeps bold and fills")
    c.ok(ws["C4"].number_format == "0.0%" and "A3:C3" in [str(m) for m in ws.merged_cells.ranges], f"ods → xlsx keeps number formats and merges: {ws['C4'].number_format}")
    wb.close()
    tbl = w / "tbl.xlsx"
    run(["sheet_create.py", tbl, "--from", fx["csv"], "--table", "--totals", "--chart", "column"])
    r = run(["sheet_convert.py", tbl, "--out", w / "tbl.ods"])
    c.ok("not written to .ods" in r.stdout and "chart" in r.stdout, f"xlsx → ods names what the built-in writer drops: {r.stdout!r}")
    j = jrun(["sheet_read.py", w / "tbl.ods", "--formulas"])
    c.ok(j["formulas"] and not any("[" in f["formula"] for f in j["formulas"]), f"structured references become ranges in .ods: {j['formulas'][:2]}")
    pdf = w / "f.pdf"
    run(["sheet_convert.py", fx["fancy"], "--out", pdf])
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf))
    try:
        sizes = {tuple(round(x) for x in doc[i].get_size()) for i in range(len(doc))}
        c.ok(sizes <= {(595, 842), (842, 595)} and len(doc) <= 3, f"built-in PDF: A4 pages, no blank ones: {len(doc)} {sizes}")
    finally:
        doc.close()


def test_ops_regressions(c: Checks, fx: dict, w: Path) -> None:
    src = w / "g.xlsx"
    run(["sheet_create.py", src, "--spec", json.dumps({"sheets": [{"name": "G", "columns": [{"header": "a"}, {"header": "b"}, {"header": "x"}, {"header": "y"}], "rows": [[1, 2], [3, 4], [5, 6], [7, 8]], "cells": {"C2": "=A2+B2", "D2": "=A2*B2"}}]})])
    out = w / "g2.xlsx"
    run(["sheet_edit.py", src, "--out", out, "--ops", '[{"op":"fill","range":"C2:D5"}]'])
    f = formulas(out)
    c.ok(f.get("C5") == "=A5+B5" and f.get("D5") == "=A5*B5", f"fill copies each column's top cell down (Ctrl+D): {f}")
    v = values(out)
    c.ok(v.get("C5") == 15 and v.get("D5") == 56, "filled formulas recalculated")
    p = w / "p.xlsx"
    rep = jrun(["sheet_edit.py", src, "--out", p, "--mode", "patch", "--ops", '[{"op":"fill","range":"C2:D5"},{"op":"style","range":"A1:D1","fill":"#DDEBF7","border":"thin"}]'])
    c.ok(rep.get("mode") == "patch" and formulas(p).get("D5") == "=A5*B5" and values(p).get("D5") == 56, "patch mode fills from the top row too")
    import openpyxl

    wb = openpyxl.load_workbook(p)
    ws = wb["G"]
    c.ok((ws["B1"].fill.fgColor.rgb or "").endswith("DDEBF7") and ws["B1"].border.left.style == "thin" and ws["B1"].font.b, "patch-mode styles keep the header's bold and add fill and border")
    wb.close()
    # overwriting the cell that holds a shared formula's text keeps the rest of the group working
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for i in range(1, 6):
        ws.cell(i, 1, i)
        ws.cell(i, 2, f"=A{i}*2")
    wb.save(w / "sh_raw.xlsx")

    def share(name: str, data: bytes):
        if name != "xl/worksheets/sheet1.xml":
            return data
        t = re.sub(r'<c r="B1"><f>A1\*2</f><v\s*/></c>', '<c r="B1"><f t="shared" ref="B1:B5" si="0">A1*2</f><v>2</v></c>', data.decode())
        for i in range(2, 6):
            t = re.sub(rf'<c r="B{i}"><f>A{i}\*2</f><v\s*/></c>', f'<c r="B{i}"><f t="shared" si="0"/><v>{2 * i}</v></c>', t)
        return t.encode()

    share.extra = {}
    _patch_zip(w / "sh_raw.xlsx", w / "shared.xlsx", share)
    rep = jrun(["sheet_edit.py", w / "shared.xlsx", "--out", w / "shared2.xlsx", "--mode", "patch", "--ops", '[{"op":"set","cell":"B1","value":100},{"op":"set","cell":"A3","value":10}]'])
    v = values(w / "shared2.xlsx")
    c.ok(v.get("B1") == 100 and v.get("B3") == 20 and v.get("B5") == 10 and formulas(w / "shared2.xlsx").get("B4") == "=A4*2", f"shared-formula group survives its master being overwritten: {v}")
    c.ok(any("shared formula" in x for x in rep["warnings"]), "and the output says so")
    made = w / "h.xlsx"
    run(["sheet_create.py", made, "--from", fx["csv"]])
    wb = openpyxl.load_workbook(made)
    ws = wb.active
    c.ok(ws["A1"].alignment.horizontal == "center", "headers are centred")
    c.ok(ws.column_dimensions["E"].width >= len("Share") + 3, f"header widths leave room for bold text and filter buttons: {ws.column_dimensions['E'].width}")
    wb.close()


def test_edit_keeps_extensions(c: Checks, fx: dict, w: Path) -> None:
    ext = (b'<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
           b'<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:sparklineGroup>'
           b"<x14:sparklines><x14:sparkline><xm:f>Fancy!B16:B19</xm:f><xm:sqref>D16</xm:sqref></x14:sparkline></x14:sparklines>"
           b"</x14:sparklineGroup></x14:sparklineGroups></ext></extLst>")
    src = w / "spark.xlsx"

    def edit(name: str, data: bytes):
        if name == "xl/worksheets/sheet1.xml":
            return data.replace(b"</worksheet>", ext + b"</worksheet>")
        return data

    edit.extra = {}
    _patch_zip(fx["fancy"], src, edit)
    auto = jrun(["sheet_edit.py", src, "--out", w / "a.xlsx", "--ops", '[{"op":"set","cell":"B16","value":20}]'])
    with zipfile.ZipFile(w / "a.xlsx") as z:
        kept = b"sparklineGroup" in z.read("xl/worksheets/sheet1.xml")
    c.ok(auto.get("mode") == "patch" and kept, "cell edits keep sparklines (patch mode chosen)")
    rep = jrun(["sheet_edit.py", src, "--out", w / "b.xlsx", "--ops", '[{"op":"insert_rows","at":2}]'])
    c.ok(any("dropped" in x and "sparklines" in x for x in rep["warnings"]), f"openpyxl mode says what it dropped: {rep['warnings']}")


def test_pivot_chart_cache(c: Checks, fx: dict, w: Path) -> None:
    sys.path.insert(0, str(HERE))
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.data_source import NumData, NumVal, StrData, StrVal

    import _grid

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["k", "v"])
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    s = ch.series[0]
    s.val.numRef.f = "pt@data 0"
    s.val.numRef.numCache = NumData(pt=[NumVal(idx=0, v=5), NumVal(idx=1, v=7)], ptCount=2)
    from openpyxl.chart.data_source import AxDataSource, StrRef

    s.cat = AxDataSource(strRef=StrRef(f="pt@categories", strCache=StrData(pt=[StrVal(idx=0, v="a"), StrVal(idx=1, v="b")], ptCount=2)))
    ws.add_chart(ch, "D2")
    p = w / "pivotchart.xlsx"
    wb.save(p)
    v = _grid.load_view(p, None, None, 30, 12, 1.0)
    got = v.charts[0]["series"][0] if v.charts else {}
    c.ok(got.get("values") == [5.0, 7.0] and got.get("categories") == ["a", "b"], f"pivot charts are drawn from their cached values: {got}")


def _near_color(img, want: tuple[int, int, int], tol: int = 14) -> int:
    """How many pixels of a PIL image are within tol of a colour."""
    return sum(n for n, p in (img.convert("RGB").getcolors(maxcolors=1 << 22) or []) if all(abs(a - b) <= tol for a, b in zip(p, want)))


def test_formulas_without_values(c: Checks, fx: dict, w: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty", "Price", "Total"])
    for i, (n, q, pr) in enumerate([("a", 1, 2), ("b", 3, 4), ("c", 5, 6), ("d", 7, 8), ("e", 9, 10)], start=2):
        ws.cell(i, 1, n), ws.cell(i, 2, q), ws.cell(i, 3, pr), ws.cell(i, 4, f"=B{i}*C{i}")
    ws["F1"], ws["G1"], ws["H1"], ws["I1"] = "=SUM(D2:D6)", "=F1*2", "=G1+1", '="x"&H1'
    ws["A8"], ws["A9"] = "=SUM(B2:B6)", "=A8/5"
    only = wb.create_sheet("Only")  # formulas and nothing else: the value reader sees an empty sheet
    only["A1"], only["A2"], only["B3"] = "=1+1", "=A1*2", "=A2+1"
    src = w / "nocache.xlsx"
    wb.save(src)
    want = {"F1", "G1", "H1", "I1", "D2:D6", "A8", "A9"}
    j = jrun(["sheet_read.py", src, "--formulas"])
    got = {f["range"] for f in j["formulas"]}
    c.ok(got == want, f"--formulas lists formula cells outside the value area (F1:I1, A8, A9): {sorted(got)}")
    c.ok(j.get("formulas_without_values") == 11 and "sheet_recalc.py" in j.get("formulas_note", ""), f"formulas without stored results are counted and point to sheet_recalc: {j.get('formulas_note')}")
    c.ok(j["row_numbers"][-1] == 9 and len(j["rows"][0]) == 9, f"the table covers the formula cells too: rows to {j['row_numbers'][-1]}, {len(j['rows'][0])} columns")
    md = run(["sheet_read.py", src, "--formulas"]).stdout
    c.ok("A9: =A8/5  → (no stored result)" in md and "no stored results)" in md and "sheet_recalc.py" in md, f"Markdown says which formulas have no result: {md[-500:]!r}")
    seen: list[str] = []
    for off in (0, 3, 6):
        seen += [f["range"] for f in jrun(["sheet_read.py", src, "--formulas", "--limit", "3", "--offset", off])["formulas"]]
    cells = []
    for rng in seen:
        m = re.fullmatch(r"([A-Z]+)(\d+)(?::[A-Z]+(\d+))?", rng)
        cells += [f"{m.group(1)}{r}" for r in range(int(m.group(2)), int(m.group(3) or m.group(2)) + 1)]
    c.ok(sorted(cells) == sorted(["F1", "G1", "H1", "I1", "A8", "A9"] + [f"D{r}" for r in range(2, 7)]), f"paged reads list every formula exactly once: {seen}")
    j = jrun(["sheet_read.py", src, "--formulas", "--range", "A1:D6"])
    c.ok([f["range"] for f in j["formulas"]] == ["D2:D6"], "--range limits the listing to the range")
    j = jrun(["sheet_read.py", src, "--sheet", "Only", "--formulas"])
    c.ok({f["range"] for f in j["formulas"]} == {"A1", "A2", "B3"} and j["row_numbers"] == [1, 2, 3] and j.get("formulas_without_values") == 3, f"a sheet of formulas only is listed and shown: {j['formulas']} {j['row_numbers']}")
    j = jrun(["sheet_read.py", src, "--sheet", "Only", "--formulas", "--rows", "2-3"])
    c.ok([f["range"] for f in j["formulas"]] == ["A2", "B3"] and j["row_numbers"] == [2, 3], f"--rows on it too: {j['formulas']} {j['row_numbers']}")
    run(["sheet_recalc.py", src, "--out", w / "calc.xlsx"])
    j = jrun(["sheet_read.py", w / "calc.xlsx", "--formulas"])
    f1 = next((f for f in j["formulas"] if f["range"] == "F1"), {})
    c.ok(f1.get("value") == 190 and not j.get("formulas_without_values"), f"after sheet_recalc the results show: {f1}")


def _wide_report(path: Path, anchor: str, width_cm: float, legend: str = "b") -> None:
    """24 columns of numbers (wide enough to need two images) and a two-series bar chart."""
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Name"] + [f"M{i}" for i in range(1, 24)])
    for r in range(2, 32):
        ws.append([f"row {r}"] + [(r * 7 + k * 3) % 50 for k in range(1, 24)])
    ws.column_dimensions["A"].width = 20
    for k in range(2, 25):
        ws.column_dimensions[get_column_letter(k)].width = 10
    ch = BarChart()
    ch.title = "Split me"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_col=3, max_row=8), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=8))
    ch.legend.position = legend
    ch.width, ch.height = width_cm, 7
    ws.add_chart(ch, anchor)
    wb.save(path)


def test_render_split_charts(c: Checks, fx: dict, w: Path) -> None:
    from PIL import Image

    blue, orange = (0x44, 0x72, 0xC4), (0xED, 0x7D, 0x31)
    src = w / "wide.xlsx"
    _wide_report(src, "Q4", 12)
    j = jrun(["sheet_render.py", src, "--out", w / "a"])
    sh = j["sheets"][0]
    ch = sh["charts"][0] if sh.get("charts") else {}
    c.ok(len(sh["images"]) == 2 and len(ch.get("images", [])) == 1 and ch["anchor"] == "Q4", f"a chart across the natural split lands whole in one image: {sh.get('pages')} {ch}")
    if len(sh["images"]) == 2 and ch.get("images"):
        other = [p for p in sh["images"] if p not in ch["images"]][0]
        with Image.open(ch["images"][0]) as im_in, Image.open(other) as im_out:
            c.ok(_near_color(im_in, orange) > 400 and _near_color(im_out, orange) == 0, "the image named for the chart holds it; the other has none of it")
    md = run(["sheet_render.py", src, "--out", w / "a2"]).stdout
    c.ok("chart 'Split me' at Q4: Report-p" in md, f"the output says which image holds each chart: {md[:300]!r}")
    big = w / "wider.xlsx"
    _wide_report(big, "B4", 48)
    j = jrun(["sheet_render.py", big, "--out", w / "b"])
    sh = j["sheets"][0]
    ch = sh["charts"][0] if sh.get("charts") else {}
    c.ok(len(ch.get("images", [])) == 2, f"a chart too wide for one image is named in both: {ch}")
    parts = []
    for p in ch.get("images", []):
        with Image.open(p) as im:
            parts.append((_near_color(im, blue), _near_color(im, orange), im.convert("RGB")))
    c.ok(len(parts) == 2 and all(b > 300 and o > 300 for b, o, _ in parts), f"each image draws its part of the chart: {[(b, o) for b, o, _ in parts]}")
    if len(parts) == 2:
        # the chart's frame runs straight across both images: consistent clipping, same size and place
        def frame_rows(im, x: int) -> set[int]:
            return {y for y in range(im.height) if all(im.getpixel((x + k, y)) == (191, 191, 191) for k in range(5))}

        a_rows, b_rows = frame_rows(parts[0][2], parts[0][2].width - 8), frame_rows(parts[1][2], 60)
        c.ok(len(a_rows) == 2 and a_rows == b_rows, f"the two parts line up (the chart frame on the same rows): {sorted(a_rows)} {sorted(b_rows)}")
    md = run(["sheet_render.py", big, "--out", w / "b2"]).stdout
    c.ok("split across Report-p1.png, Report-p2.png" in md, f"a split chart is named as split: {md[:300]!r}")


def test_chart_legend_position(c: Checks, fx: dict, w: Path) -> None:
    sys.path.insert(0, str(HERE))
    import _grid

    boxes = {}
    for pos in ("r", "l", "t", "b", "tr"):
        src = w / f"legend-{pos}.xlsx"
        _wide_report(src, "B4", 12, legend=pos)
        v = _grid.load_view(src, None, None, 40, 30, 1.0)
        chart = v.charts[0] if v.charts else {}
        c.ok(chart.get("legend_pos") == pos, f"legendPos {pos} read from the file: {chart.get('legend_pos')}")
        img = _grid.draw_chart(chart, 480, 300, 1.0)
        box = img.info.get("legend_box")
        boxes[pos] = box
        if box:
            x0, y0, x1, y1 = box
            crop = img.crop((max(0, x0), max(0, y0), min(480, x1 + 1), min(300, y1 + 1)))
            c.ok(_near_color(crop, (0xED, 0x7D, 0x31)) > 20, f"legend {pos}: the swatches are drawn inside its box {box}")
    b, t, l, r, tr = (boxes.get(k) or (0, 0, 0, 0) for k in ("b", "t", "l", "r", "tr"))
    c.ok(b[1] > 300 * 0.75 and b[0] > 60 and b[2] < 420, f"legend b sits at the bottom, centred: {b}")
    c.ok(t[3] < 300 * 0.3 and t[0] > 60, f"legend t sits at the top, under the title: {t}")
    c.ok(l[2] < 480 * 0.35 and r[0] > 480 * 0.6 and tr[0] > 480 * 0.6, f"legends l, r and tr are at the sides: {l} {r} {tr}")
    c.ok(tr[1] < r[1], f"tr sits higher than r (top right corner vs centred): {tr} {r}")
    img = _grid.draw_chart({"kind": "column", "title": "", "legend": False, "series": [{"name": "A", "values": [1, 2]}, {"name": "B", "values": [2, 1]}]}, 480, 300, 1.0)
    c.ok("legend_box" not in img.info, "a chart without a legend element draws none")
    c.ok(_grid._tick(7.5, 2.5) == "7.5" and _grid._tick(10, 2) == "10", "axis ticks keep the step's decimals (2.5, 5.0, 7.5)")


def test_odd_sheet_names(c: Checks, fx: dict, w: Path) -> None:
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    evil = "../../../evil"

    def book(path: Path, names: list[str]) -> None:
        wb = openpyxl.Workbook()
        first = True
        for k, nm in enumerate(names):
            ws = wb.active if first else wb.create_sheet(f"S{k}")
            first = False
            ws.title = f"S{k}"
            ws.append(["k", "v"])
            for i in range(1, 6):
                ws.append([f"r{i}", i * 3])
            ws["D1"] = f"=SUM(S{k}!B2:B6)"
            ch = BarChart()
            ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
            ws.add_chart(ch, "D3")
        raw = path.with_suffix(".raw.xlsx")
        wb.save(raw)

        def edit(name: str, data: bytes):
            if name == "xl/workbook.xml" or name.startswith(("xl/worksheets/sheet", "xl/charts/chart")):
                for k, nm in enumerate(names):
                    q = ("'" + nm.replace("'", "''") + "'!").replace("&", "&amp;").replace("<", "&lt;").encode()
                    data = data.replace(b'name="S%d"' % k, b'name="' + nm.replace("&", "&amp;").replace("<", "&lt;").replace('"', "&quot;").encode() + b'"')
                    data = data.replace(b"'S%d'!" % k, q).replace(b"S%d!" % k, q)
            return data

        edit.extra = {}
        _patch_zip(raw, path, edit)
        raw.unlink()

    src = w / "evil.xlsx"
    book(src, [evil, "Other"])
    out = w / "r"
    j = jrun(["sheet_render.py", src, "--out", out])
    c.ok([Path(p).name for p in j["images"]] == ["evil.png"] and j["sheets"][0]["sheet"] == evil, f"a sheet named {evil} renders, safely named: {j['images']}")
    c.ok(sorted(p.name for p in w.rglob("*.png")) == ["evil.png"], f"nothing written outside the output folder: {sorted(str(p) for p in w.rglob('*.png'))}")
    if j["images"]:
        from PIL import Image

        with Image.open(j["images"][0]) as im:
            c.ok(_near_color(im, (0x44, 0x72, 0xC4)) > 400, "its chart is drawn from the sheet's data (references by the real name)")
    j = jrun(["sheet_render.py", src, "--all", "--formulas", "--out", w / "all"])
    c.ok(sorted(Path(p).name for p in j["images"]) == ["Other.png", "evil.png"], f"--all renders it too: {j['images']}")
    conv = jrun(["sheet_convert.py", src, "--to", "csv", "--out-dir", w / "csv"])
    c.ok(sorted(Path(p).name for p in conv["outputs"]) == ["Other.csv", "evil.csv"], f"sheet_convert names files the same way: {conv['outputs']}")
    twins = w / "twins.xlsx"
    book(twins, ["a b", "a&b", "CON"])
    j = jrun(["sheet_render.py", twins, "--all", "--out", w / "t"])
    names = sorted(Path(p).name for p in j["images"])
    c.ok(names == ["_CON.png", "a_b-2.png", "a_b.png"], f"sheets whose names clean up alike get distinct files, device names are avoided: {names}")
    conv = jrun(["sheet_convert.py", twins, "--to", "csv", "--out-dir", w / "tcsv"])
    c.ok(sorted(Path(p).name for p in conv["outputs"]) == ["_CON.csv", "a_b-2.csv", "a_b.csv"], f"and so does sheet_convert: {conv['outputs']}")


if __name__ == "__main__":
    sys.exit(main())
