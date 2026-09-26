#!/usr/bin/env python3
"""Self-test for the word-documents skill: builds fixtures, runs every script as an agent would, checks results.

Runs the LibreOffice-free path always (DESK_SOFFICE=none) and the LibreOffice path too when it is installed.
Prints one summary line ("ok: N checks in S s") and exits non-zero on any failure.

Example:
  python3 scripts/selftest.py
  python3 scripts/selftest.py --keep      # keep the temp folder and print its path
"""

from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ── fixtures ────────────────────────────────────────────────────────────

CT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Default Extension="png" ContentType="image/png"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>
<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>
<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>
<Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>"""

ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""

DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" Target="settings.xml"/>
<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>
<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>
<Relationship Id="rId6" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>
<Relationship Id="rId7" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>
<Relationship Id="rId8" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com/terms" TargetMode="External"/>
<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
</Relationships>"""

NSDECL = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
          'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
          'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
          'xmlns:v="urn:schemas-microsoft-com:vml" '
          'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
          'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14"')

STYLES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles {NSDECL}>
<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="Calibri" w:cs="Calibri"/><w:sz w:val="22"/><w:szCs w:val="22"/><w:lang w:val="en-US"/></w:rPr></w:rPrDefault>
<w:pPrDefault><w:pPr><w:spacing w:after="160" w:line="259" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>
<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>
<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:pPr><w:jc w:val="center"/></w:pPr><w:rPr><w:b/><w:color w:val="1F3864"/><w:sz w:val="48"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:color w:val="2F5496"/><w:sz w:val="32"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="200" w:after="80"/><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:b/><w:color w:val="2F5496"/><w:sz w:val="26"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/><w:pPr><w:ind w:left="720"/><w:contextualSpacing/></w:pPr></w:style>
<w:style w:type="paragraph" w:styleId="Quote"><w:name w:val="Quote"/><w:basedOn w:val="Normal"/><w:pPr><w:ind w:left="864" w:right="864"/></w:pPr><w:rPr><w:i/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="FootnoteText"><w:name w:val="footnote text"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="0"/></w:pPr><w:rPr><w:sz w:val="20"/></w:rPr></w:style>
<w:style w:type="character" w:default="1" w:styleId="DefaultParagraphFont"><w:name w:val="Default Paragraph Font"/></w:style>
<w:style w:type="character" w:styleId="FootnoteReference"><w:name w:val="footnote reference"/><w:rPr><w:vertAlign w:val="superscript"/></w:rPr></w:style>
<w:style w:type="character" w:styleId="Hyperlink"><w:name w:val="Hyperlink"/><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr></w:style>
<w:style w:type="character" w:styleId="Strong"><w:name w:val="Strong"/><w:rPr><w:b/></w:rPr></w:style>
<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:tblPr><w:tblCellMar><w:left w:w="108" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>
<w:style w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:basedOn w:val="TableNormal"/><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/></w:tblBorders></w:tblPr>
<w:tblStylePr w:type="firstRow"><w:rPr><w:b/></w:rPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="D9E2F3"/></w:tcPr></w:tblStylePr></w:style>
</w:styles>"""

NUMBERING = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering {NSDECL}>
<w:abstractNum w:abstractNumId="0"><w:multiLevelType w:val="hybridMultilevel"/>
<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>
<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl>
<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:lvlText w:val="%3."/><w:lvlJc w:val="right"/><w:pPr><w:ind w:left="2160" w:hanging="180"/></w:pPr></w:lvl></w:abstractNum>
<w:abstractNum w:abstractNumId="1"><w:multiLevelType w:val="hybridMultilevel"/>
<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val=""/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl></w:abstractNum>
<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
<w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num>
</w:numbering>"""

SETTINGS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:settings {NSDECL}><w:zoom w:percent="100"/><w:defaultTabStop w:val="720"/><w:characterSpacingControl w:val="doNotCompress"/><w:compat><w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat></w:settings>"""

FOOTNOTES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes {NSDECL}>
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr><w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Payable in euros to ACME Ltd.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""

COMMENTS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments {NSDECL}>
<w:comment w:id="0" w:author="Alice Martin" w:date="2026-09-01T10:00:00Z" w:initials="AM"><w:p><w:r><w:t>Check the payment terms.</w:t></w:r></w:p></w:comment>
</w:comments>"""

HEADER = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:hdr {NSDECL}><w:p><w:pPr><w:jc w:val="right"/></w:pPr><w:r><w:t xml:space="preserve">ACME Ltd </w:t></w:r><w:r><w:t>— Confidential</w:t></w:r></w:p></w:hdr>"""

FOOTER = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:ftr {NSDECL}><w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:t xml:space="preserve">Page </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t xml:space="preserve"> of </w:t></w:r><w:fldSimple w:instr=" NUMPAGES "><w:r><w:t>2</w:t></w:r></w:fldSimple></w:p></w:ftr>"""


def _p(inner: str, style: str | None = None, extra_ppr: str = "") -> str:
    ppr = (f'<w:pStyle w:val="{style}"/>' if style else "") + extra_ppr
    return f"<w:p>{'<w:pPr>' + ppr + '</w:pPr>' if ppr else ''}{inner}</w:p>"


def _r(text: str, rpr: str = "") -> str:
    return f'<w:r>{"<w:rPr>" + rpr + "</w:rPr>" if rpr else ""}<w:t xml:space="preserve">{text}</w:t></w:r>'


def _num(level: int, num_id: int = 1) -> str:
    return f'<w:numPr><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'


def _drawing(rid: str, cx: int, cy: int, name: str, descr: str) -> str:
    return (f'<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cx}" cy="{cy}"/><wp:docPr id="1" name="{name}" descr="{descr}"/>'
            f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="{name}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def _textbox(text: str) -> str:
    inner = _p(_r(text))
    return ('<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            '<wp:extent cx="2286000" cy="457200"/><wp:wrapTopAndBottom/><wp:docPr id="5" name="Text Box 5"/>'
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2286000" cy="457200"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFF2CC"/></a:solidFill><a:ln w="9525"><a:solidFill><a:srgbClr val="BF9000"/></a:solidFill></a:ln></wps:spPr>'
            f'<wps:txbx><w:txbxContent>{inner}</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
            f'<mc:Fallback><w:pict><v:shape style="width:180pt;height:36pt"><v:textbox><w:txbxContent>{inner}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>')


def complex_document_xml() -> str:
    """A Word-like body: split runs, lists, tracked changes, comments, footnote, link, merged table, text box, image, math."""
    b = []
    b.append(_p(_r("Service Agreement"), "Title"))                                                         # 0
    b.append(_p(_r("Parties"), "Heading1"))                                                                # 1
    b.append(_p(_r("This Agreement is made between ") + _r("AC", "<w:b/>") + _r("ME", "<w:b/><w:i/>") + _r(" Ltd") + _r(" and the Client.")))  # 2
    b.append(_p(_r("First obligation"), "ListParagraph", _num(0)))                                          # 3
    b.append(_p(_r("Sub obligation"), "ListParagraph", _num(1)))                                            # 4
    b.append(_p(_r("Second obligation"), "ListParagraph", _num(0)))                                         # 5
    b.append(_p(_r("Bullet point"), "ListParagraph", _num(0, 2)))                                           # 6
    b.append(_p(_r("The term is ") + '<w:del w:id="11" w:author="Bob" w:date="2026-09-02T09:00:00Z"><w:r><w:delText>30</w:delText></w:r></w:del>'
                + '<w:ins w:id="12" w:author="Bob" w:date="2026-09-02T09:00:00Z"><w:r><w:t>60</w:t></w:r></w:ins>' + _r(" days.")))  # 7
    b.append(_p(_r("Invoices require ") + '<w:commentRangeStart w:id="0"/>' + _r("payment") + '<w:commentRangeEnd w:id="0"/>'
                + '<w:r><w:commentReference w:id="0"/></w:r>' + _r(" within 30 days.") + '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteReference w:id="1"/></w:r>'))  # 8
    b.append(_p(_r("See the ") + '<w:hyperlink r:id="rId8"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr><w:t>terms</w:t></w:r></w:hyperlink>' + _r(" online.")))  # 9
    grid = '<w:tblGrid><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/></w:tblGrid>'
    tbl = ('<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="9000" w:type="dxa"/><w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/></w:tblPr>' + grid
           + '<w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:tcPr><w:tcW w:w="6000" w:type="dxa"/><w:gridSpan w:val="2"/></w:tcPr>' + _p(_r("Item and region")) + '</w:tc><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>' + _p(_r("Price")) + '</w:tc></w:tr>'
           + '<w:tr><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/><w:vMerge w:val="restart"/></w:tcPr>' + _p(_r("Widgets")) + '</w:tc><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>' + _p(_r("North")) + '</w:tc><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/><w:shd w:val="clear" w:color="auto" w:fill="FFFF00"/></w:tcPr>' + _p(_r("1,200")) + '</w:tc></w:tr>'
           + '<w:tr><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/><w:vMerge/></w:tcPr>' + _p("") + '</w:tc><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>' + _p(_r("South")) + '</w:tc><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>' + _p(_r("950")) + '</w:tc></w:tr>'
           + '</w:tbl>')
    b.append(tbl)                                                                                        # 10
    b.append(_p(_r("Note:") + _textbox("Box text for ACME Ltd")))                                         # 11
    b.append(_p(_drawing("rId9", 914400, 457200, "Logo", "Company logo"), extra_ppr='<w:jc w:val="center"/>'))  # 12
    b.append(_p(_r("Energy: ") + '<m:oMath><m:r><m:t>E=m</m:t></m:r><m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>'))  # 13
    b.append(_p('<w:r><w:br w:type="page"/></w:r>'))                                                       # 14
    b.append(_p(_r("Payment"), "Heading1"))                                                                # 15
    b.append(_p(_r("Fees are due quarterly. ACME Ltd may revise fees yearly."), "Quote"))                # 16
    b.append(_p(_r("Tab\tseparated\tvalues")))                                                             # 17
    sect = ('<w:sectPr><w:headerReference w:type="default" r:id="rId6"/><w:footerReference w:type="default" r:id="rId7"/>'
            '<w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>')
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:document {NSDECL}><w:body>{"".join(b)}{sect}</w:body></w:document>'


def png_bytes(color: tuple[int, int, int] = (46, 116, 181), size: tuple[int, int] = (120, 60)) -> bytes:
    from PIL import Image, ImageDraw

    im = Image.new("RGB", size, color)
    d = ImageDraw.Draw(im)
    d.rectangle([4, 4, size[0] - 5, size[1] - 5], outline=(255, 255, 255), width=3)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def write_complex_docx(path: Path) -> Path:
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<dc:title>Service Agreement</dc:title><dc:creator>Alice Martin</dc:creator><cp:lastModifiedBy>Bob</cp:lastModifiedBy>'
            '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-01T08:00:00Z</dcterms:created></cp:coreProperties>')
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
           '<Application>Microsoft Office Word</Application><Pages>2</Pages><Words>90</Words><Company>ACME</Company></Properties>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", CT)
        z.writestr("_rels/.rels", ROOT_RELS)
        z.writestr("word/_rels/document.xml.rels", DOC_RELS)
        z.writestr("word/document.xml", complex_document_xml())
        z.writestr("word/styles.xml", STYLES)
        z.writestr("word/numbering.xml", NUMBERING)
        z.writestr("word/settings.xml", SETTINGS)
        z.writestr("word/footnotes.xml", FOOTNOTES)
        z.writestr("word/comments.xml", COMMENTS)
        z.writestr("word/header1.xml", HEADER)
        z.writestr("word/footer1.xml", FOOTER)
        z.writestr("word/media/image1.png", png_bytes())
        z.writestr("docProps/core.xml", core)
        z.writestr("docProps/app.xml", app)
    return path


def write_template_docx(path: Path) -> Path:
    """A template with placeholders split across runs, a repeating table row, a conditional and a repeated block."""
    b = []
    b.append(_p(_r("Invoice {{") + _r("invoice.number", "<w:b/>") + _r("}}"), "Title"))
    b.append(_p(_r("Date: {{ date | date:\"%d %B %Y\" }}")))
    b.append(_p(_r("Bill to: {{client.name | upper}}, {{client.city}}")))
    grid = '<w:tblGrid><w:gridCol w:w="4500"/><w:gridCol w:w="2000"/><w:gridCol w:w="2500"/></w:tblGrid>'
    b.append('<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="9000" w:type="dxa"/></w:tblPr>' + grid
             + '<w:tr>' + "".join(f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/></w:tcPr>' + _p(_r(t)) + "</w:tc>" for t, w in (("Item", 4500), ("Qty", 2000), ("Amount", 2500))) + "</w:tr>"
             + '<w:tr>' + '<w:tc><w:tcPr><w:tcW w:w="4500" w:type="dxa"/></w:tcPr>' + _p(_r("{{#each items}}{{name}}")) + "</w:tc>"
             + '<w:tc><w:tcPr><w:tcW w:w="2000" w:type="dxa"/></w:tcPr>' + _p(_r("{{qty}}")) + "</w:tc>"
             + '<w:tc><w:tcPr><w:tcW w:w="2500" w:type="dxa"/></w:tcPr>' + _p(_r("{{amount | number:2}}{{/each}}")) + "</w:tc></w:tr>"
             + "</w:tbl>")
    b.append(_p(_r("Total: {{total | number:2}} EUR")))
    b.append(_p(_r("{{#if paid}}")))
    b.append(_p(_r("Thank you, this invoice is paid.")))
    b.append(_p(_r("{{else}}")))
    b.append(_p(_r("Please pay within {{terms}} days.")))
    b.append(_p(_r("{{/if}}")))
    b.append(_p(_r("{{#each notes}}")))
    b.append(_p(_r("Note {{@number}}: {{this}}"), "ListParagraph"))
    b.append(_p(_r("{{/each}}")))
    b.append(_p(_r("Tags: {{#each tags}}{{this}}{{#unless @last}}, {{/unless}}{{/each}}.")))
    b.append(_p(_r("Missing: {{nothere}}")))
    sect = '<w:sectPr><w:headerReference w:type="default" r:id="rId6"/><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>'
    doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:document {NSDECL}><w:body>{"".join(b)}{sect}</w:body></w:document>'
    header = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:hdr {NSDECL}>{_p(_r("{{company}} — invoice {{invoice.number}}"))}</w:hdr>'
    rels = DOC_RELS.replace('<Relationship Id="rId7" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>', "")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", CT.replace('<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>', ""))
        z.writestr("_rels/.rels", ROOT_RELS.replace('<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>', "").replace('<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>', ""))
        z.writestr("word/_rels/document.xml.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/styles.xml", STYLES)
        z.writestr("word/numbering.xml", NUMBERING)
        z.writestr("word/settings.xml", SETTINGS)
        z.writestr("word/footnotes.xml", FOOTNOTES)
        z.writestr("word/comments.xml", COMMENTS)
        z.writestr("word/header1.xml", header)
        z.writestr("word/media/image1.png", png_bytes())
    return path


REPORT_MD = """# Introduction

Some **bold** and *italic* text with a footnote.[^n]

- one
- two
  - nested

1. first
2. second

| Name | Value |
|------|------:|
| alpha | 1 |
| beta | 22 |

![A red box](red.png){width=4cm}

$$a^2+b^2=c^2$$

\\pagebreak

# Results

```python
print("hi")
```

[^n]: The note.
"""

RTF = r"""{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}
{\pard\b Quarterly memo\b0\par}
{\pard The budget is approved.\par}
}"""


def spec_json() -> dict:
    return {
        "properties": {"title": "Spec doc", "author": "Selftest"},
        "page": {"size": "A4", "margins": "2cm"},
        "header": "Desk||Spec",
        "footer": "Page {page} of {pages}",
        "blocks": [
            {"type": "title", "text": "Spec doc", "subtitle": "Built from JSON"},
            {"type": "heading", "text": "Summary", "level": 1},
            {"type": "paragraph", "runs": ["Amount: ", {"text": "1,333.50 EUR", "bold": True, "color": "C00000"}, {"text": " due.", "footnote": "Within 30 days."}]},
            {"type": "bullets", "items": ["Fast", {"text": "Reliable", "items": ["Uptime"]}]},
            {"type": "numbered", "items": ["First", "Second"], "format": "upperRoman", "start": 3},
            {"type": "table", "header": ["Item", "Qty", "Amount"], "rows": [["Widget", 3, "1,234.50"], [{"text": "Total", "bold": True}, "", "1,333.50"]],
             "widths": ["8cm", "3cm", "4cm"], "caption": "Table 1: items", "merge": [{"from": [1, 0], "to": [1, 1]}], "shading": {"header": "DDEEFF"}},
            {"type": "image", "path": "red.png", "width": "3cm", "caption": "Figure 1: box"},
            {"type": "markdown", "text": "Some **Markdown** with a [link](https://example.com)."},
            {"type": "toc", "depth": 2},
            {"type": "section", "orientation": "landscape"},
            {"type": "heading", "text": "Annex", "level": 1},
            {"type": "paragraph", "text": "Wide.", "align": "center"},
        ],
    }


TEMPLATE_DATA = {"invoice": {"number": 42}, "date": "2026-09-24", "client": {"name": "Ada", "city": "Paris"},
                 "items": [{"name": "Apple", "qty": 1, "amount": 10}, {"name": "Banana", "qty": 2, "amount": 20.5}],
                 "total": 30.5, "paid": False, "terms": 30, "notes": ["n1", "n2"], "tags": ["x", "y", "z"], "company": "ACME"}


def build_fixtures(d: Path) -> dict[str, Path]:
    d.mkdir(parents=True, exist_ok=True)
    (d / "red.png").write_bytes(png_bytes((200, 40, 40), (400, 200)))
    (d / "blue.png").write_bytes(png_bytes((40, 60, 200), (300, 300)))
    (d / "report.md").write_text(REPORT_MD, encoding="utf-8")
    (d / "spec.json").write_text(json.dumps(spec_json(), indent=1), encoding="utf-8")
    (d / "data.json").write_text(json.dumps(TEMPLATE_DATA), encoding="utf-8")
    (d / "memo.rtf").write_text(RTF, encoding="utf-8")
    (d / "page.html").write_text("<html><body><h1>Web page</h1><p>Hello <b>world</b>.</p><ul><li>a</li><li>b</li></ul></body></html>", encoding="utf-8")
    (d / "notdocx.docx").write_text("hello, not a document\n", encoding="utf-8")
    (d / "locked.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504 + "EncryptedPackage".encode("utf-16-le") + b"\0" * 512)
    (d / "legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 1016)
    return {
        "complex": write_complex_docx(d / "complex.docx"), "template": write_template_docx(d / "template.docx"),
        "red": d / "red.png", "blue": d / "blue.png", "md": d / "report.md", "spec": d / "spec.json", "data": d / "data.json",
        "rtf": d / "memo.rtf", "html": d / "page.html", "notdocx": d / "notdocx.docx", "locked": d / "locked.docx", "legacy": d / "legacy.doc",
    }


# ── runner ──────────────────────────────────────────────────────────────

PY = sys.executable
NO_LO = {"DESK_SOFFICE": "none"}
SCRIPTS = ["docx_info.py", "docx_read.py", "docx_create.py", "docx_edit.py", "docx_template.py", "docx_compare.py", "docx_render.py", "docx_convert.py"]
HEAVY = {"docx", "lxml", "PIL", "typst", "pypdfium2", "mammoth", "docxcompose", "pypandoc", "regex"}


class Checks:
    def __init__(self) -> None:
        import threading

        self.n = 0
        self.failures: list[str] = []
        self._lock = threading.Lock()

    def ok(self, cond: object, what: str) -> None:
        with self._lock:
            self.n += 1
            if not cond:
                self.failures.append(what)
                print(f"FAIL: {what}", file=sys.stderr)


def run(args: list, env: dict | None = None, expect: int = 0, stdin: str | None = None, cwd: Path | None = None, lo: bool = False) -> subprocess.CompletedProcess:
    """Runs scripts/<args[0]> as an agent would. LibreOffice is hidden unless lo=True. An expected failure must print
    one clean `error: …` line (no traceback)."""
    cmd = [PY, str(HERE / str(args[0])), *[str(a) for a in args[1:]]]
    e = dict(os.environ)
    if not lo:
        e.update(NO_LO)
    e.update(env or {})
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env=e, input=stdin, cwd=cwd, timeout=300)
    if r.returncode != expect:
        raise AssertionError(f"{' '.join(str(a) for a in args[:4])} … exited {r.returncode} (expected {expect}):\n{r.stderr[-1500:]}\n{r.stdout[-600:]}")
    if expect != 0 and ("Traceback" in r.stderr or "error:" not in r.stderr):
        raise AssertionError(f"{args[0]}: the error is not one clean message: {r.stderr[-800:]}")
    return r


def jrun(args: list, **kw) -> dict:
    return json.loads(run([*args, "--format", "json"], **kw).stdout)


def read_md(path: Path, *extra: str) -> str:
    return run(["docx_read.py", path, *extra]).stdout


def sha(p: Path) -> str:
    import hashlib

    return hashlib.sha256(p.read_bytes()).hexdigest()


def png_size(p: Path) -> tuple[int, int]:
    from PIL import Image

    with Image.open(p) as im:
        return im.size


def pdf_pages(p: Path) -> int:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(p))
    try:
        return len(doc)
    finally:
        doc.close()


def norm(s: str) -> str:
    return " ".join(s.split())


def lo_available() -> bool:
    sys.path.insert(0, str(HERE))
    from _render import find_soffice

    return bool(find_soffice())


def main() -> int:
    import argparse
    from concurrent.futures import ThreadPoolExecutor

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--keep", action="store_true", help="keep the temp folder and print its path")
    ap.add_argument("--only", help="run only the tests whose name contains this")
    ap.add_argument("--fixtures", metavar="DIR", help="only write the fixtures into DIR")
    ap.add_argument("--no-lo", action="store_true", help="skip the LibreOffice tests even when it is installed")
    ap.add_argument("--times", action="store_true", help="print how long each test took")
    a = ap.parse_args()
    if a.fixtures:
        fx = build_fixtures(Path(a.fixtures))
        print("\n".join(str(p) for p in fx.values()))
        return 0
    t0 = time.time()
    tmp = Path(tempfile.mkdtemp(prefix="word-documents-selftest-"))
    # A private, cold file cache: the cache tests measure the first and the second call.
    os.environ["DESK_FILE_CACHE"] = str(tmp / "cache")
    os.environ.pop("DESK_NO_CACHE", None)
    c = Checks()
    try:
        fx = build_fixtures(tmp / "fx")
        before = {k: sha(p) for k, p in fx.items()}
        have_lo = lo_available() and not a.no_lo
        tests = [(n, f) for n, f in globals().items() if n.startswith("test_") and callable(f) and (not a.only or a.only in n)]

        def one(item: tuple) -> None:
            name, fn = item
            work = tmp / name
            work.mkdir()
            started = time.time()
            try:
                fn(c, fx, work)
            except AssertionError as e:
                c.ok(False, f"{name}: {e}")
            except Exception as e:  # noqa: BLE001
                c.ok(False, f"{name}: {type(e).__name__}: {e}")
            if a.times:
                print(f"{name}: {time.time() - started:.1f} s", file=sys.stderr)

        # The slowest tests start first; the LibreOffice tests run one after another (one LibreOffice profile) in a
        # worker of their own, alongside the others.
        slow = ("test_big_and_cache", "test_regress_create", "test_regress_edit_compare", "test_regress_safety", "test_create_markdown", "test_create_spec", "test_render_builtin", "test_edit", "test_edit_tracked")
        plain = sorted((t for t in tests if not t[0].startswith("test_lo_")), key=lambda t: (t[0] not in slow, slow.index(t[0]) if t[0] in slow else 0))
        lo_tests = [t for t in tests if t[0].startswith("test_lo_")] if have_lo else []

        def lo_group(_: object) -> None:
            for t in lo_tests:
                one(t)

        try:
            cap = max(1, int(os.environ.get("DESK_MAX_WORKERS") or 4))
        except ValueError:
            cap = 4
        # One thread per running test (each runs one script at a time), bounded like the scripts' own pools.
        with ThreadPoolExecutor(max_workers=max(2, min(6, cap + 1, os.cpu_count() or 2))) as pool:
            futures = [pool.submit(lo_group, None)] if lo_tests else []
            futures += [pool.submit(one, t) for t in plain]
            for f in futures:
                f.result()
        c.ok(all(sha(p) == before[k] for k, p in fx.items()), "no script modified an input file")
    finally:
        if a.keep:
            print(f"kept {tmp}", file=sys.stderr)
        else:
            shutil.rmtree(tmp, ignore_errors=True)
    dt = time.time() - t0
    if c.failures:
        print(f"FAILED: {len(c.failures)} of {c.n} checks failed in {dt:.1f} s", file=sys.stderr)
        return 1
    print(f"ok: {c.n} checks in {dt:.1f} s" + ("" if have_lo else " (LibreOffice tests skipped: not installed or --no-lo)"))
    return 0


# ── tests ───────────────────────────────────────────────────────────────


def test_help(c: Checks, fx: dict, w: Path) -> None:
    for s in SCRIPTS:
        t = time.time()
        r = run([s, "--help"])
        dt = time.time() - t
        c.ok("Examples:" in r.stdout and "python3 scripts/" + s in r.stdout, f"{s} --help shows examples")
        c.ok(dt < 3.0, f"{s} --help took {dt:.2f} s")
        imp = subprocess.run([PY, "-X", "importtime", str(HERE / s), "--help"], capture_output=True, text=True, encoding="utf-8", env=dict(os.environ, **NO_LO), timeout=60)
        loaded = {ln.rsplit("|", 1)[-1].strip().split(".")[0] for ln in imp.stderr.splitlines() if ln.startswith("import time:")}
        c.ok(not (loaded & HEAVY), f"{s} --help imports no heavy module ({sorted(loaded & HEAVY)})")
    run(["docx_read.py"], expect=2)
    run(["docx_edit.py", fx["complex"], w / "x.docx", "--ops", "[{\"op\": \"nope\"}]"], expect=2)


def test_info(c: Checks, fx: dict, w: Path) -> None:
    md = run(["docx_info.py", fx["complex"]]).stdout
    c.ok("pages: 2 as last saved" in md, "info: saved page count")
    c.ok("Tracked changes: insertions 1, deletions 1 by Bob" in md, "info: tracked changes")
    c.ok("Comments: 1 by Alice Martin" in md, "info: comments")
    j = jrun(["docx_info.py", fx["complex"]])
    c.ok(j["blocks"] == 18 and j["headings"]["count"] == 3, "info: blocks and headings")
    c.ok(j["tables"] and j["tables"][0]["merged_cells"] and j["tables"][0]["rows"] == 3, "info: merged table")
    c.ok(j["images"]["inline"] == 1 and j["objects"].get("text_boxes") == 1 and j["objects"].get("equations") == 1, "info: images, text box, equation")
    c.ok(j["footnotes"] == 1 and j["sections"][0]["paper"] == "A4", "info: footnotes, paper")
    c.ok(j["headers"][0]["text"].startswith("ACME Ltd") and "PAGE" in j["header_footer_fields"], "info: headers and fields")
    c.ok(j["properties"]["author"] == "Alice Martin" and j["application"].get("Company") == "ACME", "info: properties")
    ex = jrun(["docx_info.py", fx["complex"], "--exact"])
    c.ok(ex["pages"].get("rendered") == 2 and "built-in" in ex["pages"].get("engine", ""), f"info --exact without LibreOffice: {ex['pages']}")


def test_read(c: Checks, fx: dict, w: Path) -> None:
    doc = fx["complex"]
    md = read_md(doc)
    for want in ("# Service Agreement", "# Parties", "This Agreement is made between **AC*ME*** Ltd", "1. First obligation", "   a) Sub obligation",
                 "2. Second obligation", "- Bullet point", "The term is {--30--}{++60++} days.", "{==payment==}{>>Alice Martin: Check the payment terms.<<}",
                 "[^1]: Payable in euros to ACME Ltd.", "[terms](https://example.com/terms)", '<td colspan="2">Item and region</td>', '<td rowspan="2">Widgets</td>',
                 "Box text for ACME Ltd", "![Company logo](media/image1.png)", "$E=mc^2$", "<!-- page break -->", "> Fees are due quarterly",
                 "<!-- header: ACME Ltd — Confidential -->", "<!-- footer: Page {PAGE} of {NUMPAGES} -->"):
        c.ok(want in md, f"read: {want!r}")
    out = read_md(doc, "--outline")
    c.ok("3 headings" in out and "[15] Payment" in out, "read --outline")
    j = jrun(["docx_read.py", doc, "--find", "ACME Ltd"])
    where = {h["address"] for h in j["hits"]}
    c.ok({"block 2", "block 16", "block 11", "footnote 1"} <= where and any(a.startswith("header") for a in where), f"read --find across runs, text boxes, notes, headers: {sorted(where)}")
    found = run(["docx_read.py", doc, "--grep", r"(\d+) days", "--context", "12"]).stdout
    c.ok("[7] (Parties)" in found and "**60 days**" in found and "**30 days**" in found and "--blocks 4-10" in found, f"read --grep: addresses, section, context, next command: {found[:300]!r}")
    sec = read_md(doc, "--section", "payment")
    c.ok(sec.startswith("# Payment") and "Fees are due" in sec and "Parties" not in sec, "read --section by heading text")
    c.ok(read_md(doc, "--section", "8").startswith("# Parties"), "read --section by block index")
    csv_out = run(["docx_read.py", doc, "--blocks", "10", "--format", "csv"]).stdout.splitlines()
    c.ok(csv_out == ["Item and region,Item and region,Price", 'Widgets,North,"1,200"', "Widgets,South,950"], f"read --format csv repeats merged cells: {csv_out}")
    cutj = json.loads(run(["docx_read.py", doc, "--format", "json", "--max-chars", "2500"]).stdout)
    c.ok(cutj["truncated"]["next"].startswith("python3 scripts/docx_read.py") and "--blocks" in cutj["truncated"]["next"] and len(cutj["blocks"]) < 18, "read json over --max-chars: valid JSON with the next command")
    j = jrun(["docx_read.py", doc, "--runs"])
    b2 = j["blocks"][2]
    c.ok(j["total_blocks"] == 18 and b2["index"] == 2 and b2["text"] == "This Agreement is made between ACME Ltd and the Client.", "read json: plain text of block 2")
    c.ok(any(r["text"] == "ME" and r.get("bold") and r.get("italic") for r in b2["runs"]), "read --runs: effective formatting")
    c.ok(j["blocks"][3]["list"]["label"] == "1." and j["blocks"][4]["list"]["label"] == "a)", "read json: list labels")
    c.ok(j["blocks"][10]["type"] == "table" and j["blocks"][10]["cells"][0][0]["colspan"] == 2, "read json: table cells")
    c.ok(j["blocks"][7]["text"] == "The term is 60 days.", f"read json: text reads as now ({j['blocks'][7]['text']!r})")
    acc = read_md(doc, "--changes", "accept", "--comments", "none")
    c.ok("The term is 60 days." in acc and "{>>" not in acc and "{++" not in acc, "read --changes accept --comments none")
    rej = read_md(doc, "--changes", "reject", "--comments", "end")
    c.ok("The term is 30 days." in rej, "read --changes reject")
    c.ok("Comments:\n- [c0] Alice Martin: Check the payment terms." in rej, "read --comments end")
    txt = run(["docx_read.py", doc, "--format", "text"]).stdout
    c.ok("Service Agreement" in txt and "**" not in txt and "{--" not in txt and "1. First obligation" in txt and "The term is 60 days." in txt, "read --format text (revisions as they read now)")

    run(["docx_read.py", doc, "--extract-media", w / "media"])
    c.ok((w / "media" / "image1.png").exists(), "read --extract-media saves pictures")
    c.ok("# complex.docx: map" in read_md(doc, "--max-chars", "400"), "read: over --max-chars, a map")
    cut = read_md(doc, "--max-chars", "400", "--full")
    import re

    nxt = re.search(r"Continue with: python3 scripts/docx_read\.py \S+ --max-chars 400 --blocks (\d+)-17\]", cut)
    c.ok(nxt is not None and len(cut) < 900, f"read --max-chars cuts at a block with the next command: {cut[-200:]!r}")
    if nxt:
        rest = read_md(doc, "--max-chars", "400", "--blocks", f"{nxt.group(1)}-17")
        c.ok(rest.startswith(read_md(doc, "--blocks", nxt.group(1), "--max-chars", "0").strip()[:20]), "the continue command reads on from where the cut was")
    part = read_md(doc, "--blocks", "7-8,15", "--indexes")
    c.ok("[7] The term is" in part and "[15] # Payment" in part and "Service Agreement" not in part and "[^1]: Payable" in part, "read --blocks --indexes keeps the notes it refers to")


def test_read_other_formats(c: Checks, fx: dict, w: Path) -> None:
    r = run(["docx_read.py", fx["rtf"]])
    c.ok("Quarterly memo" in r.stdout and "The budget is approved." in r.stdout and "pandoc" in r.stderr, "read .rtf without LibreOffice (pandoc)")
    info = jrun(["docx_info.py", fx["rtf"]])
    # One pandoc conversion shared by both scripts (the LibreOffice tests may add their own entry alongside).
    markers = (Path(os.environ["DESK_FILE_CACHE"]) / "docx-convert").glob("*/*/.complete")
    conv = [m for m in markers if json.loads(m.read_text(encoding="utf-8"))["params"].get("engine") == "pandoc"]
    c.ok(info["words"] >= 5 and len(conv) == 1, f"the converted copy of an .rtf is cached and shared by the scripts ({len(conv)} entries)")
    renamed = w / "memo.docx"
    shutil.copyfile(fx["rtf"], renamed)
    r = run(["docx_read.py", renamed])
    c.ok("The budget is approved." in r.stdout and "really an RTF file" in r.stderr, "read an RTF saved as .docx")
    r = run(["docx_read.py", fx["notdocx"]], expect=1)
    c.ok("not a Word document" in r.stderr, "read: a text file named .docx fails cleanly")
    r = run(["docx_info.py", fx["locked"]], expect=1)
    c.ok("password-protected" in r.stderr, "info: an encrypted file fails with a clear message")
    r = run(["docx_read.py", w / "missing.docx"], expect=1)
    c.ok("does not exist" in r.stderr or "not found" in r.stderr, "read: missing file")
    r = run(["docx_read.py", fx["legacy"]], expect=1)
    c.ok("LibreOffice" in r.stderr, "read .doc without LibreOffice says what is needed")


def test_create_markdown(c: Checks, fx: dict, w: Path) -> None:
    out = w / "report.docx"
    r = run(["docx_create.py", fx["md"], out, "--title", "Selftest report", "--author", "Desk", "--toc", "--page-numbers", "Page {page} of {pages}",
             "--header", "Desk||Selftest"], cwd=w)
    c.ok("table of contents: 2 entries" in r.stdout and "built-in renderer" in r.stdout, f"create: TOC filled by the built-in renderer: {r.stdout.strip()}")
    md = read_md(out)
    for want in ("# Selftest report", "# Introduction", "Some **bold** and *italic* text", "- one", "   - nested", "1. first", "| Name | Value |",
                 "| beta | 22 |", "[^1]: The note.", "![A red box]", "$$", "```\nprint(\"hi\")\n```", "<!-- footer: Page {PAGE} of {NUMPAGES} -->", "Selftest -->", "<!-- page break -->"):
        c.ok(want in md, f"create md: {want!r}")
    import re

    pages = {m.group(1): int(m.group(2)) for m in re.finditer(r"- (Introduction|Results) … (\d+)", md)}
    c.ok(len(pages) == 2 and pages["Results"] > pages["Introduction"] >= 1, f"create: TOC page numbers {pages}")
    j = jrun(["docx_info.py", out])
    c.ok(j["properties"].get("title") == "Selftest report" and j["properties"].get("author") == "Desk", "create: properties")
    c.ok(j["sections"][0]["paper"] == "A4" and j["tables"] and j["images"]["inline"] == 1 and j["footnotes"] == 1, "create: A4, table, image, footnote")
    r = run(["docx_create.py", fx["md"], out], expect=1, cwd=w)
    c.ok("exists" in r.stderr and "--force" in r.stderr, "create refuses to overwrite without --force")
    run(["docx_create.py", fx["md"], out, "--force", "--style", "classic"], cwd=w)
    # inline Markdown, Letter landscape, template output
    out2 = w / "letter.dotx"
    run(["docx_create.py", "--markdown", "# Hello\n\nBody text.", out2, "--size", "Letter", "--landscape", "--margins", "2cm", "--style", "modern"])
    j = jrun(["docx_info.py", out2])
    c.ok(j["format"] == "dotx" and j["sections"][0]["paper"] == "Letter" and j["sections"][0]["orientation"] == "landscape", f"create: .dotx, Letter landscape ({j['format']}, {j['sections'][0]['paper']})")
    with zipfile.ZipFile(out2) as z:
        c.ok("template.main+xml" in z.read("[Content_Types].xml").decode(), "create: .dotx has the template content type")
    out3 = w / "stdin.docx"
    run(["docx_create.py", "-", out3, "--template", fx["complex"]], stdin="# From stdin\n\nWith the template's header.\n")
    md3 = read_md(out3)
    c.ok("# From stdin" in md3 and "<!-- header: ACME Ltd — Confidential -->" in md3, "create from stdin with --template keeps its header")
    r = run(["docx_create.py", fx["md"], fx["complex"]], expect=1)
    c.ok("input" in r.stderr.lower() or "exists" in r.stderr, "create refuses to write over an input")


def test_create_spec(c: Checks, fx: dict, w: Path) -> None:
    out = w / "spec.docx"
    r = run(["docx_create.py", "--spec", fx["spec"], out])
    c.ok("JSON spec" in r.stdout, "create --spec")
    md = read_md(out)
    for want in ("# Spec doc", "Built from JSON", "# Summary", "**1,333.50 EUR**", "[^1]: Within 30 days.", "- Fast", "   - Uptime", "III. First", "IV. Second",
                 "Table 1: items", '<td colspan="2">**Total**</td>', "![Figure 1: box]", "[link](https://example.com)", "- Summary … 1", "# Annex", "<!-- header: Desk\tSpec -->"):
        c.ok(want in md, f"create spec: {want!r}")
    j = jrun(["docx_info.py", out])
    c.ok(len(j["sections"]) == 2 and j["sections"][1]["orientation"] == "landscape" and j["sections"][0]["orientation"] == "portrait", "create spec: landscape section")
    r = run(["docx_create.py", "--spec", '{"blocks": [{"type": "heading", "text": "ok"}, {"type": "nope"}]}', w / "bad.docx"], expect=2)
    c.ok("block 1" in r.stderr and not (w / "bad.docx").exists(), "create spec: a bad block names its index and writes nothing")


EDIT_OPS = [
    {"op": "replace", "find": "ACME Ltd", "replace": "Acme Limited"},
    {"op": "replace", "regex": r"(\d+) days", "replace": "$1 calendar days", "scope": "body"},
    {"op": "insert_after", "index": 9, "markdown": "New **bold** paragraph."},
    {"op": "insert_after", "index": 9, "text": "Second inserted", "style": "Quote"},
    {"op": "delete", "index": 6},
    {"op": "move", "index": 17, "to": 15, "position": "before"},
    {"op": "set_style", "index": 9, "align": "center"},
    {"op": "format", "find": "Payment", "bold": True, "color": "C00000"},
    {"op": "table_set_cell", "table": 0, "row": 1, "col": 2, "text": "1,300", "bold": True},
    {"op": "table_add_row", "table": 0, "cells": ["Gadgets", "East", "400"]},
    {"op": "insert_image", "path": "PLACEHOLDER", "after": 12, "width": "3cm", "caption": "Figure 2"},
    {"op": "add_comment", "find": "quarterly", "text": "Why quarterly?", "author": "Desk"},
    {"op": "set_footer", "text": "Acme||Page {page} of {pages}"},
    {"op": "page_setup", "size": "Letter", "orientation": "landscape"},
    {"op": "properties", "title": "Edited agreement", "custom": {"Client": "Acme"}},
    {"op": "replace", "find": "NOT PRESENT", "replace": "x", "optional": True},
]


def test_edit(c: Checks, fx: dict, w: Path) -> None:
    ops = [dict(o) for o in EDIT_OPS]
    ops[10]["path"] = str(fx["red"])
    out = w / "edited.docx"
    j = jrun(["docx_edit.py", fx["complex"], out, "--ops", json.dumps(ops)])
    counts = [r["count"] for r in j["operations"]]
    c.ok(counts[0] == 5, f"edit: replace across runs, text box, header, footnote ({counts[0]})")
    c.ok(counts[1] == 2 and counts[-1] == 0, f"edit: regex count and optional op ({counts[1]}, {counts[-1]})")
    md = read_md(out)
    for want in ("This Agreement is made between **Acme Limited** and the Client.", "60 calendar days", "within 30 calendar days", "Box text for Acme Limited",
                 "<!-- header: Acme Limited — Confidential -->", "[^1]: Payable in euros to Acme Limited.", "online.\n\nNew **bold** paragraph.\n\n> Second inserted",
                 "Tab\tseparated\tvalues\n\n# **Payment**", "**1,300**", "<td>Gadgets</td><td>East</td><td>400</td>", "Figure 2", "{==quarterly==}{>>Desk: Why quarterly?<<}",
                 "<!-- footer: Acme\tPage {PAGE} of {NUMPAGES} -->"):
        c.ok(want in md, f"edit result: {want!r}")
    c.ok("Bullet point" not in md, "edit: delete")
    info = jrun(["docx_info.py", out])
    c.ok(info["images"]["inline"] == 2 and info["comments"]["count"] == 2, "edit: image and comment added")
    c.ok(info["sections"][0]["paper"] == "Letter" and info["sections"][0]["orientation"] == "landscape", "edit: page setup")
    c.ok(info["properties"]["title"] == "Edited agreement" and info["custom_properties"].get("Client") == "Acme", "edit: properties")
    b = jrun(["docx_read.py", out, "--runs", "--blocks", "2"])["blocks"][0]
    c.ok(any(r["text"].startswith("Acme Limited") and r.get("bold") for r in b["runs"]), "edit: replacement keeps the first run's bold")
    # failures write nothing
    bad = w / "bad.docx"
    r = run(["docx_edit.py", fx["complex"], bad, "--ops", '[{"op":"replace","find":"ACME","replace":"x"},{"op":"replace","find":"zzz","replace":"y"}]'], expect=1)
    c.ok("op 1 (replace)" in r.stderr and "matched nothing" in r.stderr and not bad.exists(), "edit: an op matching nothing fails and writes nothing")
    r = run(["docx_edit.py", fx["complex"], bad, "--ops", '[{"op":"delete","index":99}]'], expect=1)
    c.ok("does not exist" in r.stderr and not bad.exists(), "edit: bad index")
    run(["docx_edit.py", fx["complex"], "--ops", '[{"op":"replace","find":"ACME","replace":"x"}]', "--dry-run"])
    r = run(["docx_edit.py", fx["complex"], fx["complex"], "--ops", '[{"op":"replace","find":"ACME","replace":"x"}]', "--force"], expect=1)
    c.ok("input" in r.stderr.lower(), "edit refuses to overwrite its input, even with --force")
    # more operations
    out2 = w / "more.docx"
    ops2 = [{"op": "table_add_column", "table": 0, "cells": ["Tax", "5%", "7%"]}, {"op": "table_delete_row", "table": 0, "row": 2},
            {"op": "set_text", "index": 17, "text": "Replaced text"}, {"op": "remove_comments"}, {"op": "accept_changes"}, {"op": "remove_personal_info"},
            {"op": "set_style", "index": 13, "style": "Heading 2"}, {"op": "append", "markdown": "## Annex\n\n| A | B |\n|---|---|\n| 1 | 2 |"},
            {"op": "prepend", "text": "DRAFT"}, {"op": "page_break", "after": 5}, {"op": "update_fields"}, {"op": "replace_image", "image": 1, "path": str(fx["blue"])},
            {"op": "append_docx", "path": str(fx["template"])}]
    j = jrun(["docx_edit.py", fx["complex"], out2, "--ops", json.dumps(ops2)])
    c.ok(all(r["count"] >= 1 for r in j["operations"]), f"edit: second battery counts {[r['count'] for r in j['operations']]}")
    md2 = read_md(out2)
    for want in ("DRAFT\n\n# Service Agreement", "Tax", "Replaced text", "## Energy", "## Annex", "| 1 | 2 |", "The term is 60 days.", "Bill to: {{client.name | upper}}"):
        c.ok(want in md2, f"edit battery 2: {want!r}")
    c.ok("{>>" not in md2 and "{++" not in md2 and "South" not in md2, "edit: comments removed, changes accepted, row deleted")
    info2 = jrun(["docx_info.py", out2])
    c.ok(not info2["properties"].get("author") and not info2["tracked_changes"].get("insertions") and not info2["comments"].get("count"), f"edit: personal info removed {info2['properties']}")


def test_edit_tracked(c: Checks, fx: dict, w: Path) -> None:
    out = w / "tracked.docx"
    ops = [{"op": "replace", "find": "quarterly", "replace": "monthly"}, {"op": "delete", "index": 6}, {"op": "insert_after", "index": 9, "markdown": "Inserted para."},
           {"op": "set_text", "index": 17, "text": "Tab free"}, {"op": "table_add_row", "table": 0, "cells": ["Gadgets", "East", "400"]},
           {"op": "format", "find": "Parties", "italic": True}]
    run(["docx_edit.py", fx["complex"], out, "--ops", json.dumps(ops), "--track", "--author", "Tester"])
    md = read_md(out)
    for want in ("{--quarterly--}{++monthly++}", "{--Bullet point--}", "{++Inserted para.++}", "{--Tab\tseparated\tvalues--}{++Tab free++}"):
        c.ok(want in md, f"tracked edit: {want!r}")
    info = jrun(["docx_info.py", out])
    c.ok(info["tracked_changes"]["authors"].get("Tester", 0) >= 5, f"tracked edit: revisions by Tester {info['tracked_changes']}")
    acc, rej = w / "accepted.docx", w / "rejected.docx"
    run(["docx_edit.py", out, acc, "--ops", '[{"op":"accept_changes"}]'])
    run(["docx_edit.py", out, rej, "--ops", '[{"op":"reject_changes"}]'])
    a = read_md(acc)
    c.ok("monthly" in a and "quarterly" not in a and "Bullet point" not in a and "Inserted para." in a and "Tab free" in a and "Gadgets" in a and "{++" not in a, "accept all")
    rtext = run(["docx_read.py", rej, "--format", "text"]).stdout
    otext = run(["docx_read.py", fx["complex"], "--format", "text", "--changes", "reject"]).stdout
    c.ok(norm(rtext) == norm(otext), "reject all gives back the original text")
    part = w / "bob-only.docx"
    j = jrun(["docx_edit.py", out, part, "--ops", '[{"op":"accept_changes","author":"Bob"}]'])
    c.ok(j["operations"][0]["count"] == 2, f"accept by author ({j['operations'][0]['count']})")
    c.ok("{--quarterly--}{++monthly++}" in read_md(part) and "{--30--}" not in read_md(part), "accept by author leaves the others")


def test_template(c: Checks, fx: dict, w: Path) -> None:
    j = jrun(["docx_template.py", fx["template"], "--fields"])
    c.ok({"invoice.number", "client.name", "nothere", "company"} <= set(j["fields"]), f"template --fields: {sorted(j['fields'])}")
    c.ok(any(b["kind"] == "each" and b["expr"] == "items" for b in j["blocks"]) and any(b["kind"] == "if" for b in j["blocks"]), "template --fields: loops and conditions")
    out = w / "filled.docx"
    j = jrun(["docx_template.py", fx["template"], out, "--data", fx["data"]])
    c.ok(j["unfilled"] == {"nothere": 1} and j["loops"] == 3 and j["conditions"] == 4, f"template report: {j}")
    md = read_md(out)
    for want in ("<!-- header: ACME — invoice 42 -->", "# Invoice 42", "Date: 24 September 2026", "Bill to: ADA, Paris", "| Apple | 1 | 10.00 |", "| Banana | 2 | 20.50 |",
                 "Total: 30.50 EUR", "Please pay within 30 days.", "Note 1: n1", "Note 2: n2", "Tags: x, y, z.", "Missing: {{nothere}}"):
        c.ok(want in md, f"template: {want!r}")
    c.ok("Thank you" not in md and "{{#" not in md and "{{/" not in md, "template: else branch only, no control tags left")
    data = dict(TEMPLATE_DATA, paid=True, items=[])
    out2 = w / "paid.docx"
    run(["docx_template.py", fx["template"], out2, "--data", json.dumps(data), "--blank-missing"])
    md2 = read_md(out2)
    c.ok("Thank you, this invoice is paid." in md2 and "Please pay" not in md2 and "Missing:" in md2 and "{{nothere}}" not in md2, "template: if branch, --blank-missing")
    c.ok("Apple" not in md2 and "| Item | Qty | Amount |" in md2, "template: an empty list removes the loop rows")
    r = run(["docx_template.py", fx["template"], w / "strict.docx", "--data", fx["data"], "--strict"], expect=1)
    c.ok("nothere" in r.stderr and not (w / "strict.docx").exists(), "template --strict fails on unfilled placeholders")
    # nested loops, outer data, conditions inside table cells, filters
    from docx import Document

    d = Document()
    for line in ("Report {{title}}", "{{#each groups}}", "Group {{name}} of {{../title}} ({{@number}}/{{groups | count}})", "{{#each members}}", "Member {{this | upper}}", "{{/each}}", "{{/each}}"):
        d.add_paragraph(line)
    t = d.add_table(rows=3, cols=2)
    for (r_, c_), text in {(0, 0): "Name", (0, 1): "Late", (1, 0): "{{#each people}}{{name}}", (1, 1): "{{#if late}}yes{{else}}no{{/if}}{{/each}}", (2, 0): "Total", (2, 1): "{{people | count}}"}.items():
        t.cell(r_, c_).text = text
    d.add_paragraph('{{#if total > 1000 and not paid}}Big unpaid{{else}}Fine{{/if}} {{amount | currency:"€":2:after}} {{when | date:us}} {{rate | percent:1}} {{missing | default:"n/a"}}')
    d.save(str(w / "nested.docx"))
    data = {"title": "Q3", "groups": [{"name": "A", "members": ["x", "y"]}, {"name": "B", "members": ["z"]}], "people": [{"name": "Ann", "late": True}, {"name": "Bob", "late": False}],
            "total": 1500, "paid": False, "amount": 1234.5, "when": "2026-09-24", "rate": 0.125}
    r = run(["docx_template.py", w / "nested.docx", w / "nested-out.docx", "--data", json.dumps(data)])
    md = read_md(w / "nested-out.docx")
    for want in ("Group A of Q3 (1/2)", "Member Y", "Group B of Q3 (2/2)", "Member Z", "| Ann | yes |", "| Bob | no |", "| Total | 2 |", "Big unpaid 1,234.50 € September 24, 2026 12.5% n/a"):
        c.ok(want in md, f"template (nested): {want!r}")
    c.ok("warning" not in r.stdout and "{{" not in md, "template (nested): no warnings, no tags left")


def test_compare(c: Checks, fx: dict, w: Path) -> None:
    v1 = fx["complex"]
    v2 = w / "v2.docx"
    ops = [{"op": "accept_changes"}, {"op": "replace", "find": "quarterly", "replace": "monthly"}, {"op": "delete", "index": 6},
           {"op": "insert_after", "index": 9, "text": "Brand new paragraph."}, {"op": "table_set_cell", "table": 0, "row": 2, "col": 2, "text": "999"}]
    run(["docx_edit.py", v1, v2, "--ops", json.dumps(ops)])
    j = jrun(["docx_compare.py", v1, v2])
    s = j["summary"]
    c.ok(s["paragraphs_changed"] >= 1 and s["paragraphs_added"] >= 1 and s["paragraphs_deleted"] >= 1 and s["tables_changed"] == 1, f"compare summary {s}")
    md = run(["docx_compare.py", v1, v2]).stdout
    c.ok("{--quarterly--}{++monthly++}" in md and "Brand new paragraph." in md and "999" in md, "compare report shows word-level changes")
    red = w / "redline.docx"
    run(["docx_compare.py", v1, v2, "--redline", red, "--author", "Reviewer"])
    rmd = read_md(red)
    c.ok("{--quarterly--}{++monthly++}" in rmd and "{++Brand new paragraph.++}" in rmd and "{--Bullet point--}" in rmd, "redline has tracked changes")
    acc = run(["docx_read.py", red, "--format", "text", "--changes", "accept"]).stdout
    rej = run(["docx_read.py", red, "--format", "text", "--changes", "reject"]).stdout
    c.ok(norm(acc) == norm(run(["docx_read.py", v2, "--format", "text"]).stdout), "redline accepted == new version")
    c.ok(norm(rej) == norm(run(["docx_read.py", v1, "--format", "text", "--changes", "accept"]).stdout), "redline rejected == old version")
    c.ok(jrun(["docx_compare.py", v1, v1])["identical"], "compare: identical documents")
    cut = run(["docx_compare.py", v1, v2, "--redline", w / "red-cut.docx", "--max-chars", "2600"]).stdout
    nxt = cut.rsplit("Continue with: ", 1)[-1] if "Continue with: " in cut else ""
    c.ok(nxt.startswith("python3 scripts/docx_compare.py") and "--offset 1" in nxt and "--redline" not in nxt and "Redline:" in cut, f"compare: a cut report ends with the next command, which leaves the redline alone: {nxt!r}")
    rest = run(["docx_compare.py", v1, v2, "--offset", "1"]).stdout
    c.ok("(changes 2-" in rest, "compare --offset reads on")


def test_render_builtin(c: Checks, fx: dict, w: Path) -> None:
    out = w / "pages"
    j = jrun(["docx_render.py", fx["complex"], "--out", out, "--sheet", "--pdf", w / "keep.pdf"])
    c.ok(j["engine"].startswith("built-in") and j["page_count"] == 2 and j["rendered"] == 2, f"render builtin: {j['engine']}, {j['page_count']} pages")
    sizes = [png_size(Path(p)) for p in j["images"]]
    c.ok(all(max(s) <= 1568 and max(s) >= 1200 for s in sizes), f"render: sized for vision {sizes}")
    c.ok(Path(j["sheet"]).exists() and pdf_pages(w / "keep.pdf") == 2, "render: contact sheet and kept PDF")
    md = run(["docx_render.py", fx["complex"], "--out", w / "p2", "--pages", "2", "--dpi", "50"]).stdout
    c.ok("view_image" in md and "built-in" in md, "render prints the view_image hint and the engine")
    pngs = sorted((w / "p2").glob("*.png"))
    c.ok(len(pngs) == 1 and max(png_size(pngs[0])) < 700, "render --pages --dpi")
    r = run(["docx_render.py", fx["legacy"], "--out", w / "p3"], expect=1)
    c.ok("LibreOffice" in r.stderr, "render .doc without LibreOffice fails cleanly")
    # the markdown report, with its table, image, footnote and code block, renders too
    doc = w / "r.docx"
    run(["docx_create.py", fx["md"], doc, "--no-toc-pages"], cwd=w)
    j = jrun(["docx_render.py", doc, "--out", w / "p4"])
    c.ok(j["page_count"] == 2 and j["rendered"] == 2, f"render a created report ({j['page_count']} pages)")


def test_convert_builtin(c: Checks, fx: dict, w: Path) -> None:
    src = fx["complex"]
    r = run(["docx_convert.py", src, w / "c.pdf"])
    c.ok(pdf_pages(w / "c.pdf") == 2 and "built-in" in r.stdout, "convert to PDF (built-in)")
    run(["docx_convert.py", src, w / "c.html"])
    html = (w / "c.html").read_text(encoding="utf-8")
    c.ok("<table" in html and "data:image/png;base64" in html and "E=mc^2" in html and "<h1>" in html, "convert to HTML (mammoth)")
    run(["docx_convert.py", src, w / "c.md"])
    md = (w / "c.md").read_text(encoding="utf-8")
    c.ok("# Service Agreement" in md and "60 days" in md and "{++" not in md and (w / "c_media" / "image1.png").exists(), "convert to Markdown with media")
    run(["docx_convert.py", src, w / "c.txt"])
    txt = (w / "c.txt").read_text(encoding="utf-8")
    c.ok("Service Agreement" in txt and "**" not in txt, "convert to text")
    run(["docx_convert.py", src, w / "c.odt"])
    with zipfile.ZipFile(w / "c.odt") as z:
        c.ok(z.read("mimetype") == b"application/vnd.oasis.opendocument.text", "convert to ODT (pandoc)")
    c.ok("Service Agreement" in read_md(w / "c.odt"), "read the ODT back")
    run(["docx_convert.py", src, w / "c.rtf"])
    c.ok((w / "c.rtf").read_text(encoding="utf-8", errors="replace").startswith("{\\rtf"), "convert to RTF (pandoc)")
    run(["docx_convert.py", src, w / "c.epub"])
    c.ok(zipfile.is_zipfile(w / "c.epub"), "convert to EPUB")
    run(["docx_convert.py", fx["md"], w / "m.docx"], cwd=w)
    c.ok("# Introduction" in read_md(w / "m.docx"), "convert Markdown to docx")
    run(["docx_convert.py", fx["html"], w / "h.docx"])
    hmd = read_md(w / "h.docx")
    c.ok("# Web page" in hmd and "**world**" in hmd and "- a" in hmd, "convert HTML to docx")
    r = run(["docx_convert.py", fx["legacy"], w / "l.docx"], expect=1)
    c.ok("LibreOffice" in r.stderr, "convert .doc without LibreOffice fails cleanly")
    r = run(["docx_convert.py", src, w / "c.pdf"], expect=1)
    c.ok("--force" in r.stderr, "convert refuses to overwrite")
    j = jrun(["docx_convert.py", src, fx["template"], "--to", "md", "--out-dir", w / "batch"])
    c.ok(j["failed"] == 0 and (w / "batch" / "complex.md").exists() and (w / "batch" / "template.md").exists(), "batch conversion")


def big_markdown(sections: int = 30, paras: int = 8) -> str:
    import random

    rnd = random.Random(7)
    vocab = "the of and to in a is that for it as was with be by on not this are or from at which but have an they were her there been one all we their has would when if so no what can more out up into do only about other time than then its some could them these may first people like also any new over now".split()
    out = ["# Long report", ""]
    for k in range(1, sections + 1):
        out += [f"## Part {k}", ""]
        for _ in range(paras):
            out += [" ".join(rnd.choice(vocab) for _ in range(90)).capitalize() + ".", ""]
        out += [f"Part {k} ends here.", ""]
    return "\n".join(out)


def test_big_and_cache(c: Checks, fx: dict, w: Path) -> None:
    (w / "long.md").write_text(big_markdown(), encoding="utf-8")
    doc = w / "long.docx"
    run(["docx_create.py", w / "long.md", doc, "--no-toc-pages"])
    first = read_md(doc)
    c.ok("# long.docx: map" in first and "[1-10] Part 1 (" in first and "--section 'Part 1'" in first, f"a long document reads as a map with block ranges: {first[:400]!r}")
    again = read_md(doc)
    c.ok(again == first, "the cached read gives the same map")
    # The cache keeps the parsed model: the scripts time the parse itself (interpreter start-up is noise).
    cold = jrun(["docx_read.py", doc, "--blocks", "0", "--runs"])["timing"]
    hot = jrun(["docx_read.py", doc, "--blocks", "0", "--runs"])["timing"]
    c.ok(cold["cache"] == "miss" and hot["cache"] == "hit" and hot["seconds"] * 5 <= cold["seconds"], f"a cached read is at least 5x faster ({cold} → {hot})")
    part = read_md(doc, "--blocks", "11-20")
    c.ok(part.startswith("## Part 2") and "## Part 3" not in part, f"read --blocks after the map: {part[:60]!r}")
    c.ok(read_md(doc, "--section", "Part 7").startswith("## Part 7") and "## Part 8" not in read_md(doc, "--section", "Part 7"), "read --section in a long document")
    hits = jrun(["docx_read.py", doc, "--find", "part 12 ends here"])
    c.ok([h["index"] for h in hits["hits"]] == [120] and hits["hits"][0]["section"] == "Long report › Part 12", f"read --find in a long document: {hits['hits'][:1]}")
    full = read_md(doc, "--full")
    c.ok("Continue with: python3 scripts/docx_read.py" in full and full.startswith("# Long report"), "read --full pages the whole text")
    from docx import Document

    flat = Document()
    for i in range(260):
        flat.add_paragraph(f"Paragraph {i} " + "lorem ipsum dolor sit amet " * 12)
    flat.save(str(w / "flat.docx"))
    fm = read_md(w / "flat.docx")
    c.ok("no headings" in fm and "[0-99] Paragraph 0 " in fm and "[200-259] Paragraph 200" in fm and "--section" not in fm, f"a long document without headings maps as slices: {fm[:300]!r}")
    t = time.time()
    j = jrun(["docx_render.py", doc, "--out", w / "p1", "--pages", "1"])
    t_r1 = time.time() - t
    t = time.time()
    j2 = jrun(["docx_render.py", doc, "--out", w / "p2", "--pages", "last"])
    t_r2 = time.time() - t
    c.ok(j["page_count"] >= 12 and j2["page_count"] == j["page_count"] and j2["rendered"] == 1, f"render a long document ({j['page_count']} pages)")
    # The layout is what the cache keeps; the scripts' own timings leave out interpreter start-up and PNG encoding,
    # which a loaded machine makes noisy.
    c.ok(j2["layout_seconds"] * 5 <= j["layout_seconds"] or j2["layout_seconds"] <= 0.05, f"the second render reuses the cached layout (layout {j['layout_seconds']:.2f} s → {j2['layout_seconds']:.2f} s; whole call {t_r1:.2f} s → {t_r2:.2f} s)")
    c.ok(any("cache" in n for n in j2["notes"]), "render says the layout came from the cache")
    t = time.time()
    ex = jrun(["docx_info.py", doc, "--exact"])
    c.ok(ex["pages"]["rendered"] == j["page_count"] and time.time() - t < t_r1, "info --exact reuses the layout")
    # Pages by text and by block address, from the cached layout.
    f = jrun(["docx_render.py", doc, "--out", w / "pf", "--find", "Part 12 ends here"])
    b = jrun(["docx_render.py", doc, "--out", w / "pb", "--block", "120"])
    c.ok(f["rendered"] == 1 and b["pages"] == f["pages"] and 1 < f["pages"][0] < j["page_count"], f"render --find and --block agree: {f['pages']} {b['pages']} of {j['page_count']}")
    r = run(["docx_render.py", doc, "--out", w / "pn", "--find", "no such words anywhere"], expect=1)
    c.ok("not on any" in r.stderr, "render --find with no match fails cleanly")
    d = jrun(["docx_render.py", doc, "--out", w / "pd", "--dpi", "20"])
    c.ok(j["page_count"] > 20 and d["rendered"] == 20 and any("--pages 21-" in n for n in d["notes"]), f"render without --pages draws the first 20 pages of a long document ({d['rendered']} of {d['page_count']})")
    nc = jrun(["docx_render.py", fx["complex"], "--out", w / "p3", "--pages", "1"], env={"DESK_NO_CACHE": "1"})
    c.ok(nc["rendered"] == 1, "render works with the cache disabled")


STRICT = {
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main": "http://purl.oclc.org/ooxml/wordprocessingml/main",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships": "http://purl.oclc.org/ooxml/officeDocument/relationships",
    "http://schemas.openxmlformats.org/drawingml/2006/main": "http://purl.oclc.org/ooxml/drawingml/main",
    "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing": "http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing",
    "http://schemas.openxmlformats.org/drawingml/2006/picture": "http://purl.oclc.org/ooxml/drawingml/picture",
    "http://schemas.openxmlformats.org/officeDocument/2006/math": "http://purl.oclc.org/ooxml/officeDocument/math",
}


def to_strict(src: Path, dest: Path) -> Path:
    """An ISO Strict copy of a Transitional .docx, as Word writes with "Strict Open XML Document"."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename.endswith((".xml", ".rels")):
                text = data.decode("utf-8")
                for a, b in STRICT.items():
                    text = text.replace(a, b)
                text = text.replace("/relationships/extended-properties", "/relationships/extendedProperties")
                text = text.replace('w:line="259" w:lineRule="auto"', 'w:line="12.95pt" w:lineRule="auto"').replace('<w:sz w:val="22"/>', '<w:sz w:val="11pt"/>')
                data = text.encode("utf-8")
            zout.writestr(info, data)
    return dest


def test_strict_and_regressions(c: Checks, fx: dict, w: Path) -> None:
    strict = to_strict(fx["complex"], w / "strict.docx")
    r = run(["docx_read.py", strict])
    c.ok("ISO Strict" in r.stderr and "The term is {--30--}{++60++} days." in r.stdout and "| Widgets" not in r.stdout and '<td rowspan="2">Widgets</td>' in r.stdout, "read an ISO Strict document")
    j = jrun(["docx_render.py", strict, "--out", w / "sp", "--pages", "1"])
    c.ok(j["rendered"] == 1, "render an ISO Strict document with the built-in renderer")
    # Typst code generation: text starting with '(' right after formatted text, a coloured run inside a deletion.
    from docx import Document

    d = Document()
    para = d.add_paragraph("Smith, D. (2013). ")
    para.add_run("Climate Dynamics").italic = True
    para.add_run("(11–12), 3325–3338. [see] #tag $5 @x <y>")
    p2 = d.add_paragraph("Kept ")
    from docx.oxml import parse_xml

    p2._p.append(parse_xml('<w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z">'
                           '<w:r><w:rPr><w:color w:val="0000FF"/><w:u w:val="single"/></w:rPr><w:delText>blue gone</w:delText></w:r></w:del>'))
    d.save(str(w / "esc.docx"))
    j = jrun(["docx_render.py", w / "esc.docx", "--out", w / "ep", "--pdf", w / "esc.pdf"])
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(w / "esc.pdf"))
    try:
        text = " ".join(pdf[0].get_textpage().get_text_range().split())
    finally:
        pdf.close()
    c.ok("Climate Dynamics(11–12), 3325–3338. [see] #tag $5 @x <y>" in text and "blue gone" in text, f"built-in renderer escapes text after formatting and in revisions: {text[:160]!r}")


def test_lo_render_and_info(c: Checks, fx: dict, w: Path) -> None:
    j = jrun(["docx_render.py", fx["complex"], "--out", w / "pages"], lo=True)
    c.ok(j["engine"].startswith("LibreOffice") and j["page_count"] == 2 and j["rendered"] == 2, f"render with LibreOffice: {j['engine']} {j['page_count']}")
    c.ok(all(max(png_size(Path(p))) <= 1568 for p in j["images"]), "LibreOffice render sized for vision")
    ex = jrun(["docx_info.py", fx["complex"], "--exact"], lo=True)
    c.ok(ex["pages"].get("rendered") == 2 and "LibreOffice" in ex["pages"].get("engine", ""), f"info --exact with LibreOffice {ex['pages']}")


def test_lo_convert(c: Checks, fx: dict, w: Path) -> None:
    r = run(["docx_convert.py", fx["complex"], w / "c.pdf"], lo=True)
    c.ok("LibreOffice" in r.stdout and pdf_pages(w / "c.pdf") == 2, "convert to PDF with LibreOffice")
    run(["docx_convert.py", fx["complex"], w / "c.doc"], lo=True)
    r = run(["docx_read.py", w / "c.doc"], lo=True)
    c.ok("Service Agreement" in r.stdout and "LibreOffice" in r.stderr, "read a .doc through LibreOffice")
    r = run(["docx_convert.py", w / "c.doc", w / "c-back.docx"], lo=True)
    c.ok("LibreOffice" in r.stdout and "The term is" in read_md(w / "c-back.docx"), ".doc to .docx with LibreOffice")
    run(["docx_convert.py", fx["complex"], w / "c.odt"], lo=True)
    r = run(["docx_read.py", w / "c.odt"], lo=True)
    c.ok("Service Agreement" in r.stdout and "LibreOffice" in r.stderr, ".docx to .odt and back with LibreOffice")
    r = run(["docx_convert.py", fx["rtf"], w / "memo.docx"], lo=True)
    c.ok("LibreOffice" in r.stdout and "The budget is approved." in read_md(w / "memo.docx"), ".rtf to .docx with LibreOffice")
    f = jrun(["docx_render.py", fx["complex"], "--out", w / "pf", "--find", "Fees are due"], lo=True)
    c.ok(f["engine"].startswith("LibreOffice") and f["pages"] == [2], f"render --find on the LibreOffice layout: {f['pages']}")
    r = run(["docx_create.py", fx["md"], w / "toc.docx", "--toc"], lo=True, cwd=w)
    c.ok("filled by laying the document out (LibreOffice)" in r.stdout, f"create --toc fills page numbers from LibreOffice: {r.stdout.strip()[:200]}")
    # the TOC's page numbers are where LibreOffice puts the headings (read from the PDF's bookmarks)
    toc = w / "toc-lo.docx"
    r = run(["docx_create.py", "--markdown", TOC_MD, toc, "--title", "Paged", "--title-page", "--toc"], lo=True)
    c.ok(toc_pages(toc) == {"Alpha": 3, "Beta": 4, "Gamma": 5}, f"create --toc with LibreOffice: right page numbers {toc_pages(toc)}")
    # an RTF written by LibreOffice and a spec with a landscape annex both lay out as expected
    inv = w / "inv.docx"
    run(["docx_create.py", "--spec", json.dumps(LANDSCAPE_SPEC), inv], lo=True)
    j = jrun(["docx_render.py", inv, "--out", w / "inv"], lo=True)
    c.ok(j["page_count"] == 2 and j["engine"].startswith("LibreOffice"), f"spec with a landscape annex renders on 2 pages with LibreOffice ({j['page_count']})")


# ── regressions from the acceptance review ─────────────────────────────

TOC_MD = "# Alpha\n\nFirst chapter text about wholesale and retail.\n\n\\pagebreak\n\n# Beta\n\nSecond chapter; the word Alpha appears here too.\n\n\\pagebreak\n\n# Gamma\n\nThird chapter.\n"

LANDSCAPE_SPEC = {
    "properties": {"title": "Invoice 7", "author": "Billing Robot"},
    "page": {"size": "A4", "margins": "2cm"},
    "footer": "Northwind Traders AS · Karl Johans gate 1, Oslo · VAT NO 912 345 678||Page {page} of {pages}",
    "blocks": [
        {"type": "title", "text": "Invoice", "subtitle": "INV-7"},
        {"type": "table", "header": ["Description", "Qty", "Unit price", "Amount"], "rows": [["Consulting", 2, "1,200.00", "2,400.00"], ["Travel", 1, "318.50", "318.50"]],
         "shading": {"header": "1F3864"}},
        {"type": "paragraph", "runs": ["Pay within 30 days.", {"text": "", "footnote": "Late payments bear interest."}]},
        {"type": "section", "orientation": "landscape"},
        {"type": "heading", "text": "Annex", "level": 1},
        {"type": "paragraph", "text": "Time sheet."},
    ],
}


def toc_pages(path: Path) -> dict[str, int]:
    import re

    md = read_md(path)
    return {m.group(1): int(m.group(2)) for m in re.finditer(r"^\s*- (\w+) … (\d+)$", md, re.M)}


def docx_xml(path: Path, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


def pdf_text(p: Path) -> str:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(p))
    try:
        return " ".join(" ".join(doc[i].get_textpage().get_text_range().split()) for i in range(len(doc)))
    finally:
        doc.close()


def test_regress_create(c: Checks, fx: dict, w: Path) -> None:
    import re

    # TOC page numbers are the pages the headings are on (title page 1, contents 2, then one chapter per page)
    toc = w / "toc.docx"
    r = run(["docx_create.py", "--markdown", TOC_MD, toc, "--title", "Paged", "--title-page", "--toc"])
    c.ok(toc_pages(toc) == {"Alpha": 3, "Beta": 4, "Gamma": 5} and "3 of 3 filled" in r.stdout, f"create --toc: right page numbers {toc_pages(toc)} ({r.stdout.strip()[:160]})")
    app = docx_xml(toc, "docProps/app.xml")
    c.ok("Microsoft Word 12" not in app and "<Pages>5</Pages>" in app and "Desk word-documents" in app, f"create: app.xml says who made it and its real page count: {app[-300:]}")
    j = jrun(["docx_convert.py", toc, w / "toc.pdf"])
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(w / "toc.pdf"))
    try:
        marks = [(b.get_title(), b.get_dest().get_index() + 1) for b in pdf.get_toc()]
    finally:
        pdf.close()
    c.ok(marks == [("Alpha", 3), ("Beta", 4), ("Gamma", 5)], f"built-in PDF has heading bookmarks: {marks}")
    # money signs are text, not math; real math still works; template tags survive Markdown
    md = ("| Metric | Q3 |\n|:--|--:|\n| Revenue ($M) | 42.3 |\n| Gross margin | 40.8% |\n| Operating profit ($M) | 4.9 |\n\n"
          "Costs were $5 to $10; Einstein wrote $E=mc^2$.\n\n| # | Item | Price |\n|---|---|---:|\n"
          "| {{#each items}}{{@number}} | {{description}} | {{price | currency:\"$\"}}{{/each}} |\n\nMail {{client.email}}.\n")
    money = w / "money.docx"
    run(["docx_create.py", "--markdown", md, money])
    out = read_md(money)
    for want in ("| Revenue ($M) | 42.3 |", "| Gross margin | 40.8% |", "| Operating profit ($M) | 4.9 |", "Costs were $5 to $10", "$E=mc^2$",
                 '| {{#each items}}{{@number}} | {{description}} | {{price \\| currency:"$"}}{{/each}} |', "Mail {{client.email}}."):
        c.ok(want in out, f"create: literal $ and template tags: {want!r}")
    c.ok("mailto:" not in out, "create: {{@number}} after a loop tag is not an e-mail link")
    filled = w / "money-filled.docx"
    run(["docx_template.py", money, filled, "--data", json.dumps({"items": [{"description": "Tea", "price": 3.5}, {"description": "Cake", "price": 4}], "client": {"email": "a@b.c"}})])
    fmd = read_md(filled)
    c.ok("| 1 | Tea | $3.50 |" in fmd and "| 2 | Cake | $4.00 |" in fmd, f"template filled from a Markdown-made table: {fmd[-220:]!r}")
    # tables fill the text width with content-based columns; the paragraph after a table has room above it
    doc_xml = docx_xml(money)
    widths = [int(x) for x in re.findall(r'<w:gridCol w:w="(\d+)"/>', doc_xml.split("</w:tblGrid>")[0])]
    tblw = re.search(r'<w:tblW w:type="dxa" w:w="(\d+)"/>', doc_xml)
    c.ok(tblw and abs(sum(widths) - int(tblw.group(1))) <= 2 and int(tblw.group(1)) > 9000 and widths[0] > widths[1], f"create: a table spans the text width with content-based columns {widths}")
    after = doc_xml.split("</w:tbl>", 1)[1][:400]
    c.ok(re.search(r'w:before="(1[6-9]\d|[2-9]\d\d)"', after) is not None, "create: space between a table and the next paragraph")
    grid2 = [int(x) for x in re.findall(r'<w:gridCol w:w="(\d+)"/>', doc_xml.split("</w:tblGrid>")[1])]
    c.ok(len(grid2) == 3 and grid2[0] < grid2[1], f"create: a template table is sized by its placeholders, not its loop tags {grid2}")
    branches = w / "branches.docx"
    run(["docx_create.py", "--markdown", "| A | B |\n|---|---|\n| 1 | 2 |\n\n{{#if paid}}\n\nPaid.\n\n{{else}}\n\nDue.\n\n{{/if}}\n\nAfter.\n", branches])
    bx = docx_xml(branches).split("</w:tbl>", 1)[1]
    spaced = [t for t in ("Paid.", "Due.", "After.") if re.search(r'w:before="180"[^<]*/>(?:(?!</w:p>).)*' + re.escape(t), bx)]
    c.ok(spaced == ["Paid.", "Due."], f"create: the first paragraph of each template branch after a table gets the space ({spaced})")
    # JSON spec: no placeholder footnote, header text readable on a dark fill, numbers right-aligned, author not shown,
    # the landscape section's footer has its own tab stops at its own width, 'left||right' is one tab
    inv = w / "inv.docx"
    run(["docx_create.py", "--spec", json.dumps(LANDSCAPE_SPEC), inv])
    notes = docx_xml(inv, "word/footnotes.xml")
    c.ok("Footnote Text" not in notes and notes.count("<w:footnote ") - notes.count('w:type="') == 1 and "Late payments" in notes, "spec: only the real footnote in footnotes.xml")
    c.ok(jrun(["docx_info.py", inv])["footnotes"] == 1, "spec: info counts 1 footnote")
    b = jrun(["docx_read.py", inv, "--blocks", "2", "--runs"])
    cells = b["blocks"][0]["cells"]
    c.ok(cells[0][0]["text"] == "Description" and cells[1][2]["text"] == "1,200.00", "spec table read back")
    x = docx_xml(inv)
    head = x[x.index("<w:tbl>"): x.index("</w:tr>")]
    c.ok(head.count('w:color w:val="FFFFFF"') == 4, "spec: header text is white on a dark header fill")
    row1 = x[x.index("</w:tr>"): x.index("</w:tr>", x.index("</w:tr>") + 1)]
    c.ok(row1.count('<w:jc w:val="right"/>') == 3, "spec: numeric columns are right-aligned")
    c.ok("Billing Robot" not in read_md(inv), "spec: properties.author is metadata, not a visible line")
    footers = [n for n in zipfile.ZipFile(inv).namelist() if n.startswith("word/footer")]
    tabs = sorted(int(t) for f in footers for t in re.findall(r'<w:tab w:val="right" w:pos="(\d+)"/>', docx_xml(inv, f)))
    c.ok(len(tabs) == 2 and tabs[1] - tabs[0] > 4000, f"spec: portrait and landscape footers each have a right tab at their own width {tabs}")
    c.ok(all(docx_xml(inv, f).count("<w:tab/>") <= 1 for f in footers), "spec: a footer with an empty centre part has one tab")
    j = jrun(["docx_render.py", inv, "--out", w / "inv"])
    c.ok(j["page_count"] == 2, f"spec with a landscape annex renders on 2 pages ({j['page_count']})")
    # docx -> md -> docx keeps one caption per figure
    fig = w / "fig.docx"
    run(["docx_create.py", "--markdown", f"![Figure 1. Revenue]({fx['red'].as_posix()})\n\nText.\n", fig])
    run(["docx_convert.py", fig, w / "fig.md"])
    run(["docx_convert.py", w / "fig.md", w / "fig-rt.docx"], cwd=w)
    c.ok(read_md(w / "fig-rt.docx").count("Figure 1. Revenue") == 2, "md round trip: the caption is not duplicated (alt text + one caption)")


def test_regress_safety(c: Checks, fx: dict, w: Path) -> None:
    import time as _t

    # a zip bomb is refused at once
    bomb = w / "bomb.docx"
    with zipfile.ZipFile(fx["complex"]) as zin, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "word/document.xml":
                data = data.replace(b"<w:body>", b"<w:body>" + b"<w:p/>" * 4_000_000, 1)
            zout.writestr(info.filename, data)
    for script in ("docx_read.py", "docx_info.py", "docx_render.py"):
        t = _t.time()
        r = run([script, bomb] + (["--out", w / "b"] if script == "docx_render.py" else []), expect=1)
        c.ok("zip bomb" in r.stderr and _t.time() - t < 5, f"{script} refuses a zip bomb fast: {r.stderr.strip()[:120]}")
    # an XML document type declaration (entity tricks) is refused the same way everywhere
    dtd = w / "dtd.docx"
    with zipfile.ZipFile(fx["complex"]) as zin, zipfile.ZipFile(dtd, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "word/document.xml":
                data = data.replace(b"?>", b'?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]>', 1)
            zout.writestr(info, data)
    for args in (["docx_read.py", dtd], ["docx_convert.py", dtd, w / "dtd.md"], ["docx_convert.py", dtd, w / "dtd.html"], ["docx_info.py", dtd]):
        r = run(args, expect=1)
        c.ok("document type declaration" in r.stderr, f"{args[0]} refuses a DTD: {r.stderr.strip()[:100]}")
    # a missing picture part: read, info and render work (Word repairs such files too)
    miss = w / "missing.docx"
    with zipfile.ZipFile(fx["complex"]) as zin, zipfile.ZipFile(miss, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename != "word/media/image1.png":
                zout.writestr(info, zin.read(info))
    r = run(["docx_read.py", miss])
    c.ok("missing" in r.stderr and "Service Agreement" in r.stdout, "read a package with a missing picture part")
    c.ok(jrun(["docx_info.py", miss])["blocks"] == 18, "info on a package with a missing part")
    c.ok(jrun(["docx_render.py", miss, "--out", w / "m", "--pages", "1"])["rendered"] == 1, "render a package with a missing part")
    run(["docx_edit.py", miss, w / "fixed.docx", "--ops", '[{"op":"replace","find":"ACME","replace":"Acme"}]'])
    c.ok("Acme Ltd" in read_md(w / "fixed.docx"), "edit a package with a missing part (the saved copy is whole)")
    # a catastrophic regular expression stops with an error instead of hanging
    redos = w / "redos.docx"
    run(["docx_create.py", "--markdown", "a" * 60 + "!", redos])
    t = _t.time()
    r = run(["docx_read.py", redos, "--grep", "(a+)+$"])
    c.ok("0 matches" in r.stdout and _t.time() - t < 5, "--grep (a+)+$ (hangs Python's re) answers at once")
    t = _t.time()
    r = run(["docx_read.py", redos, "--grep", "(a|aa)+$"], expect=1)
    c.ok("catastrophic" in r.stderr and _t.time() - t < 15, f"--grep with a catastrophic pattern stops in time: {r.stderr.strip()[:100]}")
    r = run(["docx_edit.py", redos, w / "x.docx", "--ops", json.dumps([{"op": "replace", "regex": "(a|aa)+$", "replace": "b"}])], expect=1)
    c.ok("catastrophic" in r.stderr, "edit with a catastrophic regex fails cleanly")
    # RTF without LibreOffice: characters after \uN escapes survive (pandoc's reader used to drop them)
    rtf = w / "u.rtf"
    rtf.write_text("{\\rtf1\\ansi{\\pard Northwind\\u8217's history. \\u8220\"The hub\\u8221\" \\uc2\\u8364XXcosts.\\par}}", encoding="latin-1")
    own = {"DESK_FILE_CACHE": str(w / "cache")}  # its own cache: test_read_other_formats counts the shared one's entries
    txt = run(["docx_read.py", rtf, "--format", "text"], env=own).stdout
    c.ok("Northwind’s history. “\"The hub”" in txt or "Northwind’s history. “The hub”" in txt, f"RTF: characters after unicode escapes survive: {txt.strip()!r}")
    c.ok("€costs" in txt, f"RTF: \\uc2 fallbacks are skipped: {txt.strip()!r}")
    # docx -> rtf keeps the subtitle, header and footer; reading it back keeps apostrophes
    doc = w / "sub.docx"
    run(["docx_create.py", "--markdown", "It's Northwind's report.\n", doc, "--title", "Main title", "--subtitle", "The subtitle", "--header", "Head text", "--page-numbers"])
    run(["docx_convert.py", doc, w / "sub.rtf"])
    raw = (w / "sub.rtf").read_text(encoding="latin-1")
    c.ok("The subtitle" in raw and "{\\header" in raw and "Head text" in raw and "fldinst PAGE" in raw, "docx -> RTF keeps subtitle, header and page field")
    c.ok("Northwind’s report" in run(["docx_read.py", w / "sub.rtf"], env=own).stdout, "docx -> RTF -> read keeps apostrophes")


def test_regress_jail(c: Checks, fx: dict, w: Path) -> None:
    """Untrusted inputs: nothing from outside a document's folder reaches the output (includes, pictures, file: URIs,
    symlinks), pictures are named by a whitelist, batch names never collide, and LibreOffice runs one at a time."""
    import base64
    import re
    import threading

    if str(HERE) not in sys.path:
        sys.path.insert(0, str(HERE))

    doc, far = w / "doc", w / "far"
    (doc / "sub").mkdir(parents=True)
    far.mkdir()
    marker = "SECRETMARK-3W"
    secret = far / "secret.txt"
    secret.write_text(marker + "\n", encoding="utf-8")
    (far / "outside.png").write_bytes(png_bytes((10, 200, 10), (50, 50)))
    (doc / "pic.png").write_bytes(png_bytes((200, 10, 10), (40, 40)))
    link = ""
    try:
        (doc / "link.png").symlink_to(far / "outside.png")
        link = "\n\n![l](link.png)"
    except OSError:
        pass
    files = {
        "m.md": f"# M\n\nMDPART\n\n![x]({secret})\n\n![h](/etc/hosts)\n\n![y](../far/outside.png)\n\n![f](file://{secret.as_posix()})\n\n![ok](pic.png){link}\n",
        "c.rst": f"Title\n=====\n\n.. include:: {secret}\n\n.. include:: part.rst\n\n.. include:: /etc/hosts\n",
        "part.rst": "RSTPART\n\n.. include:: ../far/secret.txt\n",
        "o.org": f'* H\n#+INCLUDE: "{secret}"\n#+INCLUDE: "sub/p.org"\n',
        "sub/p.org": 'ORGPART\n#+INCLUDE: "../../far/secret.txt"\n',
        "l.tex": "\\documentclass{article}\n\\begin{document}\n\\include{" + secret.with_suffix("").as_posix() + "}\n\\input{sub/ch}\n\\def\\p{../far/secret.txt}\\input{\\p}\n\\end{document}\n",
        "sub/ch.tex": "TEXPART\n",
        "h.html": f'<h1>HTMLPART</h1><img src="file://{secret.as_posix()}"><img src="../far/outside.png"><img src="pic.png">',
    }
    for name, text in files.items():
        (doc / name).write_text(text, encoding="utf-8")
    bad = [marker.encode(), base64.b64encode((marker + "\n").encode()), (far / "outside.png").read_bytes()[:120]]

    def check_docx(path: Path, what: str, part: str, pictures: int | None) -> None:
        with zipfile.ZipFile(path) as z:
            data = b"".join(z.read(n) for n in z.namelist())
            media = [n for n in z.namelist() if n.startswith("word/media/")]
        c.ok(not any(x in data for x in bad), f"{what}: nothing from outside the folder")
        c.ok(part in read_md(path), f"{what}: its own text and includes kept")
        if pictures is not None:
            c.ok(len(media) == pictures, f"{what}: only the picture inside the folder ({media})")

    for name, part in (("m.md", "MDPART"), ("c.rst", "RSTPART"), ("o.org", "ORGPART"), ("l.tex", "TEXPART"), ("h.html", "HTMLPART")):
        out = w / f"{Path(name).stem}-{Path(name).suffix[1:]}.docx"
        r = run(["docx_convert.py", doc / name, out])
        check_docx(out, f"convert {name}", part, 1 if name in ("m.md", "h.html") else None)
        if name == "m.md":
            c.ok("left out (outside the document's folder)" in r.stderr and "left out (a URL pandoc must not open)" in r.stderr, f"refused pictures are reported: {r.stderr[-300:]}")
    run(["docx_create.py", doc / "m.md", w / "created.docx"])
    check_docx(w / "created.docx", "create from m.md", "MDPART", None)
    out = run(["docx_read.py", doc / "c.rst"]).stdout
    c.ok(marker not in out and "RSTPART" in out, "read an .rst: nothing from outside its folder")
    remote = w / "remote.md"
    remote.write_text("# R\n\n![far away](https://example.invalid/pic.png)\n\nREMOTEWORD\n", encoding="utf-8")
    r = run(["docx_convert.py", remote, w / "remote.docx"], env={"DESK_OFFLINE": "1"})
    c.ok("remote image not downloaded" in r.stderr and "REMOTEWORD" in read_md(w / "remote.docx"), "a remote picture offline becomes a link (pandoc never fetches it)")

    # mammoth --media-dir: file names never come from the document's content types
    for i, ctype in enumerate(("image/\\..\\..\\..\\Users\\Public\\evil.hta", "image/png:x", "image/" + "x" * 300)):
        crafted = w / f"ctype{i}.docx"
        with zipfile.ZipFile(fx["complex"]) as zin, zipfile.ZipFile(crafted, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info)
                if info.filename == "[Content_Types].xml":
                    data = re.sub(rb'(Extension="png"\s+ContentType=")[^"]*"', lambda m: m.group(1) + ctype.encode() + b'"', data)
                zout.writestr(info, data)
        media = w / f"media{i}"
        run(["docx_convert.py", crafted, w / f"ctype{i}.html", "--media-dir", media])
        names = sorted(p.name for p in media.iterdir()) if media.exists() else []
        c.ok(names == ["image1.bin"] and not list(w.rglob("*.hta")), f"--media-dir names pictures from a whitelist ({names})")
    from _reader import media_name

    names = {"word/media/C:x.png": "C_x.png", "word/media/img.png:ads": "img.png_ads", "word/media/con.png": "_con.png", "word/media/nul.tar.gz": "_nul.tar.gz", "word/media/..\\..\\x.png": "x.png", "word/media/a.png.": "a.png", "word/media/image1.png": "image1.png"}
    got = {k: media_name(k) for k in names}
    c.ok(got == names, f"docx pictures get names safe on Windows: {got}")

    # batch: two inputs with one name get two outputs
    for sub, name in (("ba", "notes.docx"), ("bb", "Notes.docx")):
        (w / sub).mkdir()
        shutil.copy(fx["complex"] if sub == "ba" else fx["template"], w / sub / name)
    j = jrun(["docx_convert.py", w / "ba" / "notes.docx", w / "bb" / "Notes.docx", "--to", "txt", "--out-dir", w / "same"])
    outs = sorted(p.name for p in (w / "same").glob("*.txt"))
    c.ok(j["failed"] == 0 and len(outs) == 2 and any(n.endswith("-2.txt") for n in outs), f"batch: same-name inputs get name and name-2 ({outs})")

    # an .odt or .epub zip bomb is refused before pandoc or LibreOffice unpacks it
    for ext in (".odt", ".epub"):
        bomb = w / f"bomb{ext}"
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.text" if ext == ".odt" else "application/epub+zip", compress_type=zipfile.ZIP_STORED)
            z.writestr("content.xml", b"<p/>" * 6_000_000)
        r = run(["docx_convert.py", bomb, w / f"bomb{ext}.docx"], expect=1)
        c.ok("zip bomb" in r.stderr, f"convert refuses a {ext} zip bomb: {r.stderr.strip()[:120]}")

    # LibreOffice steps run one at a time, even when batch jobs run on threads
    import _render
    import docx_convert

    live, peak = [0], [0]
    lock = threading.Lock()

    def fake_convert(src: Path, fmt: str, **_: object) -> Path:
        with lock:
            live[0] += 1
            peak[0] = max(peak[0], live[0])
        import time as _t

        _t.sleep(0.15)
        with lock:
            live[0] -= 1
        d = Path(tempfile.mkdtemp(dir=w))
        (d / "out.bin").write_bytes(b"x")
        return d / "out.bin"

    saved = _render.office_convert
    _render.office_convert = fake_convert
    try:
        ts = [threading.Thread(target=docx_convert._lo, args=(fx["complex"], w / f"lo{k}.bin", "pdf")) for k in range(4)]
        for t in ts:
            t.start()
        for t in ts:
            t.join()
    finally:
        _render.office_convert = saved
    c.ok(peak[0] == 1 and all((w / f"lo{k}.bin").exists() for k in range(4)), f"one LibreOffice at a time in a batch (peak {peak[0]})")


def test_regress_edit_compare(c: Checks, fx: dict, w: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    # a cell with line breaks: CSV and JSON give the lines, Markdown uses <br>
    d = Document()
    d.add_paragraph("Figures")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "2006"
    t.cell(1, 0).text = "Gross*"
    t.cell(1, 1).text = "17,542\n32,000\n3,508"
    body = d.add_paragraph()
    run_ = body.add_run("Patrick is single and earns €8.65 an hour.")
    run_.font.name, run_.font.size = "Times New Roman", Pt(11)
    d.add_paragraph("Note: to be removed.")
    d.add_paragraph("Closing words.")
    base = w / "base.docx"
    d.save(str(base))
    csv = run(["docx_read.py", base, "--blocks", "1", "--format", "csv"]).stdout
    c.ok('"17,542\n32,000\n3,508"' in csv and "Gross*" in csv and "\\ " not in csv, f"CSV: in-cell line breaks are new lines: {csv!r}")
    cell = jrun(["docx_read.py", base, "--blocks", "1"])["blocks"][0]["cells"][1][1]
    c.ok(cell["text"] == "17,542\n32,000\n3,508", f"JSON: cell text is plain text: {cell}")
    c.ok("17,542<br>32,000<br>3,508" in read_md(base, "--blocks", "1"), "Markdown: in-cell line breaks as <br>")
    # tracked edits: Markdown inserts take the neighbour's look; labels say what is (not) tracked; reject undoes the footer
    ops = [{"op": "insert_after", "index": 2, "markdown": "*Update:* the wage rises to **€8.75**."}, {"op": "set_footer", "text": "New footer||Page {page}"},
           {"op": "properties", "title": "T"}, {"op": "delete", "index": 3}, {"op": "move", "index": 0, "to": 4}]
    v2 = w / "v2.docx"
    out = run(["docx_edit.py", base, v2, "--ops", json.dumps(ops), "--track"]).stdout
    lines = out.splitlines()
    c.ok(all("(tracked)" in lines[k] for k in (0, 1, 3, 4)) and "not tracked" in lines[2], f"--track labels every op: {out!r}")
    blk = jrun(["docx_read.py", v2, "--blocks", "3", "--runs"])["blocks"][0]
    fonts = {(r.get("font"), r.get("size")) for r in blk["runs"]}
    c.ok(fonts == {("Times New Roman", 11.0)} and any(r.get("bold") for r in blk["runs"]), f"inserted Markdown looks like its neighbour: {fonts}")
    rej = w / "rej.docx"
    run(["docx_edit.py", v2, rej, "--ops", '[{"op":"reject_changes"}]'])
    cmp_ = run(["docx_compare.py", base, rej]).stdout
    c.ok("same text" in cmp_, f"reject all restores text, footer and order: {cmp_[:300]!r}")
    rep = run(["docx_compare.py", base, v2]).stdout
    c.ok("{++++}" not in rep and "deleted paragraph: {--Note: to be removed.--}" in rep, f"compare: a tracked deletion is not an empty added paragraph: {rep!r}")
    # a table whose rows changed is compared row by row; accepting or rejecting its redline leaves no empty table
    def invoice(path: Path, rows: list[tuple[str, str]]) -> Path:
        dd = Document()
        dd.add_paragraph("Invoice")
        tt = dd.add_table(rows=1, cols=2)
        tt.cell(0, 0).text, tt.cell(0, 1).text = "Item", "Amount"
        for a, b in rows:
            cells = tt.add_row().cells
            cells[0].text, cells[1].text = a, b
        dd.add_paragraph("Thanks.")
        dd.save(str(path))
        return path

    i1 = invoice(w / "i1.docx", [("Consulting", "2,400.00"), ("Travel", "318.50"), ("Writing", "1,950.00"), ("Total", "4,668.50")])
    i2 = invoice(w / "i2.docx", [("Licence", "9,900.00"), ("Total", "9,900.00")])
    rep = run(["docx_compare.py", i1, i2, "--redline", w / "red.docx"]).stdout
    c.ok("table changed" in rep and "row 4 → 2: col 1: {--4,668.50--}{++9,900.00++}" in rep and "added table" not in rep, f"compare: rows of a changed table: {rep!r}")
    for mode, other in (("accept_changes", i2), ("reject_changes", i1)):
        o = w / f"{mode}.docx"
        run(["docx_edit.py", w / "red.docx", o, "--ops", json.dumps([{"op": mode}])])
        x = docx_xml(o)
        c.ok(x.count("<w:tbl>") == 1 and all("<w:tr" in tb for tb in x.split("<w:tbl>")[1:]), f"{mode}: no table without rows")
        c.ok("same text" in run(["docx_compare.py", other, o]).stdout, f"{mode} of the redline gives the {'new' if other == i2 else 'old'} version")
    # a document without heading styles gets a map from its heading-like lines, and --section works on them
    flat = Document()
    for k in range(1, 9):
        flat.add_paragraph().add_run(f"ARTICLE {k} - TERMS {k}").bold = True
        for _ in range(4):
            flat.add_paragraph(f"Clause text of article {k}. " * 12)
    flat.save(str(w / "flat.docx"))
    m = run(["docx_read.py", w / "flat.docx", "--max-chars", "3000"]).stdout
    c.ok("no heading styles" in m and "[0-4] ARTICLE 1 - TERMS 1" in m and "[35-39] ARTICLE 8 - TERMS 8" in m, f"map of a document without heading styles: {m[:400]!r}")
    sec = run(["docx_read.py", w / "flat.docx", "--section", "ARTICLE 2", "--max-chars", "0"]).stdout
    c.ok("ARTICLE 2 - TERMS 2" in sec and "article 2." in sec and "ARTICLE 3" not in sec, "--section on a heading-like line")
    # the built-in renderer's NUMPAGES counts pages, whatever number the page numbering restarts at
    from docx.oxml import parse_xml

    rd = Document()
    rd.add_paragraph("One page.")
    sp = rd.sections[0]._sectPr
    sp.append(parse_xml('<w:pgNumType xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:start="24"/>'))
    rd.save(str(w / "restart.docx"))
    run(["docx_edit.py", w / "restart.docx", w / "restart2.docx", "--ops", '[{"op":"set_footer","text":"Page {page} of {pages}"}]'])
    run(["docx_convert.py", w / "restart2.docx", w / "restart.pdf"])
    c.ok("Page 24 of 1" in pdf_text(w / "restart.pdf"), f"built-in renderer: NUMPAGES is the page count: {pdf_text(w / 'restart.pdf')[-60:]!r}")


if __name__ == "__main__":
    sys.exit(main())
