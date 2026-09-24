"""Adds Markdown-lite blocks to a python-docx Document (shared by docx_create and docx_edit)."""
from __future__ import annotations

from docx.document import Document as DocumentType
from docx.shared import Pt

from _mdlite import Block, spans

CODE_FONT = "Menlo"


def style_or_normal(doc: DocumentType, name: str) -> str:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return "Normal"


def add_runs(paragraph, text: str) -> None:
    for chunk, fmt in spans(text):
        run = paragraph.add_run(chunk)
        run.bold = fmt.get("bold") or None
        run.italic = fmt.get("italic") or None
        if fmt.get("code"):
            run.font.name = CODE_FONT
            run.font.size = Pt(9.5)


def add_blocks(doc: DocumentType, blocks: list[Block]) -> list:
    """Appends the blocks to the end of the document; returns the new body elements in order."""
    added = []
    for b in blocks:
        if b.kind == "heading":
            p = doc.add_heading("", level=min(b.level, 9))
            add_runs(p, b.text)
            added.append(p._p)
        elif b.kind in ("bullet", "number"):
            base = "List Bullet" if b.kind == "bullet" else "List Number"
            name = base if b.level == 0 else f"{base} {min(b.level + 1, 3)}"
            p = doc.add_paragraph(style=style_or_normal(doc, name))
            add_runs(p, b.text)
            added.append(p._p)
        elif b.kind == "quote":
            p = doc.add_paragraph(style=style_or_normal(doc, "Quote"))
            add_runs(p, b.text)
            added.append(p._p)
        elif b.kind == "code":
            for line in b.text.split("\n") or [""]:
                p = doc.add_paragraph(style=style_or_normal(doc, "No Spacing"))
                run = p.add_run(line)
                run.font.name = CODE_FONT
                run.font.size = Pt(9.5)
                added.append(p._p)
        elif b.kind == "table":
            cols = max(len(r) for r in b.rows)
            table = doc.add_table(rows=len(b.rows), cols=cols)
            table.style = style_or_normal(doc, "Table Grid") if style_or_normal(doc, "Table Grid") != "Normal" else None
            for r, row in enumerate(b.rows):
                for c in range(cols):
                    cell = table.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            added.append(table._tbl)
        elif b.kind == "pagebreak":
            p = doc.add_page_break()
            added.append(p._p)
        else:
            p = doc.add_paragraph()
            add_runs(p, b.text)
            added.append(p._p)
    return added
