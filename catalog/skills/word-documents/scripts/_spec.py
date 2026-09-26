"""docx_create's JSON spec: exact control over blocks, runs, lists, tables, images and sections.

See references/json-spec.md. Every block type is validated with a clear error naming the block."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from _common import SkillError, UsageError

BLOCK_TYPES = ("heading", "paragraph", "bullets", "numbered", "table", "image", "page_break", "toc", "markdown", "quote", "code", "spacer", "section", "title")
RUN_KEYS = {"text", "bold", "italic", "underline", "strike", "color", "highlight", "size", "font", "superscript", "subscript", "link", "code", "footnote", "style", "caps", "small_caps", "shading", "break"}


def build_from_spec(spec: Any, reference: Path, dest: Path, *, base_dir: Path, metadata: dict[str, Any] | None = None) -> list[str]:
    """Builds a .docx at dest from a spec on top of the reference document's styles. Returns warnings."""
    from _docx import SECTPR, Doc

    if isinstance(spec, list):
        spec = {"blocks": spec}
    if not isinstance(spec, dict) or not isinstance(spec.get("blocks"), list):
        raise UsageError('the spec must be {"blocks": [...]} (or a list of blocks); see references/json-spec.md')
    doc = Doc(reference)
    body = doc.body
    for child in list(body):
        if child.tag != SECTPR:
            body.remove(child)
    clear_notes(doc)
    b = Builder(doc, base_dir)
    cli = {k: v for k, v in (metadata or {}).items() if v}
    b.cli = cli
    meta = dict(cli)
    meta.update({k: v for k, v in (spec.get("properties") or {}).items() if v is not None})
    # A title given on the command line is shown (as with Markdown); spec "properties" are metadata only.
    if cli.get("title") and not any(isinstance(x, dict) and x.get("type") == "title" for x in spec["blocks"]):
        b.title_block(cli)
    page = spec.get("page") or {}
    if page:
        # Before the blocks, so that "section" blocks start from this page setup (and change only what they say).
        from _markdown import page_setup

        page_setup(doc, size=page.get("size"), orientation=page.get("orientation"), margins=page.get("margins"))
    for i, blk in enumerate(spec["blocks"]):
        if not isinstance(blk, dict) or "type" not in blk:
            raise UsageError(f"block {i} needs a \"type\" ({', '.join(BLOCK_TYPES)})")
        try:
            b.block(blk, meta)
        except (UsageError, SkillError) as e:
            raise type(e)(f"block {i} ({blk.get('type')}): {e}") from e
    from _markdown import set_header_footer

    for kind in ("header", "footer"):
        v = spec.get(kind)
        if v:
            text = v if isinstance(v, str) else v.get("text", "")
            align = "right" if kind == "header" else "center"
            if isinstance(v, dict):
                align = v.get("align", align)
            set_header_footer(doc, kind, text, align=align)
    cp = doc.docx.core_properties
    for k in ("title", "subject", "author", "keywords", "category", "comments"):
        if meta.get(k):
            setattr(cp, k, str(meta[k]))
    doc.save(dest, force=True)
    return b.warnings


def clear_notes(doc: Any) -> None:
    """Empties the notes and comments a reference or template carries (pandoc's reference.docx has a sample
    footnote, which LibreOffice would show for the document's first note): only the separators stay."""
    from _docx import qn

    for root, tag in ((doc.footnotes, "w:footnote"), (doc.endnotes, "w:endnote")):
        if root is None:
            continue
        for n in root.findall(qn(tag)):
            if (n.get(qn("w:type")) or "normal") == "normal":
                root.remove(n)
    if doc.comments is not None:
        for c in doc.comments.findall(qn("w:comment")):
            doc.comments.remove(c)


def is_dark(hex6: str) -> bool:
    """True for fills dark enough that text on them must be white."""
    h = str(hex6).lstrip("#")
    if len(h) != 6:
        return False
    try:
        r, g, b = (int(h[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    return 0.299 * r + 0.587 * g + 0.114 * b < 140


NUMERIC = __import__("re").compile(r"^[\s(]*[-+\u2212]?\s?(?:[$€£¥]|[A-Z]{3}\s)?\s?[-+\u2212]?\d[\d,.\s\u00a0\u202f']*(?:%|pt|[kKmMbB]n?)?\s?(?:[$€£¥]|\s[A-Z]{3})?[\s)]*$")


def looks_numeric(v: Any) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    if isinstance(v, dict):
        v = v.get("text", "")
    return isinstance(v, str) and bool(NUMERIC.match(v)) and any(ch.isdigit() for ch in v)


class Builder:
    cli: dict[str, Any] = {}

    def __init__(self, doc: Any, base_dir: Path) -> None:
        self.doc = doc
        self.d = doc.docx
        self.base_dir = base_dir
        self.warnings: list[str] = []
        self.sections = 0
        self._abstract_ids: dict[str, str] = {}

    # ── helpers ─────────────────────────────────────────────────────────
    def style(self, name: str | None, kind: str = "paragraph") -> str | None:
        if not name:
            return None
        sid = self.doc.styles.find(name, kind)
        if not sid:
            raise SkillError(f"style '{name}' does not exist here (with --template, use one of its styles)")
        return sid

    def para(self, style: str | None = None) -> Any:
        p = self.d.add_paragraph()
        sid = self.style(style) if style else None
        if sid:
            from _docx import new_el, ppr_of

            ppr = ppr_of(p._p)
            ps = new_el("w:pStyle", val=sid)
            ppr.insert(0, ps)
        return p

    def fmt_para(self, p: Any, blk: dict[str, Any]) -> None:
        from _docx import ppr_of
        from _markdown import set_paragraph_format

        keys = {k: v for k, v in blk.items() if k in ("align", "space_before", "space_after", "line_spacing", "indent_left", "indent_right", "first_line_indent", "hanging_indent", "keep_with_next", "keep_together", "page_break_before")}
        if keys:
            set_paragraph_format(ppr_of(p._p), keys)

    def runs(self, p: Any, blk: dict[str, Any]) -> None:
        runs = blk.get("runs")
        if runs is None:
            text = blk.get("text", "")
            runs = [{"text": str(text), **{k: blk[k] for k in ("bold", "italic", "underline", "strike", "color", "highlight", "size", "font", "caps", "small_caps") if k in blk}}]
        if not isinstance(runs, list):
            raise UsageError('"runs" must be a list of {"text": …} objects')
        for r in runs:
            if isinstance(r, str):
                r = {"text": r}
            bad = set(r) - RUN_KEYS
            if bad:
                raise UsageError(f"unknown run field(s) {', '.join(sorted(bad))}; use {', '.join(sorted(RUN_KEYS))}")
            self.run(p, r)

    def run(self, p: Any, r: dict[str, Any]) -> None:
        from _docx import new_el, qn, sub_el
        from _runs import _text_nodes, apply_format

        if r.get("footnote") is not None and not r.get("text"):
            self.footnote(p, str(r["footnote"]))
            return
        el = new_el("w:r")
        for n in _text_nodes(str(r.get("text", ""))):
            el.append(n)
        if r.get("break") == "page":
            sub_el(el, "w:br", type="page")
        elif r.get("break") == "line":
            sub_el(el, "w:br")
        elif r.get("break"):
            raise UsageError('"break" is "page" or "line"')
        fmt = {k: v for k, v in r.items() if k in ("bold", "italic", "underline", "strike", "color", "highlight", "size", "font", "superscript", "subscript", "caps", "small_caps", "shading")}
        if r.get("code"):
            sid = self.doc.styles.find("Verbatim Char", "character")
            if sid:
                fmt["style"] = "Verbatim Char"
            else:
                fmt["font"] = "Consolas"
        if r.get("style"):
            fmt["style"] = r["style"]
        if fmt:
            apply_format(el, fmt, self.doc.styles)
        if r.get("link"):
            url = str(r["link"])
            if url.startswith("#"):
                h = new_el("w:hyperlink", anchor=url[1:])
            else:
                rid = self.doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
                h = new_el("w:hyperlink")
                h.set(qn("r:id"), rid)
            sid = self.doc.styles.find("Hyperlink", "character")
            if sid and "color" not in r:
                apply_format(el, {"style": "Hyperlink"}, self.doc.styles)
            h.append(el)
            p._p.append(h)
        else:
            p._p.append(el)
        if r.get("footnote") is not None:
            self.footnote(p, str(r["footnote"]))

    def footnote(self, p: Any, text: str) -> None:
        """Adds a real footnote (reference in the text, note in footnotes.xml)."""
        from _docx import new_el, qn, set_text, sub_el

        root = self.doc.footnotes
        if root is None:
            raise SkillError("this template has no footnotes part; use a Markdown block with [^1] footnotes instead")
        ids = [int(n.get(qn("w:id")) or 0) for n in root.findall(qn("w:footnote"))]
        nid = str(max(ids + [0]) + 1)
        note = sub_el(root, "w:footnote", id=nid)
        np_ = sub_el(note, "w:p")
        ppr = sub_el(np_, "w:pPr")
        fsid = self.doc.styles.find("Footnote Text") or self.doc.styles.find("footnote text")
        if fsid:
            sub_el(ppr, "w:pStyle", val=fsid)
        rsid = self.doc.styles.find("Footnote Reference", "character") or self.doc.styles.find("footnote reference", "character")
        r = sub_el(np_, "w:r")
        rpr = sub_el(r, "w:rPr")
        if rsid:
            sub_el(rpr, "w:rStyle", val=rsid)
        else:
            sub_el(rpr, "w:vertAlign", val="superscript")
        sub_el(r, "w:footnoteRef")
        r2 = sub_el(np_, "w:r")
        set_text(sub_el(r2, "w:t"), " " + text)
        ref = new_el("w:r")
        rpr = sub_el(ref, "w:rPr")
        if rsid:
            sub_el(rpr, "w:rStyle", val=rsid)
        else:
            sub_el(rpr, "w:vertAlign", val="superscript")
        sub_el(ref, "w:footnoteReference", id=nid)
        p._p.append(ref)

    # ── blocks ──────────────────────────────────────────────────────────
    def title_block(self, meta: dict[str, Any]) -> None:
        for key, style in (("title", "Title"), ("subtitle", "Subtitle"), ("author", "Author"), ("date", "Date")):
            if meta.get(key):
                p = self.para(style if self.doc.styles.find(style) else None)
                self.runs(p, {"text": meta[key]})

    def block(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        t = blk["type"]
        if t not in BLOCK_TYPES:
            raise UsageError(f"unknown type '{t}' (use {', '.join(BLOCK_TYPES)})")
        getattr(self, "b_" + t)(blk, meta)

    def b_title(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        # Shown: the block's own fields and what the command line gave; spec "properties" stay metadata.
        self.title_block({**self.cli, **{k: v for k, v in blk.items() if k in ("title", "subtitle", "author", "date")}, "title": blk.get("text") or blk.get("title") or self.cli.get("title") or meta.get("title")})

    def b_heading(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        level = int(blk.get("level", 1))
        if not 1 <= level <= 9:
            raise UsageError("heading level must be 1-9")
        p = self.para(f"Heading {level}")
        self.runs(p, blk)
        self.fmt_para(p, blk)

    def b_paragraph(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        p = self.para(blk.get("style") or ("Body Text" if self.doc.styles.find("Body Text") else None))
        self.runs(p, blk)
        self.fmt_para(p, blk)

    def b_quote(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        name = next((n for n in ("Quote", "Block Text", "Intense Quote") if self.doc.styles.find(n)), None)
        p = self.para(name)
        self.runs(p, blk)
        if not name:
            self.fmt_para(p, {"indent_left": "1cm", "indent_right": "1cm"})

    def b_code(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        name = next((n for n in ("Source Code", "HTML Preformatted", "No Spacing") if self.doc.styles.find(n)), None)
        for line in str(blk.get("text", "")).split("\n"):
            p = self.para(name)
            self.run(p, {"text": line, "code": True})

    def b_spacer(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        p = self.para()
        self.fmt_para(p, {"space_before": blk.get("height", "12pt"), "space_after": "0pt", "line_spacing": "1pt"})

    def b_page_break(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        self.d.add_page_break()

    def b_toc(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        from _markdown import TOC_MARKER

        title = blk.get("title", "Contents")
        self.d.add_paragraph(f"{TOC_MARKER}:{int(blk.get('depth', 3))}:{title or ''}")

    def b_markdown(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        from _markdown import markdown_elements

        markdown_elements(self.doc, str(blk.get("text", "")), resource_paths=[self.base_dir, Path.cwd()])

    def b_section(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        from docx.enum.section import WD_SECTION

        from _markdown import page_setup

        self.d.add_section(WD_SECTION.NEW_PAGE if blk.get("start", "new_page") != "continuous" else WD_SECTION.CONTINUOUS)
        self.sections += 1
        n = len(self.doc.section_elements())
        page_setup(self.doc, size=blk.get("size"), orientation=blk.get("orientation"), margins=blk.get("margins"), sections=[n], columns=blk.get("columns"))

    def b_image(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        from docx.shared import Emu

        from _markdown import parse_length

        path = Path(str(blk.get("path", ""))).expanduser()
        if not path.is_absolute() and not path.exists():
            path = self.base_dir / path
        if not path.exists():
            raise SkillError(f"image {blk.get('path')} not found")
        width = Emu(int(parse_length(blk["width"]) * 12700)) if blk.get("width") else None
        height = Emu(int(parse_length(blk["height"]) * 12700)) if blk.get("height") else None
        if width is None and height is None:
            from _markdown import section_text_width

            limit = section_text_width(self.doc.section_elements()[-1])
            from PIL import Image

            with Image.open(path) as im:
                dpi = (im.info.get("dpi") or (96, 96))[0] or 96
                natural = im.size[0] / dpi * 72
            if natural > limit:
                width = Emu(int(limit * 12700))
        fig = self.doc.styles.find("Figure") or self.doc.styles.find("Captioned Figure")
        p = self.para("Figure" if fig else None)
        p.add_run().add_picture(str(path), width=width, height=height)
        self.fmt_para(p, {"align": blk.get("align", "center")})
        if blk.get("alt") or blk.get("caption"):
            for d in p._p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
                d.set("descr", str(blk.get("alt") or blk.get("caption")))
        if blk.get("caption"):
            cap = next((n for n in ("Image Caption", "Caption") if self.doc.styles.find(n)), None)
            c = self.para(cap)
            self.runs(c, {"text": blk["caption"]})
            self.fmt_para(c, {"align": blk.get("align", "center")})

    # lists
    def _list_num(self, ordered: bool, fmt: str | None, start: int) -> str:
        """A fresh w:num (restarting at `start`) over an abstract list definition, created once per style."""
        from _docx import new_el, qn, sub_el

        npart = self.doc.docx.part.numbering_part
        root = npart.element
        key = f"{ordered}:{fmt}"
        if key not in self._abstract_ids:
            ids = [int(a.get(qn("w:abstractNumId"))) for a in root.findall(qn("w:abstractNum"))]
            aid = str(max(ids + [100]) + 1)
            an = new_el("w:abstractNum", abstractNumId=aid)
            sub_el(an, "w:multiLevelType", val="hybridMultilevel")
            bullets = ["•", "◦", "▪"]
            fmts = [fmt or "decimal", "lowerLetter", "lowerRoman"]
            for lvl in range(9):
                lv = sub_el(an, "w:lvl", ilvl=str(lvl))
                sub_el(lv, "w:start", val="1")
                if ordered:
                    f = fmts[lvl % 3] if lvl else (fmt or "decimal")
                    sub_el(lv, "w:numFmt", val=f)
                    sub_el(lv, "w:lvlText", val=f"%{lvl + 1}.")
                else:
                    sub_el(lv, "w:numFmt", val="bullet")
                    sub_el(lv, "w:lvlText", val=bullets[lvl % 3])
                sub_el(lv, "w:lvlJc", val="left")
                ppr = sub_el(lv, "w:pPr")
                sub_el(ppr, "w:ind", left=str(360 * (lvl + 1) + 360), hanging="360")
            first_num = root.find(qn("w:num"))
            (first_num.addprevious(an) if first_num is not None else root.append(an))
            self._abstract_ids[key] = aid
        aid = self._abstract_ids[key]
        ids = [int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))]
        nid = str(max(ids + [100]) + 1)
        num = sub_el(root, "w:num", numId=nid)
        sub_el(num, "w:abstractNumId", val=aid)
        ov = sub_el(num, "w:lvlOverride", ilvl="0")
        sub_el(ov, "w:startOverride", val=str(start))
        return nid

    def _items(self, items: list[Any], num_id: str, level: int, blk: dict[str, Any]) -> None:
        from _docx import new_el, ppr_of, sub_el

        style = "Compact" if self.doc.styles.find("Compact") else ("List Paragraph" if self.doc.styles.find("List Paragraph") else None)
        for it in items:
            if isinstance(it, (str, int, float)):
                it = {"text": str(it)}
            if not isinstance(it, dict):
                raise UsageError("list items are strings or {\"text\"/\"runs\", \"items\": [...]} objects")
            p = self.para(style)
            ppr = ppr_of(p._p)
            numpr = new_el("w:numPr")
            sub_el(numpr, "w:ilvl", val=str(level))
            sub_el(numpr, "w:numId", val=num_id)
            ps = ppr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
            (ps.addnext(numpr) if ps is not None else ppr.insert(0, numpr))
            self.runs(p, it)
            if it.get("items"):
                self._items(it["items"], num_id, min(level + 1, 8), blk)

    def b_bullets(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        self._items(blk.get("items") or [], self._list_num(False, None, 1), 0, blk)

    def b_numbered(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        fmt = blk.get("format")
        if fmt and fmt not in ("decimal", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman", "decimalZero"):
            raise UsageError("format: decimal, lowerLetter, upperLetter, lowerRoman, upperRoman or decimalZero")
        self._items(blk.get("items") or [], self._list_num(True, fmt, int(blk.get("start", 1))), 0, blk)

    # tables
    def b_table(self, blk: dict[str, Any], meta: dict[str, Any]) -> None:
        from docx.shared import Emu

        from _docx import new_el, qn
        from _markdown import parse_length
        from _runs import apply_format

        header = blk.get("header")
        rows = blk.get("rows") or []
        if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
            raise UsageError('"rows" must be a list of lists')
        all_rows = ([header] if header else []) + rows
        if not all_rows:
            raise UsageError("a table needs \"rows\" (and optionally \"header\")")
        ncols = max(len(r) for r in all_rows)
        style = blk.get("style") or next((s for s in ("Table", "Table Grid") if self.doc.styles.find(s, "table")), None)
        t = self.d.add_table(rows=len(all_rows), cols=ncols)
        if style:
            sid = self.style(style, "table")
            tblpr = t._tbl.tblPr
            ts = tblpr.find(qn("w:tblStyle"))
            if ts is None:
                ts = new_el("w:tblStyle")
                tblpr.insert(0, ts)
            ts.set(qn("w:val"), sid)
        widths = blk.get("widths")
        if widths:
            if len(widths) != ncols:
                raise UsageError(f"{len(widths)} widths for {ncols} columns")
            emus = [Emu(int(parse_length(w) * 12700)) for w in widths]
            t.autofit = False
            for ci, w in enumerate(emus):
                t.columns[ci].width = w
                for cell in t.columns[ci].cells:
                    cell.width = w
        aligns = blk.get("align") or []
        shading = blk.get("shading") or {}
        font_size = blk.get("font_size")
        # Columns of numbers (123, "1,200.00", "€ 45", "12%") are right-aligned, header included, unless "align" says.
        numeric_cols = set()
        for ci in range(ncols):
            vals = [r[ci] for r in rows if ci < len(r) and r[ci] not in (None, "") and not (isinstance(r[ci], dict) and not r[ci].get("text"))]
            if vals and all(looks_numeric(v) for v in vals):
                numeric_cols.add(ci)
        for ri, row in enumerate(all_rows):
            is_header = bool(header) and ri == 0
            for ci in range(ncols):
                val = row[ci] if ci < len(row) else ""
                cell = t.cell(ri, ci)
                spec = val if isinstance(val, dict) else {"text": "" if val is None else str(val)}
                p = cell.paragraphs[0]
                self.runs(p, spec)
                al = spec.get("align") or (aligns[ci] if ci < len(aligns) else None)
                if al is None and ci in numeric_cols:
                    al = "right"
                if al:
                    self.fmt_para(p, {"align": al})
                if font_size:
                    for r in p.runs:
                        apply_format(r._r, {"size": font_size})
                fill = spec.get("fill") or (shading.get("header") if is_header else None)
                band = shading.get("rows")
                if not fill and band and not is_header:
                    k = ri - (1 if header else 0)
                    fill = band[k % len(band)] if isinstance(band, list) and band else None
                if fill:
                    tcpr = cell._tc.get_or_add_tcPr()
                    shd = new_el("w:shd", val="clear", color="auto", fill=str(fill).lstrip("#").upper())
                    tcpr.append(shd)
                    if is_dark(str(fill)):
                        # Text on a dark fill is white unless the cell or run sets its own colour.
                        runs_spec = spec.get("runs") if isinstance(spec.get("runs"), list) else [spec]
                        for r, rs in zip(p.runs, runs_spec + [{}] * len(p.runs)):
                            if not (isinstance(rs, dict) and rs.get("color")) and not spec.get("color"):
                                apply_format(r._r, {"color": "FFFFFF"})
                if is_header and blk.get("bold_header", True) and not style:
                    for r in p.runs:
                        apply_format(r._r, {"bold": True})
        if header:
            tr = t.rows[0]._tr
            trpr = tr.get_or_add_trPr()
            trpr.append(new_el("w:tblHeader"))
        for m in blk.get("merge") or []:
            (r1, c1), (r2, c2) = m["from"], m["to"]
            off = 1 if header else 0
            a = t.cell(int(r1) + off, int(c1))
            b = t.cell(int(r2) + off, int(c2))
            keep = a.text
            merged = a.merge(b)
            # python-docx concatenates the texts of merged cells; keep the first cell's text.
            for extra in merged.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
        if blk.get("caption"):
            cap = next((n for n in ("Table Caption", "Caption") if self.doc.styles.find(n)), None)
            c = self.para(cap)
            self.runs(c, {"text": blk["caption"]})
            # Captions go above tables.
            t._tbl.addprevious(c._p)
