"""Markdown -> Word through the bundled pandoc, with Desk's reference styles, plus the layout helpers
docx_create, docx_edit and docx_convert share: page setup, headers and footers with fields, title page, TOC."""

from __future__ import annotations

import copy
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any, Sequence

from lxml import etree

from _common import SkillError, UsageError
from _docx import NS, P, PPR, R, RT_STYLES, SECTPR, Doc, get_or_add, new_el, qn, set_text, sub_el, wattr

_W = "{" + NS["w"] + "}"

# ── presets ─────────────────────────────────────────────────────────────

PRESETS: dict[str, dict[str, Any]] = {
    "report": {"body": "Calibri", "heading": "Calibri", "mono": "Consolas", "size": 11, "heading_color": "1F3864", "accent": "2E74B5", "muted": "595959", "line": 1.15, "after": 8, "title_align": "left", "h_sizes": [18, 14, 12, 11, 11, 11], "rule": "BDD7EE"},
    "classic": {"body": "Georgia", "heading": "Georgia", "mono": "Courier New", "size": 11, "heading_color": "222222", "accent": "7B2C2C", "muted": "555555", "line": 1.2, "after": 8, "title_align": "center", "h_sizes": [17, 14, 12, 11, 11, 11], "rule": "BFBFBF"},
    "modern": {"body": "Arial", "heading": "Arial", "mono": "Courier New", "size": 10.5, "heading_color": "0B3D5C", "accent": "00897B", "muted": "5F6B73", "line": 1.2, "after": 7, "title_align": "left", "h_sizes": [17, 13.5, 11.5, 10.5, 10.5, 10.5], "rule": "B2DFDB"},
}

PAGE_SIZES = {"a4": (210, 297), "letter": (215.9, 279.4), "legal": (215.9, 355.6), "a5": (148, 210), "a3": (297, 420), "executive": (184.15, 266.7), "tabloid": (279.4, 431.8), "b5": (176, 250)}


def parse_length(s: str | float | int, default_unit: str = "cm") -> float:
    """A length like '2.5cm', '1in', '20mm', '12pt' or a bare number (in default_unit) -> points."""
    if isinstance(s, (int, float)):
        s = f"{s}{default_unit}"
    m = re.fullmatch(r"\s*(-?[\d.]+)\s*(cm|mm|in|pt|px|%|)\s*", str(s))
    if not m:
        raise UsageError(f"bad length '{s}' (use e.g. 2.5cm, 1in, 20mm, 12pt)")
    n = float(m.group(1))
    unit = m.group(2) or default_unit
    return n * {"cm": 72 / 2.54, "mm": 72 / 25.4, "in": 72.0, "pt": 1.0, "px": 0.75}[unit]


def parse_page_size(s: str) -> tuple[float, float]:
    """(width, height) in points from 'A4', 'Letter', '21x29.7cm' or '8.5x11in'."""
    key = s.strip().lower()
    if key in PAGE_SIZES:
        w, h = PAGE_SIZES[key]
        return w * 72 / 25.4, h * 72 / 25.4
    m = re.fullmatch(r"([\d.]+)\s*x\s*([\d.]+)\s*(cm|mm|in|pt)?", key)
    if not m:
        raise UsageError(f"bad page size '{s}' (A4, Letter, Legal, A5, A3, or WxH like 21x29.7cm)")
    unit = m.group(3) or "cm"
    return parse_length(m.group(1) + unit), parse_length(m.group(2) + unit)


def parse_margins(s: str) -> dict[str, float]:
    """'2cm' (all), '2cm,3cm' (vertical, horizontal) or 'top,right,bottom,left'."""
    parts = [x.strip() for x in str(s).split(",") if x.strip()]
    vals = [parse_length(x) for x in parts]
    if len(vals) == 1:
        t = r = b = l = vals[0]
    elif len(vals) == 2:
        t = b = vals[0]
        r = l = vals[1]
    elif len(vals) == 4:
        t, r, b, l = vals
    else:
        raise UsageError(f"bad margins '{s}' (one, two or four lengths)")
    return {"top": t, "right": r, "bottom": b, "left": l}


# ── reference document ──────────────────────────────────────────────────


def _style_xml(sid: str, name: str, kind: str = "paragraph", based: str | None = None, nxt: str | None = None, ppr: str = "", rpr: str = "", extra: str = "", custom: bool = False, link: str | None = None, ui: int | None = None) -> str:
    parts = [f'<w:style xmlns:w="{NS["w"]}" w:type="{kind}" w:styleId="{sid}"' + (' w:customStyle="1"' if custom else "") + ">", f'<w:name w:val="{name}"/>']
    if based:
        parts.append(f'<w:basedOn w:val="{based}"/>')
    if nxt:
        parts.append(f'<w:next w:val="{nxt}"/>')
    if link:
        parts.append(f'<w:link w:val="{link}"/>')
    if ui is not None:
        parts.append(f'<w:uiPriority w:val="{ui}"/>')
    parts.append("<w:qFormat/>")
    if ppr:
        parts.append(_in_schema_order("pPr", ppr))
    if rpr:
        parts.append(_in_schema_order("rPr", rpr))
    parts.append(extra)
    parts.append("</w:style>")
    return "".join(parts)


def _in_schema_order(tag: str, inner: str) -> str:
    """<w:pPr>/<w:rPr> markup with its children sorted into the order the schema requires (Word is strict)."""
    from lxml import etree

    from _runs import RPR_ORDER

    order = PPR_ORDER if tag == "pPr" else RPR_ORDER
    el = etree.fromstring(f'<w:{tag} xmlns:w="{NS["w"]}">{inner}</w:{tag}>')
    kids = list(el)
    rank = {name: i for i, name in enumerate(order)}
    kids.sort(key=lambda c: rank.get(etree.QName(c).localname, len(order)))
    for c in kids:
        el.append(c)
    return etree.tostring(el, encoding="unicode")


def _fonts(name: str) -> str:
    return f'<w:rFonts w:ascii="{name}" w:hAnsi="{name}" w:eastAsia="{name}" w:cs="{name}"/>'


def _sz(pt: float) -> str:
    v = int(round(pt * 2))
    return f'<w:sz w:val="{v}"/><w:szCs w:val="{v}"/>'


def preset_styles(p: dict[str, Any]) -> list[str]:
    """Desk's style definitions for a preset (replacing pandoc's defaults)."""
    body, head, mono = p["body"], p["heading"], p["mono"]
    hc, accent, muted = p["heading_color"], p["accent"], p["muted"]
    line = int(round(240 * p["line"]))
    after = int(p["after"] * 20)
    hs = p["h_sizes"]
    title_jc = f'<w:jc w:val="{p["title_align"]}"/>'
    out = [
        _style_xml("Normal", "Normal", ppr=f'<w:spacing w:after="{after}" w:line="{line}" w:lineRule="auto"/>', extra="", ui=0).replace("<w:qFormat/>", "<w:qFormat/>").replace('w:styleId="Normal"', 'w:default="1" w:styleId="Normal"'),
        _style_xml("BodyText", "Body Text", based="Normal", link="BodyTextChar", ppr=f'<w:spacing w:before="0" w:after="{after + 40}"/>'),
        _style_xml("FirstParagraph", "First Paragraph", based="BodyText", nxt="BodyText", custom=True),
        _style_xml("Compact", "Compact", based="BodyText", custom=True, ppr='<w:spacing w:before="20" w:after="40"/>'),
        _style_xml("Title", "Title", based="Normal", nxt="BodyText", ui=10, ppr=f'<w:spacing w:before="0" w:after="120" w:line="240" w:lineRule="auto"/><w:contextualSpacing/>{title_jc}', rpr=f'{_fonts(head)}<w:b/><w:color w:val="{hc}"/><w:kern w:val="28"/>{_sz(28)}'),
        _style_xml("Subtitle", "Subtitle", based="Title", nxt="BodyText", ui=11, ppr=f'<w:spacing w:before="0" w:after="200"/>{title_jc}', rpr=f'<w:b w:val="0"/><w:color w:val="{muted}"/>{_sz(15)}'),
        _style_xml("Author", "Author", based="Normal", nxt="BodyText", custom=True, ppr=f'<w:keepNext/><w:spacing w:before="0" w:after="40"/>{title_jc}', rpr=f'<w:color w:val="{muted}"/>{_sz(p["size"] + 1)}'),
        _style_xml("Date", "Date", based="Author", nxt="BodyText", ppr='<w:spacing w:after="240"/>'),
        _style_xml("Abstract", "Abstract", based="Normal", nxt="BodyText", custom=True, ppr='<w:spacing w:before="120" w:after="240"/><w:ind w:left="567" w:right="567"/>', rpr=f'<w:i/><w:color w:val="{muted}"/>{_sz(p["size"] - 0.5)}'),
        _style_xml("AbstractTitle", "Abstract Title", based="Normal", nxt="Abstract", custom=True, ppr='<w:keepNext/><w:spacing w:before="240" w:after="0"/><w:jc w:val="center"/>', rpr="<w:b/>"),
    ]
    befores = [360, 280, 220, 180, 160, 160]
    afters = [120, 80, 60, 40, 40, 40]
    for i in range(6):
        lvl = i + 1
        rpr = f'{_fonts(head)}<w:b/><w:color w:val="{hc if lvl <= 2 else accent if lvl == 3 else hc}"/>{_sz(hs[i])}'
        if lvl >= 5:
            rpr += "<w:i/>"
        border = f'<w:pBdr><w:bottom w:val="single" w:sz="6" w:space="2" w:color="{p["rule"]}"/></w:pBdr>' if lvl == 1 else ""
        out.append(_style_xml(f"Heading{lvl}", f"heading {lvl}", based="Normal", nxt="BodyText", link=f"Heading{lvl}Char", ui=9, ppr=f'<w:keepNext/><w:keepLines/>{border}<w:spacing w:before="{befores[i]}" w:after="{afters[i]}" w:line="240" w:lineRule="auto"/><w:outlineLvl w:val="{i}"/>', rpr=rpr))
    for lvl in (7, 8, 9):
        out.append(_style_xml(f"Heading{lvl}", f"heading {lvl}", based="Normal", nxt="BodyText", ui=9, ppr=f'<w:keepNext/><w:keepLines/><w:spacing w:before="120" w:after="40"/><w:outlineLvl w:val="{lvl - 1}"/>', rpr=f'{_fonts(head)}<w:i/><w:color w:val="{muted}"/>'))
    out += [
        _style_xml("BlockText", "Block Text", based="BodyText", nxt="BodyText", ppr=f'<w:pBdr><w:left w:val="single" w:sz="18" w:space="8" w:color="{p["rule"]}"/></w:pBdr><w:spacing w:before="120" w:after="120"/><w:ind w:left="397" w:right="397"/>', rpr=f'<w:i/><w:color w:val="{muted}"/>'),
        _style_xml("Quote", "Quote", based="BlockText", nxt="BodyText", ui=29),
        _style_xml("IntenseQuote", "Intense Quote", based="BlockText", nxt="BodyText", ui=30, rpr=f'<w:b/><w:color w:val="{accent}"/>'),
        _style_xml("SourceCode", "Source Code", based="Normal", custom=True, ppr='<w:shd w:val="clear" w:color="auto" w:fill="F4F5F7"/><w:wordWrap w:val="off"/><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/><w:ind w:left="113" w:right="113"/>', rpr=f"{_fonts(mono)}{_sz(max(8, p['size'] - 1.5))}"),
        _style_xml("VerbatimChar", "Verbatim Char", kind="character", based="DefaultParagraphFont", custom=True, rpr=f'{_fonts(mono)}{_sz(max(8, p["size"] - 1.5))}<w:shd w:val="clear" w:color="auto" w:fill="F4F5F7"/>'),
        _style_xml("Hyperlink", "Hyperlink", kind="character", based="DefaultParagraphFont", ui=99, rpr=f'<w:color w:val="{accent}"/><w:u w:val="single"/>'),
        _style_xml("Caption", "caption", based="Normal", nxt="BodyText", ui=35, ppr='<w:spacing w:before="60" w:after="200"/>', rpr=f'<w:i/><w:color w:val="{muted}"/>{_sz(p["size"] - 1.5)}'),
        _style_xml("ImageCaption", "Image Caption", based="Caption", custom=True, ppr='<w:jc w:val="center"/>'),
        _style_xml("TableCaption", "Table Caption", based="Caption", custom=True, ppr='<w:keepNext/><w:spacing w:before="200" w:after="60"/>'),
        _style_xml("Figure", "Figure", based="Normal", custom=True, ppr='<w:keepNext/><w:spacing w:before="120" w:after="60"/><w:jc w:val="center"/>'),
        _style_xml("CaptionedFigure", "Captioned Figure", based="Figure", custom=True, ppr="<w:keepNext/>"),
        _style_xml("FootnoteText", "footnote text", based="Normal", link="FootnoteTextChar", ui=99, ppr='<w:spacing w:after="0" w:line="240" w:lineRule="auto"/>', rpr=_sz(p["size"] - 2)),
        _style_xml("FootnoteReference", "footnote reference", kind="character", based="DefaultParagraphFont", ui=99, rpr='<w:vertAlign w:val="superscript"/>'),
        _style_xml("Header", "header", based="Normal", link="HeaderChar", ui=99, ppr='<w:tabs><w:tab w:val="center" w:pos="4680"/><w:tab w:val="right" w:pos="9360"/></w:tabs><w:spacing w:after="0" w:line="240" w:lineRule="auto"/>', rpr=f'<w:color w:val="{muted}"/>{_sz(p["size"] - 2)}'),
        _style_xml("Footer", "footer", based="Header", link="FooterChar", ui=99),
        _style_xml("TOCHeading", "TOC Heading", based="Heading1", nxt="BodyText", ui=39, ppr='<w:outlineLvl w:val="9"/>'),
        _style_xml("NoSpacing", "No Spacing", based="Normal", ui=1, ppr='<w:spacing w:after="0" w:line="240" w:lineRule="auto"/>'),
        _style_xml("ListParagraph", "List Paragraph", based="Normal", ui=34, ppr='<w:ind w:left="720"/><w:contextualSpacing/>'),
        _style_xml("Strong", "Strong", kind="character", based="DefaultParagraphFont", ui=22, rpr="<w:b/><w:bCs/>"),
        _style_xml("Emphasis", "Emphasis", kind="character", based="DefaultParagraphFont", ui=20, rpr="<w:i/><w:iCs/>"),
    ]
    for lvl in (1, 2, 3, 4):
        out.append(_style_xml(f"TOC{lvl}", f"toc {lvl}", based="Normal", nxt="Normal", ui=39, ppr=f'<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9350"/></w:tabs><w:spacing w:before="{60 if lvl == 1 else 0}" w:after="{40}"/><w:ind w:left="{(lvl - 1) * 240}"/>', rpr="<w:b/>" if lvl == 1 else ""))
    cell_mar = '<w:tblCellMar><w:top w:w="57" w:type="dxa"/><w:left w:w="108" w:type="dxa"/><w:bottom w:w="57" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar>'
    table = (
        f'<w:style xmlns:w="{NS["w"]}" w:type="table" w:default="1" w:styleId="Table"><w:name w:val="Table"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="59"/><w:qFormat/>'
        '<w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        f'<w:rPr>{_sz(p["size"] - 0.5)}</w:rPr>'
        f'<w:tblPr><w:tblInd w:w="0" w:type="dxa"/><w:tblBorders><w:top w:val="single" w:sz="8" w:space="0" w:color="{hc}"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="{hc}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/></w:tblBorders>{cell_mar}</w:tblPr>'
        f'<w:tblStylePr w:type="firstRow"><w:rPr><w:b/><w:color w:val="{hc}"/></w:rPr><w:tcPr><w:tcBorders><w:bottom w:val="single" w:sz="8" w:space="0" w:color="{hc}"/></w:tcBorders><w:shd w:val="clear" w:color="auto" w:fill="F2F4F7"/><w:vAlign w:val="bottom"/></w:tcPr></w:tblStylePr>'
        "</w:style>"
    )
    grid = (
        f'<w:style xmlns:w="{NS["w"]}" w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="39"/>'
        '<w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/></w:tblBorders>'
        f"{cell_mar}</w:tblPr></w:style>"
    )
    out += [table, grid]
    return out


def build_reference(dest: Path, preset: str = "report", overrides: dict[str, Any] | None = None) -> Path:
    """Writes a reference .docx with Desk's professional styles (pandoc takes styles, page setup and headers from it)."""
    from _render import run_pandoc

    if preset not in PRESETS:
        raise UsageError(f"unknown style preset '{preset}' (choose {', '.join(PRESETS)})")
    p = dict(PRESETS[preset])
    for k, v in (overrides or {}).items():
        if v is not None:
            p[k] = v
    dest.parent.mkdir(parents=True, exist_ok=True)
    data = run_pandoc(["--print-default-data-file", "reference.docx"])
    dest.write_bytes(data)
    doc = Doc(dest)
    root = doc.rel_root(RT_STYLES)
    size = p["size"]
    dd = root.find(qn("w:docDefaults"))
    if dd is not None:
        root.remove(dd)
    dd = etree.fromstring(
        f'<w:docDefaults xmlns:w="{NS["w"]}"><w:rPrDefault><w:rPr>{_fonts(p["body"])}{_sz(size)}<w:lang w:val="{p.get("lang", "en-US")}" w:eastAsia="zh-CN" w:bidi="ar-SA"/></w:rPr></w:rPrDefault>'
        f'<w:pPrDefault><w:pPr><w:spacing w:after="{int(p["after"] * 20)}" w:line="{int(round(240 * p["line"]))}" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
    )
    root.insert(0, dd)
    for xml in preset_styles(p):
        el = etree.fromstring(xml)
        sid = el.get(_W + "styleId")
        for old in root.findall(qn("w:style")):
            if old.get(_W + "styleId") == sid:
                root.remove(old)
        root.append(el)
    # Letter or A4 by locale: A4 unless the caller asks otherwise; margins 2.5 cm.
    sect = doc.body.find(SECTPR)
    width, height = parse_page_size(p.get("page", "A4"))
    _set_page(sect, width, height, parse_margins(p.get("margins", "2.5cm")), None)
    doc.save(dest, force=True)
    return dest


# ── Markdown -> docx ────────────────────────────────────────────────────

PAGEBREAK_XML = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```'
TOC_MARKER = "DESK-TOC-PLACEHOLDER"


_TAG = re.compile(r"\{\{.*?\}\}")
_CODE = re.compile(r"(`+).*?\1")
_WORDS = re.compile(r"(?<![\\\w])[A-Za-z]{2,}(?:\s+[A-Za-z]{2,}){2,}")


def _protect_tag(m: re.Match[str]) -> str:
    """A {{template}} tag, kept literal: its @ would start an e-mail autolink, its | would split a table cell,
    its $ * _ ~ ^ would read as math or emphasis, and its quotes would turn curly."""
    return re.sub(r"(?<!\\)([@|$*_~^\"'])", r"\\\1", m.group(0))


def _balanced(s: str) -> bool:
    depth = {"(": 0, "[": 0, "{": 0}
    close = {")": "(", "]": "[", "}": "{"}
    prev = ""
    for ch in s:
        if prev == "\\":
            prev = ""
            continue
        if ch in depth:
            depth[ch] += 1
        elif ch in close:
            depth[close[ch]] -= 1
            if depth[close[ch]] < 0:
                return False
        prev = ch
    return not any(depth.values())


def _dollars(seg: str) -> str:
    """Escapes the $ signs of one line (or table cell) that are money, not TeX math.

    pandoc reads $…$ as math even across lines and table cells, so '| Revenue ($M) |' two rows above
    '| Profit ($M) |' swallowed the rows between. Math is kept only when both $ sit in the same line or cell,
    follow pandoc's rules (no space inside the delimiters, no digit right after the closing one), and the
    content has balanced brackets and is not a run of plain words."""
    out: list[str] = []
    i, n = 0, len(seg)
    while i < n:
        ch = seg[i]
        if ch == "\\" and i + 1 < n:
            out.append(seg[i : i + 2])
            i += 2
            continue
        if ch == "`":
            m = _CODE.match(seg, i)
            if m:
                out.append(m.group(0))
                i = m.end()
                continue
        if ch != "$":
            out.append(ch)
            i += 1
            continue
        if seg.startswith("$$", i):
            end = seg.find("$$", i + 2)
            if end > i + 2:
                out.append(seg[i : end + 2])
                i = end + 2
                continue
            out.append("\\$\\$")
            i += 2
            continue
        # an inline math candidate
        close = -1
        if i + 1 < n and not seg[i + 1].isspace():
            j = i + 1
            while j < n:
                if seg[j] == "\\":
                    j += 2
                    continue
                if seg[j] == "$" and not seg[j - 1].isspace() and not (j + 1 < n and seg[j + 1].isdigit()):
                    close = j
                    break
                j += 1
        if close > 0:
            body = seg[i + 1 : close]
            if _balanced(body) and not _WORDS.search(body):
                out.append(seg[i : close + 1])
                i = close + 1
                continue
        out.append("\\$")
        i += 1
    return "".join(out)


def _protect_line(line: str) -> str:
    """Template tags and money signs made literal in one Markdown line (outside fenced code)."""
    tags: list[str] = []

    def keep(m: re.Match[str]) -> str:
        tags.append(_protect_tag(m))
        return f"\x00{len(tags) - 1}\x00"

    work = _TAG.sub(keep, line) if "{{" in line else line
    if "$" in work:
        stripped = work.lstrip()
        if stripped.startswith("|") and not line.startswith("    "):
            # a pipe-table row: each cell on its own (split at | outside code spans and escapes)
            cells, buf, k = [], [], 0
            while k < len(work):
                c = work[k]
                if c == "\\" and k + 1 < len(work):
                    buf.append(work[k : k + 2])
                    k += 2
                    continue
                if c == "`":
                    m = _CODE.match(work, k)
                    if m:
                        buf.append(m.group(0))
                        k = m.end()
                        continue
                if c == "|":
                    cells.append("".join(buf))
                    buf = []
                    k += 1
                    continue
                buf.append(c)
                k += 1
            cells.append("".join(buf))
            work = "|".join(_dollars(c) for c in cells)
        elif not line.startswith(("    ", "\t")):
            work = _dollars(work)
    if tags:
        work = re.sub("\x00(\\d+)\x00", lambda m: tags[int(m.group(1))], work)
    return work


def preprocess(md: str, protect: bool = True) -> str:
    """Desk's Markdown extras: a line with \\pagebreak (or \\newpage) and a line with [[toc]] (or \\toc). With
    `protect`, {{template}} tags stay literal and money $ signs are not read as math (see _dollars)."""
    out = []
    in_fence = False
    in_meta = False
    for k, line in enumerate(md.splitlines()):
        s = line.strip()
        if k == 0 and s == "---":
            in_meta = True
            out.append(line)
            continue
        if in_meta:
            in_meta = s not in ("---", "...")
            out.append(line)
            continue
        if s.startswith("```") or s.startswith("~~~"):
            in_fence = not in_fence
        if protect and not in_fence and ("$" in line or "{{" in line):
            line = _protect_line(line)
            s = line.strip()
        if not in_fence and s in ("\\pagebreak", "\\newpage", "<!-- pagebreak -->", "<!-- page break -->"):
            out.append("")
            out.append(PAGEBREAK_XML)
            out.append("")
            continue
        if not in_fence and s.lower() in ("[[toc]]", "\\toc", "[toc]"):
            out.append("")
            out.append(f'```{{=openxml}}\n<w:p><w:r><w:t>{TOC_MARKER}</w:t></w:r></w:p>\n```')
            out.append("")
            continue
        out.append(line)
    text = "\n".join(out) + ("\n" if md.endswith("\n") else "")
    return _drop_repeated_captions(text) if protect else text


_FIGURE = re.compile(r"(?m)^[ \t]*!\[((?:\\.|[^\]\\])+)\]\([^)\n]*\)(?:\{[^}\n]*\})?[ \t]*\n[ \t]*\n([^\n]+)\n?")


def _drop_repeated_captions(text: str) -> str:
    """An image alone in its paragraph becomes a figure whose caption is its alt text. When the next paragraph
    repeats that text (docx_read writes a captioned picture that way), it is dropped: one caption, not two."""

    def plain(s: str) -> str:
        return re.sub(r"[\\*_`]", "", s).strip().lower()

    def fix(m: re.Match[str]) -> str:
        if plain(m.group(2)) == plain(m.group(1)):
            end = m.start(2)
            return m.group(0)[: end - m.start(0)].rstrip("\n \t") + "\n"
        return m.group(0)

    return _FIGURE.sub(fix, text) if "![" in text else text


def md_to_docx(md: str, dest: Path, reference: Path | None, *, resource_paths: Sequence[Path] = (), smart: bool = True, metadata: dict[str, str] | None = None, toc: bool = False, toc_depth: int = 3, number_sections: bool = False, fmt: str = "markdown", extra: Sequence[str] = ()) -> Path:
    """Runs pandoc on Markdown (or another input format) and writes dest. pandoc runs sandboxed: pictures and
    includes come only from `resource_paths` (the first is the document's folder), see _docx.safe_pandoc."""
    from _docx import safe_pandoc
    from _pandoc_safe import has_includes, inline_includes, resolved

    roots = resolved(resource_paths)
    tmpdir = Path(tempfile.mkdtemp(prefix="desk-md-"))
    try:
        src = tmpdir / "input.md"
        if roots and has_includes(md, fmt):
            warnings: list[str] = []
            try:
                md = inline_includes(md, fmt, roots[0], roots, warnings)
            except ValueError as e:
                raise SkillError(str(e)) from None
            for w in warnings:
                print(f"warning: {w}", file=sys.stderr)
        src.write_text(preprocess(md) if fmt.startswith("markdown") else md, encoding="utf-8")
        reader = fmt
        if fmt.startswith("markdown"):
            reader = "markdown+raw_attribute+pipe_tables+footnotes+tex_math_dollars+strikeout+superscript+subscript+task_lists+autolink_bare_uris" + ("" if smart else "-smart")
        args = ["-f", reader, "-t", "docx", "-o", str(dest)]
        if reference is not None:
            args += ["--reference-doc", str(reference)]
        paths = [str(p) for p in resource_paths] + ["."]
        args += ["--resource-path", _pathsep().join(paths)]
        for k, v in (metadata or {}).items():
            if v:
                args += ["-M", f"{k}={v}"]
        if toc:
            args += ["--toc", f"--toc-depth={toc_depth}"]
        if number_sections:
            args += ["--number-sections"]
        args += list(extra)
        args.append(str(src))
        safe_pandoc(args, roots=roots or resolved([Path.cwd()]), cwd=str(resource_paths[0]) if resource_paths else None)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    return dest


def _pathsep() -> str:
    import os

    return os.pathsep


# ── fields ──────────────────────────────────────────────────────────────


def field_runs(instr: str, cached: str | None = None, rpr: Any = None) -> list[Any]:
    """The runs of a complex field: begin, instruction, separate, cached result, end."""
    runs = []

    def run() -> Any:
        r = new_el("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        return r

    r = run()
    sub_el(r, "w:fldChar", fldCharType="begin")
    runs.append(r)
    r = run()
    it = sub_el(r, "w:instrText")
    set_text(it, f" {instr} ")
    runs.append(r)
    r = run()
    sub_el(r, "w:fldChar", fldCharType="separate")
    runs.append(r)
    r = run()
    if cached is None:
        cached = "1"
    if cached:
        set_text(sub_el(r, "w:t"), cached)
    runs.append(r)
    r = run()
    sub_el(r, "w:fldChar", fldCharType="end")
    runs.append(r)
    return runs


TOKEN_FIELDS = {"page": ("PAGE", "1"), "pages": ("NUMPAGES", "1"), "numpages": ("NUMPAGES", "1"), "section": ("SECTION", "1"), "sectionpages": ("SECTIONPAGES", "1"), "date": ("DATE \\@ \"d MMMM yyyy\"", None), "title": ("TITLE", None), "author": ("AUTHOR", None), "filename": ("FILENAME", None)}


def text_with_fields(p: Any, text: str, doc: Doc | None = None, rpr: Any = None) -> None:
    """Appends runs for text where {page}, {pages}, {date}, {title}, {author}, {section} become live fields."""
    import datetime as _dt

    pos = 0
    for m in re.finditer(r"\{(page|pages|numpages|section|sectionpages|date|title|author|filename)\}", text, re.I):
        if m.start() > pos:
            _plain_runs(p, text[pos : m.start()], rpr)
        key = m.group(1).lower()
        instr, cached = TOKEN_FIELDS[key]
        if cached is None:
            if key == "date":
                today = _dt.date.today()
                cached = f"{today.day} {today.strftime('%B %Y')}"
            elif doc is not None and key in ("title", "author"):
                cached = getattr(doc.docx.core_properties, key) or ""
            else:
                cached = ""
        for r in field_runs(instr, cached, rpr):
            p.append(r)
        pos = m.end()
    if pos < len(text):
        _plain_runs(p, text[pos:], rpr)


def _plain_runs(p: Any, text: str, rpr: Any = None) -> None:
    for i, part in enumerate(re.split(r"(\t|\n)", text)):
        if part == "":
            continue
        r = new_el("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        if part == "\t":
            sub_el(r, "w:tab")
        elif part == "\n":
            sub_el(r, "w:br")
        else:
            set_text(sub_el(r, "w:t"), part)
        p.append(r)


# ── headers, footers and page setup ─────────────────────────────────────


def _hdr_ftr_part(doc: Doc, sect: Any, kind: str, variant: str) -> Any:
    """The header/footer part a section uses for a variant, created (and referenced) when missing."""
    from docx.section import Section

    s = Section(sect, doc.part)
    if kind == "header":
        hf = {"default": s.header, "first": s.first_page_header, "even": s.even_page_header}[variant]
    else:
        hf = {"default": s.footer, "first": s.first_page_footer, "even": s.even_page_footer}[variant]
    hf.is_linked_to_previous = False
    return hf


def set_header_footer(doc: Doc, kind: str, text: str, *, align: str = "center", sections: Sequence[int] | None = None, variant: str = "default", tracker: Any = None) -> int:
    """Replaces the header or footer text. 'left|center|right' puts three parts on tab stops. Returns sections changed.

    Without `sections`, the whole document gets the text: section 1 holds it and later sections inherit it, except
    sections of another text width (a landscape annex), which get their own copy with tab stops at their own
    width, so the right-hand part stays at the right margin. With a `tracker` (docx_edit --track), the old
    text is kept as a tracked deletion and the new text is a tracked insertion."""
    sects = doc.section_elements()
    ref_tag = qn(f"w:{kind}Reference")
    if not sections:
        targets = []
        width = None
        for i, sect in enumerate(sects):
            w = round(section_text_width(sect))
            if i == 0 or w != width:
                targets.append(i)
                width = w
            else:
                # Same width as the section before: inherit its header (linked to previous).
                for ref in sect.findall(ref_tag):
                    if (wattr(ref, "type") or "default") == variant:
                        sect.remove(ref)
    else:
        targets = [i - 1 for i in sections]
    changed = 0
    style = "Header" if kind == "header" else "Footer"
    for i in targets:
        if i < 0 or i >= len(sects):
            raise SkillError(f"section {i + 1} does not exist (the document has {len(sects)})")
        sect = sects[i]
        if variant == "first":
            tp = get_or_add(sect, "w:titlePg", before=("w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange"))
            tp.attrib.pop(_W + "val", None)
        if variant == "even":
            from _docx import set_setting

            set_setting(doc, "evenAndOddHeaders")
        had_own = any((wattr(r, "type") or "default") == variant for r in sect.findall(ref_tag))
        hf = _hdr_ftr_part(doc, sect, kind, variant)
        el = hf._element
        old = [c for c in el if c.tag in (P, qn("w:tbl"), qn("w:sdt"))]
        if tracker is not None and had_own:
            from _docx import iter_blocks

            for blk in list(iter_blocks(el)):
                if blk.tag == P:
                    tracker.delete_paragraph(blk)
                else:
                    for tr in blk.iter(qn("w:tr")):
                        tracker.delete_row(tr)
        else:
            for child in old:
                el.remove(child)
        p = sub_el(el, "w:p")
        ppr = sub_el(p, "w:pPr")
        sid = doc.styles.find(style)
        if sid:
            sub_el(ppr, "w:pStyle", val=sid)
        if "|" in text:
            parts = [x.strip() for x in (text.split("|") + ["", ""])[:3]]
            width = section_text_width(sect)
            stops = []
            if parts[1]:
                stops.append(("center", width / 2, None))
            if parts[2]:
                stops.append(("right", width, None))
            set_tab_stops(doc, ppr, stops)
            # Only the parts that have text get a tab: 'left||right' is one tab to the right margin.
            text_with_fields(p, parts[0], doc)
            for part in parts[1:]:
                if part:
                    _plain_runs(p, "\t")
                    text_with_fields(p, part, doc)
        else:
            jc = {"left": "left", "center": "center", "right": "right", "centre": "center"}.get(align, "center")
            sub_el(ppr, "w:jc", val=jc)
            text_with_fields(p, text, doc)
        if tracker is not None:
            tracker.insert_paragraph(p)
        changed += 1
    return changed


def section_of(doc: Doc, el: Any) -> Any:
    """The w:sectPr of the section a body block belongs to (the next section end, or the body's last one)."""
    from _docx import PPR, SECTPR

    top = el
    while top.getparent() is not None and top.getparent() is not doc.body:
        top = top.getparent()
    nxt = top
    while nxt is not None:
        ppr = nxt.find(PPR) if nxt.tag == qn("w:p") else None
        if ppr is not None and ppr.find(SECTPR) is not None:
            return ppr.find(SECTPR)
        nxt = nxt.getnext()
    return doc.body.find(SECTPR)


def section_text_width(sect: Any) -> float:
    """Text width of a section in points."""
    from _docx import twips

    pg = sect.find(qn("w:pgSz"))
    mar = sect.find(qn("w:pgMar"))
    return twips(wattr(pg, "w"), 612) - twips(wattr(mar, "left"), 72) - twips(wattr(mar, "right"), 72)


SECTPR_ORDER = ("w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz", "w:pgMar", "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")


def _ordered_child(parent: Any, tag: str, order: Sequence[str]) -> Any:
    el = parent.find(qn(tag))
    if el is not None:
        return el
    idx = order.index(tag)
    el = new_el(tag)
    for later in order[idx + 1 :]:
        ref = parent.find(qn(later))
        if ref is not None:
            ref.addprevious(el)
            return el
    parent.append(el)
    return el


def _set_page(sect: Any, width: float | None, height: float | None, margins: dict[str, float] | None, orientation: str | None, header: float | None = None, footer: float | None = None, columns: int | None = None) -> None:
    pg = _ordered_child(sect, "w:pgSz", SECTPR_ORDER)
    from _docx import twips

    cur_w, cur_h = twips(wattr(pg, "w"), 595.3), twips(wattr(pg, "h"), 841.9)
    w, h = width or cur_w, height or cur_h
    if orientation:
        if (orientation == "landscape") != (w > h):
            w, h = h, w
        pg.set(_W + "orient", orientation)
    elif w > h:
        pg.set(_W + "orient", "landscape")
    pg.set(_W + "w", str(int(round(w * 20))))
    pg.set(_W + "h", str(int(round(h * 20))))
    mar = _ordered_child(sect, "w:pgMar", SECTPR_ORDER)
    defaults = {"top": 1440, "right": 1440, "bottom": 1440, "left": 1440, "header": 708, "footer": 708, "gutter": 0}
    for k, v in defaults.items():
        if mar.get(_W + k) is None:
            mar.set(_W + k, str(v))
    for k, v in (margins or {}).items():
        mar.set(_W + k, str(int(round(v * 20))))
    if header is not None:
        mar.set(_W + "header", str(int(round(header * 20))))
    if footer is not None:
        mar.set(_W + "footer", str(int(round(footer * 20))))
    if columns:
        cols = _ordered_child(sect, "w:cols", SECTPR_ORDER)
        cols.set(_W + "num", str(columns))
        cols.set(_W + "space", cols.get(_W + "space") or "720")


def page_setup(doc: Doc, *, size: str | None = None, orientation: str | None = None, margins: str | dict[str, Any] | None = None, sections: Sequence[int] | None = None, header_distance: str | None = None, footer_distance: str | None = None, columns: int | None = None) -> int:
    sects = doc.section_elements()
    targets = range(len(sects)) if not sections else [i - 1 for i in sections]
    wh = parse_page_size(size) if size else (None, None)
    if isinstance(margins, dict):
        mg = {k: parse_length(v) for k, v in margins.items() if k in ("top", "right", "bottom", "left", "gutter")}
    else:
        mg = parse_margins(margins) if margins else None
    if orientation and orientation not in ("portrait", "landscape"):
        raise UsageError("orientation must be portrait or landscape")
    n = 0
    for i in targets:
        if i < 0 or i >= len(sects):
            raise SkillError(f"section {i + 1} does not exist (the document has {len(sects)})")
        _set_page(sects[i], wh[0], wh[1], mg, orientation, parse_length(header_distance) if header_distance else None, parse_length(footer_distance) if footer_distance else None, columns)
        n += 1
    return n


# ── title page and table of contents ────────────────────────────────────

TITLE_STYLES = ("Title", "Subtitle", "Author", "Date", "Abstract", "AbstractTitle")


def make_title_page(doc: Doc) -> bool:
    """Puts the title block (Title, Subtitle, Author, Date, Abstract) on its own page, without header or footer."""
    from _docx import iter_blocks, style_id

    blocks = list(iter_blocks(doc.body))
    last = None
    for el in blocks:
        if el.tag == P and style_id(el) in TITLE_STYLES:
            last = el
        else:
            break
    if last is None:
        return False
    first = blocks[0]
    ppr = first.find(PPR)
    if ppr is not None:
        sp = ppr.find(qn("w:spacing"))
        if sp is None:
            sp = new_el("w:spacing")
            ps = ppr.find(qn("w:pStyle"))
            (ps.addnext(sp) if ps is not None else ppr.insert(0, sp))
        sp.set(_W + "before", "4000")
    br = new_el("w:p")
    r = sub_el(br, "w:r")
    sub_el(r, "w:br", type="page")
    last.addnext(br)
    sect = doc.section_elements()[0]
    tp = _ordered_child(sect, "w:titlePg", SECTPR_ORDER)
    tp.attrib.pop(_W + "val", None)
    # Empty first-page header and footer so the cover stays clean.
    for kind in ("header", "footer"):
        hf = _hdr_ftr_part(doc, sect, kind, "first")
        el = hf._element
        for child in list(el):
            el.remove(child)
        sub_el(el, "w:p")
    return True


def heading_paragraphs(doc: Doc, depth: int) -> list[tuple[Any, int, str, str | None]]:
    """(paragraph, level, text, bookmark) for headings up to depth, in order."""
    from _docx import heading_level, iter_blocks, para_text

    out = []
    for el in iter_blocks(doc.body):
        if el.tag != P:
            continue
        lvl = heading_level(el, doc.styles)
        if not lvl or lvl > depth:
            continue
        sid = el.find(PPR)
        if sid is not None and sid.find(qn("w:pStyle")) is not None and wattr(sid.find(qn("w:pStyle")), "val") == "TOCHeading":
            continue
        text = para_text(el).strip()
        if not text:
            continue
        bm = None
        prev = el.getprevious()
        while prev is not None and prev.tag == qn("w:bookmarkStart"):
            if not (wattr(prev, "name") or "").startswith("_GoBack"):
                bm = wattr(prev, "name")
                break
            prev = prev.getprevious()
        if bm is None:
            inner = el.find(qn("w:bookmarkStart"))
            bm = wattr(inner, "name") if inner is not None else None
        out.append((el, lvl, text, bm))
    return out


def fill_toc(doc: Doc, depth: int = 3, pages: dict[int, int] | None = None, title: str | None = "Contents") -> int:
    """Fills every TOC field (or [[toc]] placeholder) with entries linked to the headings. Returns the entry count."""
    from _docx import INSTR, iter_blocks, para_text

    heads = heading_paragraphs(doc, 9)
    # Headings need bookmarks for the links (and for Word's PAGEREF).
    used = {wattr(b, "name") for b in doc.body.iter(qn("w:bookmarkStart"))}
    next_id = max([int(wattr(b, "id") or 0) for b in doc.body.iter(qn("w:bookmarkStart"))] + [0]) + 1
    entries = []
    for i, (el, lvl, text, bm) in enumerate(heads):
        if not bm:
            bm = f"_Toc{100000 + i}"
            while bm in used:
                bm += "x"
            used.add(bm)
            start = new_el("w:bookmarkStart", id=str(next_id), name=bm)
            end = new_el("w:bookmarkEnd", id=str(next_id))
            next_id += 1
            ppr = el.find(PPR)
            (ppr.addnext(start) if ppr is not None else el.insert(0, start))
            el.append(end)
        entries.append((lvl, text, bm, (pages or {}).get(i)))
    all_entries = entries
    entries = [e for e in all_entries if e[0] <= depth]
    placeholders = []
    for el in list(iter_blocks(doc.body)):
        if el.tag == P and para_text(el).strip().startswith(TOC_MARKER):
            placeholders.append(("marker", el))
        elif el.tag == P and any((x.text or "").strip().upper().startswith("TOC") for x in el.iter(INSTR)):
            placeholders.append(("field", el))
    if not placeholders:
        return 0
    styles = doc.styles
    default_entries, default_title = entries, title
    for kind, el in placeholders:
        entries, title = default_entries, default_title
        toc_width = section_text_width(section_of(doc, el))
        parent = el.getparent()
        if kind == "field":
            # Replace pandoc's empty field paragraph; keep its TOC Heading sibling.
            instr = "".join(x.text or "" for x in el.iter(INSTR)).strip()
        else:
            parts = para_text(el).strip().split(":", 2)
            own_depth = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else depth
            if len(parts) > 2:
                title = parts[2] or None
            instr = f'TOC \\o "1-{own_depth}" \\h \\z \\u'
            if own_depth != depth:
                entries = [e for e in all_entries if e[0] <= own_depth]
            if title:
                hp = new_el("w:p")
                hppr = sub_el(hp, "w:pPr")
                sid = styles.find("TOC Heading") or styles.find("Heading 1")
                if sid:
                    sub_el(hppr, "w:pStyle", val=sid)
                _plain_runs(hp, title)
                el.addprevious(hp)
        new_paras = []
        for lvl, text, bm, page in entries or [(1, "(no headings)", None, None)]:
            p = new_el("w:p")
            ppr = sub_el(p, "w:pPr")
            sid = styles.find(f"TOC {lvl}") or styles.find(f"toc {lvl}")
            if sid:
                sub_el(ppr, "w:pStyle", val=sid)
            else:
                sub_el(ppr, "w:ind", left=str((lvl - 1) * 240))
            # The page numbers sit at the right margin of the section holding the TOC.
            set_tab_stops(doc, ppr, [("right", toc_width, "dot")])
            if bm:
                h = sub_el(p, "w:hyperlink", anchor=bm, history="1")
                _plain_runs(h, text)
                _plain_runs(h, "\t")
                for r in field_runs(f"PAGEREF {bm} \\h", str(page) if page else ""):
                    h.append(r)
            else:
                _plain_runs(p, text)
            new_paras.append(p)
        # field begin/separate in the first entry, end after the last
        begin = field_runs(instr)
        first = new_paras[0]
        ppr0 = first.find(PPR)
        anchor = ppr0
        for r in begin[:3]:
            anchor.addnext(r)
            anchor = r
        end_run = begin[4]
        new_paras[-1].append(end_run)
        for p in new_paras:
            el.addprevious(p)
        parent.remove(el)
    from _docx import set_setting

    set_setting(doc, "updateFields", val="true")
    return len(default_entries)


def _pdf_texts(pdf: Path) -> list[str]:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf))
    try:
        return [re.sub(r"\s+", " ", doc[i].get_textpage().get_text_range()).lower() for i in range(len(doc))]
    finally:
        doc.close()


def heading_pages_from_pdf(texts: list[str], headings: list[str], skip: set[int] = frozenset()) -> dict[int, int]:
    """Page numbers (1-based) of headings found in the pages' text, searched in document order. On the pages in
    `skip` (the table of contents) a heading counts only when it appears again after its TOC entry."""
    out: dict[int, int] = {}
    page = 0
    for i, h in enumerate(headings):
        needle = re.sub(r"\s+", " ", h).strip().lower()[:60]
        if not needle:
            continue
        for j in range(page, len(texts)):
            if texts[j].count(needle) >= (2 if j in skip else 1):
                out[i] = j + 1
                page = j
                break
    return out


def _norm_title(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def heading_pages_from_outline(pdf: Path, titles: list[str]) -> dict[int, int]:
    """Page numbers (1-based) of headings read from the PDF's bookmarks (its outline), which both LibreOffice and
    the built-in renderer write at each heading. Bookmarks are matched to the headings in order; a bookmark may
    carry the heading's list number in front ('2.1 Scope') or be cut short."""
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf))
    try:
        marks = []
        for bm in doc.get_toc(max_depth=10):
            dest = bm.get_dest()
            idx = dest.get_index() if dest is not None else None
            if idx is not None and idx >= 0:
                marks.append((_norm_title(bm.get_title()), idx + 1))
    finally:
        doc.close()
    out: dict[int, int] = {}
    k = 0
    for i, t in enumerate(titles):
        want = _norm_title(t)
        if not want:
            continue
        for j in range(k, min(len(marks), k + 60)):
            have = marks[j][0]
            if have == want or have.endswith(" " + want) or (len(have) >= 20 and want.startswith(have)) or (len(want) >= 20 and have.startswith(want)):
                out[i] = marks[j][1]
                k = j + 1
                break
    return out


def toc_pages_via_render(doc_path: Path, depth: int) -> tuple[dict[str, int], str, int]:
    """Renders a document (LibreOffice, else the built-in renderer): (heading bookmark -> page, engine, pages).

    Pages come from the PDF's bookmarks (exact: each points at its heading's page); searching the page text for
    the heading titles is only a fallback for a PDF without bookmarks."""
    from _docx import load
    from _render import find_soffice, office_convert

    doc = load(doc_path)
    heads = heading_paragraphs(doc, 9)
    titles = [h[2] for h in heads]
    tmp = Path(tempfile.mkdtemp(prefix="desk-toc-"))
    lo_dir = None
    try:
        if find_soffice():
            pdf = office_convert(doc_path, "pdf")
            lo_dir = pdf.parent
            engine = "LibreOffice"
        else:
            from _typst_render import render_pdf

            pdf = tmp / "toc.pdf"
            render_pdf(doc, pdf)
            engine = "built-in renderer"
        pages = heading_pages_from_outline(pdf, titles)
        import pypdfium2 as pdfium

        _pdf = pdfium.PdfDocument(str(pdf))
        page_count = len(_pdf)
        _pdf.close()
        if not pages:
            texts = _pdf_texts(pdf)
            pages = heading_pages_from_pdf(texts, titles, _toc_pages(texts, titles))
        # Keyed by the headings' bookmarks, which the TOC's PAGEREF fields name.
        return {heads[i][3]: pg for i, pg in pages.items() if heads[i][3]}, engine, page_count
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        if lo_dir is not None:
            shutil.rmtree(lo_dir, ignore_errors=True)


def _toc_pages(texts: list[str], titles: list[str]) -> set[int]:
    """Indexes of the pages holding the table of contents: the first run of pages that list most headings."""
    if len(titles) < 2:
        return set()
    found: set[int] = set()
    for i, text in enumerate(texts[:40]):
        hits = sum(1 for t in titles if t.lower()[:40] in text)
        if hits >= max(2, int(len(titles) * 0.6)) or (found and i - 1 in found and hits >= max(2, len(titles) // 3) and "contents" not in text):
            found.add(i)
        elif found:
            break
    return found


# ── inserting Markdown into an existing document (docx_edit) ────────────

PANDOC_ONLY_STYLES = {"BodyText": None, "FirstParagraph": None, "Compact": None, "BlockText": "Quote"}


def markdown_elements(doc: Doc, md: str, *, resource_paths: Sequence[Path] = ()) -> list[Any]:
    """Converts Markdown into body elements merged into `doc` (styles, numbering, images, footnotes, links),
    appended at the end of the body; the caller moves them where they belong."""
    from docxcompose.composer import Composer

    tmpdir = Path(tempfile.mkdtemp(prefix="desk-mdins-"))
    try:
        ref = tmpdir / "ref.docx"
        doc.flush()
        saved_ct = doc.part._content_type
        from _docx import CT_MAIN

        doc.part._content_type = CT_MAIN["docx"]
        try:
            doc.docx.save(str(ref))
        finally:
            doc.part._content_type = saved_ct
        frag = tmpdir / "frag.docx"
        md_to_docx(md, frag, ref, resource_paths=resource_paths, smart=False)
        sub = Doc(frag)
        _map_pandoc_styles(sub, doc)
        body = doc.body
        keep = list(body)  # holds the proxies alive so identity checks are reliable
        before = set(keep)
        comp = Composer(doc.docx)
        comp.restart_numbering = False
        comp.append(sub.docx)
        added = [x for x in body if x not in before and x.tag != SECTPR]
        # docxcompose parses footnotes itself; refresh our cached copy.
        doc._generic.clear()
        return added
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def _map_pandoc_styles(sub: Doc, target: Doc) -> None:
    """Pandoc-only paragraph styles become the target's Normal (or Quote) when the target lacks them."""
    tstyles = target.styles
    normal = tstyles.defaults.get("paragraph") or tstyles.find("Normal")
    for p in sub.body.iter(P):
        ppr = p.find(PPR)
        if ppr is None:
            continue
        ps = ppr.find(qn("w:pStyle"))
        if ps is None:
            continue
        sid = wattr(ps, "val")
        if sid in PANDOC_ONLY_STYLES and tstyles.get(sid) is None:
            alt = PANDOC_ONLY_STYLES[sid]
            new = tstyles.find(alt) if alt else None
            new = new or normal
            if new:
                ps.set(_W + "val", new)
            else:
                ppr.remove(ps)
        if sid == "Table" and tstyles.get("Table") is None:
            pass
    for ts in sub.body.iter(qn("w:tblStyle")):
        if wattr(ts, "val") == "Table" and tstyles.get("Table") is None and tstyles.find("Table Grid"):
            ts.set(_W + "val", tstyles.find("Table Grid"))


def plain_paragraph(doc: Doc, text: str, style: str | None = None, like: Any = None) -> Any:
    """A new paragraph with text (tabs and line breaks honoured), styled by name or like another paragraph."""
    p = new_el("w:p")
    ppr = None
    if like is not None and like.find(PPR) is not None:
        ppr = copy.deepcopy(like.find(PPR))
        for x in ppr.findall(SECTPR):
            ppr.remove(x)
        for tag in ("w:rPr",):
            for x in ppr.findall(qn(tag)):
                ppr.remove(x)
        p.append(ppr)
    if style:
        sid = doc.styles.find(style, "paragraph")
        if not sid:
            raise SkillError(f"style '{style}' is not defined in this document (docx_info.py lists the styles in use)")
        if ppr is None:
            ppr = sub_el(p, "w:pPr")
        ps = ppr.find(qn("w:pStyle"))
        if ps is None:
            ps = new_el("w:pStyle")
            ppr.insert(0, ps)
        ps.set(_W + "val", sid)
    rpr = None
    if like is not None:
        r0 = like.find(f".//{R}")
        if r0 is not None and r0.find(qn("w:rPr")) is not None:
            rpr = copy.deepcopy(r0.find(qn("w:rPr")))
    _plain_runs(p, text, rpr)
    return p


PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]


def ppr_child(ppr: Any, name: str) -> Any:
    """The w:pPr child `name`, created at its schema position when missing."""
    from _docx import new_el

    el = ppr.find(qn("w:" + name))
    if el is not None:
        return el
    el = new_el("w:" + name)
    idx = PPR_ORDER.index(name)
    for c in ppr:
        cname = c.tag.split("}")[-1]
        if cname in PPR_ORDER and PPR_ORDER.index(cname) > idx:
            c.addprevious(el)
            return el
    ppr.append(el)
    return el


def set_tab_stops(doc: Doc, ppr: Any, stops: Sequence[tuple[str, float, str | None]]) -> None:
    """Gives a paragraph exactly these tab stops ((kind, position in pt, leader)), clearing the ones its style has."""
    from _docx import sub_el

    sid = None
    ps = ppr.find(qn("w:pStyle"))
    if ps is not None:
        sid = wattr(ps, "val")
    inherited = [t.get("pos") for t in (doc.styles.para_style_ppr(sid).get("tabs") or []) if t.get("pos")]
    tabs = ppr_child(ppr, "tabs")
    for old in list(tabs):
        tabs.remove(old)
    ours = {str(int(round(pos * 20))): (kind, leader) for kind, pos, leader in stops}
    entries = [(int(float(pos)), pos, "clear", None) for pos in inherited if pos not in ours]
    entries += [(int(pos), pos, kind, leader) for pos, (kind, leader) in ours.items()]
    for _, pos, kind, leader in sorted(entries):
        attrs = {"val": kind, "pos": pos}
        if leader:
            attrs["leader"] = leader
        sub_el(tabs, "w:tab", **attrs)


def set_paragraph_format(ppr: Any, op: dict[str, Any]) -> None:
    """Paragraph formatting keys: align, space_before, space_after, line_spacing, indent_left, indent_right,
    first_line_indent, hanging_indent, keep_with_next, keep_together, page_break_before."""
    from _docx import new_el, qn
    from _markdown import parse_length

    order = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]

    def child(name: str) -> Any:
        el = ppr.find(qn("w:" + name))
        if el is not None:
            return el
        el = new_el("w:" + name)
        idx = order.index(name)
        for c in ppr:
            cname = c.tag.split("}")[-1]
            if cname in order and order.index(cname) > idx:
                c.addprevious(el)
                return el
        ppr.append(el)
        return el

    def tw(v: Any) -> str:
        return str(int(round(parse_length(v, "pt") * 20)))

    if "align" in op:
        a = {"centre": "center", "justify": "both", "justified": "both", "left": "left", "right": "right", "center": "center", "both": "both"}.get(op["align"])
        if not a:
            raise UsageError("align: left, center, right or justify")
        child("jc").set(qn("w:val"), a)
    if "space_before" in op:
        child("spacing").set(qn("w:before"), tw(op["space_before"]))
    if "space_after" in op:
        child("spacing").set(qn("w:after"), tw(op["space_after"]))
    if "line_spacing" in op:
        sp = child("spacing")
        v = op["line_spacing"]
        if isinstance(v, (int, float)) or str(v).replace(".", "", 1).isdigit():
            sp.set(qn("w:line"), str(int(round(float(v) * 240))))
            sp.set(qn("w:lineRule"), "auto")
        else:
            sp.set(qn("w:line"), tw(v))
            sp.set(qn("w:lineRule"), "exact")
    for key, attr in (("indent_left", "left"), ("indent_right", "right"), ("first_line_indent", "firstLine"), ("hanging_indent", "hanging")):
        if key in op:
            ind = child("ind")
            ind.set(qn("w:" + attr), tw(op[key]))
            if attr == "firstLine":
                ind.attrib.pop(qn("w:hanging"), None)
            if attr == "hanging":
                ind.attrib.pop(qn("w:firstLine"), None)
    for key, name in (("keep_with_next", "keepNext"), ("keep_together", "keepLines"), ("page_break_before", "pageBreakBefore")):
        if key in op:
            el = child(name)
            if not op[key]:
                el.set(qn("w:val"), "0")
            else:
                el.attrib.pop(qn("w:val"), None)


# ── tables made by docx_create ──────────────────────────────────────────

TBLPR_ORDER = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
NO_SPACE_STYLES = ("heading", "caption", "title", "subtitle", "toc")


def autosize_table(doc: Doc, tbl: Any, width: float | None = None) -> bool:
    """Gives a table column widths that follow its content and fill the text width (fixed layout, so Word and
    LibreOffice keep them): each column gets at least its longest word, and the rest of the width goes to the
    columns with the longest lines. Tables whose cells already have widths (a JSON spec's "widths") keep them.
    True when changed."""
    from _docx import TBLPR, TCPR, cell_text, table_grid
    from _styles import flatten

    tblpr = tbl.find(TBLPR)
    if tblpr is None:
        tblpr = new_el("w:tblPr")
        tbl.insert(0, tblpr)
    for w in tbl.iter(qn("w:tcW")):
        if (wattr(w, "w") or "0") not in ("0", "") and wattr(w, "type") in ("dxa", None):
            return False
    grid = table_grid(tbl)
    ncols = max((c["col"] + c["colspan"] for row in grid for c in row), default=0)
    if not ncols:
        return False
    total = width or section_text_width(section_of(doc, tbl))
    sz = (doc.styles.doc_rpr.get("sz") or {}).get("val")
    size = (float(sz) / 2 if sz and str(sz).replace(".", "", 1).isdigit() else 11.0) - 0.5
    cw = 0.47 * size  # an average character of a proportional font, in points
    mar = flatten(tblpr.find(qn("w:tblCellMar")))
    pad = sum(float((mar.get(side) or {}).get("w") or 108) for side in ("left", "right")) / 20 + 3
    mins, maxs = [0.0] * ncols, [0.0] * ncols
    for ri, row in enumerate(grid):
        for c in row:
            if c["hidden"] or c["colspan"] > 1:
                continue
            lines = _measured(cell_text(c["tc"])).split("\n") or [""]
            k = 1.1 if ri == 0 else 1.0  # header text is bold
            col = c["col"]
            # A word must never break, whatever the font's real metrics (capitals, a wider substitute): be generous.
            mins[col] = max(mins[col], max((len(w) for ln in lines for w in ln.split()), default=0) * 0.6 * size * k + pad)
            maxs[col] = max(maxs[col], max(len(ln) for ln in lines) * cw * k + pad)
    mins = [max(24.0, m) for m in mins]
    maxs = [max(a, b) for a, b in zip(mins, maxs)]
    if sum(maxs) <= total:
        extra = total - sum(maxs)
        widths = [m + extra * m / sum(maxs) for m in maxs]
    elif sum(mins) >= total:
        widths = [total * m / sum(mins) for m in mins]
    else:
        room = total - sum(mins)
        want = sum(b - a for a, b in zip(mins, maxs)) or 1.0
        widths = [a + room * (b - a) / want for a, b in zip(mins, maxs)]
    tw = [int(round(w * 20)) for w in widths]
    tw[-1] += int(round(total * 20)) - sum(tw)
    tblw = _ordered_child(tblpr, "w:tblW", TBLPR_ORDER)
    tblw.set(_W + "type", "dxa")
    tblw.set(_W + "w", str(int(round(total * 20))))
    _ordered_child(tblpr, "w:tblLayout", TBLPR_ORDER).set(_W + "type", "fixed")
    tg = tbl.find(qn("w:tblGrid"))
    if tg is None:
        tg = new_el("w:tblGrid")
        tblpr.addnext(tg)
    for old in list(tg):
        tg.remove(old)
    for w in tw:
        sub_el(tg, "w:gridCol", w=str(w))
    for row in grid:
        for c in row:
            tc = c["tc"]
            tcpr = tc.find(TCPR)
            if tcpr is None:
                tcpr = new_el("w:tcPr")
                tc.insert(0, tcpr)
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = new_el("w:tcW")
                cnf = tcpr.find(qn("w:cnfStyle"))
                (cnf.addnext(tcw) if cnf is not None else tcpr.insert(0, tcw))
            tcw.set(_W + "type", "dxa")
            tcw.set(_W + "w", str(sum(tw[c["col"] : c["col"] + c["colspan"]])))
    return True


_CONTROL_TAG = re.compile(r"\{\{\s*(?:[#/^!]|else\b)[^}]*\}\}")
_CONTROL_LINE = re.compile(r"^\s*(?:\{\{\s*(?:[#/^!]|else\b)[^}]*\}\}\s*)+$")


def _measured(text: str) -> str:
    """Cell text as it will read: template control tags ({{#each}}, {{/if}}…) take no room and a placeholder
    about as much as a short value."""
    if "{{" not in text:
        return text

    def guess(m: re.Match[str]) -> str:
        expr = m.group(1).strip()
        name, _, filters = expr.partition("|")
        name = name.strip().lstrip("@./").split(".")[-1]
        n = 3 if name in ("number", "index", "first", "last") else max(4, len(name)) + (4 if filters.strip() else 0)
        return "0" * n

    return re.sub(r"\{\{([^}]*)\}\}", guess, _CONTROL_TAG.sub("", text))


def space_after_table(doc: Doc, tbl: Any, space_pt: float = 9.0) -> int:
    """Word has no spacing below a table: the next body paragraph gets some space before it (unless it is a heading
    or caption, which have their own), so text does not sit on the table's bottom rule. In a template, the first
    paragraph of each {{#if}} / {{else}} branch that follows gets it. Returns how many paragraphs changed."""
    from _docx import para_text, style_id

    targets: list[Any] = []
    need = True
    nxt = tbl.getnext()
    while nxt is not None:
        if nxt.tag != P:
            if nxt.tag in (qn("w:tbl"), SECTPR) or targets:
                break
            nxt = nxt.getnext()
            continue
        text = para_text(nxt)
        if _CONTROL_LINE.match(text):
            if re.match(r"\s*\{\{\s*else\b", text):
                need = True
        elif need:
            targets.append(nxt)
            need = False
        else:
            break
        nxt = nxt.getnext()
    n = 0
    for p in targets:
        name = (doc.styles.name(style_id(p)) or "").lower()
        if any(k in name for k in NO_SPACE_STYLES):
            continue
        ppr = p.find(PPR)
        if ppr is None:
            ppr = new_el("w:pPr")
            p.insert(0, ppr)
        sp = ppr_child(ppr, "spacing")
        cur = sp.get(_W + "before")
        if cur is not None and cur.isdigit() and int(cur) >= space_pt * 20:
            continue
        sp.set(_W + "before", str(int(space_pt * 20)))
        sp.attrib.pop(_W + "beforeAutospacing", None)
        n += 1
    return n


def finish_tables(doc: Doc, tables: Sequence[Any] | None = None) -> int:
    """autosize_table + space_after_table for the body's top-level tables (or the given ones). Returns how many."""
    from _docx import TBL, iter_blocks

    n = 0
    for tbl in tables if tables is not None else [b for b in iter_blocks(doc.body) if b.tag == TBL]:
        if autosize_table(doc, tbl):
            n += 1
        space_after_table(doc, tbl)
    return n
