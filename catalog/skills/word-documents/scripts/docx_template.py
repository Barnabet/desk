#!/usr/bin/env python3
"""Fill a Word template with JSON data: {{placeholders}}, loops, conditionals, Word merge fields and content controls.

Placeholders keep the formatting of their first character, even when Word split them across runs.
  {{client.name}}  {{items.0.price}}  {{name | upper}}  {{total | number:2}}  {{date | date:"%d %B %Y"}}
  {{amount | currency:"€"}}  {{note | default:"none"}}  {{tags | join:", "}}  {{logo | image:4cm}}
Loops and conditionals ({{#each}} … {{/each}}, {{#if}} … {{else}} … {{/if}}, {{#unless}} … {{/unless}}):
  - both tags alone in their own paragraphs: the paragraphs and tables between them repeat (or show);
  - tags in different cells of a table row (or rows): those table rows repeat;
  - both tags inside one paragraph: that part of the paragraph repeats.
  Inside a loop: {{this}}, {{name}} (a field of the item, falling back to outer data), {{@index}} (0-based),
  {{@number}} (1-based), {{@first}}, {{@last}}. Conditions: {{#if paid}}, {{#if not paid}},
  {{#if status == "late"}}, {{#if total > 1000}}, and/or.
Word merge fields (MERGEFIELD Name) and content controls whose tag or title matches a data key are filled too.
Unfilled placeholders are reported (and kept as-is unless --blank-missing); --strict fails on any.

Examples:
  python3 scripts/docx_template.py invoice-template.docx invoice-0042.docx --data invoice.json
  python3 scripts/docx_template.py letter.dotx letter.docx --data '{"name": "Ada", "date": "2026-09-24"}'
  python3 scripts/docx_template.py template.docx --fields          # list the fields and blocks the template uses
"""

from __future__ import annotations

import copy
import datetime as dt
import re
from pathlib import Path
from typing import Any

from _common import SkillError, UsageError, add_format, emit, parser, run_main

TAG = re.compile(r"\{\{\s*(.*?)\s*\}\}", re.S)
CONTROL = re.compile(r"^(#each|#if|#unless|else|/each|/if|/unless)\b\s*(.*)$", re.S)
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


class Missing:
    def __init__(self, path: str) -> None:
        self.path = path

    def __bool__(self) -> bool:
        return False


def main() -> int:
    ap = parser(__doc__.split("\n\n")[0], __doc__[__doc__.index("Placeholders keep") :])
    ap.add_argument("template", help=".docx/.dotx/.docm template")
    ap.add_argument("output", nargs="?", help="the filled document to write (.docx)")
    ap.add_argument("--data", help="JSON data: inline, a .json file, or - for stdin")
    ap.add_argument("--fields", action="store_true", help="list the template's fields, loops and conditions, then exit")
    ap.add_argument("--strict", action="store_true", help="fail when a placeholder has no value")
    ap.add_argument("--blank-missing", action="store_true", help="replace placeholders that have no value with nothing")
    ap.add_argument("--force", action="store_true", help="overwrite an existing output")
    add_format(ap)
    args = ap.parse_args()

    from _common import load_json_arg, output_path
    from _docx import load

    doc = load(args.template)
    if args.fields:
        emit(scan_fields(doc), args.format, render_fields)
        return 0
    if not args.output or args.data is None:
        raise UsageError("give an output path and --data (or --fields to list what the template needs)")
    out = Path(args.output)
    output_path(out, [args.template], force=args.force)
    data = load_json_arg(args.data)
    if not isinstance(data, dict):
        raise UsageError("--data must be a JSON object")
    base = Path(args.data).parent if not str(args.data).strip().startswith("{") and args.data != "-" and Path(args.data).exists() else Path.cwd()
    eng = Engine(doc, data, blank_missing=args.blank_missing, base_dir=base)
    eng.run()
    if args.strict and eng.unfilled:
        raise SkillError("no value for: " + ", ".join(sorted(eng.unfilled)) + " (nothing written)")
    target = out if out.suffix.lower() in (".docx", ".docm", ".dotx", ".dotm") else out.with_suffix(".docx")
    warnings = doc.save(target, inputs=(doc.path,), force=args.force)
    report = {
        "output": str(target),
        "filled": eng.filled,
        "loops": eng.loops,
        "conditions": eng.conditions,
        "merge_fields": eng.merge_fields,
        "content_controls": eng.controls,
        "unfilled": dict(sorted(eng.unfilled.items())),
        "warnings": warnings + eng.warnings,
    }
    emit(report, args.format, render_report)
    return 0


def render_report(r: dict[str, Any]) -> str:
    lines = [f"wrote {r['output']}: {r['filled']} placeholders filled, {r['loops']} loops expanded, {r['conditions']} conditions, {r['merge_fields']} merge fields, {r['content_controls']} content controls"]
    if r["unfilled"]:
        lines.append("unfilled (no value in the data): " + ", ".join(f"{k} ×{v}" for k, v in r["unfilled"].items()))
    for w in r["warnings"]:
        lines.append(f"warning: {w}")
    lines.append("Check it: docx_read.py for the text, docx_render.py + view_image for the layout.")
    return "\n".join(lines)


# ── data access ─────────────────────────────────────────────────────────


class Scope:
    def __init__(self, value: Any, parent: "Scope | None" = None, meta: dict[str, Any] | None = None) -> None:
        self.value = value
        self.parent = parent
        self.meta = meta or {}

    def child(self, value: Any, meta: dict[str, Any]) -> "Scope":
        return Scope(value, self, meta)

    def lookup(self, path: str) -> Any:
        path = path.strip()
        if path in ("this", "."):
            return self.value
        if path.startswith("@"):
            s: Scope | None = self
            while s is not None:
                if path[1:] in s.meta:
                    return s.meta[path[1:]]
                s = s.parent
            return Missing(path)
        if path.startswith("this."):
            return _dig(self.value, path[5:].split("."), path)
        if path.startswith("../") and self.parent is not None:
            return self.parent.lookup(path[3:])
        parts = [p for p in re.split(r"\.|\[(\d+)\]", path) if p]
        s = self
        while s is not None:
            v = _dig(s.value, parts, path)
            if not isinstance(v, Missing):
                return v
            s = s.parent
        return Missing(path)


def _dig(value: Any, parts: list[str], path: str) -> Any:
    cur = value
    for p in parts:
        if isinstance(cur, dict):
            if p in cur:
                cur = cur[p]
            else:
                return Missing(path)
        elif isinstance(cur, list) and p.lstrip("-").isdigit():
            i = int(p)
            if -len(cur) <= i < len(cur):
                cur = cur[i]
            else:
                return Missing(path)
        else:
            return Missing(path)
    return cur


def split_filters(expr: str) -> tuple[str, list[tuple[str, list[str]]]]:
    """'total | number:2 | default:"0"' -> ('total', [('number', ['2']), ('default', ['0'])])."""
    parts = _split_outside_quotes(expr, "|")
    head = parts[0].strip()
    filters = []
    for f in parts[1:]:
        bits = _split_outside_quotes(f.strip(), ":")
        name = bits[0].strip()
        args = [_unquote(b.strip()) for b in bits[1:]]
        filters.append((name, args))
    return head, filters


def _split_outside_quotes(s: str, sep: str) -> list[str]:
    out, cur, q = [], [], None
    for ch in s:
        if q:
            cur.append(ch)
            if ch == q:
                q = None
        elif ch in "\"'“”":
            q = {"“": "”"}.get(ch, ch)
            cur.append(ch)
        elif ch == sep:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return out


def _unquote(s: str) -> str:
    if len(s) >= 2 and s[0] in "\"'“" and s[-1] in "\"'”":
        return s[1:-1]
    return s


def to_number(v: Any) -> float:
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace(" ", "").strip())
    except ValueError as e:
        raise SkillError(f"{v!r} is not a number") from e


def to_date(v: Any) -> dt.datetime:
    if isinstance(v, (int, float)):
        return dt.datetime.fromtimestamp(v, tz=dt.timezone.utc)
    s = str(v).strip()
    if s.lower() in ("today", "now"):
        return dt.datetime.now()
    try:
        return dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d %B %Y", "%B %d, %Y"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise SkillError(f"{v!r} is not a date (use ISO like 2026-09-24)")


def format_date(d: dt.datetime, fmt: str) -> str:
    presets = {"long": "%-d %B %Y", "us": "%B %-d, %Y", "short": "%d/%m/%Y", "iso": "%Y-%m-%d", "us-short": "%m/%d/%Y", "month": "%B %Y", "datetime": "%Y-%m-%d %H:%M"}
    fmt = presets.get(fmt, fmt)
    # Portable replacements for %-d / %-m (not supported by Windows strftime) and English month names.
    fmt = fmt.replace("%-d", str(d.day)).replace("%-m", str(d.month)).replace("%-H", str(d.hour))
    fmt = fmt.replace("%B", MONTHS[d.month - 1]).replace("%b", MONTHS[d.month - 1][:3])
    return d.strftime(fmt)


def format_number(v: Any, decimals: str = "", thousands: str = ",", point: str = ".") -> str:
    n = to_number(v)
    d = int(decimals) if decimals.strip().isdigit() else (0 if float(n).is_integer() else 2)
    s = f"{n:,.{d}f}"
    return s.replace(",", "\x00").replace(".", point).replace("\x00", thousands)


def apply_filters(value: Any, filters: list[tuple[str, list[str]]], path: str) -> Any:
    for name, args in filters:
        a0 = args[0] if args else ""
        if name == "default":
            if isinstance(value, Missing) or value in (None, ""):
                value = a0
            continue
        if isinstance(value, Missing):
            return value
        if name == "upper":
            value = str(value).upper()
        elif name == "lower":
            value = str(value).lower()
        elif name == "title":
            value = str(value).title()
        elif name == "capitalize":
            s = str(value)
            value = s[:1].upper() + s[1:]
        elif name == "trim":
            value = str(value).strip()
        elif name == "date":
            value = format_date(to_date(value), a0 or "long")
        elif name == "number":
            value = format_number(value, a0, args[1] if len(args) > 1 else ",", args[2] if len(args) > 2 else ".")
        elif name == "currency":
            sym = a0 or "$"
            num = format_number(value, args[1] if len(args) > 1 else "2")
            value = f"{num} {sym}" if len(args) > 2 and args[2] == "after" else f"{sym}{num}"
        elif name == "percent":
            value = f"{to_number(value) * 100:.{int(a0) if a0.isdigit() else 0}f}%"
        elif name == "round":
            value = round(to_number(value), int(a0) if a0.isdigit() else 0)
            if not a0 or a0 == "0":
                value = int(value)
        elif name == "int":
            value = int(to_number(value))
        elif name == "abs":
            value = abs(to_number(value))
        elif name == "join":
            value = (a0 if args else ", ").join(str(x) for x in value) if isinstance(value, list) else str(value)
        elif name in ("count", "length"):
            value = len(value) if hasattr(value, "__len__") else 0
        elif name == "first":
            value = value[0] if isinstance(value, list) and value else ""
        elif name == "last":
            value = value[-1] if isinstance(value, list) and value else ""
        elif name == "sum":
            value = sum(to_number(x) for x in value) if isinstance(value, list) else to_number(value)
        elif name == "yesno":
            opts = (a0 or "yes,no").split(",")
            value = opts[0] if value else (opts[1] if len(opts) > 1 else "")
        elif name == "truncate":
            n = int(a0) if a0.isdigit() else 50
            s = str(value)
            value = s if len(s) <= n else s[: max(0, n - 1)] + "…"
        elif name == "replace":
            value = str(value).replace(a0, args[1] if len(args) > 1 else "")
        elif name == "image":
            value = ImageValue(value, a0, args[1] if len(args) > 1 else "")
        else:
            raise SkillError(f"unknown filter '{name}' in {{{{{path}}}}}")
    return value


class ImageValue:
    def __init__(self, path: Any, width: str, height: str) -> None:
        self.path = str(path)
        self.width = width
        self.height = height


def to_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, list):
        return ", ".join(to_text(x) for x in v)
    if isinstance(v, dict):
        import json

        return json.dumps(v, ensure_ascii=False)
    return str(v)


# ── conditions ──────────────────────────────────────────────────────────


def evaluate(cond: str, scope: Scope) -> bool:
    cond = cond.strip()
    for sep in (" or ", " || "):
        if sep in cond:
            return any(evaluate(c, scope) for c in cond.split(sep))
    for sep in (" and ", " && "):
        if sep in cond:
            return all(evaluate(c, scope) for c in cond.split(sep))
    if cond.startswith("not ") or cond.startswith("!"):
        return not evaluate(cond[4:] if cond.startswith("not ") else cond[1:], scope)
    m = re.match(r"^(.+?)\s*(==|!=|>=|<=|>|<|contains)\s*(.+)$", cond)
    if m:
        left = operand(m.group(1), scope)
        right = operand(m.group(3), scope)
        op = m.group(2)
        if isinstance(left, Missing):
            return op == "!="
        try:
            if op in (">", "<", ">=", "<="):
                l, r = to_number(left), to_number(right)
                return {">": l > r, "<": l < r, ">=": l >= r, "<=": l <= r}[op]
        except SkillError:
            l, r = str(left), str(right)
            return {">": l > r, "<": l < r, ">=": l >= r, "<=": l <= r}[op]
        if op == "contains":
            return str(right) in (left if isinstance(left, (list, str)) else str(left))
        eq = to_text(left) == to_text(right)
        return eq if op == "==" else not eq
    v = operand(cond, scope)
    if isinstance(v, (list, dict, str)):
        return len(v) > 0
    return bool(v)


def operand(s: str, scope: Scope) -> Any:
    s = s.strip()
    if len(s) >= 2 and s[0] in "\"'“" and s[-1] in "\"'”":
        return s[1:-1]
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        return float(s)
    if s in ("true", "false"):
        return s == "true"
    if s in ("null", "none"):
        return None
    head, filters = split_filters(s)
    return apply_filters(scope.lookup(head), filters, s)


# ── the engine ──────────────────────────────────────────────────────────


class Engine:
    def __init__(self, doc: Any, data: dict[str, Any], blank_missing: bool = False, base_dir: Path | None = None) -> None:
        self.doc = doc
        self.root = Scope(data)
        self.blank_missing = blank_missing
        self.base_dir = base_dir or Path.cwd()
        self.filled = 0
        self.loops = 0
        self.conditions = 0
        self.merge_fields = 0
        self.controls = 0
        self.unfilled: dict[str, int] = {}
        self.warnings: list[str] = []

    def run(self) -> None:

        for scope, root, part in list(self.doc.stories(include_comments=False)):
            container = root
            self.container(container, self.root)
        self.content_controls()
        self.properties()

    # containers: body, cells, text boxes, headers …
    def container(self, parent: Any, scope: Scope) -> None:
        from _docx import P, SDT, SDTCONTENT, TBL, TXBX

        kids = [c for c in parent if c.tag in (P, TBL, SDT)]
        i = 0
        while i < len(kids):
            el = kids[i]
            if el.tag == P:
                ctl = self._block_marker(el)
                if ctl and ctl[0] in ("#each", "#if", "#unless"):
                    j, else_i = self._match(kids, i)
                    if j is None:
                        self.warnings.append(f"no closing tag for {{{{{ctl[0]} {ctl[1]}}}}}")
                        self.paragraph(el, scope)
                        i += 1
                        continue
                    self._expand_blocks(kids, i, j, else_i, ctl, scope)
                    i = j + 1
                    continue
                self.paragraph(el, scope)
                for tb in el.iter(TXBX):
                    self.container(tb, scope)
            elif el.tag == TBL:
                self.table(el, scope)
            else:
                content = el.find(SDTCONTENT)
                if content is not None:
                    self.container(content, scope)
            i += 1

    def _block_marker(self, p: Any) -> tuple[str, str] | None:
        from _runs import CharMap

        text = CharMap(p).text.strip()
        m = TAG.fullmatch(text)
        # Exactly one tag alone in the paragraph ("{{#if a}}x{{/if}}" is an inline conditional, not a block marker).
        if not m or "}}" in m.group(1) or "{{" in m.group(1):
            return None
        c = CONTROL.match(m.group(1).strip())
        return (c.group(1), c.group(2).strip()) if c else None

    def _match(self, kids: list[Any], i: int) -> tuple[int | None, int | None]:
        from _docx import P

        depth = 0
        else_i = None
        for j in range(i, len(kids)):
            if kids[j].tag != P:
                continue
            ctl = self._block_marker(kids[j])
            if not ctl:
                continue
            if ctl[0] in ("#each", "#if", "#unless"):
                depth += 1
            elif ctl[0].startswith("/"):
                depth -= 1
                if depth == 0:
                    return j, else_i
            elif ctl[0] == "else" and depth == 1:
                else_i = j
        return None, None

    def _expand_blocks(self, kids: list[Any], i: int, j: int, else_i: int | None, ctl: tuple[str, str], scope: Scope) -> None:
        start, end = kids[i], kids[j]
        parent = start.getparent()
        first_part = kids[i + 1 : else_i if else_i is not None else j]
        second_part = kids[else_i + 1 : j] if else_i is not None else []
        if ctl[0] == "#each":
            items = self._items(ctl[1], scope)
            self.loops += 1
            for k, (item, meta) in enumerate(items):
                copies = [copy.deepcopy(x) for x in first_part]
                for c in copies:
                    start.addprevious(c)
                holder = _Holder(copies)
                self._render_list(copies, scope.child(item, meta))
            chosen: list[Any] = []
        else:
            self.conditions += 1
            ok = evaluate(ctl[1], scope)
            if ctl[0] == "#unless":
                ok = not ok
            chosen = first_part if ok else second_part
            for c in chosen:
                start.addprevious(c)
            self._render_list(chosen, scope)
        for el in kids[i : j + 1]:
            if el not in chosen and el.getparent() is not None:
                el.getparent().remove(el)

    def _render_list(self, elements: list[Any], scope: Scope) -> None:
        """Renders sibling elements that were just inserted (loop copies or a chosen branch)."""
        if not elements:
            return
        parent = elements[0].getparent()
        # Render them as a small container, in place.
        tmp = _Fragment(elements)
        self.container(tmp, scope)

    def _items(self, expr: str, scope: Scope) -> list[tuple[Any, dict[str, Any]]]:
        head, filters = split_filters(expr)
        m = re.match(r"^(\S+)\s+as\s+(\w+)$", head)
        alias = None
        if m:
            head, alias = m.group(1), m.group(2)
        value = apply_filters(scope.lookup(head), filters, expr)
        if isinstance(value, Missing) or value is None:
            self.unfilled[head] = self.unfilled.get(head, 0) + 1
            return []
        if isinstance(value, dict):
            seq = [{"key": k, "value": v} for k, v in value.items()]
        elif isinstance(value, list):
            seq = value
        else:
            seq = [value]
        out = []
        for k, item in enumerate(seq):
            meta = {"index": k, "number": k + 1, "first": k == 0, "last": k == len(seq) - 1}
            if alias:
                item = {alias: item, **(item if isinstance(item, dict) else {})}
            out.append((item, meta))
        return out

    # tables
    def table(self, tbl: Any, scope: Scope) -> None:
        from _docx import P, TR, row_cells

        rows = tbl.findall(TR)
        # control tags per row, with their paragraph, in reading order
        tokens: list[tuple[int, Any, str, str]] = []
        for ri, tr in enumerate(rows):
            for tc in row_cells(tr):
                for p in tc.iter(P):
                    for kind, arg in self._controls_in(p):
                        tokens.append((ri, p, kind, arg))
        pairs = self._pair_tokens(tokens)
        row_pairs = [(o, c, e) for o, c, e in pairs if o[1] is not c[1] and self._cell_of(o[1]) is not self._cell_of(c[1])]
        handled_until = -1
        new_rows_done: set[int] = set()
        for o, c, e in row_pairs:
            if o[0] <= handled_until:
                continue
            a, b = o[0], c[0]
            self._strip_tag(o[1], o[2], o[3])
            self._strip_tag(c[1], c[2], c[3])
            if e is not None:
                self._strip_tag(e[1], e[2], e[3])
            template = rows[a : b + 1]
            if o[2] == "#each":
                self.loops += 1
                for item, meta in self._items(o[3], scope):
                    sub = scope.child(item, meta)
                    for tr in template:
                        cp = copy.deepcopy(tr)
                        rows[a].addprevious(cp)
                        self._render_row(cp, sub)
                for tr in template:
                    tbl.remove(tr)
            else:
                self.conditions += 1
                ok = evaluate(o[3], scope)
                if o[2] == "#unless":
                    ok = not ok
                if e is not None:
                    split = e[0]
                    keep = rows[a:split] if ok else rows[split : b + 1]
                else:
                    keep = template if ok else []
                for tr in template:
                    if tr in keep:
                        self._render_row(tr, scope)
                    else:
                        tbl.remove(tr)
            handled_until = b
            new_rows_done.update(range(a, b + 1))
        for ri, tr in enumerate(rows):
            if ri not in new_rows_done:
                self._render_row(tr, scope)

    def _render_row(self, tr: Any, scope: Scope) -> None:
        from _docx import row_cells

        for tc in row_cells(tr):
            self.container(tc, scope)

    def _cell_of(self, p: Any) -> Any:
        from _docx import TC

        a = p.getparent()
        while a is not None and a.tag != TC:
            a = a.getparent()
        return a

    def _controls_in(self, p: Any) -> list[tuple[str, str]]:
        from _runs import CharMap

        out = []
        for m in TAG.finditer(CharMap(p).text):
            c = CONTROL.match(m.group(1).strip())
            if c:
                out.append((c.group(1), c.group(2).strip()))
        return out

    def _pair_tokens(self, tokens: list[tuple[int, Any, str, str]]) -> list[tuple[Any, Any, Any]]:
        """Top-level open/close pairs (with their else) among a table's control tags."""
        stack: list[tuple[Any, Any]] = []
        pairs = []
        for t in tokens:
            kind = t[2]
            if kind in ("#each", "#if", "#unless"):
                stack.append((t, None))
            elif kind == "else" and stack:
                o, _ = stack[-1]
                stack[-1] = (o, t)
            elif kind.startswith("/") and stack:
                o, e = stack.pop()
                if not stack:
                    pairs.append((o, t, e))
        return pairs

    def _strip_tag(self, p: Any, kind: str, arg: str) -> None:
        from _runs import CharMap, replace_span

        cm = CharMap(p)
        for m in TAG.finditer(cm.text):
            c = CONTROL.match(m.group(1).strip())
            if c and c.group(1) == kind and c.group(2).strip() == arg:
                replace_span(cm, m.start(), m.end(), "")
                return

    # paragraphs
    def paragraph(self, p: Any, scope: Scope) -> None:
        from _runs import CharMap

        for _ in range(50):
            cm = CharMap(p)
            if "{{" not in cm.text:
                break
            if not self._inline_control(p, cm, scope):
                break
        self._placeholders(p, scope)
        self._merge_fields(p, scope)

    def _inline_control(self, p: Any, cm: Any, scope: Scope) -> bool:
        """Expands the last top-level inline #each/#if/#unless of the paragraph. False when there is none."""
        from _runs import CharMap, isolate, replace_span

        toks = []
        for m in TAG.finditer(cm.text):
            c = CONTROL.match(m.group(1).strip())
            if c:
                toks.append((m.start(), m.end(), c.group(1), c.group(2).strip()))
        if not toks:
            return False
        stack: list[Any] = []
        pairs = []
        for t in toks:
            if t[2] in ("#each", "#if", "#unless"):
                stack.append([t, None])
            elif t[2] == "else" and stack:
                stack[-1][1] = t
            elif t[2].startswith("/") and stack:
                o, e = stack.pop()
                if not stack:
                    pairs.append((o, t, e))
        if not pairs:
            # Unbalanced tags inside one paragraph: leave them (reported as unfilled text).
            return False
        o, c, e = pairs[-1]
        if o[2] == "#each":
            self.loops += 1
            items = self._items(o[3], scope)
            runs = isolate(p, o[1], c[0]) if c[0] > o[1] else []
            # Remove the closing tag first: it sits after the body, so the offsets before it stay valid.
            replace_span(CharMap(p), c[0], c[1], "")
            anchor = runs[0] if runs else None
            for item, meta in items:
                frag = _run_fragment(runs)
                self.paragraph(frag, scope.child(item, meta))
                for r in list(frag):
                    if anchor is not None:
                        anchor.addprevious(r)
            for r in runs:
                if r.getparent() is not None:
                    r.getparent().remove(r)
            replace_span(CharMap(p), o[0], o[1], "")
        else:
            self.conditions += 1
            ok = evaluate(o[3], scope)
            if o[2] == "#unless":
                ok = not ok
            if e is not None:
                keep = (o[1], e[0]) if ok else (e[1], c[0])
            else:
                keep = (o[1], c[0]) if ok else (c[0], c[0])
            # delete the tail after keep, then the head before it (right to left keeps offsets valid)
            replace_span(cm, keep[1], c[1], "")
            replace_span(CharMap(p), o[0], keep[0], "")
        return True

    def _remove_first_tag(self, p: Any, kind: str, arg: str) -> None:
        self._strip_tag(p, kind, arg)

    def _placeholders(self, p: Any, scope: Scope) -> None:
        from _runs import CharMap, replace_span

        cm = CharMap(p)
        found = [m for m in TAG.finditer(cm.text) if not CONTROL.match(m.group(1).strip()) and not m.group(1).startswith("!")]
        comments = [m for m in TAG.finditer(cm.text) if m.group(1).startswith("!")]
        for m in sorted(found + comments, key=lambda m: m.start(), reverse=True):
            expr = m.group(1).strip()
            if expr.startswith("!"):
                replace_span(CharMap(p), m.start(), m.end(), "")
                continue
            head, filters = split_filters(expr)
            try:
                value = apply_filters(scope.lookup(head), filters, expr)
            except SkillError as err:
                self.warnings.append(f"{{{{{expr}}}}}: {err}")
                continue
            if isinstance(value, Missing):
                self.unfilled[head] = self.unfilled.get(head, 0) + 1
                if self.blank_missing:
                    replace_span(CharMap(p), m.start(), m.end(), "")
                continue
            if isinstance(value, ImageValue):
                self._image(p, m.start(), m.end(), value)
            else:
                replace_span(CharMap(p), m.start(), m.end(), to_text(value))
            self.filled += 1

    def _image(self, p: Any, start: int, end: int, value: ImageValue) -> None:
        from docx.shared import Emu
        from docx.text.run import Run

        from _markdown import parse_length
        from _runs import CharMap, isolate, replace_span

        path = Path(value.path).expanduser()
        if not path.is_absolute() and not path.exists():
            path = self.base_dir / path
        if not path.exists():
            self.warnings.append(f"image {value.path} not found")
            replace_span(CharMap(p), start, end, "")
            return
        runs = isolate(p, start, end)
        first = runs[0]
        replace_span(CharMap(p), start, end, "")
        from _docx import new_el

        r = new_el("w:r")
        if first.getparent() is not None:
            first.addprevious(r)
        else:
            p.append(r)
        width = Emu(int(parse_length(value.width) * 12700)) if value.width else None
        height = Emu(int(parse_length(value.height) * 12700)) if value.height else None

        class _P:
            part = self.doc.part

        Run(r, _P()).add_picture(str(path), width=width, height=height)

    def _merge_fields(self, p: Any, scope: Scope) -> None:
        """Word mail-merge fields (MERGEFIELD Name) become their values."""
        from _docx import FLDCHAR, FLDSIMPLE, INSTR, R, RPR, T, new_el, qn, set_text

        for fs in list(p.iter(FLDSIMPLE)):
            instr = fs.get(qn("w:instr")) or ""
            name = _mergefield_name(instr)
            if not name:
                continue
            v = scope.lookup(name)
            if isinstance(v, Missing):
                self.unfilled[name] = self.unfilled.get(name, 0) + 1
                continue
            r0 = fs.find(R)
            r = new_el("w:r")
            if r0 is not None and r0.find(RPR) is not None:
                r.append(copy.deepcopy(r0.find(RPR)))
            set_text(_sub(r, "w:t"), to_text(v))
            fs.addprevious(r)
            fs.getparent().remove(fs)
            self.merge_fields += 1
        # complex fields: begin … instr … separate … result … end
        runs = [r for r in p.iter(R)]
        i = 0
        while i < len(runs):
            r = runs[i]
            fc = r.find(FLDCHAR)
            if fc is None or fc.get(qn("w:fldCharType")) != "begin":
                i += 1
                continue
            j = i
            instr = []
            end = None
            depth = 0
            for k in range(i, len(runs)):
                f = runs[k].find(FLDCHAR)
                if f is not None:
                    t = f.get(qn("w:fldCharType"))
                    depth += 1 if t == "begin" else -1 if t == "end" else 0
                    if depth == 0:
                        end = k
                        break
                for it in runs[k].iter(INSTR):
                    instr.append(it.text or "")
            name = _mergefield_name("".join(instr))
            if not name or end is None:
                i += 1
                continue
            v = scope.lookup(name)
            if isinstance(v, Missing):
                self.unfilled[name] = self.unfilled.get(name, 0) + 1
                i = end + 1
                continue
            result_rpr = None
            for k in range(i, end + 1):
                if runs[k].find(T) is not None and runs[k].find(RPR) is not None:
                    result_rpr = runs[k].find(RPR)
            new = new_el("w:r")
            if result_rpr is not None:
                new.append(copy.deepcopy(result_rpr))
            set_text(_sub(new, "w:t"), to_text(v))
            runs[i].addprevious(new)
            for k in range(i, end + 1):
                if runs[k].getparent() is not None:
                    runs[k].getparent().remove(runs[k])
            self.merge_fields += 1
            i = end + 1

    def content_controls(self) -> None:
        """Content controls whose tag or title (alias) names a data key get that value as their text."""
        from _docx import P, R, RPR, SDT, SDTCONTENT, new_el, qn, set_text

        for _, root, _ in self.doc.stories(include_comments=False):
            for sdt in list(root.iter(SDT)):
                pr = sdt.find(qn("w:sdtPr"))
                if pr is None:
                    continue
                keys = [x.get(qn("w:val")) for x in (pr.find(qn("w:tag")), pr.find(qn("w:alias"))) if x is not None and x.get(qn("w:val"))]
                value = None
                for k in keys:
                    v = self.root.lookup(k)
                    if not isinstance(v, Missing):
                        value = v
                        break
                if value is None:
                    continue
                content = sdt.find(SDTCONTENT)
                if content is None:
                    continue
                text = to_text(value)
                runs = list(content.iter(R))
                rpr = runs[0].find(RPR) if runs else None
                paras = [c for c in content if c.tag == P]
                target = paras[0] if paras else content
                for r in runs:
                    if r.getparent() is not None:
                        r.getparent().remove(r)
                for extra in paras[1:]:
                    content.remove(extra)
                r = new_el("w:r")
                if rpr is not None:
                    r.append(copy.deepcopy(rpr))
                set_text(_sub(r, "w:t"), text)
                target.append(r)
                sh = pr.find(qn("w:showingPlcHdr"))
                if sh is not None:
                    pr.remove(sh)
                self.controls += 1

    def properties(self) -> None:
        cp = self.doc.docx.core_properties
        for key in ("title", "subject", "keywords", "comments", "category"):
            v = getattr(cp, key) or ""
            if "{{" in v:
                def sub(m: re.Match[str]) -> str:
                    head, filters = split_filters(m.group(1))
                    val = apply_filters(self.root.lookup(head), filters, m.group(1))
                    return m.group(0) if isinstance(val, Missing) else to_text(val)

                setattr(cp, key, TAG.sub(sub, v))


def _sub(parent: Any, tag: str) -> Any:
    from _docx import sub_el

    return sub_el(parent, tag)


def _mergefield_name(instr: str) -> str | None:
    m = re.match(r'\s*MERGEFIELD\s+"?([^"\\\s]+)"?', instr, re.I)
    return m.group(1) if m else None


class _Fragment:
    """Iterates a list of sibling elements like a container (for rendering loop copies in place)."""

    def __init__(self, elements: list[Any]) -> None:
        self.elements = elements

    def __iter__(self):  # noqa: ANN204
        return iter(list(self.elements))


class _Holder:
    def __init__(self, items: list[Any]) -> None:
        self.items = items


def _run_fragment(runs: list[Any]) -> Any:
    """A detached paragraph holding copies of runs (rendered, then its runs are moved into place)."""
    from _docx import new_el

    p = new_el("w:p")
    for r in runs:
        p.append(copy.deepcopy(r))
    return p


# ── --fields ────────────────────────────────────────────────────────────


def scan_fields(doc: Any) -> dict[str, Any]:
    from _docx import FLDSIMPLE, INSTR, P, SDT, qn
    from _runs import CharMap

    fields: dict[str, dict[str, Any]] = {}
    blocks: list[dict[str, str]] = []
    stack: list[str] = []
    for scope, root, _ in doc.stories(include_comments=False):
        for p in root.iter(P):
            text = CharMap(p).text
            for m in TAG.finditer(text):
                expr = m.group(1).strip()
                c = CONTROL.match(expr)
                if c:
                    kind, arg = c.group(1), c.group(2).strip()
                    if kind.startswith("#"):
                        blocks.append({"kind": kind[1:], "expr": arg, "where": scope, "inside": " > ".join(stack)})
                        stack.append(f"{kind[1:]} {arg}")
                    elif kind.startswith("/") and stack:
                        stack.pop()
                    continue
                if expr.startswith("!"):
                    continue
                head, filters = split_filters(expr)
                key = head
                if key in ("this", ".") or key.startswith("@") or key.startswith("this."):
                    continue
                e = fields.setdefault(key, {"count": 0, "filters": [], "where": [], "inside": []})
                e["count"] += 1
                for f, a in filters:
                    label = f + (":" + ":".join(a) if a else "")
                    if label not in e["filters"]:
                        e["filters"].append(label)
                if scope not in e["where"]:
                    e["where"].append(scope)
                if stack and stack[-1] not in e["inside"]:
                    e["inside"].append(stack[-1])
        for fs in root.iter(FLDSIMPLE):
            n = _mergefield_name(fs.get(qn("w:instr")) or "")
            if n:
                fields.setdefault(n, {"count": 0, "filters": [], "where": [scope], "inside": [], "merge_field": True})["count"] += 1
        buf = "".join(it.text or "" for it in root.iter(INSTR))
        for m in re.finditer(r"MERGEFIELD\s+\"?([^\"\\\s]+)", buf, re.I):
            fields.setdefault(m.group(1), {"count": 0, "filters": [], "where": [scope], "inside": [], "merge_field": True})["count"] += 1
        for sdt in root.iter(SDT):
            pr = sdt.find(qn("w:sdtPr"))
            if pr is None:
                continue
            for t in ("w:tag", "w:alias"):
                x = pr.find(qn(t))
                if x is not None and x.get(qn("w:val")):
                    fields.setdefault(x.get(qn("w:val")), {"count": 0, "filters": [], "where": [scope], "inside": [], "content_control": True})["count"] += 1
                    break
    return {"template": str(doc.path), "fields": fields, "blocks": blocks}


def render_fields(r: dict[str, Any]) -> str:
    lines = [f"{len(r['fields'])} fields, {len(r['blocks'])} loops/conditions in {Path(r['template']).name}"]
    for name, f in r["fields"].items():
        bits = []
        if f.get("filters"):
            bits.append("filters " + ", ".join(f["filters"]))
        if f.get("inside"):
            bits.append("inside " + "; ".join(f["inside"]))
        if f.get("merge_field"):
            bits.append("Word merge field")
        if f.get("content_control"):
            bits.append("content control")
        if f["where"] != ["body"]:
            bits.append("in " + ", ".join(f["where"]))
        lines.append(f"- {name}" + (f" ×{f['count']}" if f["count"] > 1 else "") + (f" ({'; '.join(bits)})" if bits else ""))
    for b in r["blocks"]:
        lines.append(f"- {b['kind']} {b['expr']}" + (f" (inside {b['inside']})" if b["inside"] else "") + (f" in {b['where']}" if b["where"] != "body" else ""))
    return "\n".join(lines)


if __name__ == "__main__":
    run_main(main)
