#!/usr/bin/env python3
"""Read a .docx as Markdown (default), plain text, or JSON structure.

Examples:
  python3 scripts/docx_read.py contract.docx
  python3 scripts/docx_read.py contract.docx --format json    # paragraph indexes for docx_edit.py
  python3 scripts/docx_read.py contract.docx --format text

JSON output: {"properties": {...}, "paragraphs": [{"index", "style", "text"}], "tables": [{"index", "rows"}],
"headers": [...], "footers": [...], "comments": [...]} where paragraph "index" counts body paragraphs
(the numbering docx_edit.py uses).
"""
from __future__ import annotations

import argparse
import json
import re
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def md_runs(p: Paragraph) -> str:
    out = []
    for run in p.runs:
        t = run.text
        if not t:
            continue
        if run.bold and run.italic:
            t = f"***{t}***"
        elif run.bold:
            t = f"**{t}**"
        elif run.italic:
            t = f"*{t}*"
        out.append(t)
    text = "".join(out)
    return re.sub(r"\*\*\*\*|\*\*(\s+)\*\*", lambda m: m.group(1) or "", text)


def md_paragraph(p: Paragraph) -> str:
    style = (p.style.name if p.style is not None else "") or ""
    text = md_runs(p) if p.runs else p.text
    if not text.strip():
        return ""
    if m := re.match(r"Heading (\d)", style):
        return f"{'#' * int(m.group(1))} {p.text.strip()}"
    if style == "Title":
        return f"# {p.text.strip()}"
    if style.startswith("List Bullet"):
        depth = int(style[-1]) - 1 if style[-1].isdigit() else 0
        return f"{'  ' * depth}- {text}"
    if style.startswith("List Number"):
        depth = int(style[-1]) - 1 if style[-1].isdigit() else 0
        return f"{'  ' * depth}1. {text}"
    if "Quote" in style:
        return f"> {text}"
    return text


def md_table(t: Table) -> str:
    rows = [[c.text.replace("\n", " ").replace("|", "\\|").strip() for c in r.cells] for r in t.rows]
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path")
    ap.add_argument("--format", "-f", choices=["markdown", "text", "json"], default="markdown")
    args = ap.parse_args()
    doc = Document(args.path)

    if args.format == "json":
        cp = doc.core_properties
        data = {
            "properties": {k: (str(getattr(cp, k)) if getattr(cp, k) is not None else None) for k in ("title", "author", "subject", "keywords", "created", "modified")},
            "paragraphs": [{"index": i, "style": p.style.name if p.style is not None else None, "text": p.text} for i, p in enumerate(doc.paragraphs)],
            "tables": [{"index": i, "rows": [[c.text for c in r.cells] for r in t.rows]} for i, t in enumerate(doc.tables)],
            "headers": [p.text for s in doc.sections for p in s.header.paragraphs if p.text.strip()],
            "footers": [p.text for s in doc.sections for p in s.footer.paragraphs if p.text.strip()],
            "comments": [{"author": c.author, "text": c.text} for c in getattr(doc, "comments", [])] if hasattr(doc, "comments") else [],
        }
        print(json.dumps(data, ensure_ascii=False, indent=2))
        return 0

    out = ""
    prev_kind = None
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text if args.format == "text" else md_paragraph(item)
            style = (item.style.name if item.style is not None else "") or ""
            kind = "list" if style.startswith("List") else "code" if style == "No Spacing" else "para"
        else:
            text = "\n".join("\t".join(c.text for c in r.cells) for r in item.rows) if args.format == "text" else md_table(item)
            kind = "table"
        if not text.strip():
            continue
        # List items and code lines stay together; everything else is separated by a blank line.
        joiner = "\n" if kind == prev_kind and kind in ("list", "code") else "\n\n"
        out = f"{out}{joiner}{text}" if out else text
        prev_kind = kind
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
