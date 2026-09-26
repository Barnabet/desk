"""The deck builder behind pptx_create (and pptx_edit's add_slide).

Two modes share one code path:
- Theme mode: python-pptx's default template is turned into a designed 16:9 master (theme colours and fonts, text
  styles, layouts re-laid on a grid, decorations on the layouts), then slides are built on it.
- Template mode: the user's .pptx/.potx masters, layouts and placeholders are used as they are; drawn elements
  (cards, timelines, tables, charts) take the theme's colours by role and fit the layout's content area.

Text is fitted by measuring it (Typst, the same layout engine as the built-in renderer) in one batch for the whole
deck: text shrinks within bounds, and bullet lists too long for one slide continue on added slides.
"""

from __future__ import annotations

import copy
import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from _common import SkillError, UsageError
from _mdspec import inline_runs, plain
from _ooxml import EMU_PER_IN, EMU_PER_PT, Color, SlideCtx, TextResolver, contrast_ratio, local, pres_defaults
from _themes import Theme

IN = EMU_PER_IN
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
SLIDE_TYPES = ["title", "section", "bullets", "two-column", "comparison", "image", "image-text", "quote", "kpi", "table", "chart", "timeline", "process", "agenda", "closing", "code", "custom", "blank"]
TYPE_ALIASES = {"big-number": "kpi", "stats": "kpi", "metrics": "kpi", "numbers": "kpi", "image+text": "image-text", "text-image": "image-text", "picture": "image", "steps": "process", "cover": "title", "content": "bullets", "text": "bullets", "list": "bullets", "columns": "two-column", "compare": "comparison", "end": "closing", "thanks": "closing", "divider": "section"}


COMMON_KEYS = {"type", "title", "notes", "hidden", "background", "layout", "id"}
SLIDE_KEYS: dict[str, set[str]] = {
    "title": {"subtitle", "meta", "author", "date", "image"},
    "section": {"subtitle", "text", "number"},
    "closing": {"subtitle", "text", "contact"},
    "bullets": {"bullets", "items", "text", "columns"},
    "two-column": {"left", "right", "split", "ratio", "note"},
    "comparison": {"left", "right", "split", "ratio", "note"},
    "image": {"image", "images", "caption", "fit", "full_bleed"},
    "image-text": {"image", "image_side", "image_width", "fit", "bullets", "items", "text"},
    "quote": {"quote", "text", "attribution"},
    "kpi": {"items", "kpis", "metrics", "colorful"},
    "table": {"columns", "rows", "data", "align", "header", "widths", "font_size", "total", "highlight", "note"},
    "chart": {"chart", "categories", "series", "chart_type", "number_format", "labels", "legend", "colors", "note", "takeaway"},
    "timeline": {"items", "events", "milestones"},
    "process": {"items", "steps"},
    "agenda": {"items", "bullets"},
    "code": {"code", "language", "note"},
    "custom": {"elements", "shapes"},
    "blank": {"elements", "shapes"},
}
CHART_KEYS = {"type", "categories", "series", "number_format", "labels", "legend", "title", "colors"}
SERIES_KEYS = {"name", "values", "points", "x", "y"}
COLUMN_KEYS = {"heading", "title", "heading_color", "bullets", "items", "text", "image", "table", "chart", "code", "fit"}
ITEM_KEYS = {
    "kpi": {"value", "label", "delta", "note", "color", "good"},
    "timeline": {"label", "date", "title", "text", "color"},
    "process": {"label", "title", "text", "color"},
}
BULLET_KEYS = {"text", "title", "level", "bullet", "numbered", "bold", "color", "size", "children", "items", "bullets"}
ELEMENT_KEYS = {"kind", "type", "x", "y", "w", "h", "name", "text", "size", "color", "bold", "align", "anchor", "bullets", "fit", "font",
                "fill", "line", "radius", "width", "path", "image", "chart", "columns", "rows", "header", "align", "widths", "font_size",
                "total", "highlight", "categories", "series", "number_format", "labels", "legend", "colors", "data"}


def _check_keys(d: dict[str, Any], allowed: set[str], what: str) -> None:
    """Unknown or misspelled fields are an error (a typo must never silently drop content)."""
    import difflib

    bad = [k for k in d if k not in allowed and not str(k).startswith("_")]
    if not bad:
        return
    k = bad[0]
    near = difflib.get_close_matches(str(k), sorted(allowed), n=1, cutoff=0.6)
    hint = f" (did you mean \"{near[0]}\"?)" if near else ""
    raise UsageError(f"{what}: unknown field \"{k}\"{hint}; fields: {', '.join(sorted(allowed))}")


def validate_slide(spec: dict[str, Any], kind: str) -> None:
    _check_keys(spec, COMMON_KEYS | SLIDE_KEYS.get(kind, set()) | ({"full_bleed"} if kind in ("custom", "blank") else set()), f"a {kind} slide")
    ch = spec.get("chart")
    if isinstance(ch, dict):
        _check_keys(ch, CHART_KEYS, "\"chart\"")
    for s in (ch.get("series") if isinstance(ch, dict) else spec.get("series")) or []:
        if isinstance(s, dict):
            _check_keys(s, SERIES_KEYS, "a chart series")
    if kind in ITEM_KEYS:
        for it in spec.get("items") or spec.get("kpis") or spec.get("metrics") or spec.get("events") or spec.get("milestones") or spec.get("steps") or []:
            if isinstance(it, dict):
                _check_keys(it, ITEM_KEYS[kind], f"a {kind} item")
    for key in ("bullets", "items") if kind in ("bullets", "image-text", "agenda") else ():
        for it in spec.get(key) or []:
            if isinstance(it, dict):
                _check_keys(it, BULLET_KEYS, "a bullet")
    for side in ("left", "right"):
        if isinstance(spec.get(side), dict):
            _check_keys(spec[side], COLUMN_KEYS, f"\"{side}\"")
    for el in spec.get("elements") or spec.get("shapes") or []:
        if isinstance(el, dict):
            _check_keys(el, ELEMENT_KEYS, "an element")


def I(x: float) -> int:
    return int(round(x * IN))


@dataclass
class Rect:
    x: int
    y: int
    w: int
    h: int

    @property
    def r(self) -> int:
        return self.x + self.w

    @property
    def b(self) -> int:
        return self.y + self.h

    def cols(self, n: int, gutter: int, weights: list[float] | None = None) -> list["Rect"]:
        weights = weights or [1.0] * n
        total = sum(weights)
        avail = self.w - gutter * (n - 1)
        out, x = [], self.x
        for wt in weights:
            w = int(avail * wt / total)
            out.append(Rect(x, self.y, w, self.h))
            x += w + gutter
        return out

    def inset(self, dx: int, dy: int | None = None) -> "Rect":
        dy = dx if dy is None else dy
        return Rect(self.x + dx, self.y + dy, max(1, self.w - 2 * dx), max(1, self.h - 2 * dy))


@dataclass
class FitJob:
    slide: Any
    shape: Any
    min_scale: float
    groups: list[list[dict[str, Any]]] | None = None  # bullet groups, when the body may continue on new slides
    spec: dict[str, Any] | None = None
    kind: str = "text"
    no_grow: bool = False
    group: str | None = None  # jobs sharing a group get one common scale (KPI values, column bodies …)
    max_grow: float = 1.3


@dataclass
class Built:
    slides: list[Any] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


# ── colours by role ─────────────────────────────────────────────────────


class Palette:
    """Maps roles (bg, text, surface, muted, accent1-6, positive, negative, #hex) to python-pptx colours and RGB."""

    def __init__(self, prs: Any, theme: Theme | None):
        from _ooxml import ColorCtx, RT_THEME, Theme as XTheme, part_xml, related, _clr_map, child, DEFAULT_CLR_MAP

        master = prs.slide_masters[0]
        tp = related(master.part, RT_THEME)
        xt = XTheme(part_xml(tp)) if tp is not None else XTheme(None)
        cmap = dict(DEFAULT_CLR_MAP)
        cmap.update(_clr_map(child(master._element, "clrMap")) or {})
        self.cc = ColorCtx(xt, cmap)
        self.theme = theme
        self.heading_font = theme.heading_font if theme else xt.major
        self.body_font = theme.body_font if theme else xt.minor
        self.mono_font = theme.mono_font if theme else "Courier New"
        rgb = lambda n: Color.from_hex(self.cc.scheme_hex(n) or "000000")  # noqa: E731
        self.bg, self.text = rgb("bg1"), rgb("tx1")
        bg2, tx2 = rgb("bg2"), rgb("tx2")
        self.surface_role = ("theme", "bg2", 0.0) if contrast_ratio(bg2, self.bg) < 1.6 else ("theme", "accent1", 0.85 if self.bg.luminance() > 0.5 else -0.7)
        self.muted_role = ("theme", "tx2", 0.0) if contrast_ratio(tx2, self.bg) >= 4.0 else ("theme", "tx1", 0.35 if self.text.luminance() < 0.5 else -0.3)
        self.positive = theme.positive if theme else "15803D"
        self.negative = theme.negative if theme else "B91C1C"

    def spec(self, role: str | None) -> tuple[str, str, float]:
        r = (role or "text").strip()
        low = r.lower()
        if low in ("bg", "background", "bg1"):
            return ("theme", "bg1", 0.0)
        if low in ("text", "tx1", "fg"):
            return ("theme", "tx1", 0.0)
        if low in ("surface", "bg2", "card"):
            return self.surface_role
        if low in ("muted", "tx2", "secondary"):
            return self.muted_role
        if re.fullmatch(r"accent[1-6]", low):
            return ("theme", low, 0.0)
        if low in ("positive", "good", "up", "green"):
            return ("rgb", self.positive, 0.0)
        if low in ("negative", "bad", "down", "red"):
            return ("rgb", self.negative, 0.0)
        h = r.lstrip("#")
        if re.fullmatch(r"[0-9a-fA-F]{6}", h):
            return ("rgb", h.upper(), 0.0)
        raise UsageError(f"unknown colour '{role}' (use bg, text, surface, muted, accent1-6, positive, negative or #RRGGBB)")

    def apply(self, cf: Any, role: str | None, direct: bool = False) -> None:
        """Sets a python-pptx colour to a role. `direct` names dk1/lt1/dk2/lt2 instead of tx1/bg1/tx2/bg2 (chart
        parts: not every app maps them through the slide's colour map)."""
        from pptx.dml.color import RGBColor
        from pptx.enum.dml import MSO_THEME_COLOR

        kind, val, bright = self.spec(role)
        if kind == "rgb":
            cf.rgb = RGBColor.from_string(val)
            return
        if direct and val in ("bg1", "tx1", "bg2", "tx2"):
            val = self.cc.clr_map.get(val, val)
        cf.theme_color = {
            "bg1": MSO_THEME_COLOR.BACKGROUND_1, "tx1": MSO_THEME_COLOR.TEXT_1, "bg2": MSO_THEME_COLOR.BACKGROUND_2, "tx2": MSO_THEME_COLOR.TEXT_2,
            "dk1": MSO_THEME_COLOR.DARK_1, "lt1": MSO_THEME_COLOR.LIGHT_1, "dk2": MSO_THEME_COLOR.DARK_2, "lt2": MSO_THEME_COLOR.LIGHT_2,
            "accent1": MSO_THEME_COLOR.ACCENT_1, "accent2": MSO_THEME_COLOR.ACCENT_2, "accent3": MSO_THEME_COLOR.ACCENT_3,
            "accent4": MSO_THEME_COLOR.ACCENT_4, "accent5": MSO_THEME_COLOR.ACCENT_5, "accent6": MSO_THEME_COLOR.ACCENT_6,
        }[val]
        if bright:
            cf.brightness = bright

    def rgb(self, role: str | None) -> Color:
        kind, val, bright = self.spec(role)
        if kind == "rgb":
            return Color.from_hex(val)
        c = Color.from_hex(self.cc.scheme_hex(val) or "000000")
        if bright:
            import colorsys

            h, l, s = colorsys.rgb_to_hls(c.r / 255, c.g / 255, c.b / 255)
            l = l * (1 - bright) + bright if bright > 0 else l * (1 + bright)
            r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
            c = Color(round(r * 255), round(g * 255), round(b * 255))
        return c

    def on(self, role: str) -> str:
        """The text role that reads best on a fill of the given role."""
        fill = self.rgb(role)
        return "bg" if contrast_ratio(self.bg, fill) >= contrast_ratio(self.text, fill) else "text"


# ── the themed master ───────────────────────────────────────────────────


def themed_presentation(theme: Theme) -> Any:
    from lxml import etree
    from pptx import Presentation

    from _ooxml import RT_THEME, related

    prs = Presentation()
    old_w = prs.slide_width
    prs.slide_width, prs.slide_height = 12192000, 6858000
    sx = 12192000 / old_w
    master = prs.slide_masters[0]
    for el in [master._element] + [lay._element for lay in master.slide_layouts]:
        for xfrm in el.iter("{%s}xfrm" % A_NS):
            off, ext = xfrm.find("{%s}off" % A_NS), xfrm.find("{%s}ext" % A_NS)
            if off is not None:
                off.set("x", str(int(int(off.get("x", "0")) * sx)))
            if ext is not None:
                ext.set("cx", str(int(int(ext.get("cx", "0")) * sx)))
    # theme part: colours, fonts, name
    tp = related(master.part, RT_THEME)
    root = etree.fromstring(tp.blob)
    root.set("name", f"Desk {theme.name}")
    a = "{%s}" % A_NS
    cs = root.find(f"{a}themeElements/{a}clrScheme")
    cs.set("name", f"Desk {theme.name}")
    if theme.dark:
        roles = {"dk1": theme.bg, "lt1": theme.text, "dk2": theme.surface, "lt2": theme.muted}
    else:
        roles = {"dk1": theme.text, "lt1": theme.bg, "dk2": theme.muted, "lt2": theme.surface}
    roles.update({f"accent{i + 1}": c for i, c in enumerate(theme.accents)})
    roles["hlink"] = theme.hlink or theme.accents[0]
    roles["folHlink"] = theme.accents[5]
    for name, hexv in roles.items():
        e = cs.find(a + name)
        if e is None:
            continue
        for c in list(e):
            e.remove(c)
        etree.SubElement(e, a + "srgbClr", val=hexv)
    fs = root.find(f"{a}themeElements/{a}fontScheme")
    fs.set("name", f"Desk {theme.name}")
    for tag, font in (("majorFont", theme.heading_font), ("minorFont", theme.body_font)):
        lat = fs.find(f"{a}{tag}/{a}latin")
        if lat is not None:
            lat.set("typeface", font)
            for k in ("panose", "pitchFamily", "charset"):
                lat.attrib.pop(k, None)
    tp._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    _style_master(master, theme)
    _layout_grid(prs, theme)
    return prs


def _sub(parent: Any, tag: str, **attrs: str) -> Any:
    from lxml import etree

    ns = A_NS if tag.startswith("a:") else P_NS
    return etree.SubElement(parent, "{%s}%s" % (ns, tag[2:]), **attrs)


def _style_master(master: Any, t: Theme) -> None:
    from lxml import etree

    m = master._element
    p = "{%s}" % P_NS
    cm = m.find(p + "clrMap")
    if cm is not None and t.dark:
        cm.set("bg1", "dk1")
        cm.set("tx1", "lt1")
        cm.set("bg2", "dk2")
        cm.set("tx2", "lt2")
    # background: bg1
    csld = m.find(p + "cSld")
    bg = csld.find(p + "bg")
    if bg is not None:
        csld.remove(bg)
    bg = etree.Element(p + "bg")
    bgpr = _sub(bg, "p:bgPr")
    sf = _sub(bgpr, "a:solidFill")
    _sub(sf, "a:schemeClr", val="bg1")
    _sub(bgpr, "a:effectLst")
    csld.insert(0, bg)
    tx = m.find(p + "txStyles")
    s = t.sizes

    def lvl(parent: Any, n: int, size: float, bold: bool, font: str, color: str, bullet: str | None, marl: float, indent: float, before: float, line: float = 1.0, bucolor: str | None = None, algn: str = "l") -> None:
        e = parent.find("{%s}lvl%dpPr" % (A_NS, n))
        if e is None:
            e = _sub(parent, f"a:lvl{n}pPr")
        for c in list(e):
            e.remove(c)
        e.attrib.clear()
        e.set("marL", str(int(marl * IN)))
        e.set("indent", str(int(indent * IN)))
        e.set("algn", algn)
        e.set("defTabSz", "914400")
        e.set("rtl", "0")
        e.set("eaLnBrk", "1")
        e.set("latinLnBrk", "0")
        e.set("hangingPunct", "1")
        ls = _sub(e, "a:lnSpc")
        _sub(ls, "a:spcPct", val=str(int(line * 100000)))
        sb = _sub(e, "a:spcBef")
        _sub(sb, "a:spcPts", val=str(int(before * 100)))
        sa = _sub(e, "a:spcAft")
        _sub(sa, "a:spcPts", val="0")
        if bullet:
            bc = _sub(e, "a:buClr")
            _sub(bc, "a:schemeClr", val=bucolor or "accent1")
            _sub(e, "a:buSzPct", val="100000")
            _sub(e, "a:buFont", typeface="Arial", panose="020B0604020202020204", pitchFamily="34", charset="0")
            _sub(e, "a:buChar", char=bullet)
        else:
            _sub(e, "a:buNone")
        d = _sub(e, "a:defRPr", sz=str(int(size * 100)), b="1" if bold else "0", kern="1200")
        f = _sub(d, "a:solidFill")
        _sub(f, "a:schemeClr", val=color)
        _sub(d, "a:latin", typeface=font)
        _sub(d, "a:ea", typeface="+mn-ea")
        _sub(d, "a:cs", typeface="+mn-cs")

    ts = tx.find("{%s}titleStyle" % P_NS)
    lvl(ts, 1, s["title"], t.heading_bold, "+mj-lt", "tx1", None, 0, 0, 0, 0.95)
    bs = tx.find("{%s}bodyStyle" % P_NS)
    sizes = [s["body"], s["body2"], s["body3"], s["body3"], s["body3"]]
    chars = ["•", "–", "•", "–", "•"]
    for i in range(9):
        k = min(i, 4)
        lvl(bs, i + 1, sizes[k], False, "+mn-lt", "tx1", chars[k], 0.3 + 0.35 * i, -0.3 if i == 0 else -0.25, 12 if i == 0 else 5, 1.05, "accent1" if i == 0 else "tx2")
    os_ = tx.find("{%s}otherStyle" % P_NS)
    if os_ is not None:
        for i in range(1, 10):
            e = os_.find("{%s}lvl%dpPr" % (A_NS, i))
            if e is not None:
                d = e.find("{%s}defRPr" % A_NS)
                if d is not None:
                    d.set("sz", str(int(s["small"] * 100)))


def _layout_grid(prs: Any, t: Theme) -> None:
    """Re-lays the default layouts on the 16:9 grid and adds the theme's decorations to them."""
    from pptx.enum.text import MSO_ANCHOR

    W, H = prs.slide_width, prs.slide_height
    mx = I(0.75)
    title = Rect(mx, I(0.42), W - 2 * mx, I(0.98))
    body = Rect(mx, I(1.62), W - 2 * mx, H - I(1.62) - I(0.78))
    foot_y = H - I(0.52)
    for layout in prs.slide_masters[0].slide_layouts:
        name = layout.name
        for ph in layout.placeholders:
            pt_ = ph.placeholder_format.type
            tname = pt_.name if pt_ is not None else "OBJECT"
            idx = ph.placeholder_format.idx
            r: Rect | None = None
            if tname in ("DATE", "FOOTER", "SLIDE_NUMBER"):
                r = {"DATE": Rect(W - mx - I(4.5), foot_y, I(2.0), I(0.3)), "FOOTER": Rect(mx, foot_y, I(7.5), I(0.3)), "SLIDE_NUMBER": Rect(W - mx - I(1.2), foot_y, I(1.2), I(0.3))}[tname]
            elif name == "Title Slide":
                # clear of the accent panel (rule themes) or the corner discs (the others)
                tw = int(W * TITLE_PANEL) - I(0.9) - I(0.5) if t.decor == "rule" else W - I(1.8)
                sw = int(W * TITLE_PANEL) - I(0.9) - I(0.5) if t.decor == "rule" else int(W * 0.72) - I(0.9)
                r = Rect(I(0.9), I(1.85), tw, I(2.05)) if tname == "CENTER_TITLE" else Rect(I(0.9), I(4.15), sw, I(1.3))
            elif name == "Section Header":
                sw = int(W * 0.72) - I(0.9) if t.decor in ("bar", "band") else W - I(1.8)
                r = Rect(I(0.9), I(2.45), W - I(1.8), I(1.75)) if tname == "TITLE" else Rect(I(0.9), I(4.3), sw, I(1.2))
            elif tname in ("TITLE", "CENTER_TITLE"):
                r = title
            elif name == "Title and Content":
                r = body
            elif name == "Two Content":
                cols = body.cols(2, I(0.5))
                r = cols[0] if idx == 1 else cols[1]
            elif name == "Comparison":
                cols = body.cols(2, I(0.5))
                c = cols[0] if idx in (1, 2) else cols[1]
                r = Rect(c.x, c.y, c.w, I(0.6)) if idx in (1, 3) else Rect(c.x, c.y + I(0.7), c.w, c.h - I(0.7))
            if r is not None:
                ph.left, ph.top, ph.width, ph.height = r.x, r.y, r.w, r.h
            if tname in ("TITLE", "CENTER_TITLE") and name not in ("Title Slide", "Section Header"):
                ph.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
            if tname in ("BODY", "OBJECT") and name not in ("Title Slide", "Section Header"):
                # the stock layouts give some bodies their own sizes (28 pt in Two Content): use the theme's scale
                for d in ph._element.iter("{%s}defRPr" % A_NS):
                    d.attrib.pop("sz", None)
        _decorate_layout(layout, t, W, H, title)
    m = prs.slide_masters[0]
    for ph in m.placeholders:
        tname = ph.placeholder_format.type.name if ph.placeholder_format.type is not None else "OBJECT"
        if tname == "TITLE":
            ph.left, ph.top, ph.width, ph.height = title.x, title.y, title.w, title.h
            ph.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        elif tname == "BODY":
            ph.left, ph.top, ph.width, ph.height = body.x, body.y, body.w, body.h
        elif tname in ("DATE", "FOOTER", "SLIDE_NUMBER"):
            r = {"DATE": Rect(W - mx - I(4.5), foot_y, I(2.0), I(0.3)), "FOOTER": Rect(mx, foot_y, I(7.5), I(0.3)), "SLIDE_NUMBER": Rect(W - mx - I(1.2), foot_y, I(1.2), I(0.3))}[tname]
            ph.left, ph.top, ph.width, ph.height = r.x, r.y, r.w, r.h
    _subtitle_style(m.slide_layouts, t)


def _subtitle_style(layouts: Any, t: Theme) -> None:
    """Subtitles and section texts: muted, no bullets, left aligned."""
    for lay in layouts:
        if lay.name not in ("Title Slide", "Section Header"):
            continue
        for ph in lay.placeholders:
            tn = ph.placeholder_format.type.name if ph.placeholder_format.type is not None else ""
            tb = ph._element.find("{%s}txBody" % P_NS)
            if tb is None:
                continue
            ls = tb.find("{%s}lstStyle" % A_NS)
            if ls is None:
                ls = _sub(tb, "a:lstStyle")
                tb.insert(1, ls)
            for c in list(ls):
                ls.remove(c)
            e = _sub(ls, "a:lvl1pPr", marL="0", indent="0", algn="l")
            if tn in ("CENTER_TITLE", "TITLE"):
                ln = _sub(e, "a:lnSpc")
                _sub(ln, "a:spcPct", val="92000")
                _sub(e, "a:buNone")
                d = _sub(e, "a:defRPr", sz=str(int((t.sizes["cover"] if lay.name == "Title Slide" else t.sizes["section"]) * 100)), b="1" if t.heading_bold else "0")
                f = _sub(d, "a:solidFill")
                _sub(f, "a:schemeClr", val="tx1")
            else:
                sb = _sub(e, "a:spcBef")
                _sub(sb, "a:spcPts", val="0")
                _sub(e, "a:buNone")
                d = _sub(e, "a:defRPr", sz=str(int(t.sizes["subtitle"] * 100)), b="0")
                f = _sub(d, "a:solidFill")
                _sub(f, "a:schemeClr", val="tx2")
            bp = tb.find("{%s}bodyPr" % A_NS)
            if bp is not None:
                bp.set("anchor", "b" if tn in ("CENTER_TITLE", "TITLE") else "t")


TITLE_PANEL = 0.7  # rule-decor themes: the title slide's accent panel starts at this share of the width


def _decorate_layout(layout: Any, t: Theme, W: int, H: int, title: Rect) -> None:
    name = layout.name
    shapes: list[tuple[Rect, str, str]] = []

    def motif(big: str, small: str) -> None:
        # two overlapping discs in the bottom-right corner, cut by the slide edge
        d = int(H * 0.62)
        shapes.append((Rect(W - int(d * 0.55), H - int(d * 0.5), d, d), big, "ellipse"))
        s = int(H * 0.16)
        # the small disc sits on the big one's upper-left edge
        cx, cy = W - int(d * 0.404), H - int(d * 0.354)
        shapes.append((Rect(cx - s // 2, cy - s // 2, s, s), small, "ellipse"))

    if name in ("Title and Content", "Two Content", "Comparison", "Title Only", "Content with Caption"):
        if t.decor == "bar":
            shapes.append((Rect(title.x + I(0.1), title.b + I(0.06), I(0.8), I(0.06)), "accent1", "rect"))
        elif t.decor == "rule":
            shapes.append((Rect(title.x, title.b + I(0.08), title.w, I(0.012)), "tx2", "rect"))
        elif t.decor == "band":
            shapes.append((Rect(0, 0, I(0.16), H), "accent1", "rect"))
    elif name == "Title Slide":
        if t.decor == "band":
            shapes.append((Rect(0, 0, I(0.42), H), "accent1", "rect"))
        if t.decor == "rule":
            shapes.append((Rect(int(W * TITLE_PANEL), 0, W - int(W * TITLE_PANEL), H), "accent1", "rect"))
            shapes.append((Rect(int(W * TITLE_PANEL) - I(0.12), 0, I(0.12), H), "accent2", "rect"))
            shapes.append((Rect(I(1.0), I(4.0), I(2.4), I(0.03)), "accent1", "rect"))
        else:
            shapes.append((Rect(I(1.0), I(3.97), I(1.3), I(0.08)), "accent1", "rect"))
            motif("accent1", "accent2")
    elif name == "Section Header" and t.decor in ("bar", "band"):
        if t.section_bg == "accent":
            motif("accent2", "accent3")
        else:
            motif("accent1", "accent2")
    for r, clr, prst in shapes:
        _layout_rect(layout, r, clr, prst)


def _layout_rect(layout: Any, r: Rect, clr: str, prst: str = "rect") -> None:
    from lxml import etree

    tree = layout._element.find("{%s}cSld/{%s}spTree" % (P_NS, P_NS))
    ids = [int(e.get("id", "0")) for e in tree.iter("{%s}cNvPr" % P_NS)]
    nid = max(ids + [1]) + 1
    xml = (
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="{nid}" name="Desk decoration {nid}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{r.x}" y="{r.y}"/><a:ext cx="{r.w}" cy="{r.h}"/></a:xfrm><a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:schemeClr val="{clr}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>'
    )
    el = etree.fromstring(xml)
    # decorations sit behind the placeholders
    first_ph = None
    for c in tree:
        if local(c) in ("sp", "pic", "graphicFrame") and c.find(".//{%s}ph" % P_NS) is not None:
            first_ph = c
            break
    if first_ph is not None:
        first_ph.addprevious(el)
    else:
        tree.append(el)


# ── layouts ─────────────────────────────────────────────────────────────


class Layouts:
    """Finds layouts by name or by their placeholder signature, and the content areas they define."""

    def __init__(self, prs: Any):
        self.prs = prs
        self.all = [lay for m in prs.slide_masters for lay in m.slide_layouts]
        self.W, self.H = prs.slide_width, prs.slide_height

    def by_name(self, name: str) -> Any:
        for lay in self.all:
            if lay.name == name:
                return lay
        for lay in self.all:
            if lay.name.lower() == name.lower():
                return lay
        for lay in self.all:
            if name.lower() in lay.name.lower():
                return lay
        raise UsageError(f"no layout named '{name}'; layouts: {', '.join(lay.name for lay in self.all)}")

    @staticmethod
    def sig(lay: Any) -> dict[str, list[Any]]:
        out: dict[str, list[Any]] = {"title": [], "ctitle": [], "subtitle": [], "body": [], "pic": [], "other": []}
        for ph in lay.placeholders:
            t = ph.placeholder_format.type.name if ph.placeholder_format.type is not None else "OBJECT"
            if t in ("DATE", "FOOTER", "SLIDE_NUMBER", "HEADER"):
                continue
            key = {"TITLE": "title", "CENTER_TITLE": "ctitle", "VERTICAL_TITLE": "other", "SUBTITLE": "subtitle", "BODY": "body", "OBJECT": "body", "PICTURE": "pic", "VERTICAL_BODY": "other", "VERTICAL_OBJECT": "other"}.get(t, "body")
            out[key].append(ph)
        return out

    def pick(self, role: str) -> Any:
        prefs = {
            "title": ["title slide"], "section": ["section header", "section"], "content": ["title and content"],
            "two": ["two content", "comparison"], "comparison": ["comparison", "two content"], "title_only": ["title only"], "blank": ["blank"],
        }[role]
        for p in prefs:
            for lay in self.all:
                if lay.name.lower() == p:
                    return lay
        scored = []
        for lay in self.all:
            s = self.sig(lay)
            nb = len(s["body"])
            has_t = bool(s["title"] or s["ctitle"])
            score = -1
            if role == "title":
                score = 10 if s["ctitle"] else (6 if has_t and (s["subtitle"] or nb == 1) else -1)
            elif role == "section":
                score = 8 if "section" in lay.name.lower() else (4 if s["title"] and nb == 1 and not s["other"] else -1)
            elif role == "content":
                score = 10 if s["title"] and nb == 1 and not s["pic"] else (5 if has_t and nb >= 1 else -1)
            elif role in ("two", "comparison"):
                score = 10 if s["title"] and nb == (4 if role == "comparison" else 2) else (6 if s["title"] and nb >= 2 else -1)
            elif role == "title_only":
                score = 10 if s["title"] and nb == 0 and not s["pic"] else (3 if s["title"] else -1)
            elif role == "blank":
                score = 10 if not any(s.values()) else -1
            if score >= 0:
                scored.append((score, -self.all.index(lay), lay))
        if scored:
            return max(scored, key=lambda x: (x[0], x[1]))[2]
        if role == "blank":
            return self.pick("title_only")
        return self.all[0]

    def title_rect(self) -> Rect:
        lay = self.pick("title_only")
        for ph in lay.placeholders:
            if ph.placeholder_format.type is not None and ph.placeholder_format.type.name in ("TITLE", "CENTER_TITLE"):
                return Rect(ph.left or 0, ph.top or 0, ph.width or self.W, ph.height or I(1))
        return Rect(I(0.6), I(0.4), self.W - I(1.2), I(1.0))

    def content_rect(self) -> Rect:
        lay = self.pick("content")
        s = self.sig(lay)
        if s["body"]:
            ph = s["body"][0]
            if ph.width and ph.height:
                return Rect(ph.left, ph.top, ph.width, ph.height)
        t = self.title_rect()
        return Rect(t.x, t.b + I(0.2), t.w, self.H - t.b - I(0.8))


# ── the builder ─────────────────────────────────────────────────────────


class DeckBuilder:
    def __init__(self, theme: Theme | None, template: Path | None = None, keep_slides: bool = False, base_dir: Path | None = None, prs: Any = None):
        from _fonts import FontEnv

        self.base = base_dir or Path.cwd()
        self.theme = theme
        if prs is not None:
            self.prs = prs
            self.mode = "template"
        elif template is not None:
            from _deck import open_deck

            self.prs = open_deck(template).prs
            self.mode = "template"
            if not keep_slides:
                self._remove_all_slides()
            self._kept_slides = keep_slides
        else:
            assert theme is not None
            self.prs = themed_presentation(theme)
            self.mode = "theme"
        if not hasattr(self, "_kept_slides"):
            self._kept_slides = prs is not None
        self.lay = Layouts(self.prs)
        self.pal = Palette(self.prs, theme if self.mode == "theme" else None)
        self.fonts = FontEnv()
        self.jobs: list[FitJob] = []
        self.warnings: list[str] = []
        self.section_no = 0
        self.content_slides: list[Any] = []
        self.kinds: dict[int, str] = {}
        self.sizes = dict(theme.sizes) if theme else {"cover": 40, "section": 36, "title": 28, "subtitle": 20, "body": 20, "body2": 18, "body3": 16, "small": 14, "caption": 12, "footer": 10, "kpi": 48, "quote": 28, "min_body": 12, "min_title": 20}
        self.after: dict[int, list[Any]] = {}  # continuation slides to place after a slide (by slide id)
        self.footer_span: dict[int, tuple[int, int]] = {}  # slides whose footer must stay inside a text column
        self._tmp: list[Path] = []
        self.missing_layouts: dict[str, int] = {}

    # slides ------------------------------------------------------------

    def _remove_all_slides(self) -> None:
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            rid = sid.get("{%s}id" % R_NS)
            lst.remove(sid)
            if rid:
                self.prs.part.drop_rel(rid)

    def new_slide(self, layout: Any, kind: str) -> Any:
        slide = self.prs.slides.add_slide(layout)
        self.kinds[slide.slide_id] = kind
        # Layouts with two placeholders sharing an idx (LibreOffice's .ppt/.odp exports) make a slide placeholder
        # inherit the wrong box: give each clone its layout placeholder's position explicitly.
        src = list(layout.iter_cloneable_placeholders())
        idxs = [p.placeholder_format.idx for p in src]
        if len(set(idxs)) < len(idxs):
            for ph, lp in zip(list(slide.placeholders), src):
                if lp.width and lp.height:
                    ph.left, ph.top, ph.width, ph.height = lp.left or 0, lp.top or 0, lp.width, lp.height
        return slide

    def ph(self, slide: Any, *types: str, index: int = 0) -> Any:
        found = []
        for ph in slide.placeholders:
            t = ph.placeholder_format.type.name if ph.placeholder_format.type is not None else "OBJECT"
            if t in types:
                found.append(ph)
        found.sort(key=lambda p: (p.left or 0, p.top or 0))
        return found[index] if index < len(found) else None

    def drop_empty_placeholders(self, slide: Any) -> None:
        """Removes placeholders left empty (they would show "Click to add text" and trip the linter)."""
        for ph in list(slide.placeholders):
            if ph.has_text_frame and ph.text_frame.text.strip():
                continue
            ph._element.getparent().remove(ph._element)

    # public ------------------------------------------------------------

    def add(self, spec: dict[str, Any]) -> list[Any]:
        if not isinstance(spec, dict):
            raise UsageError("each slide must be an object like {\"type\": \"bullets\", \"title\": ...}")
        kind = str(spec.get("type") or ("bullets" if spec.get("bullets") else "title" if not self.prs.slides else "bullets")).lower()
        kind = TYPE_ALIASES.get(kind, kind)
        if kind not in SLIDE_TYPES:
            raise UsageError(f"unknown slide type '{kind}' (types: {', '.join(SLIDE_TYPES)})")
        validate_slide(spec, kind)
        fn = getattr(self, "s_" + kind.replace("-", "_"))
        slide = fn(spec)
        self._common(slide, spec, kind)
        return [slide]

    def layout_for(self, spec: dict[str, Any], role: str) -> Any:
        """The layout a slide names (warns and picks one itself when the deck has no such layout), else by role."""
        name = spec.get("layout")
        if name:
            try:
                return self.lay.by_name(str(name))
            except UsageError:
                self.missing_layouts.setdefault(str(name), 0)
                self.missing_layouts[str(name)] += 1
        return self.lay.pick(role)

    def _common(self, slide: Any, spec: dict[str, Any], kind: str) -> None:
        notes = spec.get("notes")
        if notes:
            tf = slide.notes_slide.notes_text_frame
            tf.text = plain(str(notes)) if not isinstance(notes, list) else "\n".join(plain(str(n)) for n in notes)
        if spec.get("hidden"):
            slide._element.set("show", "0")
        bg = spec.get("background")
        if bg:
            self._background(slide, bg)
        if kind not in ("title", "section", "closing") and not spec.get("full_bleed"):
            self.content_slides.append(slide)

    def _background(self, slide: Any, bg: Any) -> None:
        if isinstance(bg, dict) and bg.get("image") or isinstance(bg, str) and self._looks_like_path(bg):
            path = bg["image"] if isinstance(bg, dict) else bg
            self._picture(slide, path, Rect(0, 0, self.lay.W, self.lay.H), "cover", send_back=True)
            return
        slide.background.fill.solid()
        self.pal.apply(slide.background.fill.fore_color, str(bg if not isinstance(bg, dict) else bg.get("color", "bg")))

    @staticmethod
    def _looks_like_path(s: str) -> bool:
        return bool(re.search(r"\.(png|jpe?g|gif|bmp|tiff?|webp|svg)$", s.strip(), re.I))

    # text --------------------------------------------------------------

    def items(self, value: Any) -> list[dict[str, Any]]:
        """Normalises bullets: strings, {text, level, bullet, numbered, children} → flat list with levels."""
        out: list[dict[str, Any]] = []

        def walk(v: Any, level: int, nested: bool = False) -> None:
            if v is None:
                return
            if isinstance(v, str):
                for line in v.split("\n"):
                    if line.strip():
                        out.append({"text": line.strip(), "level": level})
                return
            if isinstance(v, list):
                # a list inside a list holds the sub-points of the item before it: ["Point", ["sub a", "sub b"]]
                for x in v:
                    walk(x, level + 1 if isinstance(x, list) else level, True)
                return
            if isinstance(v, dict):
                txt = v.get("text", v.get("title", ""))
                it = {"text": str(txt), "level": int(v.get("level", level))}
                for k in ("bullet", "numbered", "bold", "color", "size"):
                    if k in v:
                        it[k] = v[k]
                out.append(it)
                if v.get("children") or v.get("items") or v.get("bullets"):
                    walk(v.get("children") or v.get("items") or v.get("bullets"), it["level"] + 1)
                return
            out.append({"text": str(v), "level": level})

        walk(value, 0)
        for it in out:
            it["level"] = max(0, min(4, it["level"]))
        return out

    def write(self, tf: Any, items: list[dict[str, Any]], *, size: float | None = None, color: str | None = None, bold: bool | None = None, font: str | None = None, align: str | None = None, placeholder: bool = False, bullets: bool = False, line: float | None = None, space_before: float | None = None) -> None:
        """Writes paragraphs into a text frame; inline Markdown becomes bold/italic/code/link runs."""
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        tf.clear()
        first = True
        for it in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = it.get("level", 0)
            if align:
                p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[align]
            if line is not None:
                p.line_spacing = line
            if space_before is not None:
                p.space_before = Pt(space_before)
            want_bullet = it.get("bullet", True) if (placeholder or bullets) else False
            if it.get("numbered"):
                _bullet(p, "num", self.pal, it.get("level", 0), placeholder)
            elif placeholder and not want_bullet:
                _bullet(p, None, self.pal, 0, placeholder)
            elif bullets and want_bullet:
                _bullet(p, "•" if it.get("level", 0) % 2 == 0 else "–", self.pal, it.get("level", 0), placeholder)
            runs = inline_runs(it.get("text", ""))
            if not runs:
                runs = [{"text": ""}]
            for r in runs:
                run = p.add_run()
                run.text = r["text"]
                f = run.font
                if size or it.get("size"):
                    f.size = Pt(float(it.get("size") or size))  # type: ignore[arg-type]
                if r.get("bold") or bold or it.get("bold"):
                    f.bold = True
                elif bold is False:
                    f.bold = False
                if r.get("italic"):
                    f.italic = True
                if r.get("strike"):
                    run._r.get_or_add_rPr().set("strike", "sngStrike")
                if r.get("code"):
                    f.name = self.pal.mono_font
                elif font:
                    f.name = font
                if it.get("color") or color:
                    self.pal.apply(f.color, it.get("color") or color)
                if r.get("link"):
                    try:
                        run.hyperlink.address = r["link"]
                    except Exception:  # noqa: BLE001 — a bad URL keeps its text
                        pass

    def textbox(self, slide: Any, r: Rect, items: Any, *, size: float, color: str = "text", bold: bool | None = None, font: str | None = None, align: str = "left", anchor: str = "top", fit: float | None = 0.7, bullets: bool = False, line: float | None = None, margin: tuple[float, float] = (0.0, 0.0), space_before: float | None = None, name: str | None = None, group: str | None = None) -> Any:
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

        tb = slide.shapes.add_textbox(r.x, r.y, r.w, r.h)
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = I(margin[0])
        tf.margin_top = tf.margin_bottom = I(margin[1])
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        its = self.items(items) if not (isinstance(items, list) and items and isinstance(items[0], dict) and "level" in items[0]) else items
        self.write(tf, its, size=size, color=color, bold=bold, font=font, align=align, bullets=bullets, line=line, space_before=space_before)
        # Boxes are measured: text shrinks to fit (fit < 1 allows shrinking to that share of its size), and a word
        # wider than the box is caught (PowerPoint would break it mid-word). fit=None: fixed furniture (footers).
        if fit is not None:
            self.jobs.append(FitJob(slide, tb, min(1.0, fit), group=group))
        return tb

    def title(self, slide: Any, text: Any, kind: str = "title", color: str | None = None) -> Any:
        ph = self.ph(slide, "TITLE", "CENTER_TITLE")
        if text is None or str(text).strip() == "":
            return None
        if ph is None:
            # a layout without a title (Blank): a styled box that is still a real title placeholder, so outlines,
            # lint and screen readers see the slide's title
            r = self.lay.title_rect()
            tb = self.textbox(slide, r, [{"text": str(text), "level": 0}], size=self.sizes["title"], bold=True, font="+mj-lt", anchor="bottom", color=color or "text", fit=0.5)
            nv = tb._element.nvSpPr
            nv.cNvSpPr.attrib.pop("txBox", None)
            _sub(nv.nvPr, "p:ph").set("type", "title")
            tb.name = f"Title {tb.shape_id}"
            return tb
        self.write(ph.text_frame, [{"text": str(text), "level": 0}], color=color)
        base = {"title": self.sizes["title"], "cover": self.sizes["cover"], "section": self.sizes["section"]}.get(kind, self.sizes["title"])
        self.jobs.append(FitJob(slide, ph, max(0.5, self.sizes["min_title"] / base)))
        return ph

    def body_items(self, slide: Any, ph: Any, items: list[dict[str, Any]], spec: dict[str, Any] | None = None, split: bool = True, group: str | None = None, max_grow: float = 1.3) -> None:
        self.write(ph.text_frame, items, placeholder=True)
        _norm_autofit(ph)
        groups = _groups(items) if split else None
        self.jobs.append(FitJob(slide, ph, self.sizes["min_body"] / self.sizes["body"], groups, spec, "body", group=group, max_grow=max_grow))

    # pictures ----------------------------------------------------------

    def resolve_image(self, path: str) -> Path:
        p = Path(str(path)).expanduser()
        if not p.is_absolute():
            p = self.base / p
        if not p.exists():
            raise SkillError(f"image not found: {path} (paths are relative to {self.base})")
        return p

    def _picture(self, slide: Any, path: str, r: Rect, fit: str = "fit", send_back: bool = False, name: str | None = None) -> Any:
        from PIL import Image

        src = self.resolve_image(path)
        data, px = _image_bytes(src, r)
        pw, ph = px
        if fit == "cover":
            pic = slide.shapes.add_picture(io.BytesIO(data), r.x, r.y, r.w, r.h)
            box_ratio, img_ratio = r.w / r.h, pw / ph
            if img_ratio > box_ratio:
                c = (1 - box_ratio / img_ratio) / 2
                pic.crop_left = pic.crop_right = c
            else:
                c = (1 - img_ratio / box_ratio) / 2
                pic.crop_top = pic.crop_bottom = c
        else:
            scale = min(r.w / pw, r.h / ph)
            w, h = int(pw * scale), int(ph * scale)
            pic = slide.shapes.add_picture(io.BytesIO(data), r.x + (r.w - w) // 2, r.y + (r.h - h) // 2, w, h)
        pic.name = name or f"Picture {src.name}"
        try:
            pic._element.find(".//{%s}cNvPr" % P_NS).set("descr", src.stem.replace("_", " ").replace("-", " "))
        except AttributeError:
            pass
        if send_back:
            tree = slide.shapes._spTree
            tree.remove(pic._element)
            tree.insert(2, pic._element)
        del Image
        return pic

    # shapes ------------------------------------------------------------

    def rect(self, slide: Any, r: Rect, fill: str | None, line: str | None = None, radius: float | None = None, name: str | None = None, line_w: float = 1.0, shape: str | None = None) -> Any:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt

        kind = {"oval": MSO_SHAPE.OVAL, "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON}.get(shape or "", MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        sh = slide.shapes.add_shape(kind, r.x, r.y, r.w, r.h)
        st = sh._element.find("{%s}style" % P_NS)
        if st is not None:
            sh._element.remove(st)
        if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = max(0.0, min(0.5, radius * IN / max(1, min(r.w, r.h))))
        if kind == MSO_SHAPE.CHEVRON or kind == MSO_SHAPE.PENTAGON:
            sh.adjustments[0] = min(0.5, (r.h * 0.35) / max(1, r.h))
        if fill and str(fill).lower() != "none":
            sh.fill.solid()
            self.pal.apply(sh.fill.fore_color, fill)
        else:
            sh.fill.background()
        if line and str(line).lower() != "none":
            sh.line.width = Pt(line_w)
            self.pal.apply(sh.line.color, line)
        else:
            sh.line.fill.background()
        if name:
            sh.name = name
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = I(0.12)
        tf.margin_top = tf.margin_bottom = I(0.06)
        return sh

    def line(self, slide: Any, x1: int, y1: int, x2: int, y2: int, color: str, width: float = 1.5) -> Any:
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.util import Pt

        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        st = c._element.find("{%s}style" % P_NS)
        if st is not None:
            c._element.remove(st)
        c.line.width = Pt(width)
        self.pal.apply(c.line.color, color)
        return c

    # ── slide types ─────────────────────────────────────────────────────

    def s_title(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title"), "title")
        self.title(slide, spec.get("title", ""), "cover")
        sub = self.ph(slide, "SUBTITLE") or self.ph(slide, "BODY", "OBJECT")
        lines = [x for x in (spec.get("subtitle"), spec.get("meta") or " · ".join(str(x) for x in (spec.get("author"), spec.get("date")) if x)) if x]
        if sub is not None and lines:
            its = [{"text": str(lines[0]), "level": 0, "bullet": False}]
            if len(lines) > 1:
                its.append({"text": str(lines[1]), "level": 0, "bullet": False, "size": self.sizes["small"], "color": "muted"})
            self.write(sub.text_frame, its, placeholder=True)
            self.jobs.append(FitJob(slide, sub, 0.7))
        if spec.get("image"):
            r = Rect(int(self.lay.W * 0.58), 0, self.lay.W - int(self.lay.W * 0.58), self.lay.H)
            self._picture(slide, spec["image"], r, "cover")
            t = self.ph(slide, "TITLE", "CENTER_TITLE")
            for shp in (t, sub):
                if shp is not None and self.mode == "theme":
                    shp.width = int(self.lay.W * 0.58) - shp.left - I(0.4)
        self.drop_empty_placeholders(slide)
        return slide

    def s_section(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "section"), "section")
        self.section_no += 1
        fg = None
        if self.mode == "theme" and self.theme is not None and not spec.get("background"):
            role = {"surface": "surface", "accent": "accent1", "text": "text"}[self.theme.section_bg]
            if role != "surface":
                slide.background.fill.solid()
                self.pal.apply(slide.background.fill.fore_color, role)
                fg = self.pal.on(role)
            else:
                slide.background.fill.solid()
                self.pal.apply(slide.background.fill.fore_color, "surface")
        t = self.title(slide, spec.get("title", ""), "section", color=fg)
        num = spec.get("number", f"{self.section_no:02d}" if self.mode == "theme" and spec.get("number", True) is not False else None)
        if num and t is not None and spec.get("number", True) is not False:
            r = Rect(t.left, t.top - I(1.05), I(3), I(1.0))
            self.textbox(slide, r, [{"text": str(num), "level": 0}], size=self.sizes["kpi"], bold=True, color=fg or "accent1", font="+mj-lt", anchor="bottom", fit=0.6, margin=(0.1, 0.0), name="Section number")
        body = self.ph(slide, "BODY", "OBJECT", "SUBTITLE")
        text = spec.get("subtitle") or spec.get("text")
        if body is not None and text:
            self.write(body.text_frame, [{"text": str(text), "level": 0, "bullet": False}], placeholder=True, color=fg if fg else None)
            self.jobs.append(FitJob(slide, body, 0.7))
        self._section_fg = fg
        self.drop_empty_placeholders(slide)
        return slide

    def s_closing(self, spec: dict[str, Any]) -> Any:
        spec = dict(spec)
        spec.setdefault("title", "Thank you")
        slide = self.s_section({**spec, "number": False})
        self.kinds[slide.slide_id] = "closing"
        fg = getattr(self, "_section_fg", None)
        t = self.ph(slide, "TITLE", "CENTER_TITLE")
        if t is not None and self.mode == "theme":
            # a short accent bar under the title, like the title slide's
            self.rect(slide, Rect(t.left + I(0.1), t.top + t.height + I(0.02), I(1.3), I(0.08)), "accent2" if self.theme is not None and self.theme.section_bg == "accent" else "accent1", name="Desk decoration closing")
        if spec.get("contact"):
            r = Rect(I(1.0), self.lay.H - I(1.6), int(self.lay.W * 0.7) - I(1.0), I(0.6))
            self.textbox(slide, r, [{"text": str(spec["contact"]), "level": 0}], size=self.sizes["body3"], color=fg or "muted", fit=0.8, name="Contact")
        return slide

    def s_bullets(self, spec: dict[str, Any]) -> Any:
        items = self.items(spec.get("bullets") or spec.get("items") or spec.get("text") or [])
        cols = int(spec.get("columns", 1) or 1)
        if cols == 2 or (cols == 1 and spec.get("columns") is None and len(items) >= 10 and all(len(plain(i["text"])) <= 40 for i in items) and all(i["level"] == 0 for i in items)):
            half = (len(items) + 1) // 2
            return self.s_two_column({"title": spec.get("title"), "left": {"bullets": items[:half]}, "right": {"bullets": items[half:]}, "layout": spec.get("layout")})
        slide = self.new_slide(self.layout_for(spec, "content"), "bullets")
        self.title(slide, spec.get("title"))
        body = self.ph(slide, "BODY", "OBJECT")
        if body is None:
            r = self.lay.content_rect()
            tb = self.textbox(slide, r, items, size=self.sizes["body"], bullets=True, fit=self.sizes["min_body"] / self.sizes["body"])
            del tb
        elif items:
            self.body_items(slide, body, items, spec)
        self.drop_empty_placeholders(slide)
        return slide

    def _column(self, slide: Any, r: Rect, col: dict[str, Any], ph: Any = None, side: int = 1, heading_room: bool = False) -> None:
        """Fills one column area: heading, then bullets/text, image, table, chart or code."""
        if not isinstance(col, dict):
            col = {"bullets": col}
        _check_keys(col, COLUMN_KEYS, "a column")
        y = r.y
        gid = f"{slide.slide_id}"
        if col.get("heading") or heading_room:
            hh = I(0.6)
            if col.get("heading"):
                # headings a step above the body text, aligned with the title's text (0.1 in inset)
                self.textbox(slide, Rect(r.x, y, r.w, hh), [{"text": str(col["heading"]), "level": 0}], size=self.sizes["body"] * 1.2, bold=True, color=col.get("heading_color", "accent1"), anchor="bottom", fit=0.6, margin=(0.1, 0.0), name=f"Column heading {side}", group=f"colhead-{gid}")
                self.line(slide, r.x, y + hh + I(0.06), r.r, y + hh + I(0.06), "accent1", 1.25)
            y += hh + I(0.2)
        area = Rect(r.x, y, r.w, r.b - y)
        if col.get("image"):
            self._picture(slide, col["image"], area, col.get("fit", "fit"))
            if ph is not None:
                ph._element.getparent().remove(ph._element)
        elif col.get("table"):
            self._table(slide, area, col["table"], split=False)
            if ph is not None:
                ph._element.getparent().remove(ph._element)
        elif col.get("chart"):
            self._chart(slide, area, col["chart"])
            if ph is not None:
                ph._element.getparent().remove(ph._element)
        elif col.get("code"):
            self._code(slide, area, str(col["code"]))
            if ph is not None:
                ph._element.getparent().remove(ph._element)
        else:
            items = self.items(col.get("bullets") or col.get("items") or col.get("text") or [])
            if ph is not None:
                ph.left, ph.top, ph.width, ph.height = area.x, area.y, area.w, area.h
                if items:
                    # both columns take one size, and grow less than a full-width list would
                    self.body_items(slide, ph, items, None, split=False, group=f"colbody-{gid}", max_grow=1.2)
            elif items:
                self.textbox(slide, area, items, size=self.sizes["body2"], bullets=True, fit=self.sizes["min_body"] / self.sizes["body2"], space_before=8, name=f"Column text {side}", group=f"colbody-{gid}")

    def s_two_column(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "two"), "two-column")
        self.title(slide, spec.get("title"))
        area = self.lay.content_rect()
        note = spec.get("note")
        if note:
            area = Rect(area.x, area.y, area.w, area.h - I(0.75))
            self.textbox(slide, Rect(area.x, area.b + I(0.15), area.w, I(0.6)), self.items(note), size=self.sizes["body3"], color="muted", anchor="top", fit=0.7, margin=(0.1, 0.0), name="Columns note")
        ratio = spec.get("split") or spec.get("ratio")
        weights = [float(x) for x in ratio] if isinstance(ratio, list) and len(ratio) == 2 else None
        left, right = area.cols(2, I(0.5), weights)
        phs = [self.ph(slide, "BODY", "OBJECT", index=0), self.ph(slide, "BODY", "OBJECT", index=1)]
        lcol, rcol = spec.get("left") or {}, spec.get("right") or {}
        # when only one column has a heading, the other keeps the same top so the bodies line up
        either = any(isinstance(c, dict) and c.get("heading") for c in (lcol, rcol))
        self._column(slide, left, lcol, phs[0], 1, either)
        self._column(slide, right, rcol, phs[1], 2, either)
        self.drop_empty_placeholders(slide)
        return slide

    def s_comparison(self, spec: dict[str, Any]) -> Any:
        spec = dict(spec)
        for side in ("left", "right"):
            col = spec.get(side)
            if isinstance(col, dict) and "heading" not in col and col.get("title"):
                col = dict(col)
                col["heading"] = col.pop("title")
                spec[side] = col
        return self.s_two_column(spec)

    def s_image(self, spec: dict[str, Any]) -> Any:
        full = bool(spec.get("full_bleed")) or not spec.get("title")
        slide = self.new_slide(self.layout_for(spec, "title_only" if spec.get("title") else "blank"), "image")
        images = spec.get("images") or ([spec["image"]] if spec.get("image") else [])
        if not images:
            raise UsageError("an image slide needs \"image\": \"path\"")
        if full and not spec.get("title"):
            area = Rect(0, 0, self.lay.W, self.lay.H)
            fit = spec.get("fit", "cover")
        else:
            self.title(slide, spec.get("title"))
            area = self.lay.content_rect()
            fit = spec.get("fit", "fit")
        cap = spec.get("caption")
        if cap:
            if full and not spec.get("title"):
                band = Rect(0, self.lay.H - I(0.9), self.lay.W, I(0.9))
                self.rect(slide, band, "text", name="Caption band")
                self.textbox(slide, band.inset(I(0.75), I(0.1)), [{"text": str(cap), "level": 0}], size=self.sizes["small"], color=self.pal.on("text"), anchor="middle", fit=0.8, name="Caption")
            else:
                area = Rect(area.x, area.y, area.w, area.h - I(0.5))
                self.textbox(slide, Rect(area.x, area.b + I(0.1), area.w, I(0.4)), [{"text": str(cap), "level": 0}], size=self.sizes["caption"], color="muted", align="center", fit=0.8, name="Caption")
        if len(images) > 1:
            cells = area.cols(min(len(images), 4), I(0.3))
            for p_, c in zip(images, cells):
                self._picture(slide, p_, c, fit)
        else:
            pic = self._picture(slide, images[0], area, fit)
            if full:
                tree = slide.shapes._spTree
                tree.remove(pic._element)
                tree.insert(2, pic._element)
        self.drop_empty_placeholders(slide)
        return slide

    def s_image_text(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "image-text")
        side = spec.get("image_side", "left")
        W, H = self.lay.W, self.lay.H
        split = int(W * float(spec.get("image_width", 0.45)))
        img_r = Rect(0, 0, split, H) if side == "left" else Rect(W - split, 0, split, H)
        if spec.get("image"):
            self._picture(slide, spec["image"], img_r, spec.get("fit", "cover"))
        tx = split + I(0.6) if side == "left" else I(0.75)
        tw = W - split - I(0.6) - I(0.75)
        self.footer_span[slide.slide_id] = (tx, tx + tw)
        t = self.ph(slide, "TITLE", "CENTER_TITLE")
        if t is not None:
            t.left, t.top, t.width, t.height = tx, I(0.5), tw, I(1.35)
        self.title(slide, spec.get("title"))
        items = self.items(spec.get("bullets") or spec.get("items") or spec.get("text") or [])
        if items:
            self.textbox(slide, Rect(tx, I(2.05), tw, H - I(2.05) - I(0.8)), items, size=self.sizes["body2"], bullets=any(i.get("bullet", True) for i in items) and len(items) > 1, fit=self.sizes["min_body"] / self.sizes["body2"], space_before=10, name="Text")
        self.drop_empty_placeholders(slide)
        return slide

    def s_quote(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "blank"), "quote")
        W = self.lay.W
        q = str(spec.get("quote") or spec.get("text") or "")
        mark_font = "+mj-lt"
        self.textbox(slide, Rect(I(1.0), I(0.55), I(2.0), I(1.75)), [{"text": "“", "level": 0}], size=96, color="accent1", font=mark_font, fit=1.0, name="Quote mark")
        n = len(plain(q))
        size = self.sizes["quote"] if n < 120 else self.sizes["quote"] * 0.85 if n < 220 else self.sizes["quote"] * 0.72
        self.textbox(slide, Rect(I(1.6), I(2.35), W - I(3.2), I(2.9)), [{"text": q, "level": 0}], size=size, color="text", font="+mj-lt", anchor="middle", fit=0.6, line=1.1, name="Quote")
        if spec.get("attribution"):
            self.textbox(slide, Rect(I(1.6), I(5.35), W - I(3.2), I(0.8)), [{"text": "— " + str(spec["attribution"]).lstrip("—– -"), "level": 0}], size=self.sizes["body3"], color="muted", fit=0.8, name="Attribution")
        t = spec.get("title")
        tph = self.ph(slide, "TITLE")
        if t and tph is not None:
            self.title(slide, t)
        self.drop_empty_placeholders(slide)
        return slide

    def s_kpi(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "kpi")
        self.title(slide, spec.get("title"))
        items = spec.get("items") or spec.get("kpis") or spec.get("metrics") or []
        if not items:
            raise UsageError("a kpi slide needs \"items\": [{\"value\": \"42%\", \"label\": \"…\"}]")
        area = self.lay.content_rect()
        n = len(items)
        rows = 1 if n <= 4 else 2
        per_row = math.ceil(n / rows)
        card_h = min(I(2.6), (area.h - I(0.35) * (rows - 1)) // rows)
        top = area.y + (area.h - (card_h * rows + I(0.35) * (rows - 1))) // 2 if rows == 1 else area.y
        vsize = self.sizes["kpi"] * (1.0 if per_row <= 3 else 0.82)
        for i, it in enumerate(items):
            if not isinstance(it, dict):
                it = {"value": str(it)}
            row, col = divmod(i, per_row)
            count = per_row if row < rows - 1 else n - per_row * (rows - 1)
            cells = Rect(area.x, top + row * (card_h + I(0.35)), area.w, card_h).cols(count, I(0.35))
            c = cells[col]
            accent = str(it.get("color", f"accent{(i % 6) + 1}" if spec.get("colorful") else "accent1"))
            self.rect(slide, c, "surface", radius=0.12, name=f"KPI card {i + 1}")
            self.rect(slide, Rect(c.x, c.y, I(0.09), c.h), accent, name=f"KPI accent {i + 1}")
            inner = c.inset(I(0.3), I(0.22))
            vh = int(inner.h * 0.48)
            self.textbox(slide, Rect(inner.x, inner.y, inner.w, vh), [{"text": str(it.get("value", "")), "level": 0}], size=vsize, bold=True, color=accent, font="+mj-lt", anchor="bottom", fit=0.4, name=f"KPI value {i + 1}", group=f"kpiv-{slide.slide_id}")
            self.textbox(slide, Rect(inner.x, inner.y + vh + I(0.06), inner.w, int(inner.h * 0.3)), [{"text": str(it.get("label", "")), "level": 0}], size=self.sizes["body2"], color="text", fit=0.7, name=f"KPI label {i + 1}", group=f"kpil-{slide.slide_id}")
            if it.get("delta") or it.get("note"):
                d = str(it.get("delta") or it.get("note"))
                role = "positive" if d.strip().startswith(("+", "▲", "↑")) else "negative" if d.strip().startswith(("-", "−", "▼", "↓")) else "muted"
                if it.get("good") is False and role == "positive" or it.get("good") is True and role == "negative":
                    role = "negative" if role == "positive" else "positive"
                self.textbox(slide, Rect(inner.x, inner.b - int(inner.h * 0.2), inner.w, int(inner.h * 0.2)), [{"text": d, "level": 0}], size=self.sizes["small"], color=role, bold=role != "muted", fit=0.75, name=f"KPI delta {i + 1}", group=f"kpid-{slide.slide_id}")
        self.drop_empty_placeholders(slide)
        return slide

    def s_table(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "table")
        self.title(slide, spec.get("title"))
        area = self.lay.content_rect()
        if spec.get("note"):
            area = Rect(area.x, area.y, area.w, area.h - I(0.55))
            self.textbox(slide, Rect(area.x, area.b + I(0.12), area.w, I(0.42)), [{"text": str(spec["note"]), "level": 0}], size=self.sizes["caption"], color="muted", fit=0.8, name="Table note")
        table = {k: spec[k] for k in ("columns", "rows", "align", "header", "widths", "font_size", "total", "highlight") if k in spec}
        if "data" in spec and "rows" not in table:
            table["rows"] = spec["data"]
        extra = self._table(slide, area, table, split_title=spec.get("title"), spec=spec)
        self.drop_empty_placeholders(slide)
        if extra:
            self.after.setdefault(slide.slide_id, []).extend(extra)
        return slide

    def s_chart(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "chart")
        self.title(slide, spec.get("title"))
        area = self.lay.content_rect()
        chart = spec.get("chart") or {k: spec[k] for k in ("categories", "series", "chart_type", "number_format", "labels", "legend") if k in spec}
        if "chart_type" in chart:
            chart["type"] = chart.pop("chart_type")
        note = spec.get("note") or spec.get("takeaway")
        if note:
            chart_r, note_r = area.cols(2, I(0.5), [2.1, 1])
            self._chart(slide, chart_r, chart)
            panel = self.rect(slide, note_r, "surface", radius=0.12, name="Takeaway")
            del panel
            self.rect(slide, Rect(note_r.x, note_r.y, note_r.w, I(0.08)), "accent1", name="Takeaway accent")
            self.textbox(slide, note_r.inset(I(0.3), I(0.35)), self.items(note), size=self.sizes["body3"], color="text", anchor="middle", fit=0.7, space_before=8, name="Takeaway text")
        else:
            self._chart(slide, area, chart)
        self.drop_empty_placeholders(slide)
        return slide

    def s_timeline(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "timeline")
        self.title(slide, spec.get("title"))
        items = [it if isinstance(it, dict) else {"title": str(it)} for it in (spec.get("items") or spec.get("events") or spec.get("milestones") or [])]
        if not items:
            raise UsageError("a timeline slide needs \"items\": [{\"label\": \"Q1\", \"title\": \"…\", \"text\": \"…\"}]")
        area = self.lay.content_rect()
        n = len(items)
        ly = area.y + int(area.h * 0.36)
        self.line(slide, area.x, ly, area.r, ly, "muted", 2.0)
        slot = area.w // n
        dot = I(0.26)
        for i, it in enumerate(items):
            cx = area.x + slot * i + slot // 2
            accent = str(it.get("color", "accent1"))
            self.rect(slide, Rect(cx - dot // 2, ly - dot // 2, dot, dot), accent, line="bg", line_w=2.5, shape="oval", name=f"Milestone {i + 1}")
            w = slot - I(0.15)
            if it.get("label") or it.get("date"):
                self.textbox(slide, Rect(cx - w // 2, area.y, w, ly - area.y - I(0.3)), [{"text": str(it.get("label") or it.get("date")), "level": 0}], size=self.sizes["body2"], bold=True, color=accent, align="center", anchor="bottom", fit=0.6, name=f"Timeline label {i + 1}", group=f"tll-{slide.slide_id}")
            body: list[dict[str, Any]] = []
            if it.get("title"):
                body.append({"text": str(it["title"]), "level": 0, "bold": True, "size": self.sizes["body3"]})
            if it.get("text"):
                body.append({"text": str(it["text"]), "level": 0, "size": self.sizes["small"], "color": "muted"})
            if body:
                self.textbox(slide, Rect(cx - w // 2, ly + I(0.35), w, area.b - ly - I(0.35)), body, size=self.sizes["small"], align="center", fit=0.6, space_before=4, name=f"Timeline text {i + 1}", group=f"tlt-{slide.slide_id}")
        self.drop_empty_placeholders(slide)
        return slide

    def s_process(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "process")
        self.title(slide, spec.get("title"))
        items = [it if isinstance(it, dict) else {"title": str(it)} for it in (spec.get("items") or spec.get("steps") or [])]
        if not items:
            raise UsageError("a process slide needs \"items\": [{\"title\": \"…\", \"text\": \"…\"}]")
        area = self.lay.content_rect()
        n = len(items)
        cells = area.cols(n, I(0.08))
        ch = I(1.15)
        for i, (it, c) in enumerate(zip(items, cells)):
            shape = "pentagon" if i == 0 else "chevron"
            fill = str(it.get("color", "accent1"))
            sh = self.rect(slide, Rect(c.x, area.y + I(0.1), c.w + (I(0.25) if i < n - 1 else 0), ch), fill, shape=shape, name=f"Step {i + 1}")
            label = str(it.get("label") or f"{i + 1:02d}")
            tf = sh.text_frame
            tf.margin_left = I(0.35) if i else I(0.2)
            self.write(tf, [{"text": label, "level": 0}], size=self.sizes["body2"], bold=True, color=self.pal.on(fill), align="center")
            body: list[dict[str, Any]] = []
            if it.get("title"):
                body.append({"text": str(it["title"]), "level": 0, "bold": True, "size": self.sizes["body3"]})
            if it.get("text"):
                body.append({"text": str(it["text"]), "level": 0, "size": self.sizes["small"], "color": "muted"})
            if body:
                self.textbox(slide, Rect(c.x + I(0.05), area.y + ch + I(0.35), c.w - I(0.1), area.h - ch - I(0.45)), body, size=self.sizes["small"], fit=0.6, space_before=6, name=f"Step text {i + 1}", group=f"pst-{slide.slide_id}")
        self.drop_empty_placeholders(slide)
        return slide

    def s_agenda(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "agenda")
        self.title(slide, spec.get("title") or "Agenda")
        raw = spec.get("items") or spec.get("bullets") or []
        items = [plain(x["text"]) if isinstance(x, dict) and "text" in x else str(x.get("title", "")) if isinstance(x, dict) else str(x) for x in raw]
        area = self.lay.content_rect()
        cols = 2 if len(items) > 5 else 1
        per = math.ceil(len(items) / cols) or 1
        colr = area.cols(cols, I(0.6))
        row_h = min(I(0.85), area.h // per)
        for i, txt in enumerate(items):
            c, rr = divmod(i, per)
            r = colr[c]
            y = r.y + rr * row_h
            self.textbox(slide, Rect(r.x, y, I(1.0), row_h), [{"text": f"{i + 1:02d}", "level": 0}], size=self.sizes["body"] * 1.3, bold=True, color="accent1", font="+mj-lt", anchor="middle", fit=0.7, name=f"Agenda number {i + 1}", group=f"agn-{slide.slide_id}")
            self.textbox(slide, Rect(r.x + I(1.05), y, r.w - I(1.05), row_h), [{"text": txt, "level": 0}], size=self.sizes["body"], color="text", anchor="middle", fit=0.7, name=f"Agenda item {i + 1}", group=f"agi-{slide.slide_id}")
            if rr < per - 1 and i < len(items) - 1:
                self.line(slide, r.x, y + row_h, r.r, y + row_h, "surface", 1.0)
        self.drop_empty_placeholders(slide)
        return slide

    def s_code(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only"), "code")
        self.title(slide, spec.get("title"))
        area = self.lay.content_rect()
        if spec.get("note"):
            area = Rect(area.x, area.y, area.w, area.h - I(0.55))
            self.textbox(slide, Rect(area.x, area.b + I(0.12), area.w, I(0.42)), [{"text": str(spec["note"]), "level": 0}], size=self.sizes["caption"], color="muted", fit=0.8, name="Code note")
        self._code(slide, area, str(spec.get("code", "")))
        self.drop_empty_placeholders(slide)
        return slide

    def _code(self, slide: Any, area: Rect, code: str) -> None:
        lines = code.expandtabs(4).rstrip("\n").split("\n") or [""]
        longest = max(len(x) for x in lines)
        # size so the longest line fits (monospace ≈ 0.6 em per character), within 10–20 pt
        size = max(10.0, min(20.0, (area.w / EMU_PER_PT - 40) / max(1, longest) / 0.6, (area.h / EMU_PER_PT - 30) / max(1, len(lines)) / 1.2))
        # the panel hugs the code (short snippets do not sit in a big empty box)
        need = int((len(lines) * size * 1.2 + 0.4 * 72 + 8) * EMU_PER_PT)
        box_r = Rect(area.x, area.y, area.w, max(min(area.h, need), I(1.2)))
        box = self.rect(slide, box_r, "surface", radius=0.08, name="Code")
        tf = box.text_frame
        tf.word_wrap = True
        from pptx.enum.text import MSO_ANCHOR

        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = I(0.25)
        tf.margin_top = tf.margin_bottom = I(0.2)
        self.write(tf, [{"text": ln.replace("\\", "\\\\").replace("*", "\\*").replace("`", "\\`").replace("_", "\\_").replace("[", "\\[").replace("~", "\\~") or " ", "level": 0} for ln in lines], size=size, color="text", font=self.pal.mono_font, line=1.0, align="left")
        self.jobs.append(FitJob(slide, box, 0.7))

    def s_custom(self, spec: dict[str, Any]) -> Any:
        slide = self.new_slide(self.layout_for(spec, "title_only" if spec.get("title") else "blank"), "custom")
        self.title(slide, spec.get("title"))
        self.add_elements(slide, spec.get("elements") or spec.get("shapes") or [])
        self.drop_empty_placeholders(slide)
        return slide

    def add_elements(self, slide: Any, elements: list[dict[str, Any]]) -> None:
        """Free-form elements at given positions: text, rect/card/oval (with text), line, image, table, chart."""
        from _ooxml import parse_length

        W, H = self.lay.W, self.lay.H
        for i, el in enumerate(elements):
            if not isinstance(el, dict):
                raise UsageError(f"element {i + 1} must be an object")
            k = str(el.get("kind", el.get("type", "text")))
            try:
                r = Rect(parse_length(el.get("x", 0), W), parse_length(el.get("y", 0), H), parse_length(el.get("w", 2), W), parse_length(el.get("h", 1), H))
            except ValueError as e:
                raise UsageError(f"element {i + 1}: {e}") from e
            if k == "text":
                self.textbox(slide, r, el.get("text", ""), size=float(el.get("size", self.sizes["body3"])), color=str(el.get("color", "text")), bold=el.get("bold"), align=el.get("align", "left"), anchor=el.get("anchor", "top"), bullets=bool(el.get("bullets")), fit=float(el.get("fit", 0.7)), name=el.get("name"), font=el.get("font"))
            elif k in ("rect", "box", "card", "oval", "circle"):
                sh = self.rect(slide, r, el.get("fill", "surface"), el.get("line"), radius=el.get("radius", 0.12 if k == "card" else None), shape="oval" if k in ("oval", "circle") else None, name=el.get("name"))
                if el.get("text"):
                    fill = str(el.get("fill", "surface"))
                    self.write(sh.text_frame, self.items(el["text"]), size=float(el.get("size", self.sizes["body3"])), color=str(el.get("color", self.pal.on(fill) if fill != "none" else "text")), align=el.get("align", "center"), bold=el.get("bold"))
                    self.jobs.append(FitJob(slide, sh, 0.7))
            elif k == "line":
                self.line(slide, r.x, r.y, r.x + r.w, r.y + r.h, str(el.get("color", "muted")), float(el.get("width", 1.5)))
            elif k == "image":
                self._picture(slide, el["path"] if "path" in el else el["image"], r, el.get("fit", "fit"), name=el.get("name"))
            elif k == "table":
                self._table(slide, r, el, split=False)
            elif k == "chart":
                self._chart(slide, r, el.get("chart", el))
            else:
                raise UsageError(f"element {i + 1}: unknown kind '{k}' (text, rect, card, oval, line, image, table, chart)")

    s_blank = s_custom

    # ── tables ──────────────────────────────────────────────────────────

    def _table(self, slide: Any, area: Rect, t: dict[str, Any], split_title: Any = None, spec: dict[str, Any] | None = None, split: bool = True) -> list[Any]:
        """A table in `area`: the largest size whose columns fit; with split, rows that do not fit continue on
        new slides; without (a custom element or a column), the text shrinks until the table fits its box."""
        from _typtext import Measurer

        cols = [str(c) for c in (t.get("columns") or t.get("header") or [])]
        rows = [[("" if v is None else str(v)) for v in r] for r in (t.get("rows") or [])]
        if not cols and rows:
            cols, rows = rows[0], rows[1:]
        if not cols:
            raise UsageError("a table needs \"columns\": [...] and \"rows\": [[...], ...]")
        n = len(cols)
        rows = [(r + [""] * n)[:n] for r in rows]
        align = list(t.get("align") or [])
        # A cell looks numeric; blank runs are collapsed first (the adjacent \s* runs backtrack on long ones).
        numeric = []
        for j in range(n):
            vals = [r[j] for r in rows if r[j].strip()]
            isnum = bool(vals) and sum(1 for v in vals if re.fullmatch(r"[\s(+\-−$€£¥]*[\d.,]+\s*[%kKmMbB×x)]*\s*", re.sub(r"\s+", " ", plain(v)))) >= 0.7 * len(vals)
            numeric.append(isnum)
            if j >= len(align):
                align.append("right" if isnum else "left")
        # measure natural widths and heights at candidate sizes (one Typst pass)
        sizes = [float(t["font_size"])] if t.get("font_size") else ([20.0, 18.0] if len(rows) <= 8 else []) + [16.0, 14.0, 12.0, 11.0]
        from _ooxml import BodyProps, Para, ParaStyle, Run, RunStyle, TextFrame

        def frame(text: str, size: float, bold: bool) -> TextFrame:
            runs = [Run(r["text"], RunStyle(font=self.pal.mono_font if r.get("code") else self.pal.body_font, size=size, bold=bold or bool(r.get("bold")), italic=bool(r.get("italic")))) for r in inline_runs(text)] or [Run(" ", RunStyle(font=self.pal.body_font, size=size))]
            return TextFrame([Para(runs, ParaStyle(), size)], BodyProps())

        m = Measurer(self.fonts)
        for s in sizes:
            for j, c in enumerate(cols):
                m.add_frame(f"h{s}:{j}", frame(c, s, True), None)
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    m.add_frame(f"c{s}:{i}:{j}", frame(v, s, False), None)
        nat = m.run()
        pad_w = 2 * 0.1 * 72  # cell left+right margins in pt
        total_w = area.w / EMU_PER_PT
        chosen = sizes[-1]
        widths: list[float] = []
        avail_h = area.h / EMU_PER_PT
        for s in sizes:
            natw = [max([nat[f"h{s}:{j}"][0]] + [nat[f"c{s}:{i}:{j}"][0] for i in range(len(rows))]) + pad_w for j in range(n)]
            widths = _alloc_widths(natw, total_w, t.get("widths"))
            one_line_h = (len(rows) + 1) * (nat[f"h{s}:0"][1] + 2 * 0.06 * 72 + 4)
            fits_w = all(nw <= w + 0.5 for nw, w in zip(natw, widths))
            if (fits_w and (s <= 16 or one_line_h <= avail_h * 0.9)) or s == sizes[-1]:
                chosen = s
                break
        # now heights at the chosen size with wrapping
        pad_h = 2 * 0.06 * 72 + 4
        avail = area.h / EMU_PER_PT
        smaller = [x for x in sizes if x <= chosen] if not split and not t.get("font_size") else [chosen]
        for s in smaller:
            m2 = Measurer(self.fonts)
            for j, c in enumerate(cols):
                m2.add_frame(f"h:{j}", frame(c, s, True), widths[j] - pad_w)
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    m2.add_frame(f"c:{i}:{j}", frame(v, s, False), widths[j] - pad_w)
            hs = m2.run()
            head_h = max(hs[f"h:{j}"][1] for j in range(n)) + pad_h
            row_h = [max(hs[f"c:{i}:{j}"][1] for j in range(n)) + pad_h for i in range(len(rows))]
            chosen = s
            if head_h + sum(row_h) <= avail + 1:
                break
        if not split:
            if head_h + sum(row_h) > avail + 1:
                self.warnings.append(f"a table ('{plain(cols[0])[:30]}' …) is {(head_h + sum(row_h)) / 72:.1f} in tall at {chosen:g} pt, taller than its {avail / 72:.1f} in box: it runs past the box")
            self._table_shape(slide, area, cols, rows, widths, head_h, row_h, chosen, align, t)
            return []
        chunks: list[list[int]] = [[]]
        used = head_h
        for i, h in enumerate(row_h):
            if chunks[-1] and used + h > avail:
                chunks.append([])
                used = head_h
            chunks[-1].append(i)
            used += h
        extra_slides = []
        for ci, chunk in enumerate(chunks):
            target = slide
            if ci > 0:
                target = self.new_slide(slide.slide_layout, "table")
                if split_title:
                    self.title(target, f"{plain(str(split_title))} (cont.)")
                self.drop_empty_placeholders(target)
                if spec and spec.get("notes"):
                    target.notes_slide.notes_text_frame.text = plain(str(spec["notes"]))
                self.content_slides.append(target)
                extra_slides.append(target)
            self._table_shape(target, area, cols, [rows[i] for i in chunk], widths, head_h, [row_h[i] for i in chunk], chosen, align, t)
        if len(chunks) > 1:
            self.warnings.append(f"table '{plain(str(split_title or 'table'))}' continued over {len(chunks)} slides")
        return extra_slides

    def _table_shape(self, slide: Any, area: Rect, cols: list[str], rows: list[list[str]], widths: list[float], head_h: float, row_h: list[float], size: float, align: list[str], t: dict[str, Any]) -> Any:
        from pptx.enum.text import MSO_ANCHOR

        total_h = int((head_h + sum(row_h)) * EMU_PER_PT)
        gf = slide.shapes.add_table(len(rows) + 1, len(cols), area.x, area.y, area.w, min(total_h, area.h))
        gf.name = "Table"
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        sid = tblPr.find("{%s}tableStyleId" % A_NS)
        if sid is None:
            sid = _sub(tblPr, "a:tableStyleId")
        sid.text = NO_TABLE_STYLE
        tbl.first_row = True
        tbl.horz_banding = False
        acc = 0
        for j, w in enumerate(widths):
            cw = int(w * EMU_PER_PT) if j < len(widths) - 1 else area.w - acc
            tbl.columns[j].width = cw
            acc += cw
        tbl.rows[0].height = int(head_h * EMU_PER_PT)
        for i, h in enumerate(row_h):
            tbl.rows[i + 1].height = int(h * EMU_PER_PT)
        hl_rows = set(int(x) for x in (t.get("highlight") or []) if str(x).lstrip("-").isdigit())
        total_last = bool(t.get("total"))
        for i in range(len(rows) + 1):
            for j in range(len(cols)):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = I(0.1)
                cell.margin_top = cell.margin_bottom = I(0.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                header = i == 0
                is_total = total_last and i == len(rows)
                text = cols[j] if header else rows[i - 1][j]
                if header:
                    fill, color, bold = "accent1", self.pal.on("accent1"), True
                elif is_total:
                    fill, color, bold = "surface", "text", True
                elif (i - 1) in hl_rows:
                    fill, color, bold = "accent1", "text", True
                    fill = "surface"
                else:
                    fill, color, bold = ("bg" if i % 2 == 1 else "surface"), "text", False
                cell.fill.solid()
                self.pal.apply(cell.fill.fore_color, fill)
                self.write(cell.text_frame, [{"text": text, "level": 0}], size=size, color=color, bold=bold if bold else None, align=align[j] if j < len(align) else "left")
                _cell_border(cell, "B", self.pal, "accent1" if header else "surface" if not is_total else "muted", 1.5 if header else 0.75)
                if is_total:
                    _cell_border(cell, "T", self.pal, "muted", 1.0)
        return gf

    # ── charts ──────────────────────────────────────────────────────────

    def _chart(self, slide: Any, area: Rect, c: dict[str, Any]) -> Any:
        from pptx.chart.data import CategoryChartData, XyChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
        from pptx.util import Pt

        kind = str(c.get("type", "column")).lower()
        types = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "column-stacked": XL_CHART_TYPE.COLUMN_STACKED, "column-stacked-100": XL_CHART_TYPE.COLUMN_STACKED_100,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED, "bar-stacked": XL_CHART_TYPE.BAR_STACKED, "bar-stacked-100": XL_CHART_TYPE.BAR_STACKED_100,
            "line": XL_CHART_TYPE.LINE_MARKERS, "line-plain": XL_CHART_TYPE.LINE, "area": XL_CHART_TYPE.AREA, "area-stacked": XL_CHART_TYPE.AREA_STACKED,
            "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "scatter": XL_CHART_TYPE.XY_SCATTER, "scatter-lines": XL_CHART_TYPE.XY_SCATTER_LINES,
            "radar": XL_CHART_TYPE.RADAR_MARKERS,
        }
        if kind not in types:
            raise UsageError(f"unknown chart type '{kind}' (types: {', '.join(types)})")
        series = c.get("series") or []
        if not series:
            raise UsageError("a chart needs \"series\": [{\"name\": \"…\", \"values\": [...]}]")
        fmt = c.get("number_format")
        if kind.startswith("scatter"):
            cd = XyChartData()
            for s in series:
                ser = cd.add_series(str(s.get("name", "")))
                pts = s.get("points") or list(zip(s.get("x", []), s.get("values", s.get("y", []))))
                for x, y in pts:
                    if x is not None and y is not None:
                        ser.add_data_point(float(x), float(y))
        else:
            cats = c.get("categories") or []
            cd = CategoryChartData(number_format=fmt) if fmt else CategoryChartData()
            cd.categories = [str(x) for x in cats]
            for s in series:
                vals = s.get("values") or []
                if len(vals) != len(cats):
                    raise UsageError(f"chart series '{s.get('name', '')}' has {len(vals)} values for {len(cats)} categories")
                cd.add_series(str(s.get("name", "")), [None if v is None else float(v) for v in vals])
        gf = slide.shapes.add_chart(types[kind], area.x, area.y, area.w, area.h, cd)
        gf.name = "Chart"
        ch = gf.chart
        bg, tx = self.pal.rgb("bg"), self.pal.rgb("text")
        grid_rgb = RGBColor.from_string(Color(round(bg.r * 0.85 + tx.r * 0.15), round(bg.g * 0.85 + tx.g * 0.15), round(bg.b * 0.85 + tx.b * 0.15)).hex)
        ch.font.size = Pt(self.sizes["small"])
        self.pal.apply(ch.font.color, "text", direct=True)
        ch.font.name = self.pal.body_font
        if c.get("title"):
            ch.has_title = True
            ch.chart_title.text_frame.text = plain(str(c["title"]))
            f = ch.chart_title.text_frame.paragraphs[0].runs[0].font
            f.size, f.bold = Pt(self.sizes["body3"]), True
            self.pal.apply(f.color, "text", direct=True)
        else:
            ch.has_title = False
        radial = kind in ("pie", "doughnut")
        if radial and any(isinstance(v, (int, float)) and v < 0 for v in ((series[0] or {}).get("values") or [])):
            self.warnings.append(f"the {kind} chart on {_slide_label(slide)} has negative values, which a {kind} chart cannot show (PowerPoint draws them as positive slices); use a bar chart")
        legend = c.get("legend", "bottom" if (len(series) > 1 or radial) else "none")
        if legend and legend != "none":
            ch.has_legend = True
            ch.legend.position = {"bottom": XL_LEGEND_POSITION.BOTTOM, "right": XL_LEGEND_POSITION.RIGHT, "top": XL_LEGEND_POSITION.TOP, "left": XL_LEGEND_POSITION.LEFT}.get(str(legend), XL_LEGEND_POSITION.BOTTOM)
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(self.sizes["small"])
            self.pal.apply(ch.legend.font.color, "text", direct=True)
        else:
            ch.has_legend = False
        # colours by theme role (accent1-6), so re-theming the deck recolours its charts
        colors = [str(x) for x in c["colors"]] if c.get("colors") else [f"accent{i}" for i in range(1, 7)]
        for x in colors:
            self.pal.spec(x)  # validates each colour up front
        plot = ch.plots[0]
        if radial:
            ser = plot.series[0]
            for j in range(len(c.get("categories") or [])):
                pt_ = ser.points[j]
                pt_.format.fill.solid()
                self.pal.apply(pt_.format.fill.fore_color, colors[j % len(colors)], direct=True)
                self.pal.apply(pt_.format.line.color, "bg", direct=True)
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(self.sizes["small"])
            dl.font.bold = True
            if c.get("labels") == "value":
                dl.show_value = True
                if fmt:
                    dl.number_format, dl.number_format_is_linked = fmt, False
            else:
                dl.show_percentage = True
                dl.show_value = False
                dl.number_format, dl.number_format_is_linked = "0%", False
            self.pal.apply(dl.font.color, "bg", direct=True)
            if kind == "doughnut":
                hs = plot._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}holeSize")
                if hs is not None:
                    hs.set("val", "58")
            return gf
        try:
            plot.vary_by_categories = False
        except Exception:  # noqa: BLE001
            pass
        if kind.startswith(("column", "bar")):
            plot.gap_width = 60 if len(series) > 1 else 80
            if "stacked" in kind:
                plot.overlap = 100
        for i, s in enumerate(plot.series):
            col = colors[i % len(colors)]
            if kind.startswith(("line", "scatter", "radar")):
                self.pal.apply(s.format.line.color, col, direct=True)
                s.format.line.width = Pt(2.75)
                s.smooth = False
                try:
                    from pptx.enum.chart import XL_MARKER_STYLE

                    s.marker.style = XL_MARKER_STYLE.CIRCLE
                    s.marker.size = 7
                    s.marker.format.fill.solid()
                    self.pal.apply(s.marker.format.fill.fore_color, col, direct=True)
                    self.pal.apply(s.marker.format.line.color, col, direct=True)
                except Exception:  # noqa: BLE001
                    pass
                if kind == "scatter":
                    s.format.line.fill.background()
            else:
                s.format.fill.solid()
                self.pal.apply(s.format.fill.fore_color, col, direct=True)
                s.format.line.fill.background()
        if c.get("labels"):
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(self.sizes["caption"])
            self.pal.apply(dl.font.color, "text", direct=True)
            if fmt:
                dl.number_format, dl.number_format_is_linked = fmt, False
            if kind in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        if kind != "radar":
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = grid_rgb
            va.major_gridlines.format.line.width = Pt(0.75)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(self.sizes["caption"])
            self.pal.apply(va.tick_labels.font.color, "muted", direct=True)
            if fmt:
                va.tick_labels.number_format, va.tick_labels.number_format_is_linked = fmt, False
            ca = ch.category_axis
            ca.format.line.color.rgb = grid_rgb
            ca.tick_labels.font.size = Pt(self.sizes["small"])
            self.pal.apply(ca.tick_labels.font.color, "text", direct=True)
            ca.has_major_gridlines = False
            if kind.startswith("bar"):
                ca.reverse_order = True  # first category on top, as people read a list
                from pptx.enum.chart import XL_TICK_LABEL_POSITION

                va.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        return gf

    # ── fitting ─────────────────────────────────────────────────────────

    def run_fits(self) -> None:
        """Measures every registered text body once; shrinks what overflows; continues long bullet lists."""
        if not self.jobs:
            return
        jobs, self.jobs = self.jobs, []
        results = self._measure(jobs)
        followups: list[FitJob] = []
        scales: list[float | None] = []
        for job, (fits_at, frame, inner_h, group_h, words) in zip(jobs, results):
            too_wordy = bool(job.groups and len(job.groups) > 1 and words + self._title_words(job.slide) > WORDS_MAX)
            if fits_at is not None and not too_wordy:
                scales.append(fits_at)
                continue
            if job.groups and len(job.groups) > 1 and group_h:
                chunks = _chunk(job.groups, group_h, inner_h)
                if len(chunks) > 1:
                    new = self._continue(job, chunks)
                    followups.extend(new)
                    scales.append(None)
                    continue
            if fits_at is not None:
                scales.append(fits_at)
                continue
            scales.append(job.min_scale)
            self._overflow_warning(job, frame)
        self._apply(jobs, results, scales)
        if followups:
            results = self._measure(followups)
            scales = []
            for job, (fits_at, frame, _h, _g, _w) in zip(followups, results):
                scales.append(fits_at if fits_at is not None else job.min_scale)
                if fits_at is None:
                    self._overflow_warning(job, frame)
            self._apply(followups, results, scales)

    def _apply(self, jobs: list[FitJob], results: list[Any], scales: list[float | None]) -> None:
        """Writes each job's scale; jobs in a group all take the group's smallest scale, so they match."""
        group_min: dict[str, float] = {}
        for job, s in zip(jobs, scales):
            if job.group and s is not None:
                group_min[job.group] = min(group_min.get(job.group, 99.0), s)
        for job, res, s in zip(jobs, results, scales):
            if s is None:
                continue
            if job.group:
                s = group_min.get(job.group, s)
            if abs(s - 1.0) > 0.001:
                _apply_scale(job.shape, res[1], s)

    def _overflow_warning(self, job: FitJob, frame: Any) -> None:
        name = _slide_label(job.slide)
        why = getattr(frame, "_long_word", "")
        if why:
            self.warnings.append(f"the word '{why[:40]}' in '{job.shape.name}' on {name} is wider than its box even at the smallest size; shorten it (PowerPoint would break it mid-word)")
        else:
            self.warnings.append(f"text in '{job.shape.name}' on {name} still overflows at the smallest size; shorten it or split the slide")

    def _title_words(self, slide: Any) -> int:
        t = self.ph(slide, "TITLE", "CENTER_TITLE")
        return len(re.findall(r"\w+", t.text_frame.text)) if t is not None and t.has_text_frame else 0

    def _measure(self, jobs: list[FitJob]) -> list[tuple[float | None, Any, float, list[tuple[float, float]] | None, int]]:
        """Two passes: full size, the minimum and (for sparse bodies) larger sizes first; then the sizes in between
        only for text that fits at the minimum but not at full size. A scale fits when the text's height fits the
        box and its widest word fits the line (else PowerPoint breaks the word mid-word)."""
        from _typtext import Measurer

        d = pres_defaults(self.prs)
        ctxs: dict[int, SlideCtx] = {}
        m = Measurer(self.fonts)
        prepared = []
        for ji, job in enumerate(jobs):
            key = id(job.slide)
            if key not in ctxs:
                ctxs[key] = SlideCtx(self.prs, job.slide, d, 0)
            ctx = ctxs[key]
            el = job.shape._element
            tb = el.find("{%s}txBody" % P_NS)
            frame = TextResolver(ctx, el, tb).frame()
            frame.body.autofit = "none"
            w = (job.shape.width or 0) / EMU_PER_PT - frame.body.l - frame.body.r
            h = (job.shape.height or 0) / EMU_PER_PT - frame.body.t - frame.body.b
            grow = [g for g in (1.3, 1.2, 1.1) if g <= job.max_grow + 1e-6] if job.kind == "body" and not job.no_grow else []
            first = grow + [1.0] + ([round(job.min_scale, 3)] if job.min_scale < 0.999 else [])
            for s in first:
                m.add_frame(f"{ji}:{s}", _scaled(frame, s), w)
            m.add_words(f"{ji}:w", frame)
            if job.groups:
                idx = 0
                for gi, g in enumerate(job.groups):
                    sub = frame.paras[idx: idx + len(g)]
                    idx += len(g)
                    for s in (1.0, job.min_scale):
                        pad = copy.copy(frame)
                        pad.paras = [_empty_para()] + sub  # measured as a non-first paragraph, so space-before counts
                        m.add_frame(f"{ji}:g{gi}:{s}", _scaled(pad, s), w)
            prepared.append((ji, job, frame, w, h, grow))
        sizes = m.run()
        out: list[Any] = []
        second = Measurer(self.fonts)
        need_second = []
        for ji, job, frame, w, h, grow in prepared:
            ratio, word = m.word_ratio(sizes, f"{ji}:w", w)
            wmax = (0.995 / ratio) if ratio > 0 else 99.0  # the largest scale at which the widest word still fits
            fits = None
            for s in grow:
                # sparse slides read better with larger text, as long as it stays well inside the box
                if sizes[f"{ji}:{s}"][1] <= h * 0.62 and s <= wmax:
                    fits = s
                    break
            if fits is None and sizes[f"{ji}:1.0"][1] <= h + 0.5 and wmax >= 1.0:
                fits = 1.0
            gh = None
            if job.groups:
                gh = [(sizes[f"{ji}:g{gi}:1.0"][1] - 0.01, sizes[f"{ji}:g{gi}:{job.min_scale}"][1]) for gi in range(len(job.groups))]
            lo = round(job.min_scale, 3)
            if fits is None and job.min_scale < 0.999 and sizes[f"{ji}:{lo}"][1] <= h + 0.5 and wmax >= lo:
                mids = [s for s in _scales(job.min_scale) if job.min_scale + 1e-6 < s < 0.999 and s <= wmax]
                if wmax < 1.0 and wmax > lo + 0.005:
                    mids = sorted(set(mids) | {math.floor(wmax * 100) / 100}, reverse=True)
                for s in mids:
                    second.add_frame(f"{ji}:{s}", _scaled(frame, s), w)
                need_second.append((len(out), ji, h, mids, job.min_scale))
                fits = lo
            if fits is None and wmax < lo:
                frame._long_word = word  # type: ignore[attr-defined]
            words = sum(len(re.findall(r"\w+", p.text)) for p in frame.paras)
            out.append([fits, frame, h, gh, words])
        if need_second:
            sizes2 = second.run()
            for pos, ji, h, mids, lo in need_second:
                for s in mids:
                    if sizes2[f"{ji}:{s}"][1] <= h + 0.5:
                        out[pos][0] = s
                        break
        return [tuple(x) for x in out]  # type: ignore[misc]

    def _continue(self, job: FitJob, chunks: list[list[int]]) -> list[FitJob]:
        """Keeps the first chunk of groups on the slide and adds continuation slides for the rest."""
        slide = job.slide
        groups = job.groups or []
        title_ph = self.ph(slide, "TITLE", "CENTER_TITLE")
        title = title_ph.text_frame.text if title_ph is not None else ""
        items_of = lambda ch: [it for gi in ch for it in groups[gi]]  # noqa: E731
        self.write(job.shape.text_frame, items_of(chunks[0]), placeholder=True)
        _norm_autofit(job.shape)
        out = [FitJob(slide, job.shape, job.min_scale, None, job.spec, "body", no_grow=True)]
        for k, ch in enumerate(chunks[1:], start=2):
            spec = dict(job.spec or {})
            new = self.new_slide(slide.slide_layout, "bullets")
            if title:
                self.title(new, f"{title} ({k}/{len(chunks)})" if len(chunks) > 2 else f"{title} (cont.)")
            body = self.ph(new, "BODY", "OBJECT")
            if body is None:
                continue
            body.left, body.top, body.width, body.height = job.shape.left, job.shape.top, job.shape.width, job.shape.height
            self.write(body.text_frame, items_of(ch), placeholder=True)
            _norm_autofit(body)
            out.append(FitJob(new, body, job.min_scale, None, spec, "body", no_grow=True))
            self.drop_empty_placeholders(new)
            if spec.get("notes"):
                new.notes_slide.notes_text_frame.text = plain(str(spec["notes"]))
            self.content_slides.append(new)
            self.after.setdefault(slide.slide_id, []).append(new)
        if title_ph is not None and len(chunks) > 2:
            self.write(title_ph.text_frame, [{"text": f"{title} (1/{len(chunks)})", "level": 0}])
        self.warnings.append(f"'{title or 'untitled'}' was too long for one slide and continues on {len(chunks) - 1} more")
        # the new titles need fitting too
        for f in list(self.jobs):
            out.append(f)
        self.jobs = []
        return out

    # ── finishing ───────────────────────────────────────────────────────

    def finish(self, deck: dict[str, Any]) -> None:
        self.run_fits()
        self.layout_warnings()
        self._order()
        footer = deck.get("footer")
        numbers = deck.get("slide_numbers", True)
        ids = {s.slide_id for s in self.content_slides}
        for slide in self.prs.slides:
            if slide.slide_id not in ids:
                continue
            if footer:
                self._footer(slide, str(footer))
            if numbers:
                self._number(slide)
        cp = self.prs.core_properties
        import datetime as _dt

        now = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.title = plain(str(deck.get("title") or _first_title(self.prs) or ""))[:250]
        if deck.get("author"):
            cp.author = str(deck["author"])
        cp.last_modified_by = str(deck.get("author") or "Desk")
        if deck.get("subject"):
            cp.subject = str(deck["subject"])
        if deck.get("keywords"):
            cp.keywords = str(deck["keywords"]) if not isinstance(deck["keywords"], list) else ", ".join(map(str, deck["keywords"]))
        if self.mode == "theme" or not self._kept_slides:
            cp.created = now
        cp.modified = now
        cp.revision = 1

    def layout_warnings(self) -> None:
        if self.missing_layouts:
            names = ", ".join(f"'{k}'" + (f" ({v}×)" if v > 1 else "") for k, v in self.missing_layouts.items())
            have = ", ".join(sorted({lay.name for lay in self.lay.all}))
            self.warnings.append(f"no layout named {names} here, so those slides got the layout that fits their content (layouts: {have})")
            self.missing_layouts = {}

    def _order(self) -> None:
        if not self.after:
            return
        lst = self.prs.slides._sldIdLst
        by_id = {int(e.get("id")): e for e in lst}
        for sid, extra in self.after.items():
            anchor = by_id.get(sid)
            if anchor is None:
                continue
            for s in extra:
                e = by_id.get(s.slide_id)
                if e is None:
                    continue
                lst.remove(e)
                anchor.addnext(e)
                anchor = e

    def _footer(self, slide: Any, text: str) -> None:
        x0, x1 = self.footer_span.get(slide.slide_id, (I(0.75), self.lay.W - I(0.75)))
        r = Rect(x0, self.lay.H - I(0.5), min(I(7.5), x1 - x0 - I(1.4)), I(0.3))
        self.textbox(slide, r, [{"text": text, "level": 0}], size=self.sizes["footer"], color="muted", anchor="middle", fit=None, name="Footer")

    def _number(self, slide: Any) -> None:
        from lxml import etree

        x0, x1 = self.footer_span.get(slide.slide_id, (I(0.75), self.lay.W - I(0.75)))
        r = Rect(x1 - I(1.2), self.lay.H - I(0.5), I(1.2), I(0.3))
        tb = self.textbox(slide, r, [{"text": " ", "level": 0}], size=self.sizes["footer"], color="muted", align="right", anchor="middle", fit=None, name="Slide number")
        p = tb.text_frame.paragraphs[0]._p
        run = p.find("{%s}r" % A_NS)
        n = list(self.prs.slides).index(slide) + 1
        fld = etree.SubElement(p, "{%s}fld" % A_NS, id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
        if run is not None:
            rpr = run.find("{%s}rPr" % A_NS)
            if rpr is not None:
                fld.append(copy.deepcopy(rpr))
            p.remove(run)
        t = etree.SubElement(fld, "{%s}t" % A_NS)
        t.text = str(n)
        end = p.find("{%s}endParaRPr" % A_NS)
        if end is not None:
            p.remove(end)
            p.append(end)

    def cleanup(self) -> None:
        for p in self._tmp:
            try:
                p.unlink()
            except OSError:
                pass


# ── helpers ─────────────────────────────────────────────────────────────


def _slide_label(slide: Any) -> str:
    try:
        n = list(slide.part.package.presentation_part.presentation.slides).index(slide) + 1
        return f"slide {n}"
    except Exception:  # noqa: BLE001
        return "a slide"


def _first_title(prs: Any) -> str:
    from _deck import slide_title

    for s in prs.slides:
        t = slide_title(s)
        if t:
            return t
    return ""


def _scales(min_scale: float) -> list[float]:
    out = [1.0]
    s = 1.0
    while s - 0.06 >= min_scale - 1e-6:
        s = round(s - 0.06, 3)
        out.append(s)
    if out[-1] > min_scale + 1e-6:
        out.append(round(min_scale, 3))
    return out


def _scaled(frame: Any, s: float) -> Any:
    """A copy of a text frame with every size and spacing scaled by s (shallow where nothing changes)."""
    from dataclasses import replace as dc_replace

    from _ooxml import Para, Run, TextFrame

    if abs(s - 1.0) < 1e-6:
        return frame
    paras = []
    for p in frame.paras:
        ps = dc_replace(p.style, spc_before=p.style.spc_before * s, spc_after=p.style.spc_after * s, line_pts=p.style.line_pts * s if p.style.line_pts else p.style.line_pts)
        paras.append(Para([Run(r.text, dc_replace(r.style, size=r.style.size * s)) for r in p.runs], ps, p.end_size * s))
    return TextFrame(paras, frame.body)


def _empty_para() -> Any:
    from _ooxml import Para, ParaStyle

    return Para([], ParaStyle(), 0.01)


def _apply_scale(shape: Any, frame: Any, s: float) -> None:
    """Writes explicit, scaled font sizes and spacing (works in every app, unlike a stored autofit scale)."""
    from pptx.util import Pt

    tf = shape.text_frame
    for p, fp in zip(tf.paragraphs, frame.paras):
        rr = [r for r in p.runs]
        fr = [r for r in fp.runs if r.text != "\n"]
        for run, res in zip(rr, fr):
            run.font.size = Pt(round(res.style.size * s * 2) / 2)
        if fp.style.spc_before:
            p.space_before = Pt(round(fp.style.spc_before * s, 1))
        end = p._p.find("{%s}endParaRPr" % A_NS)
        if end is not None:
            end.set("sz", str(int(fp.end_size * s * 100)))


def _norm_autofit(ph: Any) -> None:
    """Lets PowerPoint keep shrinking text in a body placeholder if someone adds to it later."""
    bp = ph._element.find("{%s}txBody/{%s}bodyPr" % (P_NS, A_NS))
    if bp is None:
        return
    for c in list(bp):
        if local(c) in ("noAutofit", "normAutofit", "spAutoFit"):
            bp.remove(c)
    _sub(bp, "a:normAutofit")


def _groups(items: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    groups: list[list[dict[str, Any]]] = []
    for it in items:
        if it.get("level", 0) == 0 or not groups:
            groups.append([it])
        else:
            groups[-1].append(it)
    return groups


CHUNK_WORDS = 80  # continuation slides stay under pptx_lint's 90-word limit
WORDS_MAX = 90  # pptx_lint's default --max-words: a list whose slide would exceed it continues on another slide


def _chunk(groups: list[list[dict[str, Any]]], gh: list[tuple[float, float]], avail: float) -> list[list[int]]:
    """Splits bullet groups into slides of balanced height (and at most CHUNK_WORDS words) that fit at full size
    where possible."""
    heights = [h for h, _ in gh]
    words = [sum(len(str(it.get("text", "")).split()) for it in g) for g in groups]
    total = sum(heights)
    n = max(1, math.ceil(total / (avail * 0.98)), math.ceil(sum(words) / CHUNK_WORDS))
    target = total / n
    chunks: list[list[int]] = [[]]
    used = 0.0
    used_words = 0
    for i, h in enumerate(heights):
        if chunks[-1] and (used + h > avail or used_words + words[i] > CHUNK_WORDS or (used >= target * 0.95 and len(chunks) < n)):
            chunks.append([])
            used = 0.0
            used_words = 0
        chunks[-1].append(i)
        used += h
        used_words += words[i]
    return chunks


def _alloc_widths(natural: list[float], total: float, given: Any = None) -> list[float]:
    if isinstance(given, list) and len(given) == len(natural):
        s = sum(float(x) for x in given) or 1.0
        return [total * float(x) / s for x in given]
    n = len(natural)
    minw = min(total / n, 60.0)
    need = sum(natural)
    if need <= total:
        extra = total - need
        return [w + extra * (w / need) for w in natural]
    # shrink the widest columns first (they wrap), keeping narrow ones intact
    widths = list(natural)
    cap = max(widths)
    while sum(min(w, cap) for w in widths) > total and cap > minw:
        cap -= 2
    widths = [max(minw, min(w, cap)) for w in widths]
    scale = total / sum(widths)
    return [w * scale for w in widths]


def _bullet(p: Any, kind: str | None, pal: Palette, level: int, placeholder: bool) -> None:
    pPr = p._p.get_or_add_pPr()
    for c in list(pPr):
        if local(c).startswith("bu"):
            pPr.remove(c)
    if kind is None:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        _sub(pPr, "a:buNone")
    elif kind == "num":
        _sub(pPr, "a:buAutoNum", type="arabicPeriod")
        if not placeholder:
            pPr.set("marL", str(I(0.35 + 0.35 * level)))
            pPr.set("indent", str(-I(0.35)))
    else:
        if not placeholder:
            pPr.set("marL", str(I(0.3 + 0.35 * level)))
            pPr.set("indent", str(-I(0.3)))
            bc = _sub(pPr, "a:buClr")
            kind_, val, bright = pal.spec("accent1" if level == 0 else "muted")
            if kind_ == "rgb":
                _sub(bc, "a:srgbClr", val=val)
            else:
                _sub(bc, "a:schemeClr", val=val)
            _sub(pPr, "a:buFont", typeface="Arial")
            _sub(pPr, "a:buChar", char=kind)
    _order_ppr(pPr)


PPR_ORDER = ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont", "buNone", "buAutoNum", "buChar", "buBlip", "tabLst", "defRPr", "extLst"]


def _order_ppr(pPr: Any) -> None:
    kids = list(pPr)
    kids.sort(key=lambda c: PPR_ORDER.index(local(c)) if local(c) in PPR_ORDER else 99)
    for c in kids:
        pPr.remove(c)
        pPr.append(c)


def _cell_border(cell: Any, side: str, pal: Palette, role: str, width: float) -> None:
    from lxml import etree

    tcPr = cell._tc.get_or_add_tcPr()
    tag = "{%s}ln%s" % (A_NS, side)
    for e in tcPr.findall(tag):
        tcPr.remove(e)
    ln = etree.Element(tag, w=str(int(width * EMU_PER_PT)), cap="flat", cmpd="sng", algn="ctr")
    sf = etree.SubElement(ln, "{%s}solidFill" % A_NS)
    kind, val, bright = pal.spec(role)
    if kind == "rgb":
        etree.SubElement(sf, "{%s}srgbClr" % A_NS, val=val)
    else:
        c = etree.SubElement(sf, "{%s}schemeClr" % A_NS, val=val)
        if bright:
            if bright > 0:
                etree.SubElement(c, "{%s}lumMod" % A_NS, val=str(int((1 - bright) * 100000)))
                etree.SubElement(c, "{%s}lumOff" % A_NS, val=str(int(bright * 100000)))
            else:
                etree.SubElement(c, "{%s}lumMod" % A_NS, val=str(int((1 + bright) * 100000)))
    etree.SubElement(ln, "{%s}prstDash" % A_NS, val="solid")
    # tcPr children order: lnL lnR lnT lnB lnTlToBr lnBlToTr cell3D (fill) ...
    order = ["lnL", "lnR", "lnT", "lnB"]
    pos = 0
    for i, c in enumerate(tcPr):
        if local(c) in order and order.index(local(c)) < order.index("ln" + side):
            pos = i + 1
    tcPr.insert(pos, ln)


def _image_bytes(src: Path, r: Rect) -> tuple[bytes, tuple[int, int]]:
    """Image data python-pptx can embed (PNG/JPEG/GIF/BMP/TIFF as is; SVG via Typst, WebP and others via Pillow)."""
    from PIL import Image

    suf = src.suffix.lower()
    if suf == ".svg":
        import typst

        tmp_root = src.parent
        w_in = max(1.0, r.w / IN)
        doc = f'#set page(width: auto, height: auto, margin: 0pt)\n#image("{src.name}", width: {w_in:.2f}in)'
        try:
            pngs = typst.compile(doc.encode("utf-8"), root=str(tmp_root), format="png", ppi=200)
        except Exception as e:  # noqa: BLE001
            raise SkillError(f"could not render the SVG {src.name}: {e}") from e
        data = pngs if isinstance(pngs, (bytes, bytearray)) else pngs[0]
        with Image.open(io.BytesIO(data)) as im:
            return bytes(data), im.size
    data = src.read_bytes()
    try:
        with Image.open(io.BytesIO(data)) as im:
            fmt = (im.format or "").upper()
            size = im.size
            if fmt in ("PNG", "JPEG", "GIF", "BMP", "TIFF"):
                if fmt == "JPEG" and im.mode == "CMYK":
                    buf = io.BytesIO()
                    im.convert("RGB").save(buf, "JPEG", quality=92)
                    return buf.getvalue(), size
                return data, size
            buf = io.BytesIO()
            im.convert("RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB").save(buf, "PNG")
            return buf.getvalue(), size
    except SkillError:
        raise
    except Exception as e:  # noqa: BLE001
        raise SkillError(f"{src.name} is not an image this skill can place ({e})") from e
