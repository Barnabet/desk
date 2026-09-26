#!/usr/bin/env python3
"""Edit a presentation with JSON operations; writes a new file and never touches the input."""

from __future__ import annotations

import copy
import re
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import SkillError, UsageError, add_format, emit, load_json_arg, parser, run_main  # noqa: E402

EPILOG = """examples:
  python3 scripts/pptx_edit.py deck.pptx --out deck-v2.pptx --ops edits.json
  python3 scripts/pptx_edit.py deck.pptx --out fixed.pptx --ops '[{"op": "replace", "find": "2025", "replace": "2026"}]'

ops (a JSON list, applied in order; see references/edit-ops.md for every field):
  replace          {"find", "replace", "slides"?, "regex"?, "match_case"?, "whole_word"?, "notes"?}  keeps run formatting
  set_title        {"slide", "text"}
  set_body         {"slide", "bullets" | "text", "shape"?}
  set_notes        {"slide", "text", "append"?}
  add_slide        {"after"? | "position"?, "id"?, ...a pptx_create slide spec: "type", "title", "bullets", ...}
  delete_slide     {"slides"}          duplicate_slide {"slide", "after"?, "id"?}
  move_slide       {"slide", "after"}  reorder {"order": [3, 1, 2, ...]}   hide_slide {"slides", "hidden"?}
  replace_image    {"slide", "shape", "image", "fit"?: cover|fit|stretch}
  set_table        {"slide", "shape", "rows"? | "cells"?: [{"row", "col", "text"}], "start_row"?}
  update_chart     {"slide", "shape", "categories"?, "series"?: [{"name", "values"}], "title"? (null removes it),
                    "legend"?: bottom|right|top|left|none, "colors"?: ["accent1", "#0A7F6F", …], "labels"?}
  set_shape        {"slide", "shape", "text"?, "x"?, "y"?, "w"?, "h"?, "font_size"?, "bold"?, "italic"?, "color"?,
                    "font"?, "fill"?, "line"?, "align"?, "rotation"?, "alt_text"?, "name"?}
  delete_shape     {"slide", "shape"}
  add_shape        {"slide", "kind": text|rect|card|oval|line|image|table|chart, "x", "y", "w", "h", ...}
  set_background   {"slides"?, "color" | "image"}
  set_properties   {"title"?, "author"?, "subject"?, "keywords"?}

Unknown or misspelled fields are errors (nothing is written), and each op says which fields it takes.
Slide numbers refer to the ORIGINAL deck (as pptx_read shows it), whatever other ops do first; slides added by
add_slide/duplicate_slide can be named with "id" and referenced by that id. Shapes: a name ("Title 1"), an id
(7 or "#7"), or "title"/"body". Lengths: inches (1.5) or "4cm", "72pt", "10%". Colours: #RRGGBB or theme roles
(text, muted, bg, surface, accent1-6).
"""


def main() -> int:
    ap = parser("Edit .pptx/.pptm/.potx with JSON operations.", EPILOG)
    ap.add_argument("file")
    ap.add_argument("--out", "-o", required=True, help="output file (a new path; never the input)")
    ap.add_argument("--ops", required=True, help="JSON list of operations: inline, a .json path, or - for stdin")
    ap.add_argument("--force", action="store_true", help="overwrite the output if it exists")
    add_format(ap)
    a = ap.parse_args()
    from _deck import new_output, open_deck, save_deck

    ops = load_json_arg(a.ops)
    if isinstance(ops, dict):
        ops = ops.get("ops", [ops])
    if not isinstance(ops, list) or not ops:
        raise UsageError("--ops must be a non-empty JSON list of operations")
    for i, op in enumerate(ops, 1):  # every op is checked before the deck is opened: a typo late in the list costs nothing
        if not isinstance(op, dict) or "op" not in op:
            raise UsageError(f"op {i}: each operation needs an \"op\" field")
        try:
            Editor.check(op)
        except UsageError as e:
            raise UsageError(f"op {i} ({op['op']}): {e}") from e
    out = new_output(a.out, [a.file], a.force, {".pptx", ".pptm", ".potx", ".ppsx", ".potm", ".ppsm"})
    opened = open_deck(a.file)
    base = Path(a.ops).resolve().parent if not a.ops.strip().startswith(("[", "{")) and a.ops != "-" and Path(a.ops).exists() else Path.cwd()
    ed = Editor(opened.prs, base)
    results = []
    for i, op in enumerate(ops, 1):
        if not isinstance(op, dict) or "op" not in op:
            raise UsageError(f"op {i}: each operation needs an \"op\" field")
        try:
            results.append({"op": op["op"], **ed.apply(op)})
        except (SkillError, UsageError) as e:
            raise type(e)(f"op {i} ({op['op']}): {e}") from e
    ed.finish()
    notes = save_deck(opened.prs, out, warn=False)
    summary = {"file": str(out), "ops": results, "slides": ed.order_summary(), "warnings": ed.warnings + notes}
    if a.format == "json":
        emit(summary, "json", max_chars=None)
    else:
        print(render_md(summary))
    return 0


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
RT_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"

# op → (allowed fields, required fields, one-of groups). Unknown fields are errors: a typo must never wipe content.
_SHAPE_STYLE = {"text", "x", "y", "w", "h", "font_size", "bold", "italic", "color", "font", "fill", "line", "line_width", "align", "rotation", "alt_text", "name"}
OP_FIELDS: dict[str, tuple[set[str], set[str], list[set[str]]]] = {
    "replace": ({"find", "replace", "slides", "regex", "match_case", "whole_word", "notes", "max"}, {"find", "replace"}, []),
    "set_title": ({"slide", "text"}, {"slide", "text"}, []),
    "set_body": ({"slide", "bullets", "text", "shape"}, {"slide"}, [{"bullets", "text"}]),
    "set_notes": ({"slide", "text", "append"}, {"slide", "text"}, []),
    "add_slide": (set(), set(), []),  # the slide spec is checked by the builder
    "delete_slide": ({"slides", "slide"}, set(), [{"slides", "slide"}]),
    "duplicate_slide": ({"slide", "after", "position", "id"}, {"slide"}, []),
    "move_slide": ({"slide", "after", "position"}, {"slide"}, [{"after", "position"}]),
    "reorder": ({"order", "keep_rest"}, {"order"}, []),
    "hide_slide": ({"slides", "slide", "hidden"}, set(), [{"slides", "slide"}]),
    "unhide_slide": ({"slides", "slide"}, set(), [{"slides", "slide"}]),
    "replace_image": ({"slide", "shape", "image", "path", "fit", "alt_text"}, {"slide", "shape"}, [{"image", "path"}]),
    "set_table": ({"slide", "shape", "rows", "cells", "start_row", "exact"}, {"slide"}, [{"rows", "cells"}]),
    "set_cell": ({"slide", "shape", "row", "col", "text"}, {"slide", "row", "col", "text"}, []),
    "update_chart": ({"slide", "shape", "categories", "series", "number_format", "title", "legend", "colors", "labels"}, {"slide"}, [{"series", "categories", "title", "legend", "colors", "labels", "number_format"}]),
    "set_shape": ({"slide", "shape"} | _SHAPE_STYLE, {"slide", "shape"}, [_SHAPE_STYLE]),
    "delete_shape": ({"slide", "shape"}, {"slide", "shape"}, []),
    "add_shape": (set(), {"slide"}, []),
    "add_text": (set(), {"slide"}, []),
    "add_image": (set(), {"slide"}, []),
    "set_background": ({"slides", "slide", "color", "image"}, set(), [{"color", "image"}]),
    "set_properties": ({"title", "author", "subject", "keywords", "comments", "category", "last_modified_by"}, set(), []),
}


def check_op(name: str, op: dict[str, Any]) -> None:
    import difflib

    allowed, required, groups = OP_FIELDS.get(name, (set(), set(), []))
    if name.startswith("add_") and name != "add_slide":
        from _build import ELEMENT_KEYS

        allowed = ELEMENT_KEYS | {"slide"}
    if allowed:
        for k in op:
            if k != "op" and k not in allowed:
                near = difflib.get_close_matches(str(k), sorted(allowed), n=1, cutoff=0.6)
                raise UsageError(f"unknown field \"{k}\"" + (f" (did you mean \"{near[0]}\"?)" if near else "") + f"; {name} takes: {', '.join(sorted(allowed))}")
    miss = [k for k in sorted(required) if k not in op]
    if miss:
        extra = " (the new text; \"\" deletes the matches)" if name == "replace" and miss == ["replace"] else ""
        raise UsageError(f"{name} needs \"{miss[0]}\"{extra}")
    for g in groups:
        if not any(k in op for k in g):
            raise UsageError(f"{name} needs one of: {', '.join(sorted(g))}")


class Editor:
    def __init__(self, prs: Any, base: Path):
        self.prs = prs
        self.base = base
        self.orig = list(prs.slides)  # original numbering
        self.orig_ids = {s.slide_id for s in self.orig}
        self.deleted: set[int] = set()
        self.named: dict[str, Any] = {}
        self.warnings: list[str] = []
        self._builder = None

    # addressing --------------------------------------------------------

    def slide(self, ref: Any) -> Any:
        if isinstance(ref, str) and not ref.strip().lstrip("-").isdigit():
            if ref in self.named:
                if id(self.named[ref]) in self.deleted:
                    raise UsageError(f"slide '{ref}' was deleted by an earlier op")
                return self.named[ref]
            raise UsageError(f"no slide named '{ref}' (name added slides with \"id\")")
        try:
            n = int(ref)
        except (TypeError, ValueError):
            raise UsageError(f"bad slide reference {ref!r}") from None
        if n < 1 or n > len(self.orig):
            raise UsageError(f"slide {n} is out of range (the original deck has {len(self.orig)} slides)")
        s = self.orig[n - 1]
        if id(s) in self.deleted:
            raise UsageError(f"slide {n} was deleted by an earlier op")
        return s

    def slides(self, spec: Any) -> list[Any]:
        from _common import parse_ranges

        if spec is None or spec == "all":
            return [s for s in self.prs.slides]
        if isinstance(spec, list):
            return [self.slide(x) for x in spec]
        if isinstance(spec, int):
            return [self.slide(spec)]
        if isinstance(spec, str) and spec in self.named:
            return [self.named[spec]]
        return [self.slide(n) for n in parse_ranges(str(spec), len(self.orig))]

    def shape(self, slide: Any, ref: Any) -> Any:
        from _deck import iter_shapes, is_title

        if ref is None:
            raise UsageError("this op needs \"shape\" (a name, an id, or 'title'/'body')")
        all_shapes = [s for s, _ in iter_shapes(slide.shapes)]
        sid = None
        if isinstance(ref, int) or (isinstance(ref, str) and re.fullmatch(r"#?\d+", ref.strip())):
            sid = int(str(ref).lstrip("#"))
            for s in all_shapes:
                if s.shape_id == sid:
                    return s
            raise UsageError(f"no shape with id {sid} on this slide (see pptx_read --shapes)")
        name = str(ref)
        for s in all_shapes:
            if s.name == name:
                return s
        for s in all_shapes:
            if s.name.lower() == name.lower():
                return s
        low = name.lower()
        if low == "title":
            for s in all_shapes:
                if is_title(s):
                    return s
        if low in ("body", "content", "text"):
            for s in slide.placeholders:
                t = s.placeholder_format.type.name if s.placeholder_format.type is not None else "OBJECT"
                if t in ("BODY", "OBJECT", "SUBTITLE"):
                    return s
        names = ", ".join(f'"{s.name}"' for s in all_shapes[:30])
        raise UsageError(f"no shape '{name}' on this slide; shapes: {names}")

    @property
    def builder(self) -> Any:
        if self._builder is None:
            from _build import DeckBuilder
            from _themes import THEMES

            b = DeckBuilder(None, prs=self.prs, base_dir=self.base)
            name = b.pal.cc.theme.name
            if name.startswith("Desk "):
                t = THEMES.get(name[5:])
                if t is not None:
                    b.sizes = dict(t.sizes)
                    b.theme = t
            self._builder = b
        return self._builder

    # dispatch ----------------------------------------------------------

    @classmethod
    def check(cls, op: dict[str, Any]) -> str:
        """The op's method name, after checking the op exists and its fields are known and complete."""
        name = str(op["op"]).replace("-", "_")
        if not hasattr(cls, "op_" + name):
            import difflib

            near = difflib.get_close_matches(name, sorted(OP_FIELDS), n=1, cutoff=0.6)
            raise UsageError(f"unknown op '{op['op']}'" + (f" (did you mean '{near[0]}'?)" if near else "") + f"; ops: {', '.join(sorted(OP_FIELDS))}")
        check_op(name, op)
        return name

    def apply(self, op: dict[str, Any]) -> dict[str, Any]:
        name = self.check(op)
        return getattr(self, "op_" + name)(op) or {}

    def finish(self) -> None:
        if len(self.prs.slides) == 0:
            raise UsageError("these operations delete every slide; a deck needs at least one (hide slides with hide_slide instead)")
        if self._builder is not None:
            self._builder.run_fits()
            self._builder.layout_warnings()
            self._builder._order()
            self.warnings.extend(self._builder.warnings)
            self._footers()

    def _footers(self) -> None:
        """Slides added by add_slide get the footer and slide number the neighbouring slides carry."""
        b = self._builder
        orig_ids = self.orig_ids
        slides = list(self.prs.slides)
        for i, s in enumerate(slides):
            if s.slide_id in orig_ids or b.kinds.get(s.slide_id) in ("title", "section", "closing", None):
                continue
            if any(sh.name in ("Footer", "Slide number") for sh in s.shapes):
                continue
            donor = next((d for d in (slides[i - 1::-1] if i else []) + slides[i + 1:] if d.slide_id in orig_ids and any(sh.name in ("Footer", "Slide number") for sh in d.shapes)), None)
            if donor is None:
                continue
            for sh in donor.shapes:
                if sh.name in ("Footer", "Slide number"):
                    el = copy.deepcopy(sh._element)
                    ids = [int(e.get("id", "0")) for e in s._element.iter("{%s}cNvPr" % P_NS)]
                    el.find(".//{%s}cNvPr" % P_NS).set("id", str(max(ids + [1]) + 1))
                    s.shapes._spTree.append(el)

    def order_summary(self) -> list[str]:
        from _deck import slide_title

        out = []
        index = {id(s): i + 1 for i, s in enumerate(self.orig)}
        for n, s in enumerate(self.prs.slides, 1):
            o = index.get(id(s))
            tag = f"was {o}" if o is not None else "new"
            out.append(f"{n}. {slide_title(s) or '(no title)'} ({tag})")
        return out

    # text --------------------------------------------------------------

    def op_replace(self, op: dict[str, Any]) -> dict[str, Any]:
        find = op.get("find")
        if not find:
            raise UsageError("replace needs \"find\"")
        repl = str(op.get("replace", ""))
        flags = 0 if op.get("match_case", True) else re.I
        pat = find if op.get("regex") else re.escape(str(find))
        if op.get("whole_word"):
            pat = rf"\b{pat}\b"
        try:
            rx = re.compile(pat, flags)
        except re.error as e:
            raise UsageError(f"bad regex: {e}") from e
        limit = int(op.get("max", 0) or 0)
        total = 0
        where: list[int] = []
        for slide in self.slides(op.get("slides")):
            n = 0
            for txBody in _text_bodies(slide, bool(op.get("notes", False))):
                for p in txBody.iter("{%s}p" % A_NS):
                    k = _replace_in_paragraph(p, rx, repl, (limit - total - n) if limit else 0, bool(op.get("regex")))
                    n += k
                    if limit and total + n >= limit:
                        break
            if n:
                where.append(self._number(slide))
            total += n
            if limit and total >= limit:
                break
        if total == 0:
            self.warnings.append(f"replace: '{find}' was not found")
        return {"replaced": total, "slides": where}

    def op_set_title(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        from _deck import is_title

        t = next((s for s in slide.shapes if is_title(s)), None)
        text = str(op.get("text", ""))
        if t is None:
            return self._new_title(slide, text)
        _set_text_keep_format(t.text_frame, text)
        return {}

    def _new_title(self, slide: Any, text: str) -> dict[str, Any]:
        """Adds a real title placeholder (so outlines, lint and screen readers see a title) in the free band above
        the slide's content, sized to fit there."""
        from pptx.oxml import parse_xml

        from _build import FitJob, Rect
        from _deck import iter_shapes

        W, H = self.prs.slide_width, self.prs.slide_height
        geo = None
        for owner in (slide.slide_layout, slide.slide_layout.slide_master):
            for ph in owner.placeholders:
                if ph.placeholder_format.type is not None and ph.placeholder_format.type.name in ("TITLE", "CENTER_TITLE") and ph.width:
                    geo = Rect(ph.left or 0, ph.top or 0, ph.width, ph.height)
                    break
            if geo is not None:
                break
        if geo is None:
            geo = Rect(int(W * 0.05), int(H * 0.04), int(W * 0.9), int(H * 0.15))
        others = [(s, s.left, s.top, s.width, s.height) for s, depth in iter_shapes(slide.shapes) if depth == 0 and s.width and s.height and s.left is not None]
        blocking = [o for o in others if o[1] < geo.x + geo.w and o[1] + o[3] > geo.x and o[2] < geo.y + geo.h and o[2] + o[4] > geo.y]
        note = None
        if blocking:
            top_content = min(o[2] for o in blocking)
            y0 = min(geo.y, int(H * 0.035))
            y1 = top_content - int(H * 0.012)
            if y1 - y0 >= int(H * 0.06):
                geo = Rect(geo.x, y0, geo.w, y1 - y0)
            else:
                geo = Rect(geo.x, y0, geo.w, max(int(H * 0.08), y1 - y0))
                b = blocking[0][0]
                label = f"\"{b.name}\"" if b.name else f"the shape with id {b.shape_id}"
                note = f"there is no free room above {label}, so the new title overlaps it; move the content down (set_shape y) or shrink it"
        tree = slide.shapes._spTree
        ids = [int(e.get("id", "0")) for e in tree.iter("{%s}cNvPr" % P_NS)]
        nid = max(ids + [1]) + 1
        xml = (
            f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="{nid}" name="Title {nid}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="{geo.x}" y="{geo.y}"/><a:ext cx="{geo.w}" cy="{geo.h}"/></a:xfrm></p:spPr>'
            f'<p:txBody><a:bodyPr anchor="b"><a:normAutofit/></a:bodyPr><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" dirty="0"/><a:t></a:t></a:r></a:p></p:txBody></p:sp>'
        )
        el = parse_xml(xml)
        el.find(".//{%s}t" % A_NS).text = text
        first = next((c for c in tree if not c.tag.endswith(("}nvGrpSpPr", "}grpSpPr"))), None)
        if first is not None:
            first.addprevious(el)
        else:
            tree.append(el)
        shp = next(s for s in slide.shapes if s._element is el)
        self.builder.jobs.append(FitJob(slide, shp, 0.45))
        if note:
            self.warnings.append(f"set_title on slide {self._number(slide)}: {note}")
        return {"created": "title placeholder", "at_in": [round(geo.x / 914400, 2), round(geo.y / 914400, 2), round(geo.w / 914400, 2), round(geo.h / 914400, 2)]}

    def op_set_body(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        shp = self.shape(slide, op.get("shape", "body"))
        b = self.builder
        items = b.items(op.get("bullets") if op.get("bullets") is not None else op.get("text", ""))
        is_ph = getattr(shp, "is_placeholder", False)
        b.write(shp.text_frame, items, placeholder=is_ph, bullets=not is_ph and bool(op.get("bullets")))
        from _build import FitJob

        b.jobs.append(FitJob(slide, shp, 0.6))
        return {"paragraphs": len(items)}

    def op_set_notes(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        tf = slide.notes_slide.notes_text_frame
        text = str(op.get("text", ""))
        if op.get("append") and tf.text.strip():
            tf.text = tf.text.rstrip() + "\n" + text
        else:
            tf.text = text
        return {}

    # slides ------------------------------------------------------------

    def _place(self, new: Any, op: dict[str, Any], default_after: Any = None) -> None:
        lst = self.prs.slides._sldIdLst
        el = _sld_el(self.prs, new)
        lst.remove(el)
        pos = op.get("position")
        after = op.get("after", default_after)
        if pos is not None and str(pos) not in ("end", "last"):
            idx = max(0, int(pos) - 1) if str(pos) not in ("start", "first") else 0
            lst.insert(min(idx, len(lst)), el)
        elif after is not None and str(after) not in ("end", "last"):
            if str(after) in ("0", "start"):
                lst.insert(0, el)
            else:
                ref = self.slide(after)
                _sld_el(self.prs, ref).addnext(el)
        else:
            lst.append(el)
        if op.get("id"):
            self.named[str(op["id"])] = new

    def op_add_slide(self, op: dict[str, Any]) -> dict[str, Any]:
        spec = dict(op.get("spec") or {k: v for k, v in op.items() if k not in ("op", "after", "position", "id")})
        if not spec:
            raise UsageError("add_slide needs a slide spec (\"type\", \"title\", \"bullets\", …)")
        b = self.builder
        before = {s.slide_id for s in self.prs.slides}
        b.add(spec)
        new_slides = [s for s in self.prs.slides if s.slide_id not in before]
        main = new_slides[0]
        self._place(main, op)
        anchor = main
        for extra in new_slides[1:]:
            lst = self.prs.slides._sldIdLst
            e = _sld_el(self.prs, extra)
            lst.remove(e)
            _sld_el(self.prs, anchor).addnext(e)
            anchor = extra
        return {"layout": main.slide_layout.name, "type": b.kinds.get(main.slide_id)}

    def op_delete_slide(self, op: dict[str, Any]) -> dict[str, Any]:
        targets = self.slides(op.get("slides", op.get("slide")))
        lst = self.prs.slides._sldIdLst
        for s in targets:
            el = _sld_el(self.prs, s)
            rid = el.get("{%s}id" % R_NS)
            lst.remove(el)
            self.prs.part.drop_rel(rid)
            self.deleted.add(id(s))
        return {"deleted": len(targets)}

    def op_duplicate_slide(self, op: dict[str, Any]) -> dict[str, Any]:
        src = self.slide(op.get("slide"))
        new = duplicate_slide(self.prs, src)
        self._place(new, op, default_after=None)
        if op.get("after") is None and op.get("position") is None:
            lst = self.prs.slides._sldIdLst
            e = _sld_el(self.prs, new)
            lst.remove(e)
            _sld_el(self.prs, src).addnext(e)
        return {}

    def op_move_slide(self, op: dict[str, Any]) -> dict[str, Any]:
        s = self.slide(op.get("slide"))
        if op.get("after") is None and op.get("position") is None:
            raise UsageError("move_slide needs \"after\" (a slide number, 0 for the start) or \"position\"")
        self._place(s, op)
        return {}

    def op_reorder(self, op: dict[str, Any]) -> dict[str, Any]:
        order = op.get("order")
        if not isinstance(order, list) or not order:
            raise UsageError("reorder needs \"order\": [slide numbers in their new order]")
        slides = [self.slide(x) for x in order]
        if len({id(s) for s in slides}) != len(slides):
            raise UsageError("reorder: a slide is listed twice")
        lst = self.prs.slides._sldIdLst
        els = [_sld_el(self.prs, s) for s in slides]
        rest = [e for e in lst if e not in els]
        if rest and not op.get("keep_rest", True):
            raise UsageError("reorder: list every slide")
        for e in list(lst):
            lst.remove(e)
        for e in els + rest:
            lst.append(e)
        return {"unlisted_kept_at_end": len(rest)}

    def op_hide_slide(self, op: dict[str, Any]) -> dict[str, Any]:
        hidden = op.get("hidden", True)
        targets = self.slides(op.get("slides", op.get("slide")))
        for s in targets:
            if hidden:
                s._element.set("show", "0")
            elif "show" in s._element.attrib:
                del s._element.attrib["show"]
        return {"slides": len(targets), "hidden": bool(hidden)}

    op_unhide_slide = lambda self, op: self.op_hide_slide({**op, "hidden": False})  # noqa: E731

    # shapes ------------------------------------------------------------

    def op_replace_image(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        pic = self.shape(slide, op.get("shape"))
        if not pic._element.tag.endswith("}pic"):
            raise UsageError(f"'{pic.name}' is not a picture")
        path = self.builder.resolve_image(str(op.get("image") or op.get("path") or ""))
        from _build import Rect, _image_bytes

        import io

        data, (pw, ph) = _image_bytes(path, Rect(pic.left, pic.top, pic.width, pic.height))
        _img_part, rid = slide.part.get_or_add_image_part(io.BytesIO(data))
        blip = pic._element.find(".//{%s}blip" % A_NS)
        old = blip.get("{%s}embed" % R_NS)
        blip.set("{%s}embed" % R_NS, rid)
        for ext in list(blip.iter("{%s}ext" % A_NS)):
            if ext.find("{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip") is not None:
                ext.getparent().remove(ext)
        fit = op.get("fit", "cover")
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        if fit == "cover":
            box_r, img_r = pic.width / pic.height, pw / ph
            if img_r > box_r:
                c = (1 - box_r / img_r) / 2
                pic.crop_left = pic.crop_right = c
            else:
                c = (1 - img_r / box_r) / 2
                pic.crop_top = pic.crop_bottom = c
        elif fit == "fit":
            s = min(pic.width / pw, pic.height / ph)
            w, h = int(pw * s), int(ph * s)
            pic.left, pic.top = pic.left + (pic.width - w) // 2, pic.top + (pic.height - h) // 2
            pic.width, pic.height = w, h
        if old and old != rid:
            slide.part.drop_rel(old)
        if op.get("alt_text"):
            pic._element.find(".//{%s}cNvPr" % P_NS).set("descr", str(op["alt_text"]))
        return {"pixels": [pw, ph]}

    def op_set_table(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        gf = self.shape(slide, op.get("shape", "Table"))
        if not getattr(gf, "has_table", False):
            raise UsageError(f"'{gf.name}' is not a table")
        tbl = gf.table
        changed = 0
        if op.get("rows") is not None:
            rows = op["rows"]
            start = int(op.get("start_row", 0))
            need = start + len(rows)
            _resize_rows(gf, need if op.get("exact", True) else max(need, len(tbl.rows)))
            tbl = gf.table
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    if j < len(tbl.columns):
                        _set_text_keep_format(tbl.cell(start + i, j).text_frame, "" if v is None else str(v))
                        changed += 1
        for c in op.get("cells") or []:
            i, j = int(c["row"]), int(c["col"])
            if i >= len(tbl.rows) or j >= len(tbl.columns):
                raise UsageError(f"cell ({i}, {j}) is outside the {len(tbl.rows)}×{len(tbl.columns)} table (rows and columns count from 0)")
            _set_text_keep_format(tbl.cell(i, j).text_frame, str(c.get("text", "")))
            changed += 1
        return {"cells": changed, "rows": len(tbl.rows), "columns": len(tbl.columns)}

    def op_set_cell(self, op: dict[str, Any]) -> dict[str, Any]:
        return self.op_set_table({**op, "cells": [{"row": op.get("row"), "col": op.get("col"), "text": op.get("text", "")}]})

    def op_update_chart(self, op: dict[str, Any]) -> dict[str, Any]:
        from pptx.chart.data import CategoryChartData, XyChartData
        from pptx.enum.chart import XL_LEGEND_POSITION

        slide = self.slide(op.get("slide"))
        gf = self.shape(slide, op.get("shape", "Chart"))
        if not getattr(gf, "has_chart", False):
            raise UsageError(f"'{gf.name}' is not a chart")
        chart = gf.chart
        cspace = chart._chartSpace
        C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
        old_n = len(cspace.findall(".//" + C + "ser"))
        out: dict[str, Any] = {}
        series = op.get("series")
        ctype = chart.chart_type.name if chart.chart_type is not None else ""
        radial = "PIE" in ctype or "DOUGHNUT" in ctype
        if series is not None or op.get("categories") is not None:
            if series is None:
                series = [{"name": s.name, "values": list(s.values)} for s in chart.series]
            if not series:
                raise UsageError("update_chart needs \"series\": [{\"name\", \"values\"}]")
            for s in series:
                if isinstance(s, dict):
                    from _build import SERIES_KEYS, _check_keys

                    _check_keys(s, SERIES_KEYS, "a chart series")
            xy = "XY" in ctype or "BUBBLE" in ctype
            if xy:
                cd = XyChartData()
                for s in series:
                    ser = cd.add_series(str(s.get("name", "")))
                    for x, y in s.get("points") or []:
                        ser.add_data_point(float(x), float(y))
            else:
                cats = op.get("categories")
                if cats is None:
                    cats = list(chart.plots[0].categories)
                cd = CategoryChartData(number_format=op["number_format"]) if op.get("number_format") else CategoryChartData()
                cd.categories = [str(c) for c in cats]
                for s in series:
                    vals = s.get("values") or []
                    if len(vals) != len(cats):
                        raise UsageError(f"series '{s.get('name', '')}' has {len(vals)} values for {len(cats)} categories")
                    try:
                        cd.add_series(str(s.get("name", "")), [None if v is None else float(v) for v in vals])
                    except (TypeError, ValueError) as e:
                        raise UsageError(f"series '{s.get('name', '')}': values must be numbers or null ({e})") from None
            chart.replace_data(cd)
            out["series"] = len(series)
        new_n = len(cspace.findall(".//" + C + "ser"))
        # colours: the ones asked for, and a distinct accent for every series added (python-pptx copies the last one's)
        colors = op.get("colors")
        if colors is not None and not isinstance(colors, list):
            raise UsageError("\"colors\" must be a list like [\"accent1\", \"#0A7F6F\"]")
        recolored = self._chart_colors(slide, chart, old_n, colors, radial)
        if recolored:
            out["recolored"] = recolored
        title_el = cspace.find(C + "chart/" + C + "title")
        if "title" in op:
            t = op["title"]
            if t in (None, "", False):
                chart.has_title = False
                out["title"] = None
            else:
                chart.has_title = True
                chart.chart_title.text_frame.text = str(t)
                out["title"] = str(t)
        elif old_n == 1 and new_n > 1 and title_el is not None:
            has_text = any((x.text or "").strip() for x in title_el.iter("{*}t")) or title_el.find(".//" + C + "strRef") is not None
            if not has_text:
                chart.has_title = False
                self.warnings.append(f"update_chart on slide {self._number(slide)}: the chart's automatic title showed its one series' name; with {new_n} series PowerPoint would show 'Chart Title', so it was removed (set one with \"title\")")
        if "legend" in op:
            lg = op["legend"]
            if lg in (None, "none", False):
                chart.has_legend = False
            else:
                pos = {"bottom": XL_LEGEND_POSITION.BOTTOM, "right": XL_LEGEND_POSITION.RIGHT, "top": XL_LEGEND_POSITION.TOP, "left": XL_LEGEND_POSITION.LEFT}.get(str(lg))
                if pos is None:
                    raise UsageError("\"legend\" is bottom, right, top, left or none")
                chart.has_legend = True
                chart.legend.position = pos
                chart.legend.include_in_layout = False
        elif old_n == 1 and new_n > 1 and not chart.has_legend and not radial:
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            out["legend"] = "added (several series)"
        if "labels" in op:
            plot = chart.plots[0]
            plot.has_data_labels = bool(op["labels"])
            if op["labels"]:
                dl = plot.data_labels
                if radial and op["labels"] != "value":
                    dl.show_percentage, dl.show_value = True, False
                else:
                    dl.show_value = True
        return out

    def _chart_colors(self, slide: Any, chart: Any, old_n: int, colors: list[Any] | None, radial: bool) -> int:
        """Applies "colors" (theme roles or #hex) and gives series added by update_chart an unused accent."""
        from _chartdraw import _series_color

        from _ooxml import SlideCtx, pres_defaults

        pal = self.builder.pal
        cc = SlideCtx(self.prs, slide, pres_defaults(self.prs), 0).cc
        sers = list(chart.series)
        done = 0
        if radial:
            if colors:
                pts = sers[0].points if sers else []
                for j, role in enumerate(colors):
                    if j < len(pts):
                        pts[j].format.fill.solid()
                        pal.apply(pts[j].format.fill.fore_color, str(role), direct=True)
                        done += 1
            return done
        used = {c.hex for c in (_series_color(s._element, cc) for s in sers[:old_n]) if c is not None}
        accents = [f"accent{i}" for i in range(1, 7)]
        for i, s in enumerate(sers):
            role = str(colors[i]) if colors and i < len(colors) and colors[i] is not None else None
            if role is None and i >= old_n and old_n > 0:
                role = next((a for a in accents if (pal.rgb(a).hex not in used)), accents[i % 6])
            if role is None:
                continue
            used.add(pal.rgb(role).hex)
            line_like = s._element.getparent().tag.endswith(("lineChart", "scatterChart", "radarChart", "line3DChart"))
            if line_like:
                pal.apply(s.format.line.color, role, direct=True)
                try:
                    s.marker.format.fill.solid()
                    pal.apply(s.marker.format.fill.fore_color, role, direct=True)
                    pal.apply(s.marker.format.line.color, role, direct=True)
                except Exception:  # noqa: BLE001 — series without markers
                    pass
            else:
                s.format.fill.solid()
                pal.apply(s.format.fill.fore_color, role, direct=True)
            done += 1
        return done

    def op_set_shape(self, op: dict[str, Any]) -> dict[str, Any]:
        from pptx.util import Pt

        from _ooxml import parse_length

        slide = self.slide(op.get("slide"))
        shp = self.shape(slide, op.get("shape"))
        W, H = self.prs.slide_width, self.prs.slide_height
        b = self.builder
        done = []
        for k, total, attr in (("x", W, "left"), ("y", H, "top"), ("w", W, "width"), ("h", H, "height")):
            if k in op:
                setattr(shp, attr, parse_length(op[k], total))
                done.append(k)
        if "rotation" in op:
            shp.rotation = float(op["rotation"])
            done.append("rotation")
        if "name" in op:
            shp.name = str(op["name"])
        if "alt_text" in op:
            nv = shp._element.find(".//{%s}cNvPr" % P_NS)
            if nv is not None:
                nv.set("descr", str(op["alt_text"]))
                done.append("alt_text")
        if "text" in op:
            if not shp.has_text_frame:
                raise UsageError(f"'{shp.name}' has no text")
            _set_text_keep_format(shp.text_frame, str(op["text"]))
            done.append("text")
        if any(k in op for k in ("font_size", "bold", "italic", "color", "font", "align")):
            if not shp.has_text_frame:
                raise UsageError(f"'{shp.name}' has no text")
            from pptx.enum.text import PP_ALIGN

            for p in shp.text_frame.paragraphs:
                if "align" in op:
                    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[str(op["align"])]
                for r in p.runs:
                    if "font_size" in op:
                        r.font.size = Pt(float(op["font_size"]))
                    if "bold" in op:
                        r.font.bold = bool(op["bold"])
                    if "italic" in op:
                        r.font.italic = bool(op["italic"])
                    if "font" in op:
                        r.font.name = str(op["font"])
                    if "color" in op:
                        b.pal.apply(r.font.color, str(op["color"]))
                end = p._p.find("{%s}endParaRPr" % A_NS)
                if end is not None and "font_size" in op:
                    end.set("sz", str(int(float(op["font_size"]) * 100)))
            done.append("text style")
        if "fill" in op:
            if op["fill"] in (None, "none"):
                shp.fill.background()
            else:
                shp.fill.solid()
                b.pal.apply(shp.fill.fore_color, str(op["fill"]))
            done.append("fill")
        if "line" in op:
            if op["line"] in (None, "none"):
                shp.line.fill.background()
            else:
                b.pal.apply(shp.line.color, str(op["line"]))
                if op.get("line_width"):
                    shp.line.width = Pt(float(op["line_width"]))
            done.append("line")
        return {"shape": shp.name, "changed": done}

    def op_delete_shape(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        shp = self.shape(slide, op.get("shape"))
        el = shp._element
        rids = [v for e in el.iter() for k, v in e.attrib.items() if k.startswith("{%s}" % R_NS)]
        el.getparent().remove(el)
        for rid in rids:
            try:
                slide.part.drop_rel(rid)
            except KeyError:
                pass
        return {"deleted": shp.name}

    def op_add_shape(self, op: dict[str, Any]) -> dict[str, Any]:
        slide = self.slide(op.get("slide"))
        el = {k: v for k, v in op.items() if k not in ("op", "slide")}
        n_before = len(slide.shapes)
        self.builder.add_elements(slide, [el])
        return {"added": len(slide.shapes) - n_before}

    op_add_text = lambda self, op: self.op_add_shape({**op, "kind": "text"})  # noqa: E731
    op_add_image = lambda self, op: self.op_add_shape({**op, "kind": "image"})  # noqa: E731

    def op_set_background(self, op: dict[str, Any]) -> dict[str, Any]:
        targets = self.slides(op.get("slides", op.get("slide", "all")))
        for s in targets:
            if op.get("image"):
                self.builder._background(s, {"image": op["image"]})
            else:
                self.builder._background(s, str(op.get("color", "bg")))
        return {"slides": len(targets)}

    def op_set_properties(self, op: dict[str, Any]) -> dict[str, Any]:
        cp = self.prs.core_properties
        done = []
        for k in ("title", "author", "subject", "keywords", "comments", "category", "last_modified_by"):
            if k in op:
                setattr(cp, k, str(op[k]))
                done.append(k)
        return {"set": done}

    def _number(self, slide: Any) -> int:
        for i, s in enumerate(self.prs.slides, 1):
            if s.slide_id == slide.slide_id:
                return i
        return 0


# ── helpers ─────────────────────────────────────────────────────────────


def _sld_el(prs: Any, slide: Any) -> Any:
    for e in prs.slides._sldIdLst:
        if int(e.get("id")) == slide.slide_id:
            return e
    raise SkillError("slide is not in the deck")


def _text_bodies(slide: Any, notes: bool) -> list[Any]:
    out = list(slide._element.iter("{%s}txBody" % P_NS)) + list(slide._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}txBody"))
    if notes and slide.has_notes_slide:
        out += list(slide.notes_slide._element.iter("{%s}txBody" % P_NS))
    # chart titles and SmartArt are separate parts; they are not edited here
    return out


def _replace_in_paragraph(p: Any, rx: Any, repl: str, limit: int, is_regex: bool) -> int:
    """Cross-run replace: the replacement takes the formatting of the run where each match starts."""
    runs = [r for r in p if r.tag in ("{%s}r" % A_NS, "{%s}fld" % A_NS)]
    texts = []
    for r in runs:
        t = r.find("{%s}t" % A_NS)
        texts.append(t.text or "" if t is not None else "")
    full = "".join(texts)
    matches = list(rx.finditer(full))
    if limit:
        matches = matches[:limit]
    if not matches:
        return 0
    starts = []
    pos = 0
    for tx in texts:
        starts.append(pos)
        pos += len(tx)
    count = 0
    for m in reversed(matches):
        a, b = m.start(), m.end()
        if a == b:
            continue
        new = m.expand(repl) if is_regex else repl
        first = True
        for i, (st, tx) in enumerate(zip(starts, texts)):
            en = st + len(tx)
            if en <= a or st >= b:
                continue
            lo, hi = max(a, st) - st, min(b, en) - st
            if first:
                texts[i] = tx[:lo] + new + tx[hi:]
                first = False
            else:
                texts[i] = tx[:lo] + tx[hi:]
        # recompute starts after the edit
        starts = []
        pos = 0
        for tx in texts:
            starts.append(pos)
            pos += len(tx)
        count += 1
    for r, tx in zip(runs, texts):
        t = r.find("{%s}t" % A_NS)
        if t is not None:
            t.text = tx
    # drop runs emptied by the replacement (keep at least one)
    for r, tx in zip(runs, texts):
        if not tx and r.tag == "{%s}r" % A_NS and len([x for x in p if x.tag == "{%s}r" % A_NS]) > 1:
            p.remove(r)
    return count


def _set_text_keep_format(tf: Any, text: str) -> None:
    """Replaces a text frame's text, keeping the first run's and paragraph's formatting; \\n makes new paragraphs."""
    paras = tf.paragraphs
    first_p = paras[0]._p
    pPr = first_p.find("{%s}pPr" % A_NS)
    r0 = first_p.find("{%s}r" % A_NS)
    rPr = r0.find("{%s}rPr" % A_NS) if r0 is not None else None
    end = first_p.find("{%s}endParaRPr" % A_NS)
    if rPr is None and end is not None:
        rPr = copy.deepcopy(end)
        rPr.tag = "{%s}rPr" % A_NS
    body = first_p.getparent()
    for p in list(body.findall("{%s}p" % A_NS)):
        body.remove(p)
    from lxml import etree

    for line in text.split("\n"):
        p = etree.SubElement(body, "{%s}p" % A_NS)
        if pPr is not None:
            p.append(copy.deepcopy(pPr))
        if line:
            r = etree.SubElement(p, "{%s}r" % A_NS)
            if rPr is not None:
                r.append(copy.deepcopy(rPr))
            t = etree.SubElement(r, "{%s}t" % A_NS)
            t.text = line
        if end is not None:
            p.append(copy.deepcopy(end))


def _resize_rows(gf: Any, n: int) -> None:
    tbl = gf.table._tbl
    trs = tbl.findall("{%s}tr" % A_NS)
    if n < 1:
        raise UsageError("a table needs at least one row")
    while len(trs) < n:
        new = copy.deepcopy(trs[-1])
        for t in new.iter("{%s}t" % A_NS):
            t.text = ""
        trs[-1].addnext(new)
        trs = tbl.findall("{%s}tr" % A_NS)
    while len(trs) > n:
        tbl.remove(trs[-1])
        trs = tbl.findall("{%s}tr" % A_NS)
    gf.height = sum(int(tr.get("h", "0")) for tr in trs)


def duplicate_slide(prs: Any, src: Any) -> Any:
    """A full copy of a slide after the deck's last slide: shapes, pictures, charts (with their data), SmartArt,
    media links, hyperlinks and notes."""
    from pptx.opc.package import Part, XmlPart

    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag.rsplit("}", 1)[-1] not in ("nvGrpSpPr", "grpSpPr"):
            tree.remove(el)
    src_el = copy.deepcopy(src._element)
    # background and colour-map overrides come along with cSld
    new_csld = new._element.find("{%s}cSld" % P_NS)
    src_csld = src_el.find("{%s}cSld" % P_NS)
    bg = src_csld.find("{%s}bg" % P_NS)
    if bg is not None:
        new_csld.insert(0, bg)
    for el in list(src_csld.find("{%s}spTree" % P_NS)):
        if el.tag.rsplit("}", 1)[-1] in ("nvGrpSpPr", "grpSpPr"):
            continue
        tree.append(el)
    for k, v in src._element.attrib.items():
        new._element.set(k, v)
    for tag in ("clrMapOvr", "transition", "timing"):
        e = src_el.find("{%s}%s" % (P_NS, tag))
        if e is not None:
            old = new._element.find("{%s}%s" % (P_NS, tag))
            if old is not None:
                new._element.remove(old)
            new._element.append(e)
    mapping: dict[str, str] = {}
    cloned: dict[int, Any] = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT_LAYOUT, RT_NOTES):
            continue
        if rel.is_external:
            mapping[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            continue
        target = rel.target_part
        if _shareable(target):
            mapping[rid] = new.part.relate_to(target, rel.reltype)
        else:
            mapping[rid] = new.part.relate_to(_clone_part(prs, target, cloned), rel.reltype)
    _remap_rids(new._element, mapping)
    # SmartArt data parts name the slide's drawing relationship (dsp:dataModelExt relId)
    for part in cloned.values():
        el = getattr(part, "_element", None)
        if el is not None:
            for e in el.iter("{http://schemas.microsoft.com/office/drawing/2008/diagram}dataModelExt"):
                if e.get("relId") in mapping:
                    e.set("relId", mapping[e.get("relId")])
    if src.has_notes_slide:
        txt = src.notes_slide.notes_text_frame.text if src.notes_slide.notes_text_frame is not None else ""
        if txt:
            new.notes_slide.notes_text_frame.text = txt
    del Part, XmlPart
    return new


def _shareable(part: Any) -> bool:
    ct = getattr(part, "content_type", "") or ""
    return ct.startswith(("image/", "video/", "audio/")) or "media" in ct or ct.endswith("slideLayout+xml")


def _clone_part(prs: Any, part: Any, cloned: dict[int, Any]) -> Any:
    """Copies a part (chart, SmartArt data/drawing, embedded workbook, OLE) and, recursively, what it relates to."""
    from pptx.opc.package import PartFactory
    from pptx.opc.packuri import PackURI

    if id(part) in cloned:
        return cloned[id(part)]
    pkg = prs.part.package
    name = str(part.partname)
    m = re.match(r"^(.*?)(\d*)(\.[^./]+)$", name)
    tmpl = f"{m.group(1)}%d{m.group(3)}" if m else name + "%d"
    new_name = pkg.next_partname(tmpl)
    blob = part.blob
    new_part = PartFactory(PackURI(str(new_name)), part.content_type, pkg, blob)
    cloned[id(part)] = new_part
    mapping: dict[str, str] = {}
    for rid, rel in part.rels.items():
        if rel.is_external:
            mapping[rid] = new_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif _shareable(rel.target_part):
            mapping[rid] = new_part.relate_to(rel.target_part, rel.reltype)
        else:
            mapping[rid] = new_part.relate_to(_clone_part(prs, rel.target_part, cloned), rel.reltype)
    el = getattr(new_part, "_element", None)
    if el is not None:
        _remap_rids(el, mapping)
    return new_part


def _remap_rids(root: Any, mapping: dict[str, str]) -> None:
    if not mapping:
        return
    for e in root.iter():
        for k, v in list(e.attrib.items()):
            if k.startswith("{%s}" % R_NS) and v in mapping:
                e.set(k, mapping[v])


def render_md(s: dict[str, Any]) -> str:
    L = [f"Wrote {s['file']}."]
    for r in s["ops"]:
        details = ", ".join(f"{k} {v}" for k, v in r.items() if k != "op" and v not in (None, [], {}))
        L.append(f"- {r['op']}" + (f": {details}" if details else ""))
    L.append("Slides now: " + "; ".join(s["slides"][:40]) + (" …" if len(s["slides"]) > 40 else ""))
    for w in s["warnings"]:
        L.append(f"warning: {w}")
    L.append(f"Check it: python3 scripts/pptx_render.py {s['file']} --sheet (then view_image) and python3 scripts/pptx_lint.py {s['file']}")
    return "\n".join(L)


if __name__ == "__main__":
    run_main(main)
