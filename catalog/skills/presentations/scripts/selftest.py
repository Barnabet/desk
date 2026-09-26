#!/usr/bin/env python3
"""End-to-end self test for the presentations skill: builds fixture decks in a temp folder, runs every script the
way an agent does (python3 scripts/<name>.py …), and checks the real outputs. No network.

The built-in paths run with DESK_SOFFICE=none; the LibreOffice paths run too when LibreOffice is installed.
Prints "ok: N checks in X s" and exits non-zero on any failure.

  python3 scripts/selftest.py              # run everything
  python3 scripts/selftest.py --keep       # keep the temp folder and print its path
  python3 scripts/selftest.py --only edit  # run only the tests whose name contains "edit"
  python3 scripts/selftest.py --fixtures DIR   # only write the fixture decks into DIR
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
PY = sys.executable
NO_LO = {"DESK_SOFFICE": "none"}
SCRIPTS = ["pptx_info.py", "pptx_read.py", "pptx_create.py", "pptx_edit.py", "pptx_render.py", "pptx_lint.py", "pptx_convert.py"]
HEAVY = ("pptx", "PIL", "typst", "lxml", "pypdfium2", "xlsxwriter")


# ── fixtures ────────────────────────────────────────────────────────────


def _images(d: Path) -> None:
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (1600, 1000), (30, 64, 120))
    dr = ImageDraw.Draw(im)
    for i in range(0, 1600, 64):
        dr.rectangle([i, 1000 - (i * 5 % 700) - 100, i + 48, 980], fill=(240, 180, 60))
    im.save(d / "photo.jpg", quality=85)
    im = Image.new("RGB", (1000, 1000), (228, 232, 238))
    ImageDraw.Draw(im).rectangle([100, 100, 900, 900], outline=(20, 20, 20), width=24)
    im.save(d / "diagram.png")
    im = Image.new("RGB", (200, 100), (0, 0, 255))
    ImageDraw.Draw(im).rectangle([100, 0, 199, 99], fill=(255, 255, 0))
    im.save(d / "small.png")
    Image.new("RGB", (400, 300), (0, 160, 0)).save(d / "green.png")


def make_foreign(d: Path) -> Path:
    """A 4:3 deck made the way other tools make them: default template, placeholders, a table, a chart, a picture,
    a group, notes, a transition, a hidden slide, and deliberate lint problems."""
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    prs = Presentation()
    prs.core_properties.title = "Acme Quarterly deck"
    prs.core_properties.author = "Finance Team"
    L = prs.slide_layouts
    # 1: title slide with notes and a fade transition
    s = prs.slides.add_slide(L[0])
    s.shapes.title.text = "Acme Quarterly"
    s.placeholders[1].text = "Prepared by Finance"
    s.notes_slide.notes_text_frame.text = "Welcome everyone."
    tr = etree.SubElement(s._element, P + "transition")
    etree.SubElement(tr, P + "fade")
    # 2: bullets with levels; "Acme Corp" split across two runs
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = "Plan"
    tf = s.placeholders[1].text_frame
    tf.text = "Grow revenue"
    p = tf.add_paragraph()
    p.text = "Enter Spain"
    p.level = 1
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Partner with Ac"
    r = p.add_run()
    r.text = "me Corp"
    r.font.bold = True
    s.notes_slide.notes_text_frame.text = "Talk about Spain."
    # 3: table and chart
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = "Numbers"
    gf = s.shapes.add_table(3, 3, Inches(0.5), Inches(1.6), Inches(4.2), Inches(1.5))
    gf.name = "Numbers table"
    for ri, row in enumerate([["Region", "Q1", "Q2"], ["North", "120", "135"], ["South", "98", "101"]]):
        for ci, v in enumerate(row):
            gf.table.cell(ri, ci).text = v
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("2025", (1.5, 2.5, 3.5))
    cd.add_series("2026", (2, 3, 4))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.6), Inches(4.5), Inches(4), cd)
    ch.name = "Sales chart"
    # 4: picture (cropped, upscaled), a group with a red box, leftover placeholder text
    s = prs.slides.add_slide(L[6])
    pic = s.shapes.add_picture(str(d / "small.png"), Inches(5.5), Inches(4), Inches(4), Inches(2))
    pic.name = "Photo"
    pic.crop_left = 0.1
    grp = s.shapes.add_group_shape()
    grp.name = "Boxes"
    a = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2))
    a.name = "Box A"
    a.fill.solid()
    a.fill.fore_color.rgb = RGBColor(0xFF, 0, 0)
    a.line.fill.background()
    b = grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.5), Inches(1), Inches(1.5), Inches(1.5))
    b.name = "Box B"
    tb = s.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(0.6))
    tb.name = "Leftover"
    tb.text_frame.text = "Lorem ipsum dolor sit amet"
    # 5: problems, hidden
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = "Problems"

    def box(name: str, x: float, y: float, w: float, h: float, text: str, size: float, color: str | None = None, wrap: bool = True) -> Any:
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        t.name = name
        tf = t.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.text = text
        for run in tf.paragraphs[0].runs:
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        return t

    box("Overflow", 0.5, 1.6, 3, 0.6, "This sentence is far too long for its small box. " * 6, 18)
    box("Tiny", 0.5, 3.2, 4, 0.5, "Small print", 8)
    box("Pale", 0.5, 4.0, 4, 0.6, "Pale text", 20, "DDDDDD")
    off = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.5), Inches(5.2), Inches(3), Inches(1))
    off.name = "Off slide"
    box("Overlap A", 5, 1.6, 3, 1, "First box", 18)
    box("Overlap B", 5.5, 1.9, 3, 1, "Second box", 18)
    s._element.set("show", "0")
    # 6: empty body placeholder
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = "Empty body"
    # 7: too many words
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = "Wordy"
    tf = s.placeholders[1].text_frame
    for i in range(8):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"Point {i + 1} explains in some detail why this particular item matters to the whole team"
        for run in p.runs:
            run.font.size = Pt(14)
    out = d / "foreign.pptx"
    prs.save(str(out))
    return out


def make_pptm(src: Path, dest: Path) -> Path:
    """The foreign deck as a macro-enabled .pptm with a (dummy) VBA project part."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                t = data.decode("utf-8").replace(
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml")
                t = t.replace("</Types>", '<Override PartName="/ppt/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
                data = t.encode("utf-8")
            elif item.filename == "ppt/_rels/presentation.xml.rels":
                data = data.decode("utf-8").replace("</Relationships>", '<Relationship Id="rIdVba1" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>').encode("utf-8")
            zout.writestr(item, data)
        zout.writestr("ppt/vbaProject.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)
    return dest


OUTLINE = """---
title: Project Aurora
author: Selftest
footer: Aurora · Q3 review
---
# Project Aurora
Quarterly review for the leadership team
Selftest · September 2026

## Agenda
- Where we are
- What we shipped
- Next quarter

# Where we are
The state of the product

## Highlights
- Shipped the **new onboarding** flow
  - Activation up from 31% to 44%
- Launched in *three* new markets
- Rebuilt billing with `idempotent` jobs

::: notes
Lead with activation.
:::

<!-- type: kpi -->
## Numbers that matter
- 44% | Activation rate | +13 pts
- €2.4M | Annual recurring revenue | +38% YoY
- 2.1% | Monthly churn | -0.4 pts

<!-- chart: column -->
## Revenue by quarter
| Quarter | 2025 | 2026 |
|---|---|---|
| Q1 | 1.2 | 1.8 |
| Q2 | 1.4 | 2.0 |
| Q3 | 1.5 | 2.4 |

## Team by function
| Function | Headcount | Budget (€k) |
|---|---|---|
| Engineering | 24 | 3,120 |
| Design | 6 | 640 |

<!-- type: timeline -->
## Roadmap
- Oct | Beta | Private beta for partners
- Jan | GA | General availability
- Mar | Mobile | Native apps

## The new home screen
![The new home screen](photo.jpg)

## Why it works
![Diagram](diagram.png)
- One click to everything
- Progress visible from minute one

---

> Aurora cut our onboarding time in half.
> — Maria Lopez, Head of People

## Before and after
:::: columns
::: column
### Before
- Manual invoices
:::
::: column
### After
- Automatic billing
:::
::::

## Deploy script
```bash
make release VERSION=1.4
```

## Questions?
team@example.com
"""


def spec_json(d: Path) -> dict[str, Any]:
    long_items = [f"Item {i}: a reasonably long bullet that describes one more thing to remember" for i in range(1, 31)]
    return {
        "title": "Spec deck", "theme": "midnight", "footer": "Selftest", "author": "Desk",
        "slides": [
            {"type": "title", "title": "Spec deck", "subtitle": "Built from JSON", "meta": "Desk · 2026"},
            {"type": "section", "title": "Part one", "subtitle": "Lists and tables"},
            {"type": "bullets", "title": "Long list", "bullets": long_items},
            {"type": "comparison", "title": "Options", "left": {"heading": "Build", "bullets": ["Control", "Cost"]}, "right": {"heading": "Buy", "bullets": ["Speed", "Support"]}},
            {"type": "table", "title": "Big table", "columns": ["Item", "Units", "Price"], "rows": [[f"Item {i}", i * 3, f"{i * 1.5:.2f}"] for i in range(1, 41)]},
            {"type": "chart", "title": "Share", "chart": {"type": "pie", "categories": ["A", "B", "C"], "series": [{"name": "Share", "values": [50, 30, 20]}]}},
            {"type": "chart", "title": "Trend", "chart": {"type": "line", "categories": ["Jan", "Feb", "Mar", "Apr"], "series": [{"name": "Users", "values": [10, 14, 19, 25]}]}, "note": "Up and to the right."},
            {"type": "chart", "title": "Spread", "chart": {"type": "scatter", "series": [{"name": "Runs", "points": [[1, 2], [2, 3], [3, 5], [4, 4]]}]}},
            {"type": "process", "title": "How we work", "items": [{"title": "Plan", "text": "Scope it"}, {"title": "Build", "text": "Ship it"}, {"title": "Learn", "text": "Measure it"}]},
            {"type": "code", "title": "Snippet", "code": "def double(x):\n    return x * 2\n", "language": "python"},
            {"type": "custom", "title": "Architecture", "elements": [
                {"kind": "card", "x": 0.8, "y": 1.8, "w": 3.6, "h": 1.4, "text": "**API**\nPython, 3 services", "fill": "surface"},
                {"kind": "line", "x": 4.4, "y": 2.5, "w": 1.0, "h": 0, "color": "muted", "width": 2},
                {"kind": "card", "x": 5.4, "y": 1.8, "w": 3.6, "h": 1.4, "text": "**Database**\nPostgres", "fill": "surface"},
                {"kind": "image", "x": 9.5, "y": 1.8, "w": "3in", "h": "2.2in", "image": "photo.jpg", "fit": "cover"},
                {"kind": "text", "x": 0.8, "y": 5.6, "w": 11.7, "h": 0.6, "text": "All traffic stays in the EU", "size": 18, "color": "muted"},
            ]},
            {"type": "bullets", "title": "Backup slide", "bullets": ["Only if asked"], "hidden": True, "notes": ["Line one", "Line two"]},
            {"type": "closing", "title": "Thanks", "contact": "team@example.com"},
        ],
    }


def make_autotitle_chart(d: Path) -> Path:
    """One chart with a single series and an automatic title (it shows the series name, as many real decks do)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Sales chart"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("Sales", (3, 4, 5))
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(5), cd)
    gf.name = "Chart 1"
    ch = gf.chart
    ch.has_title = True  # an empty <c:title>: the automatic title
    ser = ch.plots[0].series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    out = d / "autotitle.pptx"
    prs.save(str(out))
    return out


def make_bombs(d: Path) -> tuple[Path, Path]:
    """A deck with a hugely compressible extra part (a zip bomb) and one whose part declares entities (an XML bomb)."""
    src = d / "foreign.pptx"
    bomb = d / "bomb.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("ppt/media/zeros.bin", b"\0" * (24 * 1024 * 1024))
    xxe = d / "xxe.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(xxe, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                data = data.replace(b"?>", b'?><!DOCTYPE x [<!ENTITY lol "lol">]>', 1)
            zout.writestr(item, data)
    return bomb, xxe


def big_outline(n: int) -> str:
    """A generated outline of n slides mixing bullets (with notes), tables, charts and KPIs."""
    L = ["---", "title: Big deck", "footer: Big deck", "---", "# Big deck", "Generated for the big-file checks", ""]
    for i in range(1, n):
        k = i % 4
        if k == 0:
            L += [f"## Topic {i}", f"- Point one about topic {i}", "  - A detail", f"- Revenue grew in region {i}", "", "::: notes", f"Speaker notes {i}.", ":::", ""]
        elif k == 1:
            L += [f"## Table {i}", "| Region | Units |", "|---|---|", f"| North | {i} |", f"| South | {i * 2} |", ""]
        elif k == 2:
            L += ["<!-- chart: line -->", f"## Trend {i}", "| Month | A |", "|---|---|", "| Jan | 1 |", "| Feb | 3 |", f"| Mar | {i} |", ""]
        else:
            L += ["<!-- type: kpi -->", f"## KPIs {i}", f"- {i}% | Activation | +3 pts", "- €2.4M | ARR", ""]
    return "\n".join(L)


def build_fixtures(d: Path) -> dict[str, Path]:
    from pptx import Presentation

    d.mkdir(parents=True, exist_ok=True)
    _images(d)
    foreign = make_foreign(d)
    (d / "outline.md").write_text(OUTLINE, encoding="utf-8")
    (d / "spec.json").write_text(json.dumps(spec_json(d), indent=1), encoding="utf-8")
    empty = d / "empty.pptx"
    Presentation().save(str(empty))
    (d / "notadeck.pptx").write_text("hello, not a deck\n", encoding="utf-8")
    (d / "legacy.pptx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 1016)
    bomb, xxe = make_bombs(d)
    return {
        "foreign": foreign, "pptm": make_pptm(foreign, d / "macro.pptm"), "outline": d / "outline.md", "spec": d / "spec.json",
        "empty": empty, "notadeck": d / "notadeck.pptx", "legacy": d / "legacy.pptx", "photo": d / "photo.jpg",
        "green": d / "green.png", "small": d / "small.png", "autotitle": make_autotitle_chart(d), "bomb": bomb, "xxe": xxe,
    }


# ── test runner ─────────────────────────────────────────────────────────


class Checks:
    def __init__(self) -> None:
        self.n = 0
        self.failures: list[str] = []

    def ok(self, cond: Any, what: str) -> None:
        self.n += 1
        if not cond:
            self.failures.append(what)
            print(f"FAIL: {what}", file=sys.stderr)


def run(args: list[Any], env: dict[str, str] | None = None, expect: int = 0, stdin: str | None = None, cwd: Path | None = None, lo: bool = False, message: bool = True) -> subprocess.CompletedProcess[str]:
    """Runs scripts/<args[0]> with the same Python. LibreOffice is hidden unless lo=True. A failure must print one
    clean `error: …` line unless message=False (a lint that exits 1 on findings)."""
    cmd = [PY, str(HERE / str(args[0])), *[str(a) for a in args[1:]]]
    e = dict(os.environ)
    if not lo:
        e.update(NO_LO)
    e.update(env or {})
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=e, input=stdin, cwd=cwd, timeout=240)
    if r.returncode != expect:
        raise AssertionError(f"{' '.join(str(a) for a in args[:3])} … exited {r.returncode} (expected {expect}):\n{r.stderr[-1500:]}\n{r.stdout[-500:]}")
    if expect != 0 and message and not clean_error(r.stderr):
        raise AssertionError(f"{args[0]}: error output is not one clean message: {r.stderr[-800:]}")
    return r


def jrun(args: list[Any], **kw: Any) -> Any:
    r = run([*args, "--format", "json"], **kw)
    return json.loads(r.stdout)


def clean_error(stderr: str) -> bool:
    """An expected failure prints `error: …`, never a traceback or a bare Python exception name."""
    return "Traceback" not in stderr and stderr.startswith("error: ") and not re.match(r"error: [A-Za-z]+(Error|Exception):", stderr)


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def png_info(p: Path) -> tuple[tuple[int, int], float]:
    """(size, standard deviation of the grey levels): a blank render has almost none."""
    from PIL import Image, ImageStat

    with Image.open(p) as im:
        return im.size, ImageStat.Stat(im.convert("L")).stddev[0]


def pixel(p: Path, x_frac: float, y_frac: float) -> tuple[int, int, int]:
    from PIL import Image

    with Image.open(p) as im:
        rgb = im.convert("RGB")
        return rgb.getpixel((int(x_frac * rgb.width), int(y_frac * rgb.height)))  # type: ignore[return-value]


def pdf_pages(p: Path) -> int:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(p))
    try:
        return len(doc)
    finally:
        doc.close()


def zip_names(p: Path) -> list[str]:
    with zipfile.ZipFile(p) as z:
        return z.namelist()


def main_content_type(p: Path) -> str:
    with zipfile.ZipFile(p) as z:
        t = z.read("[Content_Types].xml").decode("utf-8")
    m = re.search(r'PartName="/ppt/presentation.xml" ContentType="([^"]+)"', t)
    return m.group(1) if m else ""


def lo_available() -> bool:
    sys.path.insert(0, str(HERE))
    from _render import find_soffice

    return bool(find_soffice())


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--keep", action="store_true", help="keep the temp folder and print its path")
    ap.add_argument("--fixtures", help="only write the fixtures into this folder")
    ap.add_argument("--only", help="run only test functions whose name contains this")
    a = ap.parse_args()
    if a.fixtures:
        fx = build_fixtures(Path(a.fixtures))
        print("\n".join(str(p) for p in fx.values()))
        return 0
    t0 = time.time()
    tmp = Path(tempfile.mkdtemp(prefix="presentations-selftest-"))
    os.environ["DESK_FILE_CACHE"] = str(tmp / "cache")  # a private file cache: cold first calls, measurable hits
    os.environ.pop("DESK_NO_CACHE", None)
    c = Checks()
    try:
        fx = build_fixtures(tmp / "fx")
        before = {k: sha(p) for k, p in fx.items()}
        have_lo = lo_available()
        tests = [(name, fn) for name, fn in globals().items() if name.startswith("test_") and callable(fn)]
        for name, fn in tests:
            if a.only and a.only not in name:
                continue
            if name.startswith("test_lo_") and not have_lo:
                continue
            work = tmp / name
            work.mkdir()
            try:
                fn(c, fx, work)
            except AssertionError as e:
                c.ok(False, f"{name}: {e}")
            except Exception as e:  # noqa: BLE001
                c.ok(False, f"{name}: {type(e).__name__}: {e}")
        c.ok(all(sha(p) == before[k] for k, p in fx.items()), "no script modified an input file")
    finally:
        if a.keep:
            print(f"kept {tmp}", file=sys.stderr)
        else:
            shutil.rmtree(tmp, ignore_errors=True)
    dt = time.time() - t0
    if c.failures:
        print(f"FAILED: {len(c.failures)} of {c.n} checks failed in {dt:.1f} s", file=sys.stderr)
        return 1
    print(f"ok: {c.n} checks in {dt:.1f} s")
    return 0


# ── tests ───────────────────────────────────────────────────────────────


def test_help(c: Checks, fx: dict[str, Path], w: Path) -> None:
    for s in SCRIPTS:
        t = time.time()
        r = run([s, "--help"])
        dt = time.time() - t
        c.ok("examples:" in r.stdout and "python3 scripts/" in r.stdout, f"{s} --help has examples")
        c.ok(dt < 1.5, f"{s} --help took {dt:.2f} s")
        e = dict(os.environ, **NO_LO)
        imp = subprocess.run([PY, "-X", "importtime", str(HERE / s), "--help"], capture_output=True, text=True, encoding="utf-8", env=e, timeout=60)
        loaded = {ln.rsplit("|", 1)[-1].strip().split(".")[0] for ln in imp.stderr.splitlines() if ln.startswith("import time:")}
        heavy = sorted(loaded & set(HEAVY))
        c.ok(not heavy, f"{s} --help imports no heavy modules (got {heavy})")


def test_info(c: Checks, fx: dict[str, Path], w: Path) -> None:
    d = jrun(["pptx_info.py", fx["foreign"]])
    c.ok(d["slide_count"] == 7, "info: 7 slides")
    c.ok(d["aspect"] == "4:3" and d["slide_size_in"] == [10.0, 7.5], f"info: 4:3 at 10×7.5 in ({d['aspect']}, {d['slide_size_in']})")
    c.ok(d["hidden_slides"] == [5], "info: hidden slide 5")
    c.ok(d["properties"].get("title") == "Acme Quarterly deck" and d["properties"].get("author") == "Finance Team", "info: document properties")
    c.ok(d["content"].get("charts") == 1 and d["content"].get("tables") == 1 and d["content"].get("pictures") == 1, f"info: content counts {d['content']}")
    c.ok(d["slides_with_notes"] == 2, "info: 2 slides with notes")
    c.ok(d["masters"] and len(d["masters"][0]["layouts"]) == 11, "info: one master with 11 layouts")
    c.ok(d["macros"] is False, "info: no macros in .pptx")
    c.ok(any(s["title"] == "Plan" and s["layout"] == "Title and Content" for s in d["slides"]), "info: per-slide layout")
    md = run(["pptx_info.py", fx["foreign"]]).stdout
    c.ok(md.startswith("# foreign.pptx") and "Title and Content" in md, "info: Markdown summary")
    m = jrun(["pptx_info.py", fx["pptm"]])
    c.ok(m["macros"] is True and m["format"] == "pptm", "info: .pptm macros detected")
    e = jrun(["pptx_info.py", fx["empty"]])
    c.ok(e["slide_count"] == 0, "info: empty deck")


def test_read(c: Checks, fx: dict[str, Path], w: Path) -> None:
    d = jrun(["pptx_read.py", fx["foreign"]])
    s = {x["number"]: x for x in d["slides"]}
    c.ok([s[i]["title"] for i in range(1, 8)] == ["Acme Quarterly", "Plan", "Numbers", "", "Problems", "Empty body", "Wordy"], "read: titles in order")
    c.ok(s[1].get("transition") == "fade" and s[1]["notes"] == "Welcome everyone.", "read: transition and notes")
    c.ok(s[5]["hidden"] is True and s[4]["hidden"] is False, "read: hidden flag")
    body = next(sh for sh in s[2]["shapes"] if sh.get("placeholder", {}).get("idx") == 1)
    c.ok([(p["level"], p["text"]) for p in body["paragraphs"]] == [(0, "Grow revenue"), (1, "Enter Spain"), (0, "Partner with Acme Corp")], "read: paragraphs with levels")
    c.ok(body["paragraphs"][2]["markdown"] == "Partner with Ac**me Corp**", f"read: emphasis as Markdown ({body['paragraphs'][2]['markdown']})")
    tbl = next(sh for sh in s[3]["shapes"] if sh["kind"] == "table")
    c.ok(tbl["table"]["rows"][1] == ["North", "120", "135"], "read: table rows")
    ch = next(sh for sh in s[3]["shapes"] if sh["kind"] == "chart")
    c.ok(ch["chart"]["categories"] == ["Q1", "Q2", "Q3"] and [x["values"] for x in ch["chart"]["series"]] == [[1.5, 2.5, 3.5], [2, 3, 4]], "read: chart data")
    c.ok(ch["name"] == "Sales chart" and ch["x"] == 5.0 and ch["w"] == 4.5, "read: shape name and position in inches")
    pic = next(sh for sh in s[4]["shapes"] if sh["kind"] == "picture")
    c.ok(pic["image"]["pixels"] == [200, 100] and pic["image"]["crop"]["left"] == 0.1, "read: picture pixels and crop")
    kids = [sh["name"] for sh in s[4]["shapes"] if sh.get("group") is not None]
    c.ok(kids == ["Box A", "Box B"], f"read: group children ({kids})")
    only = jrun(["pptx_read.py", fx["foreign"], "--slides", "2-3,last", "--no-notes"])
    c.ok([x["number"] for x in only["slides"]] == [2, 3, 7] and not only["slides"][0].get("notes"), "read: --slides and --no-notes")
    md = run(["pptx_read.py", fx["foreign"], "--slides", "3"]).stdout
    c.ok("| Region | Q1 | Q2 |" in md and "Q1" in md and "2026" in md, "read: Markdown table and chart")
    md = run(["pptx_read.py", fx["foreign"], "--shapes", "--slides", "4"]).stdout
    c.ok('"Box A"' in md and "cropped L10%" in md, "read: --shapes lists group members and crop")
    lay = jrun(["pptx_read.py", fx["foreign"], "--layouts", "--slides", "1"])
    c.ok(len(lay.get("layouts", [])) == 11, "read: --layouts")


def test_create_markdown(c: Checks, fx: dict[str, Path], w: Path) -> None:
    out = w / "made.pptx"
    d = jrun(["pptx_create.py", out, "--input", fx["outline"], "--preview"])
    types = [s["type"] for s in d["slides"]]
    want = ["title", "agenda", "section", "bullets", "kpi", "chart", "table", "timeline", "image", "image-text", "quote", "two-column", "code", "closing"]
    c.ok(types == want, f"create md: slide types {types}")
    prev = w / "made-preview.png"
    c.ok(prev.exists() and png_info(prev)[1] > 5, "create md: --preview contact sheet")
    r = jrun(["pptx_read.py", out])
    s = {x["number"]: x for x in r["slides"]}
    c.ok(s[4]["notes"] == "Lead with activation.", "create md: notes")
    texts = " ".join(sh.get("text", "") for sh in s[5]["shapes"])
    c.ok("€2.4M" in texts and "+38% YoY" in texts, "create md: KPI values and deltas")
    ch = next(sh for sh in s[6]["shapes"] if sh["kind"] == "chart")
    c.ok(ch["chart"]["categories"] == ["Q1", "Q2", "Q3"] and [x["name"] for x in ch["chart"]["series"]] == ["2025", "2026"], "create md: native chart from the table")
    tbl = next(sh for sh in s[7]["shapes"] if sh["kind"] == "table")
    c.ok(len(tbl["table"]["rows"]) == 3 and "3,120" in tbl["table"]["rows"][1], "create md: table")
    c.ok(any(sh["kind"] == "picture" for sh in s[9]["shapes"]), "create md: image slide has the picture")
    foot = [sh.get("text") for sh in s[4]["shapes"] if sh["name"] == "Footer"]
    c.ok(foot == ["Aurora · Q3 review"], "create md: footer on content slides")
    c.ok(not any(sh["name"] == "Footer" for sh in s[1]["shapes"]), "create md: no footer on the title slide")
    c.ok("make release" in " ".join(sh.get("text", "") for sh in s[13]["shapes"]), "create md: code slide")
    info = jrun(["pptx_info.py", out])
    c.ok(info["aspect"] == "16:9" and info["theme"]["name"] == "Desk clean", "create md: 16:9 with the clean theme")
    c.ok(info["properties"].get("title") == "Project Aurora" and info["properties"].get("author") == "Selftest", "create md: properties from front matter")
    lint = jrun(["pptx_lint.py", out])
    c.ok(lint["counts"]["error"] == 0 and lint["counts"]["warning"] == 0, f"create md: lint clean ({[i['rule'] for i in lint['issues'] if i['severity'] != 'note']})")
    r2 = run(["pptx_create.py", out, "--input", fx["outline"]], expect=1)
    c.ok("--force" in r2.stderr, "create: refuses to overwrite without --force")
    md = "# One\n\n## Two\n- a\n- b\n"
    d2 = jrun(["pptx_create.py", w / "stdin.pptx", "--input", "-", "--theme", '{"base": "paper", "accent": "#0A7F6F"}', "--no-numbers"], stdin=md)
    c.ok([s["type"] for s in d2["slides"]] == ["title", "bullets"], "create: Markdown from stdin with theme overrides")
    info = jrun(["pptx_info.py", w / "stdin.pptx"])
    c.ok(info["theme"]["colors"].get("accent1", "").upper().lstrip("#") == "0A7F6F", f"create: accent override in the theme ({info['theme']['colors'].get('accent1')})")


def test_create_json(c: Checks, fx: dict[str, Path], w: Path) -> None:
    out = w / "spec.pptx"
    d = jrun(["pptx_create.py", out, "--spec", fx["spec"]])
    titles = [s["title"] for s in d["slides"]]
    parts = [t for t in titles if t.startswith("Long list")]
    c.ok(len(parts) >= 3 and all(re.search(r"\((cont\.|\d+/\d+)\)$", t) for t in parts[1:]), f"create json: long list continues ({parts})")
    big = [s for s in d["slides"] if s["title"].startswith("Big table")]
    c.ok(len(big) >= 2, "create json: long table split over slides")
    r = jrun(["pptx_read.py", out])
    tables = [sh["table"]["rows"] for s in r["slides"] if s["title"].startswith("Big table") for sh in s["shapes"] if sh["kind"] == "table"]
    c.ok(len(tables) >= 2 and all(t[0] == ["**Item**", "**Units**", "**Price**"] or [re.sub(r"\*", "", x) for x in t[0]] == ["Item", "Units", "Price"] for t in tables), "create json: header repeated on continuation")
    c.ok(sum(len(t) - 1 for t in tables) == 40, "create json: all 40 table rows kept")
    charts = [sh["chart"] for s in r["slides"] for sh in s["shapes"] if sh["kind"] == "chart"]
    kinds = [ch["types"][0] for ch in charts]
    c.ok(kinds == ["pie", "line", "scatter"], f"create json: chart types {kinds}")
    hidden = [s["title"] for s in r["slides"] if s["hidden"]]
    c.ok(hidden == ["Backup slide"], "create json: hidden slide")
    backup = next(s for s in r["slides"] if s["title"] == "Backup slide")
    c.ok(backup["notes"] == "Line one\nLine two", "create json: notes from a list")
    custom = next(s for s in r["slides"] if s["title"] == "Architecture")
    kinds = sorted(sh["kind"] for sh in custom["shapes"])
    c.ok("picture" in kinds and "line" in kinds and custom["shapes"][-1]["name"] != "", "create json: custom elements")
    img = next(sh for sh in custom["shapes"] if sh["kind"] == "picture")
    c.ok(abs(img["w"] - 3.0) < 0.01 and abs(img["h"] - 2.2) < 0.01, "create json: lengths with units")
    info = jrun(["pptx_info.py", out])
    c.ok(info["theme"]["name"] == "Desk midnight", "create json: theme from the spec")
    lint = run(["pptx_lint.py", out, "--fail-on", "warning"])
    c.ok("0 errors, 0 warnings" in lint.stdout, "create json: lint clean with --fail-on warning")
    bad = run(["pptx_create.py", w / "bad.pptx", "--spec", '{"slides": [{"type": "nonsense"}]}'], expect=2)
    c.ok("nonsense" in bad.stderr, "create json: unknown slide type is a usage error")


def test_create_template(c: Checks, fx: dict[str, Path], w: Path) -> None:
    spec = json.dumps({"slides": [
        {"type": "title", "title": "From the template", "subtitle": "Brand check"},
        {"type": "bullets", "title": "Points", "bullets": ["One", {"text": "Two", "children": ["Two a"]}]},
        {"type": "two-column", "title": "Sides", "left": {"bullets": ["Left"]}, "right": {"bullets": ["Right"]}},
        {"type": "chart", "title": "Chart", "chart": {"type": "bar", "categories": ["A", "B"], "series": [{"name": "S", "values": [1, 2]}]}},
        {"type": "bullets", "title": "Picked layout", "bullets": ["x"], "layout": "Comparison"},
    ]})
    out = w / "tpl.pptx"
    d = jrun(["pptx_create.py", out, "--template", fx["foreign"], "--spec", spec])
    layouts = [s["layout"] for s in d["slides"]]
    c.ok(layouts[0] == "Title Slide" and layouts[1] == "Title and Content" and layouts[4] == "Comparison", f"create template: template layouts used {layouts}")
    info = jrun(["pptx_info.py", out])
    c.ok(info["slide_count"] == 5 and info["aspect"] == "4:3" and info["theme"]["name"] == "Office Theme", "create template: template size and theme, template slides dropped")
    r = jrun(["pptx_read.py", out, "--slides", "2"])
    body = next(sh for sh in r["slides"][0]["shapes"] if sh.get("placeholder", {}).get("idx") == 1)
    c.ok([(p["level"], p["text"]) for p in body["paragraphs"]] == [(0, "One"), (0, "Two"), (1, "Two a")], "create template: bullets in the body placeholder")
    keep = jrun(["pptx_create.py", w / "keep.pptx", "--template", fx["foreign"], "--spec", spec, "--keep-template-slides"])
    c.ok(len(keep["slides"]) >= 5 and jrun(["pptx_info.py", w / "keep.pptx"])["slide_count"] == 12, "create template: --keep-template-slides")
    potx = w / "brand.potx"
    run(["pptx_convert.py", fx["foreign"], potx])
    c.ok(main_content_type(potx).endswith("presentationml.template.main+xml"), "convert: .potx content type")
    d = jrun(["pptx_create.py", w / "from-potx.pptx", "--template", potx, "--spec", spec])
    c.ok(len(d["slides"]) == 5 and main_content_type(w / "from-potx.pptx").endswith("presentation.main+xml"), "create template: from a .potx, saved as a .pptx")


def test_edit(c: Checks, fx: dict[str, Path], w: Path) -> None:
    shutil.copy(fx["green"], w / "green.png")
    ops = [
        {"op": "replace", "find": "Acme Corp", "replace": "Globex Inc"},
        {"op": "replace", "find": r"Q(\d)", "replace": r"Quarter \1", "regex": True, "slides": "3"},
        {"op": "set_title", "slide": 2, "text": "The plan"},
        {"op": "set_notes", "slide": 1, "text": "New notes."},
        {"op": "add_slide", "after": 2, "id": "risks", "type": "bullets", "title": "Risks", "bullets": ["Hiring", "Pricing"]},
        {"op": "set_notes", "slide": "risks", "text": "Added slide notes."},
        {"op": "duplicate_slide", "slide": 3, "id": "copy"},
        {"op": "update_chart", "slide": 3, "shape": "Sales chart", "categories": ["A", "B"], "series": [{"name": "X", "values": [5, 6]}]},
        {"op": "set_cell", "slide": 3, "shape": "Numbers table", "row": 1, "col": 1, "text": "999"},
        {"op": "set_table", "slide": "copy", "shape": "Numbers table", "rows": [["Region", "Q1", "Q2"], ["East", "1", "2"], ["West", "3", "4"], ["Mid", "5", "6"]], "exact": True},
        {"op": "hide_slide", "slides": "4"},
        {"op": "unhide_slide", "slides": "5"},
        {"op": "set_shape", "slide": 4, "shape": "Box A", "x": 3, "fill": "#00FF00", "text": "Moved"},
        {"op": "delete_shape", "slide": 4, "shape": "Leftover"},
        {"op": "add_text", "slide": 4, "x": 1, "y": 6, "w": 5, "h": 0.6, "text": "Added note", "size": 20},
        {"op": "replace_image", "slide": 4, "shape": "Photo", "image": "green.png", "fit": "cover"},
        {"op": "set_background", "slides": "2", "color": "#FFEEDD"},
        {"op": "set_properties", "title": "Edited deck", "author": "Selftest"},
        {"op": "move_slide", "slide": 6, "after": 0},
        {"op": "delete_slide", "slides": "7"},
    ]
    (w / "ops.json").write_text(json.dumps(ops), encoding="utf-8")
    out = w / "edited.pptx"
    d = jrun(["pptx_edit.py", fx["foreign"], "--out", out, "--ops", w / "ops.json"])
    res = d["ops"]
    c.ok(res[0].get("replaced") == 1 and res[1].get("replaced") == 2, f"edit: replace counts ({res[0].get('replaced')}, {res[1].get('replaced')})")
    r = jrun(["pptx_read.py", out])
    titles = [s["title"] for s in r["slides"]]
    c.ok(titles == ["Empty body", "Acme Quarterly", "The plan", "Risks", "Numbers", "Numbers", "", "Problems"], f"edit: final order {titles}")
    s = r["slides"]
    plan = s[2]
    c.ok(any("Partner with Globex Inc" == p["text"] for sh in plan["shapes"] for p in sh.get("paragraphs", [])), "edit: replace across runs")
    c.ok(s[1]["notes"] == "New notes." and s[3]["notes"] == "Added slide notes.", "edit: notes, including on an added slide by id")
    orig_chart = next(sh["chart"] for sh in s[4]["shapes"] if sh["kind"] == "chart")
    copy_chart = next(sh["chart"] for sh in s[5]["shapes"] if sh["kind"] == "chart")
    c.ok(orig_chart["categories"] == ["A", "B"] and orig_chart["series"][0]["values"] == [5, 6], "edit: update_chart")
    c.ok(copy_chart["categories"] == ["Q1", "Q2", "Q3"], "edit: the duplicate keeps its own chart data")
    orig_tbl = next(sh["table"]["rows"] for sh in s[4]["shapes"] if sh["kind"] == "table")
    copy_tbl = next(sh["table"]["rows"] for sh in s[5]["shapes"] if sh["kind"] == "table")
    c.ok(orig_tbl[0] == ["Region", "Quarter 1", "Quarter 2"] and orig_tbl[1][1] == "999", f"edit: regex replace in cells and set_cell ({orig_tbl[:2]})")
    c.ok(len(copy_tbl) == 4 and copy_tbl[3] == ["Mid", "5", "6"], "edit: set_table with an added row")
    pics = s[6]
    c.ok(pics["hidden"] is True and s[7]["hidden"] is False, "edit: hide and unhide")
    names = [sh["name"] for sh in pics["shapes"]]
    c.ok("Leftover" not in names and any(sh.get("text") == "Added note" for sh in pics["shapes"]), "edit: delete_shape and add_text")
    box = next(sh for sh in pics["shapes"] if sh["name"] == "Box A")
    c.ok(abs(box["x"] - 3.0) < 0.01 and box.get("text") == "Moved", "edit: set_shape moves and sets text")
    photo = next(sh for sh in pics["shapes"] if sh["name"] == "Photo")
    c.ok(photo["image"]["pixels"] == [400, 300] and (photo["w"], photo["h"]) == (4.0, 2.0), "edit: replace_image keeps the frame")
    info = jrun(["pptx_info.py", out])
    c.ok(info["properties"].get("title") == "Edited deck" and info["properties"].get("author") == "Selftest", "edit: set_properties")
    charts = [n for n in zip_names(out) if re.match(r"ppt/charts/chart\d+\.xml$", n)]
    c.ok(len(charts) == 2, f"edit: the chart part was cloned ({charts})")
    png = w / "rendered"
    rr = jrun(["pptx_render.py", out, "--slides", "7", "--out-dir", png])
    shot = Path(rr["slides"][0]["png"])
    g = pixel(shot, 0.32, 2.75 / 7.5)
    c.ok(g[1] > 200 and g[0] < 60 and g[2] < 60, f"edit: the moved box renders green at its new place {g}")
    bad = run(["pptx_edit.py", fx["foreign"], "--out", w / "bad.pptx", "--ops", '[{"op": "set_shape", "slide": 2, "shape": "Nope", "x": 1}]'], expect=2)
    c.ok("Nope" in bad.stderr and "Title 1" in bad.stderr, "edit: unknown shape lists the shapes")
    bad = run(["pptx_edit.py", fx["foreign"], "--out", w / "bad.pptx", "--ops", '[{"op": "explode"}]'], expect=2)
    c.ok("unknown op" in bad.stderr, "edit: unknown op")
    bad = run(["pptx_edit.py", fx["foreign"], "--out", fx["foreign"], "--ops", '[{"op": "set_notes", "slide": 1, "text": "x"}]'], expect=1)
    c.ok("input" in bad.stderr, "edit: refuses to write over the input")
    c.ok(not (w / "bad.pptx").exists(), "edit: nothing written when an op fails")
    e = run(["pptx_edit.py", fx["foreign"], "--out", w / "bad.pptx", "--ops", "[]"], expect=2)
    c.ok("non-empty" in e.stderr, "edit: empty ops list")
    m = jrun(["pptx_edit.py", fx["pptm"], "--out", w / "m2.pptm", "--ops", '[{"op": "set_title", "slide": 2, "text": "Macro plan"}]'])
    c.ok("ppt/vbaProject.bin" in zip_names(w / "m2.pptm") and not m["warnings"], "edit: .pptm keeps its macros")
    m = jrun(["pptx_edit.py", fx["pptm"], "--out", w / "m2.pptx", "--ops", '[{"op": "set_notes", "slide": 1, "text": "x"}]'])
    c.ok("ppt/vbaProject.bin" not in zip_names(w / "m2.pptx") and any("macro" in x for x in m["warnings"]), "edit: saving a .pptm as .pptx drops macros with a warning")


def test_render_builtin(c: Checks, fx: dict[str, Path], w: Path) -> None:
    out = w / "r"
    d = jrun(["pptx_render.py", fx["foreign"], "--out-dir", out])
    c.ok(d["engine"] == "builtin" and "built-in" in d["engine_label"], "render: built-in engine named")
    c.ok(len(d["slides"]) == 7 and all((s["width"], s["height"]) == (1568, 1176) for s in d["slides"]), "render: 7 PNGs at 1568×1176")
    shot4 = Path(d["slides"][3]["png"])
    c.ok(shot4.name == "foreign-04.png", "render: <stem>-NN.png names")
    red = pixel(shot4, 0.2, 2 / 7.5)
    c.ok(red[0] > 200 and red[1] < 60 and red[2] < 60, f"render: the grouped red box is drawn where it belongs {red}")
    yellow = pixel(shot4, 8.5 / 10, 5 / 7.5)
    c.ok(yellow[0] > 200 and yellow[1] > 200 and yellow[2] < 80, f"render: the cropped picture is drawn {yellow}")
    c.ok(all(png_info(Path(s["png"]))[1] > 3 for s in d["slides"]), "render: no blank slides")
    again = run(["pptx_render.py", fx["foreign"], "--out-dir", out], expect=1)
    c.ok("--force" in again.stderr, "render: refuses to overwrite without --force")
    d = jrun(["pptx_render.py", fx["foreign"], "--out-dir", w / "s", "--sheet-only", "--skip-hidden"])
    files = sorted(p.name for p in (w / "s").iterdir())
    c.ok(files == ["foreign-sheet.png"] and len(d["slides"]) == 6, f"render: --sheet-only with --skip-hidden ({files})")
    d = jrun(["pptx_render.py", fx["foreign"], "--out-dir", w / "small", "--slides", "2", "--width", "640", "--sheet"])
    c.ok((d["slides"][0]["width"], d["slides"][0]["height"]) == (640, 480) and Path(d["sheet"]).exists(), "render: --width and --sheet")
    md = run(["pptx_render.py", fx["foreign"], "--out-dir", w / "md", "--slides", "1"]).stdout
    c.ok("view_image" in md and "built-in renderer" in md, "render: Markdown names the engine and says to look")
    e = run(["pptx_render.py", fx["empty"], "--out-dir", w / "e"], expect=1)
    c.ok("no slides" in e.stderr, "render: empty deck")
    e = run(["pptx_render.py", fx["foreign"], "--engine", "libreoffice", "--out-dir", w / "e"], expect=1)
    c.ok("LibreOffice" in e.stderr, "render: --engine libreoffice without LibreOffice")


def test_render_created(c: Checks, fx: dict[str, Path], w: Path) -> None:
    shutil.copy(fx["photo"], w / "photo.jpg")
    shutil.copy(fx["outline"].parent / "diagram.png", w / "diagram.png")
    shutil.copy(fx["outline"], w / "outline.md")
    deck = w / "made.pptx"
    run(["pptx_create.py", deck, "--input", w / "outline.md", "--theme", "vivid"])
    d = jrun(["pptx_render.py", deck, "--out-dir", w / "r", "--sheet"])
    c.ok(len(d["slides"]) == 14 and all((s["width"], s["height"]) == (1568, 882) for s in d["slides"]), "render created: 14 slides at 16:9")
    c.ok(all(png_info(Path(s["png"]))[1] > 3 for s in d["slides"]), "render created: no blank slides")
    size, _ = png_info(Path(d["sheet"]))
    c.ok(max(size) <= 1568, f"render created: contact sheet fits the vision limit {size}")


def test_lint(c: Checks, fx: dict[str, Path], w: Path) -> None:
    d = jrun(["pptx_lint.py", fx["foreign"]])
    found = {(i["slide"], i["rule"], i["shape"]) for i in d["issues"]}
    rules = {(i["slide"], i["rule"]) for i in d["issues"]}
    c.ok((5, "overflow", "Overflow") in found, "lint: measured overflow")
    c.ok((5, "font-size", "Tiny") in found, "lint: tiny font")
    c.ok((5, "contrast", "Pale") in found, "lint: low contrast")
    c.ok((5, "off-slide", "Off slide") in found, "lint: off-slide shape")
    c.ok(any(r == (5, "overlap") for r in rules), "lint: overlapping text boxes")
    c.ok(any(i["slide"] == 6 and i["rule"] == "empty-placeholder" for i in d["issues"]), "lint: empty placeholder")
    c.ok((4, "leftover", "Leftover") in found, "lint: leftover Lorem ipsum")
    c.ok((7, "words") in rules, "lint: too many words")
    c.ok((4, "image-resolution", "Photo") in found, "lint: upscaled picture")
    c.ok((5, "hidden") in rules, "lint: hidden slide noted")
    c.ok(not any(i["slide"] in (1, 2, 3) and i["severity"] != "note" for i in d["issues"]), "lint: clean slides have no warnings")
    c.ok(all(i["fix"] for i in d["issues"]), "lint: every issue suggests a fix")
    run(["pptx_lint.py", fx["foreign"], "--fail-on", "warning"], expect=1, message=False)
    c.ok(True, "lint: --fail-on warning exits 1")
    only = jrun(["pptx_lint.py", fx["foreign"], "--slides", "1-3"])
    c.ok(only["slides_checked"] == 3 and only["counts"]["warning"] == 0 and only["counts"]["error"] == 0, "lint: --slides")
    md = run(["pptx_lint.py", fx["foreign"]]).stdout
    c.ok(md.startswith("# Lint: foreign.pptx") and "## Slide 5" in md and "Fix:" in md, "lint: Markdown report")
    loose = jrun(["pptx_lint.py", fx["foreign"], "--min-size", "6", "--max-words", "500"])
    c.ok(not any(i["rule"] in ("font-size", "words") for i in loose["issues"]), "lint: thresholds are options")


def test_convert(c: Checks, fx: dict[str, Path], w: Path) -> None:
    d = jrun(["pptx_convert.py", fx["foreign"], w / "deck.pdf"])
    c.ok(pdf_pages(w / "deck.pdf") == 7 and "built-in" in d["how"], "convert: PDF with every slide (built-in)")
    run(["pptx_convert.py", fx["foreign"], w / "part.pdf", "--slides", "2-3"])
    c.ok(pdf_pages(w / "part.pdf") == 2, "convert: PDF of some slides")
    run(["pptx_convert.py", fx["foreign"], w / "deck.md"])
    md = (w / "deck.md").read_text(encoding="utf-8")
    c.ok("Acme Quarterly" in md and "Talk about Spain." in md and "| North | 120 | 135 |" in md, "convert: Markdown outline with notes and tables")
    run(["pptx_convert.py", fx["foreign"], w / "deck.txt"])
    c.ok("--- Slide 3: Numbers" in (w / "deck.txt").read_text(encoding="utf-8"), "convert: plain text")
    run(["pptx_convert.py", fx["foreign"], w / "deck.json"])
    c.ok(len(json.loads((w / "deck.json").read_text(encoding="utf-8"))["slides"]) == 7, "convert: JSON")
    run(["pptx_convert.py", fx["foreign"], w / "show.ppsx"])
    c.ok(main_content_type(w / "show.ppsx").endswith("slideshow.main+xml"), "convert: .ppsx content type")
    run(["pptx_convert.py", w / "show.ppsx", w / "back.pptx"])
    c.ok(main_content_type(w / "back.pptx").endswith("presentation.main+xml") and jrun(["pptx_info.py", w / "back.pptx"])["slide_count"] == 7, "convert: .ppsx back to .pptx")
    d = jrun(["pptx_convert.py", fx["pptm"], w / "nomacro.pptx"])
    c.ok("ppt/vbaProject.bin" not in zip_names(w / "nomacro.pptx") and d.get("warnings"), "convert: .pptm → .pptx drops macros and says so")
    pngs = jrun(["pptx_convert.py", fx["foreign"], w / "pngs", "--to", "png", "--slides", "1", "--width", "800"])
    c.ok(len(pngs["pngs"]) == 1 and png_info(Path(pngs["pngs"][0]))[0] == (800, 600), "convert: PNG")
    shutil.copy(fx["outline"], w / "o.md")
    for n in ("photo.jpg", "diagram.png"):
        shutil.copy(fx["outline"].parent / n, w / n)
    run(["pptx_convert.py", w / "o.md", w / "o.pptx", "--theme", "slate"])
    c.ok(jrun(["pptx_info.py", w / "o.pptx"])["theme"]["name"] == "Desk slate", "convert: Markdown → pptx")
    e = run(["pptx_convert.py", fx["foreign"], w / "deck.pdf"], expect=1)
    c.ok("--force" in e.stderr, "convert: refuses to overwrite")
    run(["pptx_convert.py", fx["foreign"], w / "deck.pdf", "--force"])
    c.ok(True, "convert: --force overwrites")
    e = run(["pptx_convert.py", fx["foreign"], w / "x.odp"], expect=1)
    c.ok("LibreOffice" in e.stderr, "convert: .odp needs LibreOffice")
    (w / "in.odp").write_bytes(b"PK\x03\x04not really")
    e = run(["pptx_read.py", w / "in.odp"], expect=1)
    c.ok("LibreOffice" in e.stderr and "pptx" in e.stderr, "read: .odp without LibreOffice asks for a .pptx")


def test_errors(c: Checks, fx: dict[str, Path], w: Path) -> None:
    e = run(["pptx_read.py", w / "missing.pptx"], expect=1)
    c.ok("does not exist" in e.stderr, "errors: missing file")
    e = run(["pptx_info.py", fx["notadeck"]], expect=1)
    c.ok("not a PowerPoint file" in e.stderr, "errors: not a ZIP")
    e = run(["pptx_lint.py", fx["legacy"]], expect=1)
    c.ok("legacy binary" in e.stderr, "errors: legacy or encrypted file")
    e = run(["pptx_read.py", fx["foreign"], "--slides", "9"], expect=2)
    c.ok("9" in e.stderr, "errors: slide out of range")
    e = run(["pptx_create.py", w / "x.docx", "--input", fx["outline"]], expect=1)
    c.ok(".pptx" in e.stderr, "errors: create needs a presentation extension")
    r = jrun(["pptx_read.py", fx["empty"]])
    c.ok(r["slide_count"] == 0 and r["slides"] == [], "errors: an empty deck reads as empty")


def test_lo_render(c: Checks, fx: dict[str, Path], w: Path) -> None:
    t0 = time.time()
    d = jrun(["pptx_render.py", fx["foreign"], "--out-dir", w / "r", "--slides", "3-5"], lo=True)
    cold = time.time() - t0
    c.ok(d["engine"] == "libreoffice" and [s["number"] for s in d["slides"]] == [3, 4, 5], "LibreOffice: render picks LibreOffice, hidden slide included")
    c.ok(all((s["width"], s["height"]) == (1568, 1176) and png_info(Path(s["png"]))[1] > 3 for s in d["slides"]), "LibreOffice: PNG sizes")
    red = pixel(Path(d["slides"][1]["png"]), 0.2, 2 / 7.5)
    c.ok(red[0] > 200 and red[1] < 60 and red[2] < 60, f"LibreOffice: red box where the built-in renderer puts it {red}")
    t0 = time.time()
    d = jrun(["pptx_convert.py", fx["foreign"], w / "lo.pdf"], lo=True)
    warm = time.time() - t0
    c.ok("LibreOffice" in d["how"] and pdf_pages(w / "lo.pdf") == 7, "LibreOffice: PDF with hidden slides")
    c.ok(warm * 3 < cold, f"LibreOffice: the PDF made for the render is reused by the conversion ({cold:.2f} s → {warm:.2f} s)")
    b = jrun(["pptx_render.py", fx["foreign"], "--engine", "builtin", "--slides", "1", "--out-dir", w / "b"], lo=True)
    c.ok("--engine libreoffice renders exactly" in b["engine_label"], f"built-in engine with LibreOffice installed says how to render exactly ({b['engine_label']})")


def test_lo_formats(c: Checks, fx: dict[str, Path], w: Path) -> None:
    run(["pptx_convert.py", fx["foreign"], w / "deck.odp"], lo=True)
    c.ok((w / "deck.odp").stat().st_size > 1000, "LibreOffice: .pptx → .odp")
    r = jrun(["pptx_read.py", w / "deck.odp"], lo=True)
    titles = [s["title"] for s in r["slides"]]
    c.ok(titles[:3] == ["Acme Quarterly", "Plan", "Numbers"], f"LibreOffice: reads .odp ({titles[:3]})")
    run(["pptx_convert.py", w / "deck.odp", w / "legacy.ppt"], lo=True)
    info = jrun(["pptx_info.py", w / "legacy.ppt"], lo=True)
    c.ok(info["slide_count"] == 7, "LibreOffice: writes and reads .ppt")
    run(["pptx_convert.py", w / "legacy.ppt", w / "back.pptx"], lo=True)
    r = jrun(["pptx_read.py", w / "back.pptx", "--slides", "3"])
    tbl = [sh for sh in r["slides"][0]["shapes"] if sh["kind"] == "table"]
    c.ok(tbl and tbl[0]["table"]["rows"][1] == ["North", "120", "135"], "LibreOffice: .ppt → .pptx keeps the table")


# ── regressions for the acceptance review, big files, cache, round trip ─


def _texts(slide: dict[str, Any]) -> list[str]:
    """A slide's paragraph texts (not footers or slide numbers), sorted: what a rebuild must keep."""
    out = []
    for sh in slide["shapes"]:
        ph = (sh.get("placeholder") or {}).get("type")
        if ph in ("footer", "slide_number", "date") or sh["name"] in ("Footer", "Slide number"):
            continue
        for p in sh.get("paragraphs", []):
            t = " ".join(p["text"].split())
            if t:
                out.append(t)
    return sorted(out)


def _content(d: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for s in d["slides"]:
        out.append({
            "title": s["title"], "hidden": s["hidden"], "notes": s.get("notes", ""), "texts": _texts(s),
            "tables": [sh["table"]["rows"] for sh in s["shapes"] if "table" in sh],
            "charts": [(sh["chart"].get("categories"), [x["values"] for x in sh["chart"].get("series", [])]) for sh in s["shapes"] if "chart" in sh],
            "pictures": sorted(tuple(sh["image"].get("pixels") or ()) for sh in s["shapes"] if sh["kind"] == "picture"),
        })
    return out


def test_roundtrip(c: Checks, fx: dict[str, Path], w: Path) -> None:
    for n in ("photo.jpg", "diagram.png"):
        shutil.copy(fx["outline"].parent / n, w / n)
    shutil.copy(fx["outline"], w / "outline.md")
    shutil.copy(fx["spec"], w / "spec.json")
    for name, src in (("outline", ["--input", w / "outline.md", "--theme", "slate"]), ("spec", ["--spec", w / "spec.json"])):
        deck = w / f"{name}.pptx"
        run(["pptx_create.py", deck, *src])
        r = run(["pptx_read.py", deck, "--format", "spec", "--out", w / f"{name}.spec.json"])
        c.ok("Wrote" in r.stdout and (w / f"{name}.spec.json").exists(), f"round trip {name}: pptx_read --format spec --out")
        spec = json.loads((w / f"{name}.spec.json").read_text(encoding="utf-8"))
        rebuilt = w / f"{name}-rebuilt.pptx"
        made = jrun(["pptx_create.py", rebuilt, "--spec", w / f"{name}.spec.json"])
        a = _content(jrun(["pptx_read.py", deck, "--max-chars", "0"]))
        b = _content(jrun(["pptx_read.py", rebuilt, "--max-chars", "0"]))
        c.ok(len(a) == len(b), f"round trip {name}: same slide count ({len(a)} → {len(b)})")
        for i, (x, y) in enumerate(zip(a, b), 1):
            for k in ("title", "hidden", "notes", "texts", "tables", "charts", "pictures"):
                if x[k] != y[k]:
                    c.ok(False, f"round trip {name}: slide {i} {k} differs: {str(x[k])[:160]} → {str(y[k])[:160]}")
                    break
        c.ok(True, f"round trip {name}: titles, texts, tables, charts, notes and pictures compared")
        types = [s["type"] for s in spec["slides"]]
        c.ok("custom" not in types, f"round trip {name}: every slide mapped to a real slide type ({types})")
        c.ok(not [x for x in made["warnings"] if "overflow" in x], f"round trip {name}: the rebuild fits ({made['warnings']})")
    # exports fed back in
    run(["pptx_convert.py", w / "outline.pptx", w / "export.json"])
    d = jrun(["pptx_convert.py", w / "export.json", w / "from-json.pptx"])
    back = _content(jrun(["pptx_read.py", w / "from-json.pptx", "--max-chars", "0"]))
    c.ok(len(back) == 14 and back[3]["texts"] and any("export" in x for x in d.get("warnings", [])), "a pptx_read JSON export rebuilds the deck, with a note")
    run(["pptx_convert.py", w / "outline.pptx", w / "export.md"])
    e = run(["pptx_convert.py", w / "export.md", w / "from-md.pptx"], expect=2)
    c.ok("--format spec" in e.stderr and not (w / "from-md.pptx").exists(), "a pptx_read Markdown export is refused with the way to rebuild")
    d = jrun(["pptx_convert.py", w / "outline.pptx", w / "conv.spec.json"])
    c.ok((w / "conv.spec.json").exists() and (w / "conv-media").is_dir() and "spec" in d["how"], "pptx_convert → .spec.json with its media folder")


def test_bigdeck(c: Checks, fx: dict[str, Path], w: Path) -> None:
    (w / "big.md").write_text(big_outline(64), encoding="utf-8")
    deck = w / "big.pptx"
    run(["pptx_create.py", deck, "--input", w / "big.md"])
    m = jrun(["pptx_read.py", deck])
    c.ok(m.get("map") is True and len(m["slides"]) == 64 and m["slides"][11]["address"] == "slide 12", "big deck: the default read is a map with addresses")
    s12 = m["slides"][12]
    c.ok(s12["title"] == "Topic 12" and s12["notes"] and s12["words"] > 5 and s12["shapes"] >= 2, f"big deck: map rows carry title, words, shapes, notes ({s12})")
    c.ok(any(x.get("charts") for x in m["slides"]) and any(x.get("tables") for x in m["slides"]), "big deck: the map says which slides hold charts and tables")
    md = run(["pptx_read.py", deck]).stdout
    c.ok("a map, because it has more than 50 slides" in md and "| 64 |" in md and "--find" in md, "big deck: Markdown map with the drill-down commands")
    f = jrun(["pptx_read.py", deck, "--find", "region 12"])
    c.ok(f["total"] == 1 and f["hits"][0]["address"].startswith("slide 13 / shape ") and "«region 12»" in f["hits"][0]["context"], f"big deck: --find gives an address and context ({f['hits'][:1]})")
    g = jrun(["pptx_read.py", deck, "--grep", r"Speaker notes 1\d\b"])
    c.ok(g["total"] == 2 and all(h["address"].endswith("/ notes") for h in g["hits"]), f"big deck: --grep searches notes ({g['total']})")
    t = jrun(["pptx_read.py", deck, "--find", "South"])
    c.ok(t["total"] >= 10 and "row 3" in t["hits"][0]["address"], "big deck: --find reaches table rows")
    one = jrun(["pptx_read.py", deck, "--slides", "13"])
    c.ok(len(one["slides"]) == 1 and one["slides"][0]["shapes"] and "map" not in one, "big deck: --slides drills into full detail")
    part = run(["pptx_read.py", deck, "--full", "--max-chars", "4000"]).stdout
    mm = re.search(r"Next part: (python3 scripts/pptx_read\.py \S+ --slides (\d+)-64)", part)
    c.ok(bool(mm) and len(part) < 5200, "big deck: --max-chars stops at a slide boundary and gives the next command")
    if mm:
        nxt = run(["pptx_read.py", deck, "--slides", f"{mm.group(2)}-64", "--max-chars", "4000"]).stdout
        c.ok(f"## {mm.group(2)}. " in nxt, "big deck: the next-part command continues where the first part stopped")
    small = run(["pptx_read.py", deck, "--map", "--max-chars", "1500"]).stdout
    c.ok("--map --slides" in small, "big deck: a capped map says how to read the rest")
    csv_out = run(["pptx_read.py", deck, "--format", "csv", "--slides", "2-3"]).stdout
    c.ok('# slide 2 / table "Table"' in csv_out and "North,1" in csv_out and "category,A" in csv_out, "big deck: tables and chart data as CSV")
    # cache: a second render of the same slides is served from the cache
    t0 = time.time()
    a = jrun(["pptx_render.py", deck, "--out-dir", w / "r1", "--width", "480"])
    cold = time.time() - t0
    t0 = time.time()
    b = jrun(["pptx_render.py", deck, "--out-dir", w / "r2", "--width", "480"])
    warm = time.time() - t0
    c.ok(b["cached"] == 64 and a["cached"] == 0 and len(b["slides"]) == 64, f"cache: the second render reuses all 64 slides ({a['cached']} → {b['cached']})")
    c.ok(warm * 5 <= cold, f"cache: the second render is at least 5× faster ({cold:.2f} s → {warm:.2f} s)")
    c.ok(sha(Path(a["slides"][5]["png"])) == sha(Path(b["slides"][5]["png"])), "cache: cached PNGs are identical")
    # The render above already cached the deck model, so compare a read that bypasses the cache with a cached one.
    t0 = time.time()
    run(["pptx_read.py", deck, "--full", "--max-chars", "0", "--no-cache"])
    cold = time.time() - t0
    t0 = time.time()
    run(["pptx_read.py", deck, "--full", "--max-chars", "0"])
    warm = time.time() - t0
    c.ok(warm < cold, f"cache: a cached full read is faster than an uncached one ({cold:.2f} s → {warm:.2f} s)")
    cache = Path(os.environ["DESK_FILE_CACHE"])
    count = sum(1 for _ in cache.rglob(".complete"))
    run(["pptx_lint.py", deck, "--no-cache", "--min-size", "11"])
    run(["pptx_render.py", deck, "--slides", "3", "--out-dir", w / "r3", "--no-cache", "--width", "300"])
    c.ok(sum(1 for _ in cache.rglob(".complete")) == count, "cache: --no-cache neither reads nor writes the cache")
    sheets = jrun(["pptx_render.py", deck, "--out-dir", w / "sheets", "--sheet-only"])
    sizes = [png_info(Path(p))[0] for p in sheets["sheets"]]
    c.ok(len(sizes) == 4 and all(sz[0] >= 1400 and max(sz) <= 3136 for sz in sizes[:3]), f"contact sheets: 20 slides per sheet at a readable size {sizes}")


def test_regressions(c: Checks, fx: dict[str, Path], w: Path) -> None:
    # a KPI value wider than its card shrinks to fit, and all values on the slide keep one size
    spec = {"slides": [{"type": "kpi", "title": "Wide values", "items": [{"value": "$1,234,567", "label": "Revenue"}, {"value": "12,345,678", "label": "Users"}, {"value": "99.99%", "label": "Uptime"}, {"value": "€2.4M", "label": "ARR"}]}]}
    run(["pptx_create.py", w / "kpi.pptx", "--spec", json.dumps(spec)])
    r = jrun(["pptx_read.py", w / "kpi.pptx"])
    sizes = {tuple(sh.get("font_sizes", [])) for sh in r["slides"][0]["shapes"] if sh["name"].startswith("KPI value")}
    c.ok(len(sizes) == 1 and next(iter(sizes))[0] < 44, f"KPI values shrink together ({sizes})")
    lint = jrun(["pptx_lint.py", w / "kpi.pptx"])
    c.ok(lint["counts"]["error"] == 0 and lint["counts"]["warning"] == 0, f"KPI slide is lint clean ({[i['message'] for i in lint['issues']]})")
    run(["pptx_edit.py", w / "kpi.pptx", "--out", w / "kpi2.pptx", "--ops", '[{"op": "add_text", "slide": 1, "text": "Internationalization", "size": 28, "fit": 1, "x": 1, "y": 6, "w": 1.5, "h": 0.8}]'])
    lint = jrun(["pptx_lint.py", w / "kpi2.pptx"])
    c.ok(any(i["rule"] == "overflow" and "Internationalization" in i["message"] for i in lint["issues"]), "lint: a single word wider than its box is an overflow")
    png = jrun(["pptx_render.py", w / "kpi2.pptx", "--out-dir", w / "r", "--width", "1333"])["slides"][0]["png"]
    from PIL import Image

    with Image.open(png) as im:
        right = im.convert("L").crop((int(2.6 / 13.333 * im.width), int(6.0 / 7.5 * im.height), int(3.6 / 13.333 * im.width), int(6.8 / 7.5 * im.height)))
        c.ok(right.getextrema()[0] > 200, "render: the over-wide word breaks inside its box, as PowerPoint does")
    # misspelled fields are errors, never silent losses
    for ops, word in (('[{"op": "set_title", "slide": 1, "txt": "x"}]', '"txt" (did you mean "text"?)'),
                      ('[{"op": "replace", "find": "Q1", "with": "First"}]', 'unknown field "with"'),
                      ('[{"op": "replace", "find": "Q1"}]', 'needs "replace"'),
                      ('[{"op": "set_notes", "slide": 1, "note": "x"}]', 'unknown field "note"'),
                      ('[{"op": "set_shape", "slide": 3, "shape": "Sales chart", "colour": "#FF0000"}]', '"colour" (did you mean "color"?)'),
                      ('[{"op": "update_chart", "slide": 3, "shape": "Sales chart", "titel": "x"}]', '"titel" (did you mean "title"?)'),
                      ('[{"op": "delete_slide", "slides": "1-7"}]', "every slide")):
        e = run(["pptx_edit.py", fx["foreign"], "--out", w / "typo.pptx", "--ops", ops], expect=2)
        c.ok(word in e.stderr and not (w / "typo.pptx").exists(), f"edit refuses {ops[:60]}: {e.stderr.strip()[:120]}")
    e = run(["pptx_create.py", w / "t.pptx", "--spec", '{"slides": [{"type": "bullets", "title": "x", "bulets": ["a"]}]}'], expect=2)
    c.ok('"bulets" (did you mean "bullets"?)' in e.stderr, "create refuses a misspelled slide field")
    e = run(["pptx_create.py", w / "t.pptx", "--spec", '{"them": "clean", "slides": [{"title": "x"}]}'], expect=2)
    c.ok('"them" (did you mean "theme"?)' in e.stderr, "create refuses a misspelled deck field")
    # content after a columns block is kept (under the columns), with a warning
    md = "## Cols\n:::: columns\n::: column\n### Left\n- left one\n:::\n::: column\n### Right\n- right one\n:::\n::::\n\nTrailing paragraph after columns.\n"
    d = jrun(["pptx_create.py", w / "cols.pptx", "--input", "-"], stdin=md)
    c.ok(any("outside the ':::: columns'" in x for x in d["warnings"]), "columns: text after the block is reported")
    r = jrun(["pptx_read.py", w / "cols.pptx"])
    note = [sh for sh in r["slides"][0]["shapes"] if sh["name"] == "Columns note"]
    c.ok(note and "Trailing paragraph" in note[0]["text"], "columns: text after the block is kept under the columns")
    heads = [max(sh["font_sizes"]) for sh in r["slides"][0]["shapes"] if sh["name"].startswith("Column heading")]
    bodies = [max(sh["font_sizes"]) for sh in r["slides"][0]["shapes"] if sh.get("placeholder", {}).get("type") in ("body", "object") and sh.get("font_sizes")]
    c.ok(heads and bodies and min(heads) >= max(bodies) and len(set(bodies)) == 1, f"columns: headings at least as large as the bodies, both bodies alike ({heads}, {bodies})")
    order = run(["pptx_read.py", w / "cols.pptx"]).stdout
    c.ok(order.index("Left") < order.index("left one") < order.index("Right") < order.index("right one"), "read: two columns come out column by column")
    # update_chart: added series get their own colour; the automatic title is not left as 'Chart Title'
    lint0 = jrun(["pptx_lint.py", fx["autotitle"]])
    c.ok(not any(i["rule"] == "chart-title" for i in lint0["issues"]), "lint: one series with an automatic title is fine")
    d = jrun(["pptx_edit.py", fx["autotitle"], "--out", w / "chart2.pptx", "--ops", json.dumps([{"op": "update_chart", "slide": 1, "shape": "Chart 1", "categories": ["Q1", "Q2", "Q3", "Q4"], "series": [{"name": "Actual", "values": [9, 4, 2, 1]}, {"name": "Target", "values": [8, 5, 3, 2]}, {"name": "Plan", "values": [7, 6, 4, 3]}]}])])
    c.ok(any("Chart Title" in x for x in d["warnings"]), "update_chart: says it removed the automatic title")
    lint = jrun(["pptx_lint.py", w / "chart2.pptx"])
    c.ok(not any(i["rule"] in ("chart-title", "chart-colors") for i in lint["issues"]), f"update_chart: no 'Chart Title', distinct series colours ({[i['rule'] for i in lint['issues']]})")
    with zipfile.ZipFile(w / "chart2.pptx") as z:
        xml = z.read(next(n for n in z.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", n))).decode("utf-8")
    fills = re.findall(r"<c:ser>.*?<a:solidFill>(.*?)</a:solidFill>", xml, re.S)
    c.ok(len(fills) == 3 and len(set(fills)) == 3, f"update_chart: three series, three colours ({fills})")
    d = jrun(["pptx_edit.py", w / "chart2.pptx", "--out", w / "chart3.pptx", "--ops", '[{"op": "update_chart", "slide": 1, "shape": "Chart 1", "title": "Sales vs plan", "legend": "right", "colors": ["accent2", "#0A7F6F", "accent4"]}]'])
    ch = next(sh["chart"] for sh in jrun(["pptx_read.py", w / "chart3.pptx"])["slides"][0]["shapes"] if "chart" in sh)
    c.ok(ch["title"] == "Sales vs plan" and ch["legend"] == "right", f"update_chart: title and legend ({ch['title']}, {ch['legend']})")
    run(["pptx_edit.py", fx["foreign"], "--out", w / "t2.pptx", "--ops", '[{"op": "update_chart", "slide": 3, "shape": "Sales chart", "colors": ["accent1", "accent1"]}]'])
    lint = jrun(["pptx_lint.py", w / "t2.pptx", "--slides", "3"])
    c.ok(any(i["rule"] == "chart-colors" for i in lint["issues"]), "lint: two series with one colour are flagged")
    made = w / "made.pptx"
    shutil.copy(fx["outline"], w / "o.md")
    for n in ("photo.jpg", "diagram.png"):
        shutil.copy(fx["outline"].parent / n, w / n)
    run(["pptx_create.py", made, "--input", w / "o.md"])
    with zipfile.ZipFile(made) as z:
        xml = z.read(next(n for n in z.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", n))).decode("utf-8")
    c.ok('<a:schemeClr val="accent1"/>' in xml and '<a:schemeClr val="accent2"/>' in xml, "create: chart series take theme colours (re-themable)")
    # set_title on a slide without a title: a real title placeholder above the content
    d = jrun(["pptx_edit.py", fx["foreign"], "--out", w / "titled.pptx", "--ops", '[{"op": "set_title", "slide": 4, "text": "Pictures and shapes"}]'])
    r = jrun(["pptx_read.py", w / "titled.pptx", "--slides", "4"])
    t = next(sh for sh in r["slides"][0]["shapes"] if sh.get("placeholder", {}).get("type") == "title")
    c.ok(r["slides"][0]["title"] == "Pictures and shapes" and t["y"] + t["h"] <= 1.0, f"set_title: a title placeholder in the free band ({t['y']}, {t['h']})")
    lint = jrun(["pptx_lint.py", w / "titled.pptx", "--slides", "4"])
    c.ok(not any(i["rule"] in ("title", "overlap") and i["shape"] in (None, t["name"]) for i in lint["issues"]), f"set_title: lint sees a title and no overlap ({[i['message'] for i in lint['issues']]})")
    # LibreOffice-only targets and non-deck inputs say what is wrong
    e = run(["pptx_convert.py", fx["foreign"], w / "x.odp"], expect=1)
    c.ok("writing .odp needs LibreOffice" in e.stderr, f"convert: .odp output names what needs LibreOffice ({e.stderr.strip()})")
    (w / "doc.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
    e = run(["pptx_info.py", w / "doc.pdf"], expect=1)
    c.ok("is a PDF" in e.stderr and "pdf-toolkit" in e.stderr, "info: a PDF input is named as such")
    # lists stay under the lint's word limit, footers do not count
    bl = [f"Bullet {i} explains one more thing about the plan in about fourteen words total here" for i in range(1, 7)]
    d = jrun(["pptx_create.py", w / "wordy.pptx", "--spec", json.dumps({"footer": "Northwind · Board review · Confidential", "slides": [{"type": "bullets", "title": "A wordy slide about many things", "bullets": bl}]})])
    lint = jrun(["pptx_lint.py", w / "wordy.pptx"])
    c.ok(len(d["slides"]) == 2 and not any(i["rule"] == "words" for i in lint["issues"]), f"create: a list over the word limit continues on a second slide ({len(d['slides'])})")
    # silent coercions are reported
    md = "<!-- type: kpi -->\n## K\n- just a sentence\n- 44% | Activation\n\n<!-- chart: pie -->\n## P\n| A | v |\n|---|---|\n| x | abc |\n| y | 12 |\n| z | -5 |\n\n<!-- colour: red -->\n## D\n- a\n"
    d = jrun(["pptx_create.py", w / "co.pptx", "--input", "-"], stdin=md)
    ws = " | ".join(d["warnings"])
    c.ok("no 'value | label' form" in ws and "'abc'" in ws and "negative" in ws and "unknown directive" in ws, f"create: coerced content is reported ({ws[:300]})")
    run(["pptx_create.py", w / "nest.pptx", "--spec", '{"slides": [{"type": "bullets", "title": "N", "bullets": ["a", ["sub a", "sub b"], "b"]}]}'])
    body = next(sh for sh in jrun(["pptx_read.py", w / "nest.pptx"])["slides"][0]["shapes"] if sh.get("placeholder", {}).get("idx") == 1)
    c.ok([p["level"] for p in body["paragraphs"]] == [0, 1, 1, 0], "create: a nested list holds sub-points")
    # the outline leaves out footers and slide numbers
    md = run(["pptx_read.py", made, "--slides", "4"]).stdout
    c.ok("Aurora · Q3 review" not in md and not re.search(r"^4$", md, re.M), "read: footers and slide numbers are not content")
    # every leftover placeholder in a shape is reported
    run(["pptx_edit.py", fx["foreign"], "--out", w / "left.pptx", "--ops", '[{"op": "set_shape", "slide": 4, "shape": "Leftover", "text": "Hello {{customer_name}}, Lorem ipsum dolor sit amet"}]'])
    lint = jrun(["pptx_lint.py", w / "left.pptx", "--slides", "4"])
    msg = next(i["message"] for i in lint["issues"] if i["rule"] == "leftover")
    c.ok("{{customer_name}}" in msg and "Lorem ipsum" in msg, f"lint: all leftovers in a shape ({msg})")
    # Text from any deck: "{{" or "[insert" never closed once rescanned the rest of the shape from each one.
    sys.path.insert(0, str(HERE))
    from pptx_lint import LEFTOVER

    t = time.time()
    LEFTOVER.findall("{{" * 200000), LEFTOVER.findall("[insert" * 100000)
    c.ok([m.group(0) for m in LEFTOVER.finditer("Hi {{name}} [Insert logo] TBD")] == ["{{name}}", "[Insert logo]", "TBD"] and time.time() - t < 1,
         f"lint: placeholder patterns are linear ({time.time() - t:.2f}s)")
    # a 4:3 deck from another tool rebuilds with every slide, its titles, and positioned tables that do not split
    run(["pptx_read.py", fx["foreign"], "--format", "spec", "--out", w / "foreign.spec.json"])
    made = jrun(["pptx_create.py", w / "foreign-re.pptx", "--spec", w / "foreign.spec.json"])
    a = jrun(["pptx_info.py", fx["foreign"]])
    kept = [x["title"] for x in made["slides"] if not x["title"].endswith("(cont.)")]  # its overfull slide continues on another
    c.ok(len(kept) == a["slide_count"] and kept[:3] == ["Acme Quarterly", "Plan", "Numbers"] and any("continues" in x for x in made["warnings"]), f"spec round trip of a foreign 4:3 deck keeps its slides ({kept})")
    spec = {"slides": [{"type": "custom", "title": "Boxed", "elements": [{"kind": "table", "x": 1, "y": 2, "w": 6, "h": 1.2, "columns": ["a", "b"], "rows": [["1", "2"], ["3", "4"], ["5", "6"]]}]},
                       {"type": "table", "title": "On a blank layout", "layout": "Blank", "columns": ["a"], "rows": [["1"]]}]}
    made = jrun(["pptx_create.py", w / "boxed.pptx", "--spec", json.dumps(spec)])
    c.ok(len(made["slides"]) == 2 and made["slides"][1]["title"] == "On a blank layout", f"a positioned table shrinks instead of spilling onto new slides; a title on a Blank layout is still the title ({made['slides']})")
    # a layout whose placeholders share an idx (LibreOffice's .ppt exports) still gives new slides the right boxes
    from pptx import Presentation

    prs = Presentation(str(fx["foreign"]))
    lay = next(x for x in prs.slide_layouts if x.name == "Title and Content")
    body = next(p for p in lay.placeholders if p.placeholder_format.idx == 1)
    body._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}ph").set("idx", "0")
    prs.save(str(w / "dupidx.pptx"))
    run(["pptx_edit.py", w / "dupidx.pptx", "--out", w / "dupidx2.pptx", "--ops", '[{"op": "add_slide", "type": "bullets", "layout": "Title and Content", "title": "Added", "bullets": ["one", "two"]}]'])
    r = jrun(["pptx_read.py", w / "dupidx2.pptx", "--slides", "8"])
    bx = next(sh for sh in r["slides"][0]["shapes"] if sh.get("placeholder", {}).get("type") in ("body", "object"))
    c.ok(r["slide_count"] == 8 and bx["y"] > 1.2 and bx["h"] > 3, f"duplicate placeholder idx: the body keeps its own box ({bx['y']}, {bx['h']})")
    # zip and XML bombs are refused before parsing
    for key in ("bomb", "xxe"):
        for extra in ([], ["--map"], ["--find", "x"]):
            e = run(["pptx_read.py", fx[key], *extra], expect=1)
            c.ok("refusing" in e.stderr or "DOCTYPE" in e.stderr, f"{key} {extra}: refused ({e.stderr.strip()[:120]})")
        e = run(["pptx_render.py", fx[key], "--out-dir", w / f"r-{key}"], expect=1)
        c.ok("refusing" in e.stderr or "DOCTYPE" in e.stderr, f"{key}: render refuses it")


# ── catalog/shared modules (every skill has a copy; they are checked here once) ──

FAKE_SOFFICE = r'''
import json, os, sys, time
from pathlib import Path
from urllib.parse import unquote, urlparse

args = sys.argv[1:]
prof = unquote(urlparse(next(a.split("=", 1)[1] for a in args if a.startswith("-env:UserInstallation="))).path)
prof = prof[1:] if os.name == "nt" and prof.startswith("/") else prof
Path(prof, "user").mkdir(parents=True, exist_ok=True)  # LibreOffice fills its profile
out, src = Path(args[args.index("--outdir") + 1]), Path(args[-1])
t0 = time.time()
time.sleep(0.3)
(out / (src.stem + ".pdf")).write_bytes(b"%PDF-1.4 stand-in")
with open(os.environ["FAKE_SOFFICE_LOG"], "a", encoding="utf-8") as f:
    f.write(json.dumps({"profile": prof, "start": t0, "end": time.time()}) + "\n")
'''


def _fake_soffice(w: Path) -> Path:
    script = w / "fake_soffice.py"
    script.write_text(FAKE_SOFFICE, encoding="utf-8")
    if os.name == "nt":
        exe = w / "fake_soffice.cmd"
        exe.write_text(f'@"{PY}" "{script}" %*\r\n', encoding="utf-8")
    else:
        exe = w / "fake_soffice"
        exe.write_text(f'#!/bin/sh\nexec "{PY}" "{script}" "$@"\n', encoding="utf-8")
        exe.chmod(0o755)
    return exe


def _package(path: Path, parts: dict[str, bytes]) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        for name, data in parts.items():
            z.writestr(name, data)
    return path


def test_shared_modules(c: Checks, fx: dict[str, Path], w: Path) -> None:
    import threading

    sys.path.insert(0, str(HERE))
    import _cache
    import _render
    from _common import IS_WINDOWS, SkillError, check_zip, run_tool

    saved = {k: os.environ.get(k) for k in ("DESK_NO_CACHE", "DESK_FILE_CACHE", "FAKE_SOFFICE_LOG")}
    os.environ["DESK_NO_CACHE"] = "1"  # check_zip verdicts are computed, not read back from the cache
    try:
        # check_zip(refuse_dtd=True) reads each XML part up to its root element, in the part's own encoding.
        decl = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        dtd = b'<!DOCTYPE lol [<!ENTITY a "aaaaaaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
        body = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>'
        u16 = '<?xml version="1.0" encoding="UTF-16"?>' + (dtd + body).decode()
        hidden = {
            "at the start": decl + dtd + body,
            "after 3000 spaces in a comment": decl + b"<!--" + b" " * 3000 + b"-->" + dtd + body,
            "in UTF-16 with a BOM": b"\xff\xfe" + u16.encode("utf-16-le"),
            "in UTF-16 without a BOM": u16.encode("utf-16-be"),
            "in UTF-32": u16.replace("UTF-16", "UTF-32").encode("utf-32"),
            "after 500 processing instructions": decl + b"<?pi x?>\n" * 500 + dtd + body,
            "in a declared EBCDIC encoding": '<?xml version="1.0" encoding="cp037"?>'.encode("cp037") + (dtd + body).decode().encode("cp037"),
        }
        for what, data in hidden.items():
            try:
                check_zip(_package(w / "hidden.docx", {"word/document.xml": data}), refuse_dtd=True)
                c.ok(False, f"check_zip refuses a DOCTYPE {what}")
            except SkillError as e:
                c.ok("DOCTYPE" in str(e) and "word/document.xml" in str(e), f"check_zip refuses a DOCTYPE {what}: {e}")
        try:
            check_zip(_package(w / "long.docx", {"word/document.xml": decl + b"<!--" + b" " * (1100 * 1024) + b"-->" + body}), refuse_dtd=True)
            c.ok(False, "check_zip refuses a prolog longer than 1 MB")
        except SkillError as e:
            c.ok("before its first element" in str(e), f"check_zip refuses a prolog longer than 1 MB: {e}")
        fine = {
            "docx": {"word/document.xml": decl + b"<!-- made by hand -->\n" + body, "word/_rels/document.xml.rels": decl + b"<Relationships/>"},
            "docx with DOCTYPE-like text in its body": {"word/document.xml": decl + body.replace(b"<w:body/>", b"<w:body><w:t><![CDATA[<!DOCTYPE x>]]></w:t></w:body>")},
            "docx in UTF-16": {"word/document.xml": b"\xff\xfe" + ('<?xml version="1.0" encoding="UTF-16"?>' + body.decode()).encode("utf-16-le")},
            "xlsx": {"xl/workbook.xml": decl + b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets/></workbook>',
                     "xl/worksheets/sheet1.xml": decl + b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>'},
        }
        for what, parts in fine.items():
            try:
                check_zip(_package(w / "fine.zip", parts), refuse_dtd=True)
                c.ok(True, what)
            except SkillError as e:
                c.ok(False, f"check_zip accepts a normal {what}: {e}")
        try:
            check_zip(fx["foreign"], refuse_dtd=True)
            c.ok(True, "pptx")
        except SkillError as e:
            c.ok(False, f"check_zip accepts a normal pptx: {e}")
        rows = b"".join(b'<row r="%d"><c r="A%d"><v>%d</v></c></row>' % (i, i, i) for i in range(1, 300_001))
        parts = {"xl/worksheets/sheet1.xml": decl + b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + rows + b"</sheetData></worksheet>"}
        parts.update({f"xl/drawings/_rels/drawing{i}.xml.rels": decl + b"<Relationships/>" for i in range(1, 3001)})
        big = _package(w / "big.xlsx", parts)
        t = time.time()
        check_zip(big, refuse_dtd=True)
        took = time.time() - t
        c.ok(took < 2, f"check_zip reads only the prolog of each part: a {len(parts['xl/worksheets/sheet1.xml']) >> 20} MB sheet and 3000 parts in {took:.2f} s")
        os.environ.pop("DESK_NO_CACHE", None)

        # cached_dir: a build whose rename keeps failing is returned once, as a transient result (never rebuilt).
        os.environ["DESK_FILE_CACHE"] = str(w / "cache")
        src = w / "cache-src.bin"
        src.write_bytes(b"cache me" * 100)
        builds: list[int] = []

        def build(tmp: Path) -> None:
            builds.append(1)
            (tmp / "out.txt").write_text("built", encoding="utf-8")

        class _Blocked:
            """os, except that renames fail (a file held open on Windows, or a virus scanner)."""

            def __getattr__(self, name: str) -> Any:
                return getattr(os, name)

            @staticmethod
            def replace(a: Any, b: Any) -> None:
                raise PermissionError(13, "in use (simulated)")

        _cache.os = _Blocked()  # type: ignore[assignment]
        try:
            d = _cache.cached_dir(src, "selftest-rename", {}, "1", build)
        finally:
            _cache.os = os  # type: ignore[assignment]
        c.ok(len(builds) == 1 and (d / "out.txt").read_text(encoding="utf-8") == "built" and _cache.transient(d) and not (d / ".complete").exists(),
             f"cached_dir: a rename that keeps failing gives a transient result after one build ({len(builds)} builds, {d})")
        _cache.release(d)
        c.ok(not d.exists() and _cache.lookup(src, "selftest-rename", {}, "1") is None, "cached_dir: release() deletes that transient result")
        final = _cache.entry_dir(src, "selftest-leftover", {}, "1")
        (final / "junk").mkdir(parents=True)
        builds.clear()
        d = _cache.cached_dir(src, "selftest-leftover", {}, "1", build)
        c.ok(d == final and len(builds) == 1 and (d / "out.txt").exists() and not _cache.transient(d), "cached_dir: a leftover without a marker is replaced after one build")

        # office_convert: one LibreOffice at a time per process; the shared profile is released after each call and
        # private profiles are deleted (checked with a stand-in soffice that logs what it was given).
        exe = _fake_soffice(w)
        log = w / "soffice.log"
        os.environ["FAKE_SOFFICE_LOG"] = str(log)
        doc = w / "doc.pptx"
        shutil.copy(fx["foreign"], doc)
        tmp = w / "tmp"
        tmp.mkdir()
        saved_tmp = tempfile.tempdir
        tempfile.tempdir = str(tmp)  # the shared profile lives in the temp folder
        try:
            results: list[Any] = []
            threads = [threading.Thread(target=lambda: results.append(_render.office_convert(doc, "pdf", soffice=str(exe)))) for _ in range(3)]
            for th in threads:
                th.start()
            for th in threads:
                th.join()
            runs = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines()]
            spans = sorted((r["start"], r["end"]) for r in runs)
            c.ok(len(results) == 3 and all(Path(p).exists() for p in results) and len(runs) == 3, f"office_convert from 3 threads: 3 conversions ({len(results)} results, {len(runs)} runs)")
            c.ok(all(a[1] <= b[0] for a, b in zip(spans, spans[1:])), f"office_convert runs one soffice at a time per process: {spans}")
            c.ok(all(Path(r["profile"]).name == "desk-lo-profile" for r in runs) and not (tmp / "desk-lo-profile.lock").exists(),
                 f"office_convert reuses the shared profile and releases it after each call: {[r['profile'] for r in runs]}")
            log.unlink()
            (tmp / "desk-lo-profile.lock").write_text("", encoding="utf-8")  # another process holds the shared profile
            _render.office_convert(doc, "pdf", soffice=str(exe))
            _render.office_convert(doc, "pdf", soffice=str(exe))
            runs = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines()]
            left = sorted(p.name for p in tmp.glob("desk-lo-profile-*"))
            c.ok(len(runs) == 2 and all(Path(r["profile"]).name.startswith("desk-lo-profile-") for r in runs) and not left and (tmp / "desk-lo-profile.lock").exists(),
                 f"office_convert deletes private profiles and leaves another process's lock alone: {[r['profile'] for r in runs]}, left {left}")
        finally:
            tempfile.tempdir = saved_tmp

        # run_tool: a command past its timeout is stopped (on Windows with the processes it started).
        t = time.time()
        pidfile = w / "grandchild.pid"
        # (the grandchild does not inherit the pipes: on POSIX run_tool would wait for it to close them)
        parent = "import subprocess, sys, time; c = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL); open(sys.argv[1], 'w').write(str(c.pid)); time.sleep(60)"
        try:
            run_tool([PY, "-c", parent, str(pidfile)], timeout=3)
            c.ok(False, "run_tool stops a command at its timeout")
        except SkillError as e:
            c.ok("timed out" in str(e) and time.time() - t < 20, f"run_tool stops a command at its timeout ({time.time() - t:.1f} s): {e}")
        if IS_WINDOWS and pidfile.exists():
            from _common import process_footprint_mb

            time.sleep(1)
            c.ok(process_footprint_mb(int(pidfile.read_text())) is None, "run_tool's timeout on Windows also stops the processes the command started")
        elif pidfile.exists():
            import signal

            try:
                os.kill(int(pidfile.read_text()), signal.SIGTERM)  # POSIX kills only the command itself: tidy up here
            except (OSError, ValueError):
                pass
    finally:
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def test_lo_office_convert_profiles(c: Checks, fx: dict[str, Path], w: Path) -> None:
    """Two real LibreOffice conversions in one process reuse the shared profile and leave no private one behind."""
    tmp = w / "tmp"
    tmp.mkdir()
    code = "import sys; sys.path.insert(0, sys.argv[1]); from _render import office_convert; print(office_convert(sys.argv[2]).stat().st_size, office_convert(sys.argv[2]).stat().st_size)"
    env = dict(os.environ, TMPDIR=str(tmp), TEMP=str(tmp), TMP=str(tmp))
    r = subprocess.run([PY, "-c", code, str(HERE), str(fx["foreign"])], capture_output=True, text=True, encoding="utf-8", env=env, timeout=400)
    sizes = r.stdout.split()
    left = sorted(p.name for p in tmp.glob("desk-lo-profile*"))
    c.ok(r.returncode == 0 and len(sizes) == 2 and all(int(s) > 1000 for s in sizes), f"LibreOffice: two conversions in one process ({r.stdout.strip()} {r.stderr[-400:]})")
    c.ok(left == ["desk-lo-profile"], f"LibreOffice: the shared profile is released after each conversion and no private profile is left: {left}")


if __name__ == "__main__":
    sys.exit(main())
